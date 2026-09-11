import io

import pytest
from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.report import IRExcelExporter


@pytest.mark.parametrize("mode_setting", [
    "set central-nat enable",
    "set ngfw-mode policy-based",
])
def test_effective_central_nat_projects_standalone_vip_to_canonical_dnat(mode_setting):
    extraction = extract_fortigate_config(f"""
config system settings
    {mode_setting}
end
config firewall vip
    edit "STANDALONE_WEB"
        set extip 198.51.100.20
        set mappedip 10.0.0.20
        set extintf "wan"
        set portforward enable
        set extport 8443-8444
        set mappedport 443-444
        set protocol udp
        set src-filter "192.0.2.0/24"
        set srcintf-filter "wan"
    next
end
""")

    rules = extraction.canonical_ir.nat_rules
    assert len(rules) == 1
    rule = rules[0]
    assert rule.source_origin == "central-vip"
    assert rule.destination == ["198.51.100.20"]
    assert rule.translated_destinations == ["10.0.0.20"]
    assert rule.destination_protocol == "udp"
    assert (rule.original_destination_ports[0].start, rule.original_destination_ports[0].end) == (8443, 8444)
    assert (rule.translated_destination_ports[0].start, rule.translated_destination_ports[0].end) == (443, 444)
    assert rule.source_to_interfaces == ["wan"]
    assert rule.source_vip_filters == ["192.0.2.0/24"]
    assert rule.source_vip_interface_filters == ["wan"]

    workbook = load_workbook(io.BytesIO(IRExcelExporter(
        extraction.canonical_ir, extraction
    ).generate()))
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(
        row for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Name"]).value == "DNAT-CENTRAL-STANDALONE_WEB"
    )
    assert sheet.cell(row, headers["Original Destination"]).value == "198.51.100.20"
    assert sheet.cell(row, headers["Translated Destination"]).value == "10.0.0.20"
    assert sheet.cell(row, headers["Original Destination Port"]).value == "8443-8444"
    assert sheet.cell(row, headers["Translated Destination Port"]).value == "443-444"


def test_policy_based_ngfw_does_not_duplicate_vip_dnat_from_firewall_policy():
    extraction = extract_fortigate_config("""
config system settings
    set ngfw-mode policy-based
end
config firewall vip
    edit "WEB"
        set extip 198.51.100.22
        set mappedip 10.0.0.22
    next
end
config firewall policy
    edit 10
        set srcintf "wan"
        set dstintf "lan"
        set srcaddr "all"
        set dstaddr "WEB"
        set service "ALL"
        set action accept
    next
end
""")

    assert [rule.source_origin for rule in extraction.canonical_ir.nat_rules] == [
        "central-vip"
    ]


def test_disabled_standalone_vip_remains_inventory_only_under_central_nat():
    extraction = extract_fortigate_config("""
config system settings
    set central-nat enable
end
config firewall vip
    edit "DISABLED_WEB"
        set status disable
        set extip 198.51.100.21
        set mappedip 10.0.0.21
    next
end
""")

    assert extraction.canonical_ir.virtual_ips[0].enabled is False
    assert extraction.canonical_ir.nat_rules == []


def test_central_snat_is_normalized_in_source_order_with_pool_and_ports():
    extraction = extract_fortigate_config("""
config system settings
    set central-nat enable
end
config firewall ippool
    edit "pool1"
        set startip 198.51.100.10
        set endip 198.51.100.20
    next
end
config firewall central-snat-map
    edit 20
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
        set protocol 6
        set orig-port 1000-2000
        set dst-port 443
        set nat-ippool pool1
        set nat-port 40000-40010
    next
    edit 10
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
        set nat disable
    next
end
""")

    rules = extraction.canonical_ir.nat_rules
    assert [rule.sequence for rule in rules] == [1, 2]
    assert rules[0].source_pool_references == ["pool1"]
    assert rules[0].translated_sources == ["198.51.100.10-198.51.100.20"]
    assert rules[0].original_source_ports[0].start == 1000
    assert rules[0].translated_source_ports[0].end == 40010
    assert rules[1].source_translation_mode.value == "none"
    assert [rule.source_order for rule in extraction.canonical_ir.central_snat_rules] == [1, 2]


def test_central_snat_without_pool_uses_resolved_interface_address_and_exports():
    extraction = extract_fortigate_config("""
config system interface
    edit "wan"
        set ip 203.0.113.10 255.255.255.0
    next
end
config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 20
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
        set orig-port 1000-2000
        set dst-port 443
        set nat enable
    next
end
""")

    rule = extraction.canonical_ir.nat_rules[0]
    assert rule.source_translation_mode.value == "interface-address"
    assert rule.translated_sources == ["203.0.113.10"]
    assert rule.requires_manual_review is False

    workbook = load_workbook(io.BytesIO(IRExcelExporter(
        extraction.canonical_ir, extraction
    ).generate()))
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Original Source"]).value == "all"
    assert sheet.cell(4, headers["Original Destination"]).value == "all"
    assert sheet.cell(4, headers["Translated Source"]).value == "203.0.113.10"
    assert sheet.cell(4, headers["Source Translation Mode"]).value == "interface-address"
    assert sheet.cell(4, headers["Original Source Port"]).value == "1000-2000"
    assert sheet.cell(4, headers["Original Destination Port"]).value == "443"


def test_central_snat_without_pool_keeps_interface_address_review_for_ambiguous_egress():
    extraction = extract_fortigate_config("""
config system interface
    edit "wan"
        set ip 203.0.113.10 255.255.255.0
    next
end
config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 1
        set srcintf "lan"
        set dstintf "wan" "wan2"
        set orig-addr "all"
        set dst-addr "all"
        set nat enable
    next
end
""")

    rule = extraction.canonical_ir.nat_rules[0]
    assert rule.source_translation_mode.value == "interface-address"
    assert rule.translated_sources == []
    assert rule.requires_manual_review is True
    assert any("multiple possible outgoing interfaces" in reason for reason in rule.review_reasons)


def test_policy_based_ngfw_enables_effective_central_nat_and_suppresses_policy_nat():
    extraction = extract_fortigate_config("""
config system settings
    set ngfw-mode policy-based
end
config firewall central-snat-map
    edit 1
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
        set nat-ippool "pool1"
    next
end
config firewall policy
    edit 2
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    assert [rule.source_origin for rule in extraction.canonical_ir.nat_rules] == [
        "central-snat-map"
    ]


def test_central_snat_nat64_uses_ipv6_original_and_ipv4_translated_families():
    extraction = extract_fortigate_config("""
config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 1
        set type ipv6
        set orig-addr6 "SRC6"
        set dst-addr6 "DST6"
        set nat64 enable
    next
end
""")

    rule = extraction.canonical_ir.nat_rules[0]
    assert rule.nat_family.value == "nat64"
    assert rule.original_address_family == "ipv6"
    assert rule.translated_address_family == "ipv4"
