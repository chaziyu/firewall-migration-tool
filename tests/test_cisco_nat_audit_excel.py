import io

from openpyxl import load_workbook

from fwmigrate.ir import IRConfig
from fwmigrate.ir.enums import NATTranslationMode, NATType
from fwmigrate.ir.metadata import IRMetadata
from fwmigrate.ir.nat import IRNATRule
from fwmigrate.report import IRExcelExporter


def test_cisco_nat_modifiers_are_first_class_excel_columns():
    ir = IRConfig(
        metadata=IRMetadata(source_vendor="cisco_asa", source_product="Cisco ASA"),
        nat_rules=[IRNATRule(
            name="cisco-nat",
            type=NATType.SOURCE,
            source=["REAL"],
            destination=["any"],
            services=["any"],
            translated_sources=["POOL"],
            source_translation_mode=NATTranslationMode.DYNAMIC_IP_AND_PORT,
            source_attributes={
                "dns": True,
                "no_proxy_arp": True,
                "route_lookup": False,
                "unidirectional": True,
                "net_to_net": False,
                "pat_pool_options": ["round-robin", "extended"],
                "raw_options": ["interface", "unknown-option"],
            },
        )],
    )

    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()), data_only=False)
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    for header in (
        "DNS Rewrite",
        "No Proxy ARP",
        "Route Lookup",
        "Unidirectional",
        "Net-to-Net",
        "PAT Options",
        "Raw NAT Options",
    ):
        assert header in headers

    assert sheet.cell(4, headers["DNS Rewrite"]).value == "TRUE"
    assert sheet.cell(4, headers["No Proxy ARP"]).value == "TRUE"
    assert sheet.cell(4, headers["Route Lookup"]).value == "FALSE"
    assert sheet.cell(4, headers["Unidirectional"]).value == "TRUE"
    assert sheet.cell(4, headers["Net-to-Net"]).value == "FALSE"
    assert "round-robin" in sheet.cell(4, headers["PAT Options"]).value
    assert "unknown-option" in sheet.cell(4, headers["Raw NAT Options"]).value
    assert "Additional Settings" in headers
