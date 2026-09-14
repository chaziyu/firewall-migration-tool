import io

from openpyxl import load_workbook

from fwmigrate.ir.migrations_1_67 import migrate_1_66_to_1_67
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


def _config(settings: str = "") -> str:
    return f'''config firewall policy
    edit 1
        set action accept
{settings}    next
end
'''


def test_parser_types_all_fortios_746_policy_fields():
    policy = parse_fortigate_config(_config('''        set reputation-direction source
        set reputation-direction6 source
        set reputation-minimum 2
        set reputation-minimum6 3
        set policy-expiry enable
        set policy-expiry-date "2026-01-01 00:00:00"
        set policy-expiry-date-utc 1767225600
        set schedule-timeout enable
''')).policies[0]

    assert policy.reputation_direction == "source"
    assert policy.reputation_direction6 == "source"
    assert policy.reputation_minimum == 2
    assert policy.reputation_minimum6 == 3
    assert policy.policy_expiry == "enable"
    assert policy.policy_expiry_date == "2026-01-01 00:00:00"
    assert policy.policy_expiry_date_utc == "1767225600"
    assert policy.schedule_timeout == "enable"
    assert not {
        "reputation_direction", "reputation_direction6", "reputation_minimum",
        "reputation_minimum6", "policy_expiry", "policy_expiry_date",
        "policy_expiry_date_utc", "schedule_timeout",
    } & policy.extra_settings.keys()


def test_omitted_values_keep_nullable_source_and_effective_defaults():
    policy = extract_fortigate_config(_config()).canonical_ir.policies[0]

    assert policy.source_policy_expiry is None
    assert policy.source_effective_policy_expiry == "disable"
    assert policy.source_policy_expiry_date is None
    assert policy.source_policy_expiry_date_utc is None
    assert policy.source_schedule_timeout is None
    assert policy.source_effective_schedule_timeout == "disable"
    assert policy.source_reputation_direction is None
    assert policy.source_effective_reputation_direction == "destination"
    assert policy.source_reputation_direction6 is None
    assert policy.source_effective_reputation_direction6 == "destination"
    assert policy.source_reputation_minimum is None
    assert policy.source_effective_reputation_minimum == 0
    assert policy.source_reputation_minimum6 is None
    assert policy.source_effective_reputation_minimum6 == 0
    assert policy.source_match_vip is None
    assert policy.source_effective_match_vip == "enable"
    assert policy.source_match_vip_only is None
    assert policy.source_effective_match_vip_only == "disable"


def test_non_default_policy_semantics_require_review():
    policy = extract_fortigate_config(_config('''        set policy-expiry enable
        set schedule-timeout enable
        set reputation-direction source
        set reputation-direction6 source
        set reputation-minimum 2
        set reputation-minimum6 3
''')).canonical_ir.policies[0]

    assert policy.requires_manual_review is True
    assert all(
        any(term in reason for reason in policy.review_reasons)
        for term in (
            "policy-expiry", "schedule-timeout", "reputation direction",
            "reputation-minimum", "reputation-minimum6",
        )
    )


def test_malformed_ipv6_reputation_minimum_is_preserved_for_review():
    policy = parse_fortigate_config(
        _config("        set reputation-minimum6 future-value\n")
    ).policies[0]

    assert policy.reputation_minimum6 is None
    assert policy.extra_settings["unparsed_reputation_minimum6"] == "future-value"


def test_unknown_policy_commands_keep_values_and_operations_in_ir():
    result = extract_fortigate_config(_config('''        set future-traffic-setting "one" "two"
        append future-traffic-setting "B" "C"
        unset future-traffic-setting
'''))
    policy = result.canonical_ir.policies[0]

    assert policy.source_extra_settings == {
        "source_unset_settings": ["future-traffic-setting"]
    }
    assert [
        (command.operation, command.key, command.values)
        for command in policy.source_extra_setting_commands
    ] == [
        ("set", "future-traffic-setting", ["one", "two"]),
        ("append", "future-traffic-setting", ["B", "C"]),
        ("unset", "future-traffic-setting", []),
    ]
    assert policy.model_dump(mode="json")["source_extra_setting_commands"]


def test_policy_semantics_are_visible_as_configured_and_effective_excel_values():
    result = extract_fortigate_config(_config('''        set policy-expiry enable
        set policy-expiry-date "2026-01-01 00:00:00"
        set policy-expiry-date-utc 1767225600
        set reputation-minimum6 3
'''))
    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, extraction_result=result).generate())
    )
    sheet = workbook["Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    assert sheet.cell(4, headers["Policy Expiry"]).value == "enable"
    assert sheet.cell(4, headers["Effective Policy Expiry"]).value == "enable"
    assert sheet.cell(4, headers["Policy Expiry Date"]).value == "2026-01-01 00:00:00"
    assert sheet.cell(4, headers["Policy Expiry Date UTC"]).value == "1767225600"
    assert sheet.cell(4, headers["IPv6 Reputation Minimum"]).value == 3
    assert sheet.cell(4, headers["Effective IPv6 Reputation Minimum"]).value == 3


def test_ir_166_policy_and_nat_fields_migrate_without_fabricated_values():
    migrated = migrate_1_66_to_1_67({
        "schema_version": "1.66",
        "policies": [{"name": "policy"}],
        "nat_rules": [{"name": "nat"}],
    })

    assert migrated["schema_version"] == "1.67"
    assert migrated["policies"][0]["source_policy_expiry"] is None
    assert migrated["policies"][0]["source_extra_setting_commands"] == []
    assert migrated["nat_rules"][0]["source_policy_effective_match_vip"] is None
    assert migrated["nat_rules"][0]["source_policy_effective_match_vip_only"] is None
