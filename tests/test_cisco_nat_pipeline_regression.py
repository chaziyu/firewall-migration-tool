import io
import json
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.ir.enums import NATType
from fwmigrate.parsers.cisco_asa.extractor import extract_cisco_asa_config
from fwmigrate.parsers.cisco_ftd.extractor import extract_cisco_ftd_config
from fwmigrate.report.excel_exporter import IRExcelExporter


FIXTURES = Path(__file__).parent / "fixtures"


def _fixture(path: str) -> str:
    return (FIXTURES / path).read_text(encoding="utf-8")


def _nat_row(sheet, headers, name):
    name_column = headers["Name"]
    return next(row for row in range(4, sheet.max_row + 1) if sheet.cell(row, name_column).value == name)


def _assert_nat_fields(rule):
    for field in (
        "source", "destination", "translated_sources", "translated_destinations",
        "source_translation_mode", "destination_translation_mode",
        "original_source_ports", "translated_source_ports",
        "original_destination_ports", "translated_destination_ports",
        "source_from_interfaces", "source_to_interfaces", "identity",
        "exemption", "sequence", "requires_manual_review",
    ):
        assert hasattr(rule, field), field


def test_asa_nat_conformance_uses_public_extraction_and_excel():
    result = extract_cisco_asa_config(_fixture("cisco_asa/nat_pipeline_conformance.txt"))
    rules = result.canonical_ir.nat_rules

    assert len(rules) == 7
    assert any(rule.type == NATType.TWICE for rule in rules)
    assert any(rule.identity for rule in rules)
    assert {rule.source_attributes["section"] for rule in rules} >= {"manual", "after-auto"}
    for rule in rules:
        _assert_nat_fields(rule)

    twice = next(rule for rule in rules if rule.type == NATType.TWICE and rule.source == ["REAL_HOST"])
    assert twice.source == ["REAL_HOST"]
    assert twice.translated_sources == ["PUBLIC_HOST"]
    assert twice.destination == ["PUBLIC_HOST"]
    assert twice.translated_destinations == ["PRIVATE_SERVER"]
    assert twice.original_destination_ports == []

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()), data_only=False)
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = _nat_row(sheet, headers, twice.name)
    assert sheet.cell(row, headers["Sequence"]).value == twice.sequence
    assert sheet.cell(row, headers["Source Section"]).value == "manual"
    assert sheet.cell(row, headers["Original Source"]).value == "REAL_HOST"
    assert sheet.cell(row, headers["Translated Source"]).value == "PUBLIC_HOST"
    assert sheet.cell(row, headers["Original Destination"]).value == "PUBLIC_HOST"
    assert sheet.cell(row, headers["Translated Destination"]).value == "PRIVATE_SERVER"


def test_fmc_nat_conformance_preserves_twice_nat_and_order():
    result = extract_cisco_ftd_config(_fixture("cisco_ftd/fmc_nat_pipeline_conformance.json"))
    rules = result.canonical_ir.nat_rules

    assert [rule.sequence for rule in rules] == [1, 2]
    twice = rules[0]
    assert twice.type == NATType.TWICE
    assert twice.source == ["InsideHost"]
    assert twice.translated_sources == ["PublicHost"]
    assert twice.destination == ["PublicHost"]
    assert twice.translated_destinations == ["InsideHost"]
    assert twice.original_destination_ports[0].start == 80
    assert twice.translated_destination_ports[0].start == 8080
    for rule in rules:
        _assert_nat_fields(rule)


def test_fdm_nat_conformance_uses_public_entry_point():
    result = extract_cisco_ftd_config(_fixture("cisco_ftd/fdm_nat_pipeline_conformance.json"))
    rules = result.canonical_ir.nat_rules

    assert result.input_source_type == "fdm-rest-bundle"
    assert [rule.sequence for rule in rules] == [1, 2]
    assert rules[0].translated_sources == ["PublicHost"]
    assert rules[1].type == NATType.TWICE
    assert rules[1].translated_destinations == ["InsideHost"]
    for rule in rules:
        _assert_nat_fields(rule)


def test_ftd_text_unsupported_nat_never_disappears():
    result = extract_cisco_ftd_config("""
object network REAL
 host 10.0.0.10
nat (inside,outside) source dynamic REAL interface
""")

    assert result.generation_safe is False
    assert result.migration_complete is False
    assert any("nat (inside,outside)" in (item.raw_capture or "") for item in result.unsupported_items)
