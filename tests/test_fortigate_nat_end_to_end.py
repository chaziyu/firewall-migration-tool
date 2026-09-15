import io

import pytest
from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


INTERFACES = """
config system interface
    edit "LAN"
        set role lan
        set ip 10.0.0.1 255.255.255.0
    next
    edit "WAN"
        set role wan
        set ip 203.0.113.10 255.255.255.0
    next
end
config system zone
    edit "LAN"
        set interface "LAN"
    next
    edit "WAN"
        set interface "WAN"
    next
end
"""


def _workbook(result):
    return load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )


def _row(sheet, headers, column, value):
    return next(
        row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers[column]).value == value
    )


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[3] if cell.value}


def test_basic_policy_snat_survives_public_extraction_and_excel():
    result = extract_fortigate_config(INTERFACES + """
config firewall policy
    edit 10
        set name "LAN_TO_WAN"
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    rule = result.canonical_ir.nat_rules[0]
    assert rule.source_policy_reference == "10"
    assert rule.translated_sources == ["203.0.113.10"]
    assert rule.migration_status == "NORMALIZED", rule.review_reasons
    assert rule.safe_for_target_generation is True

    sheet = _workbook(result)["NAT Rules"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Source Policy ID", "10")
    assert sheet.cell(row, headers["Migration Status"]).value == "NORMALIZED"
    assert sheet.cell(row, headers["Manual Review"]).value == "FALSE"


@pytest.mark.parametrize(
    ("pool_config", "expected_status", "expected_translated_sources"),
    [
        (
            """
config firewall ippool
    edit "POOL"
        set startip 203.0.113.20
        set endip 203.0.113.30
    next
end
""",
            "NORMALIZED",
            ["203.0.113.20-203.0.113.30"],
        ),
        (
            """
config firewall ippool
    edit "POOL"
        set type port-block-allocation
        set startip 203.0.113.20
        set endip 203.0.113.30
        set exclude-ip "203.0.113.25"
    next
end
""",
            "PARTIALLY_NORMALIZED",
            [],
        ),
    ],
)
def test_pool_classification_controls_nat_translation_projection(
    pool_config, expected_status, expected_translated_sources
):
    result = extract_fortigate_config(INTERFACES + pool_config + """
config firewall policy
    edit 20
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
        set ippool enable
        set poolname "POOL"
    next
end
""")

    pool = result.canonical_ir.ip_pools[0]
    rule = result.canonical_ir.nat_rules[0]
    assert pool.migration_status == expected_status
    assert rule.migration_status == expected_status
    assert rule.translated_sources == expected_translated_sources
    assert rule.source_pool_references == ["POOL"]
    if expected_status != "NORMALIZED":
        assert rule.safe_for_target_generation is False


def test_central_nat_unknown_is_preserved_as_partial_manual_review():
    result = extract_fortigate_config("""
config firewall central-snat-map
    edit 30
        set srcintf "LAN"
        set dstintf "WAN"
        set orig-addr "all"
        set dst-addr "all"
        set nat-ippool "POOL"
    next
end
""")

    assert len(result.canonical_ir.nat_rules) == 1
    rule = result.canonical_ir.nat_rules[0]
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.requires_manual_review is True
    assert any("effective central-nat mode cannot be proven" in reason for reason in rule.review_reasons)
    assert result.generation_safe is False

    sheet = _workbook(result)["NAT Rules"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Name", "central-snat-30")
    assert sheet.cell(row, headers["Review Reasons"]).value


def test_central_nat_enabled_does_not_combine_policy_nat():
    result = extract_fortigate_config(INTERFACES + """
config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 40
        set srcintf "LAN"
        set dstintf "WAN"
        set orig-addr "all"
        set dst-addr "all"
    next
end
config firewall policy
    edit 40
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    assert len(result.canonical_ir.nat_rules) == 1
    assert result.canonical_ir.nat_rules[0].type.value == "central"
    assert result.canonical_ir.nat_rules[0].source_policy_reference == "40"


def test_vip_group_unresolved_member_taints_correlated_nat_and_excel_inventory():
    result = extract_fortigate_config(INTERFACES + """
config firewall vip
    edit "VIP_OK"
        set extip 203.0.113.80
        set mappedip "10.0.0.80"
        set extintf "WAN"
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set member "VIP_OK" "VIP_MISSING"
    next
end
config firewall policy
    edit 50
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "ALL"
        set action accept
    next
end
""")

    group = result.canonical_ir.virtual_ip_groups[0]
    rule = result.canonical_ir.nat_rules[0]
    assert group.members == ["VIP_OK", "VIP_MISSING"]
    assert group.unresolved_members == ["VIP_MISSING"]
    assert group.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.source_vip_group_reference == "VIP_GROUP"
    assert rule.source_vip_reference == "VIP_OK"
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.safe_for_target_generation is False

    workbook = _workbook(result)
    group_sheet = workbook["VIP Groups"]
    group_headers = _headers(group_sheet)
    group_row = _row(group_sheet, group_headers, "Name", "VIP_GROUP")
    assert group_sheet.cell(group_row, group_headers["Unresolved Members"]).value == "VIP_MISSING"
    nat_sheet = workbook["NAT Rules"]
    nat_headers = _headers(nat_sheet)
    nat_row = _row(nat_sheet, nat_headers, "Source Policy ID", "50")
    assert nat_sheet.cell(nat_row, nat_headers["Review Reasons"]).value


def test_vdom_scopes_same_nat_names_and_excel_exposes_context():
    result = extract_fortigate_config("""
config vdom
    edit "VDOM_A"
        config firewall ippool
            edit "POOL"
                set startip 198.51.100.10
                set endip 198.51.100.10
            next
        end
        config firewall policy
            edit 60
                set srcintf "lan"
                set dstintf "wan"
                set srcaddr "all"
                set dstaddr "all"
                set service "ALL"
                set action accept
                set nat enable
                set ippool enable
                set poolname "POOL"
            next
        end
    next
    edit "VDOM_B"
        config firewall ippool
            edit "POOL"
                set startip 198.51.100.20
                set endip 198.51.100.20
            next
        end
        config firewall policy
            edit 60
                set srcintf "lan"
                set dstintf "wan"
                set srcaddr "all"
                set dstaddr "all"
                set service "ALL"
                set action accept
                set nat enable
                set ippool enable
                set poolname "POOL"
            next
        end
    next
end
""")

    rules = {
        rule.source_context: rule
        for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "60"
    }
    assert set(rules) == {"VDOM_A", "VDOM_B"}
    assert rules["VDOM_A"].translated_sources == ["198.51.100.10"]
    assert rules["VDOM_B"].translated_sources == ["198.51.100.20"]

    pool_sheet = _workbook(result)["IP Pools"]
    headers = _headers(pool_sheet)
    rows = {
        pool_sheet.cell(row, headers["Source VDOM"]).value
        for row in range(4, pool_sheet.max_row + 1)
        if pool_sheet.cell(row, headers["Name"]).value == "POOL"
    }
    assert rows == {"VDOM_A", "VDOM_B"}


def test_ipv6_vip_ports_and_translation_are_kept_end_to_end():
    result = extract_fortigate_config("""
config firewall vip6
    edit "VIP6"
        set extip 2001:db8:1::10
        set mappedip 2001:db8:2::10
        set extintf "wan"
        set portforward enable
        set protocol tcp
        set extport 443
        set mappedport 8443
    next
end
config firewall policy6
    edit 70
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "VIP6"
        set action accept
    next
end
""")

    vip = result.canonical_ir.virtual_ips[0]
    assert vip.address_family == "ipv6"
    assert vip.external_ip == "2001:db8:1::10"
    assert vip.mapped_ips == ["2001:db8:2::10"]
    assert vip.external_port == "443"
    assert vip.mapped_port == "8443"
