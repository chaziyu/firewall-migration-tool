import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


ADDRESS6_TEMPLATE_CONFIG = '''
config firewall address6-template
    edit "site-template"
        set ip6 2001:db8:100::/56
        set subnet-segment-count 2
        set fabric-object enable
        config subnet-segment
            edit "site"
                set bits 8
                set exclusive enable
                set name "site-id"
                config values
                    edit "branch-a"
                        set value "10"
                    next
                    edit "branch-b"
                        set value "20"
                    next
                end
            next
            edit "host"
                set bits 64
                set exclusive disable
                set name "host-id"
            next
        end
    next
end
config firewall address6
    edit "site-host"
        set type template
        set host 2001:db8:100::10
        set host-type specific
        set template "site-template"
    next
    edit "missing-template-host"
        set type template
        set host-type any
        set template "missing-template"
    next
end
'''


def test_address6_template_uses_fortios_746_typed_schema():
    parsed = parse_fortigate_config(ADDRESS6_TEMPLATE_CONFIG)

    assert len(parsed.address6_templates) == 1
    template = parsed.address6_templates[0]
    assert template.name == "site-template"
    assert template.ip6 == "2001:db8:100::/56"
    assert template.subnet_segment_count == 2
    assert template.fabric_object == "enable"
    assert len(template.subnet_segments) == 2

    site = template.subnet_segments[0]
    assert site.source_id == "site"
    assert site.bits == 8
    assert site.exclusive == "enable"
    assert site.name == "site-id"
    assert [(item.source_id, item.value) for item in site.values] == [
        ("branch-a", "10"),
        ("branch-b", "20"),
    ]

    host = template.subnet_segments[1]
    assert host.source_id == "host"
    assert host.bits == 64
    assert host.exclusive == "disable"
    assert host.name == "host-id"
    assert host.values == []

    source_addresses = {item.name: item for item in parsed.addresses}
    assert source_addresses["site-host"].host_type == "specific"
    assert source_addresses["site-host"].template == "site-template"
    assert source_addresses["missing-template-host"].host_type == "any"


def test_address6_template_reference_is_context_scoped_and_auditable():
    result = extract_fortigate_config(ADDRESS6_TEMPLATE_CONFIG)

    resolved = [
        item
        for item in result.dependencies
        if item.source_path == "firewall address6"
        and item.source_field == "template"
        and item.reference == "site-template"
    ]
    assert resolved
    assert resolved[0].result == "RESOLVED"
    assert resolved[0].target_path == "firewall address6-template"

    unresolved = [
        item
        for item in result.dependencies
        if item.source_path == "firewall address6"
        and item.source_field == "template"
        and item.reference == "missing-template"
    ]
    assert unresolved
    assert unresolved[0].result == "UNRESOLVED"

    addresses = {item.name: item for item in result.canonical_ir.addresses}
    template_address = addresses["site-host"]
    assert template_address.source_attributes["template"] == "site-template"
    assert template_address.source_attributes["host"] == "2001:db8:100::10"
    assert template_address.source_attributes["host_type"] == "specific"
    assert template_address.subnet is None
    assert template_address.requires_manual_review is True


def test_address6_template_nested_source_data_reaches_excel():
    result = extract_fortigate_config(ADDRESS6_TEMPLATE_CONFIG)
    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["FortiGate Source Configuration"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    rows = [
        row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Source Path"]).value == "firewall address6-template"
        and sheet.cell(row, headers["Object"]).value == "site-template"
    ]
    settings = {
        sheet.cell(row, headers["Setting"]).value: sheet.cell(row, headers["Value"]).value
        for row in rows
    }
    assert settings["ip6"] == "2001:db8:100::/56"
    assert str(settings["subnet-segment-count"]) == "2"
    assert settings["fabric-object"] == "enable"
    assert any(
        sheet.cell(row, headers["Setting"]).value == "bits"
        and "subnet-segment" in (sheet.cell(row, headers["Parent / Subsection"]).value or "")
        for row in rows
    )
    assert any(
        sheet.cell(row, headers["Setting"]).value == "value"
        and sheet.cell(row, headers["Value"]).value in {"10", "20"}
        for row in rows
    )
