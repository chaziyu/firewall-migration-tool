from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


CONFIG = """
config firewall internet-service-addition
    edit 100
        set comment "Added service"
        config entry
            edit 1
                set addr-mode ipv4
                set protocol 6
                config port-range
                    edit 1
                        set start-port 80
                        set end-port 443
                    next
                end
            next
        end
    next
end
"""


def test_addition_survives_model_ir_coverage_and_excel():
    fg = parse_fortigate_config(CONFIG)
    addition = fg.internet_service_additions[0]
    assert addition.id == 100
    assert addition.entries[0].port_ranges[0].start_port == 80

    ir = FGToIRTransformer(fg).transform()
    assert ir.internet_service_additions[0].migration_status == "EXTRACT_ONLY"

    result = extract_fortigate_config(CONFIG)
    assert all(section.status.value == "EXTRACT_ONLY" for section in result.source_sections)
    workbook = load_workbook(BytesIO(IRExcelExporter(result.canonical_ir, result).generate()))
    assert workbook["IS Additions"].cell(4, 1).value == 100
    assert workbook["IS Addition Ports"].cell(4, 4).value == 80


def test_addition_accepts_ipv6_addr_mode():
    fg = parse_fortigate_config(CONFIG.replace("addr-mode ipv4", "addr-mode ipv6"))

    assert fg.internet_service_additions[0].entries[0].addr_mode == "ipv6"


def test_addition_rejects_both_but_preserves_source_and_ir_review_data():
    fg = parse_fortigate_config("""
config firewall internet-service-addition
    edit 100
        config entry
            edit 1
                set addr-mode both
            next
        end
    next
end
""")
    entry = fg.internet_service_additions[0].entries[0]

    assert entry.addr_mode is None
    assert entry.extra_settings["invalid_fields"]["addr_mode"] == "both"

    ir_entry = FGToIRTransformer(fg).transform().internet_service_additions[0].entries[0]
    assert ir_entry.addr_mode is None
    assert ir_entry.source_attributes["invalid_fields"]["addr_mode"] == "both"
