import io

from openpyxl import load_workbook

from fwmigrate.ir.core import IRNATRule, IRRoute
from fwmigrate.ir.enums import IRRouteNextHopType, NATTranslationMode, NATType
from fwmigrate.parsers.palo_alto import PANOSSourceParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def _nat_config(rules: str) -> str:
    return f"""<config><devices><entry name='fw'><vsys><entry name='vsys1'>
      <rulebase><nat><rules>{rules}</rules></nat></rulebase>
    </entry></vsys></entry></devices></config>"""


def _nat_match(name: str, translation: str = "") -> str:
    return f"""<entry name='{name}'><from><member>trust</member></from>
      <to><member>untrust</member></to><source><member>any</member></source>
      <destination><member>any</member></destination><service>any</service>{translation}</entry>"""


def test_new_ir_nat_and_route_fields_are_serializable():
    nat = IRNATRule(
        name="identity", type=NATType.SOURCE,
        source_translation_mode=NATTranslationMode.NONE,
        source_translation_bidirectional=True,
    )
    persistent_nat = IRNATRule(
        name="persistent", type=NATType.SOURCE,
        source_translation_mode=NATTranslationMode.PERSISTENT_DYNAMIC_IP_AND_PORT,
    )
    route = IRRoute(name="logical", next_hop_type=IRRouteNextHopType.NEXT_LR)

    assert nat.model_dump(mode="json")["source_translation_mode"] == "none"
    assert nat.model_dump(mode="json")["source_translation_bidirectional"] is True
    assert persistent_nat.model_dump(mode="json")["source_translation_mode"] == (
        "persistent-dynamic-ip-and-port"
    )
    assert route.model_dump(mode="json")["next_hop_type"] == "next-lr"


def test_pan_nat_preserves_order_identity_dynamic_ip_and_static_direction():
    result = PANOSSourceParser().extract(_nat_config("".join([
        _nat_match("identity"),
        _nat_match("dynamic", "<source-translation><dynamic-ip><translated-address><member>203.0.113.10</member></translated-address></dynamic-ip></source-translation>"),
        _nat_match("static", "<source-translation><static-ip><translated-address>203.0.113.20</translated-address><bi-directional>yes</bi-directional></static-ip></source-translation>"),
    ])))
    rules = result.canonical_ir.nat_rules

    assert [rule.name for rule in rules] == ["identity", "dynamic", "static"]
    assert rules[0].source_translation_mode == NATTranslationMode.NONE
    assert rules[1].source_translation_mode == NATTranslationMode.DYNAMIC_IP
    assert rules[2].source_translation_bidirectional is True


def test_pan_persistent_dipp_is_exported_as_distinct_translation_mode():
    result = PANOSSourceParser().extract(_nat_config(_nat_match(
        "persistent",
        "<source-translation><persistent-dynamic-ip-and-port><translated-address>203.0.113.30</translated-address></persistent-dynamic-ip-and-port></source-translation>",
    )))
    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    assert sheet.cell(4, headers["Source Translation Mode"]).value == (
        "persistent-dynamic-ip-and-port"
    )


def test_default_override_and_pbf_are_canonical_and_exported():
    xml = """<config><devices><entry name='fw'><network>
      <interface><ethernet><entry name='ethernet1/1'/></ethernet></interface>
    </network><vsys><entry name='vsys1'>
      <zone><entry name='trust'/></zone>
      <rulebase><default-security-rules><rules><entry name='interzone-default'>
        <action>deny</action><disabled>yes</disabled><log-end>yes</log-end>
      </entry></rules></default-security-rules>
      <pbf><rules><entry name='pbf'><from><zone><member>trust</member></zone></from>
        <action><forward><egress-interface>ethernet1/1</egress-interface></forward></action>
      </entry></rules></pbf></rulebase>
    </entry></vsys></entry></devices></config>"""
    result = PANOSSourceParser().extract(xml)
    ir = result.canonical_ir
    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))

    assert ir.default_security_rules[0].action.value == "deny"
    assert ir.default_security_rules[0].disabled is True
    assert ir.pbf_rules[0].migration_status == "NORMALIZED"
    assert "Default Security Rules" in workbook.sheetnames
    nat_headers = {cell.value for cell in workbook["NAT Rules"][3]}
    assert {"Destination Translation Mode", "Static NAT Bi-directional"} <= nat_headers


def test_empty_pan_zone_keeps_explicit_type_and_excel_projection():
    result = PANOSSourceParser().extract("""
    <config><vsys><entry name="vsys1"><zone>
      <entry name="empty-l3"><network><layer3 /></network></entry>
      <entry name="empty-l2"><network><layer2 /></network></entry>
    </zone></entry></vsys></config>
    """)

    zones = {zone.name: zone for zone in result.canonical_ir.zones}
    assert zones["empty-l3"].zone_type == "layer3"
    assert zones["empty-l3"].interfaces == []
    assert zones["empty-l2"].zone_type == "layer2"
    assert zones["empty-l2"].interfaces == []
    assert all(
        "Multiple effective network types" not in reason
        for zone in zones.values()
        for reason in zone.review_reasons
    )

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["Zones"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Name"]).value: sheet.cell(row, headers["Zone Type"]).value
        for row in range(4, sheet.max_row + 1)
    }
    assert rows["empty-l3"] == "layer3"
    assert rows["empty-l2"] == "layer2"


def test_pan_routing_instance_references_are_typed_and_fail_closed():
    parser = PANOSSourceParser()
    result = parser.extract("""
    <config><devices><entry name="fw"><network>
      <virtual-router><entry name="vr-good"><routing-table><ip><static-route>
        <entry name="route-vr-good"><destination>198.51.100.0/24</destination><nexthop><next-vr>vr-good</next-vr>
        </nexthop></entry>
        <entry name="route-vr-bad"><destination>198.51.101.0/24</destination><nexthop><next-vr>vr-missing</next-vr>
        </nexthop></entry>
      </static-route></ip></routing-table></entry></virtual-router>
      <logical-router><entry name="lr-good"><vrf><entry name="blue"><routing-table><ip><static-route>
        <entry name="route-lr-good"><destination>198.51.102.0/24</destination><nexthop><next-lr>lr-good</next-lr>
        </nexthop></entry>
        <entry name="route-lr-bad"><destination>198.51.103.0/24</destination><nexthop><next-lr>lr-missing</next-lr>
        </nexthop></entry>
      </static-route></ip></routing-table></entry></vrf></entry></logical-router>
    </network></entry></devices>
    <vsys><entry name="vsys1"><rulebase><pbf><rules>
      <entry name="pbf-vr-good"><action><forward><next-vr>vr-good</next-vr></forward></action></entry>
      <entry name="pbf-vr-bad"><action><forward><next-vr>vr-missing</next-vr></forward></action></entry>
    </rules></pbf></rulebase></entry></vsys></config>
    """)

    routes = {route.name: route for route in result.canonical_ir.routes}
    assert routes["route-vr-good"].migration_status == "NORMALIZED"
    assert routes["route-vr-good"].source_attributes["pan_next_hop_reference_resolution"] == "resolved"
    assert routes["route-vr-bad"].next_hop == "vr-missing"
    assert "unresolved-next-vr-reference" in routes["route-vr-bad"].review_reasons
    assert routes["route-lr-good"].migration_status == "NORMALIZED"
    assert routes["route-lr-good"].source_attributes["pan_next_hop_reference_type"] == "logical-router"
    assert routes["route-lr-bad"].next_hop == "lr-missing"
    assert "unresolved-next-lr-reference" in routes["route-lr-bad"].review_reasons

    pbf = {rule.name: rule for rule in result.canonical_ir.pbf_rules}
    assert pbf["pbf-vr-good"].migration_status == "NORMALIZED"
    assert pbf["pbf-vr-good"].next_vr == "vr-good"
    assert pbf["pbf-vr-bad"].next_vr == "vr-missing"
    assert "unresolved-next-vr-reference" in pbf["pbf-vr-bad"].review_reasons

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["Routes"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Name"]).value: row
        for row in range(4, sheet.max_row + 1)
    }
    bad_row = rows["route-vr-bad"]
    assert sheet.cell(bad_row, headers["Next Hop"]).value == "vr-missing"
    assert "unresolved-next-vr-reference" in sheet.cell(
        bad_row, headers["Review Reasons"]
    ).value


def test_pan_typed_semantics_survive_combined_parser_to_excel_flow():
    result = PANOSSourceParser().extract("""
    <config><devices><entry name="fw"><network>
      <interface><ethernet><entry name="ethernet1/1" /></ethernet></interface>
      <virtual-router><entry name="vr-main" /></virtual-router>
    </network><vsys><entry name="vsys1">
      <zone><entry name="empty-l3"><network><layer3 /></network></entry></zone>
      <schedule><entry name="office-hours"><schedule-type><recurring>
        <daily><member>08:00-17:00</member></daily>
      </recurring></schedule-type></entry></schedule>
      <rulebase>
        <nat><rules><entry name="persistent">
          <from><member>any</member></from><to><member>any</member></to>
          <source><member>any</member></source><destination><member>any</member></destination>
          <service>any</service><source-translation><persistent-dynamic-ip-and-port>
            <translated-address>203.0.113.40</translated-address>
          </persistent-dynamic-ip-and-port></source-translation>
        </entry></rules></nat>
        <pbf><rules><entry name="typed-pbf">
          <schedule>office-hours</schedule><negate-source>yes</negate-source>
          <negate-destination>no</negate-destination>
          <action><forward><nexthop><ip-address>198.51.100.1</ip-address></nexthop>
            <next-vr>vr-main</next-vr></forward></action>
          <enforce-symmetric-return><enabled>yes</enabled><nexthop-address>
            <member>198.51.100.2</member><member>198.51.100.3</member>
          </nexthop-address></enforce-symmetric-return>
        </entry></rules></pbf>
      </rulebase>
    </entry></vsys></entry></devices></config>
    """)

    zone = next(zone for zone in result.canonical_ir.zones if zone.name == "empty-l3")
    nat = result.canonical_ir.nat_rules[0]
    pbf = result.canonical_ir.pbf_rules[0]
    assert zone.zone_type == "layer3"
    assert nat.source_translation_mode == NATTranslationMode.PERSISTENT_DYNAMIC_IP_AND_PORT
    assert pbf.next_vr == "vr-main"
    assert pbf.schedule == "office-hours"
    assert pbf.symmetric_return.next_hop_addresses == ["198.51.100.2", "198.51.100.3"]

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    zone_headers = {cell.value: cell.column for cell in workbook["Zones"][3]}
    assert workbook["Zones"].cell(4, zone_headers["Zone Type"]).value == "layer3"
    pbf_headers = {cell.value: cell.column for cell in workbook["PBF Rules"][3]}
    assert workbook["PBF Rules"].cell(4, pbf_headers["Schedule"]).value == "office-hours"
    nat_headers = {cell.value: cell.column for cell in workbook["NAT Rules"][3]}
    assert workbook["NAT Rules"].cell(4, nat_headers["Source Translation Mode"]).value == (
        "persistent-dynamic-ip-and-port"
    )
