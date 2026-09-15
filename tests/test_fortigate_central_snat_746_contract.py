import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[3] if cell.value}


def test_central_snat_746_all_official_source_fields_are_typed_and_preserved():
    parsed = parse_fortigate_config("""
config firewall central-snat-map
    edit 101
        set comments "all fields"
        set dst-addr "DST4"
        set dst-addr6 "DST6"
        set dst-port 443-444
        set dstintf "wan" "wan2"
        set nat disable
        set nat-ippool "POOL4"
        set nat-ippool6 "POOL6"
        set nat-port 40000-40010
        set nat46 enable
        set nat64 disable
        set orig-addr "SRC4"
        set orig-addr6 "SRC6"
        set orig-port 1000-2000
        set port-preserve disable
        set protocol 6
        set srcintf "lan" "lan2"
        set status disable
        set type ipv6
        set uuid 11111111-2222-3333-4444-555555555555
    next
end
""")

    rule = parsed.central_snat_rules[0]
    assert rule.id == 101
    assert rule.comments == "all fields"
    assert rule.dst_addr == ["DST4"]
    assert rule.dst_addr6 == ["DST6"]
    assert rule.dst_port == "443-444"
    assert rule.dstintf == ["wan", "wan2"]
    assert rule.nat == "disable"
    assert rule.nat_ippool == ["POOL4"]
    assert rule.nat_ippool6 == ["POOL6"]
    assert rule.nat_port == "40000-40010"
    assert rule.nat46 == "enable"
    assert rule.nat64 == "disable"
    assert rule.orig_addr == ["SRC4"]
    assert rule.orig_addr6 == ["SRC6"]
    assert rule.orig_port == "1000-2000"
    assert rule.port_preserve == "disable"
    assert rule.protocol == 6
    assert rule.srcintf == ["lan", "lan2"]
    assert rule.status == "disable"
    assert rule.type == "ipv6"
    assert rule.uuid == "11111111-2222-3333-4444-555555555555"


def test_central_snat_746_ipv4_fields_reach_canonical_nat_and_excel():
    result = extract_fortigate_config("""
config system settings
    set central-nat enable
end
config firewall ippool
    edit "POOL4"
        set startip 198.51.100.10
        set endip 198.51.100.20
    next
end
config firewall central-snat-map
    edit 101
        set comments "contract e2e"
        set dst-addr "DST4"
        set dst-port 443-444
        set dstintf "wan"
        set nat enable
        set nat-ippool "POOL4"
        set nat-port 40000-40010
        set orig-addr "SRC4"
        set orig-port 1000-2000
        set port-preserve disable
        set protocol 6
        set srcintf "lan"
        set status enable
        set type ipv4
        set uuid 11111111-2222-3333-4444-555555555555
    next
end
""")

    rule = result.canonical_ir.nat_rules[0]
    assert rule.source_policy_reference == "101"
    assert rule.enabled is True
    assert rule.source_from_interfaces == ["lan"]
    assert rule.source_to_interfaces == ["wan"]
    assert rule.source == ["SRC4"]
    assert rule.destination == ["DST4"]
    assert rule.protocol_number == 6
    assert (rule.original_source_ports[0].start, rule.original_source_ports[0].end) == (1000, 2000)
    assert (rule.original_destination_ports[0].start, rule.original_destination_ports[0].end) == (443, 444)
    assert (rule.translated_source_ports[0].start, rule.translated_source_ports[0].end) == (40000, 40010)
    assert rule.source_pool_references == ["POOL4"]
    assert rule.translated_sources == ["198.51.100.10-198.51.100.20"]
    assert rule.source_port_behavior == "translate"

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["NAT Rules"]
    headers = _headers(sheet)
    row = next(
        row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Source Policy ID"]).value == "101"
    )
    assert sheet.cell(row, headers["Source Interface"]).value == "lan"
    assert sheet.cell(row, headers["Destination Interface"]).value == "wan"
    assert sheet.cell(row, headers["Original Source"]).value == "SRC4"
    assert sheet.cell(row, headers["Original Destination"]).value == "DST4"
    assert sheet.cell(row, headers["Protocol / Number"]).value in ("tcp / 6", "6", "tcp")
    assert sheet.cell(row, headers["Original Source Port"]).value == "1000-2000"
    assert sheet.cell(row, headers["Original Destination Port"]).value == "443-444"
    assert sheet.cell(row, headers["Translated Source Port"]).value == "40000-40010"
    assert sheet.cell(row, headers["IP Pool"]).value == "POOL4"
    assert sheet.cell(row, headers["Translated Source"]).value == "198.51.100.10-198.51.100.20"
    assert sheet.cell(row, headers["Source Port Behavior"]).value == "translate"
