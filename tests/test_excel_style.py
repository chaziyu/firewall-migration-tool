from io import BytesIO
from importlib import import_module
from pathlib import Path
from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

from fwmigrate.source_reporting.excel_style import (
    append_report_row, excel_safe_value, finish_report_workbook, style_fast_sheet, write_report_summary,
)


def test_fast_style_preserves_native_cells_and_hidden_evidence_after_save():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("Name", "Additional Settings"))
    sheet.append(("source object", "retained evidence"))
    sheet.column_dimensions["B"].hidden = True
    style_fast_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    restored = load_workbook(BytesIO(output.getvalue())).active
    assert restored.sheet_view.showGridLines
    assert restored.freeze_panes == "A2"
    assert restored.auto_filter.ref == "A1:B2"
    assert restored.column_dimensions["B"].hidden
    assert restored.column_dimensions["A"].width == 28
    assert restored.column_dimensions["B"].width == 44
    assert restored.row_dimensions[1].height == 30
    assert restored["A1"].alignment.wrap_text
    assert restored["A1"].fill.fgColor.rgb.endswith("0F766E")
    assert [cell.value for cell in restored[2]] == ["source object", "retained evidence"]
    assert all(not cell.has_style for cell in restored[2])


def test_report_row_neutralizes_formula_leading_source_values_in_full_and_streaming_workbooks():
    for write_only in (False, True):
        workbook = Workbook(write_only=write_only)
        sheet = workbook.create_sheet("Report") if write_only else workbook.active
        values = ("=1+1", "+SUM(A1:A2)", "-2+3", "@cmd", "\t=1+1", "\r=1+1", "\n=1+1",
                  "  =1+1", " \t@cmd", "  +cmd", "  -cmd", "safe", 0, -2, False, None)
        append_report_row(sheet, tuple(str(index) for index in range(len(values))))
        append_report_row(sheet, values)
        output = BytesIO()
        workbook.save(output)
        restored = load_workbook(BytesIO(output.getvalue()), read_only=True).active
        row = next(restored.iter_rows(min_row=2, max_row=2, max_col=len(values)))
        assert [cell.value for cell in row] == ["'" + value for value in values[:11]] + list(values[11:])
        assert all(cell.data_type != "f" for cell in row)


@pytest.mark.parametrize("write_only", [False, True])
def test_shared_dimensions_preserve_explicit_widths_and_unstyled_source_cells(write_only):
    workbook = Workbook(write_only=write_only)
    sheet = workbook.create_sheet("Objects") if write_only else workbook.active
    sheet.column_dimensions["B"].width = 37
    sheet.column_dimensions["D"].hidden = True
    append_report_row(sheet, ("Rule ID", "Name", "Description", "Additional Settings"))
    append_report_row(sheet, (0, False, "long technical value " * 100, "'=literal"))
    if not write_only:
        style_fast_sheet(sheet)
    else:
        sheet.auto_filter.ref = "A1:D2"
    output = BytesIO()
    workbook.save(output)
    restored = load_workbook(output).active
    assert [restored.column_dimensions[key].width for key in "ABCD"] == [16, 37, 44, 44]
    assert restored.column_dimensions["D"].hidden
    assert restored.freeze_panes == "A2"
    assert restored.auto_filter.ref == "A1:D2"
    assert restored.sheet_format.defaultRowHeight == 18
    assert restored.row_dimensions[1].height == 30
    assert restored["A1"].font.sz == 10
    assert restored["A1"].alignment.wrap_text
    assert [cell.value for cell in restored[2]] == [0, False, "long technical value " * 100, "'=literal"]
    assert all(not cell.has_style for cell in restored[2])


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-cmd", "@cmd", "\t=1+1", "\r=1+1", "\n=1+1",
                                   "  =1+1", " \t@cmd", "  +cmd", "  -cmd", "  '=escaped",
                                   "'=already escaped", "safe", "|plain text", "", 0, -2, 1.5, False, None])
def test_vendor_excel_guards_use_shared_policy_without_double_escaping(value):
    converters = (
        ("fortigate.export.excel_common", "_excel_safe"),
        ("palo_alto.export.excel", "_excel_safe"),
        ("cisco_asa.presentation", "excel_value"),
        ("cisco_ftd.export.excel", "_text"),
        ("checkpoint.export.excel", "_excel_safe"),
    )
    expected = excel_safe_value(value)
    assert excel_safe_value(expected) == expected
    for module, name in converters:
        convert = getattr(import_module(f"fwmigrate.vendors.{module}"), name)
        # FortiGate deliberately presents explicit boolean source values as Yes/No.
        converted = "No" if module.startswith("fortigate") and value is False else expected
        assert convert(value) == converted


@pytest.mark.parametrize("write_only", [False, True])
def test_summary_keeps_missing_metrics_safe_values_and_only_existing_links(write_only):
    workbook = Workbook(write_only=write_only)
    if not write_only:
        workbook.remove(workbook.active)
    summary = workbook.create_sheet("Summary")
    detail = workbook.create_sheet("Vendor's Objects")
    detail.column_dimensions["B"].hidden = True
    append_report_row(detail, ("Name", "Evidence"))
    append_report_row(detail, ("object", "retained"))
    write_report_summary(
        summary, title="Vendor Configuration Report", profile="FAST" if write_only else "FULL",
        metadata=(("Hostname", "\t=1+1"), ("Version", None), ("Scope", "shared")),
        validation=(("Errors", 0), ("Warnings", None), ("Unresolved References", None)),
        inventory=(("Objects", 0, detail.title), ("Unavailable", None, "Omitted")),
        coverage=(("Unknown", None), ("Source-only", 0, detail.title), ("Evidence", "+cmd")),
        indicators=(("Indicator", "@cmd"),),
    )
    finish_report_workbook(workbook, summary_title="Legacy title does not replace a modern summary")
    output = BytesIO()
    workbook.save(output)
    restored = load_workbook(BytesIO(output.getvalue()))
    summary = restored["Summary"]
    assert summary["A1"].value == "Vendor Configuration Report"
    assert summary["B4"].value == "'\t=1+1"
    assert summary["D4"].value == "Not reported"
    assert summary["B7"].value == 0
    assert summary["D7"].value == "Not reported"
    objects = next(row for row in summary if row[0].value == "Objects")
    assert objects[1].value == 0
    assert objects[1].alignment.horizontal == "right"
    assert objects[2].hyperlink.target == "#'Vendor''s Objects'!A1"
    missing = next(row for row in summary if row[0].value == "Unavailable")
    assert missing[1].value == "Not reported"
    assert missing[2].hyperlink is None
    source_only = next(row for row in summary if row[0].value == "Source-only")
    assert source_only[2].hyperlink.target == "#'Vendor''s Objects'!A1"
    assert all(cell.data_type != "f" for row in summary for cell in row)
    assert not summary.auto_filter.ref
    assert summary.freeze_panes == "A5"
    assert summary.page_setup.fitToWidth == 1
    assert restored[detail.title].column_dimensions["B"].hidden
    assert [cell.value for cell in restored[detail.title][2]] == ["object", "retained"]


SUMMARY_CASES = (
    ("fortigate", "FortiGateSourceReporter", "fortigate/palo_alto_mvp.conf"),
    ("palo_alto", "PaloAltoSourceReporter", "palo_alto/integrated_panorama.xml"),
    ("cisco_asa", "CiscoASASourceReporter", "cisco_asa/nat_pipeline_conformance.txt"),
    ("cisco_ftd", "CiscoFTDSourceReporter", "cisco_ftd/fmc_selected_domains.json"),
    ("juniper_srx", "JuniperSRXSourceReporter", None),
    ("checkpoint", "CheckPointSourceReporter", "checkpoint/multidomain_full.json"),
)


@pytest.mark.parametrize("vendor,reporter_class,fixture", SUMMARY_CASES)
def test_vendor_summaries_share_sections_links_and_profile_counts(vendor, reporter_class, fixture):
    reporter = getattr(import_module(f"fwmigrate.vendors.{vendor}.source_report"), reporter_class)()
    source = ((Path(__file__).parent / "fixtures" / fixture).read_text(encoding="utf-8") if fixture else
              "set system host-name srx\nset interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24\n"
              "set logical-systems branch interfaces ge-0/0/1 unit 0 family inet address 198.51.100.1/24\n")
    analysis = reporter.analyze_source(source)
    inventories = []
    for profile in ("fast", "full"):
        output = BytesIO()
        reporter.export_excel(analysis, output, profile=profile, source_name="review-input.conf")
        workbook = load_workbook(BytesIO(output.getvalue()))
        summary = workbook["Summary"]
        assert workbook.sheetnames[0] == "Summary"
        for detail in workbook.worksheets[1:]:
            header_row = 3 if vendor in {"fortigate", "palo_alto"} else 1
            header = detail.cell(header_row, 1)
            assert header.fill.fgColor.rgb.endswith("0F766E")
            assert header.font.sz == 10
            assert header.alignment.wrap_text
            assert detail.row_dimensions[header_row].height == 30
            assert detail.freeze_panes
        sections = [row[0].value for row in summary if row[0].value in {
            "SOURCE DETAILS", "VALIDATION", "CONFIGURATION INVENTORY", "SOURCE COVERAGE"}]
        assert sections == ["SOURCE DETAILS", "VALIDATION", "CONFIGURATION INVENTORY", "SOURCE COVERAGE"]
        assert not summary.auto_filter.ref
        assert summary["A1"].fill.fgColor.rgb.endswith("17324D")
        assert all(cell.value != "Workbook navigation" for row in summary for cell in row)
        assert any(row[index].value == "Source File" and row[index + 1].value == "review-input.conf"
                   for row in summary for index in (0, 2, 4))
        inventory = []
        inside = False
        for row in summary:
            if row[0].value == "CONFIGURATION INVENTORY":
                inside = True
            elif row[0].value == "SOURCE COVERAGE":
                break
            elif inside and row[0].value != "Section":
                inventory.append((row[0].value, row[1].value))
                if row[2].hyperlink:
                    target = row[2].hyperlink.target[2:-4].replace("''", "'")
                    assert target in workbook.sheetnames
                    assert row[2].value == "Open"
        assert inventory
        inventories.append(dict(inventory))
    # FAST may omit empty PAN-OS sheets. Every retained count must match FULL.
    assert all(inventories[1][name] == count for name, count in inventories[0].items())


@pytest.mark.parametrize("vendor,reporter_class,fixture", SUMMARY_CASES[2:])
@pytest.mark.parametrize("profile", ["fast", "full"])
def test_reporter_collection_evidence_retains_generator_values(vendor, reporter_class, fixture, profile):
    reporter = getattr(import_module(f"fwmigrate.vendors.{vendor}.source_report"), reporter_class)()
    source = ((Path(__file__).parent / "fixtures" / fixture).read_text(encoding="utf-8")
              if fixture else "set system host-name srx")
    output = BytesIO()
    reporter.export_excel(
        reporter.analyze_source(source), output, profile=profile, collection_status="PARTIAL",
        collection_parts=iter([{"name": "empty", "status": "FAILED", "complete": False, "count": 0}]),
        collection_warnings=iter(["password evidence-secret", "\t=1+1"]),
    )
    workbook = load_workbook(BytesIO(output.getvalue()))
    assert list(workbook["Collection Evidence"].values)[1:] == [
        ("empty", "FAILED", False, 0, None),
        (None, None, None, None, "password [REDACTED]"),
        (None, None, None, None, "'\t=1+1"),
    ]
    assert "evidence-secret" not in str([row for sheet in workbook for row in sheet.values])


def test_fortigate_inventory_escapes_cli_scope_after_reload():
    from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter

    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source('config firewall address\n edit host\n set subnet 192.0.2.1 255.255.255.255\n next\nend')
    records = tuple(replace(record, cli_scope="\t=1+1") for record in analysis.extracted.source_objects)
    analysis = replace(analysis, extracted=replace(analysis.extracted, source_objects=records))
    output = BytesIO()
    reporter.export_excel(analysis, output, profile="full")
    workbook = load_workbook(BytesIO(output.getvalue()))
    sheet = workbook["FortiGate Source Inventory"]
    column = next(cell.column for cell in sheet[3] if cell.value == "CLI Scope")
    assert sheet.max_row > 3
    assert all(sheet.cell(row, column).value == "'\t=1+1" for row in range(4, sheet.max_row + 1))
    assert all(sheet.cell(row, column).data_type != "f" for row in range(4, sheet.max_row + 1))


@pytest.mark.parametrize("write_only", [False, True])
def test_legacy_summary_finalizer_preserves_title_and_navigation(write_only):
    workbook = Workbook(write_only=write_only)
    if not write_only:
        workbook.remove(workbook.active)
    summary = workbook.create_sheet("Summary")
    detail = workbook.create_sheet("Objects")
    append_report_row(summary, ("Metric", "Value"))
    append_report_row(summary, ("Objects", 1))
    append_report_row(detail, ("Name",))
    append_report_row(detail, ("object",))
    finish_report_workbook(workbook, summary_title="Legacy report")
    output = BytesIO()
    workbook.save(output)
    restored = load_workbook(BytesIO(output.getvalue()))
    assert restored["Summary"]["A1"].value == ("Metric" if write_only else "Legacy report")
    assert any(cell.hyperlink and cell.hyperlink.target == "#'Objects'!A1"
               for row in restored["Summary"] for cell in row)
    assert restored["Objects"].auto_filter.ref == "A1:A2"


@pytest.mark.parametrize("vendor,reporter_class,source,expected", [
    ("cisco_asa", "CiscoASASourceReporter",
     "hostname asa\ninterface GigabitEthernet0/1\n nameif inside\n"
     "dhcpd address 192.0.2.10-192.0.2.20 inside\ndhcpd enable inside\n"
     "time-range WORK\n periodic weekdays 09:00 to 17:00\n"
     "crypto ipsec ikev2 ipsec-proposal ESP\n protocol esp encryption aes-256\n",
     {"DHCP Servers", "Time Ranges", "IKEv2 Proposals"}),
    ("juniper_srx", "JuniperSRXSourceReporter",
     "set system host-name branch\n"
     "set system services dhcp-local-server group LAN interface ge-0/0/0.0\n"
     "set access address-assignment pool LAN family inet network 192.0.2.0/24\n"
     "set security ike proposal IKE authentication-method pre-shared-keys\n",
     {"DHCP Local Servers", "DHCP Pools", "IKE Proposals"}),
])
@pytest.mark.parametrize("profile", ["fast", "full"])
def test_configuration_inventory_counts_all_exported_sections(vendor, reporter_class, source, expected, profile):
    reporter = getattr(import_module(f"fwmigrate.vendors.{vendor}.source_report"), reporter_class)()
    output = BytesIO()
    reporter.export_excel(reporter.analyze_source(source), output, profile=profile)
    workbook = load_workbook(BytesIO(output.getvalue()))
    inventory = {}
    inside = False
    for row in workbook["Summary"]:
        if row[0].value == "CONFIGURATION INVENTORY":
            inside = True
        elif row[0].value == "SOURCE COVERAGE":
            break
        elif inside and row[0].value != "Section":
            inventory[row[0].value] = row[1].value
            assert row[2].hyperlink.target == f"#'{row[0].value}'!A1"
    excluded = {"Summary", "Review Required", "Validation", "Source Inventory", "Extraction Coverage",
                "Unsupported", "Unsupported Source", "Unresolved References", "Additional Settings"}
    assert set(inventory) == set(workbook.sheetnames) - excluded
    assert all(inventory[name] > 0 for name in expected)
    for name, count in inventory.items():
        assert count == len(list(workbook[name].values)) - 1


@pytest.mark.parametrize("summary_values", [(), (("Metric", "Value"), ("Objects", 1))])
def test_legacy_streaming_summary_needs_no_report_counters(summary_values):
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    for values in summary_values:
        summary.append(values)
    detail = workbook.create_sheet("Objects")
    append_report_row(detail, ("Name", "Value"))
    append_report_row(detail, ("object", 0))
    finish_report_workbook(workbook)
    output = BytesIO()
    workbook.save(output)
    restored = load_workbook(output)
    assert restored["Summary"].auto_filter.ref is None
    assert restored["Objects"].auto_filter.ref == "A1:B2"
    assert any(cell.hyperlink and cell.hyperlink.target == "#'Objects'!A1"
               for row in restored["Summary"] for cell in row)


@pytest.mark.parametrize("module,function", [("cisco_asa.presentation", "excel_value"),
                                             ("cisco_ftd.export.excel", "_text")])
def test_cisco_formatters_preserve_structured_values_without_python_repr(module, function):
    import json
    convert = getattr(import_module(f"fwmigrate.vendors.{module}"), function)
    assert convert(("inside", "outside")) == "inside, outside"
    assert convert(()) == "[]"
    assert json.loads(convert([None, False, 0])) == [None, False, 0]
    refs = [{"name": "inside", "source_id": "if-1"}, {"source_id": "if-2"}, {"value": 0, "enabled": False}]
    assert json.loads(convert(refs)) == refs
    assert json.loads(convert({"items": refs, "missing": None})) == {"items": refs, "missing": None}
    assert convert(("=1+1", "outside")) == "'=1+1, outside"
    if module.startswith("cisco_asa"):
        assert "hidden-secret" not in convert({"items": [{"password": "hidden-secret", "name": "safe"}]})
