import io
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from fwmigrate.ir.enums import NATTranslationMode
from fwmigrate.parsers.palo_alto.nat import _source_translation_fallback
from fwmigrate.parsers.palo_alto.parser import PANOSSourceParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_pbf_resolves_scoped_match_and_ip_netmask_next_hop_and_exports_effective_order():
    xml = """<config version='11.1.0'>
      <shared>
        <address><entry name='Scoped-Web'><ip-netmask>192.0.2.10/32</ip-netmask></entry></address>
        <service><entry name='Scoped-Service'><protocol><tcp><port>9443</port></tcp></protocol></entry></service>
      </shared>
      <devices><entry name='fw'>
        <vsys><entry name='vsys1'>
          <address><entry name='Scoped-Web'><ip-netmask>10.1.0.10/32</ip-netmask></entry></address>
          <service><entry name='Scoped-Service'><protocol><tcp><port>10443</port></tcp></protocol></entry></service>
          <rulebase><pbf><rules><entry name='object-next-hop'>
            <source><member>Scoped-Web</member></source>
            <destination><member>Scoped-Web</member></destination>
            <application><member>any</member></application>
            <service><member>Scoped-Service</member></service>
            <action><forward><nexthop><ip-address>Scoped-Web</ip-address></nexthop></forward></action>
          </entry></rules></pbf></rulebase>
        </entry></vsys>
      </entry></devices>
    </config>"""

    result = PANOSSourceParser().extract(xml)
    assert len(result.canonical_ir.pbf_rules) == 1
    rule = result.canonical_ir.pbf_rules[0]
    assert rule.source == ["vsys1::Scoped-Web"]
    assert rule.destination == ["vsys1::Scoped-Web"]
    assert rule.service == ["vsys1::Scoped-Service"]
    assert rule.next_hop == "vsys1::Scoped-Web"
    assert rule.source_attributes["pan_pbf_resolved_next_hop"] == "10.1.0.10"
    assert rule.source_attributes["effective_policy_layer"] == "local-rules"
    assert rule.source_attributes["effective_policy_rank"] == 0
    assert "vsys:vsys1" in rule.source_attributes["pan_effective_order_by_context"]

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["PBF Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Effective Layer"]).value == "local-rules"
    assert sheet.cell(4, headers["Effective Rank"]).value == 0
    assert sheet.cell(4, headers["Next Hop"]).value == "vsys1::Scoped-Web"


def test_static_route_resolves_scoped_ip_netmask_next_hop_and_typed_installation():
    xml = """<config version='11.1.0'>
      <shared><address><entry name='GW'><ip-netmask>192.0.2.254/32</ip-netmask></entry></address></shared>
      <devices><entry name='fw'>
        <network><virtual-router><entry name='vr'><routing-table><ip><static-route>
          <entry name='route-object-gw'>
            <destination>10.20.0.0/16</destination>
            <nexthop><ip-address>GW</ip-address></nexthop>
            <route-table><no-install/></route-table>
          </entry>
        </static-route></ip></routing-table></entry></virtual-router></network>
        <vsys><entry name='vsys1'>
          <address><entry name='GW'><ip-netmask>10.1.0.254/32</ip-netmask></entry></address>
        </entry></vsys>
      </entry></devices>
    </config>"""

    result = PANOSSourceParser().extract(xml)
    assert len(result.canonical_ir.routes) == 1
    route = result.canonical_ir.routes[0]
    assert route.next_hop == "vsys1::GW"
    assert route.source_attributes["pan_resolved_next_hop"] == "10.1.0.254"
    assert "next-hop-address-reference" in route.review_reasons
    assert route.installation == "no-install"
    assert "route-table-installation" not in route.review_reasons

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["Routes"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Next Hop"]).value == "vsys1::GW"
    assert sheet.cell(4, headers["Installation"]).value == "no-install"


def test_source_nat_translated_address_fallback_is_dipp_not_static():
    node = ET.fromstring(
        "<fallback><translated-address><member>198.51.100.10</member></translated-address></fallback>"
    )
    fallback = _source_translation_fallback(node)
    assert fallback.mode == NATTranslationMode.DYNAMIC_IP_AND_PORT
    assert fallback.translated_addresses == ["198.51.100.10"]
