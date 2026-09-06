import io

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import load_workbook

from fwmigrate.ir.core import IRConfig, IRMetadata
from fwmigrate.report.excel_exporter import IRExcelExporter


def _workbook(source_vendor: str):
    ir = IRConfig(
        metadata=IRMetadata(
            hostname="visibility-test",
            source_vendor=source_vendor,
        )
    )
    return load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))


def _summary_labels(workbook) -> set[str]:
    summary = workbook["Summary"]
    return {
        summary.cell(row, 1).value
        for row in range(1, summary.max_row + 1)
        if summary.cell(row, 1).value
    }


def test_fortigate_excel_hides_palo_alto_only_sheets():
    workbook = _workbook("fortigate")

    assert "FortiGate Source Configuration" in workbook.sheetnames
    assert "GlobalProtect Portals" not in workbook.sheetnames
    assert "PAN Log Servers" not in workbook.sheetnames
    assert "PAN Device Settings" not in workbook.sheetnames
    assert "Management Service Routes" not in workbook.sheetnames

    summary_labels = _summary_labels(workbook)
    assert "GlobalProtect Portals" not in summary_labels
    assert "Management Service Routes" not in summary_labels


def test_palo_alto_excel_keeps_pan_sheets_and_hides_fortigate_only_sheets():
    workbook = _workbook("palo_alto")

    assert "GlobalProtect Portals" in workbook.sheetnames
    assert "PAN Log Servers" in workbook.sheetnames
    assert "PAN Device Settings" in workbook.sheetnames
    assert "Management Service Routes" in workbook.sheetnames
    assert "FortiGate Source Configuration" not in workbook.sheetnames
    assert "Firewall Policy Source Settings" not in workbook.sheetnames
    assert "Interface Nested Configuration" not in workbook.sheetnames
    assert "FortiTokens" not in workbook.sheetnames


@pytest.mark.parametrize("source_vendor", ["cisco_asa", "checkpoint", "juniper_srx"])
def test_other_known_vendors_hide_explicit_pan_and_fortigate_only_sheets(source_vendor):
    workbook = _workbook(source_vendor)

    assert "PAN Log Servers" not in workbook.sheetnames
    assert "GlobalProtect Portals" not in workbook.sheetnames
    assert "FortiGate Source Configuration" not in workbook.sheetnames
    assert "FortiTokens" not in workbook.sheetnames


def test_unknown_vendor_preserves_historical_full_workbook():
    workbook = _workbook("legacy_unknown_vendor")

    assert "PAN Log Servers" in workbook.sheetnames
    assert "GlobalProtect Portals" in workbook.sheetnames
    assert "FortiGate Source Configuration" in workbook.sheetnames
