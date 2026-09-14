import io
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.ir.enums import NATTranslationMode, NATType
from fwmigrate.parsers.palo_alto import PANOSSourceParser
from fwmigrate.report import IRExcelExporter


FIXTURE = Path(__file__).parent / "fixtures" / "palo_alto" / "nat_pipeline_conformance.xml"


def _result():
    return PANOSSourceParser().extract(FIXTURE.read_text(encoding="utf-8"))


def test_pan_nat_conformance_fixture_preserves_typed_pipeline_semantics():
    result = _result()
    rules = {rule.name: rule for rule in result.canonical_ir.nat_rules}

    assert len(rules) == 6
    assert not any(item.status.value == "PARSE_ERROR" for item in result.inventory_items if item.domain == "nat")

    dipp = rules["dipp-pool-rule"]
    assert dipp.protocol_name == "tcp"
    assert dipp.original_source_ports[0].start == 1024
    assert dipp.original_source_ports[0].end == 65535
    assert dipp.original_destination_ports[0].start == 8443
    assert dipp.translated_sources == ["dipp-pool"]

    interface = rules["dipp-interface-rule"]
    assert interface.source_translation_mode == NATTranslationMode.DYNAMIC_IP_AND_PORT
    assert interface.source_translation_address_selection.interface == "ethernet1/1"
    assert interface.translated_sources == []

    twice = rules["static-twice-rule"]
    assert twice.type == NATType.TWICE
    assert twice.source_translation_bidirectional is True
    assert twice.source_device_binding == "both"
    assert {match.protocol for match in twice.service_matches} == {"tcp", "udp"}
    assert twice.protocol_name is None

    dynamic = rules["dynamic-dnat-rule"]
    assert dynamic.destination_translation_distribution.method == "round-robin"
    assert dynamic.destination_dns_rewrite.enabled is True
    assert dynamic.destination_dns_rewrite.direction == "reverse"

    translated_port = rules["translated-port-rule"]
    assert translated_port.translated_destination_ports[0].start == 8443
    assert translated_port.translated_source_ports == []
    assert translated_port.original_destination_ports[0].start == 443

    disabled = rules["disabled-no-translation"]
    assert disabled.enabled is False
    assert disabled.source_translation_mode.value == "none"
    assert disabled.translated_destinations == []
    assert disabled.protocol_name == "any"
    assert disabled.service_matches == []
    assert disabled.original_destination_ports == []


def test_pan_nat_unresolved_service_stays_reviewable_without_invented_protocol():
    result = PANOSSourceParser().extract("""
    <config><devices><entry name="fw"><vsys><entry name="vsys1">
      <rulebase><nat><rules><entry name="unresolved-service">
        <from><member>trust</member></from><to><member>untrust</member></to>
        <source><member>any</member></source><destination><member>any</member></destination>
        <service>missing-service</service>
        <source-translation><dynamic-ip-and-port><interface-address>
          <interface>ethernet1/1</interface>
        </interface-address></dynamic-ip-and-port></source-translation>
      </entry></rules></nat></rulebase>
    </entry></vsys></entry></devices></config>
    """)

    rule = result.canonical_ir.nat_rules[0]
    assert rule.services == ["missing-service"]
    assert rule.service_matches == []
    assert rule.protocol_name is None
    assert rule.original_destination_ports == []
    assert "unresolved-service" in rule.review_reasons


def test_pan_nat_conformance_fixture_projects_ir_without_repairing_it_in_excel():
    workbook = load_workbook(io.BytesIO(IRExcelExporter(_result().canonical_ir).generate()))
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    assert {
        "Service Matches", "Destination Distribution", "DNS Rewrite", "Device Binding",
        "Rulebase Position", "Source Order", "Effective Rank",
    }.issubset(headers)

    rows = {
        sheet.cell(row, headers["Name"]).value: row
        for row in range(4, sheet.max_row + 1)
    }
    assert sheet.cell(rows["dipp-interface-rule"], headers["Source Translation Interface"]).value == "ethernet1/1"
    assert sheet.cell(rows["dipp-interface-rule"], headers["Translated Source"]).value in (None, "")
    assert sheet.cell(rows["dynamic-dnat-rule"], headers["Destination Distribution"]).value == "round-robin"
    assert "reverse" in str(sheet.cell(rows["dynamic-dnat-rule"], headers["DNS Rewrite"]).value)
    assert sheet.cell(rows["translated-port-rule"], headers["Translated Destination Port"]).value == "8443"
    assert sheet.cell(rows["translated-port-rule"], headers["Translated Source Port"]).value in (None, "")
    assert sheet.cell(rows["disabled-no-translation"], headers["Enabled"]).value == "FALSE"
