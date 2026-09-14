import io

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from fwmigrate.ir.core import IRConfig, IRMetadata
from fwmigrate.report import IRExcelExporter
from fwmigrate.report.excel_readability import ReadableFortiGateExcelExporter


class _ThresholdExporter(ReadableFortiGateExcelExporter):
    LARGE_SHEET_ROW_THRESHOLD = 3

    def __init__(self):
        super().__init__(IRConfig(metadata=IRMetadata(hostname="test", source_vendor="fortigate")))
        self.reorder_calls = 0
        self.fit_calls = 0

    def _reorder_and_group_columns(self, sheet, groups):
        self.reorder_calls += 1
        return super()._reorder_and_group_columns(sheet, groups)

    def _fit_visible_row_heights(self, sheet):
        self.fit_calls += 1
        return super()._fit_visible_row_heights(sheet)


def _sheet(title, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append([title])
    sheet.append(["Inventory"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    return workbook, sheet


def test_large_sheet_policy_uses_record_count_and_strict_threshold():
    exporter = _ThresholdExporter()
    _, below = _sheet("Addresses", ("Name",), (("a",), ("b",)))
    _, exact = _sheet("Addresses", ("Name",), (("a",), ("b",), ("c",)))
    _, above = _sheet(
        "Addresses",
        ("Name",),
        (("a",), ("b",), ("c",), ("d",)),
    )

    assert exporter._is_large_sheet(below) is False
    assert exporter._is_large_sheet(exact) is False
    assert exporter._is_large_sheet(above) is True


def test_large_grouped_sheet_keeps_cells_and_skips_copy_and_row_scan():
    workbook, sheet = _sheet(
        "Addresses",
        (
            "Name",
            "Source UUID",
            "Type",
            "Value",
            "Migration Status",
            "Manual Review",
            "Additional Settings",
            "Description",
        ),
        ((f"address-{row}", f"uuid-{row}", "ipmask", "10.0.0.1/32", "NORMALIZED", "No", "settings", "desc") for row in range(4)),
    )
    sheet["B4"].font = Font(bold=True)
    sheet["B4"].hyperlink = "https://example.test/source"
    sheet["B4"].comment = Comment("source metadata", "test")
    sheet.row_dimensions[4].height = 47

    exporter = _ThresholdExporter()
    exporter._apply_sheet_view(sheet)

    assert [sheet.cell(3, column).value for column in range(1, 9)] == [
        "Name",
        "Source UUID",
        "Type",
        "Value",
        "Migration Status",
        "Manual Review",
        "Additional Settings",
        "Description",
    ]
    assert sheet["B4"].value == "uuid-0"
    assert sheet["B4"].font.bold is True
    assert sheet["B4"].hyperlink.target == "https://example.test/source"
    assert sheet["B4"].comment.text == "source metadata"
    assert sheet.row_dimensions[4].height == 47
    assert exporter.reorder_calls == 0
    assert exporter.fit_calls == 0

    for letter in ("A", "C", "D", "E", "F", "H"):
        assert sheet.column_dimensions[letter].hidden is False
        assert sheet.column_dimensions[letter].outlineLevel == 0
    for letter in ("B", "G"):
        assert sheet.column_dimensions[letter].hidden is True
        assert sheet.column_dimensions[letter].outlineLevel == 1


def test_small_grouped_sheet_keeps_legacy_reorder_and_row_fit():
    _, sheet = _sheet(
        "Addresses",
        ("Name", "Source UUID", "Type", "Value", "Migration Status"),
        (("server-1", "uuid", "ipmask", "10.0.0.1/32", "NORMALIZED"),),
    )
    exporter = _ThresholdExporter()
    exporter._apply_sheet_view(sheet)

    assert exporter.reorder_calls == 1
    assert exporter.fit_calls == 1
    assert [sheet.cell(3, column).value for column in range(1, 6)] == [
        "Name",
        "Type",
        "Value",
        "Migration Status",
        "Source UUID",
    ]


def test_empty_column_scan_preserves_zero_false_and_headerless_state():
    _, sheet = _sheet(
        "Certificates",
        ("Name", "Empty", "Zero", "False", "Text"),
        (
            ("cert-1", None, 0, False, "value"),
            ("cert-2", "", None, None, ""),
        ),
    )
    sheet.column_dimensions["F"].hidden = True
    sheet.cell(4, 6, "headerless data")

    _ThresholdExporter()._hide_empty_columns(sheet)

    assert sheet.column_dimensions["A"].hidden is False
    assert sheet.column_dimensions["B"].hidden is True
    assert sheet.column_dimensions["C"].hidden is False
    assert sheet.column_dimensions["D"].hidden is False
    assert sheet.column_dimensions["E"].hidden is False
    assert sheet.column_dimensions["F"].hidden is True


def test_review_and_evidence_are_analyzed_in_one_workbook_pass():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Policies"
    sheet.append(["Policies"])
    sheet.append(["Inventory"])
    sheet.append(("Name", "Migration Status", "Manual Review", "Review Reasons"))
    sheet.append(("normalized", "NORMALIZED", "No", ""))
    sheet.append(("extract-only", "EXTRACT_ONLY", "Yes", "Source-only"))
    sheet.append(("unsupported", "UNSUPPORTED", "No", "Unsupported feature"))
    sheet.append(("partial", "PARTIALLY_NORMALIZED", "Yes", "Map manually"))

    exporter = _ThresholdExporter()
    analysis = exporter._analyze_review_workbook(workbook)

    assert [row[5] for row in analysis.review_rows] == [6, 7]
    assert [row[5] for row in analysis.evidence_rows] == [5]
    assert analysis.evidence_rows[0][1:4] == (
        "extract-only",
        "Source-only",
        "Extract only",
    )

    exporter._build_review_required(workbook)
    review = workbook["Review Required"]
    evidence = workbook["Extraction Evidence"]
    assert review.cell(4, 5).value == "Unsupported"
    assert review.cell(4, 6).hyperlink.target == "#'Policies'!A6"
    assert evidence.cell(4, 5).hyperlink.target == "#'Policies'!A5"


def test_debug_export_timing_is_opt_in_and_does_not_change_api(caplog):
    caplog.set_level("DEBUG", logger="fwmigrate.report.excel_optimized")
    ir = IRConfig(metadata=IRMetadata(hostname="timing", source_vendor="fortigate"))

    workbook_bytes = IRExcelExporter(ir).generate()

    assert workbook_bytes.startswith(b"PK")
    timing_logs = [
        record.getMessage()
        for record in caplog.records
        if record.name == "fwmigrate.report.excel_optimized"
    ]
    assert timing_logs
    assert "inventory workbook construction" in timing_logs[-1]
    assert "final XLSX serialization" in timing_logs[-1]


def test_public_export_still_round_trips_after_fast_path_changes():
    workbook = load_workbook(
        io.BytesIO(
            IRExcelExporter(
                IRConfig(metadata=IRMetadata(hostname="round-trip", source_vendor="fortigate"))
            ).generate()
        )
    )
    assert "Summary" in workbook.sheetnames
    assert "Review Required" in workbook.sheetnames
