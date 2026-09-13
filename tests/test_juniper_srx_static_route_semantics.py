from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.parsers.juniper_srx.coverage import build_juniper_dependencies
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_juniper_route_defaults_install_state_and_interface_dependencies():
    parser = JuniperSRXParser("""
    set interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24
    set routing-options static route default next-hop 192.0.2.254
    set routing-options static route 198.51.100.0/24 install
    set routing-options static route 203.0.113.0/24 no-install
    set routing-options static route 192.0.2.0/24 next-hop ge-0/0/0.0
    set routing-options static route 192.0.3.0/24 next-hop ge-0/0/9.0
    """)
    ir = parser.transform_to_ir()
    routes = {route.source_destination: route for route in ir.routes}

    assert routes["default"].destination == "0.0.0.0/0"
    assert routes["default"].source_attributes["junos_source_destination"] == "default"
    assert routes["198.51.100.0/24"].installation == "install"
    assert routes["203.0.113.0/24"].installation == "no-install"
    assert routes["203.0.113.0/24"].enabled is None

    dependencies = build_juniper_dependencies(parser.config)
    route_deps = [item for item in dependencies if item.source_field == "next-hop"]
    assert [(item.reference, item.result) for item in route_deps] == [
        ("ge-0/0/0.0", "RESOLVED"),
        ("ge-0/0/9.0", "UNRESOLVED"),
    ]

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["Routes"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(row for row in sheet.iter_rows(min_row=4, values_only=False) if row[headers["Source Destination"] - 1].value == "default")
    assert row[headers["Destination Prefix (Normalized)"] - 1].value == "0.0.0.0/0"
