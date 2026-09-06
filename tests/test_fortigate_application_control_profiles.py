from fwmigrate.parsers.fortigate.parser import parse_fortigate_config


def test_application_ids_filters_risk_and_overrides_are_typed():
    parsed = parse_fortigate_config('''config application list
    edit "apps"
        config entries
            edit 1
                set application 12345
                set category 28
                set risk 4
                set action block
            next
        end
        config filters
            edit 1
                set category 2
                set risk 5
            next
        end
        config overrides
            edit 1
                set application 12345
                set action allow
            next
        end
    next
end
''')
    profile = parsed.application_lists[0]
    assert profile.entries[0].application == [12345]
    assert profile.entries[0].application_id == 12345
    assert profile.entries[0].category == [28]
    assert profile.entries[0].risk == [4]
    assert profile.entries[0].action == "block"
    assert profile.filters[0].category == [2]
    assert profile.filters[0].risk == [5]
    assert profile.overrides[0].application == [12345]
    assert profile.overrides[0].action == "allow"


def test_application_entries_support_multiple_categories_and_applications():
    parsed = parse_fortigate_config('''config application list
    edit "block-high-risk"
        set unknown-application-log enable
        config entries
            edit 1
                set category 2 6 7
            next
            edit 2
                set application 11414 11767 15722
                set action pass
            next
        end
    next
end
''')
    profile = parsed.application_lists[0]
    assert profile.name == "block-high-risk"
    assert len(profile.entries) == 2
    assert profile.entries[0].category == [2, 6, 7]
    assert profile.entries[1].application == [11414, 11767, 15722]
    assert profile.entries[1].action == "pass"


def test_application_fixture_parses_successfully():
    from pathlib import Path
    fixture_path = Path(__file__).parent / "fixtures" / "fortigate" / "application_list_multiple_categories.conf"
    content = fixture_path.read_text(encoding="utf-8")
    parsed = parse_fortigate_config(content)
    profile = parsed.application_lists[0]
    assert profile.name == "block-high-risk"
    assert profile.entries[0].category == [2, 6, 7]
    assert profile.entries[1].application == [11414, 11767, 15722]
    assert profile.entries[1].action == "pass"


def test_application_list_single_category():
    config = '''config application list
    edit "Test"
        config entries
            edit 1
                set category 2
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 1
    entry = profile.entries[0]
    assert entry.category == [2]
    assert isinstance(entry.category, list)
    assert entry.category[0] == 2
    assert isinstance(entry.category[0], int)


def test_application_list_multiple_categories():
    config = '''config application list
    edit "Test"
        config entries
            edit 1
                set category 2 6 7
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 1
    entry = profile.entries[0]
    assert entry.category == [2, 6, 7]
    assert isinstance(entry.category, list)
    assert all(isinstance(val, int) for val in entry.category)


def test_application_list_multiple_applications():
    config = '''config application list
    edit "Test"
        config entries
            edit 1
                set application 1234 5678
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 1
    entry = profile.entries[0]
    assert entry.application == [1234, 5678]
    assert isinstance(entry.application, list)
    assert all(isinstance(val, int) for val in entry.application)


def test_application_list_multiple_risks():
    config = '''config application list
    edit "Test"
        config entries
            edit 1
                set risk 2 3 4
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 1
    entry = profile.entries[0]
    assert entry.risk == [2, 3, 4]
    assert isinstance(entry.risk, list)
    assert all(isinstance(val, int) for val in entry.risk)


def test_application_list_multiple_entries_independent_values():
    config = '''config application list
    edit "Test"
        config entries
            edit 1
                set category 2 6 7
                set action block
            next
            edit 2
                set application 1234 5678
                set risk 3 4
                set action pass
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 2
    entry1 = profile.entries[0]
    entry2 = profile.entries[1]
    assert entry1.category == [2, 6, 7]
    assert entry1.action == "block"
    assert entry1.application == []
    assert entry1.risk == []
    assert entry2.application == [1234, 5678]
    assert entry2.risk == [3, 4]
    assert entry2.action == "pass"
    assert entry2.category == []


def test_application_list_malformed_category_is_captured_defensively():
    from fwmigrate.parsers.fortigate.parser import FortiGateParser
    from fwmigrate.parsers.fortigate.tokenizer import FortiGateTokenizer
    from fwmigrate.extraction.models import ExtractionStatus

    config = '''config application list
    edit "TestProfile"
        config entries
            edit 1
                set category 2 invalid_token 7
            next
        end
    next
end
'''
    parser = FortiGateParser(FortiGateTokenizer(config))
    parsed = parser.parse()

    # 1. Parsing succeeds and valid integers are retained
    assert len(parsed.application_lists) == 1
    profile = parsed.application_lists[0]
    assert len(profile.entries) == 1
    entry = profile.entries[0]
    assert entry.category == [2, 7]

    # 2. Invalid tokens are preserved in extra_settings
    assert "unparsed_category" in entry.extra_settings
    assert entry.extra_settings["unparsed_category"] == ["invalid_token"]

    # 3. Source inventory item records manual review and issue note
    inv_items = [i for i in parser.source_inventory_items if i.source_path == "application list"]
    assert len(inv_items) >= 1
    app_inv = inv_items[0]
    assert app_inv.requires_manual_review is True
    assert any("invalid_token" in note for note in app_inv.notes)

    # 4. Child edit command reflects PARSE_ERROR
    entries_child = next(c for c in app_inv.children if c.name == "entries")
    entry_child = next(c for c in entries_child.children if c.name == "1")
    assert entry_child.requires_manual_review is True
    assert entry_child.status == ExtractionStatus.PARTIALLY_NORMALIZED
    category_cmd = next(c for c in entry_child.commands if c.key == "category")
    assert category_cmd.status == ExtractionStatus.PARSE_ERROR
    assert category_cmd.requires_manual_review is True


def test_application_list_malformed_application_and_risk_defensively():
    from fwmigrate.parsers.fortigate.parser import FortiGateParser
    from fwmigrate.parsers.fortigate.tokenizer import FortiGateTokenizer

    config = '''config application list
    edit "TestProfile"
        config entries
            edit 1
                set application 11414 bad_app 15722
                set risk 3 critical_risk 4
            next
        end
        config filters
            edit 1
                set category 2 unknown_cat 6
                set risk 1 extreme 5
            next
        end
        config overrides
            edit 1
                set application 999 broken 888
                set category 10 bogus
            next
        end
    next
end
'''
    parser = FortiGateParser(FortiGateTokenizer(config))
    parsed = parser.parse()

    profile = parsed.application_lists[0]
    entry = profile.entries[0]
    assert entry.application == [11414, 15722]
    assert entry.application_id == 11414
    assert entry.extra_settings["unparsed_application"] == ["bad_app"]
    assert entry.risk == [3, 4]
    assert entry.extra_settings["unparsed_risk"] == ["critical_risk"]

    filt = profile.filters[0]
    assert filt.category == [2, 6]
    assert filt.extra_settings["unparsed_category"] == ["unknown_cat"]
    assert filt.risk == [1, 5]
    assert filt.extra_settings["unparsed_risk"] == ["extreme"]

    override = profile.overrides[0]
    assert override.application == [999, 888]
    assert override.extra_settings["unparsed_application"] == ["broken"]
    assert override.category == [10]
    assert override.extra_settings["unparsed_category"] == ["bogus"]


def test_application_list_all_invalid_values_does_not_crash():
    from fwmigrate.parsers.fortigate.parser import parse_fortigate_config

    config = '''config application list
    edit "TestProfile"
        config entries
            edit 1
                set category unknown
                set application none
                set risk critical
            next
        end
    next
end
'''
    parsed = parse_fortigate_config(config)
    entry = parsed.application_lists[0].entries[0]
    assert entry.category == []
    assert entry.application == []
    assert entry.application_id is None
    assert entry.risk == []
    assert entry.extra_settings["unparsed_category"] == ["unknown"]
    assert entry.extra_settings["unparsed_application"] == ["none"]
    assert entry.extra_settings["unparsed_risk"] == ["critical"]


def test_application_list_malformed_values_allows_excel_export():
    import io
    from openpyxl import load_workbook
    from fwmigrate.web import _extract_source_config
    from fwmigrate.report.excel_exporter import IRExcelExporter

    config = '''config application list
    edit "block-high-risk"
        set unknown-application-log enable
        config entries
            edit 1
                set category 2 malformed_cat 7
            next
            edit 2
                set application 11414 11767
                set risk 3 4
                set action pass
            next
        end
    next
end
'''
    ir_config, extraction_result = _extract_source_config("fortigate", config)
    assert ir_config is not None
    assert extraction_result is not None

    exporter = IRExcelExporter(ir_config, extraction_result=extraction_result)
    excel_data = exporter.generate()
    assert isinstance(excel_data, bytes)
    assert len(excel_data) > 0

    workbook = load_workbook(io.BytesIO(excel_data))
    assert "Extraction Coverage" in workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    assert "Source Security Profiles" in workbook.sheetnames
