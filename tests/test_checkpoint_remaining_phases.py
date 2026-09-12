import io
import json

from openpyxl import load_workbook

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.parsers.checkpoint.extractor import extract_checkpoint_config
from fwmigrate.parsers.checkpoint.gaia import parse_gaia_configuration
from fwmigrate.report import IRExcelExporter


def _bundle(*responses):
    return json.dumps({"format": "checkpoint-export-v1", "responses": list(responses)})


def test_gaia_typed_interface_families_preserve_ids_peers_and_secret_presence():
    _, interfaces, _, _, inventory, _ = parse_gaia_configuration("\n".join([
        "add interface eth1 alias 10.10.99.1/24",
        "add vxlan id 100 dev eth0 remote 192.0.2.2 dstport 4789",
        "add interface eth0 6in4 55 remote 192.0.2.30 ttl 200",
        "add gre id 5 local 192.0.2.1 remote 192.0.2.2 ttl 64 ip 10.0.0.1 mask 255.255.255.252 peer 10.0.0.2",
        "add pppoe client id 1 interface eth0 user-name user password secret use-peer-dns on",
    ]))
    by_type = {item.interface_type: item for item in interfaces}
    assert by_type["alias"].parent == "eth1"
    assert by_type["vxlan"].source_attributes["vni"] == 100
    assert by_type["6in4"].remote_ip == "192.0.2.30"
    assert by_type["gre"].source_attributes["peer"] == "10.0.0.2"
    assert by_type["pppoe"].has_pppoe_password is True
    assert "secret" not in json.dumps([item.model_dump() for item in inventory])


def test_gaia_routes_and_pbr_are_canonical_but_separate():
    content = _bundle({
        "command": "gaia/show-configuration",
        "data": {"cli_text": "\n".join([
            "set interface eth0 state on",
            "set static-route 192.0.2.0/24 nexthop blackhole",
            "set static-route 198.51.100.0/24 nexthop gateway address 192.0.2.1 monitored-ip 198.51.100.1 on",
            "set pbr table WAN static-route 203.0.113.0/24 nexthop gateway address 192.0.2.1 interface eth0 on",
            "set pbr rule priority 100 match from 10.0.0.0/8 interface eth0 protocol tcp port 443",
            "set pbr rule priority 100 action table WAN",
        ])},
    })
    extraction = extract_checkpoint_config(content)
    assert len(extraction.canonical_ir.routes) == 2
    assert extraction.canonical_ir.routes[0].route_type == "blackhole"
    assert extraction.canonical_ir.routes[0].blackhole is True
    assert extraction.canonical_ir.routes[1].monitoring[0]["address"] == "198.51.100.1"
    assert len(extraction.canonical_ir.pbf_rules) == 1
    assert extraction.canonical_ir.pbf_rules[0].destination_port == "443"
    assert extraction.canonical_ir.pbf_rules[0].routing_table == "WAN"
    assert not any(item.source_type == "gaia-pbr-rule" and item.status == ExtractionStatus.NORMALIZED for item in extraction.inventory_items)


def test_checkpoint_typed_access_service_nat_ip_pool_and_dependency_excel():
    content = _bundle(
        {
            "command": "show-hosts",
            "domain": "D",
            "data": {"objects": [{"uid": "host", "name": "Host", "type": "host", "ipv4-address": "192.0.2.10"}]},
        },
        {
            "command": "show-services-tcp",
            "domain": "D",
            "data": {"objects": [{"uid": "svc", "name": "Svc", "type": "service-tcp", "port": "443"}]},
        },
        {
            "command": "show-gateways-and-servers",
            "domain": "D",
            "data": {"objects": [{"uid": "gw", "name": "GW", "type": "simple-gateway", "nat-settings": {"ip-pool": [{"uid": "pool", "name": "Pool", "address-range": "range", "gateway": "gw"}]}},]},
        },
        {
            "command": "show-address-ranges",
            "domain": "D",
            "data": {"objects": [{"uid": "range", "name": "Range", "type": "address-range", "ipv4-address-first": "198.51.100.10", "ipv4-address-last": "198.51.100.20"}, {"uid": "unrelated", "name": "Unrelated", "type": "address-range"}]},
        },
        {
            "command": "show-access-rulebase",
            "domain": "D", "package": "P", "layer": "L",
            "data": {"rulebase": [{"uid": "rule", "rule-number": 1, "name": "Inline", "source": [{"uid": "host", "name": "Host"}], "destination": ["Any"], "service": [{"uid": "svc", "name": "Svc"}], "action": "Accept", "enabled": True, "content": ["Content"], "inline-layer": {"uid": "child", "name": "Child"}}]},
        },
        {
            "command": "show-nat-rulebase", "domain": "D", "package": "P",
            "data": {"rulebase": [{"uid": "nat", "rule-number": 7, "name": "Service", "original-source": "Any", "original-destination": "Any", "original-service": "Any", "translated-source": "Original", "translated-destination": "Original", "translated-service": {"uid": "svc", "name": "Svc"}, "enabled": True}]},
        },
    )
    extraction = extract_checkpoint_config(content)
    assert extraction.canonical_ir.checkpoint_access_rules[0].inline_layer_reference == "Child"
    assert extraction.canonical_ir.nat_rules[0].type.value == "service"
    assert extraction.canonical_ir.nat_rules[0].source_rule_id == "7"
    assert [pool.name for pool in extraction.canonical_ir.ip_pools] == ["Pool"]
    assert any(dependency.source_path == "access-rules" for dependency in extraction.dependencies)
    workbook = load_workbook(io.BytesIO(IRExcelExporter(extraction.canonical_ir, extraction_result=extraction).generate()), read_only=True)
    assert "Checkpoint Access Rules" in workbook.sheetnames
    assert "Dependency Registry" in workbook.sheetnames
