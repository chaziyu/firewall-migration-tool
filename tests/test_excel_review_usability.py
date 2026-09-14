import io

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook, load_workbook

from fwmigrate.ir.core import IRConfig, IRMetadata
from fwmigrate.report.excel_exporter import IRExcelExporter


def _workbook(source_vendor: str = "fortigate"):
    ir = IRConfig(
        metadata=IRMetadata(
            hostname="review-usability-test",
            source_vendor=source_vendor,
        )
    )
    return load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))


def _summary_navigation(workbook):
    summary = workbook["Summary"]
    header_row = next(
        row
        for row in range(1, summary.max_row + 1)
        if summary.cell(row, 1).value == "Category"
        and summary.cell(row, 2).value == "Sheet"
    )
    headers = {
        summary.cell(header_row, column).value: column
        for column in range(1, summary.max_column + 1)
        if summary.cell(header_row, column).value
    }

    navigation = {}
    for row in range(header_row + 1, summary.max_row + 1):
        sheet_name = summary.cell(row, headers["Sheet"]).value
        if not sheet_name:
            break
        navigation[sheet_name] = {
            header: summary.cell(row, column).value
            for header, column in headers.items()
        }

    return navigation


def _apply_view(sheet_name, headers, row):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append([sheet_name])
    sheet.append(["Inventory"])
    sheet.append(headers)
    sheet.append(row)

    exporter = IRExcelExporter(
        IRConfig(
            metadata=IRMetadata(
                hostname="column-visibility-test",
                source_vendor="fortigate",
            )
        )
    )
    exporter._apply_sheet_view(sheet)
    return sheet


def test_review_required_is_first_human_review_sheet():
    workbook = _workbook()

    assert workbook.sheetnames[0] == "Summary"
    assert workbook.sheetnames[1] == "Review Required"
    assert workbook["Review Required"].sheet_state == "visible"
    assert workbook["Review Required"].freeze_panes == "B4"


def test_zero_record_inventory_sheets_are_retained_but_hidden():
    workbook = _workbook()

    assert "Interfaces" in workbook.sheetnames
    assert workbook["Interfaces"].sheet_state == "hidden"

    # Audit evidence remains directly visible even when there are no findings.
    assert workbook["Warnings"].sheet_state == "visible"
    assert workbook["Unsupported"].sheet_state == "visible"
    assert workbook["Unresolved References"].sheet_state == "visible"
    assert workbook["Extraction Coverage"].sheet_state == "visible"


def test_detail_evidence_is_retained_but_hidden_by_default():
    workbook = _workbook()

    assert "FortiGate Source Configuration" in workbook.sheetnames
    assert workbook["FortiGate Source Configuration"].sheet_state == "hidden"


def test_summary_navigation_includes_hidden_applicable_sheets():
    workbook = _workbook()
    navigation = _summary_navigation(workbook)

    assert navigation["Review Required"]["Visibility"] == "Visible"
    assert navigation["Interfaces"]["Visibility"] == "Hidden"
    assert navigation["Interface Secondary IPs"]["Visibility"] == "Hidden"
    assert navigation["Warnings"]["Visibility"] == "Visible"


def test_hidden_detail_sheets_keep_parent_adjacency():
    workbook = _workbook()

    assert workbook["Interface Secondary IPs"].sheet_state == "hidden"
    assert (
        workbook.sheetnames.index("Interface Secondary IPs")
        == workbook.sheetnames.index("Interfaces") + 1
    )


def test_presentation_hiding_preserves_complete_fortigate_sheet_contract():
    workbook = _workbook()

    assert workbook.sheetnames == list(IRExcelExporter.SHEET_ORDER)
    assert workbook.sheetnames[:2] == ["Summary", "Review Required"]


def test_core_sheet_hides_populated_non_core_columns_but_keeps_review_fields():
    sheet = _apply_view(
        "Addresses",
        (
            "Name",
            "Source UUID",
            "Type",
            "Value",
            "Migration Status",
            "Manual Review",
            "Additional Settings",
            "Description",
        ),
        (
            "server-1",
            "source-uuid",
            "ipmask",
            "10.0.0.10/32",
            "NORMALIZED",
            "No",
            '{"visibility":"still-preserved"}',
            "Application server",
        ),
    )

    assert sheet.column_dimensions["A"].hidden is False
    assert sheet.column_dimensions["B"].hidden is True
    assert sheet.column_dimensions["C"].hidden is False
    assert sheet.column_dimensions["D"].hidden is False
    assert sheet.column_dimensions["E"].hidden is False
    assert sheet.column_dimensions["F"].hidden is False
    assert sheet.column_dimensions["G"].hidden is True
    assert sheet.column_dimensions["H"].hidden is False


def test_non_core_sheet_only_hides_columns_that_are_completely_empty():
    sheet = _apply_view(
        "Certificates",
        ("Name", "Issuer", "Unused Optional Field"),
        ("vpn-cert", "Example CA", None),
    )

    assert sheet.column_dimensions["A"].hidden is False
    assert sheet.column_dimensions["B"].hidden is False
    assert sheet.column_dimensions["C"].hidden is True
