import io
import json

from openpyxl import load_workbook

from fwmigrate.collection import CollectedSource, CollectionStatus, source_collectors
from fwmigrate.web import create_app


SOURCE = "hostname asa\nusername admin password 0 do-not-save\ninterface GigabitEthernet0/1\n nameif outside\n ip address 203.0.113.1 255.255.255.0\n"


def _payload(password="do-not-persist"):
    return {"vendor": "cisco_asa", "connection": {
        "host": "192.0.2.1", "port": 22, "username": "admin", "password": password,
    }}


def test_connection_failure_does_not_return_credentials(monkeypatch):
    client = create_app({"TESTING": True}).test_client()
    monkeypatch.setattr(source_collectors.get("cisco_asa"), "test_connection", lambda options: (_ for _ in ()).throw(RuntimeError(options["password"])))
    response = client.post("/api/collection/test", json=_payload())
    assert response.status_code == 502
    assert b"do-not-persist" not in response.data


def test_collect_failure_does_not_return_credentials(monkeypatch):
    client = create_app({"TESTING": True}).test_client()
    monkeypatch.setattr(source_collectors.get("cisco_asa"), "collect", lambda options: (_ for _ in ()).throw(ValueError(options["password"])))
    response = client.post("/api/collection/collect", json=_payload())
    assert response.status_code == 502
    assert b"do-not-persist" not in response.data


def test_asa_collection_snapshot_preview_and_excel(monkeypatch):
    client = create_app({"TESTING": True}).test_client()
    monkeypatch.setattr(source_collectors.get("cisco_asa"), "collect", lambda options: CollectedSource(
        "cisco_asa", SOURCE, "live-cisco-asa.cfg", "ssh", CollectionStatus.SUCCESS
    ))
    response = client.post("/api/collection/collect", json=_payload())
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["collection"]["status"] == "SUCCESS"
    assert payload["preview_id"]
    assert "do-not-save" not in str(payload)
    assert "do-not-persist" not in str(payload)

    imported = client.post("/api/collection/snapshot/import", data={"file": (io.BytesIO(json.dumps(payload["snapshot"]).encode()), "snapshot.json")})
    assert imported.status_code == 200
    assert imported.get_json()["preview_id"] == payload["preview_id"]
    workbook = client.post("/api/extract/excel", data={"source_vendor": "cisco_asa", "preview_id": payload["preview_id"]})
    assert workbook.status_code == 200
    assert "Interfaces" in load_workbook(io.BytesIO(workbook.data), read_only=True).sheetnames
    assert b"do-not-save" not in workbook.data


def test_collection_validates_vendor_and_port():
    client = create_app({"TESTING": True}).test_client()
    sources = {item["vendor_id"]: item for item in client.get("/api/vendors").get_json()["sources"]}
    assert sources["juniper_srx"]["collection"]["connection_fields"][1]["default"] == 22
    assert sources["cisco_ftd"]["collection"]["connection_fields"][-1]["default"] is True
    assert sources["palo_alto"]["live_collection"] is False
    assert client.post("/api/collection/collect", json={**_payload(), "vendor": "palo_alto"}).status_code == 400
    assert client.post("/api/collection/test", json={**_payload(), "connection": {**_payload()["connection"], "port": 70000}}).status_code == 400


def test_partial_collection_keeps_preview(monkeypatch):
    client = create_app({"TESTING": True}).test_client()
    monkeypatch.setattr(source_collectors.get("cisco_asa"), "collect", lambda options: CollectedSource(
        "cisco_asa", SOURCE, "live-cisco-asa.cfg", "ssh", CollectionStatus.PARTIAL, warnings=("Some source unavailable",)
    ))
    response = client.post("/api/collection/collect", json=_payload())
    assert response.status_code == 200
    assert response.get_json()["collection"]["status"] == "PARTIAL"


def test_native_collectors_reach_preview_and_excel(monkeypatch):
    samples = {
        "juniper_srx": "set system host-name branch\nset interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24",
        "cisco_ftd": json.dumps({"format": "cisco-fmc-rest-export-v1", "source": "fmc-rest-api", "domain": {"id": "d1", "name": "Global"}, "objects": {"hosts": [{"id": "h1", "name": "host-1", "value": "192.0.2.1"}]}, "access_policies": [], "nat_policies": []}),
        "checkpoint": json.dumps({"format": "checkpoint-export-v1", "responses": [{"command": "show-hosts", "data": {"objects": [{"uid": "h1", "name": "host-1", "ipv4-address": "192.0.2.1"}], "from": 1, "to": 1, "total": 1}, "collection_status": "SUCCESS_WITH_DATA"}]}),
    }
    client = create_app({"TESTING": True}).test_client()
    for vendor, source in samples.items():
        monkeypatch.setattr(source_collectors.get(vendor), "collect", lambda options, vendor=vendor, source=source: CollectedSource(vendor, source, f"live-{vendor}.json", "test"))
        payload = _payload()
        payload["vendor"] = vendor
        response = client.post("/api/collection/collect", json=payload)
        assert response.status_code == 200, (vendor, response.get_json())
        data = response.get_json()
        assert data["preview_id"] and data["snapshot"]["vendor_id"] == vendor
        excel = client.post("/api/extract/excel", data={"source_vendor": vendor, "preview_id": data["preview_id"]})
        assert excel.status_code == 200, vendor
        assert load_workbook(io.BytesIO(excel.data), read_only=True).sheetnames
