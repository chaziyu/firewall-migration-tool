import io

from openpyxl import load_workbook

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_optimized import SinglePassIRExcelExporter


def _pool(name: str, start: str, end: str) -> str:
    return f'''\n    edit "{name}"\n        set type cgn-resource-allocation\n        set startip {start}\n        set endip {end}\n        set cgn-block-size 128\n    next'''


def _policy_config(poolname: str) -> str:
    return f'''\nconfig system settings\n    set central-nat disable\nend\nconfig system interface\n    edit "lan"\n    next\n    edit "wan"\n    next\nend\nconfig firewall address\n    edit "SRC"\n        set subnet 10.0.0.0 255.255.255.0\n    next\nend\nconfig firewall policy\n    edit 10\n        set srcintf "lan"\n        set dstintf "wan"\n        set srcaddr "SRC"\n        set dstaddr "all"\n        set service "ALL"\n        set action accept\n        set nat enable\n        set ippool enable\n        set poolname "{poolname}"\n    next\nend\n'''


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[3] if cell.value}


def test_ippool_group_set_append_and_unknown_setting_are_preserved():
    content = f'''\nconfig firewall ippool\n{_pool("CGN-A", "198.51.100.10", "198.51.100.19")}\n{_pool("CGN-B", "198.51.100.20", "198.51.100.29")}\nend\nconfig firewall ippool_grp\n    edit "CGN-GROUP"\n        set member "CGN-A"\n        append member "CGN-B"\n        set future-setting "retained"\n    next\nend\n'''

    parsed = parse_fortigate_config(content)
    assert len(parsed.ip_pool_groups) == 1
    group = parsed.ip_pool_groups[0]
    assert group.name == "CGN-GROUP"
    assert group.member == ["CGN-A", "CGN-B"]
    assert group.extra_settings["future_setting"] == "retained"

    result = extract_fortigate_config(content)
    ir_group = result.canonical_ir.ip_pool_groups[0]
    assert ir_group.members == ["CGN-A", "CGN-B"]
    assert ir_group.unresolved_members == []
    assert ir_group.migration_status == "EXTRACT_ONLY"
    assert ir_group.requires_manual_review is True
    assert ir_group.source_attributes["future_setting"] == "retained"

    inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall ippool_grp"
        and item.name == "CGN-GROUP"
    )
    assert [
        (command.operation, command.key, command.values)
        for command in inventory.commands
        if command.key == "member"
    ] == [
        ("set", "member", ["CGN-A"]),
        ("append", "member", ["CGN-B"]),
    ]

    section = next(
        section
        for section in result.source_sections
        if section.path == "firewall ippool_grp"
    )
    assert section.status == ExtractionStatus.EXTRACT_ONLY
    assert section.object_count_parsed == 1


def test_ippool_group_unset_clears_effective_members_and_keeps_command_evidence():
    content = '''
config firewall ippool_grp
    edit "CGN-GROUP"
        set member "CGN-A" "CGN-B"
        unset member
    next
end
'''

    parsed = parse_fortigate_config(content)
    assert parsed.ip_pool_groups[0].member == []

    result = extract_fortigate_config(content)
    assert result.canonical_ir.ip_pool_groups[0].members == []
    inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall ippool_grp"
    )
    assert any(
        command.operation == "unset" and command.key == "member"
        for command in inventory.commands
    )


def test_ippool_group_missing_member_is_preserved_and_unresolved():
    content = f'''\nconfig firewall ippool\n{_pool("CGN-A", "198.51.100.10", "198.51.100.19")}\nend\nconfig firewall ippool_grp\n    edit "CGN-GROUP"\n        set member "CGN-A" "MISSING"\n    next\nend\n'''

    result = extract_fortigate_config(content)
    group = result.canonical_ir.ip_pool_groups[0]
    assert group.members == ["CGN-A", "MISSING"]
    assert group.unresolved_members == ["MISSING"]
    assert any("MISSING" in reason for reason in group.review_reasons)

    inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall ippool_grp"
    )
    assert inventory.requires_manual_review is True
    assert "unresolved-reference:MISSING" in inventory.notes


def test_policy_pool_group_reference_is_separate_and_target_translation_is_withheld():
    content = f'''\nconfig firewall ippool\n{_pool("CGN-A", "198.51.100.10", "198.51.100.19")}\n{_pool("CGN-B", "198.51.100.20", "198.51.100.29")}\nend\nconfig firewall ippool_grp\n    edit "CGN-GROUP"\n        set member "CGN-A" "CGN-B"\n    next\nend\n{_policy_config("CGN-GROUP")}\n'''

    result = extract_fortigate_config(content)
    rule = next(
        rule
        for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "10"
    )

    assert rule.source_pool_group_references == ["CGN-GROUP"]
    assert rule.source_pool_references == []
    assert rule.translated_sources == []
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.requires_manual_review is True
    assert any("IP-pool group" in reason for reason in rule.review_reasons)
    assert rule.safe_for_target_generation is False


def test_same_context_pool_and_group_name_is_ambiguous_and_fails_closed():
    content = f'''\nconfig firewall ippool\n{_pool("SAME", "198.51.100.10", "198.51.100.19")}\n{_pool("MEMBER", "198.51.100.20", "198.51.100.29")}\nend\nconfig firewall ippool_grp\n    edit "SAME"\n        set member "MEMBER"\n    next\nend\n{_policy_config("SAME")}\n'''

    result = extract_fortigate_config(content)
    rule = next(
        rule
        for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "10"
    )

    assert rule.source_pool_group_references == ["SAME"]
    assert "SAME" in rule.source_pool_references
    assert rule.translated_sources == []
    assert rule.requires_manual_review is True
    assert any("both firewall ippool" in reason for reason in rule.review_reasons)

    policy_inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall policy" and str(item.source_id) == "10"
    )
    assert policy_inventory.requires_manual_review is True
    assert "unresolved-reference:SAME" in policy_inventory.notes


def test_single_pass_excel_exposes_pool_groups_and_nat_group_references():
    content = f'''\nconfig firewall ippool\n{_pool("CGN-A", "198.51.100.10", "198.51.100.19")}\nend\nconfig firewall ippool_grp\n    edit "CGN-GROUP"\n        set member "CGN-A"\n    next\nend\n{_policy_config("CGN-GROUP")}\n'''
    result = extract_fortigate_config(content)
    workbook = load_workbook(
        io.BytesIO(SinglePassIRExcelExporter(result.canonical_ir, result).generate())
    )

    assert "IP Pool Groups" in workbook.sheetnames
    group_sheet = workbook["IP Pool Groups"]
    group_headers = _headers(group_sheet)
    assert group_sheet.cell(4, group_headers["Name"]).value == "CGN-GROUP"
    assert group_sheet.cell(4, group_headers["Members"]).value == "CGN-A"
    assert group_sheet.cell(4, group_headers["Migration Status"]).value == "EXTRACT_ONLY"

    nat_sheet = workbook["NAT Rules"]
    nat_headers = _headers(nat_sheet)
    assert "Source Pool Group References" in nat_headers
    nat_row = next(
        row
        for row in range(4, nat_sheet.max_row + 1)
        if str(nat_sheet.cell(row, nat_headers["Source Policy ID"]).value) == "10"
    )
    assert (
        nat_sheet.cell(
            nat_row,
            nat_headers["Source Pool Group References"],
        ).value
        == "CGN-GROUP"
    )
