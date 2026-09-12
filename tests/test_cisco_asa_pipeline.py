import io

from openpyxl import load_workbook

from fwmigrate.ir.enums import NATType
from fwmigrate.parsers.cisco_asa.parser import CiscoASAParser
from fwmigrate.report.excel_exporter import IRExcelExporter


ASA_REGRESSION = """
hostname ASA-PIPELINE
zone WAN
interface GigabitEthernet0/0
 no shutdown
 nameif outside
 zone-member WAN
 ip address 192.0.2.1 255.255.255.0
 policy-route route-map PBR
interface GigabitEthernet0/0.20
 vlan 30 secondary 40 50,60-62
 nameif inside
 no shutdown
 ip address 10.0.20.1 255.255.255.0
access-list PBR-ACL extended permit tcp any any eq 443
access-list WAN-IN extended permit tcp any any eq 443
access-group WAN-IN in interface outside
route-map PBR permit 10
 match ip address PBR-ACL
 set ip next-hop 192.0.2.254 192.0.2.253
 set interface inside outside
route inside 0.0.0.0 0.0.0.0 10.0.20.254
object-group network-service WEB_SITES
 network-service-member domain example.com service tcp destination eq 443
access-list WEB-NS extended permit ip any object-group-network-service WEB_SITES
access-group WEB-NS in interface outside
ssh 10.0.0.0 255.255.255.0 inside
http 192.0.2.0 255.255.255.0 outside
telnet 198.51.100.10 255.255.255.255 outside
management-access inside
icmp permit any echo outside
icmp deny host 192.0.2.10 outside
object network REAL
 host 10.0.20.10
object network MAPPED
 host 198.51.100.10
object network PUBLIC
 host 203.0.113.10
object network PRIVATE
 host 10.0.20.20
nat (inside,outside) source static REAL MAPPED service tcp 1000 2000
nat (outside,inside) source static any interface destination static PUBLIC PRIVATE service tcp 443 8443
"""


def _rows(workbook, sheet):
    values = list(workbook[sheet].values)
    return values[2], values[3:]


def test_asa_parser_ir_excel_pipeline_covers_phased_regressions():
    parser = CiscoASAParser(ASA_REGRESSION)
    config = parser.parse_raw()
    ir = parser.transform_to_ir()

    subinterface = next(item for item in config.interfaces if item.name.endswith(".20"))
    assert subinterface.interface_suffix_vlan_id == 20
    assert subinterface.vlan_id == 30
    assert subinterface.secondary_vlan_ids == [40, 50]
    assert subinterface.secondary_vlan_ranges == ["60-62"]
    assert config.static_routes[0].effective_administrative_distance == 1

    assert [(zone.name, zone.interfaces) for zone in ir.zones] == [("WAN", ["GigabitEthernet0/0"])]
    assert ir.interfaces[0].zone == "WAN"
    assert ir.policy_route_rules[0].match_acls == ["PBR-ACL"]
    assert ir.policy_route_rules[0].next_hops == ["192.0.2.254", "192.0.2.253"]
    assert ir.policy_route_rules[0].output_interfaces == ["inside", "outside"]
    assert ir.policies[-1].destination == ["WEB_SITES"]
    assert len(ir.local_in_policies) == 6
    assert {rule.source_attributes["origin"] for rule in ir.local_in_policies} == {
        "asa-management-command", "asa-management-access", "asa-icmp-management",
    }

    source_pat, destination_pat = ir.nat_rules
    assert source_pat.type == NATType.SOURCE
    assert [(item.start, item.end) for item in source_pat.original_source_ports] == [(1000, 1000)]
    assert [(item.start, item.end) for item in source_pat.translated_source_ports] == [(2000, 2000)]
    assert source_pat.source_translation_bidirectional is True
    assert [(item.start, item.end) for item in destination_pat.original_destination_ports] == [(443, 443)]
    assert [(item.start, item.end) for item in destination_pat.translated_destination_ports] == [(8443, 8443)]
    assert ir.ip_pools == []
    assert ir.virtual_ips == []

    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()), data_only=False)
    interface_headers, interface_rows = _rows(workbook, "Interfaces")
    zone_headers, zone_rows = _rows(workbook, "Zones")
    local_headers, local_rows = _rows(workbook, "Local-In Policies")
    pbr_headers, pbr_rows = _rows(workbook, "Cisco PBR")
    nat_headers, nat_rows = _rows(workbook, "NAT Rules")
    assert "VLAN ID" in interface_headers and any(row[interface_headers.index("VLAN ID")] == 30 for row in interface_rows)
    assert zone_rows[0][zone_headers.index("Name")] == "WAN"
    assert len(local_rows) == 6
    assert pbr_rows[0][pbr_headers.index("Next Hops")] == "192.0.2.254\n192.0.2.253"
    assert nat_rows[0][nat_headers.index("Original Source Port")] == "1000-1000"
    assert nat_rows[1][nat_headers.index("Original Destination Port")] == "443-443"
    assert workbook["IP Pools"].max_row == 3
    assert workbook["Virtual IPs"].max_row == 3
