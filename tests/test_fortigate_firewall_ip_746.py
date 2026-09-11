import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.parser import FortiGateParser
from fwmigrate.parsers.fortigate.tokenizer import FortiGateTokenizer
from fwmigrate.parsers.fortigate.model import FGIPPool
from fwmigrate.parsers.fortigate.firewall_ip_746 import (
    FORTIOS_746_IPPOOL_FIELDS,
    FORTIOS_746_IPPOOL6_DEFAULTS,
    FORTIOS_746_IPPOOL_TYPES,
    effective_ippool_settings,
    FORTIOS_746_IPV6_EH_DEFAULTS,
    effective_ipv6_eh_filter_settings,
    validate_ipv6_eh_filter_746,
    validate_ippool6_746,
    validate_ippool_746,
)
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.generators.target_helpers import is_generation_safe_object
from fwmigrate.generators.fortigate.cli_generator import FortiGateCLIGenerator
from fwmigrate.generators.fortigate.terraform_generator import FortiGateTerraformGenerator
from fwmigrate.ir.core import IRConfig, IRIPPool, IRMetadata
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_ippool_tracks_explicit_fields_and_unset():
    config = '''
config firewall ippool
    edit "POOL1"
        set startip 203.0.113.10
        set arp-reply disable
        unset arp-reply
    next
end
'''

    pool = FortiGateParser(FortiGateTokenizer(config)).parse().ip_pools[0]

    assert pool.source_explicit_fields == {"startip"}
    assert pool.arp_reply == "enable"
    assert "type" not in pool.source_explicit_fields


def test_ippool6_tracks_explicit_fields():
    config = '''
config firewall ippool6
    edit "POOL6"
        set startip 2001:db8::10
        set nat46 enable
    next
end
'''

    pool = FortiGateParser(FortiGateTokenizer(config)).parse().ip_pools6[0]

    assert pool.source_explicit_fields == {"startip", "nat46"}


def test_fortios_746_contract_constants_are_centralized():
    assert "cgn-resource-allocation" in FORTIOS_746_IPPOOL_TYPES
    assert "startip" in FORTIOS_746_IPPOOL_FIELDS
    assert FORTIOS_746_IPPOOL6_DEFAULTS["startip"] == "::"


def test_ippool_validation_rejects_official_range_and_address_errors():
    pool = FGIPPool(
        name="POOL1",
        block_size=63,
        startip="not-an-ip",
        endip="203.0.113.10",
    )

    reasons = validate_ippool_746(pool)

    assert any("block-size value 63" in reason for reason in reasons)
    assert any("invalid IPv4 address 'not-an-ip'" in reason for reason in reasons)


def test_unknown_ippool_extra_setting_requires_review():
    pool = FGIPPool(
        name="POOL1",
        extra_settings={"future_setting": "enable"},
        source_explicit_fields={"future_setting"},
    )

    reasons = validate_ippool_746(pool)

    assert any("unmodeled source setting 'future_setting'" in reason for reason in reasons)
    assert not any("reference baseline" in reason for reason in reasons)

    versioned_reasons = validate_ippool_746(pool, source_version="7.4.6")
    assert any("not part of the FortiOS 7.4.6 reference baseline" in reason for reason in versioned_reasons)


def test_cgn_resource_allocation_is_retained_but_not_generation_safe():
    result = extract_fortigate_config('''
config firewall ippool
    edit "CGN"
        set type cgn-resource-allocation
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.pool_type == "cgn-resource-allocation"
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True
    assert "hyperscale" in pool.audit_note
    assert is_generation_safe_object(pool) is False


def test_unknown_ippool_semantics_are_preserved_but_not_generation_safe():
    result = extract_fortigate_config('''
config firewall ippool
    edit "UNKNOWN_POOL"
        set startip 203.0.113.10
        set endip 203.0.113.20
        set future-nat-behavior enable
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.source_attributes["future_nat_behavior"] == "enable"
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True
    assert "unmodeled source setting 'future_nat_behavior'" in pool.audit_note
    assert is_generation_safe_object(pool) is False


def test_unknown_ippool_extra_setting_forces_manual_review():
    result = extract_fortigate_config('''
config firewall ippool
    edit "UNKNOWN_POOL"
        set startip 203.0.113.10
        set endip 203.0.113.20
        set future-nat-behavior enable
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.source_attributes["future_nat_behavior"] == "enable"
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True
    assert "future_nat_behavior" in pool.audit_note
    assert is_generation_safe_object(pool) is False
    assert result.generation_safe is False
    assert result.canonical_ir.generation_safe is False
    assert result.migration_complete is False
    assert any(
        "traffic-affecting canonical objects require manual review" in reason
        for reason in result.blocking_reasons
    )


def test_basic_ippool_without_unknown_semantics_remains_generation_safe():
    result = extract_fortigate_config('''
config firewall ippool
    edit "SAFE_POOL"
        set type overload
        set startip 203.0.113.10
        set endip 203.0.113.20
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.migration_status == "NORMALIZED"
    assert pool.requires_manual_review is False
    assert is_generation_safe_object(pool) is True
    assert result.generation_safe is True


def test_fortigate_firewall_ip_documentation_contract():
    documentation = (
        Path(__file__).parents[1]
        / "documentation"
        / "FORTIGATE_CONFIG_EXTRACTION_REFERENCE.md"
    ).read_text(encoding="utf-8").lower()

    assert "firewall ippool6" in documentation
    assert "extract_only" in documentation
    assert "firewall ipv6-eh-filter" in documentation
    assert "fgipv6ehfilter -> extractionresult -> ipv6 eh filter" in documentation
    assert "`enable` means header blocking" in documentation


def test_malformed_numeric_values_remain_in_source_evidence():
    result = extract_fortigate_config('''
config firewall ippool
    edit "POOL1"
        set block-size malformed
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.block_size is None
    assert pool.source_attributes["unparsed_block_size"] == "malformed"
    assert pool.requires_manual_review is True
    assert "invalid source value for block-size" in pool.audit_note


def test_effective_defaults_are_separate_from_explicit_values():
    pool = FGIPPool(
        name="POOL1",
        startip="203.0.113.10",
        source_explicit_fields={"startip"},
    )

    effective = effective_ippool_settings(pool)

    assert effective["startip"] == "203.0.113.10"
    assert effective["type"] == "overload"
    assert effective["arp_reply"] == "enable"
    assert effective["add_nat64_route"] == "enable"
    assert "comments" not in effective


def test_ippool6_validation_keeps_invalid_source_value_visible():
    config = '''
config firewall ippool6
    edit "POOL6"
        set startip not-an-ipv6
        set endip 2001:db8::10
        set nat46 maybe
    next
end
'''
    parsed = FortiGateParser(FortiGateTokenizer(config)).parse()
    pool = parsed.ip_pools6[0]

    reasons = validate_ippool6_746(pool)
    assert pool.startip == "not-an-ipv6"
    assert any("invalid IPv6 address" in reason for reason in reasons)
    assert any("nat46 has invalid" in reason for reason in reasons)


def test_ippool6_is_extract_only_even_when_valid():
    result = extract_fortigate_config('''
config firewall ippool6
    edit "POOL6"
        set startip 2001:db8::10
        set endip 2001:db8::20
    next
end
''')

    pool = result.canonical_ir.ip_pools[0]
    assert pool.migration_status == "EXTRACT_ONLY"
    assert result.source_sections[0].status.value == "EXTRACT_ONLY"


def test_ipv6_eh_filter_is_typed_source_inventory_with_effective_defaults():
    result = extract_fortigate_config('''
config firewall ipv6-eh-filter
    set routing enable
    set hdopt-type 0 255
    set routing-type 255
    set future-setting abc
end
''')

    item = next(
        item for item in result.inventory_items
        if item.source_path == "firewall ipv6-eh-filter"
    )
    attrs = item.source_attributes
    assert attrs["routing"] == "enable"
    assert attrs["hdopt_type"] == [0, 255]
    assert attrs["routing_type"] == 255
    assert attrs["source_explicit_fields"] == [
        "future_setting", "hdopt_type", "routing", "routing_type"
    ]
    assert attrs["source_effective_settings"]["routing"] == "enable"
    assert attrs["source_effective_settings"]["auth"] == "disable"
    assert attrs["additional_settings"]["future_setting"] == "abc"
    assert item.status.value == "EXTRACT_ONLY"
    assert result.source_sections[0].status.value == "EXTRACT_ONLY"
    assert not hasattr(result.canonical_ir, "ipv6_eh_filter")


def test_ipv6_eh_filter_validation_preserves_invalid_values():
    config = '''
config firewall ipv6-eh-filter
    set hdopt-type 0 1 2 3 4 5 6 256
    set routing-type 256
end
'''
    parsed = FortiGateParser(FortiGateTokenizer(config)).parse()
    item = parsed.ipv6_eh_filter
    reasons = validate_ipv6_eh_filter_746(item)

    assert item.hdopt_type[-1] == 256
    assert any("at most seven" in reason for reason in reasons)
    assert any("outside range 0-255" in reason for reason in reasons)
    assert effective_ipv6_eh_filter_settings(item)["routing"] == (
        FORTIOS_746_IPV6_EH_DEFAULTS["routing"]
    )


def test_ip_pool_and_ipv6_eh_excel_sheets_show_provenance():
    result = extract_fortigate_config('''
config firewall ippool
    edit "POOL1"
        set startip 203.0.113.10
        set endip 203.0.113.20
    next
end
config firewall ipv6-eh-filter
    set auth enable
end
''')
    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )

    pools = workbook["IP Pools"]
    pool_headers = {cell.value: cell.column for cell in pools[3]}
    assert pools.cell(4, pool_headers["Source Explicit Fields"]).value == (
        "endip, startip"
    )
    assert "type=overload" in pools.cell(
        4, pool_headers["Effective Source Settings"]
    ).value

    eh = workbook["IPv6 EH Filter"]
    eh_headers = {cell.value: cell.column for cell in eh[3]}
    assert eh.cell(4, eh_headers["Authentication Header Blocking"]).value == "enable"
    assert "routing=enable" in eh.cell(
        4, eh_headers["Effective Source Settings"]
    ).value
    assert eh.cell(4, eh_headers["Extraction Status"]).value == "EXTRACT_ONLY"


@pytest.mark.parametrize(
    ("field", "value", "valid"),
    [
        ("block_size", 64, True), ("block_size", 4096, True),
        ("block_size", 63, False), ("block_size", 4097, False),
        ("cgn_client_ipv6shift", 0, True), ("cgn_client_ipv6shift", 127, True),
        ("cgn_client_ipv6shift", 128, False),
        ("utilization_alarm_clear", 40, True),
        ("utilization_alarm_clear", 100, True),
        ("utilization_alarm_clear", 39, False),
        ("utilization_alarm_clear", 101, False),
        ("utilization_alarm_raise", 50, True),
        ("utilization_alarm_raise", 100, True),
        ("utilization_alarm_raise", 49, False),
        ("utilization_alarm_raise", 101, False),
        ("startport", 5117, True), ("startport", 65533, True),
        ("startport", 5116, False), ("startport", 65534, False),
    ],
)
def test_official_ippool_numeric_boundaries(field, value, valid):
    reasons = validate_ippool_746(FGIPPool(name="POOL1", **{field: value}))
    assert any(field.replace("_", "-") in reason for reason in reasons) is (not valid)


@pytest.mark.parametrize(
    ("field", "value", "valid"),
    [("pba_interim_log", 0, True), ("pba_interim_log", 600, True),
     ("pba_interim_log", 86400, True), ("pba_interim_log", 1, False),
     ("port_per_user", 0, True), ("port_per_user", 32, True),
     ("port_per_user", 60417, True), ("port_per_user", 1, False)],
)
def test_official_ippool_zero_or_range_boundaries(field, value, valid):
    reasons = validate_ippool_746(FGIPPool(name="POOL1", **{field: value}))
    assert any(field.replace("_", "-") in reason for reason in reasons) is (not valid)


def test_complete_746_fixture_preserves_each_section_and_field():
    fixture = Path(__file__).parent / "fixtures" / "fortigate" / "firewall_ip_746.conf"
    result = extract_fortigate_config(fixture.read_text(encoding="utf-8"))
    pool = result.canonical_ir.ip_pools[0]
    source_pool = FortiGateParser(
        FortiGateTokenizer(fixture.read_text(encoding="utf-8"))
    ).parse().ip_pools[0]

    assert source_pool.source_explicit_fields >= FORTIOS_746_IPPOOL_FIELDS
    assert (pool.start_ip, pool.end_ip) == ("203.0.113.10", "203.0.113.20")
    assert (pool.source_start_ip, pool.source_end_ip) == ("10.0.0.10", "10.0.0.20")
    assert (pool.block_size, pool.cgn_block_size) == (4096, 64)
    assert (pool.cgn_port_start, pool.cgn_port_end) == (1024, 65535)
    assert pool.excluded_ips == ["203.0.113.11"]
    assert pool.source_effective_settings["type"] == "overload"

    pool6 = result.canonical_ir.ip_pools[1]
    assert (pool6.start_ip, pool6.end_ip) == ("2001:db8::10", "2001:db8::20")
    assert pool6.migration_status == "EXTRACT_ONLY"

    statuses = {section.path: section.status.value for section in result.source_sections}
    assert statuses["firewall ippool"] == "PARTIALLY_NORMALIZED"
    assert statuses["firewall ippool6"] == "EXTRACT_ONLY"
    assert statuses["firewall ipv6-eh-filter"] == "EXTRACT_ONLY"


def test_ippool_ordering_is_reported_without_repair():
    pool = FGIPPool(
        name="POOL1",
        startip="203.0.113.20",
        endip="203.0.113.10",
        startport=6000,
        endport=5000,
    )
    reasons = validate_ippool_746(pool)
    assert pool.startip == "203.0.113.20"
    assert pool.endip == "203.0.113.10"
    assert any("startip is greater than endip" in reason for reason in reasons)
    assert any("startport is greater than endport" in reason for reason in reasons)


def test_generators_withhold_advanced_cgn_pool():
    pool = IRIPPool(
        name="CGN",
        pool_type="cgn-resource-allocation",
        migration_status="PARTIALLY_NORMALIZED",
        requires_manual_review=True,
    )
    ir = IRConfig(
        metadata=IRMetadata(source_vendor="fortigate"),
        ip_pools=[pool],
    )

    cli = FortiGateCLIGenerator().generate(ir)[0].content
    terraform = next(
        artifact.content
        for artifact in FortiGateTerraformGenerator().generate(ir)
        if artifact.filename == "main.tf"
    )
    assert 'edit "CGN"' not in cli
    assert 'resource "fortios_firewall_ippool"' not in terraform
