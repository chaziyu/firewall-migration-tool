from io import BytesIO
from copy import deepcopy
from pathlib import Path

from fwmigrate.vendors.cisco_ftd.source_report import (
    CiscoFTDSourceReporter,
    extract_cisco_ftd_source,
)
from fwmigrate.vendors.cisco_ftd.derived import build_ftd_derived_views
from fwmigrate.vendors.cisco_ftd.validation import validate_ftd_config
from fwmigrate.vendors.cisco_ftd.model import CiscoFTDManualNATRule, CiscoFTDAutoNATRule
from fwmigrate.vendors.cisco_ftd.derived import FTDReferenceKind


ROOT = Path(__file__).parents[2]


def _fixture(name: str) -> str:
    return (ROOT / "fixtures" / "cisco_ftd" / name).read_text(encoding="utf-8")


def test_fmc_source_plane_keeps_acp_and_nat_native():
    result = extract_cisco_ftd_source(_fixture("fmc_nat_pipeline_conformance.json"))

    assert result.config.source_plane == "fmc-rest-bundle"
    assert all(item.source_plane == "fmc-rest-bundle" for policy in result.config.access_control_policies for item in policy.rules)
    assert result.config.nat_policies
    policy = result.config.nat_policies[0]
    before = policy.manual_rules_before_auto[0]
    auto = policy.auto_rules[0]
    after = policy.manual_rules_after_auto[0]
    assert policy.source_id == "nat-policy"
    assert isinstance(before, CiscoFTDManualNATRule) and isinstance(auto, CiscoFTDAutoNATRule)
    assert [before.section, auto.nat_type, after.section] == ["BEFORE_AUTO", "DYNAMIC", "AFTER_AUTO"]
    assert [before.position, auto.position, after.position] == [3, 1, 4]
    assert before.original_source["id"] == "inside" and before.translated_source["id"] == "public"
    assert before.original_destination["id"] == "public" and before.translated_destination["id"] == "inside"
    assert before.original_source_port["id"] == before.translated_source_port["id"] == "http"
    assert before.original_source_service["id"] == before.translated_source_service["id"] == "http"
    assert before.original_destination_port["port"] == "80" and before.translated_destination_port["port"] == "8080"
    assert before.original_destination_service["id"] == before.translated_destination_service["id"] == "http"
    assert auto.original_network["id"] == "inside" and auto.interface_pat is True
    assert before.raw_extra["translatedDestination"]["id"] == "inside"
    assert before.source_context == "fmc:Global" and before.domain_id == "fmc-domain"
    nat_inventory = [item for item in result.inventory_items if item.source_path.endswith(("nat-policies", "nat-rules"))]
    assert [item.source_type for item in nat_inventory] == [
        "nat-policy", "manual-nat-rule", "auto-nat-rule", "manual-nat-rule"]
    assert nat_inventory[1].source_attributes["policy_id"] == policy.source_id
    assert nat_inventory[1].source_attributes["section"] == "BEFORE_AUTO"
    assert not result.derived.unresolved_references
    assert [(item["rule_kind"], item["section"]) for item in result.derived.nat_relationships] == [
        ("manual", "BEFORE_AUTO"), ("auto", None), ("manual", "AFTER_AUTO")]
    snapshot = deepcopy(policy)
    build_ftd_derived_views(result.config)
    validate_ftd_config(result.config, result.derived)
    assert policy == snapshot


def test_fdm_source_plane_is_not_fmc():
    result = extract_cisco_ftd_source(_fixture("fdm_nat_pipeline_conformance.json"))

    assert result.config.source_plane == "fdm-rest-bundle"
    assert result.config.nat_policies
    assert result.config.source_metadata["source"] == "fdm-rest-api"


def test_cli_evidence_does_not_manufacture_managed_policy_or_nat():
    result = extract_cisco_ftd_source(
        "interface outside\n ip address 203.0.113.1 255.255.255.0\n"
    )

    assert result.config.source_plane == "ftd-text-evidence"
    assert result.config.access_control_policies == []
    assert result.config.nat_policies == []


def test_unresolved_native_reference_is_reported():
    result = extract_cisco_ftd_source(
        '{"source":"fmc-rest-api","objects":{"hosts":[]},'
        '"access_policies":[{"name":"p","rules":[{"name":"r",'
        '"destinationNetworks":{"objects":[{"name":"missing"}]}}]}]}'
    )

    assert result.derived.unresolved_references
    assert result.validation.issues
    assert result.derived.unresolved_references[0].field == "destination_networks"


def test_derived_validation_does_not_mutate_ftd_source_config():
    result = extract_cisco_ftd_source(_fixture("fmc_acp_pipeline_conformance.json"))
    before = deepcopy(result.config)

    derived = build_ftd_derived_views(result.config)
    validate_ftd_config(result.config, derived)

    assert result.config == before


def test_fmc_collection_states_are_preserved_in_derived_views_and_preview():
    import json
    from fwmigrate.vendors.cisco_ftd.web_report import build_ftd_preview

    bundle = {
        "format": "cisco-fmc-rest-export-v1", "source": "fmc-rest-api",
        "objects": {"hosts": [{"id": "h1", "name": "host"}], "networks": [], "ranges": [], "securityzones": []},
        "collection": {"status": "PARTIAL", "parts": [
            {"name": "hosts", "status": "SUCCESS", "complete": True, "count": 1},
            {"name": "networks", "status": "EMPTY", "complete": True, "count": 0},
            {"name": "ranges", "status": "PARTIAL", "complete": False, "count": 1},
            {"name": "securityzones", "status": "FAILED", "complete": False, "count": 0},
            {"name": "access_rules/policy-a", "status": "SUCCESS", "complete": True, "count": 1},
            {"name": "access_rules/policy-b", "status": "FAILED", "complete": False, "count": 0},
        ]},
    }
    result = extract_cisco_ftd_source(json.dumps(bundle))
    completeness = result.derived.source_plane_completeness
    assert [completeness[f"collection:{name}"] for name in ("hosts", "networks", "ranges", "securityzones")] == [
        "present", "known-empty", "partial", "failed"]
    assert completeness["security_zones"] == "failed"
    assert completeness["acp"] == "partial"
    assert result.source_sections[0].status.value == "PARTIAL"
    assert result.inventory_items[0].status.value == "EXTRACTED"
    preview = build_ftd_preview(result)
    assert preview["summary"]["security_zones"] == 0
    assert preview["source_plane_completeness"]["collection:networks"] == "known-empty"


def test_offline_fmc_bundle_without_collection_metadata_remains_unknown():
    result = extract_cisco_ftd_source(
        '{"source":"fmc-rest-api","objects":{"networks":[]}}'
    )
    assert result.config.collection_metadata.provided is False
    assert result.source_sections[0].status.value == "UNKNOWN"
    assert result.derived.source_plane_completeness["networks"] == "unknown"
    assert result.derived.source_plane_completeness["network_objects"] == "unknown"


def test_fmc_acp_preserves_native_rule_semantics():
    result = extract_cisco_ftd_source(_fixture("fmc_acp_pipeline_conformance.json"))
    policy = result.config.access_control_policies[0]
    rule, disabled = policy.rules

    assert (rule.policy_id, rule.policy_name, rule.source_id) == ("policy-1", "Corporate ACP", "rule-1")
    assert (rule.enabled, disabled.enabled) == (True, False)
    assert (rule.position, disabled.position) == (4, 5)
    assert [item.collection_order for item in policy.rules] == [1, 2]
    assert (rule.section, rule.category, rule.action, rule.comments) == (
        "Mandatory", "Web Access", "ALLOW", "Allow managed web access"
    )
    assert rule.source_zones[0].source_id == "zone-inside"
    assert rule.destination_zones[0].source_id == "zone-outside"
    assert rule.source_networks[0].source_id == "net-inside"
    assert rule.destination_networks[0].source_id == "net-web"
    assert rule.source_ports[0].source_id == "port-source"
    assert rule.destination_ports[0].source_id == "port-destination"
    assert rule.users[0].source_id == "user-alice"
    assert rule.applications[0].source_id == "app-web"
    assert rule.urls["urlCategoriesWithReputation"][0]["reputation"] == "BENIGN_SITES"
    assert rule.url_categories[0].source_id == "url-business"
    assert rule.time_range.source_id == "time-business"
    assert rule.intrusion_policy.source_id == "ips-balanced"
    assert rule.variable_set.source_id == "variables-default"
    assert rule.file_policy.source_id == "file-inspect"
    assert (rule.log_begin, rule.log_end) == (False, True)
    assert rule.logging["sendEventsToFMC"] is False
    assert rule.source_plane == "fmc-rest-bundle"
    assert rule.domain_id == "domain-1"
    assert "applications" in rule.explicit_fields
    assert rule.raw_extra["metadata"]["ruleIndex"] == 4
    assert result.derived.acp_relationships[0]["policy_id"] == "policy-1"
    assert result.derived.acp_relationships[0]["destination_port_refs"] == rule.destination_ports
    assert not result.derived.unresolved_references


def test_ftd_reporter_exports_source_plane_workbook():
    result = CiscoFTDSourceReporter().analyze_source(
        _fixture("fmc_acp_pipeline_conformance.json")
    )
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)

    assert output.getvalue()[:2] == b"PK"
    from openpyxl import load_workbook

    output.seek(0)
    sheet = load_workbook(output, read_only=True)["ACP Rules"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))))
    assert "port-source" in row["Source Ports"]
    assert "port-destination" in row["Destination Ports"]
    assert row["Enabled"] is True and row["Log Begin"] is False


def test_fmc_device_interfaces_keep_ownership_and_resolve_zone_membership():
    import json

    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "d0", "name": "Global"},
        "objects": {
            "securityzones": [{"id": "z1", "name": "inside", "interfaces": [{"name": "GigabitEthernet0/0"}]}],
            "interfacegroups": [{"id": "g1", "name": "trusted-group"}],
        },
        "devices": [{"id": "d1", "name": "ftd-1", "resources": {"ftd_interfaces": [
            {"id": "if0", "name": "GigabitEthernet0/0", "interfaceType": "PhysicalInterface", "address": "192.0.2.1"}
        ]}}],
    }))

    interface = result.config.source_interfaces[0]
    assert (interface.name, interface.source_id, interface.device_id, interface.source_attributes["device_name"]) == (
        "GigabitEthernet0/0", "if0", "d1", "ftd-1")
    assert interface.domain_id == "d0" and interface.source_plane == "fmc-rest-bundle"
    assert interface.raw_extra["interfaceType"] == "PhysicalInterface"
    assert [zone.name for zone in result.config.security_zones] == ["inside"]
    assert [group.name for group in result.config.interface_groups] == ["trusted-group"]
    assert not result.config.device_interfaces
    assert not result.derived.unresolved_references
    inventory_interface = next(item for item in result.inventory_items if item.source_path.endswith("/interfaces"))
    assert inventory_interface.source_id == "if0" and inventory_interface.source_context == "fmc:Global"
    assert inventory_interface.source_attributes["device_id"] == "d1"
    assert inventory_interface.source_attributes["device_name"] == "ftd-1"
    assert inventory_interface.source_attributes["domain_id"] == "d0"
    assert inventory_interface.source_attributes["source_plane"] == "fmc-rest-bundle"

    before = deepcopy(result.config)
    validate_ftd_config(result.config, build_ftd_derived_views(result.config))
    assert result.config == before


def test_zone_name_reference_stays_ambiguous_across_devices():
    import json

    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "objects": {"securityzones": [{"id": "z1", "name": "inside", "interfaces": ["GigabitEthernet0/0"]}]},
        "devices": [{"id": device_id, "name": device_id, "resources": {"ftd_interfaces": [
            {"id": f"{device_id}-if0", "name": "GigabitEthernet0/0"}
        ]}} for device_id in ("d1", "d2")],
    }))

    issue = next(item for item in result.derived.unresolved_references if item.field == "interfaces")
    assert issue.status == "AMBIGUOUS"
    assert any(item.category == "ambiguous-reference" for item in result.validation.issues)


def _reference_result(objects, reference, domain_id="domain-a"):
    import json
    return extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api", "domain": {"id": domain_id, "name": domain_id},
        "objects": objects,
        "access_policies": [{"id": "p", "name": "policy", "rules": [{"id": "r", "name": "rule",
            "destinationNetworks": {"objects": [reference]}}]}],
    }))


def test_ftd_reference_resolution_is_typed_and_domain_scoped():
    result = _reference_result({
        "hosts": [{"id": "host-1", "name": "shared"}],
        "protocolportobjects": [{"id": "service-1", "name": "shared"}],
    }, {"name": "shared"})

    resolved = result.derived.resolved_references
    assert [(item["kind"], item["target_id"]) for item in resolved if item["field"] == "destination_networks"] == [
        (FTDReferenceKind.NETWORK_ADDRESS.value, "host-1")]
    assert not [item for item in result.derived.unresolved_references if item.field == "destination_networks"]

    wrong_kind = _reference_result({"protocolportobjects": [{"id": "service-1", "name": "shared"}]}, {"name": "shared"})
    issue = next(item for item in wrong_kind.derived.unresolved_references if item.field == "destination_networks")
    assert issue.status == "WRONG_KIND" and issue.reason == "wrong-kind"
    assert issue.expected_kinds == (FTDReferenceKind.NETWORK_ADDRESS.value, FTDReferenceKind.NETWORK_GROUP.value)
    assert any(item.category == "wrong-kind-reference" for item in wrong_kind.validation.issues)


def test_ftd_reference_id_precedes_name_and_domains_do_not_collide():
    result = _reference_result({"hosts": [
        {"id": "host-a", "name": "server"}, {"id": "host-b", "name": "server"},
    ]}, {"id": "host-b", "name": "server"})
    match = next(item for item in result.derived.resolved_references if item["field"] == "destination_networks")
    assert match["target_id"] == "host-b"

    result = _reference_result({"hosts": [{"id": "host-a", "name": "server"}]}, {"name": "server"})
    other_domain = _reference_result({"hosts": [{"id": "host-b", "name": "server"}]}, {"name": "server"}, "domain-b")
    result.config.network_addresses.extend(other_domain.config.network_addresses)
    derived = build_ftd_derived_views(result.config)
    match = next(item for item in derived.resolved_references if item["field"] == "destination_networks")
    assert match["target_id"] == "host-a"


def test_ftd_same_kind_duplicate_name_is_ambiguous():
    result = _reference_result({"hosts": [
        {"id": "host-a", "name": "shared"}, {"id": "host-b", "name": "shared"},
    ]}, {"name": "shared"})
    issue = next(item for item in result.derived.unresolved_references if item.field == "destination_networks")
    assert issue.status == "AMBIGUOUS" and issue.reason == "ambiguous"


def test_fmc_address_service_and_group_objects_project_typed_source_fields():
    import json

    bundle = {
        "format": "cisco-fmc-rest-export-v1", "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "objects": {
            "networkaddresses": [
                {"id": "host", "name": "Host", "type": "Host", "value": "192.0.2.1", "description": "host desc", "overrides": [{"device": {"id": "d1"}, "value": "192.0.2.2"}]},
                {"id": "network", "name": "Network", "type": "Network", "value": "192.0.2.0/24"},
                {"id": "range", "name": "Range", "type": "Range", "value": "192.0.2.1-192.0.2.9"},
                {"id": "fqdn", "name": "FQDN", "type": "FQDN", "value": "example.test", "lookupType": "IPV4_ONLY", "addressFamily": "ipv4"},
                {"id": "missing", "name": "Missing", "type": "Host"},
            ],
            "networkgroups": [{"id": "ng", "name": "Network Group", "description": "group desc",
                "objects": [{"id": "host", "name": "Host", "type": "Host"}],
                "literals": [{"value": "198.51.100.1", "type": "Host"}], "overrides": [{"deviceId": "d1"}]}],
            "protocolportobjects": [
                {"id": "http", "name": "HTTP", "protocol": "6", "port": "80", "description": "web"},
                {"id": "range-port", "name": "Range", "protocol": "6", "port": "80", "endPort": "90"},
                {"id": "udp", "name": "UDP", "protocol": "17"},
                {"id": "protocol", "name": "Protocol", "protocol": "47"},
                {"id": "icmp", "name": "ICMP", "protocol": "1", "icmpType": "8"},
                {"id": "icmp-code", "name": "ICMP Code", "protocol": "1", "icmpType": "3", "icmpCode": "1", "overrides": [{"deviceId": "d1"}]},
            ],
            "portobjectgroups": [{"id": "pg", "name": "Port Group", "description": "ports",
                "objects": [{"id": "http", "name": "HTTP", "type": "ProtocolPortObject"}], "overrides": [{"deviceId": "d1"}]}],
        },
        "access_policies": [{"id": "p", "name": "Policy", "rules": [{"id": "r", "name": "Rule",
            "destinationNetworks": {"objects": [{"id": "host", "name": "Host", "type": "Host"}]},
            "destinationPorts": {"objects": [{"id": "http", "name": "HTTP", "type": "ProtocolPortObject"}]}}]}],
    }
    result = extract_cisco_ftd_source(json.dumps(bundle))
    config = result.config
    addresses = {item.name: item for item in config.network_addresses}
    assert [(addresses[key].address_type, addresses[key].value) for key in ("Host", "Network", "Range", "FQDN")] == [
        ("Host", "192.0.2.1"), ("Network", "192.0.2.0/24"), ("Range", "192.0.2.1-192.0.2.9"), ("FQDN", "example.test")]
    host = addresses["Host"]
    assert (host.source_id, host.description, host.domain_id) == ("host", "host desc", "domain-1")
    assert host.override_metadata == {"overrides": [{"device": {"id": "d1"}, "value": "192.0.2.2"}]}
    assert {"address_type", "value", "description", "override_metadata"} <= set(host.explicit_fields)
    assert host.raw_extra["overrides"] == host.override_metadata["overrides"]
    assert addresses["FQDN"].fqdn_lookup_type == "IPV4_ONLY" and addresses["FQDN"].address_family == "ipv4"
    assert addresses["Missing"].value is None and addresses["Missing"].description is None
    assert "value" not in addresses["Missing"].explicit_fields

    services = {item.name: item for item in config.protocol_port_objects}
    assert (services["HTTP"].protocol, services["HTTP"].port) == ("6", "80")
    assert (services["Range"].port, services["Range"].end_port) == ("80", "90")
    assert services["UDP"].protocol == "17" and services["UDP"].port is None
    assert services["Protocol"].protocol == "47" and services["Protocol"].port is None
    assert (services["ICMP"].protocol, services["ICMP"].icmp_type, services["ICMP"].icmp_code) == ("1", "8", None)
    assert (services["ICMP Code"].icmp_type, services["ICMP Code"].icmp_code) == ("3", "1")
    assert services["HTTP"].description == "web"
    assert "port" in services["HTTP"].explicit_fields and "icmp_type" not in services["HTTP"].explicit_fields
    assert services["ICMP Code"].override_metadata == {"overrides": [{"deviceId": "d1"}]}

    network_group = config.network_groups[0]
    assert network_group.description == "group desc" and network_group.members[0].source_id == "host"
    assert network_group.literal_members[0].value == "198.51.100.1"
    assert network_group.override_metadata == {"overrides": [{"deviceId": "d1"}]}
    port_group = config.port_object_groups[0]
    assert port_group.description == "ports" and port_group.members[0].source_id == "http"
    assert port_group.override_metadata == {"overrides": [{"deviceId": "d1"}]}

    before = deepcopy(config)
    derived = build_ftd_derived_views(config)
    validate_ftd_config(config, derived)
    assert config == before
    assert not [item for item in derived.unresolved_references if item.field in {"destination_networks", "destination_ports"}]

    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)
    from openpyxl import load_workbook
    sheet = load_workbook(output, read_only=True)["Services"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    http = next(row for row in rows if row["Name"] == "HTTP")
    icmp = next(row for row in rows if row["Name"] == "ICMP Code")
    assert (http["Protocol"], http["Port"]) == ("6", "80")
    assert (icmp["Protocol"], icmp["Port"], icmp["ICMP Type"], icmp["ICMP Code"]) == ("1", None, "3", "1")
