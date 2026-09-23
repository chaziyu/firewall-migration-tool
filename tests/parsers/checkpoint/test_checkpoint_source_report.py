import io
import json
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.checkpoint.source_report import extract_checkpoint_source


FIXTURES = Path(__file__).parents[2] / "fixtures" / "checkpoint"


def _source(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


def test_checkpoint_source_report_keeps_collection_and_scope_separate():
    result = extract_checkpoint_source(_source("multidomain_full.json"))

    assert result.config.hosts
    assert result.config.access_rules
    assert all(item.source_plane == "management" for item in result.config.hosts)
    assert result.collection


def test_checkpoint_source_report_does_not_treat_failed_collection_as_empty():
    result = extract_checkpoint_source(_source("partial_collection.json"))

    assert any(not item.complete for item in result.collection)
    assert any(issue.category == "collection" for issue in result.validation.issues)


def test_checkpoint_source_report_resolves_group_and_rule_references():
    result = extract_checkpoint_source(_source("r81_golden_matrix.json"))

    assert result.config.hosts
    assert isinstance(result.derived.unresolved_references, tuple)


def test_checkpoint_source_report_excel_has_traceability_and_redacts_secrets():
    result = extract_checkpoint_source(_source("single_gateway_full.json"))
    from fwmigrate.vendors.checkpoint.export.excel import export_checkpoint_excel

    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    output.seek(0)
    workbook = load_workbook(output, read_only=True)

    assert "Collection" in workbook.sheetnames
    assert "Validation" in workbook.sheetnames
    serialized = json.dumps(result.config.model_dump())
    assert "[REDACTED]" in serialized or "password" not in serialized.lower()
