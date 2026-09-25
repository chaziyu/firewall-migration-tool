from copy import deepcopy

from fwmigrate.vendors.juniper_srx.parser import JuniperSRXParser
from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel
from fwmigrate.vendors.juniper_srx.derived import build_juniper_derived_views
from fwmigrate.vendors.juniper_srx.source_report import extract_juniper_source
from fwmigrate.vendors.juniper_srx.web_report import build_juniper_preview
from fwmigrate.vendors.juniper_srx.transforms.inheritance import build_inheritance_view


def _view(source):
    parser = JuniperSRXParser(source)
    parser.extract_source()
    commands_before = deepcopy(parser.commands)
    result = build_inheritance_view(parser.commands)
    assert parser.commands == commands_before
    return result


def test_local_override_and_group_order_keep_candidates_and_provenance():
    view = _view("""set groups first interfaces ge-0/0/0 description first
set groups second interfaces ge-0/0/0 description second
set apply-groups [ first second ]
set interfaces ge-0/0/0 description local
""")
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    assert [item["source_group"] for item in inherited] == ["second", "first"]
    assert all(item["status"] == "SHADOWED" for item in inherited)
    assert all(item["group_chain"] and item["source_path"] for item in inherited)
    assert any(item["origin"] == "local" and item["status"] == "EFFECTIVE" and
               item["value"] == "local" for item in view["effective_statements"])


def test_nested_groups_and_apply_groups_except_retain_excluded_candidates():
    view = _view("""set groups inner system host-name inherited
set groups outer apply-groups inner
set groups excluded system host-name excluded
set apply-groups outer
set apply-groups-except excluded
""")
    assert any(item["source_group"] == "inner" for item in view["effective_statements"])
    assert any(item.get("status") == "EXCLUDED" for item in view["candidates"])


def test_inactive_group_statement_and_inactive_application_are_preserved():
    view = _view("""set groups G system host-name host
deactivate groups G system host-name
set apply-groups G
deactivate apply-groups G
""")
    assert any(item["status"] == "INACTIVE" for item in view["candidates"])
    assert "GROUP_INACTIVE" not in {item["status"] for item in view["issues"]}


def test_missing_and_cyclic_groups_are_reported_without_synthetic_targets():
    missing = _view("set apply-groups absent")
    cycle = _view("""set groups A apply-groups B
set groups B apply-groups A
set apply-groups A
""")
    assert "GROUP_NOT_FOUND" in {item["status"] for item in missing["issues"]}
    assert "GROUP_CYCLE" in {item["status"] for item in cycle["issues"]}


def test_wildcard_group_renders_only_to_existing_interfaces():
    view = _view("""set groups G interfaces <*> description inherited
set interfaces ge-0/0/0 unit 0
set apply-groups G
""")
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    assert [item["target_path"] for item in inherited] == [
        ("interfaces", "ge-0/0/0", "description", "inherited")
    ]


def test_same_group_name_isolated_between_logical_systems():
    view = _view("""set groups G system host-name root
set apply-groups G
set logical-systems L1 groups G system host-name one
set logical-systems L1 apply-groups G
set logical-systems L2 groups G system host-name two
set logical-systems L2 apply-groups G
""")
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    assert {item["context"] for item in inherited} == {"root", "logical-system L1", "logical-system L2"}
    assert {item["value"] for item in inherited} == {"root", "one", "two"}


def test_group_list_precedence_and_hierarchy_incompatibility_are_visible():
    view = _view("""set groups first system host-name first
set groups second system host-name second
set apply-groups [ first second ]
set groups G logical-systems L1 system host-name isolated
set logical-systems L2 apply-groups G
""")
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    root = [item for item in inherited if item["context"] == "root"]
    assert [(item["source_group"], item["status"]) for item in root] == [
        ("second", "SHADOWED"), ("first", "EFFECTIVE")
    ]
    assert "GROUP_HIERARCHY_INCOMPATIBLE" in {item["status"] for item in view["issues"]}


def test_tenant_scoped_group_and_recursion_limit():
    lines = ["set groups G system host-name tenant",
             "set tenants T1 groups G system host-name scoped",
             "set tenants T1 apply-groups G"]
    lines.extend(f"set groups G{i} apply-groups G{i + 1}" for i in range(65))
    lines.extend(["set groups G65 system host-name too-deep", "set apply-groups G0"])
    view = _view("\n".join(lines))
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    assert any(item["context"] == "tenant T1" and item["value"] == "scoped" for item in inherited)
    assert "GROUP_RECURSION_DEPTH_EXCEEDED" in {item["status"] for item in view["issues"]}


def test_group_candidate_evidence_redacts_secret_values():
    secret = "DO_NOT_EXPORT_GROUP_SECRET"
    view = _view(f"set groups G security ike policy I pre-shared-key ascii-text {secret}\nset apply-groups G")
    assert secret not in repr(view)


def test_validation_and_excel_keep_group_failures_and_excluded_candidates():
    from io import BytesIO

    from openpyxl import load_workbook

    result = extract_juniper_source("\n".join([
        "set groups G system host-name excluded-value",
        "set apply-groups-except G",
        "set apply-groups missing",
    ]))
    assert any(issue.category == "inheritance" for issue in result.validation.issues)
    preview = build_juniper_preview(result)
    assert {"policy_relationships", "nat_usage", "vpn_graph", "secure_connect_graph", "apbr_graph", "inheritance_view"} <= set(preview["relationships"])
    output = BytesIO()
    export_juniper_excel(result, output)
    output.seek(0)
    rows = list(load_workbook(output, read_only=True)["Inheritance"].values)
    assert any(row[0] == "candidate" and str(row[3]).endswith("EXCLUDED") for row in rows[1:])
    assert any(row[0] == "issue" and str(row[3]) == "GROUP_NOT_FOUND" for row in rows[1:])
    assert not any(issue.code == "GROUP_EXCLUDED" for issue in result.validation.issues)


def test_effective_inherited_address_resolves_without_mutating_source():
    parser = JuniperSRXParser("\n".join([
        "set groups G security address-book global address INHERITED 192.0.2.5/32",
        "set apply-groups G",
        "set security policies global policy P match source-address INHERITED",
    ]))
    config = parser.extract_source()
    before = deepcopy(config)
    derived = build_juniper_derived_views(config, parser.commands)
    dependency = next(item for item in derived.dependencies if item.reference == "INHERITED")
    assert dependency.result == "RESOLVED"
    assert "global" not in config.get_context().address_books
    assert config == before


def test_excluded_or_deactivated_inherited_address_does_not_resolve():
    for extra, expected_status in (("set apply-groups-except G", "EXCLUDED"),
                                   ("deactivate groups G security address-book global address INHERITED", "INACTIVE")):
        parser = JuniperSRXParser("\n".join([
            "set groups G security address-book global address INHERITED 192.0.2.5/32",
            "set apply-groups G", extra,
            "set security policies global policy P match source-address INHERITED",
        ]))
        config = parser.extract_source()
        derived = build_juniper_derived_views(config, parser.commands)
        dependency = next(item for item in derived.dependencies if item.reference == "INHERITED")
        assert dependency.result == "UNRESOLVED"
        candidates = derived.inheritance_view["candidates"]
        statements = derived.inheritance_view["effective_statements"]
        source = candidates
        assert any(item["status"] == expected_status for item in source), (
            expected_status, [(item["status"], item.get("target_path")) for item in source])


def test_effective_lookup_stays_with_its_logical_system():
    parser = JuniperSRXParser("\n".join([
        "set logical-systems LS1 groups G security address-book global address A 192.0.2.5/32",
        "set logical-systems LS1 apply-groups G",
        "set logical-systems LS1 security policies global policy P match source-address A",
        "set logical-systems LS2 security policies global policy P match source-address A",
    ]))
    config = parser.extract_source()
    derived = build_juniper_derived_views(config, parser.commands)
    results = {item.source_context: item.result for item in derived.dependencies
               if item.source_field == "source-address" and item.reference == "A"}
    assert results == {"logical-system LS1": "RESOLVED", "logical-system LS2": "UNRESOLVED"}, (
        results, [(item["context"], item["target_path"], item["status"], item["source_group"])
                  for item in derived.inheritance_view["effective_statements"]])


def test_hierarchical_inactive_parent_emits_one_deactivation_and_keeps_children():
    parser = JuniperSRXParser("""interfaces {
    inactive: ge-0/0/0 {
        description WAN;
        mtu 1500;
    }
    ge-0/0/1 {
        description LAN;
    }
}
""")
    config = parser.extract_source()
    assert parser.source_format == "junos_hierarchical"
    deactivations = [cmd for cmd in parser.commands if cmd.operation.value == "deactivate"]
    assert [cmd.tokens[1:] for cmd in deactivations] == [["interfaces", "ge-0/0/0"]]
    assert config.get_context().interfaces["ge-0/0/0"].description == "WAN"
    assert config.get_context().interfaces["ge-0/0/1"].description == "LAN"
    view = build_inheritance_view(parser.commands)
    assert any(item["status"] == "INACTIVE" and item["target_path"][:2] == ("interfaces", "ge-0/0/0")
               for item in view["effective_statements"])
    assert any(item["status"] == "EFFECTIVE" and item["target_path"][:2] == ("interfaces", "ge-0/0/1")
               for item in view["effective_statements"])


def test_hierarchical_nested_and_leaf_inactive_paths_are_exact():
    parser = JuniperSRXParser("""interfaces {
    ge-0/0/0 {
        unit 0 {
            inactive: family inet {
                address 192.0.2.1/24;
            }
        }
    }
    ge-0/0/1 {
        unit 0 {
            family inet {
                inactive: address 198.51.100.1/24;
            }
        }
    }
    ge-0/0/2 {
        unit 0 {
            family inet {
                address 203.0.113.1/24;
            }
        }
    }
}
""")
    parser.extract_source()
    directives = [cmd.tokens[1:] for cmd in parser.commands if cmd.operation.value == "deactivate"]
    assert directives == [
        ["interfaces", "ge-0/0/0", "unit", "0", "family", "inet"],
        ["interfaces", "ge-0/0/1", "unit", "0", "family", "inet", "address", "198.51.100.1/24"],
    ]


def test_activation_lookup_distinguishes_inactive_hierarchy_from_inactive_leaf():
    from fwmigrate.vendors.juniper_srx.transforms.effective_lookup import EffectiveJunosLookup

    parent = JuniperSRXParser("""set security ike proposal P1 encryption-algorithm aes-128-cbc
deactivate security ike proposal P1
""")
    parent.extract_source()
    parent_view = build_inheritance_view(parent.commands)
    parent_lookup = EffectiveJunosLookup((*parent_view["effective_statements"], *(
        {"context": item["context"], "target_path": item["path"], "origin": "activation"}
        for item in parent_view["inactive_hierarchies"])))
    assert parent_lookup.hierarchy_is_inactive("root", ("security", "ike", "proposal", "P1"))
    assert not parent_lookup.explicit_object_is_effective("root", ("security", "ike", "proposal", "P1"))

    leaf = JuniperSRXParser("""set security ike proposal P1 encryption-algorithm aes-128-cbc
deactivate security ike proposal P1 lifetime-seconds
""")
    config = leaf.extract_source()
    leaf_view = build_inheritance_view(leaf.commands)
    leaf_lookup = EffectiveJunosLookup((*leaf_view["effective_statements"], *(
        {"context": item["context"], "target_path": item["path"], "origin": "activation"}
        for item in leaf_view["inactive_hierarchies"])))
    from fwmigrate.vendors.juniper_srx.resolver import JuniperReferenceResolver
    assert JuniperReferenceResolver(config.get_context(), leaf_lookup).resolve_ike_proposal("P1") is not None
    assert not leaf_lookup.hierarchy_is_inactive(
        "root", ("security", "ike", "proposal", "P1"))


def test_activation_order_and_scalar_vs_member_inheritance():
    view = _view("""set groups G security ike proposal P1 encryption-algorithm aes-128-cbc
set groups G security ike policy IP1 proposals P1
set apply-groups G
set security ike proposal P1 encryption-algorithm aes-256-cbc
set security ike policy IP1 proposals P2
deactivate security ike proposal P1
activate security ike proposal P1
""")
    statements = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    assert next(item for item in statements if item["target_path"][-2:] == ("encryption-algorithm", "aes-128-cbc"))["status"] == "SHADOWED"
    inherited_member = next(item for item in statements if "proposals" in item["target_path"])
    assert inherited_member["value"] == "P1" and inherited_member["status"] == "EFFECTIVE"
    assert any(item["target_path"] == ("security", "ike", "proposal", "P1", "encryption-algorithm", "aes-256-cbc")
               and item["status"] == "EFFECTIVE" for item in view["effective_statements"])


def test_all_core_resolvers_reject_deactivated_source_hierarchies():
    source = """set security address-book global address A 192.0.2.1/32
set applications application APP protocol tcp
set schedulers scheduler SCH daily 12:00-13:00
set routing-instances RI instance-type virtual-router
set security nat source pool SP address 192.0.2.10/32
set security nat destination pool DP address 192.0.2.20/32
set firewall family inet filter FF term T then accept
set interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24
set security zones security-zone Z interfaces ge-0/0/0.0
set security ike proposal IP encryption-algorithm aes-128-cbc
set security ike policy IK proposals IP
set security ike gateway IG ike-policy IK
set security ipsec proposal SP protocol esp
set security ipsec policy SK proposals SP
set security ipsec vpn VPN ike gateway IG
set security ipsec vpn VPN ike ipsec-policy SK
deactivate security address-book global address A
deactivate applications application APP
deactivate schedulers scheduler SCH
deactivate routing-instances RI
deactivate security nat source pool SP
deactivate security nat destination pool DP
deactivate firewall family inet filter FF
deactivate interfaces ge-0/0/0
deactivate security zones security-zone Z
deactivate security ike proposal IP
deactivate security ike policy IK
deactivate security ike gateway IG
deactivate security ipsec policy SK
deactivate security ipsec vpn VPN
"""
    parser = JuniperSRXParser(source)
    config = parser.extract_source()
    view = build_inheritance_view(parser.commands)
    from fwmigrate.vendors.juniper_srx.transforms.effective_lookup import EffectiveJunosLookup
    activation = ({"context": item["context"], "target_path": item["path"], "origin": "activation"}
                  for item in view["inactive_hierarchies"])
    lookup = EffectiveJunosLookup((*view["effective_statements"], *activation))
    from fwmigrate.vendors.juniper_srx.resolver import JuniperReferenceResolver
    resolver = JuniperReferenceResolver(config.get_context(), lookup)
    assert resolver._resolve_in_book("global", "A").is_unresolved
    assert resolver.resolve_application("APP")[2] is None
    assert resolver.resolve_scheduler("SCH") is None
    assert resolver.resolve_routing_instance("RI") is None
    assert resolver.resolve_nat_pool("SP", "source") is None
    assert resolver.resolve_nat_pool("DP", "destination") is None
    assert resolver.resolve_firewall_filter("FF", "inet") is None
    assert resolver.resolve_interface("ge-0/0/0") is None
    assert resolver.resolve_interface("ge-0/0/0.0") is None
    assert resolver.resolve_zone("Z") is None
    assert resolver.resolve_ike_proposal("IP") is None
    assert resolver.resolve_ike_policy("IK") is None
    assert resolver.resolve_ike_gateway("IG") is None
    assert resolver.resolve_ipsec_policy("SK") is None
    assert resolver.resolve_ipsec_vpn("VPN") is None
    assert "A" in config.get_context().address_books["global"].addresses
    assert "APP" in config.get_context().applications
    assert "SP" in config.get_context().nat.source_pools
    assert "DP" in config.get_context().nat.destination_pools
    assert "FF" in config.get_context().firewall_filters


def test_deactivated_explicit_scheduler_stays_in_source_and_is_unresolved_everywhere():
    result = extract_juniper_source("""set schedulers scheduler SCH daily 12:00-13:00
deactivate schedulers scheduler SCH
set security policies global policy P scheduler-name SCH
""")
    assert "SCH" in result.config.get_context().schedulers
    assert result.config.activation_directives
    dependency = next(item for item in result.derived.dependencies if item.reference == "SCH")
    edge = next(item for item in result.derived.policy_relationships[0]["edges"]
                if item["source_field"] == "scheduler")
    assert dependency.result == "UNRESOLVED"
    assert edge["resolved"] is False


def test_alternative_effective_paths_are_checked_independently():
    from fwmigrate.vendors.juniper_srx.transforms.effective_lookup import EffectiveJunosLookup

    active = ("routing-instances", "RI")
    inactive = ("routing-options", "instance-import", "RI")
    lookup = EffectiveJunosLookup((
        {"context": "root", "target_path": active, "origin": "local", "status": "EFFECTIVE"},
        {"context": "root", "target_path": inactive, "origin": "activation"},
    ))
    assert lookup.contains_effective_path("root", active, inactive)
    assert lookup.path_is_effective("root", active)
    assert not lookup.path_is_effective("root", inactive)


def test_vpn_compound_scalars_shadow_inherited_values():
    view = _view("""set groups G security ipsec vpn V bind-interface st0.0
set groups G security ipsec vpn V establish-tunnels immediately
set groups G security ipsec vpn V ike gateway GW-A
set groups G security ipsec vpn V ike ipsec-policy IP-A
set apply-groups G
set security ipsec vpn V bind-interface st0.1
set security ipsec vpn V establish-tunnels on-traffic
set security ipsec vpn V ike gateway GW-B
set security ipsec vpn V ike ipsec-policy IP-B
""")
    inherited = [item for item in view["effective_statements"] if item["origin"] == "inherited-group"]
    for field, inherited_value in (("bind-interface", "st0.0"), ("establish-tunnels", "immediately"),
                                   ("gateway", "GW-A"), ("ipsec-policy", "IP-A")):
        item = next(item for item in inherited if item["target_path"][-2:] == (field, inherited_value))
        assert item["status"] == "SHADOWED"
    for value in ("st0.1", "on-traffic", "GW-B", "IP-B"):
        assert any(item["value"] == value and item["origin"] == "local"
                   and item["status"] == "EFFECTIVE" for item in view["effective_statements"])


def test_deactivated_address_set_member_is_kept_in_source_but_not_expanded():
    result = extract_juniper_source("""set security address-book global address A 192.0.2.1/32
set security address-book global address-set S address A
deactivate security address-book global address-set S address A
set security policies global policy P match source-address S
""")
    address_set = result.config.get_context().address_books["global"].address_sets["S"]
    assert [member.name for member in address_set.members] == ["A"]
    from fwmigrate.vendors.juniper_srx.transforms.effective_lookup import EffectiveJunosLookup
    view = result.derived.inheritance_view
    activation = ({"context": item["context"], "target_path": item["path"], "origin": "activation"}
                  for item in view["inactive_hierarchies"])
    lookup = EffectiveJunosLookup((*view["effective_statements"], *activation))
    from fwmigrate.vendors.juniper_srx.resolver import JuniperReferenceResolver
    resolved = JuniperReferenceResolver(result.config.get_context(), lookup)._resolve_in_book("global", "S")
    assert resolved.resolved_members == []


def test_nested_address_set_deactivation_reactivation_and_logical_system_scope():
    from fwmigrate.vendors.juniper_srx.transforms.effective_lookup import EffectiveJunosLookup
    from fwmigrate.vendors.juniper_srx.resolver import JuniperReferenceResolver

    cases = (
        ("""set security address-book global address A 192.0.2.1/32
set security address-book global address-set CHILD address A
set security address-book global address-set PARENT address-set CHILD
deactivate security address-book global address-set PARENT address-set CHILD
""", "root", []),
        ("""set security address-book global address A 192.0.2.1/32
set security address-book global address-set S address A
deactivate security address-book global address-set S address A
activate security address-book global address-set S address A
""", "root", ["A"]),
        ("""set logical-systems LS1 security address-book global address A 192.0.2.1/32
set logical-systems LS1 security address-book global address-set S address A
deactivate logical-systems LS1 security address-book global address-set S address A
""", "logical-system LS1", []),
    )
    for source, scope, expected in cases:
        result = extract_juniper_source(source)
        view = result.derived.inheritance_view
        activation = ({"context": item["context"], "target_path": item["path"], "origin": "activation"}
                      for item in view["inactive_hierarchies"])
        lookup = EffectiveJunosLookup((*view["effective_statements"], *activation))
        if scope == "root":
            context = result.config.get_context()
            set_name = "PARENT" if "PARENT" in context.address_books["global"].address_sets else "S"
        else:
            _, name = scope.split(" ", 1)
            context = result.config.get_context(name, context_type="logical-system")
            set_name = "S"
        resolved = JuniperReferenceResolver(context, lookup)._resolve_in_book("global", set_name)
        assert resolved.resolved_members == expected


def test_deactivated_policy_reference_is_not_an_effective_unresolved_dependency():
    result = extract_juniper_source("""set schedulers scheduler SCH daily 12:00-13:00
set security policies global policy P scheduler-name SCH
deactivate security policies global policy P scheduler-name
""")
    dependency = next(item for item in result.derived.dependencies if item.reference == "SCH")
    edge = next(item for item in result.derived.policy_relationships[0]["edges"]
                if item["source_field"] == "scheduler")
    assert dependency.result == "INACTIVE_SOURCE"
    assert edge["source_effective"] is False and edge["resolved"] is True
    from fwmigrate.vendors.juniper_srx.validation import validate_juniper_config
    assert not any(issue.code == "UNRESOLVED_REFERENCE"
                   for issue in validate_juniper_config(result.config, result.derived).issues)
    from io import BytesIO
    from openpyxl import load_workbook
    from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel
    output = BytesIO()
    export_juniper_excel(result, output)
    sheet = load_workbook(BytesIO(output.getvalue()), read_only=True)["Policy Reference Relationships"]
    rows = list(sheet.values)
    assert next(row for row in rows[1:] if row[3] == "scheduler")[9] == "INACTIVE_SOURCE"
