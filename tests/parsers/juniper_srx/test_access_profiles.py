import json
from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.juniper_srx.source_report import extract_juniper_source


def test_access_firewall_users_are_separate_from_admin_and_secrets_are_redacted():
    secret = "DO_NOT_EXPORT_THIS_PASSWORD"
    other_secrets = ("DO_NOT_EXPORT_LDAP", "DO_NOT_EXPORT_RADIUS", "DO_NOT_EXPORT_PSK")
    result = extract_juniper_source("\n".join([
        "set system login user alice class super-user",
        f"set system login user alice authentication encrypted-password {secret}",
        f"set access profile AUTH client alice firewall-user password {secret}",
        "set access profile AUTH client alice client-group STAFF",
        "set access profile AUTH client bob client-group STAFF",
        f"set access profile AUTH ldap-options password {other_secrets[0]}",
        f"set access profile AUTH radius-server 192.0.2.10 secret {other_secrets[1]}",
        f"set security ike policy IKE pre-shared-key ascii-text {other_secrets[2]}",
    ]))
    context = next(iter(result.config.iter_contexts()))
    assert "alice" in result.config.admin_users
    assert not hasattr(result.config, "local_users")
    assert context.access_profiles["AUTH"].clients["alice"].firewall_user.password_configured is True
    assert context.access_profiles["AUTH"].clients["alice"].client_groups == ["STAFF"]
    assert context.access_profiles["AUTH"].clients["bob"].client_groups == ["STAFF"]
    outputs = [result.config.model_dump_json()]
    from fwmigrate.vendors.juniper_srx.web_report import build_juniper_preview
    outputs.append(json.dumps(build_juniper_preview(result)))
    workbook = BytesIO()
    from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel
    export_juniper_excel(result, workbook)
    workbook.seek(0)
    sheets = load_workbook(workbook, read_only=True, data_only=True)
    outputs.extend(str(cell.value) for sheet in sheets for row in sheet.iter_rows() for cell in row)
    assert all(secret not in output for output in outputs)
    assert "radius-secret-value" not in " ".join(outputs)
    assert all(value not in " ".join(outputs) for value in other_secrets)
