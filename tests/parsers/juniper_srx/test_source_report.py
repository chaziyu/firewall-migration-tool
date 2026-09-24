from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.juniper_srx import JuniperSRXParser
from fwmigrate.vendors.juniper_srx.source_report import extract_juniper_source
from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel


def test_source_extraction_preserves_contexts_groups_and_activation():
    content = """
set system host-name edge
set groups base system time-zone UTC
set apply-groups base
set logical-systems LS1 interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24
deactivate logical-systems LS1 interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24
activate logical-systems LS1 interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24
"""
    parser = JuniperSRXParser(content)
    config = parser.extract_source()

    assert "logical-system:LS1" in config.contexts
    assert "base" in config.configuration_groups
    assert config.hostname == "edge"
    assert parser.source_format == "junos_display_set"
    assert parser.commands


def test_source_keeps_group_and_inactive_values_out_of_effective_source_state():
    result = extract_juniper_source("""set groups base interfaces ge-0/0/0 description inherited
set apply-groups base
set apply-groups-except future-group
set interfaces ge-0/0/0 description local
set system host-name branch-fw
deactivate system host-name
""")
    config = result.config
    interface = config.get_context().interfaces["ge-0/0/0"]
    assert interface.description == "local"
    assert config.hostname == "branch-fw"
    assert config.applied_group_exceptions["root"] == ["future-group"]
    assert [(item.operation, item.hierarchy_path) for item in config.activation_directives] == [
        ("deactivate", ("system", "host-name"))
    ]
    assert any(item["path"][-2:] == ("description", "inherited")
               for item in result.derived.inheritance["effective_commands"])


def test_same_named_logical_system_tenant_and_groups_remain_separate():
    parser = JuniperSRXParser("""set logical-systems SAME system host-name logical-fw
set tenants SAME security-profile tenant-profile
set groups shared system time-zone UTC
set logical-systems SAME groups shared system host-name grouped-fw
set tenants SAME groups shared security-profile grouped-profile
""")
    config = parser.extract_source()
    logical = config.get_context("SAME", "logical-system")
    tenant = config.get_context("SAME", "tenant")
    assert logical is not tenant
    assert logical.source_attributes["unsupported_system"]
    assert tenant.security_profile == "tenant-profile"
    assert config.configuration_groups["shared"].context_type == "root"
    assert config.configuration_groups["logical-system:SAME:shared"].context_name == "SAME"
    assert config.configuration_groups["tenant:SAME:shared"].context_type == "tenant"


def test_inactive_hierarchical_statement_keeps_value_and_directive():
    result = extract_juniper_source("""system {
    host-name branch-fw;
    inactive: time-zone UTC;
}
""")
    assert result.config.hostname == "branch-fw"
    assert result.config.time_zone == "UTC"
    assert any(item.operation == "deactivate" and item.hierarchy_path == ("system", "time-zone", "UTC")
               for item in result.config.activation_directives)


def test_activation_reenables_grouped_hierarchy_in_derived_view():
    result = extract_juniper_source("""set groups base interfaces ge-0/0/0 description inherited
set apply-groups base
deactivate interfaces ge-0/0/0
activate interfaces ge-0/0/0
""")
    inherited = result.derived.inheritance["effective_commands"]
    assert inherited and inherited[0]["active"] is True
    assert [item.operation for item in result.config.activation_directives] == ["deactivate", "activate"]


def test_source_report_and_excel_keep_junos_contexts_and_redact_secrets():
    result = extract_juniper_source(
        "set system host-name edge\nset security ike gateway gw pre-shared-key super-secret\n"
    )
    assert any(item.source_path == "system" for item in result.inventory_items)
    output = BytesIO()
    export_juniper_excel(result, output)
    output.seek(0)
    workbook = load_workbook(output, read_only=True)
    text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    assert "super-secret" not in text
