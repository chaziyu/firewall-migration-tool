from fwmigrate.vendors.juniper_srx.derived import build_juniper_derived_views
from fwmigrate.vendors.juniper_srx.parser import JuniperSRXParser


def test_policy_views_keep_zone_pair_and_global_order_separate():
    config = JuniperSRXParser("\n".join([
        "set security zones security-zone vpn interfaces st0.1",
        "set security policies from-zone vpn to-zone untrust policy P1 match source-identity alice",
        "set security policies from-zone vpn to-zone untrust policy P1 then permit",
        "set security policies from-zone vpn to-zone untrust policy P2 then deny",
        "set security policies global policy G1 then permit",
        "set security policies global policy G2 then deny",
    ])).extract_source()
    derived = build_juniper_derived_views(config)
    view = derived.policy_relationships[0]
    pair = view["zone_policy_sets"][0]
    assert (pair["from_zone"], pair["to_zone"]) == ("vpn", "untrust")
    assert [(item["name"], item["order"]) for item in pair["policies"]] == [("P1", 0), ("P2", 1)]
    assert [(item["name"], item["order"]) for item in view["global_policies"]] == [("G1", 0), ("G2", 1)]
    assert pair["policies"][0]["source_identities"] == ("alice",)
    assert not [edge for edge in view["edges"] if edge["relationship"] == "EXPLICIT_VPN_REFERENCE"]


def test_policy_view_resolves_named_addresses_applications_and_scheduler():
    config = JuniperSRXParser("\n".join([
        "set security address-book global address SRC 10.0.0.1/32",
        "set applications application WEB protocol tcp",
        "set schedulers scheduler S daily 12:00-13:00",
        "set security policies from-zone trust to-zone untrust policy P match source-address SRC",
        "set security policies from-zone trust to-zone untrust policy P match application WEB",
        "set security policies from-zone trust to-zone untrust policy P scheduler-name S",
    ])).extract_source()
    view = build_juniper_derived_views(config).policy_relationships[0]
    assert all(edge["resolved"] for edge in view["edges"])


def test_excel_policy_projection_keeps_zone_and_global_orders_separate():
    from io import BytesIO

    from openpyxl import load_workbook

    from fwmigrate.vendors.juniper_srx.source_report import extract_juniper_source
    from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel

    result = extract_juniper_source("\n".join([
        "set security policies from-zone a to-zone b policy P1 then permit",
        "set security policies from-zone a to-zone b policy P2 then deny",
        "set security policies from-zone c to-zone d policy P3 then permit",
        "set security policies global policy G1 then permit",
        "set security policies global policy G2 then deny",
    ]))
    output = BytesIO()
    export_juniper_excel(result, output)
    output.seek(0)
    rows = list(load_workbook(output, read_only=True)["Policies"].values)
    assert rows[0] == ("Context", "Policy Scope", "From Zone", "To Zone", "Name", "Order")
    assert [row[1:] for row in rows[1:]] == [
        ("zone", "a", "b", "P1", 0), ("zone", "a", "b", "P2", 1),
        ("zone", "c", "d", "P3", 0), ("global", None, None, "G1", 0), ("global", None, None, "G2", 1),
    ]


def test_global_policy_zone_matches_survive_derived_view_and_excel():
    from io import BytesIO
    from openpyxl import load_workbook
    from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel
    from fwmigrate.vendors.juniper_srx.source_report import extract_juniper_source

    result = extract_juniper_source("""set security policies from-zone trust to-zone untrust policy P1 then permit
set security policies global policy G1 match from-zone [ trust dmz ]
set security policies global policy G1 match to-zone untrust
""")
    rows = result.derived.policy_relationships[0]
    assert rows["zone_policy_sets"][0]["policies"][0]["from_zone"] == "trust"
    global_policy = rows["global_policies"][0]
    assert global_policy["policy_scope"] == "global"
    assert global_policy["from_zone"] == ("trust", "dmz")
    assert global_policy["to_zone"] == ("untrust",)
    assert global_policy["order"] == 0

    output = BytesIO()
    export_juniper_excel(result, output)
    output.seek(0)
    workbook_rows = list(load_workbook(output, read_only=True)["Policy Relationships"].values)
    global_row = next(row for row in workbook_rows[1:] if row[4] == "G1")
    assert global_row[2:4] == ("('trust', 'dmz')", "('untrust',)")


def test_inherited_policy_profiles_resolve_only_when_effective():
    for modifier, expected in ((None, "RESOLVED"),
                               ("set apply-groups-except G", "UNRESOLVED"),
                               ("deactivate groups G security idp idp-policy IPS", "UNRESOLVED")):
        lines = ["set groups G security idp idp-policy IPS rulebase-ips rule R match signature 1",
                 "set apply-groups G",
                 "set security policies global policy P then permit application-services idp-policy IPS"]
        if modifier:
            lines.append(modifier)
        parser = JuniperSRXParser("\n".join(lines))
        config = parser.extract_source()
        before = config.model_copy(deep=True)
        derived = build_juniper_derived_views(config, parser.commands)
        dependency = next(item for item in derived.dependencies if item.source_field == "idp-policy")
        edge = next(item for item in derived.policy_relationships[0]["edges"]
                    if item["source_field"] == "idp-policy")
        assert dependency.result == expected
        assert edge["resolved"] == (expected == "RESOLVED")
        assert config.get_context().idp_policies == {}
        assert config == before
