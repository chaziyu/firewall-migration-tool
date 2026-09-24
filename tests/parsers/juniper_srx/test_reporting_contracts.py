from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.vendors.juniper_srx.source_report import _account, extract_juniper_source
from fwmigrate.vendors.juniper_srx.tokenizer import JunosCommand, JunosOperation
from fwmigrate.vendors.juniper_srx.validation import validate_juniper_config
from fwmigrate.vendors.juniper_srx.web_report import build_juniper_preview
from fwmigrate.vendors.juniper_srx.export.excel import export_juniper_excel


def test_source_accounting_preserves_source_only_and_does_not_call_commands_objects():
    commands = [
        JunosCommand(operation=JunosOperation.SET, tokens=["set", "security", "native"],
                     raw_sanitized="set security native", line_number=1,
                     extraction_status=ExtractionStatus.SOURCE_ONLY),
    ]
    sections, inventory, _ = _account(commands)
    assert sections[0].status == ExtractionStatus.SOURCE_ONLY
    assert sections[0].object_count_source is None
    assert sections[0].source_commands == ["set security native"]
    assert len(inventory[0].commands) == 1


def test_structured_validation_preview_and_excel_preserve_boundaries():
    result = extract_juniper_source("""set security ike gateway GW ike-policy MISSING
set security ipsec vpn VPN ike gateway GW
""")
    source_before = deepcopy(result.config.model_dump(mode="python"))
    derived_before = deepcopy(result.derived)
    issue = next(issue for issue in result.validation.issues if issue.reference == "MISSING")
    assert issue.code == "UNRESOLVED_REFERENCE"
    assert issue.source_path == "security ike gateway"
    assert issue.field == "ike-policy"
    assert issue.expected_type == "ike-policy"
    validate_juniper_config(result.config, result.derived)

    preview = build_juniper_preview(result)
    assert "extraction_coverage" in preview
    assert "review_required" in preview
    assert preview["validation_summary"]["errors"] >= 1

    output = BytesIO()
    export_juniper_excel(result, output)
    output.seek(0)
    workbook = load_workbook(output, read_only=True)
    assert {"Review Required", "Extraction Coverage", "Unresolved References", "Additional Settings"} <= set(workbook.sheetnames)
    references = workbook["Unresolved References"]
    assert any(row[5].value == "MISSING" for row in references.iter_rows(min_row=2))
    assert result.config.model_dump(mode="python") == source_before
    assert result.derived == derived_before
