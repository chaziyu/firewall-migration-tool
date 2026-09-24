from io import BytesIO
from copy import deepcopy

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.export.excel_schema import SHEET_ORDER
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview
from fwmigrate.vendors.cisco_asa.derived import build_asa_derived_views
from fwmigrate.vendors.cisco_asa.validation import validate_asa_config


def test_source_pipeline_preserves_acl_nat_order_and_uses_source_model():
    content = """hostname asa
interface GigabitEthernet0/1
 nameif outside
 ip address 203.0.113.1 255.255.255.0
object network WEB
 host 10.0.0.10
object-group network SERVERS
 network-object object WEB
access-list OUT extended permit tcp object-group SERVERS any eq 443
access-list OUT extended deny ip any any
access-group OUT in interface outside
nat (inside,outside) source static WEB interface
nat (inside,outside) after-auto 99 source dynamic WEB interface
"""

    result = extract_cisco_asa_source(content)

    assert result.config.hostname == "asa"
    assert [rule.source_order for rule in result.config.access_rules] == [9, 10]
    assert [rule.source_order for rule in result.config.nat_rules] == [12, 13]
    assert result.derived.object_group_memberships["__global__:SERVERS"] == ("WEB",)
    assert result.derived.interface_topology.interfaces[0].nameif == "outside"
    assert not hasattr(result, "canonical_ir")


def test_source_pipeline_resolves_context_scoped_references_without_mutating_config():
    content = """changeto context customer-a
object network WEB
 host 10.0.0.10
object-group network SERVERS
 network-object object WEB
access-list OUT extended permit ip object-group SERVERS any
changeto system
object network WEB
 host 192.0.2.10
"""

    result = extract_cisco_asa_source(content)

    assert [item.source_context for item in result.config.network_objects] == ["customer-a", None]
    assert result.config.network_groups[0].review_reasons == []
    assert all(issue.resolved for issue in result.derived.relationship_issues)


def test_derived_and_validation_are_read_only():
    result = extract_cisco_asa_source(
        "object network WEB\n host 10.0.0.10\n"
        "object-group network SERVERS\n network-object object WEB\n"
    )
    before_derived = deepcopy(result.config)

    derived = build_asa_derived_views(result.config)

    assert result.config == before_derived

    before_validation = deepcopy(result.config)
    validate_asa_config(result.config, derived)

    assert result.config == before_validation


def test_source_preview_and_excel_contain_no_ir_or_secret_text():
    result = extract_cisco_asa_source(
        "username admin password 0 super-secret\n"
        "object network WEB\n host 10.0.0.10\n"
    )

    preview = str(build_asa_preview(result))
    assert "super-secret" not in preview

    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook.sheetnames == list(SHEET_ORDER)
    assert all("super-secret" not in str(row) for sheet in workbook for row in sheet.iter_rows(values_only=True))
