from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def _analyze(xml: str):
    return PaloAltoSourceReporter().analyze_source(xml)


def _sheet(analysis, name: str):
    output = BytesIO()
    PaloAltoSourceReporter().export_excel(analysis, output)
    return load_workbook(BytesIO(output.getvalue()), data_only=True)[name]


def test_nat_translation_branches_keep_source_type_and_report_values():
    analysis = _analyze("""<config><shared><rulebase><nat><rules>
      <entry name='dynamic'><nat-type>nat64</nat-type><to-interface>ethernet1/1</to-interface><source-translation><dynamic-ip><translated-address><member>pool-a</member></translated-address><fallback><interface-address><interface>ethernet1/2</interface></interface-address></fallback></dynamic-ip></source-translation></entry>
      <entry name='persistent'><source-translation><persistent-dynamic-ip-and-port><translated-address><member>pool-b</member></translated-address></persistent-dynamic-ip-and-port></source-translation></entry>
      <entry name='static'><source-translation><static-ip><translated-address>203.0.113.7</translated-address></static-ip></source-translation></entry>
    </rules></nat></rulebase></shared></config>""")
    dynamic, persistent, static = analysis.config.nat_rules
    assert (dynamic.nat_type, dynamic.to_interface) == ("nat64", "ethernet1/1")
    assert (dynamic.source_translation.translation_type, persistent.source_translation.translation_type, static.source_translation.translation_type) == ("dynamic-ip", "persistent-dynamic-ip-and-port", "static-ip")
    assert dynamic.source_translation.fallback["interface-address"]["interface"] == "ethernet1/2"
    assert not any(issue.message == "NAT rule has no explicit translation branch" for issue in analysis.validation.issues)
    sheet = _sheet(analysis, "NAT Rules")
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Source Translation Type"]).value == "dynamic-ip"
    assert sheet.cell(5, headers["Derived Source Translation Mode"]).value == "persistent-dynamic-ip-and-port"


def test_virtual_and_logical_routes_reach_validation_and_excel():
    analysis = _analyze("""<config><devices><entry name='fw'><network>
      <virtual-router><entry name='vr1'><routing-table><ip><static-route><entry name='r1'><destination>10.0.0.0/8</destination><nexthop><next-vr>vr2</next-vr></nexthop><admin-dist>10</admin-dist><bfd><profile>b1</profile></bfd><path-monitor><enable>yes</enable><monitor-destinations><entry name='target'><destination>192.0.2.1</destination></entry></monitor-destinations></path-monitor></entry></static-route></ip></routing-table></entry></virtual-router>
      <logical-router><entry name='lr1'><vrf><entry name='v1'><routing-table><ip><static-route><entry name='r2'><destination>bad-route</destination><nexthop><next-lr>lr2</next-lr></nexthop></entry></static-route></ip></routing-table></entry></vrf></entry></logical-router>
    </network></entry></devices></config>""")
    vr_route, lr_route = analysis.config.static_routes
    assert (vr_route.nexthop_type, vr_route.nexthop, vr_route.admin_distance, vr_route.bfd_profile) == ("next-vr", "vr2", "10", "b1")
    assert vr_route.path_monitor.targets[0].destination == "192.0.2.1"
    assert (lr_route.nexthop_type, lr_route.nexthop) == ("next-lr", "lr2")
    assert any(issue.domain == "route" and "bad-route" in issue.message for issue in analysis.validation.issues)
    sheet = _sheet(analysis, "Logical Router Routes")
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Route Name"]).value == "r2"
    assert sheet.cell(4, headers["Next Hop Type"]).value == "next-lr"


def test_nested_ike_and_sdwan_hierarchies_are_typed_without_secret_values():
    secret = "dont-export-this-psk"
    analysis = _analyze(f"""<config><shared>
      <sdwan-interface-profile><entry name='link'><link-tag>WAN</link-tag></entry></sdwan-interface-profile>
      <profiles><sdwan-path-quality><entry name='quality'><metric><latency><threshold>30</threshold></latency><pkt-loss><sensitivity>high</sensitivity></pkt-loss></metric></entry></sdwan-path-quality></profiles>
      <rulebase><sdwan><rules><entry name='steer'><action><traffic-distribution-profile>dist</traffic-distribution-profile></action></entry></rules></sdwan></rulebase>
      <network><ike><gateway><entry name='gw'><protocol><version>ikev2</version><ikev2><ike-crypto-profile>crypto</ike-crypto-profile></ikev2></protocol><authentication><pre-shared-key><key>{secret}</key></pre-shared-key></authentication></entry></gateway></ike></network>
    </shared></config>""")
    config = analysis.config
    assert config.sdwan_interface_profiles[0].link_tag == "WAN"
    assert (config.sdwan_path_quality_profiles[0].latency_threshold, config.sdwan_path_quality_profiles[0].packet_loss_sensitivity) == ("30", "high")
    assert config.sdwan_rules[0].traffic_distribution_profile == "dist"
    gateway = config.ike_gateways[0]
    assert (gateway.ike_version, gateway.ikev2_crypto_profile, gateway.pre_shared_key_configured) == ("ikev2", "crypto", True)
    assert secret not in str(config.model_dump())
    assert secret not in str(_sheet(analysis, "IKE Gateways").values)


def test_inventory_and_workbook_show_source_only_records_without_invalid_sheet_names():
    analysis = _analyze("<config><shared><unknown-section><entry name='unsupported'><foo>bar</foo></entry></unknown-section></shared></config>")
    unsupported = _sheet(analysis, "Unsupported")
    headers = {cell.value: cell.column for cell in unsupported[3]}
    assert unsupported.cell(4, headers["Status"]).value == "SOURCE_ONLY"
    output = BytesIO()
    PaloAltoSourceReporter().export_excel(analysis, output)
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert "SD-WAN Distribution Links" in workbook.sheetnames
