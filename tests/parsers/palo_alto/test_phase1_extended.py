from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.native import build_derived_views
from fwmigrate.vendors.palo_alto.source_builder import build_panos_config
from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


SOURCE = """<config><shared>
<profiles><vulnerability><entry name='v'><rules><entry name='r'><threat-name>t</threat-name></entry></rules></entry></vulnerability></profiles>
<admin-role><entry name='role'><permissions><entry><channel>web</channel><permission-path>x</permission-path></entry></permissions></entry></admin-role>
<administrators><entry name='admin'><custom-admin-role>role</custom-admin-role><password>secret</password></entry></administrators>
<network><ike><gateway><entry name='gw'><ikev1-crypto-profile>ikec</ikev1-crypto-profile><pre-shared-key>secret</pre-shared-key></entry></gateway>
<crypto-profiles><ike-crypto-profiles><entry name='ikec'><encryption><member>aes</member></encryption></entry></ike-crypto-profiles><ipsec-crypto-profiles><entry name='ipsec-c'><protocol>esp</protocol></entry></ipsec-crypto-profiles></crypto-profiles></ike>
<ipsec>
<entry name='t'><auto-key><ike-gateway><member>gw</member></ike-gateway></auto-key><ipsec-crypto-profile>ipsec-c</ipsec-crypto-profile><manual-key><key>secret</key></manual-key></entry></ipsec></network>
</shared></config>"""


def test_phase1_extracts_typed_references_and_secret_presence_only():
    config = build_panos_config(SOURCE)
    derived = build_derived_views(config)
    assert not config.unknown_paths
    assert config.administrators[0].password_configured is True
    assert config.ike_gateways[0].pre_shared_key_configured is True
    assert config.ipsec_tunnels[0].manual_key_configured is True
    assert "secret" not in config.model_dump_json()
    assert {item.status for item in derived.reference_resolutions} == {"RESOLVED"}

    result = PaloAltoSourceReporter().analyze_source(SOURCE)
    output = BytesIO()
    PaloAltoSourceReporter().export_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
    assert workbook["Administrators"].cell(4, 6).value is True
    assert not any("unresolved PAN-OS reference" in str(cell.value) for row in workbook["Validation"] for cell in row)
    coverage = workbook["Extraction Coverage"]
    assert any("ike_crypto_profile" in str(cell.value) for row in coverage for cell in row)
    assert any("ipsec_crypto_profile" in str(cell.value) for row in coverage for cell in row)
