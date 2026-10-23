from __future__ import annotations

import ast
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from fwmigrate.source_reporting.options import ExcelExportProfile

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter, PaloAltoSourceResult
from fwmigrate.vendors.palo_alto.export.excel_schema import HIDDEN_COLUMNS_BY_DEFAULT, SHEET_HEADERS, SHEET_ORDER
from fwmigrate.vendors.palo_alto.export.excel import _write_table
from fwmigrate.vendors.palo_alto.model import PANAddress, PANOSConfig
from fwmigrate.vendors.palo_alto.source_model import PANOSDerivedViews, PANScope
from fwmigrate.vendors.palo_alto.validation.models import PANOSValidationIssue, PANOSValidationResult


FIXTURES = Path(__file__).parents[2] / "fixtures" / "palo_alto"


@pytest.mark.parametrize("row_count", [0, 1, 2])
def test_fast_filter_only_covers_sheets_with_data(row_count):
    workbook = Workbook(write_only=True)
    rows = ({"Name": "object"} for _ in range(row_count))
    assert _write_table(workbook, "Review Required", rows) == row_count
    output = BytesIO()
    workbook.save(output)
    sheet = load_workbook(BytesIO(output.getvalue()))["Review Required"]
    last_column = get_column_letter(len(SHEET_HEADERS["Review Required"]))
    assert sheet.auto_filter.ref == (f"A3:{last_column}{row_count + 3}" if row_count else None)
    assert sheet["A1"].value == "Review Required"
    assert sheet.freeze_panes


def _workbook(name: str, profile=ExcelExportProfile.FULL):
    reporter = PaloAltoSourceReporter()
    result = reporter.analyze_source((FIXTURES / name).read_text(encoding="utf-8"))
    output = BytesIO()
    reporter.export_excel(result, output, source_name=name, profile=profile)
    output.seek(0)
    return load_workbook(output, data_only=True)


@pytest.mark.parametrize("name", [
    "objects.xml", "services.xml", "schedules.xml", "policies.xml",
    "nat_pipeline_conformance.xml", "interfaces_extended.xml",
    "integrated_firewall.xml", "integrated_panorama.xml", "dynamic_routing.xml",
])
def test_supported_fixtures_export_real_workbooks(name):
    workbook = _workbook(name)
    assert workbook.sheetnames[0] == "Summary"
    assert "Review Required" in workbook.sheetnames
    assert "Validation" not in workbook.sheetnames
    assert "PAN-OS Source Inventory" in workbook.sheetnames
    assert workbook.sheetnames == list(SHEET_ORDER)
    assert workbook["Summary"]["B5"].value == name


def test_complete_schema_is_retained_and_route_path_monitors_are_implemented():
    workbook = _workbook("objects.xml")
    assert "DHCP Servers" in SHEET_HEADERS
    assert "DHCP Servers" in workbook.sheetnames
    assert "Route Path Monitors" in workbook.sheetnames


def test_workbook_uses_fortigate_presentation_structure():
    workbook = _workbook("objects.xml")
    summary = workbook["Summary"]
    addresses = workbook["Addresses"]
    assert summary["A3"].value == "SOURCE DETAILS"
    assert summary.freeze_panes == "A5"
    assert summary["A1"].fill.fgColor.rgb[-6:] == "17324D"
    assert summary["C5"].value == "Export Profile"
    assert addresses["A1"].value == "Addresses"
    assert addresses["A3"].value == "Name"
    assert addresses.freeze_panes == "A4"
    assert addresses.auto_filter.ref
    assert addresses.cell(2, addresses.max_column).value == "Back to Summary"


@pytest.mark.parametrize("profile", [ExcelExportProfile.FULL, ExcelExportProfile.FAST])
def test_summary_inventory_has_inline_links_to_exported_sheets(profile):
    workbook = _workbook("objects.xml", profile=profile)
    summary = workbook["Summary"]
    links = [row for row in summary if row[2].hyperlink]
    assert links
    assert all(row[2].value == "Open" for row in links)
    assert all(row[2].hyperlink.target == f"#'{row[0].value}'!A1" for row in links)
    assert all(row[0].value in workbook.sheetnames for row in links)
    assert all(cell.value != "Workbook navigation" for row in summary for cell in row)
    assert summary["D5"].value == profile.value.upper()
    assert summary["B5"].value == "objects.xml"


def test_source_and_derived_policy_order_are_separate():
    sheet = _workbook("policies.xml")["Security Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Source Order"]).value == 14
    assert sheet.cell(4, headers["Rulebase Position"]).value == "pre"
    assert "source-only" in sheet.cell(4, headers["Effective Order"]).value


def test_interface_sheet_contains_topology_relationships():
    source = """<config><devices><entry name='fw'><network><interface>
      <ethernet><entry name='ethernet1/1'><layer3><aggregate-group>ae1</aggregate-group><units><entry name='ethernet1/1.10'/></units></layer3></entry></ethernet>
      <aggregate-ethernet><entry name='ae1'><layer3/></entry></aggregate-ethernet>
    </interface></network></entry></devices></config>"""
    reporter = PaloAltoSourceReporter()
    output = BytesIO()
    reporter.export_excel(reporter.analyze_source(source), output)
    sheet = load_workbook(BytesIO(output.getvalue()), data_only=True)["Interfaces"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Name"]).value: row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Name"]).value
    }
    assert sheet.cell(rows["ethernet1/1"], headers["Aggregate Interface"]).value == "ae1"
    assert sheet.cell(rows["ethernet1/1.10"], headers["Parent Interface"]).value == "ethernet1/1"
    assert sheet.cell(rows["ethernet1/1.10"], headers["Topology Path"]).value.endswith("ethernet1/1, ae1")


def test_nat_rows_join_derived_values_by_full_source_identity():
    sheet = _workbook("nat_pipeline_conformance.xml")["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Derived Source Translation Mode"]).value == "dynamic-ip-and-port"
    assert sheet.cell(6, headers["Derived Source Translation Mode"]).value == "static-ip"
    assert sheet.cell(6, headers["Derived Destination Translation Mode"]).value == "destination-translation"


def test_row_builders_do_not_reimplement_relationships_or_transforms():
    path = Path(__file__).parents[3] / "src" / "fwmigrate" / "vendors" / "palo_alto" / "export" / "excel_rows.py"
    calls = {node.func.id for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
    assert calls.isdisjoint({"resolve_references", "build_policy_order", "build_scope_hierarchy", "transform_nat"})


@pytest.mark.parametrize(("scopes", "target"), [
    ((PANScope(kind="vsys", name="vsys1", device_serial="A"), PANScope(kind="vsys", name="vsys1", device_serial="B")), 1),
    ((PANScope(kind="shared", name="shared"), PANScope(kind="vsys", name="vsys1")), 0),
    ((PANScope(kind="device-group", name="dg-a", device_group="dg-a"), PANScope(kind="device-group", name="dg-b", device_group="dg-b")), 1),
    ((PANScope(kind="template", name="base"), PANScope(kind="template-stack", name="base", template_stack="base")), 0),
])
def test_validation_issues_stay_on_the_matching_native_scope(scopes, target):
    addresses = [PANAddress(name="duplicate", source_path=f"addresses/{index}", scope=scope) for index, scope in enumerate(scopes)]
    issue = PANOSValidationIssue("warning", "address", "scoped warning", source_name="duplicate",
                                 source_scope=scopes[target], scope_identity=f"{scopes[target].kind}:{scopes[target].name}", object_type="address")
    analysis = PaloAltoSourceResult(PANOSConfig(addresses=addresses), PANOSDerivedViews(), PANOSValidationResult((issue,)))
    output = BytesIO()
    PaloAltoSourceReporter().export_excel(analysis, output)
    sheet = load_workbook(BytesIO(output.getvalue()), data_only=True)["Addresses"]
    reason_col = next(cell.column for cell in sheet[3] if cell.value == "Review Reasons")
    assert [sheet.cell(row, reason_col).value == "scoped warning" for row in (4, 5)] == [target == 0, target == 1]


def test_plain_writer_preserves_values_and_layout_after_reload():
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows = iter((
        {"Effective Order": "derived", "__review__": True},
        {"Effective Order": "derived", "__error__": True},
        {"Effective Order": "derived"},
        {"Effective Order": "derived"},
        {},
    ))
    assert _write_table(workbook, "Security Policies", rows, full=True) == 5
    output = BytesIO()
    workbook.save(output)
    sheet = load_workbook(BytesIO(output.getvalue()))["Security Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    derived = headers["Effective Order"]
    assert [sheet.cell(row, derived).value for row in range(4, 8)] == ["derived"] * 4
    assert sheet.sheet_view.showGridLines
    assert all(not cell.has_style for row in sheet.iter_rows(min_row=4) for cell in row)
    assert all(16 <= sheet.column_dimensions[get_column_letter(column)].width <= 44
               for column in range(1, len(SHEET_HEADERS["Security Policies"]) + 1))
    assert any(dimension.width == 44 for dimension in sheet.column_dimensions.values())
    assert sheet["A3"].font.sz == 10
    assert sheet["A3"].alignment.wrap_text
    assert sheet.row_dimensions[3].height == 30
    assert sheet["A2"].value.startswith("FULL report. Row counts are listed in Summary.")
    assert sheet.auto_filter.ref == f"A3:{get_column_letter(len(SHEET_HEADERS['Security Policies']))}8"
    assert sheet.freeze_panes == "D4"
    hidden = HIDDEN_COLUMNS_BY_DEFAULT.get("Security Policies", ())
    assert hidden
    assert all(sheet.column_dimensions[get_column_letter(column)].hidden
               for column, header in enumerate(SHEET_HEADERS["Security Policies"], 1) if header in hidden)
