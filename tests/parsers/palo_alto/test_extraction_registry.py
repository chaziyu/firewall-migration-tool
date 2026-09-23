from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.export.excel import export_panos_excel
from fwmigrate.vendors.palo_alto.export.excel_rows import ROW_BUILDERS, _PANExcelContext, _TYPED_COUNT_FIELDS, _typed_counts
from fwmigrate.vendors.palo_alto.export.excel_schema import ACTIVE_SHEET_ORDER, SHEET_HEADERS, SHEET_IMPLEMENTATION_STATUS
from fwmigrate.vendors.palo_alto.extraction.extractor import _EXTRACTORS, registered_typed_collections
from fwmigrate.vendors.palo_alto.model import PANOSConfig
from fwmigrate.vendors.palo_alto.schema_registry import registered_paths
from fwmigrate.vendors.palo_alto.source_builder import build_panos_config
from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def test_registered_typed_collections_are_unique_and_deterministic():
    collections = registered_typed_collections()
    assert collections == tuple(sorted(set(collections)))
    assert collections == tuple(sorted({collection for collection, _ in _EXTRACTORS.values()}))


def test_extractor_destinations_are_list_fields_on_panos_config():
    config = PANOSConfig()
    assert all(isinstance(getattr(config, collection), list) for collection in registered_typed_collections())


def test_registered_paths_are_represented_in_coverage_counts():
    assert set(registered_paths()) <= set(_TYPED_COUNT_FIELDS)
    assert set(registered_typed_collections()) <= set(_TYPED_COUNT_FIELDS.values())


def test_coverage_counts_include_new_typed_domains():
    analysis = PaloAltoSourceReporter().analyze_source("""<config><shared><network>
      <dhcp><entry name='dhcp-main'/></dhcp>
      <sdwan><rules><entry name='rule-main'/></rules></sdwan>
      <global-protect><portal><entry name='portal-main'/></portal></global-protect>
    </network></shared></config>""")

    counts = _typed_counts(_PANExcelContext(analysis))
    assert counts["dhcp_server"] == 1
    assert counts["sdwan_rule"] == 1
    assert counts["globalprotect_portal"] == 1


def test_source_builder_initializes_all_registered_collections():
    config = build_panos_config("<config />")
    assert all(isinstance(getattr(config, collection), list) for collection in registered_typed_collections())


def test_registered_source_path_reaches_expected_collection():
    config = build_panos_config(
        "<config><shared><tag><entry name='production'><color>red</color><comments>keep</comments></entry></tag></shared></config>"
    )
    assert config.tags[0].name == "production"
    assert config.tags[0].color == "red"


def test_active_excel_sheets_have_headers_and_row_generation_paths():
    assert set(ACTIVE_SHEET_ORDER) <= set(SHEET_HEADERS)
    assert set(ACTIVE_SHEET_ORDER) - {"Summary"} <= set(ROW_BUILDERS)


def test_tags_sheet_is_active_and_exports_rows():
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(
        "<config><shared><tag><entry name='production'><color>red</color><comments>keep</comments></entry></tag></shared></config>"
    )
    output = BytesIO()
    export_panos_excel(analysis, output)
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
    assert SHEET_IMPLEMENTATION_STATUS["Tags"] == "IMPLEMENTED"
    sheet = workbook["Tags"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Name"]).value == "production"
    assert sheet.cell(4, headers["Color"]).value == "red"
    assert sheet.cell(4, headers["Comments"]).value == "keep"
