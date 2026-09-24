import io
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.web import create_app
from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


FIXTURE = Path(__file__).parents[2] / "fixtures" / "example_palo_alto.xml"


def test_palo_alto_source_preview_and_excel_are_registered():
    client = create_app({"TESTING": True}).test_client()
    source = FIXTURE.read_bytes()

    preview = client.post(
        "/api/preview",
        data={
            "source_vendor": "palo_alto",
            "file": (io.BytesIO(source), "panos.xml"),
        },
        content_type="multipart/form-data",
    )

    assert preview.status_code == 200
    payload = preview.get_json()
    assert payload["vendor"] == "palo_alto"
    assert payload["summary"]["policies"] >= 0
    severity = payload["summary"]["validation"]["severity_counts"]
    assert severity["error"] == sum(issue["severity"] == "error" for issue in payload["validation"])
    assert severity["warning"] == sum(issue["severity"] == "warning" for issue in payload["validation"])
    assert payload["summary"]["objects"]["interfaces"] == len(payload["sections"]["interface_topology"])
    assert payload["summary"]["objects"]["policies"] == len(payload["sections"]["policies"])
    assert set(payload["sections"]) >= {"interfaces", "addresses", "address_groups", "services", "service_groups", "policies", "nat", "routes", "validation"}

    workbook = client.post(
        "/api/extract/excel",
        data={
            "source_vendor": "palo_alto",
            "preview_id": payload["preview_id"],
        },
        content_type="multipart/form-data",
    )

    assert workbook.status_code == 200
    exported = load_workbook(io.BytesIO(workbook.data), read_only=True, data_only=True)
    assert "Interfaces" in exported.sheetnames
    assert exported["Summary"]["B3"].value == "panos.xml"


def test_palo_alto_reporter_keeps_source_result_native_and_scope_aware():
    analysis = PaloAltoSourceReporter().analyze_source(FIXTURE.read_text())

    assert analysis.config.source_format == "xml"
    assert analysis.config.source_inventory
    assert analysis.config.scopes
    assert analysis.config.source_inventory[0].source_path
    assert not hasattr(analysis.config, "canonical_ir")
    assert not hasattr(analysis, "legacy_extraction")


def test_palo_alto_preview_exposes_interface_tree_and_vpn_attachment():
    source = """<config><devices><entry name='fw'><network><interface>
      <ethernet><entry name='ethernet1/1'><layer3><aggregate-group>ae1</aggregate-group><units><entry name='ethernet1/1.10'/></units></layer3></entry></ethernet>
      <aggregate-ethernet><entry name='ae1'><layer3/></entry></aggregate-ethernet>
      <tunnel><units><entry name='tunnel.1'/></units></tunnel>
      <ipsec><entry name='vpn1'><tunnel-interface>tunnel.1</tunnel-interface><auto-key><ike-gateway><member>gw1</member></ike-gateway></auto-key></entry></ipsec>
    </interface></network></entry></devices></config>"""
    payload = PaloAltoSourceReporter().build_preview(PaloAltoSourceReporter().analyze_source(source))

    rows = payload["sections"]["interface_topology"]
    assert any(row["display_name"].endswith("ethernet1/1") and row["aggregate"] == "ae1" for row in rows)
    assert any(row["parent"] == "ethernet1/1" and "ethernet1/1.10" in row["display_name"] for row in rows)
    assert payload["sections"]["vpn_tunnels"][0]["interface"] == "tunnel.1"
    assert payload["sections"]["vpn_tunnels"][0]["ike_gateways"] == ["gw1"]


def test_palo_alto_native_source_evidence_redacts_secret_values():
    source = "<config><shared><entry name='admin'><password>do-not-export</password></entry></shared></config>"

    analysis = PaloAltoSourceReporter().analyze_source(source)

    assert "do-not-export" not in str(analysis.config.model_dump())


def test_palo_alto_redacts_typed_inventory_preview_excel_and_validation():
    secret = "typed-secret-value"
    source = f"<config><shared><address><entry name='x'><future-password>{secret}</future-password></entry></address></shared></config>"
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(source)

    address = analysis.config.addresses[0]
    record = analysis.config.source_inventory[0]
    assert secret not in str(address.raw_extra)
    assert secret not in str(record.values)
    assert secret not in str(record.raw_xml)
    assert secret not in str(analysis.config.model_dump())
    assert secret not in str(analysis.validation)

    preview = reporter.build_preview(analysis)
    assert secret not in str(preview)

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)
    assert secret not in "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
