from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_custom_service_keeps_ipv4_and_port_hierarchy():
    fg = parse_fortigate_config("""
config firewall internet-service-custom
    edit "custom-web"
        set comment "Custom web"
        set reputation 3
        config entry
            edit 1
                set addr-mode ipv4
                set dst 192.0.2.0/24
                set protocol 6
                set future-entry keep
                config port-range
                    edit 1
                        set start-port 80
                        set end-port 443
                    next
                end
            next
        end
    next
end
""")
    item = fg.custom_internet_services[0]
    entry = item.entries[0]
    assert item.reputation == 3
    assert not hasattr(entry, "reputation")
    assert (entry.addr_mode, entry.dst, entry.dst6, entry.protocol) == (
        "ipv4", ["192.0.2.0/24"], [], 6
    )
    assert entry.extra_settings == {"future_entry": "keep"}
    ir_item = FGToIRTransformer(fg).transform().custom_internet_services[0]
    ir_entry = ir_item.entries[0]
    assert ir_item.reputation == 3
    assert not hasattr(ir_entry, "reputation")
    assert ir_entry.destination_ipv4 == ["192.0.2.0/24"]


def test_custom_service_accepts_ipv6_addr_mode():
    fg = parse_fortigate_config("""
config firewall internet-service-custom
    edit "custom-v6"
        config entry
            edit 1
                set addr-mode ipv6
                set dst6 2001:db8::/32
                set protocol 6
            next
        end
    next
end
""")
    entry = fg.custom_internet_services[0].entries[0]

    assert (entry.addr_mode, entry.dst6) == ("ipv6", ["2001:db8::/32"])


def test_custom_service_rejects_both_but_preserves_source_and_ir_review_data():
    fg = parse_fortigate_config("""
config firewall internet-service-custom
    edit "invalid-custom"
        config entry
            edit 1
                set addr-mode both
            next
        end
    next
end
""")
    entry = fg.custom_internet_services[0].entries[0]

    assert entry.addr_mode is None
    assert entry.extra_settings["invalid_fields"]["addr_mode"] == "both"

    ir_entry = FGToIRTransformer(fg).transform().custom_internet_services[0].entries[0]
    assert ir_entry.addr_mode is None
    assert ir_entry.source_attributes["invalid_fields"]["addr_mode"] == "both"


def test_custom_service_entry_id_uses_fortios_0_to_255_range():
    fg = parse_fortigate_config('''
config firewall internet-service-custom
    edit "custom-web"
        config entry
            edit 0
            next
            edit 255
            next
            edit 256
            next
            edit 4294967295
            next
        end
    next
end
''')
    entries = fg.custom_internet_services[0].entries

    assert [entry.id for entry in entries[:2]] == [0, 255]
    assert [entry.id for entry in entries[2:]] == [None, None]
    assert entries[2].extra_settings["invalid_fields"]["id"] == 256
    assert entries[3].extra_settings["invalid_fields"]["id"] == 4294967295


def test_custom_service_parent_reputation_validation_preserves_source_values():
    malformed = parse_fortigate_config('''
config firewall internet-service-custom
    edit "custom-web"
        set reputation invalid
    next
end
''').custom_internet_services[0]
    assert malformed.reputation is None
    assert malformed.extra_settings["unparsed_fields"]["reputation"] == "invalid"

    out_of_range = parse_fortigate_config('''
config firewall internet-service-custom
    edit "custom-web"
        set reputation 4294967296
    next
end
''').custom_internet_services[0]
    assert out_of_range.reputation is None
    assert out_of_range.extra_settings["invalid_fields"]["reputation"] == "4294967296"


def test_custom_service_reputation_is_not_typed_inside_entry():
    item = parse_fortigate_config('''
config firewall internet-service-custom
    edit "custom-web"
        config entry
            edit 1
                set reputation 3
            next
        end
    next
end
''').custom_internet_services[0]

    assert item.reputation is None
    assert not hasattr(item.entries[0], "reputation")
    assert item.entries[0].extra_settings["reputation"] == "3"


def test_custom_service_excel_reflects_parent_reputation():
    extraction = extract_fortigate_config('''
config firewall internet-service-custom
    edit "custom-web"
        set reputation 3
        config entry
            edit 1
                set protocol 6
            next
        end
    next
end
''')
    workbook = load_workbook(BytesIO(IRExcelExporter(extraction.canonical_ir).generate()))

    services = workbook["Custom Internet Services"]
    service_headers = {cell.value: cell.column for cell in services[3]}
    assert "Reputation" in service_headers
    assert services.cell(4, service_headers["Reputation"]).value == 3

    entries = workbook["Custom IS Entries"]
    assert "Reputation" not in {cell.value for cell in entries[3]}
