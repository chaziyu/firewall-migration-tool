from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.parser import FortiGateParser
from fwmigrate.parsers.fortigate.tokenizer import FortiGateTokenizer
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.ir.core import IRConfig, IRMetadata, IRNATRule
from fwmigrate.ir.enums import NATType
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_ip_translation_is_typed_and_preserves_unknown_settings():
    config = FortiGateParser(FortiGateTokenizer("""
config firewall ip-translation
    edit 7
        set type SCTP
        set startip 10.0.0.1
        set endip 10.0.0.4
        set map-startip 192.0.2.1
        set future-setting keep
    next
end
""")).parse()

    rule = config.ip_translations[0]
    assert (rule.id, rule.source_order) == (7, 1)
    assert (rule.startip, rule.endip, rule.map_startip) == (
        "10.0.0.1", "10.0.0.4", "192.0.2.1"
    )
    assert rule.extra_settings["future_setting"] == "keep"


def test_sctp_ip_translation_is_normalized_as_an_address_range_mapping():
    result = extract_fortigate_config("""
config firewall ip-translation
    edit 7
        set type SCTP
        set startip 10.0.0.1
        set endip 10.0.0.4
        set map-startip 192.0.2.1
    next
end
""")

    rule = result.canonical_ir.nat_rules[0]
    assert rule.type.value == "address-translation"
    assert rule.address_range_mappings[0].translated_end == "192.0.2.4"
    assert rule.protocol_name == "SCTP"
    assert rule.migration_status == "NORMALIZED"


def test_ip_translation_omitted_fields_use_effective_defaults_without_mutation():
    config = FortiGateParser(FortiGateTokenizer("""
config firewall ip-translation
    edit 7
    next
end
""")).parse()

    rule = config.ip_translations[0]
    result = extract_fortigate_config("""
config firewall ip-translation
    edit 7
    next
end
""")

    assert (rule.type, rule.startip, rule.endip, rule.map_startip) == (
        "SCTP", None, None, None
    )
    ir_rule = result.canonical_ir.nat_rules[0]
    assert ir_rule.address_range_mappings[0].model_dump() == {
        "original_start": "0.0.0.0",
        "original_end": "0.0.0.0",
        "translated_start": "0.0.0.0",
        "translated_end": "0.0.0.0",
    }
    assert ir_rule.migration_status == "NORMALIZED"


def test_ip_translation_partial_defaults_and_invalid_values_require_review():
    result = extract_fortigate_config("""
config firewall ip-translation
    edit 7
        set startip 10.0.0.1
        set endip 10.0.0.4
    next
    edit 8
        set startip 10.0.0.4
        set endip 10.0.0.1
    next
    edit 9
        set startip 10.0.0.1
        set endip 10.0.0.4
        set map-startip 255.255.255.254
    next
end
""")

    valid = result.canonical_ir.nat_rules[0]
    assert valid.address_range_mappings[0].translated_start == "0.0.0.0"
    assert valid.address_range_mappings[0].translated_end == "0.0.0.3"
    assert result.canonical_ir.nat_rules[1].migration_status == "PARTIALLY_NORMALIZED"
    assert result.canonical_ir.nat_rules[2].migration_status == "PARTIALLY_NORMALIZED"
    assert not result.canonical_ir.nat_rules[1].address_range_mappings
    assert not result.canonical_ir.nat_rules[2].address_range_mappings


def test_ip_translation_source_inventory_and_excel_ranges_are_retained():
    result = extract_fortigate_config("""
config firewall ip-translation
    edit 7
        set type SCTP
        set startip 10.0.0.1
        set endip 10.0.0.4
        set map-startip 192.0.2.1
        unset map-startip
        set map-startip 192.0.2.1
    next
end
""")

    inventory = [
        item for item in result.inventory_items
        if item.source_path == "firewall ip-translation"
    ]
    assert len(inventory) == 1
    assert inventory[0].source_id == "7"
    assert {command.key for command in inventory[0].commands} >= {
        "type", "startip", "endip", "map-startip",
    }
    assert any(command.operation == "unset" for command in inventory[0].commands)
    assert len(result.canonical_ir.nat_rules) == 1

    workbook = load_workbook(BytesIO(IRExcelExporter(
        result.canonical_ir,
        extraction_result=result,
    ).generate()))
    source_sheet = workbook["FortiGate Source Configuration"]
    source_rows = list(source_sheet.iter_rows(min_row=4, values_only=True))
    assert any(
        row[1] == "firewall ip-translation"
        and row[3] == "7"
        and row[6] == "map-startip"
        and row[7] == "192.0.2.1"
        for row in source_rows
    )

    nat_sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in nat_sheet[3]}
    nat_row = next(
        row for row in range(4, nat_sheet.max_row + 1)
        if nat_sheet.cell(row, headers["Name"]).value == "ip-translation-7"
    )
    assert nat_sheet.cell(nat_row, headers["Source Origin"]).value == "ip-translation"
    assert nat_sheet.cell(nat_row, headers["NAT Family"]).value == "nat44"
    assert nat_sheet.cell(nat_row, headers["Protocol / Number"]).value == "SCTP/132"
    assert nat_sheet.cell(nat_row, headers["Original Range Start"]).value == "10.0.0.1"
    assert nat_sheet.cell(nat_row, headers["Original Range End"]).value == "10.0.0.4"
    assert nat_sheet.cell(nat_row, headers["Translated Range Start"]).value == "192.0.2.1"
    assert nat_sheet.cell(nat_row, headers["Translated Range End"]).value == "192.0.2.4"


def test_nat_range_columns_stay_empty_for_ordinary_nat():
    workbook = load_workbook(BytesIO(IRExcelExporter(IRConfig(
        metadata=IRMetadata(source_vendor="fortigate"),
        nat_rules=[IRNATRule(name="ordinary-snat", type=NATType.SOURCE, sequence=1)],
    )).generate()))

    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Name"]).value == "ordinary-snat"
    assert all(
        sheet.cell(4, headers[column]).value in ("", None)
        for column in (
            "Original Range Start", "Original Range End",
            "Translated Range Start", "Translated Range End",
        )
    )
