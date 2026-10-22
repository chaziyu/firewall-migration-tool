"""Offline source-report browser acceptance against the built frontend and Flask.

Run with --fortigate/--palo-alto for local customer files; none are stored in fixtures.
Requires optional Playwright tooling and a Chromium binary. Never contacts a device.
"""

import argparse
import json
import re
import tempfile
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from fwmigrate.collection.contracts import CollectedSource, CollectionStatus
from fwmigrate.collection.snapshot import make_snapshot


FORTIGATE = b'''config firewall address
 edit server
  set subnet 192.0.2.10 255.255.255.255
 next
end
config firewall policy
 edit 17
  set name reviewed
  set srcaddr all
  set dstaddr server
  set service ALL
  set action accept
  set logtraffic all
  set comments "source comment"
 next
end
config system dns
 set primary 192.0.2.53
end
config unsupported-empty
end
'''
PANOS = b'''<config><shared><address><entry name="server"><ip-netmask>192.0.2.10/32</ip-netmask></entry></address>
<profile-group><entry name="profiles"><vulnerability><member>missing-profile</member></vulnerability></entry></profile-group>
<pre-rulebase><security><rules><entry name="reviewed"><source><member>missing-address</member></source>
<destination><member>server</member></destination><action>allow</action></entry></rules></security></pre-rulebase>
<post-rulebase><security><rules><entry name="reviewed"><source><member>any</member></source>
<destination><member>server</member></destination><action>deny</action></entry></rules></security></post-rulebase>
</shared></config>'''


def payload(vendor, path=None, name=None):
    if path:
        return str(path)
    return {"name": name or ("review.conf" if vendor == "fortigate" else "review.xml"),
            "mimeType": "text/plain", "buffer": FORTIGATE if vendor == "fortigate" else PANOS}


def upload(page, vendor, source):
    page.get_by_label("Source vendor").select_option(vendor)
    with page.expect_response(lambda response: response.url.endswith("/api/preview"), timeout=90000) as response:
        page.get_by_label("Drop your configuration here").set_input_files(source)
    result = response.value.json()
    assert response.value.status == 200 and result["success"]
    expect(page.get_by_role("heading", name="Source inventory", exact=True)).to_be_visible(timeout=90000)
    return result


def section(page, name):
    page.get_by_role("tab", name=name, exact=True).click()


def workbook_check(path, vendor, preview):
    workbook = load_workbook(path)
    sheet = workbook["Policies" if vendor == "fortigate" else "Security Policies"]
    rows = list(sheet.values)
    data = [dict(zip(rows[2], row)) for row in rows[3:]]
    policies = preview["sections"]["policies"]
    if vendor == "fortigate":
        assert [row["Rule #"] for row in data] == [row["policy_id"] for row in policies]
        for actual, source in zip(data, policies):
            assert actual["Log Setting"] == source["logtraffic"]
            assert actual["Comments"] == source["comments"]
    else:
        assert len(data) == len(policies) + len(preview["sections"]["default_security_rules"])
        assert [row["Name"] for row in data[:len(policies)]] == [row["name"] for row in policies]
        assert len(list(workbook["Unresolved References"].values)) - 3 == len(preview["sections"]["unresolved_references"])
    assert not any(cell.data_type in {"f", "e"} for ws in workbook for row in ws for cell in row)
    wrapped = [(cell, ws) for ws in workbook for row in ws.iter_rows(min_row=4) for cell in row
               if cell.alignment.wrap_text]
    assert wrapped, "Expected readable long business values"
    assert all((ws.row_dimensions[cell.row].height or 18) <= 409 for cell, ws in wrapped)
    workbook.close()


def review(page, vendor, source, directory):
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    start = perf_counter()
    preview = upload(page, vendor, source)
    load_seconds = perf_counter() - start
    references = preview["sections"]["unresolved_references"]
    expect(page.locator(".configuration-summary tr").filter(has_text="Unresolved references").locator("td").first).to_have_text(str(len(references)))
    page.screenshot(path=str(directory / f"{vendor}-overview.png"))
    for name in page.locator(".report-tabs button").all_text_contents():
        section(page, name)
    section(page, "Policies")
    page.get_by_role("button", name=re.compile("^Details for ")).first.click()
    expect(page.locator("dialog[open]")).to_be_visible()
    if vendor == "fortigate" and any(row["logtraffic"] is not None for row in preview["sections"]["policies"]):
        assert "logtraffic" in preview["sections"]["policies"][0]
    page.get_by_label("Close details", exact=True).click()

    section(page, "Extraction")
    page.locator("#subsection-tab-source_inventory").click()
    inventory = preview["sections"]["source_inventory"]
    scope_filter = page.get_by_label("Filter report scope", exact=True)
    if scope_filter.count():
        selected_scope = scope_filter.locator("option").nth(1).get_attribute("value")
        scope_filter.select_option(selected_scope)
        expected = sum((row.get("scope") or row.get("vdom") or "") == selected_scope for row in inventory)
        expect(page.locator(".report-row-count:visible")).to_contain_text(f"{expected:,} matching rows")
        scope_filter.select_option("")
    search = page.get_by_role("searchbox", name="Search report rows")
    start = perf_counter()
    line_key = "start_line_number" if vendor == "fortigate" else "source_order"
    search.fill(json.dumps({line_key: inventory[-1][line_key]}, separators=(",", ":"))[1:-1])
    expect(page.locator(".report-table tbody tr").first).to_be_visible()
    search_seconds = perf_counter() - start
    assert any(inventory[-1]["source_path"] in row for row in page.locator(".report-table tbody tr").all_text_contents())
    if vendor == "fortigate":
        assert len(inventory) == preview["summary"]["records"]
        search.fill("system dns")
        if any(row["source_path"] == "system dns" for row in inventory):
            expect(page.locator(".report-table tbody tr").first).to_be_visible()
    search.fill("")

    section(page, "Validation")
    navigable = [row for row in preview["sections"]["validation"]
                 if row.get("object_type") in {"policy", "security-profile-group", "ipsec-tunnel", "static-route", "route", "interface"}]
    if vendor == "palo_alto":
        assert navigable
        for index, finding in enumerate(preview["sections"]["validation"]):
            if finding not in navigable:
                continue
            page.locator(".validation-table tbody tr").nth(index).get_by_role("button", name=re.compile("^View object")).click()
            target = page.locator(".finding-object-target")
            expect(target).to_be_visible()
            target.locator("details > summary").first.click()
            assert finding["object_name"] in target.inner_text()
            assert "No matching object" not in target.inner_text()
            page.get_by_label("Close details", exact=True).click()

    section(page, "Policies")
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate("document.documentElement.scrollWidth") == 390
    page.get_by_role("button", name=re.compile("^Details for ")).first.click()
    assert page.locator("dialog[open]").evaluate("dialog => dialog.matches(':modal')")
    page.screenshot(path=str(directory / f"{vendor}-mobile-details.png"))
    page.get_by_label("Close details", exact=True).click()
    page.set_viewport_size({"width": 1440, "height": 1000})
    downloads = {}
    for profile in ("fast", "full"):
        page.locator('select[aria-label="Excel export profile"]:visible').select_option(profile)
        start = perf_counter()
        with page.expect_download(timeout=90000) as pending:
            page.get_by_role("button", name="Export Excel", exact=True).click()
        download = pending.value
        assert download.failure() is None
        path = directory / f"{vendor}-{profile}.xlsx"
        download.save_as(path)
        workbook_check(path, vendor, preview)
        downloads[profile] = {"seconds": round(perf_counter() - start, 3), "bytes": path.stat().st_size}

    for width in (1440, 390):
        page.set_viewport_size({"width": width, "height": 1000 if width == 1440 else 844})
        page.route("**/api/extract/excel", lambda route: route.fulfill(status=500, content_type="application/json", body='{"error":"Controlled export failure"}'))
        page.get_by_role("button", name="Export Excel", exact=True).click()
        alert = page.locator(".report-sticky-header").get_by_role("alert")
        expect(alert).to_have_text("Controlled export failure")
        expect(alert).to_be_in_viewport()
        page.screenshot(path=str(directory / f"{vendor}-error-{width}.png"))
        page.unroute("**/api/extract/excel")
        page.route("**/api/extract/excel", lambda route: route.fulfill(status=200,
                   content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", body=(directory / f"{vendor}-full.xlsx").read_bytes()))
        with page.expect_download() as retry:
            page.get_by_role("button", name="Export Excel", exact=True).click()
        assert retry.value.failure() is None
        expect(alert).to_have_count(0)
        page.unroute("**/api/extract/excel")
    assert not errors, errors
    return {"load_seconds": round(load_seconds, 3), "search_seconds": round(search_seconds, 3),
            "preview_bytes": len(json.dumps(preview)), "records": len(inventory), "references": len(references),
            "downloads": downloads, "javascript_errors": errors}


def stale_exports(browser, url, workbook_bytes):
    checks = 0
    for replacement in ("file", "vendor", "snapshot", "workspace"):
        for status in (200, 500):
            page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            page.goto(url)
            upload(page, "fortigate", payload("fortigate"))
            held, downloads = [], []
            page.on("download", lambda download: downloads.append(download))
            page.route("**/api/extract/excel", lambda route: held.append(route))
            page.get_by_role("button", name="Export Excel", exact=True).click()
            expect(page.get_by_role("button", name="Preparing workbook…", exact=True)).to_be_visible()
            if replacement == "workspace":
                page.once("dialog", lambda dialog: dialog.accept())
                page.get_by_role("button", name="New workspace", exact=False).click()
                expect(page.get_by_role("heading", name="Source inventory", exact=True)).to_have_count(0)
            else:
                page.get_by_role("button", name="Change source", exact=True).click()
                if replacement == "vendor":
                    page.get_by_label("Source vendor").select_option("palo_alto")
                elif replacement == "file":
                    upload(page, "fortigate", payload("fortigate", name="replacement.conf"))
                else:
                    snapshot = make_snapshot(CollectedSource("fortigate", FORTIGATE.decode(), "snapshot.conf", "test", CollectionStatus.SUCCESS))
                    page.get_by_role("tab", name=re.compile("Upload Snapshot")).click()
                    page.get_by_label("Drop your collection snapshot here").set_input_files(
                        {"name": "snapshot.json", "mimeType": "application/json", "buffer": json.dumps(snapshot).encode()})
                    expect(page.get_by_role("heading", name="Source inventory", exact=True)).to_be_visible()
            has_new_source = replacement in {"file", "snapshot"}
            if has_new_source:
                expect(page.get_by_role("button", name="Export Excel", exact=True)).to_be_enabled()
                page.get_by_role("button", name="Export Excel", exact=True).click()
                expect(page.get_by_role("button", name="Preparing workbook…", exact=True)).to_be_visible()
            assert held
            held[0].fulfill(status=status, content_type="application/json" if status == 500 else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            body='{"error":"Obsolete failure"}' if status == 500 else workbook_bytes)
            page.wait_for_timeout(150)
            assert not downloads, (replacement, status, "stale download")
            expect(page.get_by_role("alert").filter(has_text="Obsolete failure")).to_have_count(0)
            if has_new_source:
                expect(page.get_by_role("button", name="Preparing workbook…", exact=True)).to_be_visible()
                assert len(held) == 2
                with page.expect_download() as pending:
                    held[1].fulfill(status=200, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", body=workbook_bytes)
                assert pending.value.failure() is None
                expect(page.get_by_role("button", name="Export Excel", exact=True)).to_be_enabled()
            checks += 1
            page.close()
    return checks


def run(url, chromium, fortigate, palo_alto, output_dir):
    with sync_playwright() as playwright, tempfile.TemporaryDirectory() as temporary:
        directory = Path(output_dir or temporary)
        directory.mkdir(parents=True, exist_ok=True)
        browser = playwright.chromium.launch(executable_path=chromium)
        result = {}
        for vendor, path in (("fortigate", fortigate), ("palo_alto", palo_alto)):
            page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            page.goto(url)
            result[vendor] = review(page, vendor, payload(vendor, path), directory)
            print(json.dumps({vendor: result[vendor]}), flush=True)
            page.close()
        result["stale_export_checks"] = stale_exports(browser, url, (directory / "fortigate-fast.xlsx").read_bytes())
        browser.close()
        (directory / "browser-results.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        print(json.dumps({"flow": "PASS", "stale_export_checks": result["stale_export_checks"]}), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:5000")
    parser.add_argument("--chromium")
    parser.add_argument("--fortigate", type=Path)
    parser.add_argument("--palo-alto", type=Path)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    run(args.url, args.chromium, args.fortigate, args.palo_alto, args.output_dir)
