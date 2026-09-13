from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_stateless_firewall_filters_preserve_terms_and_attachments():
    ir = JuniperSRXParser("""
    set interfaces ge-0/0/0 unit 0 family inet filter input F
    set interfaces ge-0/0/1 unit 0 family inet filter output F
    set firewall family inet filter F term first from source-address 10.0.0.0/8
    set firewall family inet filter F term first then accept
    set firewall family inet filter F term second then discard
    set firewall family inet6 filter F6 term reject then reject
    """).transform_to_ir()

    filt = next(item for item in ir.firewall_filters if item.name == "F")
    assert [term.name for term in filt.terms] == ["first", "second"]
    assert [term.source_order for term in filt.terms] == sorted(term.source_order for term in filt.terms)
    assert {item["direction"] for item in filt.attachments} == {"input", "output"}
    assert filt.terms[0].actions[0]["action"] == ["accept"]
    assert not ir.policy_route_rules

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    assert "Firewall Filters" in workbook.sheetnames
    sheet = workbook["Firewall Filters"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = list(sheet.iter_rows(min_row=4, values_only=False))
    assert rows[0][headers["Term"] - 1].value == "first"
    assert "ge-0/0/0.0" in rows[0][headers["Interface Attachments"] - 1].value

