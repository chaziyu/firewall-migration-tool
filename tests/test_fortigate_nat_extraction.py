"""Production-path NAT regression tests. All addresses/configs are synthetic."""

import io
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

from fwmigrate.core.constants import IR_KEYWORD_ANY
from fwmigrate.extraction.models import ExtractionStatus as Status
from fwmigrate.parsers.fortigate import FortiGateSourceParser
from fwmigrate.parsers.fortigate.api_client import FortiGateAPIClient
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter
from fwmigrate.web import create_app


BASE = '''#config-version=FGVM-7.4.8-FW-build0000
config system interface
    edit "port1"
        set vdom "root"
        set ip 192.0.2.1 255.255.255.0
        set role lan
    next
    edit "wan1"
        set vdom "root"
        set mode dhcp
        set role wan
    next
end
config firewall address
    edit "CLIENTS"
        set subnet 192.0.2.0 255.255.255.0
    next
end
'''


def policy(identifier=10, extra='', destination='all', service='ALL'):
    return f'''config firewall policy
    edit {identifier}
        set name "Outbound"
        set srcintf "port1"
        set dstintf "wan1"
        set srcaddr "CLIENTS"
        set dstaddr "{destination}"
        set action accept
        set service "{service}"
        set schedule "always"
        set nat enable
        {extra}
    next
end
'''


def pool(name='PUBLIC', extra=''):
    return f'''config firewall ippool
    edit "{name}"
        set startip 203.0.113.10
        set endip 203.0.113.20
        {extra}
    next
end
'''


def vip(extra='', name='DNS'):
    return f'''config firewall vip
    edit "{name}"
        set extip 203.0.113.53
        set mappedip "192.0.2.53"
        set extintf "wan1"
        set portforward enable
        set protocol udp
        set extport 53
        set mappedport 5353
        {extra}
    next
end
'''


CENTRAL = '''config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 25
        set srcintf "port1"
        set dstintf "wan1"
        set orig-addr "CLIENTS"
        set dst-addr "all"
        set protocol 17
        set orig-port 10000-11000
        set dst-port 53
        set nat-ippool "PUBLIC"
        set nat-port 20000-21000
        set status disable
    next
end
'''


def parse(text):
    return FortiGateSourceParser().parse(text)


def workbook(ir):
    return load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))


def rows(book, name):
    sheet = book[name]
    headers = [c.value for c in sheet[3]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=4, values_only=True)]


def test_policy_interface_snat_reaches_excel_without_invented_ip_or_zones():
    ir = parse(BASE + policy())
    assert len(ir.nat_rules) == 1
    nat = ir.nat_rules[0]
    assert nat.source == ['CLIENTS']
    assert nat.destination == [IR_KEYWORD_ANY]
    assert nat.service == IR_KEYWORD_ANY
    assert nat.source_interfaces == ['port1']
    assert nat.destination_interfaces == ['wan1']
    assert nat.from_zone == nat.to_zone == []  # roles are not source zone definitions
    assert nat.interface_address and nat.translated_source is None
    assert nat.source_policy == '10'
    assert not nat.requires_manual_review
    row = rows(workbook(ir), 'NAT Rules')[0]
    assert row['Use Outgoing Interface Address'] == 'Yes'
    assert row['Source Policy ID'] == '10'
    assert row['Destination Interfaces'] == 'wan1'
    assert row['Translated Source'] is None


def test_unused_pool_is_inventory_not_rule_and_all_fields_are_retained():
    ir = parse(BASE + pool(extra='set type one-to-one\nset arp-reply disable\nset new-feature-x enable'))
    assert ir.nat_rules == []
    record = ir.extraction.objects[0]
    assert record.attributes['arp_reply'] == 'disable'
    assert record.attributes['type'] == 'one-to-one'
    assert record.attributes['new_feature_x'] == 'enable'
    assert record.blocking
    settings = {r['Setting']: r['Explicit Source Value'] for r in rows(workbook(ir), 'NAT Source Settings')}
    assert settings['type'] == 'one-to-one'
    assert settings['new-feature-x'] == 'enable'


def test_shared_pool_produces_policy_rules_with_order_and_disabled_state():
    content = BASE + pool() + policy(200, 'set ippool enable\nset poolname "PUBLIC"')
    content += policy(3, 'set ippool enable\nset poolname "PUBLIC"\nset status disable')
    ir = parse(content)
    assert len(ir.nat_rules) == 2
    assert [n.source_policy for n in ir.nat_rules] == ['200', '3']
    assert [n.enabled for n in ir.nat_rules] == [True, False]
    assert all(n.translated_source == '203.0.113.10-203.0.113.20' for n in ir.nat_rules)
    resource = next(o for o in ir.extraction.objects if o.section == 'firewall ippool')
    assert len(resource.references) == 2


@pytest.mark.parametrize('extra', ['set nat disable', 'unset nat', 'set action deny'])
def test_non_snat_policy_preserved_without_spurious_rule(extra):
    ir = parse(BASE + policy(extra=extra))
    assert ir.nat_rules == []
    assert len(ir.extraction.objects) == 1


def test_udp_vip_protocol_ports_public_ip_and_group_links():
    group = '''config firewall vipgrp
    edit "DNS_GROUP"
        set interface "wan1"
        set member "DNS"
    next
end
'''
    ir = parse(BASE + vip() + group + policy(extra='set nat disable', destination='DNS_GROUP'))
    nat = ir.nat_rules[0]
    assert nat.protocol == 'udp'
    assert nat.original_destination_port == '53' and nat.translated_port == '5353'
    assert nat.original_destination_values == ['203.0.113.53']
    assert nat.translated_destination == '192.0.2.53'
    assert next(s for s in ir.services if s.name == nat.service).ports[0].protocol.value == 'udp'
    assert next(g for g in ir.address_groups if g.name == 'DNS_GROUP').members == ['DNS']
    assert next(o for o in ir.extraction.objects if o.section == 'firewall vip').references
    row = rows(workbook(ir), 'NAT Rules')[0]
    assert row['Protocol'] == 'udp'
    assert row['Original Destination IPs'] == '203.0.113.53'


def test_vip_ranges_are_not_fabricated_as_host_addresses():
    ir = parse(BASE + vip('set extip 203.0.113.50-203.0.113.59\nset mappedip "192.0.2.50-192.0.2.59"\nset extport 8000-8009\nset mappedport 80-89'))
    address = next(a for a in ir.addresses if a.name == 'DNS')
    assert address.type.value == 'range'
    assert address.ip_range_start == '203.0.113.50'
    assert address.ip_range_end == '203.0.113.59'
    assert ir.nat_rules[0].translated_destination == '192.0.2.50-192.0.2.59'
    assert ir.nat_rules[0].requires_manual_review


def test_central_snat_mode_match_ports_and_disabled_state():
    ir = parse(BASE + pool() + CENTRAL + policy())
    assert len(ir.nat_rules) == 1  # no duplicate policy SNAT in central mode
    nat = ir.nat_rules[0]
    assert nat.name == 'CENTRAL_SNAT_root_25'
    assert not nat.enabled and nat.protocol == 'udp'
    assert nat.original_source_port == '10000-11000'
    assert nat.original_destination_port == '53'
    assert nat.translated_source_port == '20000-21000'
    assert nat.pool_references == ['PUBLIC']
    assert not nat.requires_manual_review


def test_central_no_nat_and_inactive_mode_are_distinct():
    no_nat = CENTRAL.replace('set nat-ippool "PUBLIC"', 'set nat disable')
    ir = parse(BASE + no_nat)
    assert ir.nat_rules[0].translation_mode == 'identity'
    assert ir.nat_rules[0].translated_source is None
    inactive = parse(BASE + CENTRAL.replace('set central-nat enable', 'set central-nat disable'))
    assert inactive.nat_rules == []
    assert any(o.section == 'firewall central-snat-map' for o in inactive.extraction.objects)


def test_ngfw_policy_mode_uses_central_snat():
    content = CENTRAL.replace('set central-nat enable', 'set ngfw-mode policy-based')
    ir = parse(BASE + pool() + content + policy())
    assert len(ir.nat_rules) == 1
    assert ir.nat_rules[0].source_section == 'firewall central-snat-map'


@pytest.mark.parametrize('extra', [
    'set ippool enable\nset poolname "MISSING"',
    'set ippool enable', 'set srcaddr "MISSING"', 'set srcintf "MISSING"',
    'set service "MISSING"', 'set srcaddr-negate enable', 'set nat64 enable',
])
def test_unsafe_policy_settings_are_visible_and_block_migration(extra):
    ir = parse(BASE + policy(extra=extra))
    assert ir.extraction.status == 'PARTIAL'
    assert ir.nat_rules[0].requires_manual_review
    with pytest.raises(ValueError, match='manual review'):
        ir.assert_nat_migration_ready()
    assert rows(workbook(ir), 'NAT Diagnostics')


def test_malformed_nat_object_is_retained_and_does_not_hide_other_objects():
    content = BASE + pool() + vip().replace('set mappedip "192.0.2.53"', '')
    ir = parse(content)
    broken = next(o for o in ir.extraction.objects if o.section == 'firewall vip')
    assert broken.status == Status.PARSE_ERROR and not broken.parsed
    assert len(ir.extraction.objects) == 2
    assert rows(workbook(ir), 'NAT Inventory')[1]['Status'] == 'PARSE_ERROR'


def test_unclosed_config_is_reported_and_blocks_generation():
    ir = parse(BASE + vip().rsplit('end', 1)[0])
    assert ir.extraction.blocking_issues
    assert ir.extraction.status == 'PARTIAL'
    with pytest.raises(ValueError):
        ir.assert_nat_migration_ready()


def test_multiple_vdoms_preserved_without_cross_scope_merge():
    content = 'config vdom\nedit "BLUE"\n' + pool() + 'next\nedit "RED"\n' + pool() + 'next\nend\n'
    ir = parse(content)
    assert {o.scope for o in ir.extraction.objects} == {'BLUE', 'RED'}
    assert len(ir.extraction.objects) == 2
    assert ir.nat_rules == []
    assert all(o.status == Status.UNSUPPORTED for o in ir.extraction.objects)


def test_single_nested_vdom_is_normalized_with_scope():
    ir = parse('config vdom\nedit "root"\n' + BASE + policy() + 'next\nend\n')
    assert ir.nat_rules[0].scope_id == 'root'
    assert ir.nat_rules[0].source_policy == '10'


def test_unknown_ipv6_and_nested_nat_settings_are_accounted_for_and_redacted():
    content = BASE + vip('''set api-token "NEVER_EXPORT"
        config realservers
            edit 1
                set ip 192.0.2.8
                set password "NESTED_SECRET"
            next
        end''') + '''config firewall vip6
edit "IPV6_VIP"
set extip 2001:db8::1
set mappedip 2001:db8:1::1
next
end
'''
    ir = parse(content)
    assert any(o.status == Status.UNSUPPORTED for o in ir.extraction.objects)
    serialized = ir.model_dump_json()
    assert 'NEVER_EXPORT' not in serialized and 'NESTED_SECRET' not in serialized
    assert '192.0.2.8' in serialized and '[REDACTED]' in serialized


def test_source_settings_prevent_formula_injection_and_preserve_long_values():
    long_text = 'x' * 40000
    ir = parse(BASE + vip(f'set comment "{long_text}"', name='=1+1'))
    book = workbook(ir)
    settings = rows(book, 'NAT Source Settings')
    chunks = [r for r in settings if r['Setting'] == 'comment']
    assert len(chunks) == 2
    assert ''.join(r['Explicit Source Value'] for r in chunks) == long_text
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)


def test_upload_to_excel_runs_production_nat_parser():
    client = create_app({'TESTING': True}).test_client()
    response = client.post('/api/extract/excel', data={
        'source_vendor': 'fortigate', 'file': (io.BytesIO((BASE + vip()).encode()), 'synthetic.conf'),
    })
    assert response.status_code == 200
    row = rows(load_workbook(io.BytesIO(response.data)), 'NAT Rules')[0]
    assert row['Protocol'] == 'udp' and row['Translated Port'] == '5353'


@pytest.mark.parametrize('vendor', ['fortigate', 'palo_alto', 'cisco_asa', 'checkpoint', 'juniper_srx'])
def test_generators_cannot_silently_drop_extended_nat(vendor):
    from fwmigrate.core.registry import PluginRegistry
    ir = parse(BASE + policy())
    with pytest.raises(ValueError, match='target-specific'):
        PluginRegistry.get_generator(vendor).generate(ir)


def test_api_preserves_udp_multiple_mapped_addresses_and_source_settings():
    client = FortiGateAPIClient('synthetic.invalid', api_key='test-token')
    data = {
        'cmdb/system/global': [{'hostname': 'SYNTHETIC'}],
        'cmdb/system/settings': [{'central-nat': 'disable'}],
        'cmdb/firewall/vip': [{'name': 'DNS', 'extip': '203.0.113.53',
            'mappedip': [{'q_origin_key': '192.0.2.53'}, {'q_origin_key': '192.0.2.54'}],
            'protocol': 'udp', 'portforward': 'enable', 'extport': 53, 'mappedport': 5353,
            'api-token': 'NEVER_EXPORT'}],
    }
    with patch.object(client, 'get', side_effect=lambda endpoint: data.get(endpoint, [])):
        ir = FGToIRTransformer(client.extract_config()).transform()
    nat = ir.nat_rules[0]
    assert nat.protocol == 'udp'
    assert nat.translated_destinations == ['192.0.2.53', '192.0.2.54']
    assert 'NEVER_EXPORT' not in ir.model_dump_json()


def test_api_unavailable_nat_endpoint_is_not_an_empty_success():
    client = FortiGateAPIClient('synthetic.invalid', api_key='test-token')
    def get(endpoint):
        if endpoint == 'cmdb/firewall/central-snat-map':
            raise KeyError('not available')
        return [{'hostname': 'SYNTHETIC'}] if endpoint == 'cmdb/system/global' else []
    with patch.object(client, 'get', side_effect=get):
        ir = FGToIRTransformer(client.extract_config()).transform()
    assert ir.extraction.status == 'PARTIAL'
    assert any('central-snat-map' in m for m in ir.extraction.blocking_issues)


def test_api_pagination_collects_all_pages_and_rejects_missing_cursor():
    client = FortiGateAPIClient('synthetic.invalid', api_key='test-token')
    response = MagicMock(status_code=200)
    response.json.side_effect = [
        {'results': [{'name': 'a'}], 'limit_reached': True, 'next_idx': 1},
        {'results': [{'name': 'b'}], 'limit_reached': False},
    ]
    with patch.object(client.session, 'get', return_value=response) as request:
        assert client.get('cmdb/firewall/vip') == [{'name': 'a'}, {'name': 'b'}]
        assert request.call_args.kwargs['params']['start'] == 1
    response.json.side_effect = None
    response.json.return_value = {'results': [], 'limit_reached': True}
    with patch.object(client.session, 'get', return_value=response):
        with pytest.raises(ValueError, match='pagination metadata'):
            client.get('cmdb/firewall/vip')
