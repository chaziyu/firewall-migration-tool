"""Static NAT analysis regressions; all configurations are synthetic."""

import io
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fwmigrate.parsers.fortigate import FortiGateSourceParser
from fwmigrate.ir.core import IRConfig
from fwmigrate.ir.enums import NATType
from fwmigrate.report.excel_exporter import IRExcelExporter


FIXTURE = Path(__file__).parent / "fixtures" / "fortigate" / "nat_analysis_synthetic.conf"


def parse(text=None):
    return FortiGateSourceParser().parse(FIXTURE.read_text() if text is None else text)


def analysis(text=None):
    return parse(text).extraction.nat_analysis


def findings(result, code):
    return [d for d in result.diagnostics if d.diagnostic_id == code]


def change_policy(text, identifier, old, new):
    start = text.index(f'    edit {identifier}\n')
    end = text.index('    next', start)
    return text[:start] + text[start:end].replace(old, new) + text[end:]


def prepend_policy(text, body):
    return text.replace('config firewall policy\n', 'config firewall policy\n' + body, 1)


def test_synthetic_fixture_rule_and_inventory_counts():
    ir = parse()
    assert len(ir.policies) == 12
    assert len(ir.nat_rules) == 12
    assert sum(n.effective is True for n in ir.nat_rules) == 9
    assert Counter(n.type for n in ir.nat_rules) == {NATType.SOURCE: 6, NATType.DESTINATION: 6}
    assert all(n.effective is False and n.analysis_status == 'DISABLED' for n in ir.nat_rules if not n.enabled)
    pools = [o for o in ir.extraction.nat_analysis.inventory if o.object_type == 'IP_POOL']
    vips = [o for o in ir.extraction.nat_analysis.inventory if o.object_type == 'VIP']
    assert len(pools) == 5 and len(vips) == 6
    assert Counter(o.usage_state for o in pools) == {'ACTIVE': 3, 'DISABLED_ONLY': 1, 'UNUSED': 1}
    assert Counter(o.usage_state for o in vips) == {'ACTIVE': 5, 'UNREFERENCED': 1}
    by_name = {o.name: o for o in pools + vips}
    assert by_name['POOL_UNUSED'].usage_state == 'UNUSED'
    assert by_name['POOL_PBA_LAB'].disabled_policy_refs == ['50']
    assert by_name['POOL_GENERAL'].active_policy_refs == ['5', '30']
    assert by_name['VIP_DNS'].active_policy_refs == ['130']
    assert by_name['VIP_WEB_HTTPS_DUP'].usage_state == 'UNREFERENCED'
    duplicate = next(n for n in ir.nat_rules if n.name.endswith('VIP_WEB_HTTPS_DUP'))
    assert duplicate.effective is False and duplicate.analysis_status == 'CONFIGURED_NOT_EFFECTIVE'


def test_synthetic_fixture_traffic_counts_and_semantics():
    coverage = analysis().traffic_coverage
    assert len(coverage) == 7
    assert Counter(c.status for c in coverage) == {'PASS': 5, 'WARN': 1, 'FAIL': 1}
    by_key = {(c.source[0], c.egress): c for c in coverage}
    for name in ('RANGE_FINANCE', 'NET_SERVERS', 'NET_USERS', 'NET_VOICE'):
        row = by_key[name, 'wan1']
        assert row.nat_state == 'COVERED' and row.coverage_percent == 100
    servers = by_key['NET_SERVERS', 'wan1']
    assert any('Excluding' in n for n in servers.notes)
    assert servers.source_ranges == ['10.10.20.0-10.10.20.99', '10.10.20.105-10.10.20.255']
    assert by_key['NET_GUEST', 'wan1'].nat_state == 'DISABLED_ONLY'
    assert by_key['NET_GUEST', 'wan1'].coverage_percent == 0
    assert by_key['NET_DMZ', 'wan1'].nat_state == 'NOT_TRANSLATED'
    assert by_key['NET_DMZ', 'wan1'].coverage_percent == 0
    assert by_key['NET_SERVERS', 'VPN_HQ'].nat_state == 'EXEMPT'
    assert by_key['NET_SERVERS', 'VPN_HQ'].coverage_percent == 100


@pytest.mark.parametrize('code,status,severity', [
    ('NAT-001', 'FAIL', 'CRITICAL'), ('NAT-002', 'FAIL', 'HIGH'),
    ('NAT-003', 'WARN', 'MEDIUM'), ('NAT-004', 'WARN', 'MEDIUM'),
    ('NAT-005', 'WARN', 'LOW'), ('NAT-006', 'WARN', 'LOW'), ('NAT-007', 'WARN', 'LOW'),
    ('NAT-008', 'PASS', 'INFO'), ('NAT-009', 'PASS', 'INFO'), ('NAT-010', 'PASS', 'INFO'),
])
def test_each_requested_diagnostic(code, status, severity):
    result = analysis()
    matches = findings(result, code)
    assert len(matches) == 1
    assert matches[0].status == status and matches[0].severity == severity
    assert matches[0].objects and matches[0].finding and matches[0].recommended_action
    assert Counter(d.status for d in result.diagnostics) == {'FAIL': 2, 'WARN': 5, 'PASS': 3}


@pytest.mark.parametrize('service', ['DNS', 'HTTP', 'HTTPS', 'NTP', 'SSH', 'PING', 'RDP', 'SMTP', 'ALL'])
def test_predefined_services_do_not_need_custom_definition(service):
    ir = parse(FIXTURE.read_text().replace('set service "DNS" "HTTP" "HTTPS" "NTP"', f'set service "{service}"'))
    notes = ' '.join(n for o in ir.extraction.objects for n in o.notes)
    assert 'Unresolved service reference' not in notes
    users = next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.source == ['NET_USERS'])
    assert users.nat_state == 'COVERED'


@pytest.mark.parametrize('name,kind', [('POOL_VOICE', 'fixed-port-range'), ('POOL_PBA_LAB', 'port-block-allocation')])
def test_advanced_pool_source_validity_separate_from_migration(name, kind):
    ir = parse()
    record = next(o for o in ir.extraction.objects if o.name == name)
    item = next(o for o in ir.extraction.nat_analysis.inventory if o.name == name)
    assert record.parsed and not record.blocking
    assert record.attributes['type'] == kind
    assert item.source_validity == 'VALID'
    assert item.migration_compatibility == 'SOURCE_VALID_TARGET_UNSUPPORTED'
    assert not any(n.migration_eligible for n in ir.nat_rules)


def test_unknown_custom_service_remains_unresolved():
    ir = parse(FIXTURE.read_text().replace('"DNS" "HTTP" "HTTPS" "NTP"', '"CUSTOM_MISSING"'))
    assert any('Unresolved service reference: CUSTOM_MISSING' in n for o in ir.extraction.objects for n in o.notes)
    assert next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.source == ['NET_USERS']).coverage_percent is None


def test_reversed_policy_order_warns_and_marks_shadowed_rule():
    text = FIXTURE.read_text()
    start = text.index('    edit 70\n')
    middle = text.index('    edit 5\n', start)
    end = text.index('    edit 30\n', middle)
    text = text[:start] + text[middle:end] + text[start:middle] + text[end:]
    ir = parse(text)
    assert findings(ir.extraction.nat_analysis, 'NAT-009')[0].status == 'WARN'
    finance = next(n for n in ir.nat_rules if n.source_policy == '70')
    assert finance.effective is False and finance.analysis_status == 'SHADOWED'


def test_one_to_one_cardinality_mismatch_is_capacity_warning_not_parse_error():
    ir = parse(FIXTURE.read_text().replace('set endip 198.51.100.34', 'set endip 198.51.100.33'))
    assert findings(ir.extraction.nat_analysis, 'NAT-010')[0].status == 'WARN'
    assert next(o for o in ir.extraction.objects if o.name == 'POOL_FINANCE').parsed


@pytest.mark.parametrize('replacement', [
    ('set extport 443', 'set extport 444'),
    ('set protocol tcp', 'set protocol udp'),
    ('set extintf "wan1"', 'set extintf "another-wan"'),
])
def test_nonoverlapping_vip_is_not_false_conflict(replacement):
    text = FIXTURE.read_text()
    start = text.index('    edit "VIP_WEB_HTTPS_DUP"')
    end = text.index('    next', start)
    text = text[:start] + text[start:end].replace(*replacement) + text[end:]
    assert not findings(analysis(text), 'NAT-001')


def test_vip_port_range_intersection_detected():
    text = FIXTURE.read_text().replace('set extport 443', 'set extport 440-445', 1)
    assert findings(analysis(text), 'NAT-001')[0].status == 'FAIL'


def test_serialization_and_excel_analysis_are_generic_and_complete():
    ir = parse()
    restored = IRConfig.model_validate_json(ir.model_dump_json())
    assert restored.extraction.nat_analysis == ir.extraction.nat_analysis
    book = load_workbook(io.BytesIO(IRExcelExporter(restored).generate()))
    assert tuple(book.sheetnames) == IRExcelExporter.SHEET_ORDER
    assert 'NAT Coverage' not in book.sheetnames
    assert book['NAT Traffic Coverage'].max_row == 10
    assert book['NAT Diagnostics'].max_row == 13
    assert book['NAT Policy Summary'].max_row == 15
    summary = {r[0]: r[1] for r in book['Summary'].iter_rows(values_only=True) if len(r) >= 2}
    assert summary['Configured NAT Rules'] == 12 and summary['Effective NAT Rules'] == 9
    assert summary['SNAT Rules'] == 6 and summary['DNAT Rules'] == 6
    settings = list(book['NAT Source Settings'].values)
    assert any('source-startip' in r for r in settings)
    assert any('block-size' in r for r in settings)


def test_earlier_deny_prevents_false_effective_and_coverage_claim():
    text = prepend_policy(FIXTURE.read_text(), '''    edit 999
        set srcintf "lan"
        set dstintf "wan1"
        set srcaddr "NET_USERS"
        set dstaddr "all"
        set service "ALL"
        set action deny
    next
''')
    ir = parse(text)
    users = next(n for n in ir.nat_rules if n.source_policy == '30')
    assert users.effective is False and users.analysis_status == 'SHADOWED'
    row = next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.source == ['NET_USERS'])
    assert row.nat_state == 'NO_MATCH' and row.coverage_percent == 0


def test_enabled_fallback_replaces_disabled_only_coverage():
    text = change_policy(FIXTURE.read_text(), 60, 'set status disable', 'set status enable')
    result = analysis(text)
    row = next(c for c in result.traffic_coverage if c.source == ['NET_GUEST'])
    assert row.nat_state == 'COVERED' and row.coverage_percent == 100
    assert not findings(result, 'NAT-003')


@pytest.mark.parametrize('setting', ['set schedule "BUSINESS_HOURS"', 'set srcaddr-negate enable', 'set internet-service enable', 'set users "alice"'])
def test_uncertain_match_never_gets_full_coverage(setting):
    text = change_policy(FIXTURE.read_text(), 30, 'set name "Users_Internet"', 'set name "Users_Internet"\n' + setting)
    ir = parse(text)
    row = next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.source == ['NET_USERS'])
    assert row.status == 'WARN' and row.coverage_percent is None
    assert next(n for n in ir.nat_rules if n.source_policy == '30').effective is None


def test_partial_service_overlap_does_not_subtract_entire_finance_range():
    text = change_policy(FIXTURE.read_text(), 70, 'set service "ALL"', 'set service "HTTPS"')
    result = analysis(text)
    servers = next(c for c in result.traffic_coverage if c.source == ['NET_SERVERS'] and c.egress == 'wan1')
    assert servers.nat_state == 'PARTIAL' and servers.coverage_percent is None
    assert not any('Excluding' in n for n in servers.notes)
    assert findings(result, 'NAT-009')[0].status == 'WARN'


def test_destination_restrictions_are_not_treated_as_all_internet():
    text = change_policy(FIXTURE.read_text(), 70, 'set dstaddr "all"', 'set dstaddr "NET_HQ"')
    result = analysis(text)
    servers = next(c for c in result.traffic_coverage if c.source == ['NET_SERVERS'] and c.egress == 'wan1')
    assert servers.nat_state == 'PARTIAL' and servers.coverage_percent is None


@pytest.mark.parametrize('service,effective', [('HTTPS', True), ('HTTP', False), ('DNS', False)])
def test_vip_policy_service_intersection(service, effective):
    text = change_policy(FIXTURE.read_text(), 100, 'set service "ALL"', f'set service "{service}"')
    ir = parse(text)
    rule = next(n for n in ir.nat_rules if n.name.endswith('_VIP_WEB_HTTPS'))
    assert rule.effective is effective


def test_vip_service_checks_post_dnat_port():
    text = change_policy(FIXTURE.read_text(), 110, 'set service "ALL"', 'set service "SSH"')
    ir = parse(text)
    admin = next(n for n in ir.nat_rules if n.name.endswith('VIP_ADMIN_RISKY'))
    assert admin.original_destination_port == '2222' and admin.translated_port == '22'
    assert admin.effective is True


def test_builtin_custom_override_is_respected():
    text = FIXTURE.read_text() + '''config firewall service custom
    edit "ALL"
        set tcp-portrange 80
    next
end
'''
    ir = parse(text)
    assert next(n for n in ir.nat_rules if n.name.endswith('_VIP_WEB_HTTPS')).effective is False


def test_disabled_vip_policy_reference_is_not_active():
    text = change_policy(FIXTURE.read_text(), 100, 'set action accept', 'set action accept\nset status disable')
    ir = parse(text)
    item = next(o for o in ir.extraction.nat_analysis.inventory if o.name == 'VIP_WEB_HTTPS')
    assert item.usage_state == 'DISABLED_ONLY' and item.disabled_policy_refs == ['100']
    assert next(n for n in ir.nat_rules if n.name.endswith('_VIP_WEB_HTTPS')).effective is False


def test_unknown_wan_role_is_not_inferred_from_name():
    result = analysis(FIXTURE.read_text().replace('set role wan', 'set role undefined'))
    row = next(c for c in result.traffic_coverage if c.source == ['NET_DMZ'])
    assert row.nat_state == 'UNKNOWN' and row.coverage_percent is None
    assert not findings(result, 'NAT-002')


def test_tunnel_name_alone_is_not_proof_of_ipsec_exemption():
    text = FIXTURE.read_text().replace('config vpn ipsec phase1-interface', 'config unimplemented-vpn-section')
    result = analysis(text)
    row = next(c for c in result.traffic_coverage if c.egress == 'VPN_HQ')
    assert row.nat_state == 'UNKNOWN' and not findings(result, 'NAT-008')


def test_central_nat_does_not_reuse_policy_nat_coverage():
    text = FIXTURE.read_text() + 'config system settings\nset central-nat enable\nend\n'
    result = analysis(text)
    assert all(c.coverage_percent is None for c in result.traffic_coverage)
    assert not findings(result, 'NAT-002') and not findings(result, 'NAT-008')
    assert all(p.classification == 'REVIEW_REQUIRED' for p in result.policy_summaries)


def test_fixed_port_range_outside_source_allocation_is_unknown():
    ir = parse(FIXTURE.read_text().replace('set source-endip 10.10.40.255', 'set source-endip 10.10.40.100'))
    voice = next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.source == ['NET_VOICE'])
    assert voice.coverage_percent is None
    assert next(n for n in ir.nat_rules if n.source_policy == '40').effective is None


@pytest.mark.parametrize('setting', ['set block-size 1', 'set num-blocks-per-user 0'])
def test_invalid_advanced_pool_values_remain_parse_errors(setting):
    text = FIXTURE.read_text().replace('set num-blocks-per-user 8', 'set num-blocks-per-user 8\n' + setting)
    ir = parse(text)
    record = next(o for o in ir.extraction.objects if o.name == 'POOL_PBA_LAB')
    assert not record.parsed and record.status.value == 'PARSE_ERROR'


def test_nested_address_group_cardinality_deduplicates_members():
    text = FIXTURE.read_text().replace('set srcaddr "RANGE_FINANCE"', 'set srcaddr "FINANCE_GROUP"') + '''config firewall addrgrp
    edit "FINANCE_GROUP"
        set member "RANGE_FINANCE" "FINANCE_NESTED"
    next
    edit "FINANCE_NESTED"
        set member "RANGE_FINANCE"
    next
end
'''
    diagnostic = findings(analysis(text), 'NAT-010')[0]
    assert diagnostic.status == 'PASS' and 'cardinality: 5' in diagnostic.finding


def test_service_group_missing_member_stays_unresolved():
    text = change_policy(FIXTURE.read_text(), 30, 'set service "DNS" "HTTP" "HTTPS" "NTP"', 'set service "WEB_GROUP"') + '''config firewall service group
    edit "WEB_GROUP"
        set member "HTTPS" "MISSING_CUSTOM"
    next
end
'''
    ir = parse(text)
    assert any('Unresolved service reference: MISSING_CUSTOM' in note for o in ir.extraction.objects for note in o.notes)
    assert next(n for n in ir.nat_rules if n.source_policy == '30').effective is None


def test_names_and_ids_are_not_hardcoded():
    text = FIXTURE.read_text()
    for old, new in [('NET_GUEST', 'NETWORK_X'), ('POOL_PBA_LAB', 'RESOURCE_X'), ('RANGE_FINANCE', 'RANGE_X'),
                     ('VIP_WEB_HTTPS_DUP', 'DUPLICATE_X'), ('edit 70\n', 'edit 707\n')]:
        text = text.replace(old, new)
    ir = parse(text)
    assert sum(n.effective is True for n in ir.nat_rules) == 9
    assert Counter(d.status for d in ir.extraction.nat_analysis.diagnostics) == {'FAIL': 2, 'WARN': 5, 'PASS': 3}


def test_upload_endpoint_returns_new_workbook_analysis():
    from fwmigrate.web import create_app
    response = create_app({'TESTING': True}).test_client().post('/api/extract/excel', data={
        'file': (io.BytesIO(FIXTURE.read_bytes()), 'synthetic.conf'), 'source_vendor': 'fortigate'})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.data))
    assert book['NAT Diagnostics'].max_row == 13
    assert book['NAT Traffic Coverage'].max_row == 10


def test_phase1_definition_resolves_tunnel_without_system_interface_entry():
    text = FIXTURE.read_text()
    start = text.index('    edit "VPN_HQ"')
    end = text.index('    next', start) + len('    next\n')
    ir = parse(text[:start] + text[end:])
    row = next(c for c in ir.extraction.nat_analysis.traffic_coverage if c.egress == 'VPN_HQ')
    assert row.nat_state == 'EXEMPT' and row.status == 'PASS'


def test_pool_source_mismatch_has_specific_diagnostic():
    result = analysis(FIXTURE.read_text().replace('set source-endip 10.10.40.255', 'set source-endip 10.10.40.9'))
    diagnostic = findings(result, 'NAT-011')[0]
    assert diagnostic.status == 'WARN' and 'only 10' in diagnostic.finding


def test_post_dnat_service_mismatch_has_specific_diagnostic():
    text = change_policy(FIXTURE.read_text(), 110, 'set service "ALL"', 'set service "HTTP"')
    diagnostic = findings(analysis(text), 'NAT-012')[0]
    assert diagnostic.status == 'FAIL' and 'VIP_ADMIN_RISKY' in diagnostic.objects


def test_api_and_file_nat_analysis_match_for_synthetic_fixture():
    from unittest.mock import patch
    from fwmigrate.parsers.fortigate.api_client import FortiGateAPIClient
    from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
    from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
    fg = parse_fortigate_config(FIXTURE.read_text())
    data = {'cmdb/system/global': [{'hostname': 'synthetic'}], 'cmdb/system/settings': [{'central-nat': 'disable'}]}
    for endpoint, objects in [
        ('cmdb/system/interface', fg.interfaces), ('cmdb/firewall/address', fg.addresses),
        ('cmdb/firewall/addrgrp', fg.address_groups), ('cmdb/firewall.service/custom', fg.services),
        ('cmdb/firewall.service/group', fg.service_groups), ('cmdb/vpn.ipsec/phase1-interface', fg.phase1_interfaces),
    ]:
        data[endpoint] = [{k.replace('_', '-'): v for k, v in o.model_dump(exclude={'source_attributes'}, exclude_unset=True).items()} for o in objects]
    for record in fg.extraction.objects:
        if record.section not in {'firewall policy', 'firewall ippool', 'firewall vip', 'firewall vipgrp'}:
            continue
        item = {k.replace('_', '-'): v for k, v in record.attributes.items()}
        if record.section == 'firewall policy':
            item['policyid'] = int(record.name)
            for key in ('srcintf', 'dstintf', 'srcaddr', 'dstaddr', 'service', 'poolname'):
                if key in item:
                    values = item[key] if isinstance(item[key], list) else [item[key]]
                    item[key] = [{'name': v} for v in values]
        else:
            item['name'] = record.name
        data.setdefault('cmdb/' + record.section.replace(' ', '/'), []).append(item)
    client = FortiGateAPIClient('synthetic.invalid', api_key='test-token')
    with patch.object(client, 'get', side_effect=lambda endpoint: data.get(endpoint, [])):
        api = FGToIRTransformer(client.extract_config()).transform()
    file = parse()
    assert [(n.name, n.effective, n.analysis_status) for n in api.nat_rules] == [(n.name, n.effective, n.analysis_status) for n in file.nat_rules]
    assert api.extraction.nat_analysis.traffic_coverage == file.extraction.nat_analysis.traffic_coverage
    assert api.extraction.nat_analysis.diagnostics == file.extraction.nat_analysis.diagnostics


def test_supplied_focus_fixture_optional_regression():
    """Opt-in local customer/test input: never embed private files in the repo."""
    import os
    path = os.environ.get('FWMIGRATE_NAT_TEST_CONFIG')
    if not path:
        pytest.skip('Set FWMIGRATE_NAT_TEST_CONFIG to the supplied FortiOS 7.4 focus fixture')
    ir = parse(Path(path).read_text(encoding='utf-8-sig'))
    result = ir.extraction.nat_analysis
    assert len(ir.nat_rules) == 12
    assert Counter(n.type for n in ir.nat_rules) == {NATType.SOURCE: 6, NATType.DESTINATION: 6}
    # The actual file has two post-DNAT service mismatches and one incomplete
    # deterministic source allocation; do not force the requested 9 effective.
    assert Counter(n.effective for n in ir.nat_rules) == {True: 6, False: 5, None: 1}
    assert Counter(c.status for c in result.traffic_coverage) == {'PASS': 4, 'WARN': 2, 'FAIL': 1}
    assert Counter(d.status for d in result.diagnostics) == {'FAIL': 4, 'WARN': 7, 'PASS': 3}
    assert len(findings(result, 'NAT-012')) == 2
    assert findings(result, 'NAT-011')[0].status == 'WARN'
    assert findings(result, 'NAT-008')[0].status == 'PASS'
    assert all('Unresolved service reference' not in note for o in ir.extraction.objects for note in o.notes)
