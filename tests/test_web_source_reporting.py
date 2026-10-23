import io
import json
from copy import deepcopy
from importlib import import_module
from pathlib import Path

import openpyxl
import pytest
from openpyxl import load_workbook

from fwmigrate.web import create_app, source_reporters
from fwmigrate.source_reporting.web_report import normalize_web_report
from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter


def test_malformed_secret_bearing_source_returns_safe_client_errors():
    client = create_app({"TESTING": True}).test_client()
    source = json.dumps({"format": "checkpoint-export-v1", "responses": [
        {"data": {"password": "ERROR_SECRET_SENTINEL"}},
    ]}).encode()
    for route in ("/api/preview", "/api/extract/excel"):
        response = client.post(route, data={"source_vendor": "checkpoint",
            "file": (io.BytesIO(source), "invalid.json")}, content_type="multipart/form-data")
        assert response.status_code == 400
        assert b"ERROR_SECRET_SENTINEL" not in response.data
        assert b"input_value" not in response.data


ASA_SOURCE = (
    "hostname asa\n"
    "username admin password 0 web-secret\n"
    "interface GigabitEthernet0/1\n"
    " nameif outside\n"
    " ip address 203.0.113.1 255.255.255.0\n"
    "object network WEB\n"
    " host 10.0.0.10\n"
    "access-list OUT extended permit tcp object WEB any eq 443\n"
)

PALO_SOURCE = '<config><shared><address><entry name="web"><ip-netmask>203.0.113.10</ip-netmask><future-password>palo-web-secret</future-password></entry></address></shared></config>'
CHECKPOINT_SOURCE = '{"format":"checkpoint-export-v1","domain":"SMC User","gateway":"CP-Enterprise-Gateway","selected_domain":"SMC User","selected_package":"Standard","selected_access_layer":"Network","selected_gateway":"CP-Enterprise-Gateway","responses":[{"command":"show-hosts","domain":"SMC User","data":{"objects":[{"name":"web","type":"host","ipv4-address":"203.0.113.10","password":"checkpoint-web-secret"}],"from":1,"to":1,"total":1}}]}'
FORTIGATE_SOURCE = """config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
    next
end
config firewall policy
    edit 1
        set name "web"
        set srcintf "port1"
        set dstintf "port1"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set schedule "always"
        set action accept
    next
end
config router static
    edit 1
        set dst 0.0.0.0 0.0.0.0
    next
end
"""

FORTIGATE_IPV6_POLICY_SOURCE = """config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
    next
end
config firewall address6
    edit "SRC-V6"
        set ip6 2001:db8:1::/64
    next
    edit "DST-V6"
        set ip6 2001:db8:2::/64
    next
end
config firewall policy
    edit 60
        set srcintf "port1"
        set dstintf "port1"
        set srcaddr6 "SRC-V6"
        set dstaddr6 "DST-V6"
        set service "ALL"
        set schedule always
        set action accept
    next
end
"""


EXCEL_SOURCE_CASES = (
    ("fortigate", FORTIGATE_SOURCE), ("palo_alto", PALO_SOURCE), ("cisco_asa", ASA_SOURCE),
    ("juniper_srx", "set system host-name branch\nset interfaces ge-0/0/0 unit 0 family inet address 203.0.113.10/24"),
    ("cisco_ftd", json.dumps({"format": "cisco-fmc-rest-export-v1", "source": "fmc-rest-api",
                             "domain": {"id": "d1", "name": "Global"}, "objects": {"hosts": [
                                 {"id": "h1", "name": "Host", "value": "203.0.113.10"},
                             ]}, "access_policies": [], "nat_policies": []})),
    ("checkpoint", CHECKPOINT_SOURCE),
)


@pytest.mark.parametrize("profile", ["fast", "full"])
@pytest.mark.parametrize("fixture", ["fmc_selected_domains.json", "fdm_object_pipeline_conformance.json",
                                     "fdm_nat_pipeline_conformance.json", "cli"])
def test_ftd_supported_formats_reach_preview_and_excel_through_all_transports(profile, fixture):
    from fwmigrate.collection.contracts import CollectedSource
    from fwmigrate.collection.snapshot import make_snapshot

    if fixture == "cli":
        source = ("configure network ipv4 manual 192.0.2.42 255.255.255.0 192.0.2.1\n"
                  "interface inside\n description WEB_CLI_MARKER\n ipv6 address 2001:db8::42/64\n"
                  "access-list WEB_UNSUPPORTED extended permit ip any any\n password WEB_FTD_SECRET\n")
        markers = ["WEB_CLI_MARKER", "192.0.2.42", "2001:db8::42", "WEB_UNSUPPORTED"]
    else:
        payload = json.loads((Path(__file__).parent / "fixtures/cisco_ftd" / fixture).read_text())
        objects = payload["objects"]
        hosts = objects["networkaddresses"] if "networkaddresses" in objects else objects["hosts"]
        hosts[0]["customField"] = "WEB_CUSTOM_MARKER"
        hosts[0]["apiToken"] = "WEB_FTD_SECRET"
        markers = ["WEB_CUSTOM_MARKER"]
        source = json.dumps(payload)
    client = create_app({"TESTING": True}).test_client()
    preview = client.post("/api/preview", data={"source_vendor": "cisco_ftd",
        "file": (io.BytesIO(source.encode()), "ftd-source.txt")})
    assert preview.status_code == 200, preview.get_json()
    evidence = preview.get_json()["source_evidence"]
    assert "WEB_FTD_SECRET" not in str(preview.get_json())
    snapshot = make_snapshot(CollectedSource("cisco_ftd", source, "ftd-source.txt", "offline"))
    exports = [client.post("/api/extract/excel", buffered=True, data={"source_vendor": "cisco_ftd", "excel_profile": profile,
        "file": (io.BytesIO(source.encode()), "ftd-source.txt")})]
    for item in (evidence, {"vendor": "cisco_ftd", "source_text": source}, snapshot):
        exports.append(client.post("/api/extract/excel", buffered=True, json={"source_vendor": "cisco_ftd",
            "excel_profile": profile, "source": item}))
    for response in exports:
        assert response.status_code == 200, response.get_json()
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        try:
            values = str([row for sheet in workbook for row in sheet.values])
            assert all(marker in values for marker in markers)
            assert "WEB_FTD_SECRET" not in values
        finally:
            workbook.close()


@pytest.mark.parametrize("source", ['{"format":"other"}', '{"format":\ninterface inside',
                                     '[\ninterface inside', "unrecognized source"])
def test_ftd_web_rejects_invalid_source_without_cli_fallback(source):
    client = create_app({"TESTING": True}).test_client()
    for route in ("/api/preview", "/api/extract/excel"):
        response = client.post(route, data={"source_vendor": "cisco_ftd",
            "file": (io.BytesIO(source.encode()), "invalid.json")})
        assert response.status_code == 400


@pytest.mark.parametrize('profile', ['fast', 'full'])
@pytest.mark.parametrize('vendor,source', EXCEL_SOURCE_CASES, ids=[case[0] for case in EXCEL_SOURCE_CASES])
def test_all_vendor_exports_preserve_oversized_sanitized_evidence(profile, vendor, source):
    warning = 'x' * 32767 + '=retained-tail\x01 password overflow-secret'
    reporter = source_reporters.get(vendor)
    output = io.BytesIO()
    reporter.export_excel(reporter.analyze_source(source), output, profile=profile,
                          collection_status='PARTIAL', collection_warnings=[warning])
    workbook = load_workbook(output)
    rows = [row for row in list(workbook['Text Overflow'].values)[1:] if row[0] == 'Collection Evidence']
    reconstructed = ''.join(row[4] for row in rows)
    assert reconstructed == 'x' * 32767 + '=retained-tail\\u0001 password [REDACTED]'
    assert 'overflow-secret' not in str([row for sheet in workbook for row in sheet.values])
    assert all(cell.data_type != 'f' for row in workbook['Text Overflow'] for cell in row)


@pytest.mark.parametrize('profile', ['fast', 'full'])
def test_fortigate_excel_keeps_control_characters_and_all_large_group_members(profile):
    names = [f'member-{index:04}-' + 'x' * 48 for index in range(600)]
    source = ('config firewall address\n' + ''.join(
        f'edit {name}\nset subnet 192.0.2.1 255.255.255.255\nnext\n' for name in names)
        + 'end\nconfig firewall addrgrp\nedit large\nset comment "visible\x01é"\nset member '
        + ' '.join(names) + '\nnext\nend\n')
    client = create_app({'TESTING': True}).test_client()
    response = client.post('/api/extract/excel', data={'source_vendor': 'fortigate', 'excel_profile': profile,
        'file': (io.BytesIO(source.encode()), 'large.conf')}, content_type='multipart/form-data')
    assert response.status_code == 200, response.get_json()
    workbook = load_workbook(io.BytesIO(response.data))
    rows = list(workbook['Text Overflow'].values)[1:]
    members = ''.join(row[4] for row in rows if row[:2] == ('Address Groups', 'D4'))
    assert members.splitlines() == names
    assert any(cell.value == 'visible\\u0001é' for row in workbook['Address Groups'] for cell in row)
    if profile == 'full':
        inventory = ''.join(row[4] for row in rows if row[0] == 'FortiGate Source Inventory')
        assert all(name in inventory for name in names)


@pytest.mark.parametrize("profile", ("fast", "full"))
@pytest.mark.parametrize("vendor,source", EXCEL_SOURCE_CASES)
@pytest.mark.parametrize("status", ("PARTIAL", "SUCCESS"))
def test_snapshot_excel_preserves_collection_evidence(profile, vendor, source, status, monkeypatch):
    from fwmigrate.collection.snapshot import FORMAT

    client = create_app({"TESTING": True}).test_client()
    reporter = source_reporters.get(vendor)
    analyze = reporter.analyze_source
    analyzed = []

    def capture_analysis(source, **options):
        analysis = analyze(source, **options)
        analyzed.append((analysis, deepcopy(analysis)))
        return analysis

    monkeypatch.setattr(reporter, "analyze_source", capture_analysis)
    parts = [
        {"name": "config", "status": "SUCCESS", "complete": True, "count": 1},
        {"name": "empty", "status": "EMPTY", "complete": True, "count": 0},
    ]
    if status == "PARTIAL":
        parts.append({"name": "unavailable", "status": "PERMISSION_DENIED", "complete": False, "count": None})
    snapshot = {
        "format": FORMAT, "vendor_id": vendor, "source_name": "snapshot.conf",
        "collection_method": "test", "status": status, "source_text": source,
        "parts": parts, "warnings": ["password evidence-secret", "=1+1", "+1+1", "@SUM(1)"],
    }
    before = deepcopy(snapshot)
    response = client.post("/api/extract/excel", json={
        "source_vendor": vendor, "source": snapshot, "excel_profile": profile,
    })
    assert response.status_code == 200
    assert snapshot == before
    assert len(analyzed) == 1
    assert analyzed[0][0] == analyzed[0][1]
    workbook = load_workbook(io.BytesIO(response.data))
    try:
        summary = {row[0]: row[1] for row in workbook["Summary"].values if row[0]}
        assert any(row[index] == "Source File" and row[index + 1] == "snapshot.conf"
                   for row in workbook["Summary"].values for index in (0, 2, 4))
        assert summary["Collection Status"] == status
        if status == "PARTIAL":
            assert "Missing configuration is unknown" in summary["Collection Completeness"]
        evidence = workbook["Collection Evidence"]
        rows = list(evidence.values)
        assert rows[0] == ("Part Name", "Status", "Complete", "Count", "Warnings")
        assert [row[:4] for row in rows[1:len(parts) + 1]] == [
            (part["name"], part["status"], part["complete"], part["count"]) for part in parts
        ]
        assert [row[4] for row in rows[len(parts) + 1:]] == [
            "password [REDACTED]", "'=1+1", "'+1+1", "'@SUM(1)",
        ]
        assert all(cell.data_type != "f" for row in evidence for cell in row)
        values = str([row for sheet in workbook for row in sheet.values])
        assert "evidence-secret" not in values
        assert "palo-web-secret" not in values
        assert "web-secret" not in values
        address = {"fortigate": "192.0.2.1", "cisco_asa": "203.0.113.1"}.get(vendor, "203.0.113.10")
        assert address in values
    finally:
        workbook.close()


@pytest.mark.parametrize("vendor,source", EXCEL_SOURCE_CASES)
@pytest.mark.parametrize("profile", ("fast", "full"))
def test_excel_rejects_failed_collection_snapshot(vendor, source, profile):
    from fwmigrate.collection.snapshot import FORMAT

    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/extract/excel", json={
        "source_vendor": vendor, "excel_profile": profile,
        "source": {"format": FORMAT, "vendor_id": vendor, "status": "FAILED", "source_text": source},
    })
    assert response.status_code == 400
    assert response.get_json()["error"] == "Invalid source configuration"


@pytest.mark.parametrize("vendor,source", EXCEL_SOURCE_CASES)
@pytest.mark.parametrize("profile", ("fast", "full"))
def test_uploaded_excel_preserves_filename_without_inventing_collection(vendor, source, profile):
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/extract/excel", data={
        "source_vendor": vendor, "excel_profile": profile,
        "file": (io.BytesIO(source.encode()), "=source.conf"),
    })
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    try:
        assert "Collection Evidence" not in workbook.sheetnames
        assert any(row[index] == "Source File" and row[index + 1] == "'=source.conf"
                   for row in workbook["Summary"].values for index in (0, 2, 4))
        assert all(cell.data_type != "f" for row in workbook["Summary"] for cell in row)
    finally:
        workbook.close()


def test_panos_fast_does_not_build_source_appendix(monkeypatch):
    from fwmigrate.vendors.palo_alto.export.excel_rows import ROW_BUILDERS

    def unexpected_appendix(context):
        pytest.fail("FAST must skip XML appendix processing")

    monkeypatch.setitem(ROW_BUILDERS, "PAN-OS Source Appendix", unexpected_appendix)
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/extract/excel", data={
        "source_vendor": "palo_alto", "excel_profile": "fast",
        "file": (io.BytesIO(PALO_SOURCE.encode()), "pan.xml"),
    })
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    try:
        summary = str(list(workbook["Summary"].values))
        for name in ("PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"):
            assert name not in workbook.sheetnames
            assert name in summary
    finally:
        workbook.close()


def test_panos_excel_defaults_to_full_and_explicit_fast_is_preserved():
    client = create_app({"TESTING": True}).test_client()
    full = client.post("/api/extract/excel", buffered=True, data={
        "source_vendor": "palo_alto", "file": (io.BytesIO(PALO_SOURCE.encode()), "pan.xml"),
    })
    fast = client.post("/api/extract/excel", buffered=True, data={
        "source_vendor": "palo_alto", "excel_profile": "fast",
        "file": (io.BytesIO(PALO_SOURCE.encode()), "pan.xml"),
    })
    other = client.post("/api/extract/excel", buffered=True, data={
        "source_vendor": "cisco_asa", "file": (io.BytesIO(ASA_SOURCE.encode()), "asa.cfg"),
    })
    assert full.status_code == fast.status_code == other.status_code == 200
    full_book, fast_book, other_book = [load_workbook(io.BytesIO(item.data), read_only=True) for item in (full, fast, other)]
    try:
        evidence_sheets = {"PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"}
        assert evidence_sheets <= set(full_book.sheetnames)
        assert evidence_sheets.isdisjoint(fast_book.sheetnames)
        assert other_book["Summary"]["A1"].value == "Cisco ASA Configuration Report"
    finally:
        full_book.close()
        fast_book.close()
        other_book.close()


def _post(client, path, **fields):
    data = {"source_vendor": "cisco_asa", **fields}
    if path == "/api/preview" or "file" not in data:
        data["file"] = (io.BytesIO(ASA_SOURCE.encode()), "asa.cfg")
    return client.post(path, data=data, content_type="multipart/form-data")


def test_source_preview_and_excel_use_opaque_cached_asa_analysis():
    client = create_app({"TESTING": True}).test_client()

    preview = _post(client, "/api/preview")
    assert preview.status_code == 200
    payload = preview.get_json()
    assert payload["vendor"] == "cisco_asa"
    assert payload["summary"]["interfaces"] == 1
    assert "interfaces" in payload["sections"]
    assert "vpn_phase2" in payload["sections"]
    assert "canonical_ir" not in str(payload)

    workbook = client.post("/api/extract/excel", json={"source_vendor": "cisco_asa", "source": payload["source_evidence"]})
    assert workbook.status_code == 200
    assert workbook.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sheets = load_workbook(io.BytesIO(workbook.data), read_only=True).sheetnames
    assert "Interfaces" in sheets
    assert "web-secret" not in workbook.data.decode("latin1", errors="ignore")


def test_source_is_reanalyzed_for_each_request_and_survives_restart(monkeypatch):
    client = create_app({'TESTING': True}).test_client()
    reporter = source_reporters.get('fortigate')
    original = reporter.analyze_source
    calls = []
    def analyze(source, **options):
        calls.append(source)
        return original(source, **options)
    monkeypatch.setattr(reporter, 'analyze_source', analyze)
    preview = client.post('/api/preview', data={'file': (io.BytesIO(FORTIGATE_SOURCE.encode()), 'source.conf')}).get_json()
    assert 'preview_id' not in preview
    restarted = create_app({'TESTING': True}).test_client()
    response = restarted.post('/api/extract/excel', json={'source_vendor': 'fortigate', 'source': preview['source_evidence']})
    assert response.status_code == 200
    assert len(calls) == 2
    assert restarted.post('/api/migration/requirements', json={'source': preview['source_evidence']}).status_code == 200


def test_source_excel_rejects_missing_input_and_unknown_vendor():
    client = create_app({"TESTING": True}).test_client()

    missing = client.post("/api/extract/excel", data={"source_vendor": "cisco_asa"})
    assert missing.status_code == 400

    unknown = _post(client, "/api/preview", source_vendor="unknown")
    assert unknown.status_code == 400


@pytest.mark.parametrize('route', ['/api/preview', '/api/extract/excel'])
def test_source_decode_errors_do_not_expose_exception_text(monkeypatch, route):
    import fwmigrate.web_api.source as source_api
    from fwmigrate.web_support.reporting import ConfigurationDecodeError

    monkeypatch.setattr(source_api, 'analyze_request_source',
                        lambda *args, **kwargs: (_ for _ in ()).throw(ConfigurationDecodeError('SYNTHETIC_DECODE_SECRET')))
    client = create_app({'TESTING': True}).test_client()
    response = client.post(route, data={'source_vendor': 'fortigate',
        'file': (io.BytesIO(b'valid input'), 'source.conf')})
    assert response.status_code == 400
    payload = response.get_json()
    assert payload['error'] == 'Configuration file is not valid UTF-8.'
    assert payload['stage'] == 'decode'
    assert 'SYNTHETIC_DECODE_SECRET' not in response.get_data(as_text=True)


def test_excel_unavailable_error_does_not_expose_exception_text(monkeypatch):
    import fwmigrate.web_api.source as source_api
    from fwmigrate.source_reporting import ExcelExportUnavailableError

    reporter = source_api.source_reporters.get('fortigate')
    monkeypatch.setattr(reporter, 'export_excel',
                        lambda *args, **kwargs: (_ for _ in ()).throw(ExcelExportUnavailableError('SYNTHETIC_EXPORT_SECRET')))
    response = create_app({'TESTING': True}).test_client().post('/api/extract/excel', data={
        'source_vendor': 'fortigate', 'file': (io.BytesIO(FORTIGATE_SOURCE.encode()), 'source.conf'),
    })
    assert response.status_code == 503
    assert response.get_json()['error'] == 'Excel export is unavailable for this source.'
    assert 'SYNTHETIC_EXPORT_SECRET' not in response.get_data(as_text=True)


def test_source_excel_vendor_mismatch_does_not_use_cached_preview():
    client = create_app({"TESTING": True}).test_client()
    preview = _post(client, "/api/preview").get_json()

    response = client.post(
        "/api/extract/excel",
        data={"source_vendor": "unknown", "source": preview["source_evidence"]},
        content_type="multipart/form-data",
    )
    assert response.status_code == 422


def test_palo_alto_upload_cached_preview_and_excel_redact_secrets():
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/preview", data={"source_vendor": "palo_alto", "file": (io.BytesIO(PALO_SOURCE.encode()), "pan.xml")}, content_type="multipart/form-data")
    assert response.status_code == 200
    payload = response.get_json()
    workbook = client.post("/api/extract/excel", json={"source_vendor": "palo_alto", "source": payload["source_evidence"]})
    assert workbook.status_code == 200
    assert b"palo-web-secret" not in response.data + workbook.data


@pytest.mark.parametrize('profile', ('fast', 'full'))
def test_palo_alto_excel_preserves_secret_like_object_names_after_upload(profile):
    client = create_app({'TESTING': True}).test_client()
    source = b'''<config><password>xml-secret</password><shared><address>
      <entry name="token object"><ip-netmask>192.0.2.1</ip-netmask></entry>
    </address></shared></config>'''
    preview = client.post('/api/preview', data={
        'source_vendor': 'palo_alto', 'file': (io.BytesIO(source), 'pan.xml'),
    })
    assert preview.status_code == 200
    evidence = preview.get_json()['source_evidence']
    response = client.post('/api/extract/excel', json={
        'source_vendor': 'palo_alto', 'source': evidence, 'excel_profile': profile,
    })
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        names = [row[0] for row in workbook['Addresses'].iter_rows(min_row=4, values_only=True)]
        assert names == ['token object']
        values = str([row for sheet in workbook for row in sheet.values])
        assert 'xml-secret' not in values + str(evidence)
    finally:
        workbook.close()


def test_check_point_upload_cached_preview_and_excel_redact_secrets():
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/preview", data={"source_vendor": "checkpoint", "file": (io.BytesIO(CHECKPOINT_SOURCE.encode()), "cp.json")}, content_type="multipart/form-data")
    assert response.status_code == 200
    payload = response.get_json()
    workbook = client.post("/api/extract/excel", json={"source_vendor": "checkpoint", "source": payload["source_evidence"]})
    assert workbook.status_code == 200
    assert b"checkpoint-web-secret" not in response.data + workbook.data


def test_fortigate_preview_exposes_policy_fields_and_overview_sections():
    client = create_app({"TESTING": True}).test_client()
    response = client.post(
        "/api/preview",
        data={"source_vendor": "fortigate", "file": (io.BytesIO(FORTIGATE_SOURCE.encode()), "fortigate.conf")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    report = response.get_json()
    policy = report["sections"]["policies"][0]
    assert policy["source_addresses"] == ["all"]
    assert policy["destination_addresses"] == ["all"]
    assert policy["schedule"] == "always"
    assert len(report["sections"]["routes"]) == 1
    assert report["sections"]["vpn_tunnels"] == []
    assert report["sections"]["vpn_phase2"] == []
    assert report["summary"]["unsupported_count"] == report["summary"]["source_only_count"]
    assert report["summary"]["records"] == len(report["sections"]["source_inventory"])


def test_fortigate_preview_preserves_ipv6_policy_addresses():
    source_policy = FortiGateSourceReporter().analyze_source(FORTIGATE_IPV6_POLICY_SOURCE).extracted.config.policies[0]
    assert source_policy.srcaddr is None
    assert source_policy.dstaddr is None
    client = create_app({"TESTING": True}).test_client()
    response = client.post(
        "/api/preview",
        data={"source_vendor": "fortigate", "file": (io.BytesIO(FORTIGATE_IPV6_POLICY_SOURCE.encode()), "fortigate.conf")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    policy = response.get_json()["sections"]["policies"][0]
    assert policy["source_addresses"] is None
    assert policy["source_addresses_ipv6"] == ["SRC-V6"]
    assert policy["destination_addresses"] is None
    assert policy["destination_addresses_ipv6"] == ["DST-V6"]


def test_all_registered_sources_advertise_view_report():
    client = create_app({"TESTING": True}).test_client()
    payload = client.get("/api/vendors").get_json()
    sources = payload["sources"]
    assert "targets" not in payload
    assert {item["vendor_id"] for item in sources} >= {
        "fortigate", "palo_alto", "cisco_asa", "cisco_ftd", "checkpoint", "juniper_srx",
    }
    assert all(item["web_report"] is True for item in sources)


def test_removed_source_diagnostics_endpoint_is_not_exposed():
    client = create_app({"TESTING": True}).test_client()

    assert client.post("/api/diagnostics").status_code == 404


@pytest.mark.parametrize("transport", ("upload", "preview"))
@pytest.mark.parametrize(("vendor", "fixture", "excluded"), (
    ("fortigate", "example_fortigate.conf", {"FortiGate Source Inventory", "Extraction Coverage"}),
    ("palo_alto", "example_palo_alto.xml", {"PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"}),
    ("cisco_asa", "example_cisco_asa.cfg", {"Source Inventory", "Extraction Coverage"}),
    ("cisco_ftd", "cisco_ftd/fmc_selected_domains.json", set()),
    ("checkpoint", "checkpoint/multidomain_full.json", set()),
    ("juniper_srx", "example_juniper_srx.set", {"Source Inventory", "Extraction Coverage"}),
))
def test_excel_selector_reaches_real_vendor_profiles(vendor, fixture, excluded, transport, monkeypatch):
    client = create_app({"TESTING": True}).test_client()
    path = Path(__file__).parent / "fixtures" / fixture
    source = path.read_bytes()
    evidence = None
    if transport == "preview":
        preview = client.post("/api/preview", data={
            "source_vendor": vendor, "file": (io.BytesIO(source), path.name),
        }, content_type="multipart/form-data")
        assert preview.status_code == 200
        evidence = preview.get_json()["source_evidence"]
    original = openpyxl.Workbook
    write_modes = []

    def workbook_factory(*args, **kwargs):
        workbook = original(*args, **kwargs)
        write_modes.append(workbook.write_only)
        return workbook

    monkeypatch.setattr(openpyxl, "Workbook", workbook_factory)
    module = import_module(f"fwmigrate.vendors.{vendor}.export.excel")
    if hasattr(module, "Workbook"):
        monkeypatch.setattr(module, "Workbook", workbook_factory)
    workbooks = {}
    for profile in ("fast", "full"):
        if transport == "preview":
            response = client.post("/api/extract/excel", json={
                "source_vendor": vendor, "excel_profile": profile, "source": evidence,
            })
        else:
            response = client.post("/api/extract/excel", data={
                "source_vendor": vendor, "excel_profile": profile,
                "file": (io.BytesIO(source), path.name),
            }, content_type="multipart/form-data")
        assert response.status_code == 200, (vendor, profile, response.get_json())
        workbooks[profile] = load_workbook(io.BytesIO(response.data))
    assert write_modes == [True, vendor == "palo_alto"]
    fast, full = workbooks["fast"], workbooks["full"]
    if vendor in {"fortigate", "palo_alto"}:
        for workbook in (fast, full):
            assert "Collection Evidence" not in workbook.sheetnames
            assert "Collection Status" not in str(list(workbook["Summary"].values))
        assert "CONFIGURATION INVENTORY" in {row[0] for row in full["Summary"].values}
    if vendor == "palo_alto":
        for workbook in (fast, full):
            assert "Review Required" in workbook.sheetnames
            assert "Validation" not in workbook.sheetnames
    assert excluded <= set(full.sheetnames)
    assert excluded.isdisjoint(fast.sheetnames)
    assert set(fast.sheetnames) <= set(full.sheetnames)
    assert fast.sheetnames[0] == full.sheetnames[0] == "Summary"
    header_row = 3 if vendor in {"fortigate", "palo_alto"} else 1
    for name in fast.sheetnames:
        if name == "Summary":
            continue
        fast_sheet, full_sheet = fast[name], full[name]
        assert fast_sheet.sheet_view.showGridLines is not False
        assert full_sheet.sheet_view.showGridLines is not False
        fast_rows = list(fast_sheet.iter_rows(min_row=header_row, values_only=True))
        full_rows = list(full_sheet.iter_rows(min_row=header_row, values_only=True))
        if vendor == "fortigate" and name == "Unsupported":
            # FortiGate redirects its inventory reference in FAST exports.
            fast_rows = [row[:-1] for row in fast_rows]
            full_rows = [row[:-1] for row in full_rows]
        assert fast_rows == full_rows, (vendor, name)
        assert fast_sheet.freeze_panes == full_sheet.freeze_panes
        if fast_sheet.max_row > header_row:
            assert fast_sheet.auto_filter.ref == full_sheet.auto_filter.ref, (vendor, name)
        for row in fast_sheet.iter_rows(min_row=header_row + 1):
            for cell in row:
                if vendor in {"fortigate", "palo_alto"} and cell.has_style:
                    assert cell.alignment.wrap_text and cell.alignment.vertical == "top"
                else:
                    assert not cell.has_style
        for column, dimension in full_sheet.column_dimensions.items():
            if dimension.hidden:
                assert fast_sheet.column_dimensions[column].hidden, (vendor, name, column)
    assert any(cell.hyperlink is not None for row in fast["Summary"] for cell in row)
    assert any(cell.hyperlink is not None for row in full["Summary"] for cell in row)


@pytest.mark.parametrize("vendor", ("fortigate", "palo_alto", "cisco_asa", "cisco_ftd", "checkpoint", "juniper_srx"))
def test_excel_selector_rejects_invalid_profiles_for_every_vendor(vendor):
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/extract/excel", json={"source_vendor": vendor, "excel_profile": "invalid"})
    assert response.status_code == 400
    assert response.get_json()["error"] == "excel_profile must be one of: fast, full"


def test_api_preview_contract_covers_every_registered_vendor():
    fixtures = Path(__file__).parent / "fixtures"
    cases = (
        ("fortigate", fixtures / "example_fortigate.conf"),
        ("palo_alto", fixtures / "example_palo_alto.xml"),
        ("cisco_asa", fixtures / "example_cisco_asa.cfg"),
        ("cisco_ftd", fixtures / "cisco_ftd" / "fmc_selected_domains.json"),
        ("checkpoint", fixtures / "checkpoint" / "multidomain_full.json"),
        ("juniper_srx", fixtures / "example_juniper_srx.set"),
    )
    required_sections = {"interfaces", "addresses", "address_groups", "services", "service_groups",
                         "schedules", "policies", "nat", "routes", "vpn_tunnels", "vpn_phase2",
                         "validation", "unresolved_references"}
    client = create_app({"TESTING": True}).test_client()
    for vendor, path in cases:
        response = client.post(
            "/api/preview",
            data={"source_vendor": vendor, "file": (io.BytesIO(path.read_bytes()), path.name)},
            content_type="multipart/form-data",
        )
        assert response.status_code == 200, (vendor, response.get_json())
        report = response.get_json()
        assert report["vendor"] == vendor
        assert isinstance(report["summary"]["objects"], dict)
        assert isinstance(report["summary"]["validation"], dict)
        assert required_sections <= set(report["sections"])
        assert all(isinstance(report["sections"][name], list) for name in required_sections)
        json.dumps(report)
        reporter = source_reporters.get(vendor)
        analysis = reporter.analyze_source(path.read_text(encoding="utf-8"))
        before = deepcopy(analysis)
        copied = normalize_web_report(reporter.build_preview(analysis), vendor)
        owned = normalize_web_report(reporter.build_preview(analysis), vendor, copy=False)
        assert owned == copied, vendor
        assert json.loads(json.dumps(owned)) == {key: value for key, value in report.items() if key not in {"success", "source_evidence", "source_digest"}}, vendor
        assert analysis == before, vendor


def test_source_reporting_rejects_oversized_upload_and_inline_evidence(monkeypatch):
    request_source = import_module("fwmigrate.web_support.request_source")
    monkeypatch.setattr(request_source, "MAX_BYTES", 64)
    client = create_app({"TESTING": True}).test_client()
    oversized = b"<config>" + (b"x" * 80) + b"</config>"

    for route in ("/api/preview", "/api/extract/excel"):
        response = client.post(
            route,
            data={
                "source_vendor": "palo_alto",
                "file": (io.BytesIO(oversized), "oversized.xml"),
            },
            content_type="multipart/form-data",
        )
        assert response.status_code == 413

    inline = client.post(
        "/api/extract/excel",
        json={
            "source_vendor": "palo_alto",
            "source": {
                "vendor": "palo_alto",
                "source_text": "x" * 65,
            },
        },
    )
    assert inline.status_code == 413


@pytest.mark.parametrize('route', ['/api/preview', '/api/extract/excel'])
@pytest.mark.parametrize('exception_type', [ValueError, KeyError, RuntimeError])
def test_source_failures_keep_synthetic_secrets_private(monkeypatch, caplog, route, exception_type):
    import fwmigrate.web_api.source as source_routes

    def fail(*args, **kwargs):
        raise exception_type('SYNTHETIC_SOURCE_SECRET')

    monkeypatch.setattr(source_routes, 'analyze_request_source', fail)
    response = create_app({'TESTING': True}).test_client().post(route, data={
        'file': (io.BytesIO(b'password SYNTHETIC_SOURCE_SECRET'), 'source.conf'),
    })
    assert response.status_code in {400, 422, 500}
    assert 'SYNTHETIC_SOURCE_SECRET' not in response.get_data(as_text=True) + caplog.text


def test_fortigate_mixed_case_secrets_are_safe_in_snapshot_preview_excel_and_logs(caplog):
    from fwmigrate.collection.contracts import CollectedSource, CollectionStatus
    from fwmigrate.collection.snapshot import make_snapshot, parse_snapshot
    source = ('config system admin\nedit admin\nSET password ENC UPPER_SECRET_SENTINEL\n'
              'ApPeNd private-key "MULTILINE_SECRET_SENTINEL\nTAIL_SECRET_SENTINEL"\n'
              'unset password\nnext\nend\n')
    secrets = ('UPPER_SECRET_SENTINEL', 'MULTILINE_SECRET_SENTINEL', 'TAIL_SECRET_SENTINEL')
    snapshot = make_snapshot(CollectedSource('fortigate', source, 'safe.conf', 'test', CollectionStatus.SUCCESS))
    assert all(secret not in json.dumps(snapshot) for secret in secrets)
    imported = parse_snapshot(json.dumps(snapshot).encode())
    assert all(secret not in imported.source_text for secret in secrets)
    client = create_app({'TESTING': True}).test_client()
    for route in ('/api/preview', '/api/extract/excel'):
        response = client.post(route, data={'source_vendor': 'fortigate',
            'file': (io.BytesIO(json.dumps(snapshot).encode()), 'safe.json')})
        assert response.status_code == 200
        if route.endswith('excel'):
            workbook = load_workbook(io.BytesIO(response.data))
            contents = repr([list(sheet.values) for sheet in workbook])
        else:
            contents = response.get_data(as_text=True)
        assert all(secret not in contents for secret in secrets)
    assert all(secret not in caplog.text for secret in secrets)
