import io
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.parsers.palo_alto import PANOSSourceParser
from fwmigrate.report.excel_exporter import IRExcelExporter


FIXTURES = Path(__file__).parent / "fixtures" / "palo_alto"


def test_pan_fidelity_fields_survive_parse_and_excel_export():
    parser = PANOSSourceParser()

    interfaces = parser.extract((FIXTURES / "interfaces_extended.xml").read_text())
    interface = next(item for item in interfaces.canonical_ir.interfaces if item.name == "ethernet1/1")
    assert [item.address for item in interface.additional_ipv4_addresses] == ["192.0.2.10/24"]
    interface_workbook = load_workbook(io.BytesIO(IRExcelExporter(interfaces.canonical_ir).generate()))
    interface_headers = {cell.value: cell.column for cell in interface_workbook["Interfaces"][3]}
    interface_row = next(
        row for row in interface_workbook["Interfaces"].iter_rows(min_row=4)
        if row[interface_headers["Name"] - 1].value == "ethernet1/1"
    )
    assert interface_row[interface_headers["Additional IPv4 Addresses"] - 1].value == "192.0.2.10/24"

    conformance = parser.extract((FIXTURES / "phase13_conformance.xml").read_text())
    route = next(item for item in conformance.canonical_ir.routes if item.name == "phase13-default")
    pbf = next(item for item in conformance.canonical_ir.pbf_rules if item.name == "phase13-pbf")
    nat = next(item for item in conformance.canonical_ir.nat_rules if item.name == "phase13-snat")
    assert route.next_hop_type.value == "ip-address"
    assert pbf.next_hop_type.value == "ip-address"
    assert nat.source_translation_mode.value == "dynamic-ip-and-port"

    fallback_xml = (FIXTURES / "phase13_conformance.xml").read_text().replace(
        "<interface-address><interface>ethernet1/2</interface></interface-address>",
        "<fallback><interface-address><interface>ethernet1/2</interface></interface-address></fallback>",
    )
    fallback = parser.extract(fallback_xml)
    fallback_nat = fallback.canonical_ir.nat_rules[0]
    assert fallback_nat.source_translation_mode.value == "dynamic-ip-and-port"
    assert fallback_nat.source_translation_fallback.interface == "ethernet1/2"

    families = parser.extract((FIXTURES / "policy_families.xml").read_text())
    sdwan = next(item for item in families.canonical_ir.pan_sdwan_rules if item.name == "sdwan-rule")
    assert sdwan.path_quality_profile == "quality"
    assert sdwan.traffic_distribution_profile == "balanced"

    workbook = load_workbook(io.BytesIO(IRExcelExporter(conformance.canonical_ir).generate()))
    assert {"Routes", "PBF Rules", "NAT Rules"}.issubset(workbook.sheetnames)
    route_headers = {cell.value: cell.column for cell in workbook["Routes"][3]}
    route_row = next(row for row in workbook["Routes"].iter_rows(min_row=4) if row[route_headers["Name"] - 1].value == "phase13-default")
    assert route_row[route_headers["Next Hop Type"] - 1].value == "ip-address"
    pbf_headers = {cell.value: cell.column for cell in workbook["PBF Rules"][3]}
    pbf_row = next(row for row in workbook["PBF Rules"].iter_rows(min_row=4) if row[pbf_headers["Name"] - 1].value == "phase13-pbf")
    assert pbf_row[pbf_headers["Next Hop"] - 1].value == "198.51.100.253"

    fallback_workbook = load_workbook(io.BytesIO(IRExcelExporter(fallback.canonical_ir).generate()))
    nat_headers = {cell.value: cell.column for cell in fallback_workbook["NAT Rules"][3]}
    nat_row = next(row for row in fallback_workbook["NAT Rules"].iter_rows(min_row=4) if row[nat_headers["Name"] - 1].value == "phase13-snat")
    assert "ethernet1/2" in nat_row[nat_headers["Source Translation Fallback"] - 1].value

    sdwan_workbook = load_workbook(io.BytesIO(IRExcelExporter(families.canonical_ir).generate()))
    sdwan_headers = {cell.value: cell.column for cell in sdwan_workbook["PAN SD-WAN Rules"][3]}
    sdwan_row = next(row for row in sdwan_workbook["PAN SD-WAN Rules"].iter_rows(min_row=4) if row[sdwan_headers["Name"] - 1].value == "sdwan-rule")
    assert sdwan_row[sdwan_headers["Path Quality Profile"] - 1].value == "quality"
