from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_host_inbound_exclusions_survive_ir_and_excel():
    parser = JuniperSRXParser("""
    set security zones security-zone trust interfaces ge-0/0/0.0
    set security zones security-zone trust host-inbound-traffic system-services all except ssh
    set security zones security-zone trust host-inbound-traffic protocols all except ospf
    set security zones security-zone trust interfaces ge-0/0/0.0 host-inbound-traffic system-services all except telnet
    """)
    ir = parser.transform_to_ir()

    zone_rules = {
        rule.source_attributes["junos_host_inbound_kind"]: rule
        for rule in ir.local_in_policies
        if rule.interface is None
    }
    assert zone_rules["system-services"].service == ["all"]
    assert zone_rules["system-services"].service_exclusions == ["ssh"]
    assert zone_rules["protocols"].protocol == "all"
    assert zone_rules["protocols"].protocol_exclusions == ["ospf"]

    interface_rule = next(rule for rule in ir.local_in_policies if rule.interface == "ge-0/0/0.0")
    assert interface_rule.service_exclusions == ["telnet"]

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["Local-In Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(row for row in sheet.iter_rows(min_row=4, values_only=False) if row[headers["Interface"] - 1].value is None)
    assert row[headers["Service Exclusions"] - 1].value == "ssh"
