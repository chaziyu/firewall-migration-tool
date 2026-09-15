import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[3] if cell.value}


def test_ipv4_vip_group_expands_all_valid_members_in_source_order_and_exports_relationships():
    result = extract_fortigate_config("""
config firewall vip
    edit "VIP_A"
        set extip 203.0.113.80
        set mappedip "10.0.0.80"
        set extintf "wan"
    next
    edit "VIP_B"
        set extip 203.0.113.81
        set mappedip "10.0.0.81"
        set extintf "wan"
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set member "VIP_A" "VIP_B"
    next
end
config firewall policy
    edit 50
        set srcintf "wan"
        set dstintf "lan"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "ALL"
        set action accept
    next
end
""")

    group = result.canonical_ir.virtual_ip_groups[0]
    assert group.members == ["VIP_A", "VIP_B"]
    assert group.unresolved_members == []

    rules = [
        rule
        for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "50"
    ]
    assert [rule.source_vip_reference for rule in rules] == ["VIP_A", "VIP_B"]
    assert [rule.source_vip_group_reference for rule in rules] == [
        "VIP_GROUP",
        "VIP_GROUP",
    ]
    assert [rule.destination for rule in rules] == [
        ["203.0.113.80"],
        ["203.0.113.81"],
    ]
    assert [rule.translated_destinations for rule in rules] == [
        ["10.0.0.80"],
        ["10.0.0.81"],
    ]

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    nat_sheet = workbook["NAT Rules"]
    nat_headers = _headers(nat_sheet)
    nat_rows = [
        row
        for row in range(4, nat_sheet.max_row + 1)
        if nat_sheet.cell(row, nat_headers["Source Policy ID"]).value == "50"
    ]
    assert len(nat_rows) == 2
    assert [nat_sheet.cell(row, nat_headers["VIP"]).value for row in nat_rows] == [
        "VIP_A",
        "VIP_B",
    ]
    assert [
        nat_sheet.cell(row, nat_headers["VIP Group"]).value for row in nat_rows
    ] == ["VIP_GROUP", "VIP_GROUP"]
    assert [
        nat_sheet.cell(row, nat_headers["Translated Destination"]).value
        for row in nat_rows
    ] == ["10.0.0.80", "10.0.0.81"]
