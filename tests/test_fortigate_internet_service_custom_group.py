from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_custom_group_members_are_preserved_as_extract_only_data():
    fg = parse_fortigate_config("""
config firewall internet-service-custom-group
    edit "web-group"
        set comment "Web services"
        set member "custom-web" "custom-api"
    next
end
""")
    group = fg.custom_internet_service_groups[0]
    assert group.members == ["custom-web", "custom-api"]
    ir = FGToIRTransformer(fg).transform()
    assert ir.custom_internet_service_groups[0].members == group.members
    assert ir.custom_internet_service_groups[0].migration_status == "EXTRACT_ONLY"


def test_custom_group_is_exported_to_its_dedicated_sheet():
    result = extract_fortigate_config("""
config firewall internet-service-custom-group
    edit "web-group"
        set comment "Web services"
        set member "custom-web" "custom-api"
    next
end
""")
    workbook = load_workbook(BytesIO(IRExcelExporter(result.canonical_ir, result).generate()))
    sheet = workbook["Custom Internet Service Groups"]
    headers = {cell.value: cell.column for cell in sheet[3]}

    assert set(headers) == {
        "Name", "Comment", "Members", "Status", "Manual Review", "Additional Settings"
    }
    assert sheet.cell(4, headers["Name"]).value == "web-group"
    assert sheet.cell(4, headers["Comment"]).value == "Web services"
    assert sheet.cell(4, headers["Members"]).value == "custom-web, custom-api"
    assert sheet.cell(4, headers["Status"]).value == "EXTRACT_ONLY"
    assert sheet.cell(4, headers["Manual Review"]).value == "Yes"
    assert sheet.cell(4, headers["Additional Settings"]).value in (None, "")
