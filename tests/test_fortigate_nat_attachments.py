import io
import json

from openpyxl import load_workbook

from fwmigrate.ir.io import dump_ir_json, load_ir_payload
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


BASE_CONFIG = """\
config system interface
    edit "lan"
        set ip 10.0.0.1 255.255.255.0
    next
    edit "wan"
        set ip 203.0.113.1 255.255.255.0
    next
end
config firewall address
    edit "src"
        set subnet 10.0.0.0 255.255.255.0
    next
    edit "dst"
        set subnet 192.0.2.0 255.255.255.0
    next
end
config firewall service custom
    edit "HTTPS"
        set tcp-portrange 443
    next
end
{extras}
config firewall policy
    edit 1
        set srcintf "{srcintf}"
        set dstintf "{dstintf}"
        set srcaddr "src"
        set dstaddr "{dstaddr}"
        set service "HTTPS"
        set action accept
        set schedule "always"
        {policy_settings}
    next
end
"""


def _extract(
    *,
    srcintf="lan",
    dstintf="wan",
    dstaddr="dst",
    policy_settings="set nat enable",
    extras="",
):
    return extract_fortigate_config(
        BASE_CONFIG.format(
            srcintf=srcintf,
            dstintf=dstintf,
            dstaddr=dstaddr,
            policy_settings=policy_settings,
            extras=extras,
        )
    ).canonical_ir


def _nat(ir):
    return next(rule for rule in ir.nat_rules if rule.source_origin == "firewall-policy")


def test_interface_attachments_are_preserved_without_inventing_zones():
    ir = _extract()
    policy = ir.policies[0]
    rule = _nat(ir)

    assert [(item.name, item.kind, item.resolved) for item in policy.source_attachments] == [
        ("lan", "interface", True)
    ]
    assert [(item.name, item.kind, item.resolved) for item in rule.destination_attachments] == [
        ("wan", "interface", True)
    ]
    assert policy.from_zone == []
    assert policy.to_zone == []
    assert rule.migration_status == "VENDOR_EXTENSION"
    assert rule.requires_manual_review is False


def test_zone_attachments_populate_canonical_zones():
    ir = _extract(
        srcintf="LAN_ZONE",
        dstintf="WAN_ZONE",
        extras="""\
config system zone
    edit "LAN_ZONE"
        set interface "lan"
    next
    edit "WAN_ZONE"
        set interface "wan"
    next
end
config firewall ippool
    edit "SNAT_POOL"
        set startip 203.0.113.10
        set endip 203.0.113.10
    next
end
""",
        policy_settings='set nat enable\n        set ippool enable\n        set poolname "SNAT_POOL"',
    )
    rule = _nat(ir)

    assert rule.from_zone == ["LAN_ZONE"]
    assert rule.to_zone == ["WAN_ZONE"]
    assert [item.kind for item in rule.source_attachments] == ["zone"]
    assert [item.kind for item in rule.destination_attachments] == ["zone"]
    assert rule.migration_status == "NORMALIZED"


def test_any_and_missing_attachments_are_distinguished():
    any_ir = _extract(srcintf="any")
    any_rule = _nat(any_ir)
    assert [(item.name, item.kind, item.resolved) for item in any_rule.source_attachments] == [
        ("any", "any", True)
    ]
    assert any_rule.migration_status == "VENDOR_EXTENSION"

    missing_ir = _extract(srcintf="missing")
    missing_rule = _nat(missing_ir)
    assert [(item.name, item.kind, item.resolved) for item in missing_rule.source_attachments] == [
        ("missing", "unknown", False)
    ]
    assert missing_rule.migration_status == "PARTIALLY_NORMALIZED"
    assert any("unresolved source NAT attachment 'missing'" in reason for reason in missing_rule.review_reasons)


def test_sdwan_and_interface_address_resolution_keep_distinct_states():
    sdwan_ir = _extract(
        srcintf="WAN",
        extras="""\
config system sdwan
    set status enable
    config zone
        edit "WAN"
        next
    end
    config members
        edit 1
            set interface "wan"
            set zone "WAN"
        next
    end
end
""",
    )
    sdwan_rule = _nat(sdwan_ir)
    assert [item.kind for item in sdwan_rule.source_attachments] == ["sdwan_zone"]
    assert sdwan_rule.from_zone == []

    dynamic_ir = _extract(
        dstintf="wan-dhcp",
        extras="""\
config system interface
    edit "wan-dhcp"
        set mode dhcp
    next
end
""",
    )
    dynamic_rule = _nat(dynamic_ir)
    assert dynamic_rule.migration_status == "PARTIALLY_NORMALIZED"
    assert any("interface-address SNAT is unresolved" in reason for reason in dynamic_rule.review_reasons)


def test_unrelated_policy_review_does_not_downgrade_nat():
    ir = _extract(
        policy_settings="set nat enable\n        set inspection-mode proxy",
    )
    rule = _nat(ir)

    assert ir.policies[0].requires_manual_review is True
    assert rule.migration_status == "VENDOR_EXTENSION"
    assert rule.requires_manual_review is False
    assert "source policy semantics require manual review" not in rule.review_reasons


def test_vip_dnat_preserves_both_policy_attachments():
    ir = _extract(
        srcintf="wan",
        dstintf="lan",
        dstaddr="VIP",
        policy_settings="set nat disable",
        extras="""\
config firewall vip
    edit "VIP"
        set extip 203.0.113.80
        set mappedip "10.0.0.80"
        set extintf "wan"
    next
end
""",
    )
    rule = next(item for item in ir.nat_rules if item.source_vip_reference == "VIP")

    assert [item.kind for item in rule.source_attachments] == ["interface"]
    assert [item.kind for item in rule.destination_attachments] == ["interface"]
    assert rule.migration_status == "VENDOR_EXTENSION"
    assert rule.requires_manual_review is False


def test_nat_attachments_round_trip_and_excel_columns():
    ir = _extract()
    restored = load_ir_payload(json.loads(dump_ir_json(ir)))
    rule = _nat(restored)
    assert [(item.name, item.kind, item.resolved) for item in rule.source_attachments] == [
        ("lan", "interface", True)
    ]

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(restored).generate())
    )
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Source Attachments"]).value == "lan"
    assert sheet.cell(4, headers["Source Attachment Types"]).value == "interface"
    assert sheet.cell(4, headers["Canonical Source Zones"]).value is None
    assert sheet.cell(4, headers["Destination Attachments"]).value == "wan"
    assert sheet.cell(4, headers["Destination Attachment Types"]).value == "interface"
    assert sheet.cell(4, headers["Canonical Destination Zones"]).value is None
    assert sheet.cell(4, headers["Attachment Resolution Status"]).value == "RESOLVED"
    assert sheet.cell(4, headers["Migration Status"]).value == "VENDOR_EXTENSION"
    assert "unresolved canonical NAT zones" not in str(
        sheet.cell(4, headers["Review Reasons"]).value or ""
    ).lower()
