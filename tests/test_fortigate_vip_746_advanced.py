import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter
from fwmigrate.report.excel_optimized import SinglePassIRExcelExporter


def _parse(body: str, name: str = "VIP_746"):
    return parse_fortigate_config(
        f'''config firewall vip
    edit "{name}"
{body}
    next
end
'''
    )


def test_vip_746_nested_source_and_typed_projection_are_lossless():
    parsed = _parse(
        '''        set extip 203.0.113.10
        set mappedip 10.0.0.10
        set portforward enable
        set h3-support enable
        set future-setting "keep-me"
        config quic
            set max-idle-timeout 30000
            set active-migration enable
            set future-quic-value "preserve"
        end
        config gslb-public-ips
            edit 2
                set ip 198.51.100.2
            next
        end
        config ssl-cipher-suites
            edit 1
                set priority 1
                set cipher TLS-AES-128-GCM-SHA256
                set versions tls-1.2 tls-1.3
            next
        end
        config ssl-server-cipher-suites
            edit 2
                set priority 2
                set cipher TLS-AES-256-GCM-SHA384
            next
        end
'''
    )

    source = parsed.vips[0]
    assert source.extra_settings["future_setting"] == "keep-me"
    assert [node.name for node in source.nested_configs] == [
        "quic", "gslb-public-ips", "ssl-cipher-suites", "ssl-server-cipher-suites"
    ]
    assert source.quic.max_idle_timeout == 30000
    assert source.quic.extra_settings["future_quic_value"] == "preserve"
    assert source.gslb_public_ips[0].index == 2
    assert source.ssl_cipher_suites[0].versions == ["tls-1.2", "tls-1.3"]
    assert source.ssl_server_cipher_suites[0].priority == 2

    vip = FGToIRTransformer(parsed).transform().virtual_ips[0]
    assert vip.source_quic.max_idle_timeout == 30000
    assert vip.source_gslb_public_ips[0].ip == "198.51.100.2"
    assert vip.source_ssl_cipher_suites[0].priority == 1
    assert vip.nested_source_configs[0].commands[0].key == "max-idle-timeout"
    assert vip.requires_manual_review is True
    assert vip.migration_status == "PARTIALLY_NORMALIZED"
    assert "unmodeled VIP source settings" in vip.audit_note
    assert "advanced nested VIP configuration" in vip.audit_note


def test_vip_defaults_keep_explicit_provenance_and_static_vip_normalized():
    parsed = _parse("""        set extip 203.0.113.11
        set mappedip 10.0.0.11
""", name="Static_Defaults")
    source = parsed.vips[0]
    vip = FGToIRTransformer(parsed).transform().virtual_ips[0]

    assert source.source_explicit_fields == {"extip", "mappedip"}
    assert vip.source_explicit_fields == ["extip", "mappedip"]
    assert vip.source_effective_settings == {
        "type": "static-nat",
        "protocol": "tcp",
        "portmapping_type": "1-to-1",
        "nat44": True,
        "nat46": False,
        "add_nat46_route": True,
    }
    assert vip.protocol is None
    assert vip.port_mapping_type is None
    assert vip.migration_status == "NORMALIZED"
    assert vip.requires_manual_review is False


def test_vip_nested_source_keeps_set_append_unset_and_multivalue_commands():
    parsed = _parse(
        """        config quic
            set active-migration enable
            append future-values one two
            unset active-migration
        end
"""
    )
    node = parsed.vips[0].nested_configs[0]
    assert [(command.operation, command.key, command.values) for command in node.commands] == [
        ("set", "active-migration", ["enable"]),
        ("append", "future-values", ["one", "two"]),
        ("unset", "active-migration", []),
    ]


def test_vip_nested_malformed_integer_is_preserved_for_review():
    parsed = _parse(
        """        config quic
            set max-ack-delay malformed
        end
"""
    )
    quic = parsed.vips[0].quic
    assert quic.max_ack_delay is None
    assert quic.extra_settings["unparsed_max_ack_delay"] == "malformed"


def test_vip_nested_configuration_is_exported_by_both_excel_paths():
    parsed = _parse(
        """        set extip 203.0.113.12
        config quic
            set max-ack-delay 25
        end
""",
        name="Excel_Nested",
    )
    ir = FGToIRTransformer(parsed).transform()

    for exporter in (IRExcelExporter, SinglePassIRExcelExporter):
        workbook = load_workbook(io.BytesIO(exporter(ir).generate()))
        sheet = workbook["VIP Nested Configuration"]
        headers = {cell.value: cell.column for cell in sheet[3]}
        rows = list(sheet.iter_rows(min_row=4, values_only=True))
        assert rows
        assert rows[0][headers["VIP Name"] - 1] == "Excel_Nested"
        assert rows[0][headers["Key"] - 1] == "max-ack-delay"
        assert rows[0][headers["Values"] - 1] == "25"
