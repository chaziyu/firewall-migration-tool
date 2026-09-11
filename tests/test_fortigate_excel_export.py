"""End-to-end FortiGate Excel export regression tests.

Verifies the complete pipeline:
FortiGate configuration -> parser -> extraction result -> IRConfig -> IRExcelExporter -> valid XLSX.
Specifically ensures multi-value Application Control profiles (multiple categories and multiple application IDs)
are parsed accurately, retained without silent loss, and exported into structurally valid workbooks.
"""

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import FortiGateParser
from fwmigrate.parsers.fortigate.tokenizer import FortiGateTokenizer
from fwmigrate.report.excel_exporter import IRExcelExporter, XLSX_MIMETYPE
from fwmigrate.web import _extract_source_config, create_app


FIXTURE_PATH = (
    Path(__file__).parent
    / "fixtures"
    / "fortigate"
    / "application_list_multiple_categories.conf"
)


@pytest.fixture
def client():
    return create_app({"TESTING": True}).test_client()


def test_fortigate_excel_export_supports_multiple_application_categories():
    """Verify that multi-category Application Control entries survive extraction and export to Excel."""
    content = FIXTURE_PATH.read_text(encoding="utf-8")

    # Step 1: FortiGate extraction via public extractor
    extraction_result = extract_fortigate_config(content)
    ir_config = extraction_result.canonical_ir
    assert ir_config is not None

    # Step 2: Assert Application Control source data was parsed and retained
    parser = FortiGateParser(FortiGateTokenizer(content))
    fg_config = parser.parse()
    assert len(fg_config.application_lists) >= 1
    profile = fg_config.application_lists[0]
    assert profile.name == "block-high-risk"
    entry1 = profile.entries[0]
    assert entry1.category == [2, 6, 7]

    # Source inventory commands retain original multi-value tokens
    inv_items = [i for i in parser.source_inventory_items if i.source_path == "application list"]
    assert len(inv_items) >= 1
    entries_child = next(c for c in inv_items[0].children if c.name == "entries")
    entry1_item = next(c for c in entries_child.children if c.name == "1")
    cat_cmd = next(c for c in entry1_item.commands if c.key == "category")
    assert cat_cmd.values == ["2", "6", "7"]

    # Source sections recorded in extraction result
    sections = {s.path: s for s in extraction_result.source_sections}
    assert "application list" in sections
    assert sections["application list"].status == ExtractionStatus.NORMALIZED
    assert "application list entries" in sections
    assert sections["application list entries"].status == ExtractionStatus.EXTRACT_ONLY

    # Step 3: Excel export completes without error
    excel_data = IRExcelExporter(
        ir_config,
        extraction_result=extraction_result,
    ).generate()
    assert isinstance(excel_data, bytes)
    assert len(excel_data) > 0

    # Step 4: Verify openpyxl can load the generated workbook
    workbook = load_workbook(BytesIO(excel_data))
    assert workbook.sheetnames
    assert len(workbook.sheetnames) > 0
    assert "Summary" in workbook.sheetnames
    assert "Extraction Coverage" in workbook.sheetnames

    # Check that Extraction Coverage accurately reflects application list sections
    coverage_sheet = workbook["Extraction Coverage"]
    coverage_headers = {cell.value: cell.column for cell in coverage_sheet[3]}
    section_rows = {
        coverage_sheet.cell(r, 1).value: r
        for r in range(4, coverage_sheet.max_row + 1)
    }
    assert "application list" in section_rows
    app_list_row = section_rows["application list"]
    assert coverage_sheet.cell(app_list_row, coverage_headers["Status"]).value == "NORMALIZED"
    assert coverage_sheet.cell(app_list_row, coverage_headers["Found"]).value == "Yes"


def test_fortigate_excel_export_supports_multiple_applications():
    """Verify that multiple application IDs survive extraction and export to Excel."""
    content = FIXTURE_PATH.read_text(encoding="utf-8")

    # Step 1: FortiGate extraction
    ir_config, extraction_result = _extract_source_config("fortigate", content)
    assert ir_config is not None
    assert extraction_result is not None

    # Step 2: Assert multiple application IDs are parsed as typed list of integers
    parser = FortiGateParser(FortiGateTokenizer(content))
    fg_config = parser.parse()
    profile = fg_config.application_lists[0]
    entry2 = profile.entries[1]
    assert entry2.application == [11414, 11767, 15722]
    assert entry2.action == "pass"

    # Step 3: Exporter produces valid workbook
    excel_data = IRExcelExporter(
        ir_config,
        extraction_result=extraction_result,
    ).generate()
    workbook = load_workbook(BytesIO(excel_data))
    assert workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    assert "Source Security Profiles" in workbook.sheetnames


def test_fortigate_excel_export_regression_multi_value_application_control():
    """End-to-end regression proving multi-category and multi-application entries don't raise Pydantic or Excel errors."""
    content = FIXTURE_PATH.read_text(encoding="utf-8")

    # Full web extraction helper
    ir_config, extraction_result = _extract_source_config("fortigate", content)
    assert ir_config is not None

    exporter = IRExcelExporter(ir_config, extraction_result=extraction_result)

    # Verify list serialization helper converts collections safely without causing openpyxl errors
    assert exporter._safe_value([2, 6, 7]) == "2\n6\n7"
    assert exporter._safe_value([11414, 11767, 15722]) == "11414\n11767\n15722"

    excel_bytes = exporter.generate()
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = load_workbook(BytesIO(excel_bytes))
    assert "Summary" in wb.sheetnames
    assert "Extraction Coverage" in wb.sheetnames

    # Workbook must contain populated Summary data
    summary = wb["Summary"]
    assert summary["A1"].value == "Firewall Source Inventory"


def test_fortigate_excel_export_includes_canonical_multicast_policies():
    content = """config firewall multicast-policy
 edit 1
  set name mcast-v4
  set protocol 17
  set start-port 5000
  set end-port 5001
  set dnat 198.51.100.10
 next
end
config firewall multicast-policy6
 edit 2
  set name mcast-v6
  set protocol 17
 next
end
"""
    result = extract_fortigate_config(content)
    workbook = load_workbook(BytesIO(IRExcelExporter(
        result.canonical_ir,
        extraction_result=result,
    ).generate()))

    sheet = workbook["Multicast Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Name"]).value: row
        for row in range(4, sheet.max_row + 1)
    }
    assert set(rows) == {"mcast-v4", "mcast-v6"}
    assert sheet.cell(rows["mcast-v4"], headers["Address Family"]).value == "ipv4"
    assert sheet.cell(rows["mcast-v6"], headers["Address Family"]).value == "ipv6"
    assert sheet.cell(rows["mcast-v4"], headers["Destination Start Port"]).value == 5000
    assert sheet.cell(rows["mcast-v4"], headers["Destination End Port"]).value == 5001

    nat = workbook["NAT Rules"]
    original_destination_port_column = next(
        cell.column for cell in nat[3]
        if cell.value == "Original Destination Port"
    )
    assert nat.cell(4, original_destination_port_column).value == "5000-5001"


def test_fortigate_excel_export_uses_canonical_nat_ports_and_typed_ngfw_profiles():
    result = extract_fortigate_config("""
config firewall vip
    edit "WEB"
        set extip 198.51.100.10
        set mappedip 10.0.0.10
        set portforward enable
        set extport 8443-8444
        set mappedport 443-444
    next
end
config firewall policy
    edit 1
        set srcintf "wan"
        set dstintf "lan"
        set srcaddr "all"
        set dstaddr "WEB"
        set service "ALL"
        set action accept
    next
end
config firewall security-policy
    edit 2
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set casb-profile "casb"
        set diameter-filter-profile "diameter"
        set virtual-patch-profile "virtual-patch"
        set waf-profile "waf"
    next
end
""")
    workbook = load_workbook(BytesIO(IRExcelExporter(
        result.canonical_ir,
        extraction_result=result,
    ).generate()))

    nat = workbook["NAT Rules"]
    headers = [cell.value for cell in nat[3]]
    nat_headers = {value: index + 1 for index, value in enumerate(headers)}
    nat_row = next(
        row for row in range(4, nat.max_row + 1)
        if nat.cell(row, nat_headers["Name"]).value == "DNAT-P1-WEB"
    )
    assert headers.count("Original Destination Port") == 1
    assert "Legacy Original Destination Port" in headers
    assert nat.cell(nat_row, nat_headers["Original Destination Port"]).value == "8443-8444"
    assert nat.cell(nat_row, nat_headers["Translated Destination Port"]).value == "443-444"

    ngfw = workbook["NGFW Security Policies"]
    ngfw_headers = {cell.value: cell.column for cell in ngfw[3]}
    profiles = ngfw.cell(4, ngfw_headers["Security Profile References"]).value
    assert "casb-profile=casb" in profiles
    assert "diameter-filter-profile=diameter" in profiles
    assert "virtual-patch-profile=virtual-patch" in profiles
    assert "waf-profile=waf" not in profiles
    assert "waf-profile=waf" in ngfw.cell(4, ngfw_headers["Additional Settings"]).value


def test_fortigate_excel_export_preserves_746_ip_pool_ir_and_review_values():
    result = extract_fortigate_config("""# config-version = 7.4.6
config firewall ippool
    edit "ONE_TO_ONE"
        set type one-to-one
        set startip 203.0.113.10
        set endip 203.0.113.20
        set source-startip 10.0.0.10
        set source-endip 10.0.0.20
        set startport 5117
        set endport 65533
        set future-pool-setting retained
    next
end
""")
    pool = result.canonical_ir.ip_pools[0]
    assert pool.pool_type == "one-to-one"
    assert (pool.start_ip, pool.end_ip) == ("203.0.113.10", "203.0.113.20")
    assert (pool.source_start_ip, pool.source_end_ip) == ("10.0.0.10", "10.0.0.20")
    assert (pool.start_port, pool.end_port) == (5117, 65533)
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True

    workbook = load_workbook(BytesIO(IRExcelExporter(
        result.canonical_ir,
        extraction_result=result,
    ).generate()))
    sheet = workbook["IP Pools"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Type"]).value == "one-to-one"
    assert sheet.cell(4, headers["Start IP"]).value == "203.0.113.10"
    assert sheet.cell(4, headers["End IP"]).value == "203.0.113.20"
    assert sheet.cell(4, headers["Source Start IP"]).value == "10.0.0.10"
    assert sheet.cell(4, headers["Source End IP"]).value == "10.0.0.20"
    assert sheet.cell(4, headers["Start Port"]).value == 5117
    assert sheet.cell(4, headers["End Port"]).value == 65533
    assert sheet.cell(4, headers["Extraction Status"]).value == "PARTIALLY_NORMALIZED"
    assert sheet.cell(4, headers["Manual Review"]).value == "TRUE"
    assert "future-pool-setting=retained" in sheet.cell(
        4, headers["Additional Settings"]
    ).value


def test_fortigate_excel_export_web_route_extract_excel(client):
    """Verify the /api/extract/excel Flask route accepts the FortiGate fixture and returns a valid XLSX file."""
    content = FIXTURE_PATH.read_bytes()

    response = client.post(
        "/api/extract/excel",
        data={
            "source_vendor": "fortigate",
            "file": (BytesIO(content), "application_list_multiple_categories.conf"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert response.mimetype == XLSX_MIMETYPE

    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    assert "Extraction Coverage" in workbook.sheetnames


def test_fortigate_excel_export_web_route_migrate_zip(client):
    """Verify the /api/migrate Flask route includes a valid source inventory workbook in the generated ZIP."""
    import zipfile

    content = FIXTURE_PATH.read_bytes()

    response = client.post(
        "/api/migrate",
        data={
            "source_vendor": "fortigate",
            "target_vendor": "palo_alto",
            "file": (BytesIO(content), "application_list_multiple_categories.conf"),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert response.mimetype == "application/zip"

    with zipfile.ZipFile(BytesIO(response.data)) as archive:
        assert "source_inventory_fortigate.xlsx" in archive.namelist()
        inventory_bytes = archive.read("source_inventory_fortigate.xlsx")
        workbook = load_workbook(BytesIO(inventory_bytes))
        assert workbook.sheetnames
        assert "Summary" in workbook.sheetnames
        assert "Extraction Coverage" in workbook.sheetnames
