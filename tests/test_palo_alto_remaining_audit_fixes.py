import io
from textwrap import dedent

from openpyxl import load_workbook

from fwmigrate.parsers.palo_alto.parser import PANOSSourceParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_static_route_uses_official_monitor_destinations_and_pan_bounds():
    result = PANOSSourceParser().extract(dedent("""
        <config><devices><entry name="fw-a">
          <network><virtual-router><entry name="vr1"><routing-table><ip><static-route>
            <entry name="default-monitor">
              <destination>0.0.0.0/0</destination>
              <path-monitor>
                <enable>yes</enable>
                <monitor-destinations>
                  <entry name="primary">
                    <destination>198.51.100.1</destination><interval>1</interval><count>3</count>
                  </entry>
                  <entry name="backup">
                    <destination>198.51.100.2</destination><interval>60</interval><count>10</count>
                  </entry>
                  <entry name="invalid-count">
                    <destination>198.51.100.3</destination><interval>5</interval><count>2</count>
                  </entry>
                </monitor-destinations>
              </path-monitor>
            </entry>
          </static-route></ip></routing-table></entry></virtual-router></network>
          <vsys><entry name="vsys1"/></vsys>
        </entry></devices></config>
    """))

    route = result.canonical_ir.routes[0]
    monitor = route.path_monitor
    assert monitor is not None
    assert [item.name for item in monitor.destinations] == [
        "primary", "backup", "invalid-count"
    ]
    assert [(item.interval, item.count) for item in monitor.destinations[:2]] == [
        (1, 3), (60, 10)
    ]
    assert monitor.destinations[2].count is None
    assert any(
        "./count must be between 3 and 10" in reason
        for reason in monitor.destinations[2].review_reasons
    )


def test_pbf_accepts_direct_address_scalars_and_official_symmetric_return_list():
    result = PANOSSourceParser().extract(dedent("""
        <config><vsys><entry name="vsys1"><rulebase><pbf><rules>
          <entry name="direct-scalars">
            <source><member>192.0.2.0/24</member><member>app.example.com</member></source>
            <destination><member>2001:db8::/64</member></destination>
            <application><member>any</member></application>
            <service><member>any</member></service>
            <action><forward><nexthop><ip-address>198.51.100.1</ip-address></nexthop></forward></action>
            <enforce-symmetric-return>
              <enabled>yes</enabled>
              <nexthop-address-list>
                <entry name="198.51.100.2"/>
                <entry name="198.51.100.3"/>
              </nexthop-address-list>
            </enforce-symmetric-return>
          </entry>
        </rules></pbf></rulebase></entry></vsys></config>
    """))

    assert len(result.canonical_ir.pbf_rules) == 1
    rule = result.canonical_ir.pbf_rules[0]
    assert rule.source == ["192.0.2.0/24", "app.example.com"]
    assert rule.destination == ["2001:db8::/64"]
    assert "unresolved-address-reference" not in rule.review_reasons
    assert rule.symmetric_return is not None
    assert rule.symmetric_return.enabled is True
    assert rule.symmetric_return.next_hop_addresses == [
        "198.51.100.2", "198.51.100.3"
    ]


def test_nat_original_literal_matches_are_not_unresolved_object_references():
    result = PANOSSourceParser().extract(dedent("""
        <config><vsys><entry name="vsys1"><rulebase><nat><rules>
          <entry name="literal-match">
            <from><member>any</member></from><to><member>any</member></to>
            <source><member>192.0.2.0/24</member></source>
            <destination><member>198.51.100.10</member></destination>
            <service>any</service>
            <source-translation><dynamic-ip-and-port>
              <translated-address><member>203.0.113.10</member></translated-address>
            </dynamic-ip-and-port></source-translation>
          </entry>
        </rules></nat></rulebase></entry></vsys></config>
    """))

    assert len(result.canonical_ir.nat_rules) == 1
    rule = result.canonical_ir.nat_rules[0]
    assert rule.source == ["192.0.2.0/24"]
    assert rule.destination == ["198.51.100.10"]
    assert "unresolved-source" not in rule.review_reasons
    assert "unresolved-destination" not in rule.review_reasons
    assert "pan_unresolved_sources" not in rule.source_attributes
    assert "pan_unresolved_destinations" not in rule.source_attributes


def test_palo_alto_security_rule_type_is_visible_in_policy_excel():
    result = PANOSSourceParser().extract(dedent("""
        <config><vsys><entry name="vsys1"><rulebase><security><rules>
          <entry name="interzone-rule">
            <from><member>trust</member></from><to><member>untrust</member></to>
            <source><member>any</member></source><destination><member>any</member></destination>
            <application><member>any</member></application><service><member>any</member></service>
            <rule-type>interzone</rule-type><action>allow</action>
          </entry>
        </rules></security></rulebase></entry></vsys></config>
    """))

    policy = result.canonical_ir.policies[0]
    assert policy.source_extra_settings["pan_rule_type"] == "interzone"
    assert policy.source_extra_settings["pan_rule_type_valid"] is True

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir).generate()))
    sheet = workbook["Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Rule Type"]).value == "interzone"
