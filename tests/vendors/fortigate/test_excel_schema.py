from __future__ import annotations
import io
import unittest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from fwmigrate.vendors.fortigate.config import ExtractionConfig
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.export.excel_schema import SHEET_HEADERS, SHEET_ORDER
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.security.extraction import sanitize_source_attributes, sanitize_source_value
from fwmigrate.vendors.fortigate.validation.validator import validate_config

_SAMPLE_CONFIG = r"""
config system global
    set hostname "FG-TEST"
end

config system interface
    edit "wan-dhcp"
        set mode dhcp
    next
    edit "wan-pppoe"
        set mode pppoe
        set pppoe-username "source-user"
    next
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
        set allowaccess ping https
        config secondaryip
            edit 1
                set ip 192.0.2.2 255.255.255.0
                set ha-priority 20
            next
        end
    next
    edit "port2"
    next
    edit "agg1"
        set type aggregate
        set member "port1" "port2"
    next
    edit "vlan100"
        set type vlan
        set interface "agg1"
        set vlanid 100
        set ip 10.100.0.1 255.255.255.0
        set ip6 2001:db8:100::1/64
    next
end

config firewall service custom
    edit "multi-port"
        set tcp-portrange 80 443
        set udp-portrange 53
    next
    edit "single-port"
        set protocol TCP
        set tcp-portrange 8443
    next
end

config firewall service group
    edit "web-services"
        set member "multi-port" "single-port"
    next
end

config firewall address
    edit "inside"
        set subnet 10.100.0.0 255.255.255.0
    next
end

config firewall ippool
    edit "pool1"
        set startip 198.51.100.10
        set endip 198.51.100.20
    next
end

config firewall vip
    edit "lb-vip"
        set type server-load-balance
        set extip 198.51.100.30
        set extintf "wan-dhcp"
        set ldb-method round-robin
        config realservers
            edit 1
                set ip 10.100.0.10
                set port 443
            next
            edit 2
                set ip 10.100.0.11
                set port 443
            next
        end
    next
end

config vpn ipsec phase1-interface
    edit "VPN-HQ"
        set interface "vlan100"
        set remote-gw 203.0.113.2
        set psksecret "do-not-export-this-secret"
    next
end

config vpn ipsec phase2-interface
    edit "VPN-subnet"
        set phase1name "VPN-HQ"
        set src-subnet 10.100.0.0 255.255.255.0
        set dst-subnet 10.200.0.0 255.255.255.0
    next
    edit "VPN-range"
        set phase1name "VPN-HQ"
        set src-start-ip 10.100.0.10
        set src-end-ip 10.100.0.20
        set dst-start-ip 10.200.0.10
        set dst-end-ip 10.200.0.20
    next
end

config firewall policy
    edit 1
        set name "A policy name that is deliberately much longer than thirty-two characters"
        set srcintf "vlan100"
        set dstintf "port1"
        set srcaddr "inside"
        set dstaddr "all"
        set service "web-services"
        set action accept
        set schedule always
        set nat enable
    next
    edit 2
        set name "Collision policy name with a shared long prefix A"
        set srcintf "vlan100"
        set dstintf "port1"
        set srcaddr "inside"
        set dstaddr "all"
        set service "multi-port"
        set action accept
        set schedule always
        set nat enable
        set ippool enable
        set poolname "pool1"
    next
    edit 3
        set name "Collision policy name with a shared long prefix B"
        set srcintf "vlan100"
        set dstintf "port1" "port2"
        set srcaddr "inside"
        set dstaddr "all"
        set service "multi-port"
        set action accept
        set schedule always
        set nat enable
    next
    edit 4
        set name "short policy"
    next
end

config firewall unsupported-section
    edit "source-only"
        set unsupported-setting "preserve-me"
        append unsupported-list "first" "second"
        unset unsupported-setting
    next
end

config user local
    edit "alice"
        set type password
        set passwd "another-do-not-export-secret"
    next
end
"""

class ExcelReportTest(unittest.TestCase):
    def _workbook(self, config_text: str = '''config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
    next
end
'''):
        extracted = extract_fortigate_config(
            parse_fortigate_config(config_text),
            config=ExtractionConfig(),
        )
        derived = build_derived_views(extracted.config)
        validation = validate_config(extracted.config, derived=derived)
        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
            source_name="sample.conf",
        )
        output.seek(0)
        return load_workbook(output, data_only=False)
    @staticmethod
    def _rows(sheet):
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        return headers, list(sheet.iter_rows(min_row=4, values_only=True))
    def test_workbook_order_headers_and_removed_sheets(self):
        workbook = self._workbook()
        self.assertEqual(list(SHEET_ORDER), workbook.sheetnames)
        removed = {
            "System Settings", "Session Helpers", "Router Settings", "IPS Settings",
            "Service Categories", "Protocol Options", "Per-IP Shapers",
            "SSL VPN Host Checks", "SSL VPN Host Check Items", "User Group Guests",
            "NAC Policies", "Kerberos Keytabs", "KMIP Servers", "SDN Proxies",
            "IPv6 Address Templates", "On-Demand Sniffers", "Affinity Interrupts",
            "Serial Ports", "Firewall Regions", "Vendor MACs",
            "Schedules", "Schedule Groups",
            "SSL VPN Bookmark Owners", "SSL VPN Bookmarks",
            "Address Group Tags", "Local-In Policies",
            "Multicast Policies", "Policy Routes", "Routing Protocol Settings",
            "Session TTL Settings", "Session TTL Overrides", "SD-WAN SLAs", "SD-WAN Duplication",
            "SD-WAN Neighbors", "SD-WAN Rule SLAs",
            "SSL VPN Bookmark Groups", "LDAP Servers", "RADIUS Servers",
            "TACACS+ Servers", "SAML Servers", "FSSO Servers", "FortiTokens", "Authentication Rules",
            "Identity Server Endpoints", "Extraction Evidence", "Firewall Policy Source Settings",
            "Interface Source Settings", "Interface Nested Configuration",
        }
        self.assertTrue(removed.isdisjoint(SHEET_ORDER))
        self.assertTrue(removed.isdisjoint(SHEET_HEADERS))
        self.assertTrue(removed.isdisjoint(workbook.sheetnames))
        summary = workbook["Summary"]
        navigation = {summary.cell(row, 5).value for row in range(4, summary.max_row + 1)}
        inventory = {summary.cell(row, 1).value for row in range(12, summary.max_row + 1)}
        self.assertTrue(removed.isdisjoint(navigation))
        self.assertTrue(removed.isdisjoint(inventory))
        self.assertIn("FortiGate Source Inventory", workbook.sheetnames)
        self.assertNotIn("Source Inventory", workbook.sheetnames)
        self.assertNotIn("FortiGate Source Configuration", workbook.sheetnames)
        for sheet_name in SHEET_ORDER:
            if sheet_name == "Summary":
                continue
            sheet = workbook[sheet_name]
            self.assertEqual(
                list(SHEET_HEADERS[sheet_name]),
                [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)],
                sheet_name,
            )
    def test_presentation_and_review_contract(self):
        workbook = self._workbook()
        interfaces = workbook["Interfaces"]
        self.assertEqual("D4", interfaces.freeze_panes)
        self.assertEqual("17324D", interfaces["A1"].fill.fgColor.rgb[-6:])
        self.assertEqual("0F766E", interfaces["A3"].fill.fgColor.rgb[-6:])
        self.assertTrue(interfaces.column_dimensions["V"].hidden)
        self.assertTrue(interfaces.column_dimensions["W"].hidden)
        self.assertTrue(interfaces.auto_filter.ref.startswith("A3:"))
        self.assertEqual("#'Summary'!A1", interfaces[2][interfaces.max_column - 1].hyperlink.target)

        for sheet_name, headers in SHEET_HEADERS.items():
            if sheet_name == "Summary":
                continue
            sheet = workbook[sheet_name]
            for column, header in enumerate(headers, start=1):
                if header in {"Source Explicit Fields", "Additional Settings"}:
                    self.assertTrue(sheet.column_dimensions[get_column_letter(column)].hidden)

        review, rows = self._rows(workbook["Review Required"])
        for row in rows:
            self.assertTrue(row[review.index("VDOM")])
            self.assertTrue(row[review.index("Field")])
    def test_removed_legacy_target_columns(self):
        workbook = self._workbook()
        self.assertNotIn("System Settings", workbook.sheetnames)
        policy_headers, _ = self._rows(workbook["Policies"])
        self.assertNotIn("Source Address (Original)", policy_headers)
        pool_headers, _ = self._rows(workbook["IP Pools"])
        self.assertNotIn("Check Point Pool Object Type", pool_headers)
        security_headers, _ = self._rows(workbook["Security Profiles"])
        self.assertNotIn("WildFire", security_headers)
