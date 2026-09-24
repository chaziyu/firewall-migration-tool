import io

from openpyxl import load_workbook

from fwmigrate.vendors.fortigate.config import ExtractionConfig
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.extraction.coverage import typed_source_paths
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter
from fwmigrate.vendors.fortigate.section_registry import registered_sections
from fwmigrate.vendors.fortigate.validation.validator import validate_config


def test_static6_is_typed_and_reaches_reports():
    source = """config router static6
    edit 21
        set dst 2001:db8:1::/64
        set device port1
        set gateway 2001:db8::1
        set distance 7
        set sdwan-zone virtual-wan-link
        set weight 3
        set vendor-option preserve-me
    next
end
"""
    extracted = extract_fortigate_config(parse_fortigate_config(source), config=ExtractionConfig())
    route = extracted.config.static_routes[0]

    assert route.address_family == "ipv6"
    assert route.dst == "2001:db8:1::/64"
    assert route.device == "port1"
    assert route.distance == 7
    assert route.sdwan_zone == ["virtual-wan-link"]
    assert route.raw_extra["vendor-option"] == "preserve-me"
    assert "router static6" in typed_source_paths()

    derived = build_derived_views(extracted.config)
    analysis = FortiGateSourceReporter().analyze_source(source)
    preview = FortiGateSourceReporter().build_preview(analysis)
    assert preview["sections"]["routes"][0]["address_family"] == "ipv6"

    output = io.BytesIO()
    export_excel(extracted=extracted, derived=derived, validation=validate_config(extracted.config, derived=derived), output=output)
    sheet = load_workbook(io.BytesIO(output.getvalue()), read_only=True)["Routes"]
    headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
    values = next(sheet.iter_rows(min_row=4, values_only=True))
    assert values[headers.index("Address Family")] == "ipv6"


def test_ippool6_has_independent_namespace_and_policy_reference():
    from fwmigrate.vendors.fortigate.model.ippool import FGIPPool
    from fwmigrate.vendors.fortigate.model.ippool6 import FGIPPool6
    from fwmigrate.vendors.fortigate.model.policy import FGPolicy
    from fwmigrate.vendors.fortigate.model.source import FGConfig
    from fwmigrate.vendors.fortigate.relationships.references import (
        ReferenceKind,
        build_reference_index,
        collect_broken_references,
    )

    config = FGConfig(
        ip_pools=[FGIPPool(name="shared")],
        ip_pools6=[FGIPPool6(name="shared", startip="2001:db8::1", nat46="enable")],
        policies=[FGPolicy(policy_id=1, poolname6=["shared", "missing-v6"])],
    )
    index = build_reference_index(config)
    assert index.get(ReferenceKind.IP_POOL, vdom="root", name="shared") is config.ip_pools[0]
    assert index.get(ReferenceKind.IP_POOL6, vdom="root", name="shared") is config.ip_pools6[0]
    assert [(item.source_field, item.reference, item.expected_kinds) for item in collect_broken_references(config)] == [
        ("poolname6", "missing-v6", (ReferenceKind.IP_POOL6,))
    ]


def test_ippool6_extraction_preserves_unknowns_and_reports_family():
    source = """config firewall ippool6
    edit v6-pool
        set startip 2001:db8::10
        set endip 2001:db8::20
        set nat46 enable
        set future-pool-setting keep-me
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    pool = analysis.extracted.config.ip_pools6[0]
    assert pool.startip == "2001:db8::10"
    assert pool.nat46 == "enable"
    assert pool.raw_extra["future-pool-setting"] == "keep-me"
    assert pool.address_family == "ipv6"

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    sheet = load_workbook(io.BytesIO(output.getvalue()), read_only=True)["IP Pools"]
    headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
    values = next(sheet.iter_rows(min_row=4, values_only=True))
    assert values[headers.index("Address Family")] == "ipv6"
    assert values[headers.index("NAT46")] == "enable"


def test_vip6_and_vipgrp6_keep_separate_namespaces_and_report():
    from fwmigrate.vendors.fortigate.model.vip import FGVIP, FGVIPGroup
    from fwmigrate.vendors.fortigate.model.vip6 import FGVIP6, FGVIPGroup6
    from fwmigrate.vendors.fortigate.model.source import FGConfig
    from fwmigrate.vendors.fortigate.relationships.references import (
        ReferenceKind,
        build_reference_index,
        collect_broken_references,
    )

    config = FGConfig(
        vips=[FGVIP(name="same")], vips6=[FGVIP6(name="same")],
        vip_groups=[FGVIPGroup(name="group")],
        vip_groups6=[FGVIPGroup6(name="group", members=["same", "missing-v6"])],
    )
    index = build_reference_index(config)
    assert index.get(ReferenceKind.VIP, vdom="root", name="same") is config.vips[0]
    assert index.get(ReferenceKind.VIP6, vdom="root", name="same") is config.vips6[0]
    assert not index.duplicates
    assert [(item.source_field, item.reference) for item in collect_broken_references(config)] == [
        ("members", "missing-v6")
    ]

    source = """config firewall vip6
    edit same
        set extip 2001:db8::1
        set mappedip 2001:db8:1::1
        set nat64 enable
        set vendor-option preserve-me
    next
end
config firewall vipgrp6
    edit group
        set member same
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    assert analysis.extracted.config.vips6[0].extip == "2001:db8::1"
    assert analysis.extracted.config.vips6[0].raw_extra["vendor-option"] == "preserve-me"
    assert analysis.extracted.config.vip_groups6[0].members == ["same"]

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    vip_headers = [workbook["Virtual IPs"].cell(3, column).value for column in range(1, workbook["Virtual IPs"].max_column + 1)]
    vip_values = next(workbook["Virtual IPs"].iter_rows(min_row=4, values_only=True))
    assert vip_values[vip_headers.index("Address Family")] == "ipv6"
    assert vip_values[vip_headers.index("NAT64")] == "enable"
    group_headers = [workbook["VIP Groups"].cell(3, column).value for column in range(1, workbook["VIP Groups"].max_column + 1)]
    group_values = next(workbook["VIP Groups"].iter_rows(min_row=4, values_only=True))
    assert group_values[group_headers.index("Address Family")] == "ipv6"


def test_policy_based_ipsec_phase1_is_separate_and_secret_safe():
    source = """config vpn ipsec phase1
    edit policy-vpn
        set interface wan1
        set remote-gw 192.0.2.1
        set proposal aes256-sha256
        set psksecret do-not-retain-this
        set authpasswd also-do-not-retain
        unset authpasswd
    next
end
config firewall policy
    edit 9
        set action ipsec
        set vpntunnel policy-vpn
    next
end
config vpn ipsec phase2
    edit policy-child
        set phase1name policy-vpn
        set proposal aes256-sha256
        set pfs enable
        set src-addr-type subnet
        set src-subnet 10.1.0.0 255.255.255.0
        set dst-addr-type subnet
        set dst-subnet 10.2.0.0 255.255.255.0
        set protocol 6
        set src-port 443
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    assert analysis.extracted.config.ipsec_phase1 == []
    assert analysis.extracted.config.ipsec_phase2 == []
    phase1 = analysis.extracted.config.ipsec_policy_phase1[0]
    assert phase1.interface == "wan1"
    assert phase1.psk_configured
    assert not phase1.auth_password_configured
    assert "vpn ipsec phase1" in typed_source_paths()
    assert "do-not-retain-this" not in repr(analysis)
    assert "also-do-not-retain" not in repr(analysis)
    assert not analysis.derived.topology.vpns
    assert not analysis.derived.broken_references
    policy_phase2 = analysis.extracted.config.ipsec_policy_phase2[0]
    assert policy_phase2.protocol == 6
    assert policy_phase2.src_port == 443
    normalized = analysis.derived.vpn.phase2[0]
    assert normalized.vpn_type == "policy-based"
    assert normalized.source_range == "10.1.0.0-10.1.0.255"
    preview = reporter.build_preview(analysis)
    assert preview["sections"]["policy_vpn_phase1"][0]["name"] == "policy-vpn"
    assert preview["sections"]["policy_vpn_phase2"][0]["source_range"] == "10.1.0.0-10.1.0.255"

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    assert "policy-vpn" in [workbook["Policy IPsec Phase 1"].cell(row, 1).value for row in (4, 5)]
    values = next(workbook["Policy IPsec Phase 1"].iter_rows(min_row=4, values_only=True))
    headers = [workbook["Policy IPsec Phase 1"].cell(3, column).value for column in range(1, workbook["Policy IPsec Phase 1"].max_column + 1)]
    assert values[headers.index("PSK Configured")] == "Yes"
    policy_phase2_headers = [workbook["Policy IPsec Phase 2"].cell(3, column).value for column in range(1, workbook["Policy IPsec Phase 2"].max_column + 1)]
    policy_phase2_values = next(workbook["Policy IPsec Phase 2"].iter_rows(min_row=4, values_only=True))
    assert policy_phase2_values[policy_phase2_headers.index("Source Range")] == "10.1.0.0-10.1.0.255"


def test_security_policy_protocol_options_and_per_ip_shaper_are_typed_and_linked():
    source = """config firewall profile-protocol-options
    edit proto
        set comment explicit-profile
        config http
            set ports 8080
            set future-http-option preserve-me
        end
    next
end
config firewall shaper per-ip-shaper
    edit client-limit
        set max-bandwidth 2000
        set max-concurrent-session 5
    next
end
config firewall policy
    edit 1
        set profile-protocol-options proto
        set per-ip-shaper client-limit
    next
end
config firewall security-policy
    edit 41
        set action accept
        set app-category 2 3
        set future-policy-option preserve-me-too
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    assert analysis.extracted.config.security_policies[0].policy_id == 41
    assert analysis.extracted.config.security_policies[0].app_category == [2, 3]
    assert analysis.extracted.config.security_policies[0].raw_extra["future-policy-option"] == "preserve-me-too"
    assert analysis.extracted.config.protocol_options[0].comment == "explicit-profile"
    assert analysis.extracted.config.per_ip_shapers[0].max_bandwidth == 2000
    assert not analysis.derived.broken_references
    preview = reporter.build_preview(analysis)
    assert preview["sections"]["security_policies"][0]["policy_id"] == 41
    assert preview["sections"]["policies"][0]["per_ip_shaper"] == "client-limit"

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    headers = [workbook["Security Policies"].cell(3, column).value for column in range(1, workbook["Security Policies"].max_column + 1)]
    values = next(workbook["Security Policies"].iter_rows(min_row=4, values_only=True))
    assert values[headers.index("Policy ID")] == 41
    assert "client-limit" in [workbook["Policies"].cell(4, column).value for column in range(1, workbook["Policies"].max_column + 1)]


def test_session_helper_is_typed_and_reported():
    source = """config system session-helper
    edit 18
        set name ftp
        set protocol 6
        set port 2121
        set future-setting preserve-me
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    helper = analysis.extracted.config.session_helpers[0]
    assert (helper.id, helper.name, helper.protocol, helper.port) == (18, "ftp", 6, 2121)
    assert helper.raw_extra["future-setting"] == "preserve-me"
    assert "system session-helper" in typed_source_paths()

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    sheet = load_workbook(io.BytesIO(output.getvalue()), read_only=True)["Session Helpers"]
    assert next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))[:4] == (18, "ftp", 2121, 6)


def test_ssl_vpn_realms_and_clients_are_typed_and_secret_safe():
    source = """config vpn ssl web realm
    edit /partner
        set login-page <html>custom login page</html>
        set radius-server radius1
        set radius-port 1812
        set virtual-host-only enable
    next
end
config vpn ssl client
    edit remote
        set server vpn.example.test
        set psk never-export-this
        set status enable
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    realm = analysis.extracted.config.ssl_vpn_realms[0]
    client = analysis.extracted.config.ssl_vpn_clients[0]
    assert realm.url_path == "/partner"
    assert realm.radius_port == 1812
    assert client.psk_configured
    assert "never-export-this" not in repr(analysis)
    assert "vpn ssl web realm" in typed_source_paths()
    assert "vpn ssl client" in typed_source_paths()
    preview = reporter.build_preview(analysis)
    assert preview["sections"]["ssl_vpn_realms"][0]["login_page_length"] > 0

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    realm_headers = [workbook["SSL VPN Realms"].cell(3, col).value for col in range(1, workbook["SSL VPN Realms"].max_column + 1)]
    realm_values = next(workbook["SSL VPN Realms"].iter_rows(min_row=4, values_only=True))
    assert realm_values[realm_headers.index("Login Page Length")] > 0
    assert "<html>" not in repr(realm_values)
    client_headers = [workbook["SSL VPN Clients"].cell(3, col).value for col in range(1, workbook["SSL VPN Clients"].max_column + 1)]
    client_values = next(workbook["SSL VPN Clients"].iter_rows(min_row=4, values_only=True))
    assert client_values[client_headers.index("PSK Configured")] == "Yes"
    all_cells = repr([[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets])
    assert "never-export-this" not in all_cells


def test_ssl_vpn_user_and_group_bookmarks_share_nested_models_and_redact_passwords():
    source = """config vpn ssl web user-bookmark
    edit alice
        set custom-lang en
        config bookmarks
            edit app
                set apptype rdp
                set host 192.0.2.20
                set logon-user alice
                set logon-password never-export-bookmark-secret
                set sso-password another-never-export-secret
                config form-data
                    edit token
                        set value form-field
                    next
                end
            next
        end
    next
end
config vpn ssl web user-group-bookmark
    edit remote-users
        config bookmarks
            edit web
                set apptype web
                set url https://example.test
            next
        end
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    user_bookmark = analysis.extracted.config.ssl_vpn_user_bookmarks[0].bookmarks[0]
    group_bookmark = analysis.extracted.config.ssl_vpn_user_group_bookmarks[0].bookmarks[0]
    assert user_bookmark.form_data[0].value == "form-field"
    assert user_bookmark.logon_password_configured and user_bookmark.sso_password_configured
    assert group_bookmark.apptype == "web"
    assert "never-export-bookmark-secret" not in repr(analysis)
    assert "another-never-export-secret" not in repr(analysis)
    preview = reporter.build_preview(analysis)
    assert preview["sections"]["ssl_vpn_bookmarks"][0]["owner_type"] == "user"
    assert "vpn ssl web user-bookmark bookmarks form-data" in typed_source_paths()

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    item_rows = list(workbook["SSL VPN Bookmarks"].iter_rows(min_row=4, values_only=True))
    assert {row[0] for row in item_rows} == {"User", "Group"}
    assert "never-export-bookmark-secret" not in repr(item_rows)
    assert "another-never-export-secret" not in repr(item_rows)
    all_cells = repr([[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets])
    assert "never-export-bookmark-secret" not in all_cells
    assert "another-never-export-secret" not in all_cells


def test_remaining_selected_sections_extract_and_report_with_secret_redaction():
    source = """config firewall address6-template
    edit lan-template
        set ip6 2001:db8::/48
        config subnet-segment
            edit 1
                set bits 16
                set name site
                config values
                    edit main
                        set value 10
                    next
                end
            next
        end
    next
end
config user nac-policy
    edit known-device
        set category device
        set firewall-address addr
        set user-group staff
        set mac aa:bb:cc:dd:ee:ff
    next
end
config firewall address
    edit addr
        set subnet 192.0.2.1 255.255.255.255
    next
end
config user group
    edit staff
        set group-type firewall
    next
end
config ips settings
    set ips-packet-quota 100
    set packet-log-history 30
end
config router setting
    set hostname edge-a
    set show-filter routes
end
config vpn kmip-server
    edit kmip
        set interface port1
        set password kmip-secret-value
        config server-list
            edit 1
                set server kmip.example.test
                set port 5696
            next
        end
    next
end
config user krb-keytab
    edit service-account
        set principal svc/example
        set ldap-server ldap1
        set keytab keytab-secret-value
    next
end
config system sdn-proxy
    edit cloud
        set server proxy.example.test
        set server-port 443
        set password proxy-secret-value
    next
end
config firewall on-demand-sniffer
    edit capture
        set hosts 192.0.2.10
        set interface port1
        set max-packet-count 100
        set ports 80 443
        set protocols 6 17
    next
end
config system interface
    edit port1
        set ip 192.0.2.1 255.255.255.0
    next
end
config system affinity-interrupt
    edit 5
        set interrupt irq5
        set affinity-cpumask 0x03
    next
end
config system serial-port
    edit console
    next
end
config firewall region
    edit 7
        set name region-a
        set city 3 4
    next
end
config firewall vendor-mac
    edit 100
        set name vendor-a
        set mac-number 200
        set obsolete 0
    next
end
"""
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(source)
    config = analysis.extracted.config
    assert config.addresses == [] or all(item.name != "lan-template" for item in config.addresses)
    assert config.address6_templates[0].segments[0].values[0].value == "10"
    assert config.ips_settings[0].ips_packet_quota == 100
    assert config.router_settings[0].hostname == "edge-a"
    assert config.kmip_servers[0].server_list[0].port == 5696
    assert config.kerberos_keytabs[0].keytab_configured
    assert config.sdn_proxies[0].password_configured
    assert config.on_demand_sniffers[0].ports == [80, 443]
    assert config.affinity_interrupts[0].id == 5
    assert config.serial_ports[0].name == "console"
    assert config.firewall_regions[0].city == [3, 4]
    assert config.vendor_macs[0].mac_number == 200
    assert not analysis.derived.broken_references
    assert reporter.build_preview(analysis)["sections"]["nac_policies"][0]["firewall_address"] == "addr"
    for section in (
        "user nac-policy", "firewall address6-template", "firewall address6-template subnet-segment",
        "firewall address6-template subnet-segment values", "ips settings", "router setting",
        "vpn kmip-server", "vpn kmip-server server-list", "user krb-keytab", "system sdn-proxy",
        "firewall on-demand-sniffer", "system affinity-interrupt", "system serial-port",
        "firewall region", "firewall vendor-mac",
    ):
        assert section in typed_source_paths()
    planned_sections = {
        "router static6", "firewall ippool6", "firewall vip6", "firewall vipgrp6",
        "vpn ipsec phase1", "vpn ipsec phase2", "firewall security-policy",
        "firewall profile-protocol-options", "firewall shaper per-ip-shaper",
        "system session-helper", "vpn ssl web realm", "vpn ssl web user-bookmark",
        "vpn ssl web user-group-bookmark", "vpn ssl client", "user nac-policy",
        "firewall address6-template", "ips settings", "vpn kmip-server", "user krb-keytab",
        "router setting", "system sdn-proxy", "firewall on-demand-sniffer",
        "system affinity-interrupt", "system serial-port", "firewall region", "firewall vendor-mac",
    }
    assert planned_sections <= set(registered_sections())
    assert planned_sections <= typed_source_paths()
    assert "kmip-secret-value" not in repr(analysis)
    assert "keytab-secret-value" not in repr(analysis)
    assert "proxy-secret-value" not in repr(analysis)

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    assert workbook["IPv6 Address Templates"].cell(4, 1).value == "lan-template"
    assert workbook["Kerberos Keytabs"].cell(4, 5).value == "Yes"
    assert workbook["KMIP Servers"].cell(4, 4).value == "Yes"
    assert workbook["SDN Proxies"].cell(4, 6).value == "Yes"
    assert workbook["Firewall Regions"].cell(4, 4).value == "Yes"
    assert workbook["Vendor MACs"].cell(4, 5).value == "Yes"
    all_cells = repr([[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets])
    for secret in ("kmip-secret-value", "keytab-secret-value", "proxy-secret-value"):
        assert secret not in all_cells
