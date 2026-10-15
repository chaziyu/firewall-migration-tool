from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.fortigate.model.policy import FGPolicy
from fwmigrate.vendors.fortigate.model.admin import FGAdministrator
from fwmigrate.vendors.fortigate.model.schedule import FGScheduleGroup, FGOneTimeSchedule, FGRecurringSchedule
from fwmigrate.vendors.fortigate.model.source import FGConfig
from fwmigrate.vendors.fortigate.relationships.references import collect_broken_references
from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter
from fwmigrate.source_reporting import ExcelExportProfile


FIXTURES = Path(__file__).parents[2] / "fixtures" / "fortigate"


def test_all_schedule_forms_are_extracted_from_source():
    reporter = FortiGateSourceReporter()
    result = reporter.analyze_source((FIXTURES / "p0_migration_critical.conf").read_text())
    config = result.extracted.config

    assert [item.name for item in config.recurring_schedules] == ["WORK_HOURS"]
    assert config.recurring_schedules[0].days == ["monday", "tuesday", "wednesday", "thursday", "friday"]
    assert [item.name for item in config.one_time_schedules] == ["MAINT_WINDOW"]
    assert config.one_time_schedules[0].start == "01:00 2026/10/01"
    assert config.schedule_groups[0].members == ["WORK_HOURS", "MAINT_WINDOW"]
    preview = reporter.build_preview(result)
    assert preview["summary"]["objects"]["schedules"] == 2
    assert preview["sections"]["schedule_groups"][0]["members"] == ["WORK_HOURS", "MAINT_WINDOW"]


def test_schedule_references_are_vdom_scoped_and_groups_resolve():
    config = FGConfig(
        one_time_schedules=[
            FGOneTimeSchedule(name="maintenance", vdom="vdom-a"),
            FGOneTimeSchedule(name="maintenance", vdom="root"),
        ],
        recurring_schedules=[FGRecurringSchedule(name="business", vdom="vdom-b")],
        schedule_groups=[FGScheduleGroup(name="group", vdom="vdom-a", members=["maintenance"])],
        administrators=[FGAdministrator(name="admin", schedule="maintenance")],
        policies=[
            FGPolicy(policy_id=1, vdom="vdom-a", schedule="group"),
            FGPolicy(policy_id=2, vdom="vdom-a", schedule="business"),
            FGPolicy(policy_id=3, vdom="vdom-a", schedule="missing"),
            FGPolicy(policy_id=4, vdom="vdom-a", schedule="always"),
        ],
    )

    broken = collect_broken_references(config)
    assert [(item.source_name, item.reference) for item in broken] == [("2", "business"), ("3", "missing")]


def test_unknown_schedule_fields_remain_in_raw_extra():
    result = FortiGateSourceReporter().analyze_source(
        """config firewall schedule recurring
    edit custom
        set day monday
        set start 08:00
        set end 18:00
        set future-setting retain-me
    next
end
"""
    )
    schedule = result.extracted.config.recurring_schedules[0]
    assert schedule.raw_extra["future-setting"] == "retain-me"


def test_native_schedule_sheets_and_policy_references_are_exported():
    analysis = FortiGateSourceReporter().analyze_source((FIXTURES / "p0_migration_critical.conf").read_text())
    for profile in ExcelExportProfile:
        output = BytesIO()
        FortiGateSourceReporter().export_excel(analysis, output, profile=profile)
        workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
        assert "Schedules" not in workbook.sheetnames
        for sheet, expected in (("Schedule Groups", "BUSINESS_SCHEDULES"),
                                ("One-Time Schedules", "MAINT_WINDOW"),
                                ("Recurring Schedules", "WORK_HOURS")):
            assert workbook[sheet].cell(4, 1).value == expected
        if profile == ExcelExportProfile.FULL:
            policies = workbook["Policies"]
            headers = {cell.value: cell.column for cell in policies[3]}
            assert policies.cell(4, headers["Schedule"]).value == "BUSINESS_SCHEDULES"


def test_repeated_schedule_edits_preserve_final_state_scope_and_inventory():
    reporter = FortiGateSourceReporter()
    blocks = '''config firewall schedule group
edit group
set member first
set future-setting retained
next
edit group
append member second
unset member
set member final
next
end
config firewall schedule onetime
edit once
set start "01:00 2026/10/01"
next
edit once
unset start
set end "02:00 2026/10/01"
set expiration-days 0
next
end
config firewall schedule recurring
edit recurring
set day monday
next
edit recurring
append day tuesday
unset day
set day friday
next
end
'''
    source = 'config vdom\n' + ''.join(
        f'edit {vdom}\n{blocks}next\n' for vdom in ('a', 'b')
    ) + 'end\n'
    analysis = reporter.analyze_source(source)
    config = analysis.extracted.config
    assert [(item.vdom, item.members) for item in config.schedule_groups] == [('a', ['final']), ('b', ['final'])]
    assert all(item.raw_extra == {'future-setting': 'retained'} for item in config.schedule_groups)
    assert all(item.explicit_fields == {'members'} for item in config.schedule_groups)
    assert [(item.vdom, item.start, item.end, item.expiration_days) for item in config.one_time_schedules] == [
        (vdom, None, '02:00 2026/10/01', 0) for vdom in ('a', 'b')
    ]
    assert all(item.days == ['friday'] and item.explicit_fields == {'days'} for item in config.recurring_schedules)
    assert not any('Multiple explicit objects' in issue.message for issue in analysis.validation.issues)
    output = BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(output, data_only=True)
    inventory = list(workbook['FortiGate Source Inventory'].values)
    assert any('first' in str(row) for row in inventory)
    assert any('second' in str(row) for row in inventory)


def test_native_schedule_fields_and_navigation_match_full_and_fast():
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source('''config firewall schedule onetime
edit once
set start "01:00 2026/10/01"
set end "02:00 2026/10/01"
set start-utc 1790816400
set end-utc 1790820000
set expiration-days 0
set color 0
set fabric-object disable
set future-setting retained
next
end
config firewall schedule recurring
edit recurring
set day monday tuesday
set start 08:00
set end 18:00
next
end
config firewall schedule group
edit group
set member once recurring
next
end
''')
    shared = {}
    for profile in ExcelExportProfile:
        output = BytesIO()
        reporter.export_excel(analysis, output, profile=profile)
        workbook = load_workbook(output, data_only=True)
        for name in ('One-Time Schedules', 'Recurring Schedules', 'Schedule Groups'):
            sheet = workbook[name]
            headers = [cell.value for cell in sheet[3]]
            row = dict(zip(headers, next(sheet.iter_rows(min_row=4, values_only=True))))
            if name in shared:
                assert row == shared[name]
            shared[name] = row
            if profile == ExcelExportProfile.FULL:
                assert sheet.cell(2, sheet.max_column).hyperlink.target == "#'Summary'!A1"
        once = shared['One-Time Schedules']
        assert once['Start'] == '01:00 2026/10/01'
        assert once['End'] == '02:00 2026/10/01'
        assert once['Start UTC'] == '1790816400'
        assert once['End UTC'] == '1790820000'
        assert once['Expiration Days'] == once['Color'] == 0
        assert once['Fabric Object'] == 'disable'
        assert 'retained' in once['Additional Settings']
        assert shared['Recurring Schedules']['Days'] == 'monday\ntuesday'
        assert shared['Recurring Schedules']['Color'] is None
        assert shared['Schedule Groups']['Members'] == 'once\nrecurring'
        summary = {row[0]: row[1] for row in workbook['Summary'].values}
        assert summary['One-Time Schedules'] == 1
        assert summary['Recurring Schedules'] == 1
        assert summary['Schedule Groups'] == 1
