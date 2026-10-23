from __future__ import annotations
import io
import unittest
from openpyxl import load_workbook
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.validation.validator import validate_config
from fwmigrate.source_reporting import source_reporters
from fwmigrate.vendors.fortigate.source_report import (
    FortiGateSourceReporter,
)

SOURCE = """config system global
    set hostname fg-report
end
config vdom
    edit root
        config firewall address
            edit web
                set subnet 192.0.2.10 255.255.255.255
            next
        end
    next
end
"""
UNSUPPORTED_SOURCE = """config firewall unsupported-section
    edit source-only
        set unsupported-setting preserve-me
    next
end
"""
SOURCE_INVENTORY = """config vdom
    edit tenant-a
        config firewall unsupported-section
            edit source-only
                set known value
                append member one
                unset comment
                mystery raw
                config tagging
                    edit tag-a
                        set color blue
                    next
                end
            next
        end
    next
end
"""
SECRET_SOURCE = """config vpn ipsec phase1-interface
    edit VPN-HQ
        set psksecret do-not-export-this-secret
    next
end
config user local
    edit alice
        set type password
        set passwd another-do-not-export-secret
    next
end
"""

def test_unnamed_policy_nat_review_stays_with_policy_id():
    source = """config firewall policy
    edit 1
        set nat enable
        set ippool enable
        set poolname missing
    next
    edit 2
    next
end
"""
    reporter = source_reporters.get("fortigate")
    analysis = reporter.analyze_source(source)
    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    headers = [cell.value for cell in workbook["Policies"][3]]
    rows = {row[headers.index("Rule #")]: row for row in workbook["Policies"].iter_rows(min_row=4, values_only=True)}
    assert rows[1][headers.index("Analysis Status")] == "REVIEW_REQUIRED"
    assert rows[2][headers.index("Analysis Status")] == "EXTRACTED"


def test_cli_constraint_warnings_reach_scoped_source_rows_and_review_sheets():
    from copy import deepcopy
    from fwmigrate.source_reporting import ExcelExportProfile

    reporter = FortiGateSourceReporter()
    source = '''config vdom
edit invalid
config firewall ippool
edit pool
set startport 5116
set endport 65534
next
end
config firewall schedule onetime
edit schedule
set expiration-days 101
next
end
config firewall schedule recurring
edit recurring
set day everyday
next
end
next
edit valid
config firewall ippool
edit pool
set startport 5117
set endport 65533
next
end
config firewall schedule onetime
edit schedule
set expiration-days 0
next
end
config firewall schedule recurring
edit recurring
set day none
next
end
next
end
'''
    analysis = reporter.analyze_source(source)
    before = deepcopy(analysis.extracted.config)
    preview = reporter.build_preview(analysis)
    assert all('FortiOS 7.4.6' in row['review'] for row in preview['sections']['schedules'] if row['vdom'] == 'invalid')
    assert all(row['review'] is None for row in preview['sections']['schedules'] if row['vdom'] == 'valid')
    shared = {}
    for profile in ExcelExportProfile:
        output = io.BytesIO()
        reporter.export_excel(analysis, output, profile=profile)
        workbook = load_workbook(output, data_only=True)
        for name in ('IP Pools', 'One-Time Schedules', 'Recurring Schedules', 'Review Required'):
            sheet = workbook[name]
            headers = [cell.value for cell in sheet[3]]
            rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=4, values_only=True)]
            if name in shared:
                assert rows == shared[name]
            shared[name] = rows
            if name == 'Review Required':
                assert len(rows) == 4
                assert {row['Source Sheet'] for row in rows} == {'IP Pools', 'One-Time Schedules', 'Recurring Schedules'}
                assert all(row['VDOM'] == 'invalid' for row in rows)
            else:
                scoped = {row['VDOM']: row for row in rows}
                assert scoped['invalid']['Analysis Status'] == 'REVIEW_REQUIRED'
                assert 'FortiOS 7.4.6' in scoped['invalid']['Review Reasons']
                assert scoped['valid']['Analysis Status'] == 'EXTRACTED'
        assert shared['IP Pools'][0]['Start Port'] == 5116
        assert shared['IP Pools'][0]['End Port'] == 65534
        assert shared['One-Time Schedules'][0]['Expiration Days'] == 101
        assert shared['Recurring Schedules'][0]['Days'] == 'everyday'
        if profile == ExcelExportProfile.FULL:
            assert 'FortiGate Source Inventory' in workbook.sheetnames
            assert 'Extraction Coverage' in workbook.sheetnames
        else:
            assert 'FortiGate Source Inventory' not in workbook.sheetnames
            assert 'Extraction Coverage' not in workbook.sheetnames
    assert analysis.extracted.config == before


class ExcelReportTest(unittest.TestCase):
    def _workbook(self, config_text: str = ""):
        extracted = extract_fortigate_config(
            parse_fortigate_config(config_text),
        )
        derived = build_derived_views(extracted.config)
        validation = validate_config(extracted.config, derived=derived)
        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
            source_name="sample.conf",
        )
        output.seek(0)
        return load_workbook(output, data_only=False)
    @staticmethod
    def _rows(sheet):
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        return headers, list(sheet.iter_rows(min_row=4, values_only=True))
    def test_analysis_status_is_scoped_by_validation_domain(self):
        workbook = self._workbook(
            r'''
config system interface
    edit "BCA_INF"
        set ip 192.0.2.1 255.255.255.0
    next
end

config vpn ipsec phase2-interface
    edit "BCA_INF"
        set src-addr-type range
        set src-start-ip 192.0.2.10
    next
end
''',
        )

        interfaces, interface_rows = self._rows(workbook["Interfaces"])
        interface = next(row for row in interface_rows if row[interfaces.index("Name")] == "● BCA_INF")
        self.assertEqual("EXTRACTED", interface[interfaces.index("Analysis Status")])
        self.assertFalse(interface[interfaces.index("Review Reasons")])

        phase2, phase2_rows = self._rows(workbook["VPN Phase 2"])
        phase2_row = next(row for row in phase2_rows if row[phase2.index("Name")] == "BCA_INF")
        self.assertEqual("REVIEW_REQUIRED", phase2_row[phase2.index("Analysis Status")])
        self.assertIn("Selector has only one range endpoint.", phase2_row[phase2.index("Review Reasons")])
    def test_same_name_vip_warning_does_not_mark_service_for_review(self):
        workbook = self._workbook(
            r'''
config firewall service custom
    edit "shared-name"
        set protocol TCP
        set tcp-portrange 443
    next
end

config firewall vip
    edit "shared-name"
        set type server-load-balance
        set extintf "missing-interface"
    next
end
''',
        )

        services, service_rows = self._rows(workbook["Services"])
        service = next(row for row in service_rows if row[services.index("Name")] == "shared-name")
        self.assertEqual("EXTRACTED", service[services.index("Analysis Status")])
        self.assertFalse(service[services.index("Review Reasons")])

        vips, vip_rows = self._rows(workbook["Virtual IPs"])
        vip = next(row for row in vip_rows if row[vips.index("Name")] == "shared-name")
        self.assertEqual("REVIEW_REQUIRED", vip[vips.index("Analysis Status")])
        self.assertIn("missing-interface", vip[vips.index("Review Reasons")])
    def test_new_reference_issues_reach_their_excel_rows(self):
        workbook = self._workbook(
            r'''
config router static
    edit 1
        set dstaddr "missing-address"
        set device "missing-route-interface"
    next
end

config system dhcp server
    edit 7
        set interface "missing-dhcp-interface"
    next
end

config system admin
    edit "operator"
        set accprofile "missing-profile"
    next
end

config firewall vip
    edit "broken-vip"
        set extintf "missing-vip-interface"
    next
end
''',
        )
        for sheet_name, identity_header, identity, expected in (
            ("Routes", "Route ID", 1, ("missing-address", "missing-route-interface")),
            ("DHCP Servers", "Server ID", 7, ("missing-dhcp-interface",)),
            ("Administrators", "Name", "operator", ("missing-profile",)),
            ("Virtual IPs", "Name", "broken-vip", ("missing-vip-interface",)),
        ):
            headers, rows = self._rows(workbook[sheet_name])
            row = next(item for item in rows if item[headers.index(identity_header)] == identity)
            self.assertEqual("REVIEW_REQUIRED", row[headers.index("Analysis Status")], sheet_name)
            reasons = str(row[headers.index("Review Reasons")] or "")
            for reference in expected:
                self.assertIn(reference, reasons, sheet_name)
    def test_malformed_nat_warning_reaches_nat_and_review_sheets(self):
        source = "source static any any destination static " + "103.230.127.102 " * 40
        workbook = self._workbook(
            f'''
config system interface
    edit "wan1"
        set ip {source}
    next
end

config firewall policy
    edit 1
        set name "bad nat"
        set dstintf "wan1"
        set nat enable
    next
end
'''
        )

        nat, nat_rows = self._rows(workbook["NAT Rules"])
        nat_row = nat_rows[0]
        reason = nat_row[nat.index("Review Reasons")]
        self.assertTrue(reason)
        self.assertNotIn(source, reason)

        review, review_rows = self._rows(workbook["Review Required"])
        review_row = next(row for row in review_rows if row[review.index("Category")] == "nat")
        self.assertEqual(reason, review_row[review.index("Issue / Review Reason")])
        severity_column = review.index("Severity") + 1
        self.assertEqual("solid", workbook["Review Required"].cell(4, severity_column).fill.fill_type)
    def test_any_nat_warning_reaches_nat_and_review_sheets(self):
        source = r'''
config firewall policy
    edit 1
        set name "any egress nat"
        set dstintf "any"
        set nat enable
    next
end
'''
        workbook = self._workbook(
            source,
        )

        nat, nat_rows = self._rows(workbook["NAT Rules"])
        nat_row = nat_rows[0]
        self.assertEqual("REVIEW_REQUIRED", nat_row[nat.index("Analysis Status")])
        reason = nat_row[nat.index("Review Reasons")]
        self.assertTrue(reason)

        review, review_rows = self._rows(workbook["Review Required"])
        review_row = next(row for row in review_rows if row[review.index("Category")] == "nat")
        self.assertEqual(reason, review_row[review.index("Issue / Review Reason")])
