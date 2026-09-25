from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from fwmigrate.vendors.fortigate.config import ExtractionConfig
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.export.excel_schema import SHEET_HEADERS, SHEET_ORDER
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.security.extraction import sanitize_source_attributes, sanitize_source_value
from fwmigrate.vendors.fortigate.validation.validator import validate_config


_SAMPLE_CONFIG = r"""
config system global
    set hostname "FG-TEST"
end

config system interface
    edit "wan-dhcp"
        set mode dhcp
    next
    edit "wan-pppoe"
        set mode pppoe
        set pppoe-username "source-user"
    next
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
        set allowaccess ping https
        config secondaryip
            edit 1
                set ip 192.0.2.2 255.255.255.0
                set ha-priority 20
            next
        end
    next
    edit "port2"
    next
    edit "agg1"
        set type aggregate
        set member "port1" "port2"
    next
    edit "vlan100"
        set type vlan
        set interface "agg1"
        set vlanid 100
        set ip 10.100.0.1 255.255.255.0
        set ip6 2001:db8:100::1/64
    next
end

config firewall service custom
    edit "multi-port"
        set tcp-portrange 80 443
        set udp-portrange 53
    next
    edit "single-port"
        set protocol TCP
        set tcp-portrange 8443
    next
end

config firewall service group
    edit "web-services"
        set member "multi-port" "single-port"
    next
end

config firewall address
    edit "inside"
        set subnet 10.100.0.0 255.255.255.0
    next
end

config firewall ippool
    edit "pool1"
        set startip 198.51.100.10
        set endip 198.51.100.20
    next
end

config firewall vip
    edit "lb-vip"
        set type server-load-balance
        set extip 198.51.100.30
        set extintf "wan-dhcp"
        set ldb-method round-robin
        config realservers
            edit 1
                set ip 10.100.0.10
                set port 443
            next
            edit 2
                set ip 10.100.0.11
                set port 443
            next
        end
    next
end

config vpn ipsec phase1-interface
    edit "VPN-HQ"
        set interface "vlan100"
        set remote-gw 203.0.113.2
        set psksecret "do-not-export-this-secret"
    next
end

config vpn ipsec phase2-interface
    edit "VPN-subnet"
        set phase1name "VPN-HQ"
        set src-subnet 10.100.0.0 255.255.255.0
        set dst-subnet 10.200.0.0 255.255.255.0
    next
    edit "VPN-range"
        set phase1name "VPN-HQ"
        set src-start-ip 10.100.0.10
        set src-end-ip 10.100.0.20
        set dst-start-ip 10.200.0.10
        set dst-end-ip 10.200.0.20
    next
end

config firewall policy
    edit 1
        set name "A policy name that is deliberately much longer than thirty-two characters"
        set srcintf "vlan100"
        set dstintf "port1"
        set srcaddr "inside"
        set dstaddr "all"
        set service "web-services"
        set action accept
        set schedule always
        set nat enable
    next
    edit 2
        set name "Collision policy name with a shared long prefix A"
        set srcintf "vlan100"
        set dstintf "port1"
        set srcaddr "inside"
        set dstaddr "all"
        set service "multi-port"
        set action accept
        set schedule always
        set nat enable
        set ippool enable
        set poolname "pool1"
    next
    edit 3
        set name "Collision policy name with a shared long prefix B"
        set srcintf "vlan100"
        set dstintf "port1" "port2"
        set srcaddr "inside"
        set dstaddr "all"
        set service "multi-port"
        set action accept
        set schedule always
        set nat enable
    next
    edit 4
        set name "short policy"
    next
end

config firewall unsupported-section
    edit "source-only"
        set unsupported-setting "preserve-me"
        append unsupported-list "first" "second"
        unset unsupported-setting
    next
end

config user local
    edit "alice"
        set type password
        set passwd "another-do-not-export-secret"
    next
end
"""


class ExcelReportTest(unittest.TestCase):
    def _workbook(self, config_text: str = _SAMPLE_CONFIG):
        extracted = extract_fortigate_config(
            parse_fortigate_config(config_text),
            config=ExtractionConfig(),
        )
        derived = build_derived_views(extracted.config)
        validation = validate_config(extracted.config, derived=derived)
        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
            source_name="sample.conf",
        )
        output.seek(0)
        return load_workbook(output, data_only=False)

    @staticmethod
    def _rows(sheet):
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        return headers, list(sheet.iter_rows(min_row=4, values_only=True))

    def test_predefined_admin_profile_and_web_proxy_are_not_flagged(self):
        workbook = self._workbook('''config system admin
    edit "built-in"
        set accprofile "super_admin"
    next
    edit "custom"
        set accprofile "missing-custom-profile"
    next
end
config firewall service custom
    edit "webproxy"
        set proxy enable
        set protocol ALL
        set tcp-portrange 0-65535:0-65535
    next
    edit "invalid-service"
        set protocol IP
        set tcp-portrange 443
    next
end
''')

        for sheet_name, identity_header, identities in (
            ("Administrators", "Name", {"built-in", "custom"}),
            ("Services", "Name", {"webproxy", "invalid-service"}),
        ):
            headers, rows = self._rows(workbook[sheet_name])
            by_name = {row[headers.index(identity_header)]: row for row in rows}
            for name in identities - {"custom", "invalid-service"}:
                self.assertEqual("EXTRACTED", by_name[name][headers.index("Analysis Status")])
                self.assertFalse(by_name[name][headers.index("Review Reasons")])

        admin_headers, admins = self._rows(workbook["Administrators"])
        custom = next(row for row in admins if row[admin_headers.index("Name")] == "custom")
        self.assertEqual("REVIEW_REQUIRED", custom[admin_headers.index("Analysis Status")])
        self.assertIn("missing-custom-profile", custom[admin_headers.index("Review Reasons")])

        service_headers, services = self._rows(workbook["Services"])
        invalid = next(row for row in services if row[service_headers.index("Name")] == "invalid-service")
        self.assertEqual("REVIEW_REQUIRED", invalid[service_headers.index("Analysis Status")])
        self.assertIn("inactive for protocol IP", invalid[service_headers.index("Review Reasons")])

        _, review_rows = self._rows(workbook["Review Required"])
        review_text = " ".join(str(row) for row in review_rows)
        self.assertNotIn("super_admin", review_text)
        self.assertNotIn("webproxy", review_text)
        self.assertIn("missing-custom-profile", review_text)

    def test_repeated_interface_edit_exports_one_extracted_interface(self):
        workbook = self._workbook('''config system interface
    edit "INFUAT-BIBDUAT"
        set type tunnel
        set snmp-index 34
        set interface "wan1"
    next
    edit "INFUAT-BIBDUAT"
        set snmp-index 11
    next
end
''')

        headers, rows = self._rows(workbook["Interfaces"])
        interfaces = [row for row in rows if str(row[headers.index("Name")]).endswith("INFUAT-BIBDUAT")]
        self.assertEqual(len(interfaces), 1)
        self.assertNotIn("Multiple explicit objects", str(interfaces[0][headers.index("Review Reasons")]))

        _, review_rows = self._rows(workbook["Review Required"])
        interface_reviews = [row for row in review_rows if "INFUAT-BIBDUAT" in str(row)]
        self.assertFalse(any("Multiple explicit objects" in str(row) for row in interface_reviews))

    def test_additional_settings_are_stable_across_repeated_exports(self):
        source = _SAMPLE_CONFIG + r'''
config firewall policy
    edit 101
        set name "deterministic-policy"
        set custom-policy-option "policy-marker"
    next
end
config vpn ipsec phase1-interface
    edit "deterministic-tunnel"
        set custom-tunnel-option "tunnel-marker"
    next
end
config vpn ipsec phase2-interface
    edit "deterministic-phase2"
        set phase1name "deterministic-tunnel"
        set custom-phase2-option "phase2-marker"
    next
end
config firewall vip
    edit "deterministic-vip"
        config realservers
            edit 1
                set ip 10.0.0.1
                set custom-realserver-option "realserver-marker"
            next
        end
    next
end
config system accprofile
    edit "deterministic-profile"
        set custom-profile-option "profile-marker"
        config fwgrp-permission
            set policy read-write
            set custom-permission-option "permission-marker"
        end
    next
end
config ips sensor
    edit "deterministic-sensor"
        config entries
            edit 9
                set rule 100
                set custom-entry-option "entry-marker"
            next
        end
    next
end
'''
        first = self._workbook(source)
        second = self._workbook(source)
        sheets = (
            "Policies", "VPN Tunnels", "VPN Phase 2", "VIP Real Servers",
            "Admin Profiles", "Admin Profile Permissions", "IPS Sensor Entries",
        )
        for name in sheets:
            self.assertEqual(
                list(first[name].iter_rows(values_only=True)),
                list(second[name].iter_rows(values_only=True)),
                name,
            )
        expected_markers = {
            "Policies": "policy-marker",
            "VPN Tunnels": "tunnel-marker",
            "VPN Phase 2": "phase2-marker",
            "VIP Real Servers": "realserver-marker",
            "Admin Profiles": "profile-marker",
            "Admin Profile Permissions": "permission-marker",
            "IPS Sensor Entries": "entry-marker",
        }
        for name, marker in expected_markers.items():
            values = [
                str(cell.value)
                for row in first[name].iter_rows()
                for cell in row
                if cell.value is not None
            ]
            self.assertTrue(any(marker in value for value in values), name)

    def test_additional_settings_preserve_unknown_source_order(self):
        workbook = self._workbook('''config firewall policy
    edit 101
        set name "ordered"
        set future-z last
        set future-a first
    next
end
''')
        headers, rows = self._rows(workbook["Policies"])
        row = next(row for row in rows if row[headers.index("Policy Name")] == "ordered")
        settings = str(row[headers.index("Additional Settings")])

        self.assertLess(settings.index("future_z"), settings.index("future_a"))

    def test_formula_like_source_text_is_exported_as_literal(self):
        workbook = self._workbook('''config firewall address
    edit "=1+1"
        set subnet 192.0.2.1 255.255.255.255
    next
end
''')
        sheet = workbook["Addresses"]
        headers, rows = self._rows(sheet)
        name_index = headers.index("Name")
        row_index = next(index for index, row in enumerate(rows, start=4) if str(row[name_index]).lstrip("'") == "=1+1")
        cell = sheet.cell(row_index, name_index + 1)

        self.assertNotEqual("f", cell.data_type)
        self.assertEqual("=1+1", str(cell.value).lstrip("'"))

    def test_ipv6_source_families_reach_their_workbook_rows(self):
        workbook = self._workbook('''config router static6
    edit 21
        set dst 2001:db8:1::/64
    next
end
config firewall ippool6
    edit v6-pool
        set startip 2001:db8::10
        set nat46 enable
    next
end
config firewall vip6
    edit v6-vip
        set extip 2001:db8::1
        set nat64 enable
    next
end
config firewall vipgrp6
    edit v6-group
        set member v6-vip
    next
end
''')
        for sheet_name, identity_header, identity in (
            ("Routes", "Route ID", 21),
            ("IP Pools", "Name", "v6-pool"),
            ("Virtual IPs", "Name", "v6-vip"),
            ("VIP Groups", "Name", "v6-group"),
        ):
            headers, rows = self._rows(workbook[sheet_name])
            row = next(row for row in rows if row[headers.index(identity_header)] == identity)
            self.assertEqual("ipv6", row[headers.index("Address Family")], sheet_name)

    def test_ssl_vpn_source_flags_and_nested_bookmarks_reach_workbook(self):
        workbook = self._workbook('''config vpn ssl client
    edit remote
        set psk client-secret
    next
end
config vpn ssl web user-bookmark
    edit alice
        config bookmarks
            edit app
                set apptype rdp
                set logon-password bookmark-secret
            next
        end
    next
end
''')
        clients, client_rows = self._rows(workbook["SSL VPN Clients"])
        self.assertEqual("Yes", client_rows[0][clients.index("PSK Configured")])
        bookmarks = list(workbook["SSL VPN Bookmarks"].iter_rows(min_row=4, values_only=True))
        self.assertEqual("User", bookmarks[0][0])

    def test_policy_based_ipsec_rows_keep_their_source_family(self):
        workbook = self._workbook('''config vpn ipsec phase1
    edit policy-vpn
        set interface wan1
        set psksecret policy-secret
    next
end
config vpn ipsec phase2
    edit policy-child
        set phase1name policy-vpn
        set src-addr-type subnet
        set src-subnet 10.1.0.0 255.255.255.0
    next
end
''')
        phase1_headers, phase1_rows = self._rows(workbook["Policy IPsec Phase 1"])
        phase1 = next(row for row in phase1_rows if row[phase1_headers.index("Name")] == "policy-vpn")
        self.assertEqual("Yes", phase1[phase1_headers.index("PSK Configured")])
        phase2_headers, phase2_rows = self._rows(workbook["Policy IPsec Phase 2"])
        phase2 = next(row for row in phase2_rows if row[phase2_headers.index("Name")] == "policy-child")
        self.assertEqual("10.1.0.0-10.1.0.255", phase2[phase2_headers.index("Source Range")])


    def test_typed_sheets_and_phase2_lifetimes_are_exported(self):
        source = _SAMPLE_CONFIG + r'''
config firewall service category
    edit "Applications"
        set comment "application services"
    next
end
config vpn ipsec phase2-interface
    edit "P2-LIFETIME"
        set phase1name "VPN-HQ"
        set keylifeseconds 3600
        set keylifekbs 10240
        set protocol 17
        set src-port 500
        set dst-port 4500
    next
end
config system dhcp server
    edit 1
        set interface "port1"
        config exclude-range
            edit 10
                set start-ip 192.0.2.100
                set end-ip 192.0.2.110
                set lease-time 600
            next
        end
    next
end
config vpn ssl web host-check-software
    edit "endpoint-av"
        set type av
        set os-type windows
        set version "1.2"
        set guid "host-check-guid"
        config check-item-list
            edit 1
                set action require
                set type file
                set target "C:\\Program Files\\agent.exe"
                set md5s "abc123"
                set version "1.0"
            next
            edit 2
                set action deny
                set type registry
                set target "HKLM\\Software\\Agent"
                set md5s "def456"
                set version "2.0"
            next
        end
    next
end
'''
        workbook = self._workbook(source)

        headers, rows = self._rows(workbook["VPN Phase 2"])
        row = next(row for row in rows if row[headers.index("Name")] == "P2-LIFETIME")
        self.assertEqual(3600, row[headers.index("Key Lifetime Seconds")])
        self.assertEqual(10240, row[headers.index("Key Lifetime KB")])
        self.assertNotIn("Keylife Seconds", headers)
        self.assertNotIn("Keylife KB", headers)
        self.assertEqual(17, row[headers.index("Protocol")])
        self.assertEqual(500, row[headers.index("Source Port")])
        self.assertEqual(4500, row[headers.index("Destination Port")])

        headers, rows = self._rows(workbook["Service Categories"])
        category = next(row for row in rows if row[headers.index("Name")] == "Applications")
        self.assertEqual("application services", category[headers.index("Description")])
        self.assertEqual("root", category[headers.index("VDOM")])

        headers, rows = self._rows(workbook["DHCP Exclude Ranges"])
        excluded = next(row for row in rows if row[headers.index("Range ID")] == 10)
        self.assertEqual((1, "port1", "192.0.2.100", "192.0.2.110", 600, "root"), tuple(
            excluded[headers.index(column)]
            for column in ("Server ID", "Interface", "Start IP", "End IP", "Lease Time", "VDOM")
        ))

        headers, rows = self._rows(workbook["SSL VPN Host Checks"])
        software = next(row for row in rows if row[headers.index("Name")] == "endpoint-av")
        self.assertEqual(2, software[headers.index("Check Item Count")])
        self.assertEqual("root", software[headers.index("VDOM")])

        headers, rows = self._rows(workbook["SSL VPN Host Check Items"])
        self.assertEqual(2, len(rows))
        by_id = {row[headers.index("ID")]: row for row in rows}
        self.assertEqual("endpoint-av", by_id[1][headers.index("Host Check")])
        self.assertEqual("endpoint-av", by_id[2][headers.index("Host Check")])
        for item_id, action, kind, target, md5, version in (
            (1, "require", "file", r"C:\Program Files\agent.exe", "abc123", "1.0"),
            (2, "deny", "registry", r"HKLM\Software\Agent", "def456", "2.0"),
        ):
            row = by_id[item_id]
            self.assertEqual((action, kind, target, md5, version, "root"), tuple(
                row[headers.index(column)]
                for column in ("Action", "Type", "Target", "MD5s", "Version", "VDOM")
            ))

        for sheet in ("Service Categories", "DHCP Exclude Ranges", "SSL VPN Host Checks", "SSL VPN Host Check Items"):
            self.assertIn(sheet, SHEET_ORDER)
            self.assertIn(sheet, SHEET_HEADERS)
            self.assertIn(sheet, workbook.sheetnames)

    def test_vpn_phase1_ike_fields_are_exported_without_defaults_or_secrets(self):
        source = r'''
config vpn ipsec phase1-interface
    edit "VPN-IKE-TEST"
        set interface "wan1"
        set type static
        set remote-gw 203.0.113.10
        set ike-version 2
        set mode main
        set authmethod signature
        set authmethod-remote psk
        set proposal aes256-sha256 aes128-sha256
        set dhgrp 14 19
        set keylife 28800
        set nattraversal forced
        set dpd on-idle
        set dpd-retrycount 5
        set dpd-retryinterval 10
        set local-gw 192.0.2.1
        set localid "site-a"
        set localid-type fqdn
        set peerid "site-b"
        set certificate "vpn-cert"
        set eap enable
        set psksecret-remote "TEST-SECRET-DO-NOT-EXPORT"
    next
end
'''
        extracted = extract_fortigate_config(
            parse_fortigate_config(source),
            config=ExtractionConfig(),
        )
        item = extracted.config.ipsec_phase1[0]
        self.assertEqual("2", item.ike_version)
        self.assertEqual("main", item.mode)
        self.assertEqual("signature", item.authmethod)
        self.assertEqual("psk", item.authmethod_remote)
        self.assertEqual(["aes256-sha256", "aes128-sha256"], item.proposal)
        self.assertEqual([14, 19], item.dhgrp)
        self.assertEqual(28800, item.keylife)
        self.assertEqual("forced", item.nattraversal)
        self.assertEqual("on-idle", item.dpd)
        self.assertEqual(5, item.dpd_retrycount)
        self.assertEqual("10", item.dpd_retryinterval)
        self.assertEqual("site-a", item.localid)
        self.assertEqual("fqdn", item.localid_type)
        self.assertEqual("site-b", item.peerid)
        self.assertEqual(["vpn-cert"], item.certificate)
        self.assertTrue(item.psk_configured)
        self.assertNotIn("TEST-SECRET-DO-NOT-EXPORT", str(item.model_dump()))
        self.assertNotIn("TEST-SECRET-DO-NOT-EXPORT", str(item.raw_extra))

        workbook = self._workbook(source)
        headers, rows = self._rows(workbook["VPN Tunnels"])
        self.assertNotIn("Proposal", headers)
        self.assertNotIn("DH Groups", headers)
        self.assertNotIn("Key Lifetime", headers)
        self.assertNotIn("DPD", headers)
        for header in (
            "IKE Version", "IKE Mode", "Authentication Method",
            "Remote Authentication Method", "PSK Configured",
            "Phase 1 Proposal", "Phase 1 DH Groups", "Key Lifetime (Seconds)",
            "NAT Traversal", "DPD Mode", "DPD Retry Count", "DPD Retry Interval",
            "Local ID", "Local ID Type", "Peer ID", "Certificate",
        ):
            self.assertIn(header, headers)
        self.assertTrue(workbook["VPN Tunnels"].column_dimensions[
            get_column_letter(headers.index("Additional Settings") + 1)
        ].hidden)
        row = rows[0]
        self.assertEqual("2", row[headers.index("IKE Version")])
        self.assertEqual("main", row[headers.index("IKE Mode")])
        self.assertEqual("signature", row[headers.index("Authentication Method")])
        self.assertEqual("psk", row[headers.index("Remote Authentication Method")])
        self.assertEqual("Yes", row[headers.index("PSK Configured")])
        self.assertEqual("aes256-sha256\naes128-sha256", row[headers.index("Phase 1 Proposal")])
        self.assertEqual("14\n19", row[headers.index("Phase 1 DH Groups")])
        self.assertEqual(28800, row[headers.index("Key Lifetime (Seconds)")])
        self.assertEqual(5, row[headers.index("DPD Retry Count")])
        self.assertEqual("10", row[headers.index("DPD Retry Interval")])
        self.assertIn("eap", str(row[headers.index("Additional Settings")]))
        self.assertNotIn("TEST-SECRET-DO-NOT-EXPORT", str(row))

    def test_vpn_phase1_missing_ike_fields_remain_blank(self):
        workbook = self._workbook(
            """
config vpn ipsec phase1-interface
    edit "MINIMAL"
        set interface "wan1"
    next
end
"""
        )
        headers, rows = self._rows(workbook["VPN Tunnels"])
        row = rows[0]
        for header in (
            "IKE Version", "Key Lifetime (Seconds)", "NAT Traversal",
            "DPD Retry Count",
        ):
            self.assertIsNone(row[headers.index(header)])
        self.assertEqual("No", row[headers.index("PSK Configured")])

    def test_ips_source_fields_and_nested_exemptions(self):
        source = r'''
config ips sensor
    edit "ips-sensor"
        config entries
            edit 42
                set rule 1001
                set cve CVE-2026-0001
                set default-action pass
                set default-status enable
                set action block
                set status enable
                set last-modified "2026-01-02 03:04:05"
                set custom-option "preserve-me"
                config exempt-ip
                    edit 7
                        set src-ip 192.0.2.10
                        set dst-ip 198.51.100.10
                    next
                end
            next
        end
    next
end
'''
        extracted = extract_fortigate_config(
            parse_fortigate_config(source),
            config=ExtractionConfig(),
        )
        entry = extracted.config.ips_sensors[0].entries[0]
        self.assertEqual("2026-01-02 03:04:05", entry.last_modified)
        self.assertEqual("pass", entry.default_action)
        self.assertEqual("enable", entry.default_status)
        self.assertEqual("block", entry.action)
        self.assertEqual("enable", entry.status)
        self.assertEqual([(7, "192.0.2.10", "198.51.100.10")], [
            (item.id, item.src_ip, item.dst_ip) for item in entry.exempt_ips
        ])

        workbook = self._workbook(source)
        headers, rows = self._rows(workbook["IPS Sensor Entries"])
        self.assertIn("Default Action Filter", headers)
        self.assertIn("Default Status Filter", headers)
        self.assertIn("Last Modified Filter", headers)
        self.assertNotIn("Default Action", headers)
        self.assertNotIn("Default Status", headers)
        self.assertIn("Action", headers)
        self.assertIn("Status", headers)
        self.assertIn("Additional Settings", headers)
        self.assertTrue(workbook["IPS Sensor Entries"].column_dimensions[
            get_column_letter(headers.index("Additional Settings") + 1)
        ].hidden)
        row = rows[0]
        self.assertEqual("pass", row[headers.index("Default Action Filter")])
        self.assertEqual("enable", row[headers.index("Default Status Filter")])
        self.assertEqual("block", row[headers.index("Action")])
        self.assertEqual("enable", row[headers.index("Status")])
        self.assertEqual("2026-01-02 03:04:05", row[headers.index("Last Modified Filter")])
        self.assertEqual("root", row[headers.index("VDOM")])
        self.assertIn("preserve-me", str(row[headers.index("Additional Settings")]))

        headers, rows = self._rows(workbook["IPS Exempt IPs"])
        self.assertEqual(1, len(rows))
        row = rows[0]
        self.assertEqual("ips-sensor", row[headers.index("Sensor")])
        self.assertEqual(42, row[headers.index("Entry ID")])
        self.assertEqual(7, row[headers.index("Exempt IP ID")])
        self.assertEqual("192.0.2.10", row[headers.index("Source IP")])
        self.assertEqual("198.51.100.10", row[headers.index("Destination IP")])
        self.assertEqual("root", row[headers.index("VDOM")])

    def test_ntp_config_settings_are_separate_from_servers(self):
        workbook = self._workbook(
            """
            config system ntp
                set ntpsync enable
                set server-mode enable
                set interface "port1" "port4"
            end
            """
        )

        settings_headers, settings_rows = self._rows(workbook["NTP Settings"])
        self.assertNotIn("Object", settings_headers)
        self.assertEqual(
            {row[settings_headers.index("Setting")] for row in settings_rows},
            {"ntpsync", "server_mode", "interface"},
        )
        self.assertTrue(all(row[settings_headers.index("Analysis Status")] == "EXTRACTED" for row in settings_rows))
        self.assertTrue(all(not row[settings_headers.index("Review Reasons")] for row in settings_rows))
        server_headers, server_rows = self._rows(workbook["NTP Servers"])
        self.assertEqual(server_rows, [])
        self.assertEqual(
            list(SHEET_HEADERS["NTP Servers"]),
            server_headers,
        )

    def test_unnamed_settings_do_not_inherit_vdom_validation_issues(self):
        workbook = self._workbook('''config system global
    set hostname firewall
end
config system dns
    set primary 192.0.2.53
    set secondary 192.0.2.54
end
config system ntp
    set ntpsync enable
    set server-mode enable
end
config firewall service custom
    edit "invalid-service"
        set protocol ALL
        set tcp-portrange 80
    next
end
''')

        for sheet_name in ("System Settings", "DNS Settings", "NTP Settings"):
            headers, rows = self._rows(workbook[sheet_name])
            self.assertTrue(rows, sheet_name)
            self.assertTrue(all(row[headers.index("Analysis Status")] == "EXTRACTED" for row in rows), sheet_name)
            self.assertTrue(all(not row[headers.index("Review Reasons")] for row in rows), sheet_name)

        headers, rows = self._rows(workbook["DNS Settings"])
        self.assertEqual({"primary", "secondary"}, {row[headers.index("Setting")] for row in rows})

        headers, rows = self._rows(workbook["Services"])
        service = next(row for row in rows if row[headers.index("Name")] == "invalid-service")
        self.assertEqual("REVIEW_REQUIRED", service[headers.index("Analysis Status")])

        headers, rows = self._rows(workbook["Review Required"])
        self.assertTrue(any("invalid-service" in str(row) for row in rows))

    def test_ntp_server_objects_use_server_ids(self):
        workbook = self._workbook(
            """
            config system ntp
                config ntpserver
                    edit 1
                        set server "192.0.2.10"
                    next
                    edit 2
                        set server "192.0.2.20"
                    next
                end
            end
            """
        )

        headers, rows = self._rows(workbook["NTP Servers"])
        server_id = headers.index("Server ID")
        setting = headers.index("Setting")
        value = headers.index("Value")
        source_path = headers.index("Source Path")
        self.assertEqual(
            {(row[server_id], row[setting], row[value], row[source_path]) for row in rows},
            {
                ("1", "server", "192.0.2.10", "system ntp ntpserver"),
                ("2", "server", "192.0.2.20", "system ntp ntpserver"),
            },
        )

    def test_ntp_settings_and_servers_do_not_duplicate_rows(self):
        workbook = self._workbook(
            """
            config system ntp
                set ntpsync enable
                config ntpserver
                    edit 1
                        set server "192.0.2.10"
                    next
                end
            end
            """
        )

        settings_headers, settings_rows = self._rows(workbook["NTP Settings"])
        server_headers, server_rows = self._rows(workbook["NTP Servers"])
        self.assertEqual(
            [(row[settings_headers.index("Setting")], row[settings_headers.index("Value")]) for row in settings_rows],
            [("ntpsync", "enable")],
        )
        self.assertEqual(
            [(row[server_headers.index("Setting")], row[server_headers.index("Value")]) for row in server_rows],
            [("server", "192.0.2.10")],
        )

    def test_source_and_topology_contract(self):
        workbook = self._workbook()
        interfaces, rows = self._rows(workbook["Interfaces"])
        name = interfaces.index("Name")
        names = [row[name] for row in rows]
        self.assertEqual(
            {"◆ agg1", "├─ ● port1", "├─ ● port2", "└─ ▣ vlan100", "   └─ ◈ VPN-HQ", "● wan-dhcp", "● wan-pppoe"},
            set(names),
        )
        by_name = {row[name]: row for row in rows}
        self.assertEqual("dhcp", by_name["● wan-dhcp"][interfaces.index("Addressing Mode")])
        self.assertIsNone(by_name["├─ ● port2"][interfaces.index("Addressing Mode")])
        self.assertEqual("agg1", by_name["└─ ▣ vlan100"][interfaces.index("Parent Interface")])
        self.assertEqual("port1\nport2", by_name["└─ ▣ vlan100"][interfaces.index("Physical Interfaces")])
        self.assertLess(names.index("◆ agg1"), names.index("├─ ● port1"))
        self.assertLess(names.index("├─ ● port1"), names.index("├─ ● port2"))
        self.assertLess(names.index("├─ ● port2"), names.index("└─ ▣ vlan100"))
        self.assertLess(names.index("└─ ▣ vlan100"), names.index("   └─ ◈ VPN-HQ"))
        vpn = by_name["   └─ ◈ VPN-HQ"]
        self.assertEqual("vlan100", vpn[interfaces.index("Parent Interface")])
        self.assertEqual("agg1", vpn[interfaces.index("Aggregate")])
        self.assertEqual("port1\nport2", vpn[interfaces.index("Physical Interfaces")])
        self.assertEqual("VPN-HQ\nvlan100\nagg1", vpn[interfaces.index("Topology Path")])

        secondary, rows = self._rows(workbook["Interface Secondary IPs"])
        row = rows[0]
        self.assertEqual("192.0.2.2 255.255.255.0", row[secondary.index("IP / Prefix")])
        self.assertEqual(20, row[secondary.index("HA Priority")])

        summary = workbook["Summary"]
        summary_values = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(1, summary.max_row + 1)}
        self.assertEqual("FG-TEST", summary_values["Hostname"])
        self.assertEqual("Yes", summary_values["IPv6 Explicit Configuration Present"])
        extracted = extract_fortigate_config(
            parse_fortigate_config(_SAMPLE_CONFIG),
            config=ExtractionConfig(),
        )
        self.assertEqual(
            len(extracted.config.interfaces),
            summary_values["Source Interfaces"],
        )
        self.assertNotIn("Interfaces", summary_values)
        source_interfaces_row = next(
            row for row in range(1, summary.max_row + 1)
            if summary.cell(row, 1).value == "Source Interfaces"
        )
        self.assertEqual(
            "#'Interfaces'!A1",
            summary.cell(source_interfaces_row, 3).hyperlink.target,
        )

    def test_topology_edge_rows_remain_visible(self):
        workbook = self._workbook(
            r'''
config system interface
    edit "physical"
    next
    edit "agg-missing"
        set type aggregate
        set member "missing-member"
    next
    edit "orphan"
        set type vlan
        set interface "missing-parent"
    next
    edit "cycle-a"
        set interface "cycle-b"
    next
    edit "cycle-b"
        set interface "cycle-a"
    next
end

config vpn ipsec phase1-interface
    edit "VPN-physical"
        set interface "physical"
    next
end
''',
        )
        headers, rows = self._rows(workbook["Interfaces"])
        name = headers.index("Name")
        by_name = {row[name]: row for row in rows}
        self.assertIn("▣ orphan", by_name)
        self.assertIn("◆ agg-missing", by_name)
        self.assertIn("◇ cycle-a", by_name)
        self.assertIn("└─ ◇ cycle-b", by_name)
        self.assertTrue(by_name["◆ agg-missing"][headers.index("Topology Issues")])
        self.assertEqual("REVIEW_REQUIRED", by_name["◇ cycle-a"][headers.index("Analysis Status")])
        self.assertEqual("physical", by_name["└─ ◈ VPN-physical"][headers.index("Parent Interface")])
        self.assertLess(
            [row[name] for row in rows].index("● physical"),
            [row[name] for row in rows].index("└─ ◈ VPN-physical"),
        )

    def test_semantics_and_source_preservation(self):
        workbook = self._workbook()
        services, rows = self._rows(workbook["Services"])
        service_rows = {row[services.index("Name")]: row for row in rows}
        generated = [row for row in rows if row[services.index("Source Service")] == "multi-port"]
        self.assertTrue(all(row[services.index("Generated")] == "Yes" for row in generated))
        self.assertIn("80", str(service_rows["multi-port-1"][services.index("Destination Port")]))

        groups, rows = self._rows(workbook["Service Groups"])
        group = {row[groups.index("Name")]: row for row in rows}
        self.assertEqual("Yes", group["multi-port"][groups.index("Generated")])
        self.assertEqual("No", group["web-services"][groups.index("Generated")])

        policies, rows = self._rows(workbook["Policies"])
        self.assertIsNotNone(rows[0][policies.index("Source Name")])
        self.assertEqual(rows[0][policies.index("Source Name")], rows[0][policies.index("Policy Name")])
        self.assertIsNotNone(rows[1][policies.index("Source Name")])
        self.assertIsNotNone(rows[2][policies.index("Source Name")])
        self.assertNotEqual(
            rows[1][policies.index("Policy Name")],
            rows[2][policies.index("Policy Name")],
        )
        self.assertTrue(all(rows[index][policies.index("Policy Name")] == rows[index][policies.index("Source Name")] for index in (1, 2)))
        self.assertNotIn(
            "Normalized policy name collision",
            rows[1][policies.index("Review Reasons")] or "",
        )
        self.assertNotIn(
            "Normalized policy name collision",
            rows[2][policies.index("Review Reasons")] or "",
        )
        self.assertEqual(rows[3][policies.index("Policy Name")], rows[3][policies.index("Source Name")])
        self.assertEqual("pool1", rows[1][policies.index("IP Pool Name")])
        self.assertIsNone(rows[2][policies.index("SNAT Address")])

        nat, rows = self._rows(workbook["NAT Rules"])
        nat_by_rule = {row[nat.index("Rule #")]: row for row in rows}
        self.assertEqual("192.0.2.1", nat_by_rule[1][nat.index("SNAT Address")])
        self.assertEqual("198.51.100.10-198.51.100.20", nat_by_rule[2][nat.index("SNAT Address")])
        self.assertIsNone(nat_by_rule[3][nat.index("SNAT Address")])

        vips, rows = self._rows(workbook["Virtual IPs"])
        vip = next(row for row in rows if row[vips.index("Name")] == "lb-vip")
        self.assertEqual("round-robin", vip[vips.index("Load Balance Method")])
        self.assertEqual(2, vip[vips.index("Real Server Count")])

        phase2, rows = self._rows(workbook["VPN Phase 2"])
        by_name = {row[phase2.index("Name")]: row for row in rows}
        self.assertEqual("10.100.0.0-10.100.0.255", by_name["VPN-subnet"][phase2.index("Source Range")])
        self.assertEqual("10.100.0.10-10.100.0.20", by_name["VPN-range"][phase2.index("Source Range")])

        inventory = workbook["FortiGate Source Inventory"]
        inventory_headers, inventory_rows = self._rows(inventory)
        self.assertEqual(
            list(SHEET_HEADERS["FortiGate Source Inventory"]),
            inventory_headers,
        )
        self.assertEqual("root", inventory_rows[0][inventory_headers.index("VDOM")])
        self.assertIn(
            "port1",
            str(next(row for row in inventory_rows if row[inventory_headers.index("Setting")] == "ha-priority")[inventory_headers.index("Parent / Subsection")]),
        )
        operations = {row[inventory_headers.index("Operation")] for row in inventory_rows}
        self.assertTrue({"set", "append", "unset"}.issubset(operations))
        self.assertEqual(
            "TYPED",
            next(row for row in inventory_rows if row[inventory_headers.index("Object")] == "port2")[inventory_headers.index("Extraction Status")],
        )
        self.assertEqual(
            "SOURCE_ONLY",
            next(row for row in inventory_rows if row[inventory_headers.index("Object")] == "source-only")[inventory_headers.index("Extraction Status")],
        )
        self.assertIsNone(
            next(row for row in inventory_rows if row[inventory_headers.index("Object")] == "port2")[inventory_headers.index("Operation")],
        )
        unsupported, unsupported_rows = self._rows(workbook["Unsupported"])
        source_only = next(row for row in unsupported_rows if row[unsupported.index("Section")] == "firewall unsupported-section")
        self.assertEqual("FortiGate Source Inventory", source_only[unsupported.index("Raw Capture Location")])
        self.assertNotIn("do-not-export-this-secret", inventory_values := [cell.value for row in inventory.iter_rows() for cell in row])
        self.assertNotIn("another-do-not-export-secret", inventory_values)
        self.assertIn("2001:db8:100::1/64", inventory_values)
        self.assertIn("preserve-me", inventory_values)

        extracted = extract_fortigate_config(
            parse_fortigate_config(_SAMPLE_CONFIG),
            config=ExtractionConfig(),
        )
        self.assertEqual(
            sum(max(1, len(record.commands)) for record in extracted.source_objects),
            len(inventory_rows),
        )
        self.assertEqual("A4", inventory.freeze_panes)
        self.assertEqual(f"A3:{get_column_letter(inventory.max_column)}{len(inventory_rows) + 3}", inventory.auto_filter.ref)
        self.assertEqual(60, inventory.column_dimensions["I"].width)
        self.assertEqual("#'Summary'!A1", inventory[2][inventory.max_column - 1].hyperlink.target)
        self.assertTrue(
            all(
                cell.fill.fill_type is None
                for row in inventory.iter_rows(min_row=4)
                for cell in row
            )
        )

        unknown_workbook = self._workbook(
            r'''
config firewall unsupported-section
    edit "unknown"
        mystery-setting "opaque"
    next
end
''',
        )
        unknown_headers, unknown_rows = self._rows(unknown_workbook["FortiGate Source Inventory"])
        self.assertIn("unknown", {row[unknown_headers.index("Operation")] for row in unknown_rows})

    def test_policy_reports_preserve_ipv4_and_ipv6_address_families(self):
        source = _SAMPLE_CONFIG + r'''
config firewall address6
    edit "SRC-V6"
        set ip6 2001:db8:1::/64
    next
    edit "DST-V6"
        set ip6 2001:db8:2::/64
    next
end
config firewall policy
    edit 60
        set srcintf "port1"
        set dstintf "port1"
        set srcaddr6 "SRC-V6"
        set dstaddr6 "DST-V6"
        set srcaddr6-negate enable
        set dstaddr6-negate disable
        set service "multi-port"
        set schedule always
        set action accept
        set nat enable
    next
    edit 61
        set srcintf "port1"
        set dstintf "port1"
        set srcaddr "inside"
        set srcaddr6 "SRC-V6"
        set dstaddr "all"
        set dstaddr6 "DST-V6"
        set service "multi-port"
        set schedule always
        set action accept
        set nat enable
    next
end
'''
        workbook = self._workbook(source)
        headers, rows = self._rows(workbook["Policies"])
        by_rule = {row[headers.index("Rule #")]: row for row in rows}
        for column in ("Source IPv6 Addresses", "Destination IPv6 Addresses"):
            self.assertIn(column, headers)
        ipv6_only = by_rule[60]
        self.assertIn(ipv6_only[headers.index("Source Addresses")], (None, ""))
        self.assertEqual("SRC-V6", ipv6_only[headers.index("Source IPv6 Addresses")])
        self.assertIn(ipv6_only[headers.index("Destination Addresses")], (None, ""))
        self.assertEqual("DST-V6", ipv6_only[headers.index("Destination IPv6 Addresses")])
        self.assertEqual("enable", ipv6_only[headers.index("Source IPv6 Address Negate")])
        self.assertEqual("disable", ipv6_only[headers.index("Destination IPv6 Address Negate")])

        mixed = by_rule[61]
        self.assertEqual("inside", mixed[headers.index("Source Addresses")])
        self.assertEqual("SRC-V6", mixed[headers.index("Source IPv6 Addresses")])
        self.assertEqual("all", mixed[headers.index("Destination Addresses")])
        self.assertEqual("DST-V6", mixed[headers.index("Destination IPv6 Addresses")])

        nat_headers, nat_rows = self._rows(workbook["NAT Rules"])
        nat_by_rule = {row[nat_headers.index("Rule #")]: row for row in nat_rows}
        self.assertEqual("SRC-V6", nat_by_rule[60][nat_headers.index("Source IPv6 Addresses")])
        self.assertEqual("DST-V6", nat_by_rule[60][nat_headers.index("Destination IPv6 Addresses")])









    def test_summary_uses_explicit_fortios_version_header(self):
        workbook = self._workbook(
            "#config-version=FG100E-7.2.13-FW-build1762-260128:opmode=0:vdom=0\n"
        )
        summary = workbook["Summary"]
        summary_values = {
            summary.cell(row, 1).value: summary.cell(row, 2).value
            for row in range(1, summary.max_row + 1)
        }
        self.assertEqual("7.2.13", summary_values["FortiOS Version"])

        workbook = self._workbook("# no config version\n")
        self.assertIsNone(workbook["Summary"]["B5"].value)




if __name__ == "__main__":
    unittest.main()

