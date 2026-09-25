from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook
from pydantic import BaseModel

from fwmigrate.vendors.fortigate.config import ExtractionConfig
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.export.excel_schema import SHEET_HEADERS
from fwmigrate.vendors.fortigate.extraction.coverage import (
    _SPECS,
    build_typed_source_inventory,
    extraction_status,
    find_typed_source_object,
    typed_source_paths,
)
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.extraction.source_inventory import SourceObjectRecord
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.section_registry import registered_sections
from fwmigrate.vendors.fortigate.validation.validator import validate_config


class CoverageAccountingTest(unittest.TestCase):
    def _report(self, source: str):
        extracted = extract_fortigate_config(parse_fortigate_config(source), config=ExtractionConfig())
        derived = build_derived_views(extracted.config)
        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validate_config(extracted.config, derived=derived),
            output=output,
        )
        output.seek(0)
        return extracted, load_workbook(output, data_only=False)

    @staticmethod
    def _rows(sheet):
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        return headers, list(sheet.iter_rows(min_row=4, values_only=True))

    def test_registry_collects_unique_pydantic_objects_without_mutating_source(self):
        extracted, _ = self._report('''
config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
    next
end
''')
        before = extracted.config.model_dump()
        inventory = build_typed_source_inventory(extracted.config)
        self.assertEqual(len(_SPECS), len(typed_source_paths()))
        objects = [item for items in inventory.values() for item in items]
        self.assertTrue(all(isinstance(item.model, BaseModel) for item in objects))
        self.assertTrue(all(item.identity[1] == path for path, items in inventory.items() for item in items))
        identities = [item.identity for item in objects]
        self.assertEqual(len(identities), len(set(identities)))
        self.assertEqual(before, extracted.config.model_dump())

    def test_workbook_keeps_registration_model_coverage_and_raw_extra_separate(self):
        extracted, workbook = self._report('''
config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
        set mtu 1500
        set future-option preserve-me
    next
    edit "port2"
    next
end
config system dhcp server
    edit 1
        config exclude-range
            edit 4
                set start-ip 192.0.2.20
                set end-ip 192.0.2.30
            next
        end
    next
end
config firewall unsupported-section
    edit "source-only"
        set future-setting preserve-me
    next
end
''')
        coverage_headers, coverage_rows = self._rows(workbook["Extraction Coverage"])
        coverage = {row[coverage_headers.index("Source Section")]: row for row in coverage_rows}
        interface = coverage["system interface"]
        self.assertEqual("REGISTERED", interface[coverage_headers.index("Primitive Registration")])
        self.assertEqual("TYPED", interface[coverage_headers.index("Model Coverage")])
        self.assertEqual(len(extracted.config.interfaces), interface[coverage_headers.index("Typed Objects")])
        self.assertEqual(1, interface[coverage_headers.index("Raw Extra Objects")])
        self.assertGreaterEqual(interface[coverage_headers.index("Raw Extra Entries")], 2)
        self.assertEqual(0, interface[coverage_headers.index("Source-only Records")])

        inventory_headers, inventory_rows = self._rows(workbook["FortiGate Source Inventory"])
        port1 = next(row for row in inventory_rows if row[inventory_headers.index("Object")] == "port1")
        self.assertEqual("TYPED_WITH_RAW_EXTRA", port1[inventory_headers.index("Extraction Status")])
        self.assertGreaterEqual(port1[inventory_headers.index("Raw Extra Entries")], 2)
        excluded = next(row for row in inventory_rows if row[inventory_headers.index("Source Path")] == "system dhcp server exclude-range")
        self.assertEqual("TYPED", excluded[inventory_headers.index("Extraction Status")])
        generic = next(row for row in inventory_rows if row[inventory_headers.index("Object")] == "source-only")
        self.assertEqual("GENERIC", generic[inventory_headers.index("Primitive Registration")])
        self.assertEqual("SOURCE_ONLY", generic[inventory_headers.index("Extraction Status")])
        self.assertEqual(list(SHEET_HEADERS["FortiGate Source Inventory"]), inventory_headers)
        self.assertEqual(f"A3:L{len(inventory_rows) + 3}", workbook["FortiGate Source Inventory"].auto_filter.ref)

        unsupported_headers, unsupported_rows = self._rows(workbook["Unsupported"])
        source_only = next(row for row in unsupported_rows if row[unsupported_headers.index("Section")] == "firewall unsupported-section")
        self.assertEqual("SOURCE_ONLY", source_only[unsupported_headers.index("Model Coverage")])
        self.assertNotIn("system interface", {row[unsupported_headers.index("Section")] for row in unsupported_rows})
        source_only_coverage = coverage["firewall unsupported-section"]
        self.assertEqual(1, source_only_coverage[coverage_headers.index("Source-only Records")])
        self.assertEqual(0, source_only_coverage[coverage_headers.index("Typed Objects")])
        coverage_values = [cell.value for row in workbook["Extraction Coverage"].iter_rows() for cell in row]
        self.assertNotIn("preserve-me", coverage_values)

    def test_fully_typed_input_counts_model_instances_not_source_records(self):
        extracted, workbook = self._report('''
config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
    next
    edit "port2"
        set ip 192.0.2.2 255.255.255.0
    next
end
''')
        headers, rows = self._rows(workbook["Extraction Coverage"])
        row = next(item for item in rows if item[headers.index("Source Section")] == "system interface")
        self.assertEqual("REGISTERED", row[headers.index("Primitive Registration")])
        self.assertEqual("TYPED", row[headers.index("Model Coverage")])
        self.assertEqual(len(extracted.config.interfaces), row[headers.index("Typed Objects")])
        self.assertEqual(0, row[headers.index("Raw Extra Objects")])
        self.assertEqual(0, row[headers.index("Raw Extra Entries")])
        self.assertEqual(0, row[headers.index("Source-only Records")])

    def test_model_gap_is_not_source_only_and_registry_does_not_mean_typed(self):
        empty = SourceObjectRecord(
            vdom="root", source_path="system interface", object_name="missing-model",
            parent_objects=(), values={}, explicit_fields=(), unset_fields=(), commands=(),
            start_line_number=None, end_line_number=None,
        )
        self.assertIsNone(find_typed_source_object({}, empty))
        self.assertEqual("MODEL_GAP", extraction_status(True, None))
        self.assertEqual("SOURCE_ONLY", extraction_status(False, None))
        self.assertTrue(registered_sections())


if __name__ == "__main__":
    unittest.main()
