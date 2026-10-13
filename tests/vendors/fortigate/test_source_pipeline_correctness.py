import io

import pytest
from openpyxl import load_workbook

from fwmigrate.source_reporting import ExcelExportProfile
from fwmigrate.vendors.fortigate.extraction.coverage import (
    build_typed_source_identity_index,
    build_typed_source_inventory,
    find_typed_source_object,
)
from fwmigrate.vendors.fortigate.extraction.section_index import SectionIndex
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.relationships.references import ReferenceKind
from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter
from fwmigrate.vendors.fortigate.validation.index import ValidationIssueIndex


def test_global_scope_has_canonical_paths_and_original_provenance():
    source = '''config global
    config system interface
        edit port1
            set vdom tenant-a
            config secondaryip
                edit 1
                    set ip 192.0.2.2 255.255.255.0
                next
            end
        next
    end
    config system admin
        edit operator
            set password never-export-this
        next
    end
end
config vdom
    edit tenant-a
        config firewall address
            edit same
            next
        end
    next
    edit root
        config firewall address
            edit same
            next
        end
    next
end
'''
    index = SectionIndex.build(parse_fortigate_config(source))
    assert [(item.source_path, item.vdom, item.cli_scope) for item in index.entries] == [
        ("system interface", None, "global"),
        ("system interface secondaryip", None, "global"),
        ("system admin", None, "global"),
        ("firewall address", "tenant-a", "vdom"),
        ("firewall address", "root", "vdom"),
    ]
    analysis = FortiGateSourceReporter().analyze_source(source)
    records = analysis.extracted.source_objects
    assert records[1].cli_path == ("global", "system interface", "port1", "secondaryip", "1")
    assert records[0].start_line_number == 3
    assert records[0].end_line_number == 10
    inventory = build_typed_source_inventory(analysis.extracted.config)
    identity_index = build_typed_source_identity_index(inventory)
    assert all(find_typed_source_object(identity_index, record) is not None for record in records)
    assert "never-export-this" not in repr(records)


@pytest.mark.parametrize("ownership,expected,explicit,unset", [
    ("set vdom tenant-a", "tenant-a", True, False),
    ("set vdom root", "root", True, False),
    ("", None, False, False),
    ("set vdom tenant-a\n            unset vdom", None, False, True),
    ("set vdom tenant-a\n            set vdom tenant-b", "tenant-b", True, False),
])
def test_global_interface_ownership_is_explicit(ownership, expected, explicit, unset):
    analysis = FortiGateSourceReporter().analyze_source(f'''config global
    config system interface
        edit wan1
            {ownership}
            set ip 192.0.2.1 255.255.255.0
        next
    end
end
''')
    interface = analysis.extracted.config.interfaces[0]
    assert interface.vdom == expected
    assert ("vdom" in interface.explicit_fields) == explicit
    assert ("vdom" in interface.unset_fields) == unset
    assert interface.cli_scope == "global" and interface.cli_vdom is None
    assert bool([issue for issue in analysis.validation.issues if issue.field == "vdom"]) == (expected is None)
    before = analysis.extracted.config.model_dump()
    preview = FortiGateSourceReporter().build_preview(analysis)
    assert preview["sections"]["interfaces"][0]["vdom"] == expected
    assert analysis.extracted.config.model_dump() == before


def test_global_assignment_drives_same_vdom_references_and_nat():
    source = '''config global
    config system interface
        edit wan1
            set vdom tenant-a
            set ip 192.0.2.1 255.255.255.0
        next
    end
end
config vdom
    edit tenant-a
        config firewall policy
            edit 1
                set dstintf wan1
                set nat enable
            next
        end
    next
    edit tenant-b
        config firewall policy
            edit 2
                set dstintf wan1
                set nat enable
            next
        end
    next
end
'''
    analysis = FortiGateSourceReporter().analyze_source(source)
    references = analysis.derived.references
    assert references.get(ReferenceKind.INTERFACE, vdom="tenant-a", name="wan1") is not None
    assert references.get(ReferenceKind.INTERFACE, vdom="tenant-b", name="wan1") is None
    assert [item.source_vdom for item in analysis.derived.broken_references if item.reference == "wan1"] == ["tenant-b"]
    assert analysis.derived.nat[0].translated_addresses == ("192.0.2.1",)
    assert analysis.derived.nat[1].translated_addresses == ()


def test_repeated_objects_keep_final_state_nested_tagging_and_all_evidence():
    source = '''config firewall address
    edit a
        set subnet 192.0.2.0 255.255.255.0
        set comment old
        set future-option preserved
        config tagging
            edit tag
                set category group
                set tags one
            next
        end
    next
    edit b
    next
    edit a
        unset comment
        set subnet 198.51.100.0 255.255.255.0
        config tagging
            edit tag
                append tags two
            next
            edit second
                set category other
            next
        end
    next
end
config firewall addrgrp
    edit group
        set member a
        config tagging
            edit tag
                set category group
                set tags one
            next
        end
    next
    edit group
        append member b
        config tagging
            edit tag
                unset category
                append tags two
            next
        end
    next
end
config firewall service custom
    edit svc
        set tcp-portrange 80
    next
    edit svc
        set tcp-portrange 443
    next
end
config firewall service group
    edit services
        set member svc
    next
    edit services
        unset member
    next
end
config firewall policy
    edit 1
        set srcaddr a
    next
    edit 2
    next
    edit 1
        append srcaddr b
    next
end
config vdom
    edit tenant
        config firewall address
            edit a
                set subnet 203.0.113.0 255.255.255.0
            next
        end
    next
end
'''
    analysis = FortiGateSourceReporter().analyze_source(source)
    config = analysis.extracted.config
    assert [(item.vdom, item.name) for item in config.addresses] == [("root", "a"), ("root", "b"), ("tenant", "a")]
    address = config.addresses[0]
    assert address.subnet == "198.51.100.0 255.255.255.0" and address.comment is None
    assert address.raw_extra == {"future-option": "preserved"}
    assert address.tagging[0].tags == ["one", "two"]
    assert [item.name for item in address.tagging] == ["tag", "second"]
    assert config.address_groups[0].members == ["a", "b"]
    assert config.address_groups[0].tagging[0].tags == ["one", "two"]
    assert config.address_groups[0].tagging[0].category is None
    assert len(config.services) == 1 and config.services[0].tcp_portrange == "443"
    assert len(config.service_groups) == 1 and "members" not in config.service_groups[0].explicit_fields
    assert [item.policy_id for item in config.policies] == [1, 2]
    assert config.policies[0].srcaddr == ["a", "b"]
    assert not analysis.derived.references.duplicates
    assert len([record for record in analysis.extracted.source_objects if record.source_path == "firewall address" and record.object_name == "a"]) == 3


@pytest.mark.parametrize("tail,expected,malformed", [
    ("append dhgrp bad 19", None, True),
    ("set dhgrp 14 bad 19", None, True),
    ("set dhgrp 14 bad\n        append dhgrp 19", None, True),
    ("set dhgrp 14 bad\n        set dhgrp 19 20", [19, 20], False),
    ("set dhgrp 14 bad\n        unset dhgrp", [], False),
])
def test_integer_lists_fail_closed_and_recover_in_source_order(tail, expected, malformed):
    analysis = FortiGateSourceReporter().analyze_source(f'''config vpn ipsec phase2-interface
    edit tunnel
        set dhgrp 14
        {tail}
    next
end
''')
    item = analysis.extracted.config.ipsec_phase2[0]
    assert item.dhgrp == expected
    assert ("unparsed_dhgrp" in item.raw_extra) == malformed
    assert bool([issue for issue in analysis.validation.issues if issue.field == "dhgrp"]) == malformed
    if malformed:
        assert "bad" in item.raw_extra["unparsed_dhgrp"]
        assert "dhgrp" not in item.explicit_fields
        assert analysis.derived.vpn.phase2[0].dhgrp is None
        assert all("bad" not in issue.message for issue in analysis.validation.issues)
        issues = ValidationIssueIndex(analysis.validation).issues_for("root", ("tunnel",), ("source_value",))
        assert any(issue.field == "dhgrp" for issue in issues)
    FortiGateSourceReporter().build_preview(analysis)


def test_repeated_vpn_edit_keeps_malformed_append_and_secret_state():
    analysis = FortiGateSourceReporter().analyze_source('''config vpn ipsec phase1-interface
    edit tunnel
        set dhgrp 14
        set psksecret never-export-this
    next
    edit tunnel
        append dhgrp bad 19
        unset psksecret
    next
end
''')
    item = analysis.extracted.config.ipsec_phase1[0]
    assert len(analysis.extracted.config.ipsec_phase1) == 1
    assert item.dhgrp is None
    assert item.raw_extra["unparsed_dhgrp"] == [14, "bad", "19"]
    assert not item.psk_configured
    assert "never-export-this" not in repr(analysis.extracted)


def test_policy_auto_allocation_is_not_a_repeated_object_identity():
    analysis = FortiGateSourceReporter().analyze_source('''config firewall policy
    edit 0
        set name first
    next
    edit 0
        set name second
    next
end
''')
    assert [item.name for item in analysis.extracted.config.policies] == ["first", "second"]


def test_repeated_ipv6_objects_do_not_merge_with_ipv4():
    analysis = FortiGateSourceReporter().analyze_source('''config firewall address
    edit a
        set subnet 192.0.2.0 255.255.255.0
    next
end
config firewall address6
    edit a
        set ip6 2001:db8:1::/64
    next
    edit a
        set ip6 2001:db8:2::/64
    next
end
config firewall addrgrp6
    edit group
        set member a
    next
    edit group
        append member b
    next
end
''')
    assert [(item.address_family, item.name) for item in analysis.extracted.config.addresses] == [("ipv4", "a"), ("ipv6", "a")]
    assert analysis.extracted.config.addresses[1].ip6 == "2001:db8:2::/64"
    assert analysis.extracted.config.address_groups[0].members == ["a", "b"]


def test_empty_sections_are_evidence_without_model_gaps_in_full_and_fast():
    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source('''config global
    config system interface
    end
    config unknown empty
    end
end
config vdom
    edit tenant
        config firewall address
        end
    next
end
''')
    assert not analysis.extracted.config.interfaces and not analysis.extracted.config.addresses
    assert [(record.cli_scope, record.vdom, record.source_path) for record in analysis.extracted.source_objects] == [
        ("global", None, "system interface"), ("global", None, "unknown empty"),
        ("vdom", "tenant", "firewall address"),
    ]
    for profile in (ExcelExportProfile.FULL, ExcelExportProfile.FAST):
        output = io.BytesIO()
        reporter.export_excel(analysis, output, profile=profile)
        workbook = load_workbook(io.BytesIO(output.getvalue()))
        if profile is ExcelExportProfile.FAST:
            assert "FortiGate Source Inventory" not in workbook.sheetnames
            continue
        sheet = workbook["FortiGate Source Inventory"]
        headers = [cell.value for cell in sheet[3]]
        rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=4, values_only=True)]
        assert len(rows) == 3 and all(row["Extraction Status"] == "EMPTY_SECTION" for row in rows)
        assert rows[0]["CLI Scope"] == "global" and rows[0]["Line Start"] == 2
        coverage = workbook["Extraction Coverage"]
        headers = [cell.value for cell in coverage[3]]
        assert all(row[headers.index("Model Gaps")] == 0 for row in coverage.iter_rows(min_row=4, values_only=True))


def test_global_unknown_ownership_and_malformed_values_reach_web_and_excel():
    from fwmigrate.web import create_app

    client = create_app({"TESTING": True}).test_client()
    source = b'''config global
    config system interface
        edit unknown
            set ip 192.0.2.1 255.255.255.0
        next
        edit assigned
            set vdom tenant
        next
    end
end
config vpn ipsec phase2-interface
    edit tunnel
        set dhgrp 14 invalid 19
    next
end
'''
    for route in ("/api/preview", "/api/extract/excel"):
        response = client.post(route, data={
            "source_vendor": "fortigate", "excel_profile": "full",
            "file": (io.BytesIO(source), "scope.conf"),
        }, content_type="multipart/form-data")
        assert response.status_code == 200
        if route.endswith("excel"):
            workbook = load_workbook(io.BytesIO(response.data))
            assert "FortiGate Source Inventory" in workbook.sheetnames
            assert "Review Required" in workbook.sheetnames
