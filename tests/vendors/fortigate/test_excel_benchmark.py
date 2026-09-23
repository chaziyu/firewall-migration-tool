from __future__ import annotations

import hashlib
import io
import unittest

from openpyxl import load_workbook

from fwmigrate.vendors.fortigate.benchmark import BenchmarkResult, BenchmarkTimings
from fwmigrate.vendors.fortigate.config import ExtractionConfig
from fwmigrate.vendors.fortigate.derived import build_derived_views
from fwmigrate.vendors.fortigate.export import export_excel
from fwmigrate.vendors.fortigate.export.excel_schema import SHEET_ORDER
from fwmigrate.vendors.fortigate.extraction.extractor import extract_fortigate_config
from fwmigrate.vendors.fortigate.parser import parse_fortigate_config
from fwmigrate.vendors.fortigate.validation.validator import validate_config


_SOURCE = '''
config system admin
    edit "benchmark-admin"
        set password "BENCHMARK_SECRET_MARKER"
    next
end
config firewall policy
    edit 7
        set name "BENCHMARK_RAW_CLI_MARKER"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
    next
end
config vpn ipsec phase1-interface
    edit "benchmark-phase1"
        set psksecret "BENCHMARK_PSK_MARKER"
    next
end
config system api-user
    edit "benchmark-api-user"
        set api-key "BENCHMARK_API_TOKEN_MARKER"
    next
end
'''


def _result(source: bytes) -> BenchmarkResult:
    return BenchmarkResult(
        timestamp_utc="2026-09-23 01:02:03",
        scenario="test",
        git_branch="codex/test",
        git_commit="abc123",
        python_version="3.14.0",
        input_size_bytes=len(source),
        input_sha256=hashlib.sha256(source).hexdigest(),
        runs=3,
        warmup_runs=1,
        timings=BenchmarkTimings(1, 2, 3, 4, 5, 6, 21),
        peak_memory_bytes=2 * 1024 * 1024,
        top_level_sections=2,
        source_object_count=5,
        validation_issue_count=1,
        object_counts={"Policies": 1, "Interfaces": 2},
    )


def _workbook(benchmark: BenchmarkResult | None):
    extracted = extract_fortigate_config(
        parse_fortigate_config(_SOURCE), config=ExtractionConfig()
    )
    derived = build_derived_views(extracted.config)
    validation = validate_config(extracted.config, derived=derived)
    output = io.BytesIO()
    export_excel(
        extracted=extracted,
        derived=derived,
        validation=validation,
        output=output,
        source_name="sanitized.conf",
        benchmark=benchmark,
    )
    output.seek(0)
    return load_workbook(output, data_only=True)


class ExcelBenchmarkTest(unittest.TestCase):
    def test_benchmark_sheet_is_opt_in_and_appended_once(self):
        normal = _workbook(None)
        self.assertEqual(list(SHEET_ORDER), normal.sheetnames)
        self.assertNotIn("Performance Benchmark", normal.sheetnames)

        measured = _workbook(_result(_SOURCE.encode()))
        self.assertEqual(list(SHEET_ORDER), measured.sheetnames[:-1])
        self.assertEqual("Performance Benchmark", measured.sheetnames[-1])
        self.assertEqual(1, measured.sheetnames.count("Performance Benchmark"))

    def test_benchmark_values_and_exact_input_identity_are_written(self):
        raw = _SOURCE.encode()
        workbook = _workbook(_result(raw))
        values = {
            cell.value
            for row in workbook["Performance Benchmark"].iter_rows()
            for cell in row
            if cell.value is not None
        }
        self.assertIn("Git Commit", values)
        self.assertIn("abc123", values)
        self.assertIn("Input SHA-256", values)
        self.assertIn(hashlib.sha256(raw).hexdigest(), values)
        self.assertIn(len(raw), values)
        for expected in ("Parse", "Extraction", "Excel Export", "Peak Traced Memory", "Policies", "Source Objects"):
            self.assertIn(expected, values)

    def test_benchmark_sheet_contains_no_source_configuration(self):
        workbook = _workbook(_result(_SOURCE.encode()))
        values = " ".join(
            str(cell.value)
            for row in workbook["Performance Benchmark"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertNotIn("BENCHMARK_SECRET_MARKER", values)
        self.assertNotIn("BENCHMARK_PSK_MARKER", values)
        self.assertNotIn("BENCHMARK_API_TOKEN_MARKER", values)
        self.assertNotIn("BENCHMARK_RAW_CLI_MARKER", values)
        self.assertNotIn("set password", values)


if __name__ == "__main__":
    unittest.main()
