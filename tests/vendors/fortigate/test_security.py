import io
import json
import unittest
from openpyxl import load_workbook
from fwmigrate.source_reporting import source_reporters
from fwmigrate.web import create_app
from fwmigrate.vendors.fortigate.security.extraction import sanitize_source_attributes, sanitize_source_value
from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter, FortiGateSourceResult

SECRET_SOURCE = """config vpn ipsec phase1-interface
    edit VPN-HQ
        set psksecret do-not-export-this-secret
    next
end
config user local
    edit alice
        set type password
        set passwd another-do-not-export-secret
    next
end
"""

def test_fortigate_analysis_does_not_retain_parser_tree_or_secrets():
    analysis = source_reporters.get("fortigate").analyze_source(SECRET_SOURCE)

    assert isinstance(analysis, FortiGateSourceResult)
    assert not hasattr(analysis, "tree")
    assert analysis.top_level_sections == 2
    assert "do-not-export-this-secret" not in repr(analysis)
    assert "another-do-not-export-secret" not in repr(analysis)

def test_unknown_command_secret_evidence_is_redacted():
    source = """config firewall unsupported-section
    edit source-only
        mystery password do-not-export-unknown
        mystery authpasswd do-not-export-auth
    next
end
"""
    record = source_reporters.get("fortigate").analyze_source(source).extracted.source_objects[0]
    assert "do-not-export-unknown" not in str(record)
    assert "do-not-export-auth" not in str(record)
    assert record.commands[0].values[-1] == "[REDACTED]"

def test_fortigate_web_preview_and_excel_redact_secrets():
    client = create_app({"TESTING": True}).test_client()
    preview = client.post(
        "/api/preview",
        data={"file": (io.BytesIO(SECRET_SOURCE.encode()), "secrets.conf")},
        content_type="multipart/form-data",
    )
    assert preview.status_code == 200
    assert "do-not-export-this-secret" not in json.dumps(preview.get_json())
    assert "another-do-not-export-secret" not in json.dumps(preview.get_json())

    workbook = client.post(
        "/api/extract/excel",
        data={"preview_id": preview.get_json()["preview_id"]},
    )
    assert workbook.status_code == 200
    assert b"do-not-export-this-secret" not in workbook.data
    assert b"another-do-not-export-secret" not in workbook.data

def test_phase1_unset_secret_has_no_presence_flag():
    source = """config vpn ipsec phase1-interface
    edit tunnel
        set ppk-secret do-not-export-ppk
        unset ppk-secret
    next
end
"""
    analysis = source_reporters.get("fortigate").analyze_source(source)
    assert not analysis.extracted.config.ipsec_phase1[0].ppk_secret_configured
    assert "do-not-export-ppk" not in str(analysis.extracted.source_objects)


def test_vpn_ipv6_selectors_and_credential_evidence_reach_reports_safely():
    source = """config firewall address6
    edit ipv6-peer
        set ip6 2001:db8:2::1/128
    next
end
config vpn ipsec phase1-interface
    edit tunnel
        set ppk-secret do-not-export-ppk
        set authpasswd do-not-export-auth
        set group-authentication-secret do-not-export-group
    next
end
config vpn ipsec phase2-interface
    edit selectors
        set phase1name tunnel
        set src-addr-type subnet6
        set src-subnet6 2001:db8:1::/126
        set dst-addr-type name6
        set dst-name6 ipv6-peer
    next
    edit mixed
        set src-addr-type range6
        set src-start-ip6 2001:db8:3::1
        set src-end-ip6 2001:db8:3::5
        set dst-addr-type subnet
        set dst-subnet 192.0.2.0 255.255.255.0
    next
end
"""
    reporter = source_reporters.get("fortigate")
    analysis = reporter.analyze_source(source)
    phase1 = analysis.extracted.config.ipsec_phase1[0]
    phase2 = analysis.extracted.config.ipsec_phase2[0]
    assert phase1.ppk_secret_configured
    assert phase1.auth_password_configured
    assert phase1.group_authentication_secret_configured
    assert phase2.src_subnet6 == "2001:db8:1::/126"
    assert phase2.dst_name6 == "ipv6-peer"
    assert analysis.extracted.config.addresses[0].ip6 == "2001:db8:2::1/128"
    assert "do-not-export" not in str(phase1.model_dump())
    assert "do-not-export" not in str(analysis.extracted.config.model_dump())
    assert "do-not-export" not in str(analysis.extracted.source_objects)

    preview = reporter.build_preview(analysis)
    assert preview["sections"]["addresses"][0]["value"] == "2001:db8:2::1/128"
    row = preview["sections"]["vpn_phase2"][0]
    assert row["source_range6"] == "2001:db8:1::-2001:db8:1::3"
    assert row["destination_range6"] == "2001:db8:2::1-2001:db8:2::1"
    assert row["source_range"] is None
    mixed = preview["sections"]["vpn_phase2"][1]
    assert mixed["source_range6"] == "2001:db8:3::1-2001:db8:3::5"
    assert mixed["destination_range"] == "192.0.2.0-192.0.2.255"
    assert "do-not-export" not in json.dumps(preview)

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    address_headers = [cell.value for cell in workbook["Addresses"][3]]
    address_values = [cell.value for cell in workbook["Addresses"][4]]
    assert address_values[address_headers.index("Value")] == "2001:db8:2::1/128"
    headers = [cell.value for cell in workbook["VPN Phase 2"][3]]
    values = [cell.value for cell in workbook["VPN Phase 2"][4]]
    assert values[headers.index("Source Range IPv6")] == row["source_range6"]
    assert values[headers.index("Destination Range IPv6")] == row["destination_range6"]
    mixed_values = [cell.value for cell in workbook["VPN Phase 2"][5]]
    assert mixed_values[headers.index("Source Range IPv6")] == mixed["source_range6"]
    assert mixed_values[headers.index("Destination Range")] == mixed["destination_range"]
    headers = [cell.value for cell in workbook["VPN Tunnels"][3]]
    values = [cell.value for cell in workbook["VPN Tunnels"][4]]
    assert values[headers.index("PPK Secret Configured")] == "Yes"


class SecurityTest(unittest.TestCase):
    def test_source_value_sanitization_matches_attribute_sanitization(self):
        values = {
            "passwd": "secret",
            "psksecret": "secret",
            "api-key": "secret",
            "some-shared-secret": "secret",
            "passwd-time": "metadata",
            "has-psk": "Yes",
            "ordinary-setting": "value",
        }
        sanitized = sanitize_source_attributes(values)
        self.assertEqual("[REDACTED]", sanitize_source_value("passwd", "secret"))
        self.assertEqual("[REDACTED]", sanitize_source_value("psksecret", "secret"))
        self.assertEqual("[REDACTED]", sanitize_source_value("api-key", "secret"))
        self.assertEqual("[REDACTED]", sanitize_source_value("some-shared-secret", "secret"))
        self.assertEqual("metadata", sanitize_source_value("passwd-time", "metadata"))
        self.assertEqual("Yes", sanitize_source_value("has-psk", "Yes"))
        self.assertEqual("value", sanitize_source_value("ordinary-setting", "value"))
        self.assertEqual(
            sanitized,
            {key.replace("-", "_"): sanitize_source_value(key, value) for key, value in values.items()},
        )

    def test_secret_families_are_redacted_through_all_reporting_layers(self):
        source = '''config vpn ipsec phase1-interface
    edit tunnel
        set psksecret secret-psk
        set ppk-secret secret-ppk
        set authpasswd secret-auth
        set group-authentication-secret secret-group
    next
end
config user local
    edit alice
        set type password
        set passwd secret-local
    next
end
config vpn ssl client
    edit remote
        set psk secret-ssl
    next
end
config vpn ssl web user-bookmark
    edit alice
        config bookmarks
            edit app
                set logon-password secret-logon
                set sso-password secret-sso
            next
        end
    next
end
config vpn kmip-server
    edit kmip
        set password secret-kmip
    next
end
config user krb-keytab
    edit svc
        set keytab secret-keytab
    next
end
config system sdn-proxy
    edit cloud
        set password secret-proxy
    next
end
config firewall unsupported-section
    edit hidden
        mystery authpasswd secret-unknown
    next
end
'''
        reporter = FortiGateSourceReporter()
        analysis = reporter.analyze_source(source)
        secrets = (
            "secret-psk", "secret-ppk", "secret-auth", "secret-group", "secret-local",
            "secret-ssl", "secret-logon", "secret-sso", "secret-kmip", "secret-keytab",
            "secret-proxy", "secret-unknown",
        )

        for value in (analysis, analysis.extracted.config, analysis.extracted.source_objects,
                      analysis.derived, analysis.validation, reporter.build_preview(analysis)):
            rendered = repr(value)
            for secret in secrets:
                self.assertNotIn(secret, rendered)

        output = io.BytesIO()
        reporter.export_excel(analysis, output)
        workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
        cells = repr([[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets])
        for secret in secrets:
            self.assertNotIn(secret, cells)
