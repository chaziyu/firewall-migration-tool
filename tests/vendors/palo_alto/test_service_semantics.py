import io

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def test_service_override_preserves_yes_no_presence_through_excel():
    source = """<config><shared><service>
      <entry name='missing'><protocol><tcp><port>443</port></tcp></protocol></entry>
      <entry name='disabled'><protocol><tcp><override><no/></override></tcp></protocol></entry>
      <entry name='enabled'><protocol><tcp><override><yes/></override></tcp></protocol></entry>
      <entry name='timeout'><protocol><tcp><override><yes><timeout>300</timeout></yes></override></tcp></protocol></entry>
    </service></shared></config>"""
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(source)
    values = {item.name: item.tcp.override for item in analysis.config.services}
    assert values["missing"] is None
    assert values["disabled"].enabled == "no" and values["disabled"].explicit_fields == {"enabled"}
    assert values["enabled"].enabled == "yes"
    assert (values["timeout"].enabled, values["timeout"].timeout) == ("yes", "300")

    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)
    headers = next(workbook["Services"].iter_rows(min_row=3, max_row=3, values_only=True))
    rows = [dict(zip(headers, row)) for row in workbook["Services"].iter_rows(min_row=4, values_only=True)]
    enabled = {row["Name"]: row["Override Enabled"] for row in rows}
    assert enabled["disabled"] == "no"
    assert enabled["enabled"] == "yes"
