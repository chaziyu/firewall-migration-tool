import io

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def test_secret_leaves_are_redacted_from_result_preview_and_workbook():
    secret = "pan-secret-regression"
    result = PaloAltoSourceReporter().analyze_source(f"<config><shared><address><entry name='x'><future-password>{secret}</future-password></entry></address></shared></config>")
    output = io.BytesIO()
    reporter = PaloAltoSourceReporter()
    reporter.export_excel(result, output)
    values = " ".join(str(cell.value) for sheet in load_workbook(io.BytesIO(output.getvalue()), read_only=True).worksheets for row in sheet.iter_rows() for cell in row)
    assert secret not in str(result.config.model_dump()) + str(reporter.build_preview(result)) + values
