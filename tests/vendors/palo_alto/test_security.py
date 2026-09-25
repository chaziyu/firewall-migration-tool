import io

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def test_known_secret_classes_never_reach_source_or_report_surfaces():
    secrets = ("LOCAL-PASSWORD", "LOCAL-PHASH", "ADMIN-PASSWORD", "IPSEC-KEY", "PORTAL-PASSWORD", "PORTAL-PASSCODE", "GATEWAY-PASSWORD")
    source = """<config><shared>
      <local-user><entry name='local'><password>LOCAL-PASSWORD</password></entry></local-user>
      <local-user-database><user><entry name='hashed'><phash>LOCAL-PHASH</phash></entry></user></local-user-database>
      <mgt-config><users><entry name='admin'><password>ADMIN-PASSWORD</password></entry></users></mgt-config>
      <network><ipsec><entry name='vpn'><manual-key><key>IPSEC-KEY</key></manual-key></entry></ipsec>
        <global-protect><portal><entry name='portal'><client-config><entry name='client'>
          <authentication-override><cookie-encrypt-decrypt><password>PORTAL-PASSWORD</password></cookie-encrypt-decrypt></authentication-override>
          <agent-configuration><passcode>PORTAL-PASSCODE</passcode></agent-configuration>
        </entry></client-config></entry></portal>
        <gateway><entry name='gateway'><client-authentication><entry name='auth'><password>GATEWAY-PASSWORD</password></entry></client-authentication></entry></gateway></global-protect>
      </network>
    </shared></config>"""
    reporter = PaloAltoSourceReporter()
    result = reporter.analyze_source(source)
    preview = reporter.build_preview(result)
    output = io.BytesIO()
    reporter.export_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)
    workbook_values = " ".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    surfaces = (
        result.config.model_dump_json(), str(result.config.source_inventory), str(result.config.extraction_issues),
        str(result.derived), str(result.validation), str(preview), workbook_values,
    )
    assert all(secret not in " ".join(surfaces) for secret in secrets)
    assert result.config.local_users[0].password_configured is True
    assert result.config.administrators[0].password_configured is True
    assert result.config.ipsec_tunnels[0].manual_key_configured is True


def test_unknown_password_leaf_is_redacted_end_to_end():
    secret = "pan-secret-regression"
    source = f"""<config><shared><address><entry name='x'><future-password>{secret}</future-password></entry></address>
      <network><sdwan><traffic-distribution-profile><entry name='bad'><future-password>{secret}</future-password><link><entry><weight><member>not-a-scalar</member></weight></entry></link></entry></traffic-distribution-profile></sdwan></network>
    </shared></config>"""
    reporter = PaloAltoSourceReporter()
    result = reporter.analyze_source(source)
    output = io.BytesIO()
    reporter.export_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    values = " ".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    assert secret not in str(result.config.model_dump()) + str(reporter.build_preview(result)) + values
    assert "[REDACTED]" in values
