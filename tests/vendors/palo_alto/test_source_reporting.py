import io
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


ROOT = Path(__file__).parents[2]
FIXTURES = ROOT / "fixtures" / "palo_alto"
REPORTER = PaloAltoSourceReporter()


def _analysis(name="integrated_panorama.xml"):
    return REPORTER.analyze_source((FIXTURES / name).read_text(encoding="utf-8"))


def test_reporter_returns_native_opaque_result_and_all_reporting_layers():
    result = _analysis()
    assert result.config.source_format == "xml"
    assert not hasattr(result.config, "canonical_ir")
    assert REPORTER.build_preview(result)["vendor"] == "palo_alto"
    output = io.BytesIO()
    REPORTER.export_excel(result, output)
    assert "Summary" in load_workbook(io.BytesIO(output.getvalue()), read_only=True).sheetnames
