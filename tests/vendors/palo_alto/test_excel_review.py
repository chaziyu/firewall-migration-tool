import io

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def _workbook(source):
    output = io.BytesIO()
    reporter = PaloAltoSourceReporter()
    reporter.export_excel(reporter.analyze_source(source), output)
    return load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)


def test_validation_issue_is_exported_to_validation_sheet():
    workbook = _workbook("<config><shared><address><entry name='bad'><ip-netmask>999.999.999.999</ip-netmask></entry></address></shared></config>")
    rows = list(workbook["Validation"].iter_rows(min_row=4, values_only=True))
    assert any("malformed ip_netmask" in str(row) for row in rows)


def test_same_named_scoped_rows_keep_review_reasons_in_their_scope():
    source = """<config><devices><entry name='panorama'><device-group>
      <entry name='dg-a'><pre-rulebase><security><rules><entry name='Allow-Web'><from><member>any</member></from><to><member>any</member></to><source><member>missing-a</member></source><destination><member>any</member></destination><action>allow</action></entry></rules></security></pre-rulebase></entry>
      <entry name='dg-b'><pre-rulebase><security><rules><entry name='Allow-Web'><from><member>any</member></from><to><member>any</member></to><source><member>any</member></source><destination><member>any</member></destination><action>allow</action></entry></rules></security></pre-rulebase></entry>
    </device-group></entry></devices></config>"""
    workbook = _workbook(source)
    sheet = workbook["Security Policies"]
    headers = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    rows = [row for values in sheet.iter_rows(min_row=4, values_only=True) if (row := dict(zip(headers, values)))["Name"] == "Allow-Web"]
    assert len(rows) == 2
    by_scope = {row["Scope Name"]: row["Review Reasons"] for row in rows}
    assert "missing-a" in by_scope["dg-a"]
    assert not by_scope["dg-b"]
