import io
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter


def _sheet_rows(source, name):
    output = io.BytesIO()
    PaloAltoSourceReporter().export_excel(PaloAltoSourceReporter().analyze_source(source), output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)
    headers = next(workbook[name].iter_rows(min_row=3, max_row=3, values_only=True))
    return [dict(zip(headers, row)) for row in workbook[name].iter_rows(min_row=4, values_only=True)]


def test_excel_report_contains_native_domain_sheets():
    source = (Path(__file__).parents[2] / "fixtures" / "example_palo_alto.xml").read_text()
    output = io.BytesIO()
    reporter = PaloAltoSourceReporter()
    reporter.export_excel(reporter.analyze_source(source), output)
    sheetnames = set(load_workbook(io.BytesIO(output.getvalue()), read_only=True).sheetnames)
    assert {"Summary", "Review Required", "Interfaces", "Security Policies"}.issubset(sheetnames)
    assert "Validation" not in sheetnames


def test_tags_sheet_exports_rows():
    rows = _sheet_rows("<config><shared><tag><entry name='production'><color>red</color><comments>keep</comments></entry></tag></shared></config>", "Tags")
    assert rows[0]["Name"] == "production"
    assert rows[0]["Color"] == "red"
    assert rows[0]["Comments"] == "keep"


def test_dhcp_rows_are_exported():
    rows = _sheet_rows("<config><shared><network><dhcp><interface><entry name='ethernet1/1'><server><mode>enabled</mode></server></entry></interface></dhcp></network></shared></config>", "DHCP Servers")
    assert (rows[0]["Interface"], rows[0]["Mode"]) == ("ethernet1/1", "enabled")


def test_sdwan_rows_are_exported():
    rows = _sheet_rows("<config><shared><network><sdwan><rules><entry name='prefer-primary'><description>keep</description></entry></rules></sdwan></network></shared></config>", "SD-WAN Rules")
    assert rows[0]["Name"] == "prefer-primary"


def test_identity_rows_are_exported():
    rows = _sheet_rows("<config><shared><local-user><entry name='alice'><disabled>no</disabled></entry></local-user></shared></config>", "Local Users")
    assert rows[0]["Name"] == "alice"


def test_globalprotect_rows_are_exported():
    rows = _sheet_rows("<config><shared><network><global-protect><portal><entry name='portal-main'/></portal></global-protect></network></shared></config>", "GlobalProtect Portals")
    assert rows[0]["Name"] == "portal-main"


def test_security_policy_columns_are_exported():
    rows = _sheet_rows("""<config><shared><rulebase><security><rules><entry name='allow'>
      <source-hip><member>hip-a</member></source-hip><destination-hip><member>hip-b</member></destination-hip>
      <icmp-unreachable>yes</icmp-unreachable><group-tag>group-a</group-tag>
    </entry></rules></security></rulebase></shared></config>""", "Security Policies")
    row = rows[0]
    assert (row["Source HIP"], row["Destination HIP"], row["ICMP Unreachable"], row["Group Tag"]) == ("hip-a", "hip-b", "yes", "group-a")


def test_nat_translation_columns_are_exported():
    rows = _sheet_rows("""<config><shared><rulebase><nat><rules><entry name='dnat'>
      <destination-translation><translated-address>192.0.2.5</translated-address><translated-port>443</translated-port><dns-rewrite><direction>reverse</direction></dns-rewrite></destination-translation>
    </entry></rules></nat></rulebase></shared></config>""", "NAT Rules")
    row = rows[0]
    assert row["Destination Translated Address"] == "192.0.2.5"
    assert row["Destination Translated Port"] == "443"
    assert row["DNS Rewrite Direction"] == "reverse"


def test_zone_columns_are_exported():
    rows = _sheet_rows("""<config><shared><zone><entry name='inside'><network><zone-protection-profile>protect</zone-protection-profile>
      <enable-packet-buffer-protection>yes</enable-packet-buffer-protection><net-inspection>yes</net-inspection>
    </network><user-acl><include-list><member>u1</member></include-list></user-acl></entry></zone></shared></config>""", "Zones")
    row = rows[0]
    assert (row["Zone Protection Profile"], row["Packet Buffer Protection"], row["Network Inspection"], row["User ACL Include"]) == ("protect", "yes", "yes", "u1")


def test_vulnerability_columns_are_exported():
    rows = _sheet_rows("""<config><shared><profiles><vulnerability><entry name='vuln'><rules><entry name='threat'>
      <action><block-ip><track-by>source</track-by><duration>60</duration></block-ip></action>
    </entry></rules></entry></vulnerability></profiles></shared></config>""", "Vulnerability Rules")
    assert (rows[0]["Rule Name"], rows[0]["Block IP Track By"], rows[0]["Block IP Duration"]) == ("threat", "source", "60")


def test_route_path_monitor_rows_preserve_router_ownership():
    source = """<config><shared><network>
      <virtual-router><entry name='vr-main'><routing-table><ip><static-route><entry name='vr-route'><path-monitor><monitor-destinations><entry name='probe-vr'><destination>192.0.2.1</destination></entry></monitor-destinations></path-monitor></entry></static-route></ip></routing-table></entry></virtual-router>
      <logical-router><entry name='lr-main'><vrf><entry name='production'><routing-table><ip><static-route><entry name='lr-route'><path-monitor><monitor-destinations><entry name='probe-lr'><destination-fqdn>probe.example</destination-fqdn></entry></monitor-destinations></path-monitor></entry></static-route></ip></routing-table></entry></vrf></entry></logical-router>
    </network></shared></config>"""
    rows = _sheet_rows(source, "Route Path Monitors")
    assert {(row["Router Type"], row["Router"], row["VRF"], row["Route Name"], row["Monitor Name"]) for row in rows} == {
        ("virtual-router", "vr-main", None, "vr-route", "probe-vr"),
        ("logical-router", "lr-main", "production", "lr-route", "probe-lr"),
    }


def test_service_override_values_are_exported():
    rows = _sheet_rows("""<config><shared><service>
      <entry name='disabled'><protocol><tcp><port>443</port><override><no/></override></tcp></protocol></entry>
      <entry name='enabled'><protocol><tcp><port>443</port><override><yes/></override></tcp></protocol></entry>
    </service></shared></config>""", "Services")
    enabled = {row["Name"]: row["Override Enabled"] for row in rows}
    assert enabled == {"disabled": "no", "enabled": "yes"}


def test_sdwan_interface_settings_are_exported():
    source = """<config><shared><network><sdwan><interface-profile><entry name='wan-primary'/></interface-profile></sdwan></network></shared>
      <devices><entry name='fw'><network><interface><ethernet><entry name='ethernet1/1'><layer3><sdwan-link-settings>
        <enable>yes</enable><ipv6-enable>no</ipv6-enable><sdwan-interface-profile>wan-primary</sdwan-interface-profile><upstream-nat><enable>yes</enable></upstream-nat>
      </sdwan-link-settings></layer3></entry></ethernet></interface></network></entry></devices></config>"""
    rows = _sheet_rows(source, "Interfaces")
    row = next(item for item in rows if item["Name"] == "ethernet1/1")
    assert (row["SD-WAN Enabled"], row["IPv6 SD-WAN Enabled"], row["SD-WAN Interface Profile"], row["Upstream NAT"]) == ("yes", "no", "wan-primary", "yes")


def test_formula_like_source_text_is_exported_as_literal():
    source = "<config><shared><tag><entry name='formula'><comments>=1+1</comments></entry></tag></shared></config>"
    output = io.BytesIO()
    reporter = PaloAltoSourceReporter()
    reporter.export_excel(reporter.analyze_source(source), output)
    sheet = load_workbook(io.BytesIO(output.getvalue()), data_only=False)["Tags"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    cell = sheet.cell(4, headers["Comments"])
    assert cell.value == "'=1+1"
    assert cell.data_type != "f"


def test_partial_extraction_status_and_source_appendix_evidence():
    source = """<config><shared><network><dhcp><interface><entry name='ethernet1/1'><mode>auto</mode><server><future-setting>retain-me</future-setting></server></entry></interface></dhcp>
      <service><entry name='web'><protocol><tcp><port>443</port><future-tcp-setting>retain-protocol-source</future-tcp-setting></tcp></protocol></entry></service>
      <sdwan><traffic-distribution-profile><entry name='malformed'><link><entry><weight><member>not-a-scalar</member></weight></entry></link></entry></traffic-distribution-profile></sdwan>
    </network></shared></config>"""
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(source)
    output = io.BytesIO()
    reporter.export_excel(analysis, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=True)
    inventory = list(workbook["PAN-OS Source Inventory"].iter_rows(min_row=4, values_only=True))
    dhcp = next(row for row in inventory if row[7] == "ethernet1/1")
    inventory_headers = next(workbook["PAN-OS Source Inventory"].iter_rows(min_row=3, max_row=3, values_only=True))
    assert dhcp[inventory_headers.index("Extraction Status")] == "PARTIAL"
    dhcp_sheet = workbook["DHCP Servers"]
    headers = next(dhcp_sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    dhcp_rows = [dict(zip(headers, row)) for row in dhcp_sheet.iter_rows(min_row=4, values_only=True)]
    assert dhcp_rows[0]["Analysis Status"] == "NO_VALIDATION_ISSUES"
    coverage = list(workbook["Extraction Coverage"].iter_rows(min_row=4, values_only=True))
    assert next(row for row in coverage if row[0] == "dhcp_interface")[4] == "PARTIAL"
    appendix = list(workbook["PAN-OS Source Appendix"].iter_rows(min_row=4, values_only=True))
    assert any(row[7].endswith("server/future-setting") and row[9] == "retain-me" for row in appendix)
    assert any(row[7].endswith("protocol/tcp/future-tcp-setting") and row[9] == "retain-protocol-source" for row in appendix)
    assert "Services" in workbook.sheetnames and list(workbook["Services"].iter_rows(min_row=4, values_only=True))
    validation = list(workbook["Review Required"].iter_rows(min_row=4, values_only=True))
    assert any(row[0] == "warning" and row[1] == "extraction" and "ValidationError" in str(row) for row in validation)
    unsupported = list(workbook["Unsupported"].iter_rows(min_row=4, values_only=True))
    assert any("not-a-scalar" in str(row) for row in unsupported)


def test_interface_units_only_link_to_parents_in_the_same_scope():
    source = """<config><devices>
      <entry name='fw-a'><vsys><entry name='vsys1'><import><network><interface><member>ethernet1/1</member></interface></network></import></entry></vsys>
        <network><interface><ethernet><entry name='ethernet1/1'><layer3/></entry></ethernet></interface></network>
      </entry>
      <entry name='fw-b'><vsys><entry name='vsys1'><import><network><interface><member>ethernet1/1</member></interface></network></import></entry></vsys>
        <network><interface><ethernet><entry name='ethernet1/1'><layer3><units><entry name='ethernet1/1.10'><tag>10</tag><layer3/></entry></units></layer3></entry></ethernet></interface></network>
      </entry>
    </devices></config>"""
    rows = _sheet_rows(source, "Interfaces")
    assert len(rows) == 3
    unit = next(row for row in rows if row["Name"] == "ethernet1/1.10")
    assert unit["Parent Interface"] == "ethernet1/1"
    assert unit["Device"] == "fw-b"


def test_source_appendix_does_not_repeat_nested_entry_evidence():
    source = """<config><shared><profiles><vulnerability><entry name='parent'>
      <rules><entry name='child'><severity>high</severity></entry></rules>
    </entry></vulnerability></profiles></shared></config>"""
    rows = _sheet_rows(source, "PAN-OS Source Appendix")
    entry_name_values = [row["Evidence Value"] for row in rows if row["Evidence Path"].endswith("/@name")]
    assert entry_name_values.count("parent") == 1
    assert entry_name_values.count("child") == 1


def test_full_streaming_export_preserves_all_rows_layout_and_collection_evidence():
    from copy import deepcopy
    from openpyxl.utils import get_column_letter
    from fwmigrate.vendors.palo_alto.export.excel_rows import ROW_BUILDERS, _PANExcelContext
    from fwmigrate.vendors.palo_alto.export.excel_schema import SHEET_HEADERS, SHEET_ORDER, HIDDEN_COLUMNS_BY_DEFAULT
    from fwmigrate.source_reporting.excel_style import excel_safe_value

    source = "<config><shared><address>" + "".join(
        f'<entry name="host-{index}"><ip-netmask>192.0.2.1</ip-netmask><description>  =host-{index}</description></entry>'
        for index in range(96)) + "</address></shared></config>"
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(source)
    before = deepcopy(analysis.config.model_dump(mode="python"))
    context = _PANExcelContext(analysis, None, collection_status="PARTIAL")
    output = io.BytesIO()
    reporter.export_excel(analysis, output, profile="full", collection_status="PARTIAL",
                          collection_parts=[{"name": "config", "status": "PARTIAL", "complete": False, "count": 96}],
                          collection_warnings=["Incomplete source"])
    workbook = load_workbook(output)
    assert workbook.sheetnames == [*SHEET_ORDER, "Collection Evidence"]
    summary_values = [cell.value for row in workbook["Summary"] for cell in row]
    assert "FULL" in summary_values and "FAST" not in summary_values
    assert "FULL includes PAN-OS Source Inventory, Extraction Coverage, and PAN-OS Source Appendix." in summary_values
    for name in SHEET_ORDER:
        if name == "Summary":
            continue
        headers = SHEET_HEADERS[name]
        expected = [tuple(None if (value := excel_safe_value(row.get(header))) == "" else value for header in headers)
                    for row in ROW_BUILDERS[name](context)]
        sheet = workbook[name]
        assert tuple(cell.value for cell in sheet[3]) == headers
        assert list(sheet.iter_rows(min_row=4, max_col=len(headers), values_only=True)) == expected
        assert sheet["A2"].value.startswith("FULL report. Row counts are listed in Summary.")
        assert f"A1:{get_column_letter(len(headers))}1" in sheet.merged_cells
        assert sheet.cell(2, len(headers)).hyperlink.target == "#'Summary'!A1"
        assert sheet.auto_filter.ref == (f"A3:{get_column_letter(len(headers))}{len(expected) + 3}" if expected else None)
        assert all(sheet.column_dimensions[get_column_letter(column)].hidden
                   for column, header in enumerate(headers, 1) if header in HIDDEN_COLUMNS_BY_DEFAULT.get(name, ()))
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert workbook["PAN-OS Source Appendix"].max_row == 96 * 3 + 3
    assert workbook["Collection Evidence"]["C2"].value is False
    assert workbook["Collection Evidence"]["E3"].value == "Incomplete source"
    workbook.close()
    assert analysis.config.model_dump(mode="python") == before
