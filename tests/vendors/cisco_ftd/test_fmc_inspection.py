import json
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

from fwmigrate.vendors.cisco_ftd.fmc.fmc_adapter import CiscoFMCBundleParser
from fwmigrate.vendors.cisco_ftd.source_report import CiscoFTDSourceReporter, extract_cisco_ftd_source
from fwmigrate.vendors.cisco_ftd.derived import build_ftd_derived_views
from fwmigrate.vendors.cisco_ftd.validation import validate_ftd_config
from fwmigrate.vendors.cisco_ftd.export.excel import export_ftd_excel
from fwmigrate.vendors.cisco_ftd.web_report import build_ftd_preview


FIXTURE = Path(__file__).parents[2] / "fixtures" / "cisco_ftd" / "fmc_selected_domains.json"

def test_inspection_policy_rules_are_typed_ordered_resolved_and_complete():
    payload = {"format": "cisco-fmc-rest-export-v1", "domain": {"id": "d1"},
        "objects": {
            "networkaddresses": [{"id": "net1", "name": "Inside"}],
            "protocolportobjects": [{"id": "svc1", "name": "HTTPS", "protocol": "TCP", "port": "443"}],
            "securityzones": [{"id": "zone1", "name": "Inside Zone"}],
            "vlanobjects": [{"id": "vlan1", "name": "VLAN 20"}],
            "internalcertificates": [{"id": "cert1", "name": "CA"}],
            "siurllists": [{"id": "si1", "name": "Threat URLs", "url": "https://feed.test/list"}],
            "siurlfeeds": [{"id": "feed1", "name": "Live Feed", "url": "https://live.test/feed",
                "entries": ["runtime.example"]}],
            "filepolicies": [{"id": "f1", "name": "Files", "description": "Files", "rules": [
                {"id": "fr2", "name": "second", "metadata": {"ruleIndex": 2}, "action": "BLOCK",
                 "protocol": "HTTP", "direction": "DOWNLOAD",
                 "fileTypes": [{"id": "pdf", "name": "PDF"}], "analysis": ["spero"], "storeFiles": ["MALWARE"]},
                {"id": "fr1", "name": "first", "order": 1, "action": "MONITOR"}]} ,
                {"id": "f2", "name": "Empty", "rules": []}, {"id": "f3", "name": "Unknown"}],
            "decryptionpolicies": [{"id": "dec1", "name": "Decrypt", "defaultAction": "Do Not Decrypt",
                "undecryptableActions": {"decryptionErrors": "BLOCK"}, "advancedOptions": {"strict": True}, "rules": [{
                    "id": "dr1", "name": "Decrypt web", "metadata": {"ruleIndex": 4}, "ruleAction": "DECRYPT_RESIGN",
                    "sourceNetworks": {"objects": [{"id": "net1", "name": "Inside"}]},
                    "sourcePorts": {"objects": [{"id": "svc1", "name": "HTTPS"}]},
                    "decryptionCerts": {"objects": [{"id": "cert1", "name": "CA"}]},
                    "tlsVersions": {"tls13": True}, "certStatuses": {"revoked": False}, "privateKey": "must-not-leak"}]}],
            "dnspolicies": [{"id": "dns1", "name": "DNS", "umbrellaSettings": {"enabled": True},
                "block_rules": [{"id": "dnsr1", "name": "Block feed", "metadata": {"ruleIndex": 3}, "ruleAction": "BLOCK",
                    "sourceZones": [{"id": "zone1", "name": "Inside Zone"}],
                    "sourceNetworks": [{"id": "net1", "name": "Inside"}],
                    "vlanTags": [{"id": "vlan1", "name": "VLAN 20"}],
                    "dnsLists": {"objects": [{"id": "si1", "name": "Threat URLs"}]},
                    "dnsFeeds": {"objects": [{"id": "feed1", "name": "Live Feed"}]}}]}],
        },
        "collection": {"status": "PARTIAL", "parts": [
            {"name": "filepolicies/f3/rules", "status": "FAILED", "complete": False}]},
        "access_policies": [{"id": "acp1", "name": "ACP", "decryptionPolicy": {"id": "dec1", "name": "Decrypt"},
            "dnsPolicy": {"id": "dns1", "name": "DNS"}, "rules": [{"id": "acpr1", "name": "Rule",
                "filePolicy": {"id": "f1", "name": "Files"}}]}]}
    result = extract_cisco_ftd_source(json.dumps(payload))
    config = result.config
    assert [(rule.position, rule.collection_order, rule.action) for rule in config.file_policies[0].rules] == [
        (2, 1, "BLOCK"), (1, 2, "MONITOR")]
    assert "ruleIndex" not in config.file_policies[0].rules[0].raw_extra.get("metadata", {})
    assert config.file_policies[1].rules == [] and config.file_policies[2].rules is None
    assert config.file_policies[0].rules[0].malware_inspection == ["spero"]
    decrypt = config.decryption_policies[0]
    assert decrypt.default_action == "Do Not Decrypt" and decrypt.undecryptable_action == {"decryptionErrors": "BLOCK"}
    assert decrypt.rules[0].action == "DECRYPT_RESIGN"
    dns = config.dns_policies[0]
    assert dns.rules[0].action == "BLOCK" and dns.rules[0].position == 3 and dns.rules[0].collection_order == 1
    assert "must-not-leak" not in json.dumps(config.model_dump())
    assert "runtime.example" not in json.dumps(config.model_dump())
    kinds = {item["kind"] for item in result.derived.resolved_references}
    assert {"NETWORK_ADDRESS", "SERVICE_OBJECT", "CERTIFICATE", "SECURITY_ZONE", "VLAN_OBJECT",
            "SECURITY_INTELLIGENCE_SOURCE", "DECRYPTION_POLICY", "DNS_POLICY", "FILE_POLICY"} <= kinds
    assert result.derived.source_plane_completeness["file_rules"] == "failed"
    assert not {"filepolicyrules", "decryptionpolicyrules", "block_rules"} & {
        item.source_attributes.get("resource_type") for item in config.native_resources}
    paths = {item.source_path for item in result.inventory_items}
    assert {"fmc-rest-bundle/file-policy-rules", "fmc-rest-bundle/decryption-policy-rules",
            "fmc-rest-bundle/dns-policy-rules"} <= paths
    assert result.derived.inspection_relationships
    preview = CiscoFTDSourceReporter().build_preview(result)
    assert preview["summary"]["dns_rules"] == 1
    assert "must-not-leak" not in json.dumps(preview)
    workbook_bytes = BytesIO()
    export_ftd_excel(result, workbook_bytes)
    workbook_bytes.seek(0)
    exported = load_workbook(workbook_bytes, read_only=True)
    assert all("must-not-leak" not in str(cell.value)
               for sheet in exported.worksheets for row in sheet.iter_rows() for cell in row)
    inspection_rows = list(exported["Inspection Policies"].values)
    assert ("File", "Empty", "f2", None, None, "KNOWN EMPTY") == inspection_rows[3][:6]
    assert ("File", "Unknown", "f3", None, None, "UNKNOWN") == inspection_rows[4][:6]
    assert any(row[0] == "Decryption" and row[9] == "DECRYPT_RESIGN" for row in inspection_rows[1:])

def test_intrusion_behavior_group_membership_conflicts_and_acp_references():
    payload = {
        "format": "cisco-fmc-rest-export-v1", "domain": {"id": "domain-1", "name": "Global"},
        "objects": {
            "variablesets": [{"id": "vars-1", "name": "Variables"}],
            "intrusionpolicies": [{"id": "ips-1", "name": "IPS", "variableSet": {"id": "vars-1", "name": "Variables"},
                "rules": [{"id": "behavior-1", "ruleId": "sig-1", "state": "enabled", "action": "alert"}],
                "rule_groups": [{"id": "group-1", "name": "Group", "rules": [
                    {"id": "behavior-1", "ruleId": "sig-1", "state": "disabled", "action": "drop"},
                    {"id": "behavior-2", "ruleId": "sig-2"}]}]}],
        },
        "access_policies": [{"id": "acp-1", "name": "ACP", "rules": [{"id": "acp-rule-1", "name": "Allow",
            "ipsPolicy": {"id": "ips-1", "name": "IPS"}, "variableSet": {"id": "vars-1", "name": "Variables"}}]}],
        "collection": {"status": "SUCCESS", "parts": [
            {"name": "intrusionpolicies", "status": "SUCCESS", "complete": True},
            {"name": "intrusionpolicies/ips-1/rule_groups", "status": "SUCCESS", "complete": True},
            {"name": "intrusionpolicies/ips-1/rules", "status": "SUCCESS", "complete": True}]},
    }
    config = CiscoFMCBundleParser(json.dumps(payload)).parse_source()
    assert len(config.intrusion_policies) == len(config.intrusion_rule_groups) == 1
    assert [(item.source_id, item.rule_id, item.state, item.action) for item in config.intrusion_rule_behaviors] == [
        ("behavior-1", "sig-1", "enabled", "alert")]
    assert not config.intrusion_rule_overrides
    behavior = config.intrusion_rule_behaviors[0]
    assert behavior.source_attributes["group_membership_evidence"][0]["rule_group_id"] == "group-1"
    assert behavior.source_attributes["conflicting_group_payload"][0]["conflicting_fields"]["action"] == {
        "behavior": "alert", "group": "drop"}
    assert not [item for item in config.native_resources
                if item.source_attributes.get("parent_policy_type") == "intrusionpolicies"]

    derived = build_ftd_derived_views(config)
    assert {("intrusion_rule_groups", "present"), ("intrusion_rule_behaviors", "present")} <= set(
        derived.source_plane_completeness.items())
    assert any(item["rule_group_id"] == "group-1" and item["rule_behavior_id"] == "behavior-1"
               for item in derived.intrusion_relationships)
    resolved = {(item["owner"], item["field"], item["kind"]) for item in derived.resolved_references}
    assert ("IPS", "variable_set", "VARIABLE_SET") in resolved
    assert ("Allow", "intrusion_policy", "INTRUSION_POLICY") in resolved
    assert ("Allow", "variable_set", "VARIABLE_SET") in resolved
    assert any(item.category == "intrusion-rule-group-conflict" for item in validate_ftd_config(config, derived).issues)
    build_ftd_derived_views(config)
    validate_ftd_config(config, derived)

    analysis = extract_cisco_ftd_source(json.dumps(payload))
    preview = CiscoFTDSourceReporter().build_preview(analysis)["summary"]
    assert preview["intrusion_policies"] == 1
    assert preview["intrusion_rule_groups"] == 1
    assert preview["intrusion_rule_behaviors"] == 1
    assert preview["intrusion_rule_overrides"] == 0
    assert {item.source_type for item in analysis.inventory_items} >= {
        "intrusion_rule_group", "intrusion_rule_behavior"}
    workbook_bytes = BytesIO()
    export_ftd_excel(analysis, workbook_bytes)
    workbook_bytes.seek(0)
    native_rows = list(load_workbook(workbook_bytes, read_only=True)["Native Sources"].values)
    assert {tuple(row[:2]) for row in native_rows[1:]} >= {
        ("intrusion_rule_groups", "Group"), ("intrusion_rule_behaviors", "behavior-1")}

def test_intrusion_group_only_and_missing_vs_empty_child_collections():
    payload = {"format": "cisco-fmc-rest-export-v1", "domain": {"id": "d", "name": "Global"},
        "objects": {"intrusionpolicies": [{"id": "p", "name": "IPS", "rule_groups": [
            {"id": "g", "name": "Group", "rules": [{"id": "sig", "name": "Signature"}]}]}]},
        "collection": {"status": "PARTIAL", "parts": [
            {"name": "intrusionpolicies", "status": "SUCCESS", "complete": True},
            {"name": "intrusionpolicies/p/rule_groups", "status": "SUCCESS", "complete": True},
            {"name": "intrusionpolicies/p/rules", "status": "FAILED", "complete": False}]}}
    config = CiscoFMCBundleParser(json.dumps(payload)).parse_source()
    assert len(config.intrusion_rule_groups) == 1
    assert not config.intrusion_rule_behaviors and not config.intrusion_rule_overrides
    derived = build_ftd_derived_views(config)
    assert derived.source_plane_completeness["intrusion_rule_groups"] == "present"
    assert derived.source_plane_completeness["intrusion_rule_behaviors"] == "failed"
    assert any(item["rule_group_id"] == "g" and item["rule_id"] == "sig" and
               item["rule_behavior_id"] is None for item in derived.intrusion_relationships)

    payload["objects"]["intrusionpolicies"][0]["rules"] = []
    payload["collection"]["parts"] = [*payload["collection"]["parts"][:2],
        {"name": "intrusionpolicies/p/rules", "status": "EMPTY", "complete": True}]
    derived = build_ftd_derived_views(CiscoFMCBundleParser(json.dumps(payload)).parse_source())
    assert derived.source_plane_completeness["intrusion_rule_behaviors"] == "known-empty"
    payload["collection"]["parts"] = [payload["collection"]["parts"][0]]
    derived = build_ftd_derived_views(CiscoFMCBundleParser(json.dumps(payload)).parse_source())
    assert derived.source_plane_completeness["intrusion_rule_groups"] == "unknown"
    assert derived.source_plane_completeness["intrusion_rule_behaviors"] == "unknown"


def test_prefilter_and_network_analysis_children_are_typed_and_owned():
    result = extract_cisco_ftd_source(json.dumps({"source": "fmc-rest-api", "objects": {
        "prefilterpolicies": [{"id": "prefilter-1", "name": "Prefilter", "rules": [
            {"id": "rule-1", "name": "Fast path", "position": 1, "action": "FASTPATH"}],
            "default_actions": [{"id": "default-1", "name": "Default"}]}],
        "networkanalysispolicies": [{"id": "analysis-1", "name": "Analysis",
            "inspectorconfigs": [{"id": "inspector-1", "name": "Inspector"}],
            "inspectoroverrideconfigs": [{"id": "override-1", "name": "Override"}]}],
    }, "access_policies": [{"id": "acp-1", "name": "ACP", "prefilterPolicy": {
        "id": "prefilter-1", "name": "Prefilter"}, "networkAnalysisPolicy": {
        "id": "analysis-1", "name": "Analysis"}}]}))
    config = result.config
    assert config.prefilter_rules[0].position == 1 and config.prefilter_rules[0].action == "FASTPATH"
    assert config.inspector_configs[0].source_attributes["parent_policy_id"] == "analysis-1"
    assert {item.source_id for item in config.inspector_override_configs} == {"override-1"}
    assert not {"prefilterpolicies", "networkanalysispolicies"} & {
        item.source_attributes.get("resource_type") for item in config.native_resources}
    assert any(item["field"] == "prefilter_policy" and item["target_id"] == "prefilter-1"
               for item in result.derived.resolved_references)
