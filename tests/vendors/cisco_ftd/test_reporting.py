import json
from copy import deepcopy
import pytest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_ftd.source_report import CiscoFTDSourceReporter, extract_cisco_ftd_source
from fwmigrate.vendors.cisco_ftd.web_report import build_ftd_preview


FIXTURE = Path(__file__).parents[2] / "fixtures" / "cisco_ftd" / "fmc_selected_domains.json"


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_cli_excel_preserves_complete_records_and_unsupported_payload(profile):
    source = ("configure network ipv4 manual 192.0.2.42 255.255.255.0 192.0.2.1\n"
              "interface GigabitEthernet0/1\n description UNIQUE_DESCRIPTION\n"
              " ipv6 address 2001:db8::42/64\n security-level 0\n no shutdown\n"
              "access-list UNIQUE_UNSUPPORTED extended permit ip any any\n"
              " password CLI_SECRET\n")
    result = extract_cisco_ftd_source(source)
    before = deepcopy(result.config)
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output, profile=profile)
    workbook = load_workbook(output, read_only=True)
    rows = list(workbook["CLI Source Evidence"].values)
    records = [dict(zip(rows[0], row)) for row in rows[1:]]
    interface = next(row for row in records if row["Record Kind"] == "Interface")
    values = json.loads(interface["Source Values"])
    assert values["description"] == "UNIQUE_DESCRIPTION"
    assert values["ipv6_addresses"][0]["address"] == "2001:db8::42"
    assert values["shutdown"] is False and values["security_level"] == 0
    assert values["management_only"] is None
    assert (interface["Line Start"], interface["Line End"]) == (2, 6)
    all_values = str([row for sheet in workbook for row in sheet.values])
    assert "192.0.2.42" in all_values and "UNIQUE_UNSUPPORTED" in all_values
    assert "CLI_SECRET" not in all_values
    evidence = list(workbook["Source Evidence"].values)
    assert "UNIQUE_UNSUPPORTED" in dict(zip(evidence[0], evidence[1]))["Raw Source"]
    summary = list(workbook["Summary"].values)
    assert any(len(row) > 1 and row[0] == "CLI Source Evidence" and row[1] == len(records) for row in summary)
    workbook.close()
    assert result.config == before


@pytest.mark.parametrize("profile", ["fast", "full"])
@pytest.mark.parametrize(("bundle_format", "collection"), [
    ("cisco-fmc-rest-export-v1", "hosts"), ("cisco-fdm-rest-export-v1", "hosts"),
])
def test_managed_object_raw_fields_survive_excel(profile, bundle_format, collection):
    payload = {"format": bundle_format, "objects": {collection: [{
        "id": "host", "name": "Host", "value": "192.0.2.1", "customField": "UNIQUE_CUSTOM",
        "nested": {"apiToken": "OBJECT_SECRET"},
    }]}}
    if "fdm" in bundle_format:
        payload["unknownResource"] = {"marker": "UNKNOWN_RESOURCE", "password": "UNKNOWN_SECRET"}
    result = extract_cisco_ftd_source(json.dumps(payload))
    before = deepcopy(result.config)
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output, profile=profile)
    workbook = load_workbook(output, read_only=True)
    native = str(list(workbook["Native Sources"].values))
    assert "network_addresses" in native and "UNIQUE_CUSTOM" in native
    all_values = str([row for sheet in workbook for row in sheet.values])
    assert "OBJECT_SECRET" not in all_values and "UNKNOWN_SECRET" not in all_values
    if "fdm" in bundle_format:
        assert "UNKNOWN_RESOURCE" in str(list(workbook["Source Evidence"].values))
    workbook.close()
    assert result.config == before


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_source_evidence_uses_shared_overflow_and_literal_text_guards(profile):
    marker = "=UNIQUE_START" + "x" * 33000 + "UNIQUE_END\x01 password OVERFLOW_SECRET"
    result = extract_cisco_ftd_source(json.dumps({"format": "cisco-fdm-rest-export-v1", "unknown": marker}))
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output, profile=profile)
    workbook = load_workbook(output, read_only=True)
    chunks = [row[4] for row in list(workbook["Text Overflow"].values)[1:] if row[0] == "Source Evidence"]
    reconstructed = "".join(chunks)
    assert "UNIQUE_START" in reconstructed and "UNIQUE_END" in reconstructed
    assert "OVERFLOW_SECRET" not in reconstructed and "\\u0001" in reconstructed
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_native_collections_are_visible_in_preview_and_excel():
    result = CiscoFTDSourceReporter().analyze_source(FIXTURE.read_text(encoding="utf-8"))
    preview = build_ftd_preview(result)
    assert preview["summary"]["s2s_vpn_topologies"] == 1
    assert preview["summary"]["objects"]["policies"] > 0
    assert preview["source_plane"] == "fmc-rest-bundle"
    assert {"interfaces", "addresses", "address_groups", "services", "service_groups", "schedules",
            "policies", "nat", "routes", "vpn_tunnels", "vpn_phase2", "validation",
            "unresolved_references"} <= set(preview["sections"])
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)
    workbook = load_workbook(output, read_only=True)
    assert "Native Sources" in workbook.sheetnames
    assert "Source Inventory" in workbook.sheetnames
    assert "Collection Completeness" in workbook.sheetnames
    assert "Semantic Completeness" in workbook.sheetnames
    assert workbook["Native Sources"].max_row > 1
    assert workbook["Source Inventory"].max_row > 1
    assert result.source_sections[0].object_count_source == len(result.inventory_items)
    assert result.source_sections[0].object_count_parsed == len(result.inventory_items)
    assert result.source_sections[0].object_count_extracted == len(result.inventory_items)


def test_report_preserves_ftd_source_plane_and_partial_coverage():
    fdm = CiscoFTDSourceReporter().build_preview(extract_cisco_ftd_source(
        (FIXTURE.parent / "fdm_object_pipeline_conformance.json").read_text(encoding="utf-8")))
    cli = CiscoFTDSourceReporter().build_preview(extract_cisco_ftd_source(
        "interface outside\n ip address 203.0.113.1 255.255.255.0\n"))
    partial = CiscoFTDSourceReporter().build_preview(extract_cisco_ftd_source(
        '{"format":"cisco-fmc-rest-export-v1","source":"fmc-rest-api","access_policies":[{"id":"p1","name":"Policy A","rules":[{"id":"r1","name":"Rule"}]},{"id":"p2","name":"Policy B"}]}'))
    assert fdm["source_plane"] == "fdm-rest-bundle"
    assert cli["source_plane"] == "ftd-text-evidence" and cli["sections"]["policies"] == []
    assert partial["source_plane_completeness"]["acp"] == "partial"
    assert partial["sections"]["policies"][0]["source_plane"] == "fmc-rest-bundle"


def test_excel_formula_like_source_text_is_literal():
    result = extract_cisco_ftd_source(json.dumps({"source": "fmc-rest-api", "objects": {
        "networkaddresses": [{"id": "net-1", "name": "=1+1", "type": "Host", "value": "192.0.2.1"}]}}))
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)
    workbook = load_workbook(output, read_only=True)
    sheet = workbook["Managed Objects"]
    row = next(sheet.iter_rows(min_row=2, max_row=2))
    assert row[0].value == "'=1+1"
    assert row[0].data_type != "f"

def test_fmc_missing_name_uses_id_only_as_internal_fallback():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "objects": {"networkaddresses": [
            {"id": "host-without-name", "type": "Host", "value": "192.0.2.10"}
        ]},
    }))

    record = result.config.network_addresses[0]
    inventory = next(item for item in result.inventory_items if item.source_id == "host-without-name")
    assert record.name is None
    assert record.source_attributes["source_name_explicit"] is False
    assert record.source_attributes["internal_record_key"] == "id:host-without-name"
    assert inventory.name is None
    assert inventory.source_id == "host-without-name"
    assert inventory.source_record_id is not None


def test_preview_preserves_structured_fmc_scope():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "objects": {},
        "devices": [{"id": "device-1", "name": "FTD-A", "resources": {
            "ftd_interfaces": [{"id": "if-1", "name": "GigabitEthernet0/0"}]
        }}],
    }))

    preview = build_ftd_preview(result)
    row = preview["sections"]["interfaces"][0]
    assert row["scope_details"]["domain_id"] == "domain-1"
    assert row["scope_details"]["device_id"] == "device-1"
    assert row["scope_details"]["device_name"] == "FTD-A"


def test_fmc_missing_domain_name_is_not_inferred_as_global():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1"},
        "objects": {"networkaddresses": [
            {"id": "host-1", "name": "Host", "type": "Host", "value": "192.0.2.10"}
        ]},
    }))
    assert result.config.source_metadata["domain_name"] is None
    assert result.config.network_addresses[0].source_context is None
    assert result.config.network_addresses[0].domain_id == "domain-1"


def test_device_only_fmc_bundle_uses_fmc_adapter():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "devices": [{"id": "device-1", "name": "FTD-A", "resources": {
            "ftd_interfaces": [{"id": "if-1", "name": "GigabitEthernet0/0"}]
        }}],
    }))
    assert result.config.source_plane == "fmc-rest-bundle"
    assert result.config.source_interfaces[0].source_id == "if-1"


def test_preview_keeps_parent_policy_and_rule_identity_distinct():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "access_policies": [{
            "id": "acp-1", "name": "Corporate ACP",
            "rules": [{"id": "rule-1", "name": "Allow DNS", "position": 7, "action": "ALLOW"}],
        }],
        "nat_policies": [{
            "id": "nat-1", "name": "Edge NAT",
            "manual_rules_before_auto": [
                {"id": "nat-rule-1", "name": "Static NAT", "position": 3, "section": "BEFORE_AUTO"}
            ],
        }],
    }))
    preview = build_ftd_preview(result)
    policy = preview["sections"]["policies"][0]
    assert (policy["policy_id"], policy["policy_name"], policy["rule_id"], policy["position"]) == (
        "acp-1", "Corporate ACP", "rule-1", 7,
    )
    nat = preview["sections"]["nat"][0]
    assert (nat["policy_id"], nat["policy_name"], nat["rule_id"], nat["position"]) == (
        "nat-1", "Edge NAT", "nat-rule-1", 3,
    )


def test_excel_source_inventory_exposes_structured_ownership_columns():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "devices": [{"id": "device-1", "name": "FTD-A", "resources": {
            "virtual_routers": [{"id": "vr-1", "name": "VR-A", "resources": {}}]
        }}],
    }))
    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)
    output.seek(0)
    workbook = load_workbook(output, read_only=True)
    headers = [cell.value for cell in workbook["Source Inventory"][1]]
    assert {"Domain", "Device", "Device Name", "Virtual Router ID", "Virtual Router",
            "Parent Policy ID", "Parent Policy"} <= set(headers)


def test_preview_keeps_fmc_policy_and_rule_identity_separate():
    result = CiscoFTDSourceReporter().analyze_source(FIXTURE.read_text(encoding="utf-8"))
    policy = result.config.access_control_policies[0]
    rule = policy.rules[0]
    row = build_ftd_preview(result)["sections"]["policies"][0]

    assert row["policy_id"] == policy.source_id == "acp-1"
    assert row["policy_name"] == policy.name == "Access"
    assert row["rule_id"] == rule.source_id == "r-1"
    assert row["name"] == rule.name == "Allow updates"
    assert row["position"] == rule.position
    assert row["policy_id"] != row["position"]


def test_validation_findings_keep_stable_fmc_scope_identity():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "access_policies": [{
            "id": "policy-1", "name": "Policy",
            "rules": [
                {"id": "rule-1", "name": "First", "position": 4, "action": "ALLOW"},
                {"id": "rule-2", "name": "Second", "position": 4, "action": "BLOCK"},
            ],
        }],
    }))
    issue = next(item for item in result.validation.issues if item.category == "duplicate-acp-position")
    assert issue.source_id == "policy-1"
    assert issue.source_context == "fmc:Global"
    assert issue.domain_id == "domain-1"

    row = next(item for item in build_ftd_preview(result)["sections"]["validation"]
               if item["domain"] == "duplicate-acp-position")
    assert row["source_id"] == "policy-1"
    assert row["source_context"] == "fmc:Global"
    assert row["domain_id"] == "domain-1"


def test_source_inventory_exports_policy_ownership_and_native_types():
    result = CiscoFTDSourceReporter().analyze_source(FIXTURE.read_text(encoding="utf-8"))
    rule_item = next(item for item in result.inventory_items if item.source_id == "r-1")
    assert rule_item.source_attributes["policy_id"] == "acp-1"
    assert rule_item.source_record_id.startswith("fmc-rest-bundle/domain=domain-1/")
    assert "/parent=acp-1/" in rule_item.source_record_id

    address_item = next(item for item in result.inventory_items
                        if item.source_path.endswith("/network-addresses"))
    assert address_item.source_type == "network_address"

    output = BytesIO()
    CiscoFTDSourceReporter().export_excel(result, output)
    workbook = load_workbook(output, read_only=True)
    headers = [cell.value for cell in workbook["Source Inventory"][1]]
    for column in ("Virtual Router ID", "Virtual Router", "Parent Policy ID", "Parent Policy"):
        assert column in headers
    rows = [dict(zip(headers, row)) for row in
            workbook["Source Inventory"].iter_rows(min_row=2, values_only=True)]
    exported_rule = next(row for row in rows if row["Source ID"] == "r-1")
    assert exported_rule["Parent Policy ID"] == "acp-1"


def test_unsupported_evidence_does_not_cross_associate_duplicate_scoped_names():
    from fwmigrate.vendors.cisco_ftd.model import CiscoFTDConfig, CiscoFTDNetworkAddress
    from fwmigrate.vendors.cisco_ftd.source_inventory import build_ftd_source_inventory

    first = CiscoFTDNetworkAddress(
        name="duplicate", source_id="id-1", source_plane="fmc-rest-bundle", domain_id="domain-1")
    second = CiscoFTDNetworkAddress(
        name="duplicate", source_id="id-2", source_plane="fmc-rest-bundle", domain_id="domain-2")
    config = CiscoFTDConfig(
        input_source_type="fmc-rest-bundle", source_plane="fmc-rest-bundle",
        network_addresses=[first, second],
        unsupported_evidence=[{"source_name": "duplicate", "reason": "ambiguous name-only evidence"}],
    )
    items = [item for item in build_ftd_source_inventory(config)
             if item.source_path.endswith("/network-addresses")]
    assert [item.requires_manual_review for item in items] == [False, False]

    config.unsupported_evidence = [{
        "source_id": "id-2", "domain_id": "domain-2", "reason": "scoped evidence",
    }]
    items = [item for item in build_ftd_source_inventory(config)
             if item.source_path.endswith("/network-addresses")]
    assert [item.requires_manual_review for item in items] == [False, True]
    assert len({item.source_record_id for item in items}) == 2


def test_fmc_unnamed_policy_does_not_manufacture_child_policy_name():
    result = extract_cisco_ftd_source(json.dumps({
        "source": "fmc-rest-api",
        "domain": {"id": "domain-1", "name": "Global"},
        "access_policies": [{
            "id": "policy-without-name",
            "rules": [{"id": "rule-1", "name": "Rule", "action": "ALLOW"}],
        }],
    }))
    policy = result.config.access_control_policies[0]
    rule = policy.rules[0]
    assert policy.name is None
    assert rule.policy_id == "policy-without-name"
    assert rule.policy_name is None
    assert rule.source_attributes["policy_name"] is None

    preview = build_ftd_preview(result)["sections"]["policies"][0]
    assert preview["policy_id"] == "policy-without-name"
    assert preview["policy_name"] is None

    inventory = next(item for item in result.inventory_items if item.source_id == "rule-1")
    assert inventory.source_attributes["policy_id"] == "policy-without-name"
    assert inventory.source_attributes["policy_name"] is None


def test_excel_preserves_interfaces_with_partial_and_scoped_topology():
    from copy import deepcopy
    from dataclasses import replace
    from fwmigrate.vendors.cisco_ftd.relationships.interface_topology import (
        FTDInterfaceTopology, FTDInterfaceTopologyEntry,
    )

    result = extract_cisco_ftd_source(
        "interface GigabitEthernet0/0\n ip address 192.0.2.1 255.255.255.0\n"
        "interface Management0/0\n ip address 198.51.100.1 255.255.255.0\n")
    before = deepcopy(result.config.model_dump(mode="python"))
    for entries in (result.derived.interface_topology.interfaces,
                    result.derived.interface_topology.interfaces[:1], (),
                    (FTDInterfaceTopologyEntry("Management0/0", "subinterface", "wrong-parent", device_id="other"),)):
        analysis = replace(result, derived=replace(result.derived, interface_topology=FTDInterfaceTopology(entries)))
        for profile in ("full", "fast"):
            output = BytesIO()
            CiscoFTDSourceReporter().export_excel(analysis, output, profile=profile)
            workbook = load_workbook(output, read_only=True)
            headers, *rows = list(workbook["Interfaces"].iter_rows(max_col=9, values_only=True))
            assert [row[0] for row in rows] == [item.name for item in result.config.interfaces]
            for row, source in zip(rows, result.config.interfaces):
                assert row[3] == source.ip
                expected = next((item for item in entries if item.name == source.name and not item.device_id), None)
                assert row[6:8] == ((expected.kind, expected.parent) if expected else (None, None))
            completeness = list(workbook["Collection Completeness"].iter_rows(max_col=5, values_only=True))
            assert completeness[1] == ("Collection metadata", "NOT_PROVIDED", None, None, None)
            workbook.close()
    assert result.config.model_dump(mode="python") == before


def test_ftd_zone_references_are_json_with_names_and_ids_in_both_profiles():
    result = extract_cisco_ftd_source(json.dumps({"source": "fmc-rest-api", "objects": {
        "securityzones": [{"id": "zone-1", "name": "trust", "interfaces": [
            {"id": "if-1", "name": "inside"}, {"id": "if-2"}]}]}}))
    for profile in ("full", "fast"):
        output = BytesIO()
        CiscoFTDSourceReporter().export_excel(result, output, profile=profile)
        workbook = load_workbook(output, read_only=True)
        row = next(workbook["Zones"].iter_rows(min_row=2, values_only=True))
        refs = json.loads(row[1])
        assert refs[0]["name"] == "inside" and refs[0]["source_id"] == "if-1"
        assert refs[1]["source_id"] == "if-2"
        workbook.close()
