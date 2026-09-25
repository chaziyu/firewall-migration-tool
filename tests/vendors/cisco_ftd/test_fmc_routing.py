import json
from copy import deepcopy
from copy import deepcopy
from types import SimpleNamespace
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

from fwmigrate.vendors.cisco_ftd.fmc.fmc_adapter import CiscoFMCBundleParser
from fwmigrate.vendors.cisco_ftd.source_report import CiscoFTDSourceReporter, extract_cisco_ftd_source
from fwmigrate.vendors.cisco_ftd.derived import build_ftd_derived_views
from fwmigrate.vendors.cisco_ftd.validation import validate_ftd_config
from fwmigrate.vendors.cisco_ftd.model import CiscoFTDReference
from fwmigrate.vendors.cisco_ftd.export.excel import export_ftd_excel
from fwmigrate.vendors.cisco_ftd.web_report import build_ftd_preview


FIXTURE = Path(__file__).parents[2] / "fixtures" / "cisco_ftd" / "fmc_selected_domains.json"

def test_fmc_dhcp_interface_reference_is_typed_scoped_and_keeps_missing_state():
    payload = json.loads(FIXTURE.read_text(encoding="utf-8"))
    payload["devices"][0]["resources"]["dhcp_servers"][0]["vendorField"] = "preserved"
    config = CiscoFMCBundleParser(json.dumps(payload)).parse_source()
    server = config.dhcp_servers[0]

    assert server.interface.name == "inside"
    assert server.explicit_fields == ["interface"]
    assert server.raw_extra["interfaceName"] == "inside"
    assert server.raw_extra["vendorField"] == "preserved"
    assert server.source_attributes["device_id"] == "device-1"
    derived = build_ftd_derived_views(config)
    assert any(issue.owner == server.name and issue.field == "interface"
               for issue in derived.unresolved_references)
    assert any(issue.source_object == server.name and issue.category == "unresolved-reference"
               and "in interface" in issue.message for issue in validate_ftd_config(config, derived).issues)

    del payload["devices"][0]["resources"]["dhcp_servers"][0]["interfaceName"]
    missing = CiscoFMCBundleParser(json.dumps(payload)).parse_source().dhcp_servers[0]
    assert missing.interface is None
    assert "interface" not in missing.explicit_fields

def test_fmc_override_acp_assignment_ips_and_dhcp_source_state_remain_separate():
    payload = {
        "format": "cisco-fmc-rest-export-v1", "domain": {"id": "d1", "name": "Global"},
        "access_policies": [
            {"id": "base-acp", "name": "Base", "rules": []},
            {"id": "child-acp", "name": "Child", "description": "Child policy", "metadata": {
                "inherit": False, "parentPolicy": {"id": "base-acp", "name": "Base", "type": "AccessPolicy"}},
             "defaultAction": {"id": "default-1", "name": "Default", "type": "AccessPolicyDefaultAction"},
             "identityPolicy": {"id": "identity-1", "name": "Identity", "type": "IdentityPolicy"},
             "logging_settings": {"logAtBeginning": True},
             "default_actions": [{"id": "default-1", "name": "Default", "action": "BLOCK"}], "rules": []},
        ],
        "objects": {"networkaddresses": [{"id": "net-1", "name": "Shared", "type": "Host", "value": "10.0.0.1"}],
            "identitypolicies": [{"id": "identity-1", "name": "Identity"}],
            "network_address_overrides": [
                {"id": "ov-1", "name": "Shared", "type": "Host", "value": "10.0.0.2",
                 "overridable": True, "overrides": {"parent": {"id": "net-1", "name": "Shared", "type": "Host"},
                    "target": {"id": "dev-1", "name": "FTD-A", "type": "Device"}}, "password": "must-not-leak"},
                {"id": "ov-2", "name": "Shared", "type": "Host", "value": "10.0.0.3",
                 "overrides": {"parent": {"id": "net-1", "name": "Shared", "type": "Host"},
                    "target": {"id": "dev-2", "name": "FTD-B", "type": "Device"}}}],
            "intrusionpolicies": [{"id": "ips-1", "name": "IPS", "rules": [
                {"id": "beh-1", "ruleId": "sig-1", "state": "enabled", "action": "alert"}],
                "overrides": [{"id": "ov-rule-1", "ruleId": "sig-1", "state": "disabled"}]}],
            "policy_assignments": [{"id": "assignment-1", "name": "ACP assignment", "policy":
                {"id": "child-acp", "name": "Child", "type": "AccessPolicy"}, "targets": [
                {"id": "dev-1", "name": "FTD-A", "type": "Device"}, {"id": "dev-2", "name": "FTD-B", "type": "Device"}]}]},
        "devices": [{"id": "dev-1", "name": "FTD-A", "resources": {
            "dhcp_servers": [{"id": "dhcp-1", "name": "DHCP", "interfaceName": "inside", "addressPool": "10.0.0.10-10.0.0.20"}],
            "dhcp_relay_settings": [{"id": "relay-1", "name": "Relay", "dhcpRelayAgent": [], "dhcpRelayServers": [],
                "ipv4TimeoutInSec": "20", "trustAllInformation": True}],
            "ftd_interfaces": [{"id": "if-1", "name": "GigabitEthernet0/0.10", "interfaceType": "SubInterface"}],
            "static_routes": [{"id": "route-1", "name": "inside-route", "network": {"id": "net-1", "name": "Shared", "type": "Host"}, "addressFamily": "IPv4"}],
        }}],
    }
    result = extract_cisco_ftd_source(json.dumps(payload))
    config = result.config

    assert config.network_addresses[0].value == "10.0.0.1"
    assert [item.value for item in config.network_address_overrides] == ["10.0.0.2", "10.0.0.3"]
    assert config.network_address_overrides[0].parent.source_id == "net-1"
    assert config.network_address_overrides[0].target.name == "FTD-A"
    assert config.access_control_policies[1].inherit is False
    assert {"inherit", "base_policy"} <= set(config.access_control_policies[1].explicit_fields)
    assert config.access_control_policies[1].logging_settings == {"logAtBeginning": True}
    assert config.access_control_policies[1].base_policy.source_id == "base-acp"
    assert len(config.access_control_default_actions) == 1
    assert len(config.policy_assignments[0].targets) == 2
    assert any(item["field"] == "identity_policy" and item["target_id"] == "identity-1"
               for item in result.derived.resolved_references)
    assert [(item.rule_id, item.state) for item in config.intrusion_rule_overrides] == [("sig-1", "disabled")]
    assert config.dhcp_relay_settings[0].ipv4_timeout_seconds == "20"
    assert "dhcpRelayAgent" not in config.dhcp_servers[0].raw_extra
    assert not any(item.source_attributes.get("resource_type") == "dhcp_relay_settings" for item in config.native_resources)

    assert any(item["field"] == "parent" and item["target_id"] == "net-1"
               for item in result.derived.resolved_references)
    assert any(item.source_name == "inside-route" and item.normalized_destination == "10.0.0.1/32"
               for item in result.derived.normalized_routes), result.derived.normalized_routes
    assert any(item.category == "dhcp-server-relay-conflict" for item in result.validation.issues)
    assert any(item.device_id == "dev-1" and item.name == "GigabitEthernet0/0.10"
               for item in result.derived.interface_topology.interfaces)

    preview = result
    output = BytesIO()
    export_ftd_excel(preview, output)
    output.seek(0)
    workbook = load_workbook(output, read_only=True)
    assert workbook["Object Overrides"].max_row == 3
    assert workbook["ACP Policies"].max_row == 3
    assert workbook["DHCP"].max_row == 3
    assert "must-not-leak" not in json.dumps(config.model_dump())

def test_fmc_routes_preserve_selected_networks_family_provenance_and_ownership():
    payload = {"format": "cisco-fmc-rest-export-v1", "domain": {"id": "d1"},
        "collection": {"status": "PARTIAL", "parts": [
            {"name": "device/dev1/ipv4_static_routes", "status": "SUCCESS", "complete": True, "count": 2},
            {"name": "device/dev1/ipv6_static_routes", "status": "SUCCESS", "complete": True, "count": 1},
            {"name": "device/dev1/virtual_router/vr1/ipv4_static_routes", "status": "SUCCESS", "complete": True, "count": 1},
            {"name": "device/dev1/virtual_router/vr1/ipv6_static_routes", "status": "FAILED", "complete": False, "count": 0},
        ]},
        "objects": {"networkaddresses": [
            {"id": "net4", "name": "net4", "type": "Network", "value": "192.0.2.0/24"},
            {"id": "net4b", "name": "net4b", "type": "Network", "value": "198.51.100.0/24"},
            {"id": "net6", "name": "net6", "type": "Network", "value": "2001:db8::/64"},
            {"id": "gw", "name": "gateway", "type": "Host", "value": "192.0.2.1"}],
            "slamonitors": [{"id": "sla1", "name": "WAN SLA"}]},
        "devices": [{"id": "dev1", "name": "FTD-A", "resources": {
            "ipv4_static_routes": [
                {"id": "multi", "name": "same", "selectedNetworks": [
                    {"id": "net4", "name": "net4"}, {"id": "net4b", "name": "net4b"}],
                 "gateway": {"literal": {"value": "192.0.2.1"}},
                 "routeTracking": {"slaMonitor": {"id": "sla1", "name": "WAN SLA"}}},
                {"id": "drop", "name": "drop", "network": {"id": "net4", "name": "net4"}, "interfaceName": "Null0"}],
            "ipv6_static_routes": [{"id": "v6", "name": "v6", "selectedNetworks": [{"id": "net6", "name": "net6"}],
                "routeTracking": {"slaMonitor": {"id": "sla1", "name": "WAN SLA"}}}],
            "virtual_routers": [{"id": "vr1", "name": "VR1", "resources": {
                "ipv4_static_routes": [{"id": "vr-route", "name": "same", "network": {"id": "net4", "name": "net4"},
                    "gateway": {"object": {"id": "gw", "name": "gateway", "type": "Host"}}}],
                "ipv6_static_routes": []}}]}}]}
    result = extract_cisco_ftd_source(json.dumps(payload))
    global_route = next(route for route in result.config.routes if route.source_id == "multi")
    assert global_route.destination is None
    assert [item.source_id for item in global_route.selected_networks] == ["net4", "net4b"]
    assert global_route.address_family is None and "address_family" not in global_route.explicit_fields
    assert global_route.source_attributes["collection_address_family"] == "IPv4"
    assert result.derived.source_plane_completeness["routing"] == "partial"
    normalized = [item for item in result.derived.normalized_routes if item.source_id == "multi"]
    assert [item.normalized_destination for item in normalized] == ["192.0.2.0/24", "198.51.100.0/24"]
    assert all(item.destination_index in {1, 2} for item in normalized)
    assert not any(item.owner == "same" and item.field == "gateway" for item in result.derived.unresolved_references)
    assert not any(item.owner == "drop" and item.field == "interface" for item in result.derived.unresolved_references)
    assert any(item.source_object == "v6" and item.category == "ipv6-route-tracking"
               for item in result.validation.issues)
    vr_route = next(route for route in result.config.routes if route.source_id == "vr-route")
    assert vr_route.virtual_router == "VR1" and vr_route.source_attributes["device_id"] == "dev1"
    build_ftd_derived_views(result.config)

def test_route_sla_monitor_is_a_typed_resolvable_source_reference():
    payload = {
        "domain": {"id": "domain-1", "name": "Global"},
        "objects": {"slamonitors": [{"id": "sla-1", "name": "WAN-Monitor"}]},
        "devices": [{"id": "device-1", "name": "FTD-A", "resources": {"static_routes": [
            {"id": "route-1", "name": "tracked-route", "slaMonitor": {
                "id": "sla-1", "name": "WAN-Monitor", "type": "SLAMonitor"}},
            {"id": "route-2", "name": "untracked-route"},
            {"id": "route-3", "name": "null-monitor-route", "slaMonitor": None},
            {"id": "route-4", "name": "missing-monitor-route", "slaMonitor": {
                "id": "sla-missing", "name": "missing"}},
        ]}}],
    }
    config = CiscoFMCBundleParser(json.dumps(payload)).parse_source()
    tracked, untracked, null_monitor, missing = config.routes

    assert isinstance(tracked.sla_monitor, CiscoFTDReference)
    assert (tracked.sla_monitor.source_id, tracked.sla_monitor.name, tracked.sla_monitor.source_type) == (
        "sla-1", "WAN-Monitor", "SLAMonitor")
    assert config.sla_monitors[0].source_id == "sla-1"
    assert not any(item.source_attributes.get("resource_type") == "slamonitors" for item in config.native_resources)
    assert untracked.sla_monitor is None
    assert null_monitor.sla_monitor is None
    assert isinstance(missing.sla_monitor, CiscoFTDReference)

    derived = build_ftd_derived_views(config)
    assert any(item["owner"] == tracked.name and item["field"] == "sla_monitor"
               and item["kind"] == "SLA_MONITOR" for item in derived.resolved_references)
    issues = validate_ftd_config(config, derived).issues
    assert not any(item.source_object == tracked.name and "sla_monitor" in item.message for item in issues)
    assert not any(item.source_object == untracked.name and "sla_monitor" in item.message for item in issues)
    assert any(item.source_object == missing.name and "sla_monitor" in item.message
               and item.category == "unresolved-reference" for item in issues)

    output = BytesIO()
    export_ftd_excel(SimpleNamespace(config=config, validation=SimpleNamespace(issues=issues)), output)
    sheet = load_workbook(output, read_only=True)["Routes"]
    assert [cell.value for cell in sheet[1]] == ["Name", "Device", "Virtual Router", "Interface", "Destination", "Gateway", "SLA Monitor",
        "Route Tracking", "Tunneled", "Source Plane", "Address Family", "Mask", "Normalized Destination(s)"]
    assert sheet[2][6].value == "WAN-Monitor"

def test_fmc_route_normalization_and_interface_topology_keep_device_scope():
    payload = {"format": "cisco-fmc-rest-export-v1", "domain": {"id": "d1"},
        "objects": {"networkaddresses": [
            {"id": "net4", "name": "IPv4 network", "type": "Network", "value": "10.0.0.0/24"},
            {"id": "net6", "name": "IPv6 network", "type": "Network", "value": "2001:db8::/64"}],
            "network_address_overrides": []},
        "devices": [
            {"id": "dev1", "name": "FTD-A", "resources": {"ftd_interfaces": [
                {"id": "if1", "name": "GigabitEthernet0/0", "interfaceType": "PhysicalInterface"},
                {"id": "if2", "name": "GigabitEthernet0/0.10", "interfaceType": "SubInterface"}],
                "static_routes": [
                    {"id": "r4", "name": "route4", "network": {"id": "net4", "name": "IPv4 network"}, "addressFamily": "iPv4"},
                    {"id": "r6", "name": "route6", "network": {"id": "net6", "name": "IPv6 network"}, "addressFamily": "IPv6"},
                    {"id": "null", "name": "null-route", "network": {"id": "net4", "name": "IPv4 network"}, "gateway": "Null0", "addressFamily": "IPv4"},
                    {"id": "missing", "name": "unresolved", "network": {"id": "missing", "name": "Missing"}, "addressFamily": "IPv4"},
                    {"id": "selected", "name": "selected-route", "selectedNetworks": [{"id": "net4", "name": "IPv4 network"}],
                     "gateway": {"literal": {"value": "10.0.0.1"}}, "routeTracking": {"slaMonitor": {"id": "sla1", "name": "SLA"}},
                     "isTunneled": True, "addressFamily": "IPv4"},
                    {"id": "mismatch", "name": "mismatch-route", "selectedNetworks": [{"id": "net6", "name": "IPv6 network"}],
                     "addressFamily": "IPv4"}]}},
            {"id": "dev2", "name": "FTD-B", "resources": {"ftd_interfaces": [
                {"id": "if3", "name": "GigabitEthernet0/0", "interfaceType": "PhysicalInterface"},
                {"id": "if4", "name": "GigabitEthernet0/0.10", "interfaceType": "SubInterface"}]}}]}
    result = extract_cisco_ftd_source(json.dumps(payload))
    normalized = {item.source_name: item for item in result.derived.normalized_routes}
    assert normalized["route4"].normalized_destination == "10.0.0.0/24"
    assert normalized["route6"].normalized_destination == "2001:db8::/64"
    assert normalized["null-route"].gateway == "Null0"
    assert normalized["unresolved"].normalized_destination is None
    assert normalized["selected-route"].normalized_destination == "10.0.0.0/24"
    assert normalized["selected-route"].gateway == "10.0.0.1"
    assert normalized["selected-route"].issue is None
    assert normalized["mismatch-route"].issue
    assert any(item.source_object == "mismatch-route" and item.category == "invalid-route"
               for item in result.validation.issues)
    subinterfaces = [item for item in result.derived.interface_topology.interfaces
                     if item.name == "GigabitEthernet0/0.10"]
    assert {(item.device_id, item.parent, item.kind) for item in subinterfaces} == {
        ("dev1", "GigabitEthernet0/0", "subinterface"),
        ("dev2", "GigabitEthernet0/0", "subinterface")}


def test_fmc_virtual_router_pbr_and_ecmp_keep_device_scope():
    config = CiscoFMCBundleParser(json.dumps({"source": "fmc-rest-api", "devices": [{"id": "d1", "name": "FTD",
        "resources": {"pbr_policies": [{"id": "global-pbr", "name": "Global PBR"}],
            "ecmp_zones": [{"id": "global-ecmp", "name": "Global ECMP"}],
            "virtual_routers": [{"id": "vr-a", "name": "VR-A", "resources": {
                "pbr_policies": [{"id": "vr-pbr", "name": "VR PBR"}],
                "ecmp_zones": [{"id": "vr-ecmp", "name": "VR ECMP"}]}}]}}]})).parse_source()
    assert [item.source_id for item in config.virtual_routers] == ["vr-a"]
    pbr = {item.source_id: item for item in config.policy_based_routes}
    assert set(pbr) == {"global-pbr", "vr-pbr"}
    assert (pbr["vr-pbr"].device_id, pbr["vr-pbr"].source_attributes["virtual_router_id"],
            pbr["vr-pbr"].source_attributes["virtual_router_name"]) == ("d1", "vr-a", "VR-A")
    assert {item.source_id for item in config.ecmp_zones} == {"global-ecmp", "vr-ecmp"}
