import json
from io import BytesIO
from openpyxl import load_workbook

from fwmigrate.extraction.sanitize import sanitize_source_attributes
from fwmigrate.vendors.cisco_ftd.derived import build_ftd_derived_views
from fwmigrate.vendors.cisco_ftd.export.excel import export_ftd_excel
from fwmigrate.vendors.cisco_ftd.source_report import extract_cisco_ftd_source
from fwmigrate.vendors.cisco_ftd.validation import validate_ftd_config
from fwmigrate.vendors.cisco_ftd.web_report import build_ftd_preview


def test_fmc_secret_fields_are_removed_and_presence_metadata_survives():
    safe = sanitize_source_attributes({
        "password": "pw", "token": "tok", "bindPassword": "ldap", "radiusSecret": "radius",
        "preSharedKey": "psk", "vpnSecret": "vpn", "privateKey": "key",
        "credential": "credential", "eventHubsConnString": "connection-string",
        "preSharedKeyConfigured": True, "privateKeyPresent": True,
    })
    for secret in ("pw", "tok", "ldap", "radius", "psk", "vpn", "key", "credential", "connection-string"):
        assert secret not in [value for key, value in safe.items() if key not in {"preSharedKeyConfigured", "privateKeyPresent"}]
    assert safe["preSharedKeyConfigured"] is True
    assert safe["privateKeyPresent"] is True


def test_fmc_secrets_stay_redacted_through_ftd_reports():
    secrets = ("realm-secret", "local-secret", "private-secret", "p12-secret", "s2s-secret", "ra-secret")
    payload = {"format": "cisco-fmc-rest-export-v1", "domain": {"id": "d1", "name": "Global"},
        "objects": {
            "realms": [{"id": "realm-1", "name": "Directory", "directoryConfigurations": [
                {"hostname": "ldap.example.test", "dirPassword": secrets[0]}]}],
            "localrealmusers": [{"id": "local-1", "name": "alice", "password": secrets[1]}],
            "internalcertificates": [{"id": "cert-1", "name": "VPN cert", "privateKey": secrets[2],
                "pkcs12Password": secrets[3]}],
            "s2svpns": [{"id": "s2s-1", "name": "Branch", "ike_settings": [{"id": "ike-1", "name": "IKE",
                "preSharedKey": secrets[4]}]}],
            "ravpns": [{"id": "ra-1", "name": "Remote", "secure_client_customization_settings": [
                {"id": "secure-1", "name": "Secure Client", "token": secrets[5]}]}],
        }}
    result = extract_cisco_ftd_source(json.dumps(payload))
    derived = build_ftd_derived_views(result.config)
    validation = validate_ftd_config(result.config, derived)
    preview = build_ftd_preview(result)
    output = BytesIO()
    export_ftd_excel(result, output)
    workbook = load_workbook(output, read_only=True)
    workbook_data = [[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets]
    evidence = json.dumps((result.config.model_dump(), result.inventory_items, result.config.native_resources,
        derived, validation, preview, workbook_data), default=str)
    assert all(secret not in evidence for secret in secrets)
    assert result.config.local_realm_users[0].password_configured is True
    assert result.config.certificates[0].private_key_present is True
    assert result.config.s2s_ike_settings[0].psk_present is True
