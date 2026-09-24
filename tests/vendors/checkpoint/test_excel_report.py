import io
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.checkpoint.source_report import extract_checkpoint_source
from fwmigrate.vendors.checkpoint.export.excel import export_checkpoint_excel
from fwmigrate.vendors.checkpoint.export.excel_schema import SHEET_ORDER


def test_excel_report_uses_typed_source_and_derived_sheets():
    source = (Path(__file__).parents[2] / "fixtures" / "checkpoint" / "multidomain_full.json").read_text()
    result = extract_checkpoint_source(source)
    before = (result.config.model_dump(), deepcopy(result.derived), deepcopy(result.validation))
    output = io.BytesIO(); export_checkpoint_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    assert {"Hosts", "Networks", "Access Sections", "Access Rules", "NAT Rules", "Gaia Interfaces",
            "NAT Migration Views", "Policy Traversal", "Interface Views", "VPN Views",
            "Check Point Source Inventory", "Review Required"} <= set(workbook.sheetnames)
    assert "Network Objects" not in workbook.sheetnames and "Gaia" not in workbook.sheetnames
    assert "IP Pools" not in SHEET_ORDER and "VIPs" not in SHEET_ORDER and "SD-WAN" not in SHEET_ORDER
    assert (result.config.model_dump(), result.derived, result.validation) == before


def test_excel_formula_like_source_text_is_safe():
    result = extract_checkpoint_source('{"objects":[{"type":"host","name":"=1+1"}]}')
    output = io.BytesIO(); export_checkpoint_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    sheet = workbook["Hosts"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    name_col = headers.index("name")
    assert list(sheet.iter_rows(min_row=2, values_only=True))[0][name_col] == "'=1+1"
