import io
from copy import deepcopy
from pathlib import Path
from openpyxl import load_workbook
from fwmigrate.vendors.checkpoint.source_report import extract_checkpoint_source
from fwmigrate.vendors.checkpoint.export.excel import export_checkpoint_excel
from fwmigrate.vendors.checkpoint.export.excel_schema import SHEET_ORDER
from fwmigrate.vendors.checkpoint.web_report import build_checkpoint_preview

def test_excel_report_uses_typed_source_and_derived_sheets():
    source = (Path(__file__).parents[2] / "fixtures" / "checkpoint" / "multidomain_full.json").read_text()
    result = extract_checkpoint_source(source)
    before = (result.config.model_dump(), deepcopy(result.derived), deepcopy(result.validation))
    output = io.BytesIO(); export_checkpoint_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    assert {"Hosts", "Networks", "Access Sections", "Access Rules", "NAT Rules", "Gaia Interfaces",
            "Derived Completeness", "NAT Migration Views", "Policy Traversal", "Interface Views", "VPN Views",
            "Check Point Source Inventory", "Review Required"} <= set(workbook.sheetnames)
    assert "Network Objects" not in workbook.sheetnames and "Gaia" not in workbook.sheetnames
    assert "IP Pools" not in SHEET_ORDER and "VIPs" not in SHEET_ORDER and "SD-WAN" not in SHEET_ORDER
    assert (result.config.model_dump(), result.derived, result.validation) == before

def test_excel_formula_like_source_text_is_safe():
    result = extract_checkpoint_source('{"objects":[{"type":"host","name":"=1+1"}]}')
    output = io.BytesIO(); export_checkpoint_excel(result, output)
    workbook = load_workbook(io.BytesIO(output.getvalue()), read_only=True)
    sheet = workbook["Hosts"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    name_col = headers.index("Name")
    assert list(sheet.iter_rows(min_row=2, values_only=True))[0][name_col] == "'=1+1"

def test_web_report_has_typed_source_derived_validation_and_traceability():
    source = (Path(__file__).parents[2] / "fixtures" / "checkpoint" / "multidomain_full.json").read_text()
    result = extract_checkpoint_source(source)
    before = (result.config.model_dump(), deepcopy(result.derived), deepcopy(result.validation))
    preview = build_checkpoint_preview(result)
    assert preview["vendor"] == "checkpoint"
    assert "access_rules" in preview["source"]
    assert {"nat", "policy_traversal", "interfaces", "vpn", "completeness"} <= set(preview["derived"])
    assert {"scope", "collection", "source_inventory", "unsupported", "validation"} <= set(preview)
    assert "No direct R81.00 equivalent" == preview["summary"]["capabilities"]["SD-WAN"]
    assert {"nat", "policy", "interfaces", "vpn"} <= set(preview["summary"]["derived_source_complete"])
    assert preview["summary"]["scopes"]
    assert {"interfaces", "addresses", "address_groups", "services", "service_groups", "schedules",
            "policies", "nat", "routes", "vpn_tunnels", "vpn_phase2", "validation",
            "unresolved_references"} <= set(preview["sections"])
    assert all("package" in row and "layer" in row and "section" in row
               for row in preview["sections"]["policies"])
    if preview["source"]["nat_rules"]:
        assert "original_source" in preview["source"]["nat_rules"][0]
    assert (result.config.model_dump(), result.derived, result.validation) == before

def test_nat_transform_is_exposed_in_preview_and_excel():
    result = extract_checkpoint_source(
        '{"responses":[{"command":"show-hosts","data":{"objects":[{"type":"host","uid":"h","name":"Web","nat-settings":{"auto-rule":true}}]}}]}'
    )

    assert result.config.hosts[0].nat_settings == {"auto-rule": True}
    preview = build_checkpoint_preview(result)
    assert preview["derived"]["nat"][0]["source_kind"] == "automatic_object_settings"

    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    output.seek(0)
    sheet = load_workbook(output, read_only=True)["NAT Migration Views"]
    headers = next(sheet.iter_rows(values_only=True))
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Source Kind")] == "automatic_object_settings"
    assert row[headers.index("Resolved Translation")] in (None, "")


def test_dual_stack_preview_keeps_both_explicit_families_and_excel_labels_are_readable():
    result = extract_checkpoint_source('{"responses":[{"command":"show-hosts","data":{"objects":[{"uid":"dual","name":"dual","type":"host","ipv4-address":"192.0.2.4","ipv6-address":"2001:db8::4"}]}}]}')
    address = next(row for row in build_checkpoint_preview(result)["sections"]["addresses"] if row["name"] == "dual")
    assert address["ipv4_value"] == "192.0.2.4"
    assert address["ipv6_value"] == "2001:db8::4"
    assert address["address_family"] == "Dual Stack"

    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    headers = next(load_workbook(io.BytesIO(output.getvalue()), read_only=True)["Hosts"].iter_rows(values_only=True))
    assert {"Name", "Domain", "UID", "IPv4 Address", "IPv6 Address", "Additional Settings"} <= set(headers)


def test_fast_and_full_exports_keep_the_same_checkpoint_source_rows():
    result = extract_checkpoint_source('{"responses":[{"command":"show-hosts","collection_status":"SUCCESS_WITH_DATA","data":{"objects":[{"uid":"h","name":"web","type":"host","ipv4-address":"192.0.2.8"}]}}],"collection_completeness":{"show-hosts":{"command":"show-hosts","status":"SUCCESS_WITH_DATA","complete":true}}}')
    books = []
    for profile in ("fast", "full"):
        output = io.BytesIO()
        export_checkpoint_excel(result, output, profile=profile)
        books.append(load_workbook(io.BytesIO(output.getvalue()), read_only=True))
    fast, full = books
    assert next(fast["Hosts"].iter_rows(values_only=True)) == next(full["Hosts"].iter_rows(values_only=True))
    width = len(next(fast["Hosts"].iter_rows(values_only=True)))
    fast_rows = [row + (None,) * (width - len(row)) for row in fast["Hosts"].iter_rows(min_row=2, values_only=True)]
    full_rows = [row + (None,) * (width - len(row)) for row in full["Hosts"].iter_rows(min_row=2, values_only=True)]
    assert fast_rows == full_rows



def test_vpn_community_name_is_not_redacted_in_preview():
    result = extract_checkpoint_source(
        '{"responses":[{"command":"show-vpn-communities-star","data":{"objects":[{"uid":"vpn1","name":"Partner VPN","participating-gateways":[]}]}}]}'
    )
    preview = build_checkpoint_preview(result)

    assert preview["derived"]["vpn"][0]["vpn_community"] == "Partner VPN"
    assert preview["sections"]["vpn_tunnels"][0]["vpn_community"] == "Partner VPN"
