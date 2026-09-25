import io

from openpyxl import load_workbook

from fwmigrate.vendors.checkpoint.source_report import extract_checkpoint_source
from fwmigrate.vendors.checkpoint.source_report import CheckPointSourceReporter
from fwmigrate.vendors.checkpoint.export.excel import export_checkpoint_excel


def test_secret_bearing_source_is_redacted_from_preview_and_excel():
    secret = "checkpoint-secret-regression"
    result = extract_checkpoint_source(f'{{"objects":[{{"type":"host","name":"x","password":"{secret}"}}]}}')
    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    values = " ".join(str(cell.value) for sheet in load_workbook(io.BytesIO(output.getvalue()), read_only=True).worksheets for row in sheet.iter_rows() for cell in row)
    preview = CheckPointSourceReporter().build_preview(result)
    evidence = str(result.config.model_dump()) + str(result.derived) + str(result.validation) + str(preview) + str(result.source_inventory) + values
    assert secret not in evidence


def test_gaia_secret_is_redacted_everywhere():
    secret = "gaia-secret-regression"
    result = extract_checkpoint_source(f"set user admin password {secret}\nset user admin shell /bin/bash\n")
    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    values = " ".join(str(cell.value) for sheet in load_workbook(io.BytesIO(output.getvalue()), read_only=True).worksheets for row in sheet.iter_rows() for cell in row)
    evidence = str(result.config.model_dump()) + str(result.source_inventory) + str(result.derived) + str(result.validation) + str(CheckPointSourceReporter().build_preview(result)) + values
    assert secret not in evidence


def test_management_user_and_admin_secrets_are_redacted_everywhere():
    secret = "management-identity-secret"
    source = ('{"responses": ['
        '{"command":"show-users","data":{"objects":[{"uid":"u","name":"alice","type":"user","password":"' + secret + '","password-hash":"' + secret + '","token":"' + secret + '","future":"keep"}] }},'
        '{"command":"show-administrators","data":{"objects":[{"uid":"a","name":"admin","type":"administrator","shared-secret":"' + secret + '","challenge-response":"' + secret + '"}] }}'
        ']}'
    )
    result = extract_checkpoint_source(source)
    output = io.BytesIO()
    export_checkpoint_excel(result, output)
    values = " ".join(str(cell.value) for sheet in load_workbook(io.BytesIO(output.getvalue()), read_only=True).worksheets for row in sheet.iter_rows() for cell in row)
    evidence = str(result.config.model_dump()) + str(result.source_inventory) + str(result.derived) + str(CheckPointSourceReporter().build_preview(result)) + values
    assert secret not in evidence
    assert result.config.users[0].raw_extra["future"] == "keep"
