from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview


def test_remote_access_view_keeps_context_and_source_layers_separate():
    result = extract_cisco_asa_source(
        "changeto context blue\n"
        "interface GigabitEthernet0/1\n nameif outside\n"
        "webvpn\n enable outside\n"
        "vpn-addr-assign local reuse-delay 10\n"
        "ip local pool POOL 10.0.0.1-10.0.0.20 mask 255.255.255.0\n"
        "aaa-server AUTH protocol radius\n"
        "group-policy PARENT internal\n"
        "group-policy PARENT attributes\n dns-server value 8.8.8.8\n"
        "group-policy GP internal from PARENT\n"
        "group-policy GP attributes\n vpn-tunnel-protocol ssl-client\n vpn-filter value FILTER\n"
        "tunnel-group RA type remote-access\n"
        "tunnel-group RA general-attributes\n default-group-policy GP\n address-pool POOL\n authentication-server-group AUTH\n"
        "tunnel-group RA webvpn-attributes\n authentication saml\n"
        "changeto context green\n"
        "webvpn\n enable vpn-out\n"
        "vpn-addr-assign dhcp\n"
        "tunnel-group RA type remote-access\n"
    )

    blue, green = result.derived.vpn.remote_access
    assert blue.source_context == "blue"
    assert blue.group_policy.name == "GP"
    assert blue.inherited_group_policy.name == "PARENT"
    assert blue.dns_servers == ()
    assert blue.authentication_server_group.name == "AUTH"
    assert blue.address_assignment_methods == ("local",)
    assert [getattr(pool, "name", pool) for pool in blue.address_pools] == ["POOL"]
    assert blue.vpn_protocols == ("ssl-client",)
    assert blue.vpn_filter_acl == "FILTER"
    assert "Unresolved acl reference" in blue.issues
    assert blue.resolution_status == "PARTIAL"

    assert green.source_context == "green"
    assert green.group_policy is None
    assert green.address_assignment_methods == ("dhcp",)
    assert green.dhcp_servers == ()
    assert "DHCP address-assignment source is unresolved" in green.issues
    assert green.enabled_interfaces == ("vpn-out",)
    assert result.derived.vpn.ipsec_topologies == ()

    preview = build_asa_preview(result)["derived"]["vpn"]
    assert len(preview["ipsec"]) == 0
    assert [row["source_context"] for row in preview["remote_access"]] == ["blue", "green"]
    blue_preview = preview["remote_access"][0]
    assert blue_preview["webvpn_attributes"]["tunnel_group"] == ("authentication",)
    assert "saml" not in str(blue_preview["webvpn_attributes"])

    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["Remote Access VPN"].max_row == 3
