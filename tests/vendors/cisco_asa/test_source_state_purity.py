from pathlib import Path
from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.model import CiscoInterface, CiscoNATRule, CiscoStaticRoute
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview


def test_missing_hostname_remains_unconfigured():
    config = extract_cisco_asa_source("interface Ethernet0/0\n").config
    assert config.hostname is None
    assert config.system_settings.hostname is None
    assert "hostname" not in config.explicit_fields


def test_explicit_hostname_is_tracked():
    config = extract_cisco_asa_source("hostname edge\n").config
    assert config.hostname == config.system_settings.hostname == "edge"
    assert "hostname" in config.explicit_fields
    assert "hostname" in config.system_settings.explicit_fields


def test_interface_without_shutdown_has_no_effective_state():
    interface = extract_cisco_asa_source("interface Ethernet0/0\n").config.interfaces[0]
    assert interface.shutdown is None
    assert interface.administrative_state is None
    assert "shutdown" not in interface.explicit_fields


def test_shutdown_and_no_shutdown_remain_distinct_explicit_source():
    config = extract_cisco_asa_source(
        "interface Ethernet0/0\n shutdown\ninterface Ethernet0/1\n no shutdown\n"
    ).config
    down, up = config.interfaces
    assert (down.shutdown, down.administrative_state) == (True, "down")
    assert (up.shutdown, up.administrative_state) == (False, "up")
    assert {"shutdown", "administrative_state"} <= set(down.explicit_fields)
    assert {"shutdown", "administrative_state"} <= set(up.explicit_fields)


def test_explicit_no_command_is_distinct_from_missing_field():
    config = extract_cisco_asa_source(
        "interface Ethernet0/0\n no nameif\ninterface Ethernet0/1\n"
    ).config
    explicit_no, missing = config.interfaces
    assert explicit_no.nameif is missing.nameif is None
    assert "nameif" in explicit_no.explicit_fields
    assert "nameif" not in missing.explicit_fields


def test_route_without_distance_does_not_store_default():
    route = extract_cisco_asa_source("route outside 192.0.2.0 255.255.255.0 192.0.2.1\n").config.static_routes[0]
    assert route.administrative_distance is None
    assert not hasattr(route, "effective_administrative_distance")


def test_explicit_route_distance_is_preserved():
    route = extract_cisco_asa_source("route outside 192.0.2.0 255.255.255.0 192.0.2.1 7\n").config.static_routes[0]
    assert route.administrative_distance == 7
    assert "administrative_distance" in route.explicit_fields


def test_nat_model_contains_only_source_order():
    fixture = Path("tests/fixtures/cisco_asa/nat_pipeline_conformance.txt").read_text()
    rules = extract_cisco_asa_source(fixture).config.nat_rules
    assert [rule.source_order for rule in rules] == sorted(rule.source_order for rule in rules)
    assert all(rule.section in {"manual", "object", "after-auto"} for rule in rules)
    assert not {"effective_source_order", "object_nat_precedence", "object_nat_specificity", "effective_order_inputs"} & set(CiscoNATRule.model_fields)


def test_unknown_source_evidence_is_preserved_in_raw_extra():
    config = extract_cisco_asa_source(
        "interface Ethernet0/0\n vendor-knob keep-me\n"
        "route outside 192.0.2.0 255.255.255.0 192.0.2.1 unknown-route-option\n"
        "nat (inside,outside) source static 10.0.0.1 192.0.2.1 vendor-nat-option\n"
    ).config
    assert "vendor-knob keep-me" in config.interfaces[0].raw_extra["unmodeled_lines"]
    assert "unknown-route-option" in config.static_routes[0].raw_extra["unparsed_options"]
    assert "vendor-nat-option" in config.nat_rules[0].raw_extra["unparsed_tokens"]


def test_removed_effective_fields_are_not_in_source_schema():
    assert not {"administrative_state_explicit", "administrative_state_effective"} & set(CiscoInterface.model_fields)
    assert not {"effective_administrative_distance"} & set(CiscoStaticRoute.model_fields)
    assert not {"effective_source_order", "object_nat_precedence", "object_nat_specificity", "effective_order_inputs"} & set(CiscoNATRule.model_fields)


def test_preview_does_not_manufacture_missing_hostname():
    result = extract_cisco_asa_source("interface Ethernet0/0\n")
    assert result.config.hostname is None
    assert build_asa_preview(result)["hostname"] is None


def test_excel_keeps_missing_hostname_and_labels_nat_source_order():
    result = extract_cisco_asa_source(
        "nat (inside,outside) source static 10.0.0.1 192.0.2.1\n"
    )
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    nat_rows = list(workbook["NAT Rules"].iter_rows(values_only=True))
    assert summary["Hostname"] is None
    assert nat_rows[0][1] == "Source Order"
    assert nat_rows[1][1] == result.config.nat_rules[0].source_order
