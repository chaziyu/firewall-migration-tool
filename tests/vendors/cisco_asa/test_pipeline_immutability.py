from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa.derived import build_asa_derived_views
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.source_report import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa import CiscoASASourceReporter
from fwmigrate.vendors.cisco_asa.validation import validate_asa_config

from .helpers import assert_secret_absent, assert_source_unchanged, snapshot_source, snapshot_value


def test_every_report_stage_leaves_authoritative_source_and_derived_state_read_only():
    result = extract_cisco_asa_source(
        "interface Ethernet0/0\n nameif outside\n"
        "object network WEB\n host 10.0.0.10\n"
        "object-group network SERVERS\n network-object object WEB\n"
        "access-list OUT extended permit ip object-group SERVERS any\n"
        "access-group OUT in interface outside\n"
        "nat (inside,outside) source static WEB interface\n"
        "route outside 192.0.2.0 255.255.255.0 192.0.2.1\n"
        "interface Ethernet0/1\n vendor-option retained\n"
    )
    source = snapshot_source(result.config)

    derived = build_asa_derived_views(result.config)
    assert_source_unchanged(result.config, source)
    derived_snapshot = snapshot_value(derived)
    validation = validate_asa_config(result.config, derived)
    assert_source_unchanged(result.config, source)

    from dataclasses import replace
    analysis = replace(result, derived=derived, validation=validation)

    output = BytesIO()
    export_asa_excel(analysis, output)
    assert output.getvalue()
    assert_source_unchanged(result.config, source)
    assert snapshot_value(derived) == derived_snapshot


def test_public_reporter_preview_and_export_leave_source_and_derived_state_unchanged():
    reporter = CiscoASASourceReporter()
    fixture = Path(__file__).parent / "fixtures" / "architecture_regression.cfg"
    analysis = reporter.analyze_source(fixture.read_text(encoding="utf-8"))
    assert {context.name for context in analysis.config.contexts} == {"customer-a", "customer-b"}
    source = snapshot_source(analysis.config)
    derived = snapshot_value(analysis.derived)

    preview = reporter.build_preview(analysis)
    assert preview["vendor"] == "cisco_asa"
    assert not hasattr(analysis, "canonical_ir")
    assert_secret_absent(preview, "LOCAL_PASSWORD_SENTINEL", "ENABLE_SECRET_SENTINEL")
    assert preview["derived"]["nat"]["rules"][0]["section"] == "object"
    assert preview["derived"]["nat"]["rules"][0]["effective_order"] == 2
    assert preview["derived"]["routes"][0]["effective_administrative_distance"] == 1
    assert_source_unchanged(analysis.config, source)
    assert snapshot_value(analysis.derived) == derived

    output = BytesIO()
    reporter.export_excel(analysis, output)
    assert output.getvalue()
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["NAT Rules"]["C2"].value == 2
    assert workbook["NAT Rules"]["E2"].value == "object"
    assert workbook["Routes"]["F2"].value is None
    assert workbook["Routes"]["G2"].value == 1
    assert_secret_absent([list(sheet.iter_rows(values_only=True)) for sheet in workbook],
                         "LOCAL_PASSWORD_SENTINEL", "ENABLE_SECRET_SENTINEL")
    assert_source_unchanged(analysis.config, source)
    assert snapshot_value(analysis.derived) == derived
