from fwmigrate.vendors.cisco_asa.relationships.references import ASAReferenceKind, ASAReferenceStatus, build_asa_reference_index

from fwmigrate.vendors.cisco_asa.source_report import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.parser import CiscoASAParser
import pytest
from io import BytesIO
from openpyxl import load_workbook
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview


@pytest.mark.parametrize("profile", ["fast", "full"])
@pytest.mark.parametrize("root", ["", "hostname ASA_SYSTEM\ndomain-name system.example\n"])
def test_management_settings_keep_context_ownership_in_reports(profile, root):
    result = extract_cisco_asa_source(root +
        "changeto context blue\nhostname ASA_BLUE\ndomain-name blue.example\n"
        "dns-group BLUE_DNS\ndns domain-lookup inside\nhttp server enable\n"
        "changeto context green\nhostname ASA_GREEN\ndomain-name green.example\n"
        "dns-group GREEN_DNS\nno http server enable\n"
        "changeto context blue\nhttp server idle-timeout 10\nno http server idle-timeout\n")
    config = result.config
    before = config.model_dump(mode="json")
    assert config.hostname == ("ASA_SYSTEM" if root else None)
    blue, green = config.contexts
    assert blue.system_settings.hostname == "ASA_BLUE"
    assert green.system_settings.hostname == "ASA_GREEN"
    assert blue.dns_settings.domain_name == "blue.example"
    assert green.dns_settings.domain_name == "green.example"
    assert blue.dns_settings.lookup_interfaces == ["inside"]
    assert blue.http_server.enabled is True
    assert green.http_server.enabled is False
    assert blue.http_server.idle_timeout is None
    preview = build_asa_preview(result)["source"]["other_source"]
    assert [(r["context"], r["hostname"]) for r in preview["system_settings"]] == (
        ([(None, "ASA_SYSTEM")] if root else []) + [("blue", "ASA_BLUE"), ("green", "ASA_GREEN")])
    assert [(r["context"], r["enabled"]) for r in preview["http_server"]] == [("blue", True), ("green", False)]
    output = BytesIO()
    export_asa_excel(result, output, profile=profile)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    for sheet, field, expected in [
        ("System Settings", "Hostname", ["ASA_BLUE", "ASA_GREEN"]),
        ("DNS Settings", "Domain Name", ["blue.example", "green.example"]),
        ("HTTP Server", "Enabled", [True, False]),
    ]:
        rows = list(workbook[sheet].values)
        records = [dict(zip(rows[0], row)) for row in rows[1:]]
        scoped = [row for row in records if row["Context"] is not None]
        assert [row["Context"] for row in scoped] == ["blue", "green"]
        assert [row[field] for row in scoped] == expected
    workbook.close()
    assert config.model_dump(mode="json") == before


def parse(text):
    return CiscoASAParser(text).parse_raw()

from copy import deepcopy

from fwmigrate.vendors.cisco_asa.derived import build_asa_derived_views



from fwmigrate.vendors.cisco_asa.model import CiscoASAConfig, CiscoInterface, CiscoNetworkObject


def test_duplicate_names_and_nat_references_stay_inside_their_context():
    result = extract_cisco_asa_source(
        "changeto context customer-a\nobject network WEB\n host 10.0.0.1\n"
        "object network ONLY_A\n host 10.0.0.2\n"
        "nat (inside,outside) source static WEB interface\n"
        "changeto context customer-b\nobject network WEB\n host 192.0.2.1\n"
        "object network ONLY_B\n host 192.0.2.2\n"
        "nat (inside,outside) source static WEB interface\n"
        "changeto context customer-a\nnat (inside,outside) source static ONLY_B interface\n"
    )
    refs = build_asa_reference_index(result.config)

    a = refs.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "WEB")
    b = refs.resolve("customer-b", ASAReferenceKind.NETWORK_OBJECT, "WEB")
    assert a.target.value == "10.0.0.1" and b.target.value == "192.0.2.1"
    assert refs.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "ONLY_B").status is ASAReferenceStatus.UNRESOLVED
    contexts = [row.source_context for row in result.derived.nat_relationships.rules]
    assert contexts.count("customer-a") == 2
    assert contexts.count("customer-b") == 1
    assert [(row.source_context, row.effective_order) for row in result.derived.nat.rules] == [
        ("customer-a", 1), ("customer-b", 1), ("customer-a", 2)]
    assert any(issue.source_context == "customer-a" and issue.reference_name == "ONLY_B"
               for issue in result.derived.relationship_issues)

def test_dhcp_and_context_interface_resolution_stays_in_derived_relationships():
    config = parse("\n".join([
        "interface GigabitEthernet0/0", " nameif outside", " dhcprelay server 192.0.2.1",
        "context blue", " allocate-interface GigabitEthernet0/0 mapped",
    ]))
    before = deepcopy(config)

    assert config.dhcp_relays[0].server_entries[0].interface == "GigabitEthernet0/0"
    assert config.dhcp_relays[0].server_entries[0].explicit_fields == {"server", "interface"}
    assert config.contexts[0].allocated_interface_entries[0].physical_interface == "GigabitEthernet0/0"
    assert config.contexts[0].allocated_interface_entries[0].explicit_fields == {"physical_interface", "mapped_name"}

    derived = build_asa_derived_views(config)
    routing = derived.routing_relationships
    assert routing.dhcp_relay_interfaces[0].interface is config.interfaces[0]
    assert routing.allocated_interfaces[0].interface is config.interfaces[0]
    assert config == before

def test_named_references_are_context_local_and_duplicates_are_ambiguous():
    config = CiscoASAConfig(network_objects=[
        CiscoNetworkObject(name="WEB", source_context="customer-a"),
        CiscoNetworkObject(name="WEB", source_context="customer-b"),
        CiscoNetworkObject(name="ONLY-IN-B", source_context="customer-b"),
        CiscoNetworkObject(name="DUP", source_context="customer-a"),
        CiscoNetworkObject(name="DUP", source_context="customer-a"),
        CiscoNetworkObject(name="SYSTEM", source_context=None),
    ])
    references = build_asa_reference_index(config)

    assert references.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "WEB").status is ASAReferenceStatus.RESOLVED
    assert references.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "ONLY-IN-B").status is ASAReferenceStatus.UNRESOLVED
    assert references.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "SYSTEM").status is ASAReferenceStatus.UNRESOLVED
    assert references.resolve("customer-a", ASAReferenceKind.NETWORK_OBJECT, "DUP").status is ASAReferenceStatus.AMBIGUOUS

def test_interface_aliases_and_topology_are_context_local():
    config = CiscoASAConfig(interfaces=[
        CiscoInterface(name="GigabitEthernet0/1", source_context="customer-a", nameif="inside"),
        CiscoInterface(name="GigabitEthernet0/1", source_context="customer-b", nameif="inside"),
        CiscoInterface(name="GigabitEthernet0/3", source_context="customer-a", nameif="inside"),
        CiscoInterface(name="GigabitEthernet0/2", source_context="customer-a", parent_interface="GigabitEthernet0/1"),
    ])
    references = build_asa_reference_index(config)

    assert references.resolve("customer-a", ASAReferenceKind.INTERFACE, "inside").status is ASAReferenceStatus.AMBIGUOUS
    assert references.resolve("customer-b", ASAReferenceKind.INTERFACE, "inside").target.source_context == "customer-b"
    assert references.resolve(None, ASAReferenceKind.INTERFACE, "inside").status is ASAReferenceStatus.UNRESOLVED
    topology = build_asa_derived_views(config).interface_topology
    child = next(item for item in topology.interfaces if item.name.endswith("0/2"))
    assert child.parent.source_context == "customer-a"
    assert not topology.issues

def test_source_inventory_and_diagnostics_preserve_execution_context():
    result = extract_cisco_asa_source(
        "changeto context blue\n"
        "object network WEB\n"
        " host 10.0.0.1\n"
        "ssh not-an-address not-a-mask outside\n"
        "changeto context green\n"
        "future-command retained\n"
    )

    blue_inventory = [
        item for item in result.inventory_items
        if item.source_context == "blue"
    ]
    green_inventory = [
        item for item in result.inventory_items
        if item.source_context == "green"
    ]
    assert any(item.source_path == "object network" for item in blue_inventory)
    assert any(item.source_path == "other" for item in green_inventory)
    assert any(item.source_context == "blue" for item in result.config.diagnostics)
    assert any(
        issue.source_context == "blue" and issue.category == "parse"
        for issue in result.validation.issues
    )
    assert any(item.source_context == "green" for item in result.unsupported_items)

    from io import BytesIO
    from openpyxl import load_workbook
    from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel

    output = BytesIO()
    export_asa_excel(result, output)
    sheet = load_workbook(BytesIO(output.getvalue()), read_only=True)["Source Inventory"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    context_index = headers.index("Context")
    path_index = headers.index("Source Path")
    assert any(
        row[context_index] == "blue" and row[path_index] == "object network"
        for row in rows[1:]
    )
    assert any(
        row[context_index] == "green" and row[path_index] == "other"
        for row in rows[1:]
    )

