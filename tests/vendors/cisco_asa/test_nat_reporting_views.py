from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa import extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview


def test_nat_reporting_views_are_traceable_and_do_not_promote_manual_source_nat_to_vip():
    result = extract_cisco_asa_source(
        "object network WEB\n host 10.0.0.2\n nat (inside,outside) static 192.0.2.2\n"
        "object network CLIENTS\n subnet 10.1.0.0 255.255.255.0\n nat (inside,outside) dynamic interface\n"
        "nat (inside,outside) source static 10.2.0.1 192.0.2.10\n"
        "nat (inside,outside) source static 10.2.0.10 192.0.2.20 destination static 203.0.113.50 10.2.0.10 service tcp 443 8443\n"
        "nat (inside,outside) source dynamic any pat-pool PATPOOL\n"
    )

    pools = {row.source_rule.name: row for row in result.derived.nat.source_nat_pools}
    vips = result.derived.nat.vips

    assert len(pools) == 5
    interface_pat = next(row for row in pools.values() if row.pool_type == "interface_pat")
    assert getattr(interface_pat.mapped_interface, "name", interface_pat.mapped_interface) == "outside"
    assert any(row.pool_type == "dynamic_pat_pool" for row in pools.values())
    assert len(vips) == 2
    object_vip = next(row for row in vips if row.source_rule.owning_object == "WEB")
    twice_nat_vip = next(row for row in vips if row.source_rule.destination_mode == "static")
    assert (object_vip.real_address, object_vip.mapped_address) == ("10.0.0.2", "192.0.2.2")
    assert (twice_nat_vip.real_address, twice_nat_vip.mapped_address) == ("10.2.0.10", "203.0.113.50")
    assert (twice_nat_vip.real_service, twice_nat_vip.mapped_service, twice_nat_vip.protocol) == ("443", "8443", "tcp")
    assert all(row.source_rule.syntax_family == "object" or row.source_rule.destination_mode == "static" for row in vips)

    nat_preview = build_asa_preview(result)["derived"]["nat"]
    assert len(nat_preview["source_nat_pools"]) == 5
    assert len(nat_preview["vips"]) == 2
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert "Source NAT Pools" in workbook.sheetnames
    assert "Published Services - VIPs" in workbook.sheetnames
    assert workbook["Published Services - VIPs"].max_row == 3


def test_identity_nat_does_not_create_comparison_rows():
    result = extract_cisco_asa_source(
        "nat (inside,outside) source static 10.0.0.1 10.0.0.1\n"
        "nat (inside) 0 access-list NO_NAT\n"
    )
    assert not result.derived.nat.source_nat_pools
    assert not result.derived.nat.vips


def test_source_nat_pools_keep_duplicate_names_separate_by_context():
    result = extract_cisco_asa_source(
        "changeto context customer-a\nobject network WEB\n host 10.0.0.1\n"
        "nat (inside,outside) source static WEB interface\n"
        "changeto context customer-b\nobject network WEB\n host 192.0.2.1\n"
        "nat (inside,outside) source static WEB interface\n"
    )
    assert [row.source_context for row in result.derived.nat.source_nat_pools] == ["customer-a", "customer-b"]
    assert [row.mapped_interface for row in result.derived.nat.source_nat_pools] == ["outside", "outside"]


def test_unresolved_nat_objects_stay_traceable_and_report_issues():
    result = extract_cisco_asa_source("nat (inside,outside) source static MISSING 192.0.2.1\n")
    row = result.derived.nat.source_nat_pools[0]
    assert row.source_rule.real_source == "MISSING"
    assert any("Unresolved" in issue for issue in row.issues)
    assert any(issue.category == "nat" and "Unresolved" in issue.message
               for issue in result.derived.transform_issues)
