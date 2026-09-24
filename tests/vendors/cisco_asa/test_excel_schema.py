from io import BytesIO
import ast

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel
from fwmigrate.vendors.cisco_asa.export.excel_schema import SHEET_HEADERS, SHEET_ORDER
from fwmigrate.vendors.cisco_asa.source_report import extract_cisco_asa_source


def test_selected_asa_domains_have_stable_report_sheets_and_headers():
    expected = {
        "Review Required", "Source Inventory", "Interfaces", "Zones", "Network Objects",
        "Network Groups", "Service Objects", "Service Groups", "Time Ranges", "ACL Rules",
        "ACL Bindings", "NAT Rules", "Source NAT Pools", "Published Services - VIPs",
        "Class Maps", "Policy Maps", "Service Policies", "DHCP Servers", "DHCP Reservations",
        "DHCP Relays", "Routes", "Route Maps", "Policy Routing", "SLA Monitors", "Tracks",
        "Local Users", "AAA Server Groups", "AAA Server Hosts", "Command Privileges",
        "IKE Policies", "IKEv2 Proposals", "IPsec Transform Sets", "Crypto Maps", "Tunnel Groups",
        "Group Policies", "VPN Address Pools", "IPsec VPN", "Remote Access VPN", "DNS", "NTP",
        "Management Access", "SNMP", "Logging", "Failover", "Contexts", "Validation",
        "Unsupported", "Extraction Coverage",
    }
    assert expected <= set(SHEET_ORDER)

    result = extract_cisco_asa_source("hostname asa\ninterface Ethernet0/0\n nameif outside\n")
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    for name, headers in SHEET_HEADERS.items():
        assert tuple(workbook[name].iter_rows(min_row=1, max_row=1, values_only=True))[0] == headers


def test_nat_workbook_keeps_source_and_derived_order_columns_distinct():
    result = extract_cisco_asa_source("nat (inside,outside) source static any any\n")
    output = BytesIO()
    export_asa_excel(result, output)
    sheet = load_workbook(BytesIO(output.getvalue()), read_only=True)["NAT Rules"]
    headers = next(sheet.iter_rows(values_only=True))
    assert headers[1:5] == ("Source Order", "Effective Order", "Order Status", "Section")


def test_presentation_modules_do_not_import_parsers_or_calculate_semantics():
    from pathlib import Path

    root = Path(__file__).parents[3] / "src/fwmigrate/vendors/cisco_asa"
    for path in (root / "export/excel.py", root / "web_report.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        imports = [node.module or "" for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)]
        assert not any("parser" in module.casefold() for module in imports)
        calls = [node.func.id if isinstance(node.func, ast.Name) else node.func.attr
                 for node in ast.walk(tree) if isinstance(node, ast.Call)]
        assert not {"parse_raw", "resolve", "transform_nat", "transform_routes"} & set(calls)
