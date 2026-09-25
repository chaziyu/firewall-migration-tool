from io import BytesIO

import ast

from openpyxl import load_workbook

from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel

from fwmigrate.vendors.cisco_asa.export.excel_schema import SHEET_HEADERS, SHEET_ORDER

from fwmigrate.vendors.cisco_asa.source_report import extract_cisco_asa_source

from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview

from .helpers import assert_source_unchanged, snapshot_source

import io

from fwmigrate.source_reporting import source_reporters

from fwmigrate.web import create_app

from fwmigrate.vendors.cisco_asa import CiscoASASourceReporter, ASASourceResult

from pathlib import Path

from fwmigrate.vendors.cisco_asa import extract_cisco_asa_source

from fwmigrate.vendors.cisco_asa.model import CiscoInterface, CiscoNATRule, CiscoStaticRoute


def test_selected_asa_domains_have_stable_report_sheets_and_headers():
    expected = {
        "Review Required", "Source Inventory", "Interfaces", "Zones", "Network Objects",
        "Network Groups", "Service Objects", "Service Groups", "Time Ranges", "ACL Rules",
        "ACL Remarks", "DHCP Global Settings", "IPsec Profiles",
        "ACL Bindings", "NAT Rules", "Source NAT Pools", "Published Services - VIPs",
        "Class Maps", "Policy Maps", "Service Policies", "DHCP Servers", "DHCP Reservations",
        "DHCP Relays", "Routes", "Route Maps", "Policy Routing", "SLA Monitors", "Tracks",
        "Local Users", "AAA Server Groups", "AAA Server Hosts", "Command Privileges",
        "IKE Policies", "IKEv2 Proposals", "IPsec Transform Sets", "Crypto Maps", "Tunnel Groups",
        "Group Policies", "VPN Address Pools", "IPsec VPN", "Remote Access VPN", "DNS", "NTP",
        "Management Access", "SNMP", "Logging", "Failover", "Contexts", "Validation",
        "Unsupported", "Extraction Coverage",
    }
    assert expected <= set(SHEET_ORDER)

    result = extract_cisco_asa_source("hostname asa\ninterface Ethernet0/0\n nameif outside\n")
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    for name, headers in SHEET_HEADERS.items():
        assert tuple(workbook[name].iter_rows(min_row=1, max_row=1, values_only=True))[0] == headers

def test_nat_workbook_keeps_source_and_derived_order_columns_distinct():
    result = extract_cisco_asa_source("nat (inside,outside) source static any any\n")
    output = BytesIO()
    export_asa_excel(result, output)
    sheet = load_workbook(BytesIO(output.getvalue()), read_only=True)["NAT Rules"]
    headers = next(sheet.iter_rows(values_only=True))
    assert headers[1:5] == ("Source Order", "Effective Order", "Order Status", "Section")

def test_every_workbook_row_matches_its_header_and_acl_binding_cells_keep_their_meaning():
    result = extract_cisco_asa_source(
        "interface Ethernet0/0\n nameif inside\n security-level 100\n"
        "access-list ACL line 10 remark before\n"
        "access-list ACL line 20 extended permit ip any any\n"
        "access-list ACL line 30 remark trailing\n"
        "access-group ACL in interface inside\n"
        "object network WEB\n host 10.0.0.2\n nat (inside,outside) static 192.0.2.2\n"
        "dhcpd dns 8.8.8.8\n dhcpd address 192.0.2.10-192.0.2.20 inside\n dhcpd enable inside\n"
        "route null0 198.51.100.0 255.255.255.0\n"
        "crypto ipsec profile VTI\n set ikev2 ipsec-proposal P2\n"
        "interface Tunnel1\n tunnel protection ipsec profile VTI\n"
    )
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)

    for name, headers in SHEET_HEADERS.items():
        for row in workbook[name].iter_rows(min_row=2, values_only=True):
            assert len(row) == len(headers), (name, row)
    binding = next(workbook["ACL Bindings"].iter_rows(min_row=2, values_only=True))
    columns = dict(zip(SHEET_HEADERS["ACL Bindings"], binding))
    assert columns["Resolved Interface"] == "Ethernet0/0"
    assert columns["Issues"] == "()"

def test_derived_workbook_rows_match_semantic_columns():
    result = extract_cisco_asa_source(
        "interface Ethernet0/0\n nameif inside\n security-level 100\n"
        "object network REAL\n host 10.0.0.1\n"
        "object network MAPPED\n host 192.0.2.1\n"
        "nat (inside,outside) source static REAL MAPPED\n"
        "access-list SELECTOR extended permit ip any any\n access-group SELECTOR in interface inside\n"
        "crypto ipsec ikev1 transform-set TS esp-aes esp-sha-hmac\n"
        "crypto ipsec profile VTI\n set ikev1 transform-set TS\n"
        "interface Tunnel1\n tunnel source outside\n tunnel destination 203.0.113.5\n"
        " tunnel protection ipsec profile VTI\n tunnel protection ipsec policy SELECTOR\n"
        "ip local pool CLIENTS 10.10.0.1-10.10.0.10\n"
        "group-policy GP attributes\n vpn-tunnel-protocol ssl-client\n"
        "tunnel-group RA type remote-access\n tunnel-group RA general-attributes\n default-group-policy GP\n address-pool CLIENTS\n"
        "dhcpd reserve-address 192.0.2.15 0011.2233.4455 inside\n"
    )
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    pool = dict(zip(SHEET_HEADERS["Source NAT Pools"], next(workbook["Source NAT Pools"].iter_rows(min_row=2, values_only=True))))
    assert pool["Mapped Object"] == "MAPPED" and pool["Mapped Interface"] is None
    ipsec = dict(zip(SHEET_HEADERS["IPsec VPN"], next(workbook["IPsec VPN"].iter_rows(min_row=2, values_only=True))))
    assert ipsec["Tunnel Source"] == "outside"
    assert ipsec["VTI Policy ACL"] == "SELECTOR"
    access = dict(zip(SHEET_HEADERS["Remote Access VPN"], next(workbook["Remote Access VPN"].iter_rows(min_row=2, values_only=True))))
    assert access["Connection Profile Type"] == "remote-access"
    assert access["Group Policy"] == "GP" and access["Address Pools"] == "('CLIENTS',)"
    reservation = next(workbook["DHCP Reservations"].iter_rows(min_row=2, values_only=True))
    assert reservation == ("192.0.2.15", "0011.2233.4455", "inside", None, 26)
    binding = dict(zip(SHEET_HEADERS["ACL Bindings"], next(workbook["ACL Bindings"].iter_rows(min_row=2, values_only=True))))
    assert binding["ACL"] == "SELECTOR" and binding["Resolved Interface"] == "Ethernet0/0"


def test_excel_formula_like_source_text_is_literal():
    result = extract_cisco_asa_source("hostname =1+1\n")
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True, data_only=False)

    cell = workbook["Summary"]["B3"]
    assert cell.value == "'=1+1"
    assert cell.data_type != "f"

def test_empty_management_singletons_are_not_reported_as_configured_source():
    empty = extract_cisco_asa_source("")
    output = BytesIO()
    export_asa_excel(empty, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert all(workbook[name].max_row == 1 for name in ("DNS Settings", "System Settings", "HTTP Server", "Failover"))
    management = build_asa_preview(empty)["source"]["management"]
    assert management["dns"] == []
    assert management["system"] is None
    assert management["http_server"] is None
    assert build_asa_preview(empty)["source"]["failover"]["config"] is None
    assert build_asa_preview(empty)["source"]["other_source"]["multi_context_system"] is None
    relationship_keys = build_asa_preview(empty)["relationships"]
    assert {"acl_bindings", "vpn", "vpn_relationships"} <= set(relationship_keys)

def test_explicit_management_singletons_remain_visible():
    result = extract_cisco_asa_source("hostname asa\nhttp server enable\n")
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["System Settings"].max_row == 2
    assert workbook["HTTP Server"].max_row == 2

def test_excel_preserves_missing_and_explicit_no_source_values():
    result = extract_cisco_asa_source(
        "interface Ethernet0/0\n ip address 192.0.2.1 255.255.255.0\n"
        "interface Ethernet0/1\n no shutdown\n no nameif\n no security-level\n no ip address\n"
        "route outside 192.0.2.0 255.255.255.0 192.0.2.1\n"
        "route outside 198.51.100.0 255.255.255.0 192.0.2.1 10\n"
    )
    before = snapshot_source(result.config)

    output = BytesIO()
    export_asa_excel(result, output)

    assert_source_unchanged(result.config, before)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["Routes"]["A2"].value == "outside"
    rows = list(workbook["Interfaces"].iter_rows(values_only=True))
    headers = rows[0]
    missing, explicit_no = (dict(zip(headers, row)) for row in rows[1:3])
    assert missing["Shutdown"] is None
    assert explicit_no["Shutdown"] is False
    assert explicit_no["Nameif"] is None
    assert explicit_no["Security Level"] is None
    assert explicit_no["Ip"] is None


def test_nat_presentation_exposes_derived_rows():
    result = extract_cisco_asa_source(
        "object network WEB\n host 10.0.0.2\n nat (inside,outside) static 192.0.2.2\n"
    )

    preview = build_asa_preview(result)["derived"]["nat"]
    assert len(preview["vips"]) == 1
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert "Source NAT Pools" in workbook.sheetnames
    assert "Published Services - VIPs" in workbook.sheetnames
    assert workbook["Published Services - VIPs"].max_row == 2


def test_remote_access_presentation_exposes_preview_and_excel_rows():
    result = extract_cisco_asa_source(
        "tunnel-group RA type remote-access\n"
        "tunnel-group RA general-attributes\n"
    )

    preview = build_asa_preview(result)["derived"]["vpn"]
    assert len(preview["remote_access"]) == 1
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["Remote Access VPN"].max_row == 2


def test_nat_presentation_exposes_derived_rows():
    result = extract_cisco_asa_source(
        "object network WEB\n host 10.0.0.2\n nat (inside,outside) static 192.0.2.2\n"
    )

    preview = build_asa_preview(result)["derived"]["nat"]
    assert len(preview["vips"]) == 1
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert "Source NAT Pools" in workbook.sheetnames
    assert "Published Services - VIPs" in workbook.sheetnames
    assert workbook["Published Services - VIPs"].max_row == 2


def test_remote_access_presentation_exposes_preview_and_excel_rows():
    result = extract_cisco_asa_source(
        "tunnel-group RA type remote-access\n"
        "tunnel-group RA general-attributes\n"
    )

    preview = build_asa_preview(result)["derived"]["vpn"]
    assert len(preview["remote_access"]) == 1
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert workbook["Remote Access VPN"].max_row == 2

SOURCE = """hostname asa
interface GigabitEthernet0/1
 nameif outside
 ip address 203.0.113.1 255.255.255.0
object network WEB
 host 10.0.0.10
access-list OUT extended permit tcp object network WEB any eq 443
access-group OUT in interface outside
nat (inside,outside) source static WEB interface
"""

def test_asa_reporter_is_registered_with_an_opaque_native_result():
    reporter = source_reporters.get("CISCO_ASA")

    assert isinstance(reporter, CiscoASASourceReporter)
    analysis = reporter.analyze_source(SOURCE)
    assert isinstance(analysis, ASASourceResult)
    assert analysis.config.hostname == "asa"
    assert [rule.source_order for rule in analysis.config.access_rules] == [7]
    assert not hasattr(analysis, "canonical_ir")

def test_asa_shared_web_flow_uses_vendor_preview_and_excel():
    client = create_app({"TESTING": True}).test_client()
    preview = client.post(
        "/api/preview",
        data={"source_vendor": "cisco_asa", "file": (io.BytesIO(SOURCE.encode()), "asa.cfg")},
        content_type="multipart/form-data",
    )

    assert preview.status_code == 200
    preview_id = preview.get_json()["preview_id"]

    workbook = client.post(
        "/api/extract/excel",
        data={"source_vendor": "cisco_asa", "preview_id": preview_id},
    )

    assert workbook.status_code == 200
    assert workbook.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def test_preview_does_not_manufacture_missing_hostname():
    result = extract_cisco_asa_source("interface Ethernet0/0\n")
    assert result.config.hostname is None
    assert build_asa_preview(result)["hostname"] is None

def test_excel_keeps_missing_hostname_and_labels_nat_source_order():
    result = extract_cisco_asa_source(
        "nat (inside,outside) source static 10.0.0.1 192.0.2.1\n"
    )
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    nat_rows = list(workbook["NAT Rules"].iter_rows(values_only=True))
    assert summary["Hostname"] is None
    assert nat_rows[0][1] == "Source Order"
    assert nat_rows[1][1] == result.config.nat_rules[0].source_order

def test_preview_separates_source_relationship_and_derived_and_redacts_secrets():
    result = extract_cisco_asa_source(
        "hostname asa\n"
        "interface Ethernet0/0\n nameif outside\n"
        "object network WEB\n host 192.0.2.10\n"
        "access-list OUT extended permit ip any any\n"
        "access-group OUT in interface outside\n"
        "nat (inside,outside) source static WEB interface\n"
    )
    preview = build_asa_preview(result)
    assert "interfaces" in preview["source"]
    assert "addresses" in preview["source"]
    assert "acl" in preview["relationships"]
    assert "nat" in preview["derived"]
    assert "inventory" in preview and "unsupported" in preview and "coverage" in preview
    assert preview["summary"]["objects"]["policies"] == len(result.config.access_rules)
    assert preview["summary"]["scopes"] == ["root"]
    assert {"interfaces", "addresses", "policies", "nat", "routes", "vpn_tunnels", "vpn_phase2",
            "validation", "unresolved_references"} <= set(preview["sections"])
    acl_row = preview["sections"]["policies"][0]
    assert acl_row["acl_name"] == "OUT" and acl_row["interface"] == "outside"
    assert acl_row["source_order"] is not None
