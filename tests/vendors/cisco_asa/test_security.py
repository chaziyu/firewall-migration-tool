from copy import deepcopy

from io import BytesIO

from fwmigrate.vendors.cisco_asa.model import CiscoASAConfig, CiscoInterface

from fwmigrate.vendors.cisco_asa.source_report import _sanitize, extract_cisco_asa_source
from fwmigrate.vendors.cisco_asa.web_report import build_asa_preview

from .helpers import assert_secret_absent

def test_nested_sensitive_source_attribute_keys_are_redacted():
    secret = "NESTED_TOKEN_SENTINEL"
    config = CiscoASAConfig(interfaces=[CiscoInterface(
        name="Ethernet0/0",
        source_attributes={"future": {"api_token": secret, "safe_option": "keep"}},
        raw_extra={"future": {"api_token": secret, "safe_option": "keep"}},
    )])
    safe = _sanitize(deepcopy(config))

    assert_secret_absent(safe.model_dump(mode="python"), secret)
    assert safe.interfaces[0].source_attributes["future"]["safe_option"] == "keep"
    assert safe.interfaces[0].raw_extra["future"]["safe_option"] == "keep"

def test_cli_password_sentinel_does_not_reach_source_or_accounting():
    from openpyxl import load_workbook

    from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel

    sentinel = "LOCAL_PASSWORD_SENTINEL"
    result = extract_cisco_asa_source(f"username admin password 0 {sentinel}\n")
    assert_secret_absent(result.config.model_dump(mode="python"), sentinel)
    assert_secret_absent(result.inventory_items, sentinel)
    assert_secret_absent(result.unsupported_items, sentinel)
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert_secret_absent([list(sheet.iter_rows(values_only=True)) for sheet in workbook], sentinel)

def test_secret_sentinels_stay_out_of_every_report_surface():
    from io import BytesIO
    from openpyxl import load_workbook

    from fwmigrate.vendors.cisco_asa.export.excel import export_asa_excel

    sentinels = ("LOCAL_PASSWORD_SENTINEL", "ENABLE_SECRET_SENTINEL", "AAA_KEY_SENTINEL",
                 "IKE_PSK_SENTINEL", "FAILOVER_KEY_SENTINEL")
    source = (
        f"username admin password 0 {sentinels[0]}\n"
        f"enable password {sentinels[1]}\n"
        "aaa-server RAD protocol radius\naaa-server RAD host 192.0.2.5\n"
        f" key {sentinels[2]}\n"
        "tunnel-group 192.0.2.1 type ipsec-l2l\ntunnel-group 192.0.2.1 ipsec-attributes\n"
        f" ikev1 pre-shared-key {sentinels[3]}\n"
        f"failover key {sentinels[4]}\n"
    )
    result = extract_cisco_asa_source(source)
    for surface in (result.config.model_dump(mode="python"), result.inventory_items,
                    result.unsupported_items, result.source_sections, result.derived,
                    result.validation, build_asa_preview(result)):
        assert_secret_absent(surface, *sentinels)
    output = BytesIO()
    export_asa_excel(result, output)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    assert_secret_absent([list(sheet.iter_rows(values_only=True)) for sheet in workbook], *sentinels)
