from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_application_services_children_are_typed_and_exported():
    ir = JuniperSRXParser("""
    set security idp idp-policy idp1 rulebase-ips rule r action drop
    set security utm utm-policy utm1
    set services ssl proxy profile ssl1
    set security policies from-zone trust to-zone untrust policy p match source-address 10.0.0.0/8
    set security policies from-zone trust to-zone untrust policy p match destination-address 0.0.0.0/0
    set security policies from-zone trust to-zone untrust policy p match application any
    set security policies from-zone trust to-zone untrust policy p then permit application-services idp-policy idp1
    set security policies from-zone trust to-zone untrust policy p then permit application-services utm-policy utm1
    set security policies from-zone trust to-zone untrust policy p then permit application-services ssl-proxy-profile ssl1
    set security policies from-zone trust to-zone untrust policy p application-services unknown-child
    """).transform_to_ir()

    policy = next(policy for policy in ir.policies if policy.name == "p")
    assert policy.source_security_profile_references == {
        "idp-policy": "idp1",
        "utm-policy": "utm1",
        "ssl-proxy-profile": "ssl1",
    }
    assert policy.security_profile_reference_statuses == {
        "idp-policy": "RESOLVED",
        "utm-policy": "RESOLVED",
        "ssl-proxy-profile": "RESOLVED",
    }
    assert "unknown-child" in policy.source_extra_settings["junos_application_services"]
    assert policy.security_profile_semantics_review is True

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(row for row in sheet.iter_rows(min_row=4, values_only=False) if row[headers["Name"] - 1].value == "p")
    assert "idp1" in row[headers["Security Profile References"] - 1].value

