import io

from openpyxl import load_workbook

from fwmigrate.ir.enums import AddressType
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


CONFIG = '''
config firewall address
    edit "default4"
    next
    edit "fsso-multi"
        set type dynamic
        set sub-type fsso
        set fsso-group "Finance Users" "Remote Users"
    next
end
config firewall address6
    edit "default6"
    next
end
config firewall multicast-address
    edit "default-mcast4"
    next
    edit "default-broadcast"
        set type broadcastmask
    next
end
config firewall multicast-address6
    edit "default-mcast6"
    next
end
'''


def _by_name(items):
    return {item.name: item for item in items}


def test_fortios_746_address_defaults_are_section_specific():
    parsed = _by_name(parse_fortigate_config(CONFIG).addresses)

    default4 = parsed["default4"]
    assert default4.type == "ipmask"
    assert default4.subnet == "0.0.0.0 0.0.0.0"
    assert default4.source_effective_defaults == {
        "type": "ipmask",
        "subnet": "0.0.0.0 0.0.0.0",
    }

    default6 = parsed["default6"]
    assert default6.type == "ipprefix"
    assert default6.ip6 == "::/0"
    assert default6.source_effective_defaults == {
        "type": "ipprefix",
        "ip6": "::/0",
    }

    multicast4 = parsed["default-mcast4"]
    assert multicast4.type == "multicastrange"
    assert multicast4.start_ip == "0.0.0.0"
    assert multicast4.end_ip == "0.0.0.0"
    assert multicast4.source_effective_defaults == {
        "type": "multicastrange",
        "start_ip": "0.0.0.0",
        "end_ip": "0.0.0.0",
    }

    broadcast = parsed["default-broadcast"]
    assert broadcast.type == "broadcastmask"
    assert broadcast.subnet == "0.0.0.0 0.0.0.0"
    assert broadcast.source_effective_defaults == {
        "subnet": "0.0.0.0 0.0.0.0",
    }

    multicast6 = parsed["default-mcast6"]
    assert multicast6.type is None
    assert multicast6.ip6 == "::/0"
    assert multicast6.source_effective_defaults == {"ip6": "::/0"}


def test_fortios_746_effective_defaults_reach_ir_without_becoming_source_commands():
    addresses = _by_name(FGToIRTransformer(parse_fortigate_config(CONFIG)).transform().addresses)

    default4 = addresses["default4"]
    assert default4.type == AddressType.NETWORK
    assert default4.value == "0.0.0.0/0"
    assert default4.source_type == "ipmask"
    assert default4.source_attributes == {}
    assert default4.source_effective_defaults == {
        "type": "ipmask",
        "subnet": "0.0.0.0 0.0.0.0",
    }

    default6 = addresses["default6"]
    assert default6.type == AddressType.NETWORK
    assert default6.value == "::/0"
    assert default6.source_type == "ipprefix"
    assert default6.source_attributes == {}
    assert default6.source_effective_defaults == {
        "type": "ipprefix",
        "ip6": "::/0",
    }

    multicast4 = addresses["default-mcast4"]
    assert multicast4.type == AddressType.HOST
    assert multicast4.value == "0.0.0.0/32"
    assert multicast4.source_type == "multicastrange"
    assert multicast4.source_attributes == {}
    assert multicast4.source_effective_defaults == {
        "type": "multicastrange",
        "start_ip": "0.0.0.0",
        "end_ip": "0.0.0.0",
    }

    broadcast = addresses["default-broadcast"]
    assert broadcast.type == AddressType.HOST
    assert broadcast.value == "255.255.255.255/32"
    assert broadcast.source_type == "broadcastmask"
    assert broadcast.source_attributes == {}
    assert broadcast.source_effective_defaults == {
        "subnet": "0.0.0.0 0.0.0.0",
    }

    multicast6 = addresses["default-mcast6"]
    assert multicast6.type == AddressType.NETWORK
    assert multicast6.value == "::/0"
    assert multicast6.source_type is None
    assert multicast6.source_attributes == {}
    assert multicast6.source_effective_defaults == {"ip6": "::/0"}


def test_multi_value_fsso_group_preserves_boundaries_through_excel():
    parsed = _by_name(parse_fortigate_config(CONFIG).addresses)
    assert parsed["fsso-multi"].fsso_group == ["Finance Users", "Remote Users"]

    ir = FGToIRTransformer(parse_fortigate_config(CONFIG)).transform()
    addresses = _by_name(ir.addresses)
    fsso = addresses["fsso-multi"]
    assert fsso.source_fsso_group == ["Finance Users", "Remote Users"]
    assert fsso.source_attributes["fsso_group"] == ["Finance Users", "Remote Users"]

    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["Addresses"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Name"]).value: row
        for row in range(4, sheet.max_row + 1)
    }
    row = rows["fsso-multi"]
    fsso_cell = sheet.cell(row, headers["FSSO Group"]).value
    assert "Finance Users" in fsso_cell
    assert "Remote Users" in fsso_cell
    assert sheet.cell(row, headers["Effective Defaults"]).value in {None, ""}
