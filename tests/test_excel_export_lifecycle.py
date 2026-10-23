from importlib import import_module
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from fwmigrate.collection.contracts import CollectedSource, CollectionPart, CollectionStatus
from fwmigrate.collection.snapshot import make_snapshot
from fwmigrate.source_reporting.excel_lifecycle import report_workbook, save_report_workbook
from fwmigrate.source_reporting.excel_style import append_readable_report_row, check_report_dimensions
from fwmigrate.source_reporting.options import ExcelReportLimitError
from fwmigrate.web import create_app


SOURCE = 'config firewall address\n edit host\n set subnet 192.0.2.1 255.255.255.255\n next\nend'
EVIDENCE = {'vendor': 'fortigate', 'source_text': SOURCE, 'source_name': 'review.conf'}
CASES = (
    ('fortigate', 'FortiGateSourceReporter', None),
    ('palo_alto', 'PaloAltoSourceReporter', 'palo_alto/integrated_panorama.xml'),
    ('cisco_asa', 'CiscoASASourceReporter', 'cisco_asa/nat_pipeline_conformance.txt'),
    ('cisco_ftd', 'CiscoFTDSourceReporter', 'cisco_ftd/fmc_selected_domains.json'),
    ('juniper_srx', 'JuniperSRXSourceReporter', 'juniper_srx/reporting_contract.set'),
    ('checkpoint', 'CheckPointSourceReporter', 'checkpoint/multidomain_full.json'),
)


@pytest.mark.parametrize('vendor,reporter_name,fixture', CASES)
@pytest.mark.parametrize('failure', ['rows', 'save'])
def test_real_reporters_clean_up_failed_streams_without_touching_another_export(vendor, reporter_name, fixture, failure):
    reporter = getattr(import_module(f'fwmigrate.vendors.{vendor}.source_report'), reporter_name)()
    if vendor == 'juniper_srx':
        source = 'set system host-name srx\nset interfaces ge-0/0/0 unit 0 family inet address 192.0.2.1/24'
    else:
        source = (Path(__file__).parent / 'fixtures' / fixture).read_text() if fixture else SOURCE
    analysis = reporter.analyze_source(source)
    with report_workbook(Workbook(write_only=True)) as other:
        other.create_sheet('Still writing').append(['retained'])
        existing = set(ALL_TEMP_FILES)
        if failure == 'rows':
            from openpyxl.worksheet._write_only import WriteOnlyWorksheet
            original = WriteOnlyWorksheet.append
            calls = 0

            def broken(sheet, values):
                nonlocal calls
                calls += 1
                if calls == 5:
                    raise RuntimeError('row failure')
                return original(sheet, values)

            injection = patch.object(WriteOnlyWorksheet, 'append', broken)
        else:
            injection = patch('fwmigrate.source_reporting.excel_lifecycle._ReportExcelWriter.write_worksheet',
                              side_effect=RuntimeError('save failure'))
        with injection, pytest.raises(RuntimeError, match='failure'):
            reporter.export_excel(analysis, BytesIO(), profile='fast')
        assert set(ALL_TEMP_FILES) == existing
        assert all(Path(path).exists() for path in existing)
        output = BytesIO()
        save_report_workbook(other, output)
        assert load_workbook(output).active['A1'].value == 'retained'


@pytest.mark.parametrize('write_only', [False, True])
def test_zip_write_failure_closes_archive_and_worksheet_files(write_only, monkeypatch):
    before = set(ALL_TEMP_FILES)
    archives = []
    from zipfile import ZipFile
    original = ZipFile.write

    def fail(archive, *args, **kwargs):
        archives.append(archive)
        raise OSError('storage failure')

    with pytest.raises(OSError, match='storage failure'):
        with report_workbook(Workbook(write_only=write_only)) as workbook:
            sheet = workbook.create_sheet('Rows') if write_only else workbook.active
            sheet.append(['row'])
            monkeypatch.setattr(ZipFile, 'write', fail)
            save_report_workbook(workbook, BytesIO())
    monkeypatch.setattr(ZipFile, 'write', original)
    assert set(ALL_TEMP_FILES) == before
    assert archives and all(archive.fp is None for archive in archives)


def test_streamed_row_dimensions_are_bounded_and_persisted_in_xml():
    with report_workbook(Workbook(write_only=True)) as workbook:
        sheet = workbook.create_sheet('Rows')
        sheet.column_dimensions['A'].width = 20
        for row in range(1, 1001):
            append_readable_report_row(sheet, row, ['line one\nline two\nline three'])
            assert not sheet.row_dimensions
        output = BytesIO()
        save_report_workbook(workbook, output)
    restored = load_workbook(output).active
    assert restored.row_dimensions[1].height == restored.row_dimensions[1000].height == 49
    assert restored['A1000'].value == 'line one\nline two\nline three'
    assert restored['A1000'].alignment.wrap_text


@pytest.mark.parametrize('profile', ['fast', 'full'])
def test_http_collection_failure_reason_is_preserved(profile):
    snapshot = make_snapshot(CollectedSource('fortigate', SOURCE, 'review.conf', 'ssh', CollectionStatus.PARTIAL,
        parts=(CollectionPart('routes', 'FAILED', False, 0, 'TIMEOUT'),)))
    response = create_app({'TESTING': True}).test_client().post('/api/extract/excel', buffered=True,
        json={'source_vendor': 'fortigate', 'excel_profile': profile, 'source': snapshot})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert list(workbook['Collection Evidence'].values)[1] == ('routes', 'FAILED', False, 0, 'TIMEOUT')


@pytest.mark.parametrize('name', [[], {}, 1, 'x' * 256])
def test_malformed_metadata_is_rejected_before_analysis(monkeypatch, name):
    import fwmigrate.web_support.request_source as request_source
    with patch.object(request_source.source_reporters, 'get', side_effect=AssertionError('analysis should not start')):
        app = create_app({'TESTING': True})
        with app.test_request_context('/api/extract/excel', json={}):
            with pytest.raises(ValueError, match='Source name'):
                request_source.analyze_request_source({'source': {**EVIDENCE, 'source_name': name}}, 'fortigate')


@pytest.mark.parametrize('route', ['/api/preview', '/api/extract/excel', '/api/collection/snapshot/import'])
@pytest.mark.parametrize('transport', ['json', 'multipart'])
def test_oversized_requests_return_json_413(route, transport):
    client = create_app({'TESTING': True, 'MAX_CONTENT_LENGTH': 64}).test_client()
    kwargs = {'json': {'padding': 'x' * 100}} if transport == 'json' else {
        'data': {'file': (BytesIO(b'x' * 100), 'large.conf')}}
    response = client.post(route, **kwargs)
    assert response.status_code == 413
    assert 'size limit' in response.json['error']


@pytest.mark.parametrize('close_early', [False, True])
def test_export_capacity_covers_slow_responses_and_releases_on_completion_or_disconnect(close_early):
    app = create_app({'TESTING': True, 'MAX_CONCURRENT_EXPORTS': 1})
    client = app.test_client()
    payload = {'source_vendor': 'fortigate', 'source': EVIDENCE}
    response = client.post('/api/extract/excel', json=payload)
    assert response.status_code == 200
    busy = client.post('/api/extract/excel', json=payload)
    assert busy.status_code == 429 and busy.json['error_code'] == 'export_busy'
    assert client.get('/api/vendors').status_code == 200
    if not close_early:
        assert response.data.startswith(b'PK')
    response.close()
    response.close()  # Duplicate close must not over-release a bounded semaphore.
    assert client.post('/api/extract/excel', json=payload, buffered=True).status_code == 200
    assert client.post('/api/extract/excel', json={}, buffered=True).status_code == 400
    assert client.post('/api/extract/excel', json=payload, buffered=True).status_code == 200


def test_metrics_include_analysis_without_retaining_source_values(caplog):
    import logging
    client = create_app({'TESTING': True}).test_client()
    with caplog.at_level(logging.DEBUG, logger='fwmigrate.web_api.source'):
        response = client.post('/api/extract/excel', buffered=True,
            json={'source_vendor': 'fortigate', 'collect_metrics': True, 'source': EVIDENCE})
    assert response.status_code == 200
    assert 'source analysis' in caplog.text and 'fortigate_validation' in caplog.text
    assert '192.0.2.1' not in caplog.text and 'review.conf' not in caplog.text


def test_excel_limits_fail_explicitly_and_accept_boundary():
    check_report_dimensions(1_048_576, 16_384)
    for row, columns in [(1_048_577, 1), (1, 16_385)]:
        with pytest.raises(ExcelReportLimitError):
            check_report_dimensions(row, columns)


def test_panos_full_builds_each_worksheet_once(monkeypatch):
    from collections import Counter
    from fwmigrate.vendors.palo_alto.export.excel_rows import ROW_BUILDERS
    from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter
    calls = Counter()

    def counted(name, original):
        def rows(context):
            calls[name] += 1
            yield from original(context)
        return rows

    for name, builder in list(ROW_BUILDERS.items()):
        monkeypatch.setitem(ROW_BUILDERS, name, counted(name, builder))
    reporter = PaloAltoSourceReporter()
    output = BytesIO()
    reporter.export_excel(reporter.analyze_source('<config><shared><address><entry name="host">'
        '<ip-netmask>192.0.2.1/32</ip-netmask></entry></address></shared></config>'), output, profile='full')
    assert set(calls) == set(ROW_BUILDERS)
    assert set(calls.values()) == {1}
    workbook = load_workbook(output)
    assert workbook['Addresses']['A4'].value == 'host'
    assert any(row[:2] == ('Addresses', 1) for row in workbook['Summary'].values)
