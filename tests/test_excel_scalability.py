import io
import logging

from openpyxl import Workbook, load_workbook

from fwmigrate.ir.core import AddressType, IRAddress, IRConfig, IRMetadata
from fwmigrate.report import ExcelExportOptions, ExcelExportProfile, IRExcelExporter
from fwmigrate.report.excel_exporter import IRExcelExporter as BaseExporter
from fwmigrate.report.excel_audit import ExcelAuditAccumulator
from fwmigrate.report.excel_optimized import SinglePassIRExcelExporter
from fwmigrate.report.excel_serialization import profile_xlsx


def _ir() -> IRConfig:
    return IRConfig(
        metadata=IRMetadata(hostname="scalability", source_vendor="fortigate"),
        addresses=[
            IRAddress(
                name="review-address",
                type=AddressType.HOST,
                value="192.0.2.1/32",
                migration_status="UNSUPPORTED",
                requires_manual_review=True,
                audit_note="manual mapping",
            )
        ],
    )


def test_table_sheet_consumes_rows_once_and_partitions_deterministically():
    exporter = BaseExporter(
        IRConfig(metadata=IRMetadata(hostname="partition", source_vendor="fortigate")),
        options=ExcelExportOptions(max_rows_per_sheet=2),
    )
    workbook = Workbook()
    consumed = []

    def rows():
        for value in range(5):
            consumed.append(value)
            yield (value,)

    exporter._table_sheet(workbook, "Items", ("Value",), rows())

    assert consumed == [0, 1, 2, 3, 4]
    assert workbook.sheetnames == ["Sheet", "Items 1", "Items 2", "Items 3"]
    assert [workbook[name].cell(4, 1).value for name in workbook.sheetnames[1:]] == [0, 2, 4]


def test_profiles_preserve_values_and_keep_audit_sheets():
    outputs = {
        profile: load_workbook(
            io.BytesIO(
                IRExcelExporter(
                    _ir(),
                    options=ExcelExportOptions(profile=profile),
                ).generate()
            )
        )
        for profile in ExcelExportProfile
    }

    assert all("Review Required" in workbook.sheetnames for workbook in outputs.values())
    assert all("Extraction Evidence" in workbook.sheetnames for workbook in outputs.values())
    assert [outputs[profile]["Addresses"].cell(4, 1).value for profile in outputs] == [
        "review-address"
    ] * len(outputs)

    profile = profile_xlsx(
        io.BytesIO(
            IRExcelExporter(_ir()).generate()
        ).getvalue()
    )
    assert profile["zip_entries"] > 0
    assert profile["styles_xml_bytes"] > 0


class _LargeTableExporter(BaseExporter):
    LARGE_TABLE_ROW_THRESHOLD = 1


def _sized_table(profile, exporter_type=BaseExporter):
    exporter = exporter_type(
        IRConfig(metadata=IRMetadata(hostname="sizing", source_vendor="fortigate")),
        options=ExcelExportOptions(profile=profile),
    )
    workbook = Workbook()
    exporter._table_sheet(
        workbook,
        "Items",
        ("Value",),
        (("x" * 100 + "\nsecond line",), ("second",)),
    )
    return workbook["Items"]


def test_table_sizing_and_profile_cosmetics_follow_their_cost_policy():
    full_small = _sized_table(ExcelExportProfile.FULL)
    assert full_small.row_dimensions[4].height == 35
    assert full_small["A5"].fill.fill_type == "solid"
    assert full_small["A5"].font.name == "Aptos"

    full_large = _sized_table(ExcelExportProfile.FULL, _LargeTableExporter)
    assert full_large.row_dimensions[4].height is None
    assert full_large.column_dimensions["A"].width == 32

    fast = _sized_table(ExcelExportProfile.FAST)
    assert fast.row_dimensions[4].height is None
    assert fast["A5"].fill.fill_type is None
    assert fast["A5"].font.name != "Aptos"
    assert fast.column_dimensions["A"].width == 32

    data_only = _sized_table(ExcelExportProfile.DATA_ONLY)
    assert data_only.row_dimensions[4].height is None
    assert data_only["A1"].font.bold is not True


def test_audit_classifier_is_built_once_and_partition_names_are_preserved():
    calls = []
    accumulator = ExcelAuditAccumulator()
    category = lambda title: calls.append(title) or "Inventory"
    headers = ("Name", "Status")
    accumulator.add_row("Items", headers, ("one", "NORMALIZED"), 4, category)
    accumulator.add_row("Items", headers, ("two", "NORMALIZED"), 5, category)
    assert calls == ["Items"]

    exporter = BaseExporter(
        IRConfig(metadata=IRMetadata(hostname="partition-audit", source_vendor="fortigate")),
        options=ExcelExportOptions(max_rows_per_sheet=1),
    )
    exporter._audit_accumulator = accumulator
    workbook = Workbook()
    exporter._table_sheet(
        workbook,
        "Reviewable Items",
        ("Name", "Status"),
        (("one", "UNSUPPORTED"), ("two", "UNSUPPORTED")),
    )
    assert {row[4] for row in accumulator.review_rows} == {
        "Reviewable Items 1",
        "Reviewable Items 2",
    }


def test_vendor_registered_builders_only_run_for_active_sheets(monkeypatch):
    builders = (
        "_build_fortigate_source_configuration",
        "_build_pan_phase9_sheets",
        "_build_cisco_acp",
        "_build_checkpoint_access_rule_sheet",
    )
    expected = {
        "fortigate": {builders[0]},
        "palo_alto": {builders[1]},
        "cisco_asa": {builders[2]},
        "checkpoint": {builders[3]},
    }

    for vendor, expected_builders in expected.items():
        exporter = SinglePassIRExcelExporter(
            IRConfig(metadata=IRMetadata(hostname=vendor, source_vendor=vendor))
        )
        calls = []
        for builder in builders:
            monkeypatch.setattr(
                exporter,
                builder,
                lambda _workbook, builder=builder: calls.append(builder),
            )

        for builder in builders:
            exporter._build_registered_if_active(
                None,
                set(exporter._active_sheet_order()),
                builder,
            )

        assert set(calls) == expected_builders


def test_debug_metrics_are_recorded_during_table_writes():
    logger = logging.getLogger("fwmigrate.report.excel_optimized")
    previous_level = logger.level
    logger.setLevel(logging.DEBUG)
    try:
        exporter = SinglePassIRExcelExporter(_ir())
        exporter.generate()
    finally:
        logger.setLevel(previous_level)

    metrics = exporter._last_export_metrics
    assert metrics.worksheet_metrics
    assert metrics.rows_written == sum(item.rows for item in metrics.worksheet_metrics)
    assert metrics.cells_written == sum(item.cells for item in metrics.worksheet_metrics)
    assert metrics.nonempty_cells == sum(
        item.nonempty_cells for item in metrics.worksheet_metrics
    )
