import io

from openpyxl import load_workbook

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.ir.enums import NATType
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter
from fwmigrate.report.excel_optimized import SinglePassIRExcelExporter


ADVANCED_VIP6_CONFIG = '''
config firewall vip6
    edit "ADVANCED-VIP6"
        set extip 2001:db8::50
        set mappedip 2001:db8:1::50
        set http-redirect enable
        set max-embryonic-connections 250
        set ssl-certificate "CERT_A" "CERT_B"
        config quic
            set status enable
            set max-age 60
        end
        config ssl-cipher-suites
            set cipher-suite "TLS_AES_128_GCM_SHA256"
            append cipher-suite "TLS_AES_256_GCM_SHA384"
        end
        config ssl-server-cipher-suites
            set cipher-suite "TLS_CHACHA20_POLY1305_SHA256"
        end
    next
end
'''


def test_vip6_semantics_and_policy_dnat_are_preserved() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "WEB-VIP6"
        set extip 2001:db8::10
        set extintf "wan"
        set mappedip 2001:db8:1::10 2001:db8:1::11
        set src-filter 2001:db8:2::/64
        set portforward enable
        set protocol tcp
        set extport 443
        set mappedport 8443
    next
end
config firewall policy
    edit 100
        set srcaddr6 "all"
        set dstaddr6 "WEB-VIP6"
        set service "ALL"
    next
    end
    ''')
    vip = result.canonical_ir.virtual_ips[0]
    assert (vip.source_context, vip.external_interface) == ("root", "wan")
    assert vip.mapped_ips == ["2001:db8:1::10", "2001:db8:1::11"]
    assert vip.source_filters == ["2001:db8:2::/64"]
    assert vip.migration_status == "PARTIALLY_NORMALIZED"
    rule = result.canonical_ir.nat_rules[0]
    assert rule.type == NATType.DESTINATION
    assert (rule.source_policy_reference, rule.source_vip_reference) == ("100", "WEB-VIP6")
    assert rule.destination == ["2001:db8::10"]
    assert rule.translated_destinations == ["2001:db8:1::10", "2001:db8:1::11"]
    assert (rule.original_destination_port, rule.translated_port) == ("443", "8443")
    assert rule.original_destination_ports[0].start == 443
    assert rule.translated_destination_ports[0].start == 8443
    assert any(d.source_field == "dstaddr6-vip" and d.reference == "WEB-VIP6" and d.result == "RESOLVED" for d in result.dependencies)


def test_advanced_vip6_does_not_become_static_dnat() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "LB-VIP6"
        set type server-load-balance
        set extip 2001:db8::20
        set mappedip 2001:db8:1::20
    next
end
config firewall policy
    edit 101
        set srcaddr6 "all"
        set dstaddr6 "LB-VIP6"
        set service "ALL"
    next
end
''')

    assert result.canonical_ir.virtual_ips[0].vip_type == "server-load-balance"
    assert result.canonical_ir.nat_rules == []
    assert any("canonical DNAT was withheld" in entry.message for entry in result.canonical_ir.audit_entries)


def test_nat64_uses_ipv4_pool_and_vip6_translation_fields_end_to_end() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "NAT64_VIP6"
        set extip 2001:db8:64::150
        set nat66 disable
        set nat64 enable
        set ipv4-mappedip 192.0.2.156
        set portforward enable
        set protocol tcp
        set extport 443
        set ipv4-mappedport 8443
    next
end
config firewall ippool
    edit "NAT64_POOL4"
        set startip 192.0.2.101
        set endip 192.0.2.101
        set nat64 enable
        set add-nat64-route enable
    next
end
config firewall policy
    edit 202
        set name "NAT64_POLICY"
        set srcintf "LAN6"
        set dstintf "WAN6"
        set action accept
        set nat64 enable
        set srcaddr "all"
        set dstaddr "all"
        set srcaddr6 "all"
        set dstaddr6 "NAT64_VIP6"
        set schedule "always"
        set service "ALL"
        set ippool enable
        set poolname "NAT64_POOL4"
    next
end
''')

    rules = [
        rule for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "202"
    ]
    assert len(rules) == 1
    rule = rules[0]
    assert rule.nat_family.value == "nat64"
    assert (rule.original_address_family, rule.translated_address_family) == (
        "ipv6",
        "ipv4",
    )
    assert rule.destination == ["2001:db8:64::150"]
    assert rule.translated_destinations == ["192.0.2.156"]
    assert rule.translated_sources == []
    assert rule.source_pool_references == ["NAT64_POOL4"]
    assert rule.source_vip_reference == "NAT64_VIP6"
    assert rule.original_destination_ports[0].start == 443
    assert rule.translated_destination_ports[0].start == 8443
    assert any(
        dependency.reference == "NAT64_POOL4"
        and dependency.expected_type == "firewall ippool"
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )
    assert any(
        dependency.reference == "NAT64_VIP6"
        and dependency.expected_type in {"firewall vip6", "firewall vipgrp6"}
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(
        row for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Name"]).value == rule.name
    )
    assert sheet.cell(row, headers["Original Destination"]).value == "2001:db8:64::150"
    assert sheet.cell(row, headers["Translated Destination"]).value == "192.0.2.156"
    assert sheet.cell(row, headers["Translated Source"]).value is None
    assert sheet.cell(row, headers["NAT Family"]).value == "nat64"
    assert sheet.cell(row, headers["IP Pool"]).value == "NAT64_POOL4"
    assert sheet.cell(row, headers["VIP"]).value == "NAT64_VIP6"
    assert sheet.cell(row, headers["Original Destination Port"]).value == "443"
    assert sheet.cell(row, headers["Translated Destination Port"]).value == "8443"


def test_nat64_vip6_group_preserves_member_and_group_references() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "NAT64_GROUP_MEMBER"
        set extip 2001:db8:64::151
        set nat66 disable
        set nat64 enable
        set ipv4-mappedip 192.0.2.157
    next
end
config firewall vipgrp6
    edit "NAT64_GROUP"
        set member "NAT64_GROUP_MEMBER"
    next
end
config firewall ippool
    edit "NAT64_GROUP_POOL4"
        set startip 192.0.2.102
        set endip 192.0.2.102
        set nat64 enable
    next
end
config firewall policy
    edit 204
        set action accept
        set nat64 enable
        set srcaddr "all"
        set dstaddr "all"
        set srcaddr6 "all"
        set dstaddr6 "NAT64_GROUP"
        set service "ALL"
        set ippool enable
        set poolname "NAT64_GROUP_POOL4"
    next
end
''')

    rules = [rule for rule in result.canonical_ir.nat_rules if rule.source_policy_reference == "204"]
    assert len(rules) == 1
    rule = rules[0]
    assert rule.source_vip_reference == "NAT64_GROUP_MEMBER"
    assert rule.source_vip_group_reference == "NAT64_GROUP"
    assert rule.translated_destinations == ["192.0.2.157"]
    assert any(
        dependency.reference == "NAT64_GROUP"
        and dependency.target_path == "firewall vipgrp6"
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )


def test_vipgrp6_expands_in_order_and_reports_missing_members() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "A"
        set extip 2001:db8::1
        set mappedip 2001:db8:1::1
    next
    edit "B"
        set extip 2001:db8::2
        set mappedip 2001:db8:1::2
    next
end
config firewall vipgrp6
    edit "G"
        set member "B" "MISSING" "A"
    next
end
config firewall policy
    edit 1
        set srcaddr6 "all"
        set dstaddr6 "G"
        set service "ALL"
    next
end
''')
    assert [rule.source_vip_reference for rule in result.canonical_ir.nat_rules] == ["B", "A"]
    assert all(rule.source_vip_group_reference == "G" for rule in result.canonical_ir.nat_rules)
    group = result.canonical_ir.virtual_ip_groups[0]
    assert group.members == ["B", "MISSING", "A"] and group.requires_manual_review
    assert any(d.source_path == "firewall vipgrp6" and d.reference == "MISSING" and d.result == "UNRESOLVED" for d in result.dependencies)


def test_vip6_correlation_is_vdom_scoped_and_combines_snat6() -> None:
    result = extract_fortigate_config('''
config vdom
    edit "tenant-a"
        config firewall vip6
            edit "WEB"
                set extip 2001:db8::10
                set mappedip 2001:db8:1::10
            next
        end
        config firewall ippool6
            edit "POOL6"
                set startip 2001:db8:2::1
                set endip 2001:db8:2::1
            next
        end
        config firewall policy
            edit 10
                set srcaddr6 "all"
                set dstaddr6 "WEB"
                set service "ALL"
                set nat enable
                set ippool enable
                set poolname6 "POOL6"
            next
        end
    next
    edit "tenant-b"
        config firewall policy
            edit 20
                set srcaddr6 "all"
                set dstaddr6 "WEB"
                set service "ALL"
            next
        end
    next
end
''')
    assert len(result.canonical_ir.nat_rules) == 1
    rule = result.canonical_ir.nat_rules[0]
    assert (rule.type, rule.source_context) == (NATType.TWICE, "tenant-a")
    assert rule.translated_sources == []


def test_vip6_source_model_preserves_typed_fields_lists_and_nested_nodes() -> None:
    fg = parse_fortigate_config(ADVANCED_VIP6_CONFIG)

    vip = fg.vips6[0]
    assert vip.http_redirect == "enable"
    assert vip.max_embryonic_connections == 250
    assert vip.ssl_certificate == "CERT_A CERT_B"
    assert vip.mappedip == ["2001:db8:1::50"]
    assert [node.name for node in vip.nested_configs] == [
        "quic",
        "ssl-cipher-suites",
        "ssl-server-cipher-suites",
    ]
    assert vip.nested_configs[0].commands[0].key == "status"
    assert vip.nested_configs[1].commands[0].values == [
        "TLS_AES_128_GCM_SHA256"
    ]
    assert vip.nested_configs[1].commands[1].operation == "append"


def test_vip6_effective_defaults_do_not_overwrite_explicit_values() -> None:
    result = extract_fortigate_config('''
config firewall vip6
    edit "DEFAULTS"
        set extip 2001:db8::60
        set mappedip 2001:db8:1::60
    next
    edit "ENABLED"
        set extip 2001:db8::61
        set mappedip 2001:db8:1::61
        set nat66 enable
        set ndp-reply enable
        set add-nat64-route enable
        set portforward enable
    next
    edit "DISABLED"
        set extip 2001:db8::62
        set mappedip 2001:db8:1::62
        set nat66 disable
        set ndp-reply disable
        set add-nat64-route disable
        set portforward disable
    next
end
''')

    by_name = {vip.name: vip for vip in result.canonical_ir.virtual_ips}
    assert by_name["DEFAULTS"].nat66 is None
    assert by_name["DEFAULTS"].ndp_reply is None
    assert by_name["DEFAULTS"].add_nat64_route is None
    assert by_name["DEFAULTS"].source_effective_settings == {
        "protocol": "tcp",
        "portforward": "disable",
        "nat66": "enable",
        "ndp_reply": "enable",
        "add_nat64_route": "enable",
    }
    assert by_name["ENABLED"].source_effective_settings == {
        "protocol": "tcp",
        "portforward": "enable",
        "nat66": "enable",
        "ndp_reply": "enable",
        "add_nat64_route": "enable",
    }
    assert by_name["DISABLED"].source_effective_settings == {
        "protocol": "tcp",
        "portforward": "disable",
        "nat66": "disable",
        "ndp_reply": "disable",
        "add_nat64_route": "disable",
    }
    assert by_name["ENABLED"].nat66 is True
    assert by_name["DISABLED"].nat66 is False


def test_advanced_vip6_nested_config_requires_review_and_withholds_dnat() -> None:
    result = extract_fortigate_config(ADVANCED_VIP6_CONFIG + '''
config firewall policy
    edit 250
        set srcaddr6 "all"
        set dstaddr6 "ADVANCED-VIP6"
        set service "ALL"
    next
end
''')

    vip = result.canonical_ir.virtual_ips[0]
    assert vip.nested_source_configs
    assert vip.http_redirect is True
    assert vip.max_embryonic_connections == 250
    assert vip.requires_manual_review
    assert vip.migration_status == "PARTIALLY_NORMALIZED"
    assert result.canonical_ir.nat_rules == []
    assert any("canonical DNAT was withheld" in entry.message for entry in result.canonical_ir.audit_entries)

    statuses = {
        section.path: section.status
        for section in result.source_sections
        if section.path in {"firewall vip6", "firewall vip6 realservers"}
    }
    assert statuses["firewall vip6"] == ExtractionStatus.PARTIALLY_NORMALIZED


def test_vip6_excel_exports_typed_effective_and_nested_source_data() -> None:
    result = extract_fortigate_config(ADVANCED_VIP6_CONFIG)

    for exporter in (
        IRExcelExporter(result.canonical_ir, result),
        SinglePassIRExcelExporter(result.canonical_ir, result),
    ):
        workbook = load_workbook(io.BytesIO(exporter.generate()))
        assert "VIP Nested Configuration" in workbook.sheetnames

        virtual_ips = workbook["Virtual IPs"]
        headers = {cell.value: cell.column for cell in virtual_ips[3]}
        row = next(
            row for row in range(4, virtual_ips.max_row + 1)
            if virtual_ips.cell(row, headers["Name"]).value == "ADVANCED-VIP6"
        )
        assert virtual_ips.cell(row, headers["HTTP Redirect"]).value == "TRUE"
        assert virtual_ips.cell(row, headers["Max Embryonic Connections"]).value == 250
        assert "nat66" in virtual_ips.cell(row, headers["Effective Source Settings"]).value

        nested = workbook["VIP Nested Configuration"]
        flattened = "\n".join(
            " | ".join(str(value or "") for value in values)
            for values in nested.iter_rows(values_only=True)
        )
        assert "quic" in flattened
        assert "ssl-cipher-suites" in flattened
        assert "TLS_AES_256_GCM_SHA384" in flattened
        assert "ssl-server-cipher-suites" in flattened
        assert "VIP Real Servers" in workbook.sheetnames
