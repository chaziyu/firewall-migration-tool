import io

from openpyxl import load_workbook

from fwmigrate.ir import IRConfig
from fwmigrate.ir.enums import NATType
from fwmigrate.ir.metadata import IRMetadata
from fwmigrate.ir.nat import IRNATRule
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report.excel_exporter import IRExcelExporter


def _workbook(result):
    return load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[3] if cell.value}


def _row(sheet, headers, column, value):
    return next(
        row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers[column]).value == value
    )


def test_nat_excel_exposes_central_snat_source_and_enabled_mode():
    result = extract_fortigate_config("""
config system settings
    set central-nat enable
end
config firewall central-snat-map
    edit 1
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
    next
end
""")

    sheet = _workbook(result)["NAT Rules"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Name", "central-snat-1")

    assert sheet.cell(row, headers["NAT Source Mechanism"]).value == "central-snat-map"
    assert sheet.cell(row, headers["Effective Central NAT Mode"]).value == "enable"
    assert sheet.column_dimensions[
        sheet.cell(3, headers["NAT Source Mechanism"]).column_letter
    ].hidden is False
    assert sheet.column_dimensions[
        sheet.cell(3, headers["Effective Central NAT Mode"]).column_letter
    ].hidden is False


def test_nat_excel_marks_unproven_central_nat_mode_unknown():
    result = extract_fortigate_config("""
config firewall central-snat-map
    edit 2
        set srcintf "lan"
        set dstintf "wan"
        set orig-addr "all"
        set dst-addr "all"
    next
end
""")

    sheet = _workbook(result)["NAT Rules"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Name", "central-snat-2")

    assert sheet.cell(row, headers["NAT Source Mechanism"]).value == "central-snat-map"
    assert sheet.cell(row, headers["Effective Central NAT Mode"]).value == "unknown"


def test_nat_excel_exposes_policy_nat_source_and_disabled_central_mode():
    result = extract_fortigate_config("""
config system settings
    set central-nat disable
end
config system interface
    edit "lan"
    next
    edit "wan"
    next
end
config firewall address
    edit "source"
        set subnet 10.0.0.0 255.255.255.0
    next
end
config firewall vip
    edit "web-vip"
        set extip 203.0.113.10
        set mappedip "10.0.0.10"
    next
end
config firewall policy
    edit 20
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "source"
        set dstaddr "web-vip"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    rule = result.canonical_ir.nat_rules[0]
    assert rule.source_origin == "firewall-policy"

    sheet = _workbook(result)["NAT Rules"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Source Policy ID", "20")

    assert sheet.cell(row, headers["NAT Source Mechanism"]).value == "firewall-policy"
    assert sheet.cell(row, headers["Effective Central NAT Mode"]).value == "disable"


def test_non_fortigate_nat_leaves_central_nat_mode_blank():
    ir = IRConfig(
        metadata=IRMetadata(source_vendor="checkpoint"),
        nat_rules=[
            IRNATRule(
                name="checkpoint-nat",
                type=NATType.SOURCE,
                source_origin="checkpoint-nat-rulebase",
            )
        ],
    )
    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["NAT Rules"]
    headers = _headers(sheet)

    assert sheet.cell(4, headers["NAT Source Mechanism"]).value == "checkpoint-nat-rulebase"
    assert sheet.cell(4, headers["Effective Central NAT Mode"]).value is None


def test_vip_group_set_append_survives_source_ir_and_excel():
    content = """
config firewall vip
    edit "VIP1"
        set extip 203.0.113.10
        set mappedip "10.0.0.10"
    next
    edit "VIP2"
        set extip 203.0.113.11
        set mappedip "10.0.0.11"
    next
end
config firewall vipgrp
    edit "GROUP"
        set member "VIP1"
        append member "VIP2"
    next
end
"""
    parsed = parse_fortigate_config(content)
    assert parsed.vip_groups[0].member == ["VIP1", "VIP2"]

    result = extract_fortigate_config(content)
    group = result.canonical_ir.virtual_ip_groups[0]
    assert group.members == ["VIP1", "VIP2"]

    inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall vipgrp" and item.name == "GROUP"
    )
    assert [
        (command.operation, command.key, command.values)
        for command in inventory.commands
        if command.key == "member"
    ] == [
        ("set", "member", ["VIP1"]),
        ("append", "member", ["VIP2"]),
    ]

    sheet = _workbook(result)["VIP Groups"]
    headers = _headers(sheet)
    row = _row(sheet, headers, "Name", "GROUP")
    assert sheet.cell(row, headers["Members"]).value == "VIP1\nVIP2"


def test_vip_group_unset_clears_effective_members_but_preserves_command_evidence():
    content = """
config firewall vip
    edit "VIP1"
        set extip 203.0.113.10
        set mappedip "10.0.0.10"
    next
    edit "VIP2"
        set extip 203.0.113.11
        set mappedip "10.0.0.11"
    next
end
config firewall vipgrp
    edit "GROUP"
        set member "VIP1" "VIP2"
        unset member
    next
end
"""
    parsed = parse_fortigate_config(content)
    assert parsed.vip_groups[0].member == []

    result = extract_fortigate_config(content)
    assert result.canonical_ir.virtual_ip_groups[0].members == []

    inventory = next(
        item
        for item in result.inventory_items
        if item.source_path == "firewall vipgrp" and item.name == "GROUP"
    )
    assert any(
        command.operation == "unset" and command.key == "member"
        for command in inventory.commands
    )
