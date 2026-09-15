from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_repaired_junos_fields_survive_the_full_excel_pipeline():
    ir = JuniperSRXParser("""
    set interfaces ge-0/0/0 description uplink
    set interfaces ge-0/0/0 mtu 1500
    set interfaces ge-0/0/0 unit 0 family inet address 10.0.0.1/24
    set interfaces ge-0/0/0 unit 0 family inet6 address 2001:db8::1/64
    set interfaces ge-0/0/0 unit 0 family inet filter input F
    set security zones security-zone trust interfaces ge-0/0/0.0
    set security zones security-zone trust host-inbound-traffic system-services all except ssh
    set routing-options static route default next-hop 192.0.2.1 install
    set firewall family inet filter F term t from source-address 10.0.0.0/24
    set firewall family inet filter F term t then next-interface ge-0/0/0.0
    set routing-instances RI instance-type virtual-router
    set security idp idp-policy idp1 rulebase-ips rule r action drop
    set security utm utm-policy utm1
    set services ssl proxy profile ssl1
    set security address-book global address internal_srv 172.16.1.50/32
    set security policies from-zone trust to-zone untrust policy p match source-address 10.0.0.0/24
    set security policies from-zone trust to-zone untrust policy p match destination-address 0.0.0.0/0
    set security policies from-zone trust to-zone untrust policy p match application any
    set security policies from-zone trust to-zone untrust policy p then permit application-services idp-policy idp1
    set security nat destination pool dst address 10.0.0.10/32 port 8080
    set security nat destination pool dst routing-instance RI
    set security nat destination rule-set d from zone untrust
    set security nat destination rule-set d rule d1 match destination-address 198.51.100.10/32
    set security nat destination rule-set d rule d1 then destination-nat pool dst
    set security nat static rule-set s from zone untrust
    set security nat static rule-set s rule s1 match destination-address 203.0.113.10/32
    set security nat static rule-set s rule s1 then static-nat prefix-name internal_srv
    set security nat static rule-set s rule s1 then static-nat mapped-port 8443
    """).transform_to_ir()

    physical = next(item for item in ir.interfaces if item.name == "ge-0/0/0")
    logical = next(item for item in ir.interfaces if item.name == "ge-0/0/0.0")
    assert physical.description == "uplink" and physical.mtu == 1500
    assert logical.parent == "ge-0/0/0"
    assert logical.ipv6_address == "2001:db8::1/64" and not logical.secondary_ips
    assert ir.routes[0].destination == "0.0.0.0/0" and ir.routes[0].installation == "install"
    assert ir.local_in_policies[0].service_exclusions == ["ssh"]
    assert ir.policy_route_rules[0].next_interface == "ge-0/0/0.0"
    policy = next(item for item in ir.policies if item.name == "p")
    assert policy.source_security_profile_references == {"idp-policy": "idp1"}
    assert ir.ip_pools[0].routing_instance == "RI"
    assert any(item.name == "s1" and item.translated_destinations == ["172.16.1.50/32"] for item in ir.nat_rules)

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    assert {"Interfaces", "Local-In Policies", "Routes", "Cisco PBR", "Firewall Filters"} <= set(workbook.sheetnames)
    headers = {cell.value: cell.column for cell in workbook["Routes"][3]}
    row = next(workbook["Routes"].iter_rows(min_row=4, values_only=False))
    assert row[headers["Destination Prefix (Normalized)"] - 1].value == "0.0.0.0/0"
    assert row[headers["Installation"] - 1].value == "install"

    interface_headers = {cell.value: cell.column for cell in workbook["Interfaces"][3]}
    interface_rows = list(workbook["Interfaces"].iter_rows(min_row=4, values_only=False))
    interface_names = {row[interface_headers["Name"] - 1].value for row in interface_rows}
    assert {"ge-0/0/0", "ge-0/0/0.0"} <= interface_names
    logical_row = next(row for row in interface_rows if row[interface_headers["Name"] - 1].value == "ge-0/0/0.0")
    assert logical_row[interface_headers["Parent / Underlay Interface"] - 1].value == "ge-0/0/0"

    local_headers = {cell.value: cell.column for cell in workbook["Local-In Policies"][3]}
    local_row = next(workbook["Local-In Policies"].iter_rows(min_row=4, values_only=False))
    assert local_row[local_headers["Service Exclusions"] - 1].value == "ssh"

    pbr_headers = {cell.value: cell.column for cell in workbook["Cisco PBR"][3]}
    pbr_row = next(workbook["Cisco PBR"].iter_rows(min_row=4, values_only=False))
    assert pbr_row[pbr_headers["Next Interface"] - 1].value == "ge-0/0/0.0"

    filter_headers = {cell.value: cell.column for cell in workbook["Firewall Filters"][3]}
    filter_row = next(workbook["Firewall Filters"].iter_rows(min_row=4, values_only=False))
    assert filter_row[filter_headers["From Conditions"] - 1].value

    policy_headers = {cell.value: cell.column for cell in workbook["Policies"][3]}
    policy_row = next(row for row in workbook["Policies"].iter_rows(min_row=4, values_only=False) if row[policy_headers["Name"] - 1].value == "p")
    assert "idp1" in policy_row[policy_headers["Security Profile References"] - 1].value

    pool_headers = {cell.value: cell.column for cell in workbook["IP Pools"][3]}
    pool_row = next(row for row in workbook["IP Pools"].iter_rows(min_row=4, values_only=False) if row[pool_headers["Name"] - 1].value == "dst")
    assert pool_row[pool_headers["Routing Instance"] - 1].value == "RI"
    assert pool_row[pool_headers["Addresses"] - 1].value == "10.0.0.10/32"

    nat_headers = {cell.value: cell.column for cell in workbook["NAT Rules"][3]}
    nat_row = next(row for row in workbook["NAT Rules"].iter_rows(min_row=4, values_only=False) if row[nat_headers["Name"] - 1].value == "s1")
    assert "8443" in str(nat_row[nat_headers["Translated Destination Port"] - 1].value)
    assert nat_row[nat_headers["Rule Set"] - 1].value == "s"
    assert nat_row[nat_headers["Rule Position"] - 1].value == 1
    assert nat_row[nat_headers["Type"] - 1].value == "static"
    assert nat_row[nat_headers["Static NAT Bi-directional"] - 1].value == "TRUE"


def test_junos_address_set_topology_is_exported():
    ir = JuniperSRXParser("""
    set security address-book global address host 10.0.0.1/32
    set security address-book global address-set child address host
    set security address-book global address-set parent address-set child
    """).extract().canonical_ir

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    headers = {cell.value: cell.column for cell in workbook["Address Groups"][3]}
    row = next(
        row for row in workbook["Address Groups"].iter_rows(min_row=4, values_only=False)
        if row[headers["Name"] - 1].value == "parent"
    )
    assert row[headers["Direct Address Members"] - 1].value in (None, "")
    assert row[headers["Nested Address-Set Members"] - 1].value == "child"
    assert row[headers["Resolved Members"] - 1].value == "host"
