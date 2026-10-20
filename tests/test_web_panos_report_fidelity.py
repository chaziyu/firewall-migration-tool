"""Reference-shaped PAN-OS XML through the same preview/download path as the UI."""

import io
import json
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from fwmigrate.vendors.palo_alto.collection_source import PANOSCollectedSourceSanitizer
from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter
from fwmigrate.web import create_app


# PAN-OS 11.1 configure hierarchy: network tunnel ipsec auto-key and tunnel-monitor.
SOURCE = """<config version="11.1.0"><devices><entry name="fw"><network>
  <ike><gateway><entry name="gw"><authentication><pre-shared-key><key>VPN_SECRET_SENTINEL</key></pre-shared-key></authentication></entry></gateway>
    <crypto-profiles><ike-crypto-profiles><entry name="ike-profile"/></ike-crypto-profiles>
    <ipsec-crypto-profiles><entry name="ipsec-profile"/></ipsec-crypto-profiles></crypto-profiles></ike>
  <tunnel><ipsec><entry name="vpn"><tunnel-interface>tunnel.1</tunnel-interface>
    <auto-key><ike-gateway><entry name="gw"/></ike-gateway><ipsec-crypto-profile>ipsec-profile</ipsec-crypto-profile>
      <proxy-id><entry name="v4"><local>192.0.2.0/24</local><remote>198.51.100.0/24</remote><protocol><any/></protocol></entry></proxy-id>
      <proxy-id-v6><entry name="v6"><local>2001:db8:1::/64</local><remote>2001:db8:2::/64</remote><protocol><any/></protocol></entry></proxy-id-v6>
    </auto-key><tunnel-monitor><enable>no</enable><destination-ip>198.51.100.1</destination-ip>
      <tunnel-monitor-profile>monitor-profile</tunnel-monitor-profile><proxy-id>v4</proxy-id><future-setting>keep-monitor</future-setting></tunnel-monitor>
  </entry></ipsec></tunnel>
  <interface><ethernet><entry name="ethernet1/1"><layer3><units><entry name="ethernet1/1.10"><tag>10</tag><ip><entry name="192.0.2.1/24"/></ip></entry></units></layer3></entry></ethernet><tunnel><units><entry name="tunnel.1"/></units></tunnel></interface>
  <virtual-router><entry name="vr"><interface><member>ethernet1/1.10</member></interface></entry></virtual-router>
</network><vsys><entry name="vsys1"><import><network><interface><member>ethernet1/1.10</member></interface></network></import>
  <tag><entry name="tag1"><color>color1</color></entry></tag>
  <zone><entry name="trust"><network><layer3><member>ethernet1/1.10</member></layer3></network></entry></zone>
  <profile-group><entry name="profiles"><virus><member>default</member></virus></entry></profile-group>
  <schedule><entry name="office"><schedule-type><recurring><daily><member>08:00-17:00</member></daily></recurring></schedule-type></entry></schedule>
  <application><entry name="source-only-app"><future-setting>keep-app</future-setting></entry></application>
  <rulebase><security><rules><entry name="disabled-rule" uuid="11111111-2222-3333-4444-555555555555">
    <from><member>trust</member></from><to><member>any</member></to><source><member>192.0.2.0/24</member></source>
    <destination><member>198.51.100.0/24</member></destination><application><member>ssl</member></application>
    <service><member>application-default</member></service><action>allow</action><disabled>yes</disabled><log-end>no</log-end>
    <profile-setting><group><member>profiles</member></group></profile-setting>
  </entry><entry name="unspecified-state"><action>deny</action></entry></rules></security>
  <default-security-rules><rules><entry name="intrazone-default"><action>allow</action></entry></rules></default-security-rules>
  <nat><rules><entry name="nat-disabled"><from><member>trust</member></from><to><member>any</member></to>
    <source><member>192.0.2.0/24</member></source><destination><member>198.51.100.1</member></destination>
    <service>any</service><disabled>yes</disabled><source-translation><static-ip><translated-address>203.0.113.1</translated-address></static-ip></source-translation>
    <destination-translation><translated-address>198.51.100.2</translated-address><translated-port>8443</translated-port></destination-translation>
  </entry></rules></nat></rulebase>
</entry></vsys></entry></devices></config>"""


def sheet_rows(workbook, name):
    rows = iter(workbook[name].values)
    next(rows)
    next(rows)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


@pytest.mark.parametrize("profile", ["full", "fast"])
def test_preview_and_excel_preserve_reference_shaped_vpn_and_rule_state(profile):
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/preview", data={"source_vendor": "palo_alto", "file": (io.BytesIO(SOURCE.encode()), "firewall.xml")})
    assert response.status_code == 200
    preview = response.get_json()
    assert "VPN_SECRET_SENTINEL" not in json.dumps(preview)
    sections = preview["sections"]
    policy, unspecified = sections["policies"]
    assert policy["rule_order"] == 1 and unspecified["rule_order"] == 2
    assert policy["rule_uuid"] == "11111111-2222-3333-4444-555555555555"
    assert policy["source_order"] != policy["rule_order"]
    assert policy["disabled"] == "yes" and unspecified["disabled"] is None
    assert policy["source_addresses"] == ["192.0.2.0/24"]
    assert policy["destination_addresses"] == ["198.51.100.0/24"]
    assert policy["application"] == ["ssl"] and policy["log_end"] == "no"
    assert policy["profile_setting"]["groups"] == ["profiles"]
    nat = sections["nat"][0]
    assert nat["source_addresses"] == ["192.0.2.0/24"] and nat["disabled"] == "yes"
    assert nat["source_translation"]["translated_address"] == "203.0.113.1"
    assert nat["destination_translation"]["translated_address"] == "198.51.100.2"
    assert nat["destination_translation"]["translated_port"] == "8443"
    assert sections["schedules"][0]["recurring"]["daily"] == ["08:00-17:00"]
    for name in ("tags", "zones", "security_profile_groups", "default_security_rules", "ike_crypto_profiles", "ipsec_crypto_profiles"):
        assert len(sections[name]) == 1
    tunnel = sections["ipsec_tunnels"][0]
    assert tunnel["ike_gateways"] == ["gw"] and tunnel["ipsec_crypto_profile"] == "ipsec-profile"
    assert tunnel["tunnel_monitor"] == "no"
    assert tunnel["tunnel_monitor_destination_ip"] == "198.51.100.1"
    assert tunnel["tunnel_monitor_profile"] == "monitor-profile" and tunnel["tunnel_monitor_proxy_id"] == "v4"
    assert tunnel["additional_settings"]["tunnel-monitor"]["future-setting"] == "keep-monitor"
    assert {row["address_family"] for row in sections["vpn_phase2"]} == {"ipv4", "ipv6"}
    assert sections["ike_gateways"][0]["pre_shared_key_configured"] is True
    interface = next(row for row in sections["interfaces"] if row["name"] == "ethernet1/1.10")
    assert interface["tag"] == "10" and interface["virtual_routers"] == ["vr"]
    assert not preview["validation"]
    assert any(row["domain"] == "application" and row["status"] == "SOURCE_ONLY" for row in sections["extraction_coverage"])
    assert len(sections["source_inventory"]) == preview["records_total"]
    response = client.post("/api/extract/excel", json={"source_vendor": "palo_alto", "source": preview["source_evidence"], "excel_profile": profile})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        assert len(sheet_rows(workbook, "IPsec Proxy IDs")) == 2
        tunnel = sheet_rows(workbook, "IPsec Tunnels")[0]
        assert tunnel["IKE Gateways"] == "gw" and tunnel["IPsec Crypto Profile"] == "ipsec-profile"
        assert tunnel["Tunnel Monitor"] == "no" and tunnel["Tunnel Monitor Profile"] == "monitor-profile"
        assert tunnel["Tunnel Monitor Destination IP"] == "198.51.100.1" and tunnel["Tunnel Monitor Proxy ID"] == "v4"
        assert sheet_rows(workbook, "Security Policies")[0]["Disabled"] == "yes"
        assert all("VPN_SECRET_SENTINEL" not in str(cell.value) for sheet in workbook for row in sheet for cell in row)
        if profile == "full":
            appendix = sheet_rows(workbook, "PAN-OS Source Appendix")
            assert any(row["Evidence Value"] == "ipsec-profile" for row in appendix)
    finally:
        workbook.close()


def test_only_structural_ipsec_key_containers_are_preserved_and_secrets_stay_redacted():
    source = """<config><shared><address><entry name="x"><auto-key>FAKE_SECRET</auto-key></entry></address>
      <network><ipsec><entry name="vpn"><manual-key><esp><encryption><algorithm>aes-256-cbc</algorithm>
      <key>MANUAL_SECRET</key></encryption></esp></manual-key></entry></ipsec></network></shared></config>"""
    sanitizer = PANOSCollectedSourceSanitizer()
    safe = sanitizer.sanitize(source)
    assert sanitizer.sanitize(safe) == safe
    root = ET.fromstring(safe)
    assert root.findtext(".//manual-key/esp/encryption/algorithm") == "aes-256-cbc"
    reporter = PaloAltoSourceReporter()
    analysis = reporter.analyze_source(source)
    assert analysis.config.ipsec_tunnels[0].manual_key_configured is True
    surfaces = safe + analysis.config.model_dump_json() + json.dumps(reporter.build_preview(analysis))
    assert "FAKE_SECRET" not in surfaces and "MANUAL_SECRET" not in surfaces
    assert any("aes-256-cbc" in item.raw_xml for item in analysis.config.source_inventory)


@pytest.mark.parametrize("profile", ["full", "fast"])
def test_panos_non_secret_authentication_metadata_survives_preview_and_excel(profile):
    # PAN-OS 11.1: GlobalProtect client-auth labels/options and credential-enforcement mode.
    source = """<config><shared><certificate><entry name="public-cert">
      <public-key>PUBLIC_CERTIFICATE</public-key><private-key>PRIVATE_SENTINEL</private-key>
      </entry></certificate></shared><devices><entry name="fw"><vsys><entry name="vsys1">
      <profiles><url-filtering><entry name="url"><credential-enforcement><mode>
      <domain-credentials/></mode></credential-enforcement></entry></url-filtering></profiles>
      <global-protect><global-protect-portal><entry name="portal"><portal-config><client-auth>
      <entry name="client"><password-label>Password prompt</password-label>
      <user-credential-or-client-cert-required>no</user-credential-or-client-cert-required>
      </entry></client-auth></portal-config><client-config><agent-user-override-key>OVERRIDE_SENTINEL</agent-user-override-key>
      <configs><entry name="agent"><save-user-credentials>1</save-user-credentials></entry></configs>
      </client-config></entry></global-protect-portal><global-protect-gateway><entry name="gateway">
      <client-auth><entry name="client"><os>Any</os><auto-retrieve-passcode>no</auto-retrieve-passcode>
      <password-label>Password prompt</password-label><password>PASSWORD_SENTINEL</password></entry>
      <entry name="absent"/><entry name="empty"><auto-retrieve-passcode/></entry></client-auth>
      </entry></global-protect-gateway></global-protect>
      </entry></vsys></entry></devices></config>"""
    client = create_app({"TESTING": True}).test_client()
    response = client.post("/api/preview", data={"source_vendor": "palo_alto", "file": (io.BytesIO(source.encode()), "metadata.xml")})
    assert response.status_code == 200
    preview = response.get_json()
    sanitized = preview["source_evidence"]["source_text"]
    sanitizer = PANOSCollectedSourceSanitizer()
    assert sanitizer.sanitize(sanitized) == sanitized
    root = ET.fromstring(sanitized)
    assert root.findtext(".//certificate/entry/public-key") == "PUBLIC_CERTIFICATE"
    assert root.findtext(".//save-user-credentials") == "1"
    assert root.findtext(".//domain-credentials") == ""
    assert all(node.text == "Password prompt" for node in root.iter("password-label"))
    gateway = preview["sections"]["globalprotect_gateways"][0]
    reported, absent, empty = gateway["client_authentication"]
    assert reported["auto_retrieve_passcode"] == "no"
    assert reported["raw_extra"]["password-label"] == "Password prompt"
    assert reported["additional_settings"]["password-label"] == "Password prompt"
    assert absent["auto_retrieve_passcode"] is None and empty["auto_retrieve_passcode"] is None
    assert "auto_retrieve_passcode" not in absent["explicit_fields"]
    assert "auto_retrieve_passcode" in empty["explicit_fields"]
    assert preview["sections"]["globalprotect_portals"][0]["client_configs"][0]["raw_extra"]["save-user-credentials"] == "1"
    assert preview["sections"]["globalprotect_portals"][0]["client_configs"][0]["additional_settings"]["save-user-credentials"] == "1"
    assert "PUBLIC_CERTIFICATE" in json.dumps(preview["sections"]["source_inventory"])
    response = client.post("/api/extract/excel", json={"source_vendor": "palo_alto", "source": preview["source_evidence"], "excel_profile": profile})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        auth = sheet_rows(workbook, "GP Gateway Client Auth")[0]
        assert auth["Auto Retrieve Passcode"] == "no"
        assert "Password prompt" in auth["Additional Settings"]
        assert "save-user-credentials: 1" in sheet_rows(workbook, "GP Portal Client Configs")[0]["Additional Settings"]
        if profile == "full":
            appendix = sheet_rows(workbook, "PAN-OS Source Appendix")
            assert any(row["Evidence Value"] == "PUBLIC_CERTIFICATE" for row in appendix)
            assert any(str(row["Evidence Path"]).endswith("domain-credentials") and row["Evidence Type"] == "presence" for row in appendix)
        surfaces = json.dumps(preview) + " ".join(str(value) for sheet in workbook for row in sheet.values for value in row)
        for secret in ("PRIVATE_SENTINEL", "OVERRIDE_SENTINEL", "PASSWORD_SENTINEL"):
            assert secret not in surfaces
    finally:
        workbook.close()
