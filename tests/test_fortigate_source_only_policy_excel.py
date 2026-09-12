import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.report import IRExcelExporter


def test_source_only_policy_families_export_from_canonical_ir():
    result = extract_fortigate_config(
        """config router policy
    edit 10
        set src "10.0.0.0/24"
    next
end
config router policy6
    edit 20
        set src "2001:db8::/64"
    next
end
config firewall local-in-policy
    edit 30
        set intf "wan1" "wan2"
        set srcaddr "ADMIN-NET"
        set dstaddr "FGT-IP"
        set service "HTTPS"
        set action accept
    next
end
config firewall local-in-policy6
    edit 40
        set intf "wan6"
        set srcaddr "ADMIN6-NET"
        set dstaddr "FGT6-IP"
        set service "HTTPS"
    next
end
config system settings
    set ngfw-mode policy-based
end
config firewall security-policy
    edit 50
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "SRC"
        set srcaddr6 "SRC6"
        set dstaddr "DST"
        set dstaddr6 "DST6"
        set service "HTTPS"
        set application 12345
        set av-profile "default"
        set ssl-ssh-profile "certificate-inspection"
        set logtraffic all
        set action deny
        set comments "NGFW review"
    next
end
"""
    )
    ir = result.canonical_ir
    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir, result).generate()))

    assert workbook.sheetnames.index("Policies") < workbook.sheetnames.index("Local-In Policies")
    assert workbook.sheetnames.index("Routes") < workbook.sheetnames.index("Policy Routes")
    assert [cell.value for cell in workbook["Policy Routes"][3]][:4] == [
        "Source VDOM", "Rule ID", "Address Family", "Source Order",
    ]
    pbr = workbook["Policy Routes"]
    pbr_headers = {cell.value: cell.column for cell in pbr[3]}
    assert [pbr.cell(row, pbr_headers["Rule ID"]).value for row in range(4, 6)] == ["10", "20"]
    assert pbr.cell(4, pbr_headers["Protocol"]).value is None
    assert pbr.cell(4, pbr_headers["Effective Protocol"]).value == 0
    assert pbr.cell(5, pbr_headers["Effective Source Port Start"]).value == 1

    local = workbook["Local-In Policies"]
    local_headers = {cell.value: cell.column for cell in local[3]}
    assert [local.cell(row, local_headers["Rule ID"]).value for row in range(4, 6)] == ["30", "40"]
    assert local.cell(4, local_headers["Interface"]).value == "wan1\nwan2"
    assert local.cell(5, local_headers["Address Family"]).value == "ipv6"

    ngfw = workbook["NGFW Security Policies"]
    ngfw_headers = {cell.value: cell.column for cell in ngfw[3]}
    assert ngfw.cell(4, ngfw_headers["Rule ID"]).value == "50"
    assert ngfw.cell(4, ngfw_headers["Applications"]).value == "12345"
    assert ngfw.cell(4, ngfw_headers["SSL Inspection Reference"]).value == "certificate-inspection"
    assert ngfw.cell(4, ngfw_headers["Logging"]).value == "all"
    assert ngfw.cell(4, ngfw_headers["Action"]).value == "deny"
    assert ngfw.cell(4, ngfw_headers["Effective Action"]).value == "deny"
    assert ngfw.cell(4, ngfw_headers["Enabled"]).value == "Yes"
    assert ngfw.cell(4, ngfw_headers["Manual Review"]).value == "Yes"

    assert pbr.max_row == 5
    assert local.max_row == 5
    assert ngfw.max_row == 4
    assert workbook["Routes"].max_row == 3
    assert workbook["Policies"].max_row == 3


def test_ngfw_effective_defaults_and_explicit_actions_export():
    result = extract_fortigate_config(
        """config system settings
    set ngfw-mode policy-based
end
config firewall security-policy
    edit 1
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
    next
    edit 2
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
    next
    edit 3
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action deny
    next
end
"""
    )

    policies = result.canonical_ir.security_policies
    assert [policy.enabled for policy in policies] == [True, True, True]
    assert [policy.effective_action for policy in policies] == ["deny", "accept", "deny"]

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate()))
    sheet = workbook["NGFW Security Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert [sheet.cell(row, headers["Action"]).value for row in range(4, 7)] == [None, "accept", "deny"]
    assert [sheet.cell(row, headers["Effective Action"]).value for row in range(4, 7)] == ["deny", "accept", "deny"]


def test_ngfw_webfilter_dependencies_survive_ir_and_excel():
    content = """# config-version = 7.4.6
config webfilter profile
    edit "standard"
    next
end
config system settings
    set ngfw-mode policy-based
end
config firewall security-policy
    edit 50
        set webfilter-profile "standard"
    next
    edit 51
        set webfilter-profile "missing-webfilter"
    next
end
"""
    parsed = parse_fortigate_config(content)
    assert [policy.webfilter_profile for policy in parsed.security_policies] == [
        "standard", "missing-webfilter",
    ]

    result = extract_fortigate_config(content)
    dependencies = {
        (item.source_object, item.reference): item
        for item in result.dependencies
        if item.source_path == "firewall security-policy"
    }
    assert dependencies[("50", "standard")].result == "RESOLVED"
    assert dependencies[("50", "standard")].target_path == "webfilter profile"
    assert dependencies[("51", "missing-webfilter")].result == "UNRESOLVED"
    assert dependencies[("51", "missing-webfilter")].target_path is None

    policies = {
        policy.source_id: policy for policy in result.canonical_ir.security_policies
    }
    assert policies["50"].source_attributes["webfilter_profile"] == "standard"
    assert policies["51"].source_attributes["webfilter_profile"] == "missing-webfilter"
    assert len(result.canonical_ir.security_policies) == 2
    assert any("missing-webfilter" in entry.message for entry in result.canonical_ir.audit_entries)

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate()))
    sheet = workbook["NGFW Security Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    rows = {
        sheet.cell(row, headers["Rule ID"]).value: row
        for row in range(4, sheet.max_row + 1)
    }
    assert "webfilter-profile=standard" in sheet.cell(
        rows["50"], headers["Security Profile References"]
    ).value
    assert "webfilter-profile=missing-webfilter" in sheet.cell(
        rows["51"], headers["Security Profile References"]
    ).value
    assert sheet.cell(rows["51"], headers["Manual Review"]).value == "Yes"


def test_ngfw_typed_fields_survive_parser_ir_and_excel_projection():
    content = """# config-version = 7.4.6
config system settings
    set ngfw-mode policy-based
end
config firewall security-policy
    edit "NGFW-1"
        set status disable
        set action accept
        set srcintf "lan" "lan2"
        set dstintf "wan"
        set srcaddr "SRC1" "SRC2"
        set dstaddr "DST1"
        set srcaddr6 "SRC6"
        set dstaddr6 "DST6"
        set srcaddr-negate enable
        set dstaddr-negate disable
        set srcaddr6-negate enable
        set dstaddr6-negate disable
        set service "HTTPS" "DNS"
        set service-negate enable
        set schedule "business-hours"
        set users "alice" "bob"
        set groups "engineering"
        set fsso-groups "ad-users"
        set profile-type single
        set profile-group "default"
        set profile-protocol-options "default"
        set internet-service enable
        set internet-service-negate disable
        set internet-service-custom "CUSTOM-IS"
        set internet-service-custom-group "CUSTOM-GROUP"
        set internet-service-group "IS-GROUP"
        set internet-service-name "Microsoft-Office365"
        set internet-service-src enable
        set internet-service-src-negate enable
        set internet-service-src-custom "SRC-CUSTOM-IS"
        set internet-service-src-custom-group "SRC-CUSTOM-GROUP"
        set internet-service-src-group "SRC-IS-GROUP"
        set internet-service-src-name "Source-Service"
        set internet-service6 enable
        set internet-service6-negate disable
        set internet-service6-custom "CUSTOM-IS6"
        set internet-service6-custom-group "CUSTOM-GROUP6"
        set internet-service6-group "IS-GROUP6"
        set internet-service6-name "IPv6-Service"
        set internet-service6-src enable
        set internet-service6-src-negate disable
        set internet-service6-src-custom "SRC-CUSTOM-IS6"
        set internet-service6-src-custom-group "SRC-CUSTOM-GROUP6"
        set internet-service6-src-group "SRC-IS-GROUP6"
        set internet-service6-src-name "IPv6-Source-Service"
        set future-ngfw-setting "retained"
    next
end
"""

    parsed = parse_fortigate_config(content)
    policy = parsed.security_policies[0]
    assert policy.name == "NGFW-1"
    assert policy.schedule == "business-hours"
    assert policy.users == ["alice", "bob"]
    assert policy.groups == ["engineering"]
    assert policy.fsso_groups == ["ad-users"]
    assert policy.srcaddr_negate == "enable"
    assert policy.profile_protocol_options == "default"
    assert policy.internet_service == "enable"
    assert policy.internet_service_src_name == ["Source-Service"]
    assert policy.internet_service6 == "enable"
    assert policy.internet_service6_src_name == ["IPv6-Source-Service"]
    assert policy.extra_settings["future_ngfw_setting"] == "retained"

    result = extract_fortigate_config(content)
    assert len(result.canonical_ir.security_policies) == 1
    ir_policy = result.canonical_ir.security_policies[0]
    assert ir_policy.source_attributes["users"] == ["alice", "bob"]
    assert ir_policy.source_attributes["internet_service6_src_group"] == ["SRC-IS-GROUP6"]
    assert ir_policy.source_attributes["future_ngfw_setting"] == "retained"

    workbook = load_workbook(io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate()))
    sheet = workbook["NGFW Security Policies"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Name"]).value == "NGFW-1"
    assert sheet.cell(4, headers["Schedule"]).value == "business-hours"
    assert sheet.cell(4, headers["Users"]).value == "alice\nbob"
    assert sheet.cell(4, headers["Groups"]).value == "engineering"
    assert sheet.cell(4, headers["FSSO Groups"]).value == "ad-users"
    assert sheet.cell(4, headers["Source Address Negate"]).value == "enable"
    assert sheet.cell(4, headers["Profile Protocol Options"]).value == "default"
    assert sheet.cell(4, headers["Internet Service Enable/Status"]).value == "enable"
    assert sheet.cell(4, headers["Internet Service Source Names"]).value == "Source-Service"
    assert sheet.cell(4, headers["IPv6 Internet Service Source Names"]).value == "IPv6-Source-Service"
    assert "future-ngfw-setting=retained" in sheet.cell(4, headers["Additional Settings"]).value
