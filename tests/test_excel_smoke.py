import io

from openpyxl import load_workbook

from fwmigrate.ir import IRAuditEntry, IRConfig, IRInterface, IRMetadata
from fwmigrate.ir.enums import MigrationConfidence
from fwmigrate.report import IRExcelExporter


def test_excel_workbook_contains_review_evidence_and_no_secrets():
    secret = "do-not-export-this-secret"
    ir = IRConfig(
        metadata=IRMetadata(hostname="review-test", source_vendor="fortigate"),
        interfaces=[IRInterface(
            name="wan1",
            has_pppoe_password=True,
            pppoe_password_format="encrypted",
            source_attributes={"password": secret},
            requires_manual_review=True,
            review_reasons=["source setting needs review"],
        )],
        audit_entries=[IRAuditEntry(
            id="review-1",
            category="Parser Review",
            message="source setting needs review",
            confidence=MigrationConfidence.MANUAL,
        )],
    )
    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir).generate()))
    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )

    assert "source setting needs review" in text
    assert "Configured / Redacted" in text or "Password Configured" in text
    assert secret not in text
