from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


CONFIG = """
config firewall internet-service-extension
    edit 100
        set comment "Extension"
        config disable-entry
            edit 1
                set addr-mode ipv4
                config ip-range
                    edit 1
                        set start-ip 192.0.2.1
                        set end-ip 192.0.2.2
                        set future-range retained
                    next
                    edit 2
                        set start-ip 192.0.2.10
                        set end-ip 192.0.2.20
                    next
                end
                config ip6-range
                    edit 3
                        set start-ip6 2001:db8::1
                        set end-ip6 2001:db8::ffff
                    next
                end
                set protocol 6
            next
        end
        config entry
            edit 2
                set addr-mode ipv6
                set dst6 2001:db8::/32
                set protocol 17
                config port-range
                    edit 4
                        set start-port 443
                        set end-port 8443
                    next
                end
            next
        end
    next
end
"""


def _extension_config(child: str, child_id: int, *settings: str) -> str:
    return f'''config firewall internet-service-extension
    edit 100
        config {child}
            edit {child_id}
                {' '.join(settings)}
            next
        end
    next
end
'''


def test_extension_entry_ids_use_uint8_range_and_preserve_invalid_values():
    for entry_id, valid in ((0, True), (255, True), (256, False), (4294967295, False)):
        entry = parse_fortigate_config(_extension_config("entry", entry_id)).internet_service_extensions[0].entries[0]
        assert (entry.id is not None) is valid
        if not valid:
            assert entry.extra_settings["invalid_fields"]["id"] == entry_id


def test_extension_nested_ids_remain_uint32():
    config = '''config firewall internet-service-extension
    edit 100
        config disable-entry
            edit 4294967295
                config ip-range
                    edit 4294967295
                    next
                end
                config ip6-range
                    edit 4294967295
                    next
                end
                config port-range
                    edit 4294967295
                    next
                end
            next
        end
    next
end
'''
    extension = parse_fortigate_config(config).internet_service_extensions[0]
    assert extension.disable_entries[0].id == 4294967295
    assert extension.disable_entries[0].ip_range[0].id == 4294967295
    assert extension.disable_entries[0].ip6_range[0].id == 4294967295
    assert extension.disable_entries[0].port_ranges[0].id == 4294967295


def test_extension_child_addr_modes_reject_both_and_preserve_source_value():
    for child in ("disable-entry", "entry"):
        for mode, valid in (("ipv4", True), ("ipv6", True), ("both", False), ("future", False)):
            item = parse_fortigate_config(
                _extension_config(child, 1, f"set addr-mode {mode}")
            ).internet_service_extensions[0]
            typed = getattr(item, "disable_entries" if child == "disable-entry" else "entries")[0]
            assert (typed.addr_mode is not None) is valid
            if not valid:
                assert typed.extra_settings["invalid_fields"]["addr_mode"] == mode


def test_invalid_extension_values_remain_in_ir_and_block_generation():
    result = extract_fortigate_config(
        _extension_config("entry", 256, "set addr-mode both")
    )
    entry = result.canonical_ir.internet_service_extensions[0].entries[0]
    assert entry.source_id is None
    assert entry.addr_mode is None
    assert entry.source_attributes["invalid_fields"] == {"id": 256, "addr_mode": "both"}
    assert result.generation_safe is False


def test_extension_keeps_typed_disable_ranges_and_additional_entries_separate():
    fg = parse_fortigate_config(CONFIG)
    item = fg.internet_service_extensions[0]
    disable_entry = item.disable_entries[0]

    assert len(disable_entry.ip_range) == 2
    assert [(item.id, item.start_ip, item.end_ip) for item in disable_entry.ip_range] == [
        (1, "192.0.2.1", "192.0.2.2"),
        (2, "192.0.2.10", "192.0.2.20"),
    ]
    assert disable_entry.ip_range[0].extra_settings == {"future_range": "retained"}
    assert [(item.id, item.start_ip6, item.end_ip6) for item in disable_entry.ip6_range] == [
        (3, "2001:db8::1", "2001:db8::ffff"),
    ]
    assert len(item.entries) == 1
    assert item.entries[0].port_ranges[0].start_port == 443

    ir = FGToIRTransformer(fg).transform().internet_service_extensions[0]
    disable_ir = ir.disable_entries[0]
    assert [(item.source_id, item.start_ip, item.end_ip) for item in disable_ir.ipv4_ranges] == [
        (1, "192.0.2.1", "192.0.2.2"),
        (2, "192.0.2.10", "192.0.2.20"),
    ]
    assert disable_ir.ipv4_ranges[0].source_attributes == {"future_range": "retained"}
    assert [(item.source_id, item.start_ip6, item.end_ip6) for item in disable_ir.ipv6_ranges] == [
        (3, "2001:db8::1", "2001:db8::ffff"),
    ]
    assert ir.entries[0].destination_ipv6 == ["2001:db8::/32"]

    workbook = load_workbook(BytesIO(IRExcelExporter(FGToIRTransformer(fg).transform()).generate()))
    sheet = workbook["IS Extension Disabled"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["IPv4 Ranges"]).value == (
        "1: 192.0.2.1 - 192.0.2.2, 2: 192.0.2.10 - 192.0.2.20"
    )
    assert sheet.cell(4, headers["IPv6 Ranges"]).value == "3: 2001:db8::1 - 2001:db8::ffff"


def test_invalid_ports_are_retained_and_block_unsafe_generation():
    result = extract_fortigate_config("""
config firewall internet-service-extension
    edit 100
        config entry
            edit 1
                config port-range
                    edit 1
                        set start-port 500
                        set end-port 400
                    next
                end
            next
        end
    next
end
""")
    port = result.canonical_ir.internet_service_extensions[0].entries[0].port_ranges[0]
    assert (port.start_port, port.end_port) == (500, 400)
    assert result.generation_safe is False


def test_coverage_counts_append_sections_and_nested_ranges_as_objects():
    result = extract_fortigate_config("""
config firewall internet-service-append
    set addr-mode ipv4
    set append-port 8080
end
config firewall internet-service-extension
    edit 100
        config disable-entry
            edit 1
                config ip-range
                    edit 1
                        set start-ip 192.0.2.1
                        set end-ip 192.0.2.2
                    next
                end
                config ip6-range
                    edit 2
                        set start-ip6 2001:db8::1
                        set end-ip6 2001:db8::2
                    next
                end
            next
        end
    next
end
""")
    sections = {section.path: section for section in result.source_sections}

    assert sections["firewall internet-service-append"].object_count_parsed == 1
    assert sections["firewall internet-service-extension disable-entry ip-range"].object_count_parsed == 1
    assert sections["firewall internet-service-extension disable-entry ip6-range"].object_count_parsed == 1
