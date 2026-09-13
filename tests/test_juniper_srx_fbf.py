from io import BytesIO

from openpyxl import load_workbook

from fwmigrate.parsers.juniper_srx import JuniperSRXParser
from fwmigrate.report.excel_exporter import IRExcelExporter


def test_fbf_actions_project_to_policy_routes_with_order_and_dependencies():
    ir = JuniperSRXParser("""
    set interfaces ge-0/0/0 unit 0 family inet filter input F
    set interfaces ge-0/0/1 unit 0 family inet6 filter output F6
    set firewall family inet filter F term ri from source-address 10.0.0.0/8
    set firewall family inet filter F term ri then routing-instance RI
    set firewall family inet filter F term iface then next-interface ge-0/0/1.0
    set firewall family inet filter F term ip then next-ip 192.0.2.1
    set firewall family inet6 filter F6 term ip6 then next-ip6 2001:db8::1
    set routing-instances RI instance-type virtual-router
    """).transform_to_ir()

    rules = {rule.source_rule_id: rule for rule in ir.policy_route_rules}
    assert rules["F:ri"].action == "routing-instance"
    assert rules["F:iface"].next_interface == "ge-0/0/1.0"
    assert rules["F:ip"].next_hop == "192.0.2.1"
    assert "F6:ip6" not in rules
    assert rules["F:ri"].source_order < rules["F:iface"].source_order < rules["F:ip"].source_order

    workbook = load_workbook(BytesIO(IRExcelExporter(ir).generate()))
    sheet = workbook["Cisco PBR"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(row for row in sheet.iter_rows(min_row=4, values_only=False) if row[headers["Source Rule ID"] - 1].value == "F:iface")
    assert row[headers["Next Interface"] - 1].value == "ge-0/0/1.0"


def test_unresolved_fbf_references_require_review():
    ir = JuniperSRXParser("""
    set interfaces ge-0/0/0 unit 0 family inet filter input F
    set firewall family inet filter F term ri then routing-instance MISSING
    set firewall family inet filter F term iface then next-interface ge-0/0/9.0
    """).transform_to_ir()
    assert all(rule.requires_manual_review for rule in ir.policy_route_rules)
    assert any("routing-instance" in reason for reason in ir.policy_route_rules[0].review_reasons)
    assert any("next-interface" in reason for reason in ir.policy_route_rules[1].review_reasons)
