import io

import pytest
from openpyxl import load_workbook

from fwmigrate.generators.palo_alto.transformer import IRToPANOSTransformer
from fwmigrate.generators.palo_alto.terraform_generator import PANOSTerraformGenerator
from fwmigrate.generators.palo_alto.xml_generator import PANOSXMLGenerator
from fwmigrate.ir.enums import MigrationConfidence, NATTranslationMode, NATType, ServiceProtocol
from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.parsers.fortigate.parser import parse_fortigate_config
from fwmigrate.parsers.fortigate.transformer import FGToIRTransformer
from fwmigrate.report.excel_exporter import IRExcelExporter


POLICY_BASE = """
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
"""

INTERFACES = """
config system interface
    edit "LAN"
        set role lan
        set ip 10.0.0.1 255.255.255.0
    next
    edit "WAN"
        set role wan
        set ip 203.0.113.10 255.255.255.0
    next
end
config system zone
    edit "LAN"
        set interface "LAN"
    next
    edit "WAN"
        set interface "WAN"
    next
end
"""


def _transform(config: str):
    return FGToIRTransformer(parse_fortigate_config(INTERFACES + config)).transform()


def _transform_raw(config: str):
    return FGToIRTransformer(parse_fortigate_config(config)).transform()


def _main_tf(ir):
    return next(
        artifact.content
        for artifact in PANOSTerraformGenerator().generate(ir)
        if artifact.filename == "main.tf"
    )


def test_interface_address_snat_is_correlated_to_policy():
    ir = _transform(f"""
config firewall policy
    edit 10
{POLICY_BASE}
        set nat enable
    next
end
""")

    assert len(ir.nat_rules) == 1
    rule = ir.nat_rules[0]
    assert rule.type == NATType.SOURCE
    assert rule.source_policy_reference == "10"
    assert rule.source_translation_mode == NATTranslationMode.INTERFACE_ADDRESS
    assert rule.source_from_interfaces == ["LAN"]
    assert rule.source_to_interfaces == ["WAN"]
    assert rule.source == ["LAN_NET"]
    assert rule.source_pool_references == []
    assert rule.translated_sources == ["203.0.113.10"]
    assert rule.requires_manual_review is False
    xml = PANOSXMLGenerator().generate(ir)[0].content
    assert "<interface-address>" in xml
    assert "<interface>WAN</interface>" in xml
    hcl = _main_tf(ir)
    assert "original_packet {" in hcl
    assert "translated_packet {" in hcl
    assert "interface_address {" in hcl
    assert 'interface = "WAN"' in hcl


def test_nat46_uses_ipv6_pool_and_vip_translation_fields_end_to_end():
    config = f"""
config firewall vip
    edit "NAT46_VIP"
        set extip 10.1.100.150
        set nat44 disable
        set nat46 enable
        set extintf "WAN"
        set arp-reply enable
        set ipv6-mappedip 2001:db8:200::156
        set portforward enable
        set protocol tcp
        set extport 443
        set ipv6-mappedport 8443
    next
end
config firewall ippool6
    edit "NAT46_POOL6"
        set startip 2001:db8:101::1
        set endip 2001:db8:101::1
        set nat46 enable
        set add-nat46-route enable
    next
end
config firewall policy
    edit 200
        set name "NAT46_POLICY"
        set srcintf "LAN"
        set dstintf "WAN"
        set action accept
        set nat46 enable
        set srcaddr "all"
        set dstaddr "NAT46_VIP"
        set srcaddr6 "all"
        set dstaddr6 "all"
        set schedule "always"
        set service "ALL"
        set ippool enable
        set poolname6 "NAT46_POOL6"
    next
end
    """

    parsed = parse_fortigate_config(INTERFACES + config)
    assert parsed.vips[0].ipv6_mappedip == "2001:db8:200::156"
    assert parsed.policies[0].poolname6 == ["NAT46_POOL6"]

    result = extract_fortigate_config(INTERFACES + config)
    rules = [
        rule for rule in result.canonical_ir.nat_rules
        if rule.source_policy_reference == "200"
    ]
    assert len(rules) == 1
    rule = rules[0]
    assert rule.nat_family.value == "nat46"
    assert (rule.original_address_family, rule.translated_address_family) == (
        "ipv4",
        "ipv6",
    )
    assert rule.destination == ["10.1.100.150"]
    assert rule.translated_destinations == ["2001:db8:200::156"]
    assert rule.translated_sources == []
    assert rule.source_pool_references == ["NAT46_POOL6"]
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.requires_manual_review is True
    assert rule.source_vip_reference == "NAT46_VIP"
    assert rule.source_vip_group_reference is None
    assert rule.original_destination_ports[0].start == 443
    assert rule.translated_destination_ports[0].start == 8443
    assert any(
        dependency.reference == "NAT46_POOL6"
        and dependency.expected_type == "firewall ippool6"
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )
    assert any(
        dependency.reference == "NAT46_VIP"
        and dependency.target_path in {"firewall vip", "firewall vipgrp"}
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    row = next(
        row for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Name"]).value == rule.name
    )
    assert sheet.cell(row, headers["Original Destination"]).value == "10.1.100.150"
    assert sheet.cell(row, headers["Translated Destination"]).value == "2001:db8:200::156"
    assert sheet.cell(row, headers["Translated Source"]).value is None
    assert sheet.cell(row, headers["NAT Family"]).value == "nat46"
    assert sheet.cell(row, headers["IP Pool"]).value == "NAT46_POOL6"
    assert sheet.cell(row, headers["VIP"]).value == "NAT46_VIP"
    assert sheet.cell(row, headers["Original Destination Port"]).value == "443"
    assert sheet.cell(row, headers["Translated Destination Port"]).value == "8443"


def test_nat46_vip_group_preserves_member_and_group_references():
    result = extract_fortigate_config('''
config firewall vip
    edit "NAT46_GROUP_MEMBER"
        set extip 10.1.100.151
        set nat44 disable
        set nat46 enable
        set ipv6-mappedip 2001:db8:200::157
    next
end
config firewall vipgrp
    edit "NAT46_GROUP"
        set member "NAT46_GROUP_MEMBER"
    next
end
config firewall ippool6
    edit "NAT46_GROUP_POOL6"
        set startip 2001:db8:101::2
        set endip 2001:db8:101::2
        set nat46 enable
    next
end
config firewall policy
    edit 203
        set action accept
        set nat46 enable
        set srcaddr "all"
        set dstaddr "NAT46_GROUP"
        set srcaddr6 "all"
        set dstaddr6 "all"
        set service "ALL"
        set ippool enable
        set poolname6 "NAT46_GROUP_POOL6"
    next
end
''')

    rules = [rule for rule in result.canonical_ir.nat_rules if rule.source_policy_reference == "203"]
    assert len(rules) == 1
    rule = rules[0]
    assert rule.source_vip_reference == "NAT46_GROUP_MEMBER"
    assert rule.source_vip_group_reference == "NAT46_GROUP"
    assert rule.translated_destinations == ["2001:db8:200::157"]
    assert any(
        dependency.reference == "NAT46_GROUP"
        and dependency.target_path == "firewall vipgrp"
        and dependency.result == "RESOLVED"
        for dependency in result.dependencies
    )


@pytest.mark.parametrize(
    ("family", "policy_fields", "expected"),
    [
        ("NAT44", 'set srcaddr "all"\n        set dstaddr "all"', ("nat44", "ipv4", "ipv4")),
        ("NAT46", 'set srcaddr "all"\n        set dstaddr "all"\n        set nat46 enable', ("nat46", "ipv4", "ipv6")),
        ("NAT64", 'set srcaddr6 "all"\n        set dstaddr6 "all"\n        set nat64 enable', ("nat64", "ipv6", "ipv4")),
        ("NAT66", 'set srcaddr6 "all"\n        set dstaddr6 "all"', ("nat66", "ipv6", "ipv6")),
    ],
)
def test_policy_nat_family_direction_contract(family, policy_fields, expected):
    ir = _transform(f"""
config firewall policy
    edit 201
        set srcintf "LAN"
        set dstintf "WAN"
        set action accept
        set service "ALL"
        set nat enable
        {policy_fields}
    next
end
""")

    rules = [rule for rule in ir.nat_rules if rule.source_policy_reference == "201"]
    assert len(rules) == 1, family
    rule = rules[0]
    assert (rule.nat_family.value, rule.original_address_family, rule.translated_address_family) == expected


def test_nat_with_unzoned_interfaces_preserves_evidence_and_requires_review():
    ir = _transform_raw("""
config system interface
    edit "port1"
        set ip 10.0.0.1 255.255.255.0
    next
    edit "port2"
        set ip 203.0.113.10 255.255.255.0
    next
end
config firewall policy
    edit 11
        set srcintf "port1"
        set dstintf "port2"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_from_interfaces == ["port1"]
    assert rule.source_to_interfaces == ["port2"]
    assert rule.from_zone == []
    assert rule.to_zone == []
    assert rule.requires_manual_review is True
    assert any(
        "unresolved canonical zones" in entry.message
        for entry in ir.audit_entries
    )


def test_advanced_pool_semantics_survive_and_withhold_correlated_nat():
    ir = _transform(f"""
config firewall ippool
    edit "ADVANCED_POOL"
        set type port-block-allocation
        set startip 203.0.113.20
        set endip 203.0.113.30
        set source-startip 10.0.0.10
        set source-endip 10.0.0.20
        set startport 5117
        set endport 65533
        set exclude-ip "203.0.113.25" "203.0.113.26"
        set permit-any-host enable
        set block-size 128
        set num-blocks-per-user 4
        set pba-timeout 60
        set pba-interim-log 600
        set nat64 enable
        set cgn-block-size 256
        set cgn-client-startip 10.0.0.10
        set cgn-client-endip 10.0.0.20
        set cgn-port-start 1024
        set cgn-port-end 65535
        set utilization-alarm-clear 70
        set utilization-alarm-raise 90
        set future-pool-setting "retained"
    next
end
config firewall policy
    edit 100
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "ADVANCED_POOL"
    next
end
""")

    pool = ir.ip_pools[0]
    assert pool.pool_type == "port-block-allocation"
    assert pool.source_start_ip == "10.0.0.10"
    assert pool.source_end_ip == "10.0.0.20"
    assert pool.excluded_ips == ["203.0.113.25", "203.0.113.26"]
    assert pool.permit_any_host is True
    assert (pool.block_size, pool.blocks_per_user, pool.pba_timeout) == (128, 4, 60)
    assert pool.pba_interim_log == 600
    assert pool.nat64 is True
    assert pool.cgn_block_size == 256
    assert pool.cgn_client_start_ip == "10.0.0.10"
    assert pool.cgn_client_end_ip == "10.0.0.20"
    assert (pool.cgn_port_start, pool.cgn_port_end) == (1024, 65535)
    assert (pool.utilization_alarm_clear, pool.utilization_alarm_raise) == (70, 90)
    assert pool.source_attributes == {"future_pool_setting": "retained"}
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True

    rule = ir.nat_rules[0]
    assert rule.source_pool_type == "port-block-allocation"
    assert rule.source_pool_excluded_ips == ["203.0.113.25", "203.0.113.26"]
    assert rule.source_pool_permit_any_host is True
    assert rule.source_pool_original_start_ip == ["203.0.113.20"]
    assert rule.source_pool_original_end_ip == ["203.0.113.30"]
    assert rule.requires_manual_review is True
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.review_reasons
    assert "original_packet {" not in _main_tf(ir)
    assert IRToPANOSTransformer(ir).transform().vsys.nat_rules == []


def test_explicit_pba_settings_on_basic_pool_withhold_correlated_nat():
    ir = _transform(f"""
config firewall ippool
    edit "PBA_SETTINGS"
        set type overload
        set startip 203.0.113.31
        set endip 203.0.113.31
        set block-size 128
        set num-blocks-per-user 4
        set pba-timeout 60
    next
end
config firewall policy
    edit 100
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "PBA_SETTINGS"
    next
end
""")

    pool = ir.ip_pools[0]
    rule = ir.nat_rules[0]
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert any("PBA settings" in reason for reason in rule.review_reasons)
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert IRToPANOSTransformer(ir).transform().vsys.nat_rules == []


def test_unknown_pool_semantics_make_correlated_nat_unsafe():
    ir = _transform(f'''
config firewall ippool
    edit "POOL_UNKNOWN"
        set startip 203.0.113.10
        set endip 203.0.113.20
        set future-nat-behavior enable
    next
end
config firewall policy
    edit 101
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "POOL_UNKNOWN"
    next
end
''')

    pool = ir.ip_pools[0]
    rule = ir.nat_rules[0]
    assert pool.migration_status == "PARTIALLY_NORMALIZED"
    assert pool.requires_manual_review is True
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert rule.requires_manual_review is True
    assert any("future_nat_behavior" in reason for reason in rule.review_reasons)


def test_disabled_and_restricted_vip_is_preserved_but_withheld():
    ir = _transform(f"""
config firewall vip
    edit "RESTRICTED_VIP"
        set status disable
        set type server-load-balance
        set extip 203.0.113.80
        set mappedip "10.0.0.80"
        set extintf "WAN"
        set src-filter "10.0.0.0/24"
        set srcintf-filter "WAN"
        set service "HTTPS"
        set nat-source-vip enable
        set portmapping-type m-to-n
    next
end
config firewall policy
    edit 101
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "RESTRICTED_VIP"
        set service "HTTPS"
        set action accept
    next
end
""")

    vip = ir.virtual_ips[0]
    assert vip.vip_type == "server-load-balance"
    assert vip.enabled is False
    assert vip.source_filters == ["10.0.0.0/24"]
    assert vip.source_interface_filters == ["WAN"]
    assert vip.services == ["HTTPS"]
    assert vip.nat_source_vip is True
    assert vip.requires_manual_review is True

    assert ir.nat_rules == []
    assert any("canonical DNAT was withheld" in entry.message for entry in ir.audit_entries)


def test_policy_nat_controls_are_preserved_as_canonical_runtime_behavior():
    ir = _transform(f"""
config firewall policy
    edit 102
{POLICY_BASE}
        set nat enable
        set fixedport enable
        set nat46 enable
        set nat64 enable
        set natip "198.51.100.10 198.51.100.20"
        set natinbound enable
        set natoutbound enable
        set match-vip enable
        set match-vip-only enable
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_policy_fixed_port == "enable"
    assert rule.source_policy_nat46 == "enable"
    assert rule.source_policy_nat64 == "enable"
    assert rule.source_policy_nat_ip == "198.51.100.10 198.51.100.20"
    assert rule.source_policy_nat_inbound == "enable"
    assert rule.source_policy_nat_outbound == "enable"
    assert rule.source_policy_match_vip == "enable"
    assert rule.source_policy_match_vip_only == "enable"
    assert rule.requires_manual_review is True
    assert rule.nat_family.value == "nat46"
    assert rule.source_port_behavior.value == "preserve-strict"
    assert rule.runtime_behavior.fixed_port is True
    assert rule.runtime_behavior.nat_inbound is True
    assert rule.runtime_behavior.nat_outbound is True
    assert rule.runtime_behavior.nat_ip == "198.51.100.10 198.51.100.20"
    assert any("only valid for policy-based IPsec" in reason for reason in rule.review_reasons)
    assert "original_packet {" not in _main_tf(ir)


def test_policy_ipsec_nat_directions_are_separate_and_withheld_from_panos():
    ir = _transform("""
config firewall policy
    edit 200
        set name "IPSEC_POLICY"
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "REMOTE_NET"
        set service "ALL"
        set action ipsec
        set vpntunnel "IPSEC_TUNNEL"
        set natip 198.51.100.25
        set natinbound enable
        set natoutbound enable
    next
end
""")

    rules = [rule for rule in ir.nat_rules if rule.source_policy_reference == "200"]
    assert {rule.name for rule in rules} == {
        "IPSEC-DNAT-P200", "IPSEC-SNAT-P200",
    }
    inbound = next(rule for rule in rules if rule.type == NATType.DESTINATION)
    outbound = next(rule for rule in rules if rule.type == NATType.SOURCE)
    assert inbound.type == NATType.DESTINATION
    assert inbound.translated_destinations == []
    assert inbound.runtime_behavior.nat_inbound is True
    assert inbound.runtime_behavior.nat_outbound is True
    assert outbound.type == NATType.SOURCE
    assert outbound.translated_sources == ["198.51.100.25"]
    assert outbound.source_translation_mode == NATTranslationMode.STATIC
    assert outbound.runtime_behavior.nat_inbound is True
    assert outbound.runtime_behavior.nat_outbound is True
    assert all(
        rule.source_origin == "firewall-policy-ipsec"
        and rule.migration_status == "PARTIALLY_NORMALIZED"
        and rule.requires_manual_review
        for rule in rules
    )
    assert all(rule.source_attributes["vpntunnel"] == "IPSEC_TUNNEL" for rule in rules)
    assert IRToPANOSTransformer(ir).transform().vsys.nat_rules == []
    assert "original_packet {" not in _main_tf(ir)


def test_nat_correlation_uses_policy_identity_when_ir_policies_are_skipped_and_reordered():
    config = f"""
config firewall ippool
    edit "POOL20"
        set startip 203.0.113.20
        set endip 203.0.113.20
    next
    edit "POOL30"
        set startip 203.0.113.30
        set endip 203.0.113.30
    next
end
config firewall policy
    edit 10
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
    edit 20
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
        set ippool enable
        set poolname "POOL20"
    next
    edit 30
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "all"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
        set ippool enable
        set poolname "POOL30"
    next
end
"""
    transformer = FGToIRTransformer(parse_fortigate_config(INTERFACES + config))
    original_transform_policies = transformer._transform_policies

    def transform_policies_with_gap():
        original_transform_policies()
        transformer.ir.policies = [
            policy
            for policy in reversed(transformer.ir.policies)
            if policy.source_rule_id != "20"
        ]

    transformer._transform_policies = transform_policies_with_gap
    ir = transformer.transform()

    rules = {
        rule.source_policy_reference: rule
        for rule in ir.nat_rules
        if rule.source_origin == "firewall-policy"
    }
    assert set(rules) == {"10", "30"}
    assert rules["30"].source_pool_references == ["POOL30"]
    assert rules["30"].translated_sources == ["203.0.113.30"]
    assert not any("POOL20" in rule.translated_sources for rule in rules.values())
    assert any("policy 20" in entry.message.lower() for entry in ir.audit_entries)


@pytest.mark.parametrize("action", ["accept", "deny"])
def test_policy_ipsec_nat_controls_on_normal_policy_are_reviewed_without_ipsec_rows(action):
    ir = _transform(f"""
config firewall policy
    edit 201
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action {action}
        set nat enable
        set natoutbound enable
        set natip 198.51.100.26
    next
end
""")

    rules = [rule for rule in ir.nat_rules if rule.source_policy_reference == "201"]
    assert len(rules) == 1
    rule = rules[0]
    assert rule.name == "SNAT-P201"
    assert rule.source_origin == "firewall-policy"
    assert rule.source_policy_nat_outbound == "enable"
    assert rule.source_policy_nat_ip == "198.51.100.26"
    assert any("only valid for policy-based IPsec" in reason for reason in rule.review_reasons)


@pytest.mark.parametrize(
    ("natip", "expected_sources", "expected_mode"),
    [
        ("198.51.100.27", ["198.51.100.27"], NATTranslationMode.STATIC),
        ('"198.51.100.27 198.51.100.28"', [], None),
        ("0.0.0.0", [], None),
        ("198.51.100.999", [], None),
    ],
)
def test_policy_ipsec_natip_is_conservative(natip, expected_sources, expected_mode):
    ir = _transform(f"""
config firewall policy
    edit 202
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "REMOTE_NET"
        set service "ALL"
        set action ipsec
        set vpntunnel "IPSEC_TUNNEL"
        set natoutbound enable
        set natip {natip}
    next
end
""")

    rule = next(rule for rule in ir.nat_rules if rule.source_policy_reference == "202")
    assert rule.translated_sources == expected_sources
    assert rule.source_translation_mode == expected_mode
    assert rule.source_attributes["natip"] == natip.strip('"')
    assert rule.requires_manual_review is True


def test_policy_ipsec_nat_correlates_disagreeing_phase2_use_natip_values():
    ir = _transform("""
config vpn ipsec phase1
    edit "IPSEC_TUNNEL"
    next
end
config vpn ipsec phase2
    edit 1
        set name "PHASE2_ENABLE"
        set phase1name "IPSEC_TUNNEL"
        set use-natip enable
    next
    edit 2
        set name "PHASE2_DISABLE"
        set phase1name "IPSEC_TUNNEL"
        set use-natip disable
    next
end
config firewall policy
    edit 203
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "REMOTE_NET"
        set service "ALL"
        set action ipsec
        set vpntunnel "IPSEC_TUNNEL"
        set natoutbound enable
        set natip 198.51.100.29
    next
end
""")

    rule = next(rule for rule in ir.nat_rules if rule.source_policy_reference == "203")
    evidence = rule.source_attributes["ipsec_phase2"]
    assert {item["name"] for item in evidence} == {
        "PHASE2_ENABLE", "PHASE2_DISABLE",
    }
    assert {item["use_natip"] for item in evidence} == {"enable", "disable"}
    assert any("disagree on use-natip" in reason for reason in rule.review_reasons)


def test_policy_ipsec_nat_preserves_unresolved_tunnel_reference_and_withholds():
    ir = _transform("""
config firewall policy
    edit 204
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "REMOTE_NET"
        set service "ALL"
        set action ipsec
        set vpntunnel "MISSING_TUNNEL"
        set natoutbound enable
    next
end
""")

    rule = next(rule for rule in ir.nat_rules if rule.source_policy_reference == "204")
    assert rule.source_attributes["vpntunnel"] == "MISSING_TUNNEL"
    assert rule.source_attributes["ipsec_phase2"] == []
    assert rule.translated_sources == []
    assert rule.requires_manual_review is True
    assert any("missing Phase 1" in reason for reason in rule.review_reasons)
    assert any("no matching Phase 2" in reason for reason in rule.review_reasons)


def test_policy_nat_uses_effective_port_preserve_with_fixedport_precedence():
    cases = (
        ("", "preserve-if-available"),
        ("        set port-preserve enable\n", "preserve-if-available"),
        ("        set port-preserve disable\n", "dynamic"),
        ("        set fixedport enable\n        set port-preserve disable\n", "preserve-strict"),
    )
    for settings, expected in cases:
        ir = _transform(f"""
config firewall policy
    edit 104
{POLICY_BASE}
        set nat enable
{settings}    next
end
""")

        assert ir.nat_rules[0].source_port_behavior.value == expected

    result = extract_fortigate_config(f"""
{INTERFACES}
config firewall policy
    edit 105
{POLICY_BASE}
        set nat enable
    next
end
""")
    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["NAT Rules"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Source Port Behavior"]).value == "preserve-if-available"


def test_unknown_policy_port_preserve_is_preserved_and_withheld():
    ir = _transform(f"""
config firewall policy
    edit 106
{POLICY_BASE}
        set nat enable
        set port-preserve future-mode
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_port_behavior is None
    assert rule.requires_manual_review is True
    assert rule.source_attributes == {
        "source_policy_port_preserve": "future-mode",
        "source_policy_effective_port_preserve": "future-mode",
    }
    assert any("port-preserve" in reason for reason in rule.review_reasons)


def test_one_to_one_pool_source_range_survives_and_requires_review():
    ir = _transform(f"""
config firewall ippool
    edit "ONE_TO_ONE"
        set type one-to-one
        set startip 203.0.113.40
        set endip 203.0.113.49
        set source-startip 10.0.0.40
        set source-endip 10.0.0.49
    next
end
config firewall policy
    edit 103
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "ONE_TO_ONE"
    next
end
""")

    pool = ir.ip_pools[0]
    assert (pool.source_start_ip, pool.source_end_ip) == (
        "10.0.0.40", "10.0.0.49"
    )
    rule = ir.nat_rules[0]
    assert rule.source_pool_original_start_ip == ["203.0.113.40"]
    assert rule.source_pool_original_end_ip == ["203.0.113.49"]
    assert rule.requires_manual_review is True
    assert any("source-range" in reason for reason in rule.review_reasons)
    assert "original_packet {" not in _main_tf(ir)
    assert any(
        entry.category == "PAN-OS Terraform NAT"
        and "withheld" in entry.message
        for entry in ir.audit_entries
    )


def test_interface_address_snat_preserves_host_ip_not_network_or_cidr():
    ir = _transform_raw("""
config system interface
    edit "port10"
        set ip 192.168.42.30 255.255.255.0
    next
end
config system zone
    edit "LAN_ZONE"
        set interface "LAN"
    next
    edit "WAN_ZONE"
        set interface "port10"
    next
end
config firewall policy
    edit 11
        set srcintf "LAN"
        set dstintf "port10"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.translated_sources == ["192.168.42.30"]
    assert rule.translated_source == "192.168.42.30"
    assert rule.requires_manual_review is False


def test_dynamic_interface_address_snat_is_unresolved():
    for mode in ("pppoe", "dhcp"):
        ir = _transform_raw(f"""
config system interface
    edit "wan1"
        set mode {mode}
        set ip 198.51.100.99 255.255.255.0
    next
end
config firewall policy
    edit 12
        set srcintf "LAN"
        set dstintf "wan1"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

        rule = ir.nat_rules[0]
        assert rule.translated_sources == []
        assert rule.requires_manual_review is True
        assert any(
            mode in entry.message and "dynamic interface address" in entry.message
            for entry in ir.audit_entries
        )


def test_sdwan_zone_interface_address_snat_does_not_select_a_member():
    ir = _transform_raw("""
config system interface
    edit "wan1"
        set ip 203.0.113.10 255.255.255.0
    next
    edit "wan2"
        set ip 198.51.100.10 255.255.255.0
    next
end
config system sdwan
    set status enable
    config zone
        edit "Internet-Zone"
        next
    end
    config members
        edit 1
            set interface "wan1"
            set zone "Internet-Zone"
        next
        edit 2
            set interface "wan2"
            set zone "Internet-Zone"
        next
    end
end
config firewall policy
    edit 13
        set srcintf "LAN"
        set dstintf "Internet-Zone"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_to_interfaces == ["Internet-Zone"]
    assert rule.translated_sources == []
    assert rule.requires_manual_review is True
    assert any("runtime-selected SD-WAN member" in entry.message for entry in ir.audit_entries)


def test_ambiguous_interface_address_snat_is_unresolved():
    cases = (
        ('set dstintf "wan1" "wan2"', "multiple possible outgoing interfaces"),
        ('set dstintf "any"', "does not identify an egress interface"),
        ('set dstintf "missing"', "was not found"),
    )
    for dstintf, expected_reason in cases:
        ir = _transform_raw(f"""
config system interface
    edit "wan1"
        set ip 203.0.113.10 255.255.255.0
    next
    edit "wan2"
        set ip 198.51.100.10 255.255.255.0
    next
end
config firewall policy
    edit 14
        set srcintf "LAN"
        {dstintf}
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

        rule = ir.nat_rules[0]
        assert rule.translated_sources == []
        assert rule.requires_manual_review is True
        assert any(expected_reason in entry.message for entry in ir.audit_entries)


def test_missing_or_unconfigured_static_interface_ip_is_unresolved():
    for ip_setting in ("", "set ip 0.0.0.0 0.0.0.0"):
        ir = _transform_raw(f"""
config system interface
    edit "wan1"
        {ip_setting}
    next
end
config firewall policy
    edit 15
        set srcintf "LAN"
        set dstintf "wan1"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "ALL"
        set action accept
        set nat enable
    next
end
""")

        rule = ir.nat_rules[0]
        assert rule.translated_sources == []
        assert rule.requires_manual_review is True
        assert any("has no usable static primary IP" in entry.message for entry in ir.audit_entries)


def test_ip_pool_snat_is_correlated_and_inventory_remains():
    ir = _transform(f"""
config firewall ippool
    edit "PUBLIC_POOL"
        set startip 203.0.113.10
        set endip 203.0.113.20
    next
end
config firewall policy
    edit 20
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "PUBLIC_POOL"
    next
end
""")

    assert len(ir.ip_pools) == 1
    assert ir.ip_pools[0].name == "PUBLIC_POOL"
    assert len(ir.nat_rules) == 1
    rule = ir.nat_rules[0]
    assert rule.source_policy_reference == "20"
    assert rule.source_translation_mode == NATTranslationMode.POOL
    assert rule.source_pool_references == ["PUBLIC_POOL"]
    assert rule.source_pool_type == "overload"
    assert rule.translated_sources == ["203.0.113.10-203.0.113.20"]
    xml = PANOSXMLGenerator().generate(ir)[0].content
    assert "<translated-address>" in xml
    assert "<member>203.0.113.10-203.0.113.20</member>" in xml
    hcl = _main_tf(ir)
    assert 'translated_addresses = ["203.0.113.10-203.0.113.20"]' in hcl


def test_unreferenced_ip_pool_does_not_create_nat_rule():
    ir = _transform("""
config firewall ippool
    edit "UNUSED_POOL"
        set startip 203.0.113.10
        set endip 203.0.113.20
    next
end
""")

    assert [pool.name for pool in ir.ip_pools] == ["UNUSED_POOL"]
    assert ir.nat_rules == []


def test_direct_vip_dnat_uses_policy_match_and_vip_translation():
    ir = _transform("""
config firewall vip
    edit "VIP_WEB"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set extintf "WAN"
    next
end
config firewall policy
    edit 30
        set uuid "policy-30-uuid"
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_WEB"
        set service "HTTPS"
        set action accept
    next
end
""")

    assert len(ir.virtual_ips) == 1
    assert len(ir.nat_rules) == 1
    rule = ir.nat_rules[0]
    assert rule.type == NATType.DESTINATION
    assert rule.source_policy_reference == "30"
    assert rule.source_policy_uuid == "policy-30-uuid"
    assert rule.source_vip_reference == "VIP_WEB"
    assert rule.source_from_interfaces == ["WAN"]
    assert rule.source_to_interfaces == ["LAN"]
    assert rule.from_zone == ["WAN"]
    assert rule.to_zone == ["LAN"]
    assert rule.destination == ["198.51.100.10"]
    assert rule.translated_destinations == ["10.0.0.10"]


def test_policy_vip_extintf_mismatch_is_reviewed_without_changing_policy_egress():
    ir = _transform("""
config firewall vip
    edit "VIP_WEB"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set extintf "LAN"
    next
end
config firewall policy
    edit 31
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_WEB"
        set service "HTTPS"
        set action accept
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_from_interfaces == ["WAN"]
    assert rule.source_to_interfaces == ["LAN"]
    assert rule.from_zone == ["WAN"]
    assert rule.to_zone == ["LAN"]
    assert rule.requires_manual_review is True
    assert any("conflicts with the policy source interface" in reason for reason in rule.review_reasons)


def test_unreferenced_vip_does_not_create_nat_rule():
    ir = _transform("""
config firewall vip
    edit "UNUSED_VIP"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
    next
end
""")

    assert [vip.name for vip in ir.virtual_ips] == ["UNUSED_VIP"]
    assert ir.nat_rules == []


def test_vip_group_expands_deterministically_and_preserves_group_reference():
    ir = _transform("""
config firewall vip
    edit "VIP_A"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
    next
    edit "VIP_B"
        set extip 198.51.100.11
        set mappedip "10.0.0.11"
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set interface "any"
        set member "VIP_A" "VIP_B"
    next
end
config firewall policy
    edit 40
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "ALL"
        set action accept
    next
end
""")

    assert [rule.name for rule in ir.nat_rules] == ["DNAT-P40-VIP_A", "DNAT-P40-VIP_B"]
    assert [rule.source_vip_reference for rule in ir.nat_rules] == ["VIP_A", "VIP_B"]
    assert all(rule.source_vip_group_reference == "VIP_GROUP" for rule in ir.nat_rules)
    assert all(rule.source_policy_reference == "40" for rule in ir.nat_rules)


def test_vip_group_port_forward_preserves_inventory_dependencies_nat_and_excel():
    result = extract_fortigate_config(INTERFACES + """
config firewall vip
    edit "VIP_PAT"
        set extip 198.51.100.40
        set mappedip "10.0.0.40"
        set extintf "WAN"
        set portforward enable
        set protocol tcp
        set extport 8443-8444
        set mappedport 443-444
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set interface "WAN"
        set member "VIP_PAT"
    next
end
config firewall policy
    edit 40
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "ALL"
        set action accept
    next
end
""")

    vip = result.canonical_ir.virtual_ips[0]
    assert (vip.external_ip, vip.mapped_ips) == ("198.51.100.40", ["10.0.0.40"])
    assert (vip.external_port, vip.mapped_port) == ("8443-8444", "443-444")
    assert result.canonical_ir.virtual_ip_groups[0].members == ["VIP_PAT"]

    dependencies = [
        item for item in result.dependencies
        if item.source_path in {"firewall vipgrp", "firewall policy"}
    ]
    assert any(
        item.source_path == "firewall vipgrp"
        and item.source_field == "member"
        and item.reference == "VIP_PAT"
        and item.result == "RESOLVED"
        for item in dependencies
    )
    assert any(
        item.source_path == "firewall policy"
        and item.source_field == "dstaddr"
        and item.reference == "VIP_GROUP"
        and item.result == "RESOLVED"
        for item in dependencies
    )

    nat = next(rule for rule in result.canonical_ir.nat_rules if rule.source_vip_reference == "VIP_PAT")
    assert nat.source_vip_group_reference == "VIP_GROUP"
    assert nat.destination == ["198.51.100.40"]
    assert nat.translated_destinations == ["10.0.0.40"]
    assert (nat.original_destination_ports[0].start, nat.original_destination_ports[0].end) == (8443, 8444)
    assert (nat.translated_destination_ports[0].start, nat.translated_destination_ports[0].end) == (443, 444)

    workbook = load_workbook(io.BytesIO(IRExcelExporter(
        result.canonical_ir, result
    ).generate()))
    virtual_ips = workbook["Virtual IPs"]
    vip_headers = {cell.value: cell.column for cell in virtual_ips[3]}
    assert virtual_ips.cell(4, vip_headers["External IP"]).value == "198.51.100.40"
    assert virtual_ips.cell(4, vip_headers["Mapped IPs"]).value == "10.0.0.40"
    assert virtual_ips.cell(4, vip_headers["External Port"]).value == "8443-8444"
    assert virtual_ips.cell(4, vip_headers["Mapped Port"]).value == "443-444"

    nat_sheet = workbook["NAT Rules"]
    nat_headers = {cell.value: cell.column for cell in nat_sheet[3]}
    nat_row = next(
        row for row in range(4, nat_sheet.max_row + 1)
        if nat_sheet.cell(row, nat_headers["VIP"]).value == "VIP_PAT"
    )
    assert nat_sheet.cell(nat_row, nat_headers["Original Destination"]).value == "198.51.100.40"
    assert nat_sheet.cell(nat_row, nat_headers["Translated Destination"]).value == "10.0.0.40"
    assert nat_sheet.cell(nat_row, nat_headers["Original Destination Port"]).value == "8443-8444"
    assert nat_sheet.cell(nat_row, nat_headers["Translated Destination Port"]).value == "443-444"


def test_vip_group_and_member_interface_conflict_requires_review():
    ir = _transform("""
config firewall vip
    edit "VIP_A"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set extintf "WAN"
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set interface "OTHER_WAN"
        set member "VIP_A"
    next
end
config firewall policy
    edit 41
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "ALL"
        set action accept
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_vip_group_reference == "VIP_GROUP"
    assert rule.requires_manual_review is True
    assert any("conflicts" in reason for reason in rule.review_reasons)
    assert "original_packet {" not in _main_tf(ir)


def test_vip_port_forward_preserves_original_and_translated_ports():
    ir = _transform("""
config firewall vip
    edit "VIP_PAT"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set portforward enable
        set protocol tcp
        set extport 8443
        set mappedport 443
    next
end
config firewall policy
    edit 50
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_PAT"
        set service "ALL"
        set action accept
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.destination_protocol == "tcp"
    assert rule.original_destination_port == "8443"
    assert rule.translated_port == "443"
    assert any(service.name == "svc_nat_tcp_8443" for service in ir.services)
    xml = PANOSXMLGenerator().generate(ir)[0].content
    assert "<service>svc_nat_tcp_8443</service>" in xml
    assert "<translated-port>443</translated-port>" in xml
    hcl = _main_tf(ir)
    assert 'service               = "svc_nat_tcp_8443"' in hcl
    assert "static_translation {" in hcl
    assert "port    = 443" in hcl


def test_vip_sctp_port_forward_preserves_protocol_and_ports():
    ir = _transform("""
config firewall vip
    edit "VIP_SCTP"
        set extip 198.51.100.12
        set mappedip "10.0.0.12"
        set extintf "WAN"
        set portforward enable
        set protocol sctp
        set extport 3868
        set mappedport 13868
    next
end
config firewall policy
    edit 52
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_SCTP"
        set service "ALL"
        set action accept
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.destination_protocol == "sctp"
    assert rule.original_destination_port == "3868"
    assert rule.translated_port == "13868"
    service = next(service for service in ir.services if service.name == "svc_nat_sctp_3868")
    assert service.ports[0].protocol == ServiceProtocol.SCTP


@pytest.mark.parametrize(
    "vip_type",
    ["load-balance", "server-load-balance", "dns-translation", "fqdn", "access-proxy"],
)
def test_advanced_vip_types_are_inventory_only_and_not_static_dnat(vip_type):
    ir = _transform(f"""
config firewall vip
    edit "ADVANCED_VIP"
        set type {vip_type}
        set extip 198.51.100.30
        set mappedip "10.0.0.30"
        set extintf "WAN"
    next
end
config firewall policy
    edit 53
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "ADVANCED_VIP"
        set service "ALL"
        set action accept
    next
end
""")

    assert ir.virtual_ips[0].vip_type == vip_type
    assert ir.nat_rules == []
    assert any("canonical DNAT was withheld" in entry.message for entry in ir.audit_entries)


def test_vip_omitted_protocol_defaults_to_tcp_and_populates_canonical_port_ranges():
    ir = _transform("""
config firewall vip
    edit "VIP_DEFAULT_TCP"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set portforward enable
        set extport 8443-8444
        set mappedport 443-444
    next
end
config firewall policy
    edit 51
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_DEFAULT_TCP"
        set service "ALL"
        set action accept
    next
end
""")

    vip = ir.virtual_ips[0]
    rule = ir.nat_rules[0]
    assert vip.protocol == rule.destination_protocol == "tcp"
    assert (rule.original_destination_ports[0].start, rule.original_destination_ports[0].end) == (8443, 8444)
    assert (rule.translated_destination_ports[0].start, rule.translated_destination_ports[0].end) == (443, 444)


def test_disabled_policy_stays_disabled_in_ir_and_panos_xml():
    ir = _transform(f"""
config firewall policy
    edit 60
{POLICY_BASE}
        set status disable
        set nat enable
    next
end
""")

    assert ir.nat_rules[0].enabled is False
    assert ir.nat_rules[0].translated_sources == ["203.0.113.10"]
    xml = PANOSXMLGenerator().generate(ir)[0].content
    assert '<entry name="SNAT-P60">' in xml
    assert "<disabled>yes</disabled>" in xml
    assert "disabled    = true" in _main_tf(ir)


def test_multiple_services_survive_ir_and_split_into_panos_rules():
    ir = _transform("""
config firewall policy
    edit 70
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "all"
        set service "HTTPS" "port_8081"
        set action accept
        set nat enable
    next
end
""")

    assert ir.nat_rules[0].services == ["HTTPS", "port_8081"]
    pan = IRToPANOSTransformer(ir).transform()
    assert [rule.service for rule in pan.vsys.nat_rules] == ["HTTPS", "port_8081"]
    assert [rule.name for rule in pan.vsys.nat_rules] == [
        "SNAT-P70-HTTPS", "SNAT-P70-port_8081",
    ]
    hcl = _main_tf(ir)
    assert 'name = "SNAT-P70-HTTPS"' in hcl
    assert 'name = "SNAT-P70-port_8081"' in hcl


def test_snat_and_vip_create_one_twice_nat_rule():
    ir = _transform("""
config firewall vip
    edit "VIP_WEB"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
        set extintf "WAN"
    next
end
config firewall policy
    edit 80
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_WEB"
        set service "HTTPS"
        set action accept
        set nat enable
    next
end
""")

    assert len(ir.nat_rules) == 1
    rule = ir.nat_rules[0]
    assert rule.type == NATType.TWICE
    assert rule.source_translation_mode == NATTranslationMode.INTERFACE_ADDRESS
    assert rule.translated_destinations == ["10.0.0.10"]
    xml = PANOSXMLGenerator().generate(ir)[0].content
    assert xml.count("<entry name=\"TWICE-P80-VIP_WEB\">") == 1
    assert "<source-translation>" in xml
    assert "<destination-translation>" in xml
    hcl = _main_tf(ir)
    assert hcl.count('name = "TWICE-P80-VIP_WEB"') == 1
    assert "interface_address {" in hcl
    assert "static_translation {" in hcl


def test_missing_pool_never_falls_back_and_is_withheld_from_target():
    ir = _transform(f"""
config firewall policy
    edit 90
{POLICY_BASE}
        set nat enable
        set ippool enable
        set poolname "MISSING_POOL"
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.source_translation_mode == NATTranslationMode.POOL
    assert rule.source_pool_references == ["MISSING_POOL"]
    assert rule.translated_sources == []
    assert rule.requires_manual_review is True
    assert any(
        entry.confidence == MigrationConfidence.MANUAL and "missing IP pool" in entry.message
        for entry in ir.audit_entries
    )
    assert IRToPANOSTransformer(ir).transform().vsys.nat_rules == []
    assert 'resource "panos_nat_rule_group"' not in _main_tf(ir)


def test_internet_service_nat_preserves_reference_without_any_fallback():
    ir = _transform("""
config firewall policy
    edit 100
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set internet-service enable
        set internet-service-name "Microsoft-Office365"
        set action accept
        set nat enable
    next
end
""")

    rule = ir.nat_rules[0]
    assert rule.internet_services == ["Microsoft-Office365"]
    assert rule.destination == []
    assert rule.services == []
    assert rule.requires_manual_review is True


def test_mixed_vip_and_ordinary_destinations_partition_snat_match():
    ir = _transform("""
config firewall vip
    edit "VIP_WEB"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
    next
end
config firewall policy
    edit 110
        set srcintf "LAN"
        set dstintf "WAN"
        set srcaddr "LAN_NET"
        set dstaddr "VIP_WEB" "ORDINARY_SERVER"
        set service "HTTPS"
        set action accept
        set nat enable
    next
end
""")

    assert [rule.type for rule in ir.nat_rules] == [NATType.TWICE, NATType.SOURCE]
    assert ir.nat_rules[0].destination == ["198.51.100.10"]
    assert ir.nat_rules[1].destination == ["ORDINARY_SERVER"]


def test_vip_group_unresolved_member_taints_resolved_nat_without_dropping_inventory():
    result = extract_fortigate_config("""
config firewall vip
    edit "VIP_OK"
        set extip 198.51.100.10
        set mappedip "10.0.0.10"
    next
end
config firewall vipgrp
    edit "VIP_GROUP"
        set member "VIP_OK" "VIP_MISSING"
    next
end
config firewall policy
    edit 10
        set srcintf "WAN"
        set dstintf "LAN"
        set srcaddr "all"
        set dstaddr "VIP_GROUP"
        set service "HTTPS"
        set action accept
    next
end
""")

    group = result.canonical_ir.virtual_ip_groups[0]
    assert group.members == ["VIP_OK", "VIP_MISSING"]
    assert group.unresolved_members == ["VIP_MISSING"]
    assert group.migration_status == "PARTIALLY_NORMALIZED"

    rule = result.canonical_ir.nat_rules[0]
    assert rule.source_vip_reference == "VIP_OK"
    assert rule.source_vip_group_reference == "VIP_GROUP"
    assert rule.translated_destinations == ["10.0.0.10"]
    assert rule.migration_status == "PARTIALLY_NORMALIZED"
    assert any("VIP_MISSING" in reason for reason in rule.review_reasons)
