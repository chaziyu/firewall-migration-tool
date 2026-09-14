import io

from openpyxl import load_workbook

from fwmigrate.parsers.fortigate.extractor import extract_fortigate_config
from fwmigrate.report import IRExcelExporter


def test_static6_devindex_survives_canonical_ir_and_excel():
    result = extract_fortigate_config(
        """config router static6
    edit 1
        set dst 2001:db8:100::/64
        set gateway 2001:db8::1
        set device "wan1"
        set devindex 42
    next
end
"""
    )

    route = result.canonical_ir.routes[0]
    assert route.address_family == "ipv6"
    assert route.source_route_id == 1
    assert route.source_attributes["devindex"] == 42

    workbook = load_workbook(
        io.BytesIO(IRExcelExporter(result.canonical_ir, result).generate())
    )
    sheet = workbook["Routes"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Device Index"]).value == 42


def test_policy_based_firewall_policy_exports_only_as_ngfw_pre_match():
    result = extract_fortigate_config(
        """config system settings
    set ngfw-mode policy-based
end
config firewall policy
    edit 7
        set name "TLS pre-match"
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "HTTPS"
        set ssl-ssh-profile "certificate-inspection"
    next
end
config firewall security-policy
    edit 70
        set name "Application enforcement"
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "HTTPS"
        set application 12345
        set action accept
    next
end
"""
    )

    ir = result.canonical_ir
    assert ir.policies == []
    pre_match = [
        rule
        for rule in ir.source_only_rules
        if rule.family == "ngfw-pre-match-policy"
    ]
    assert len(pre_match) == 1
    assert pre_match[0].source_id == "7"
    assert pre_match[0].effective_action is None
    assert "action" not in pre_match[0].source_attributes
    assert len(ir.security_policies) == 1
    assert result.generation_safe is False

    workbook = load_workbook(io.BytesIO(IRExcelExporter(ir, result).generate()))
    assert "NGFW Pre-Match Policies" in workbook.sheetnames
    pre_match_sheet = workbook["NGFW Pre-Match Policies"]
    headers = {cell.value: cell.column for cell in pre_match_sheet[3]}
    assert "Action" not in headers
    assert pre_match_sheet.cell(4, headers["Rule ID"]).value == "7"
    assert (
        pre_match_sheet.cell(4, headers["SSL Inspection Reference"]).value
        == "certificate-inspection"
    )
    assert workbook["Policies"].max_row == 3
    assert workbook["NGFW Security Policies"].max_row == 4


def test_profile_based_omitted_action_still_defaults_to_deny():
    result = extract_fortigate_config(
        """config system settings
    set ngfw-mode profile-based
end
config firewall policy
    edit 10
        set srcintf "lan"
        set dstintf "wan"
        set srcaddr "all"
        set dstaddr "all"
        set service "HTTPS"
    next
end
"""
    )

    assert len(result.canonical_ir.policies) == 1
    policy = result.canonical_ir.policies[0]
    assert policy.source_rule_id == "10"
    assert policy.action.value == "deny"
