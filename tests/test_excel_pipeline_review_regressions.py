"""Source-to-workbook regressions for the Palo Alto / Junos export review."""

from copy import deepcopy
from ast import literal_eval
from io import BytesIO
import json
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from fwmigrate.collection.contracts import CollectedSource, CollectionStatus
from fwmigrate.collection.snapshot import make_snapshot, parse_snapshot, sanitize_source
from fwmigrate.vendors.juniper_srx.source_report import JuniperSRXSourceReporter
from fwmigrate.vendors.palo_alto.source_report import PaloAltoSourceReporter
from fwmigrate.web import create_app


PAN = PaloAltoSourceReporter()
JUNOS = JuniperSRXSourceReporter()


def export(reporter, source, profile="full"):
    analysis = reporter.analyze_source(source)
    before = deepcopy(analysis.config.model_dump())
    stream = BytesIO()
    reporter.export_excel(analysis, stream, profile=profile)
    assert analysis.config.model_dump() == before
    return analysis, load_workbook(BytesIO(stream.getvalue()))


def rows(workbook, sheet, header=None):
    header = header or (3 if sheet in {"Security Policies", "NAT Rules", "Virtual Router Routes", "Interfaces",
                                      "PAN-OS Source Inventory", "Extraction Coverage"} else 1)
    values = list(workbook[sheet].values)
    return [dict(zip(values[header - 1], value)) for value in values[header:]]


@pytest.fixture
def client():
    return create_app({"TESTING": True, "API_RATE_LIMIT_PER_MINUTE": 1000}).test_client()


PAN_SNMP = '''<config><devices><entry name="fw"><deviceconfig><system><snmp-setting>
<access-setting><version><v3><users><entry name="monitor">
<authpwd>AUTH_SENTINEL</authpwd><privpwd>PRIV_SENTINEL</privpwd><view>ordinary-view</view>
</entry></users></v3></version></access-setting></snmp-setting></system></deviceconfig></entry></devices></config>'''


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_pan_credential_aliases_are_redacted_in_direct_snapshot_and_http_exports(client, profile):
    safe = sanitize_source("palo_alto", PAN_SNMP)
    assert "SENTINEL" not in safe
    assert "ordinary-view" in safe and 'name="monitor"' in safe
    assert sanitize_source("palo_alto", safe) == safe
    snapshot = make_snapshot(CollectedSource("palo_alto", PAN_SNMP, "synthetic.xml", "test", CollectionStatus.SUCCESS))
    restored = parse_snapshot(json.dumps(snapshot).encode())
    assert "SENTINEL" not in restored.source_text
    analysis, workbook = export(PAN, PAN_SNMP, profile)
    assert "SENTINEL" not in json.dumps(analysis.config.model_dump(mode="json"))
    assert "SENTINEL" not in json.dumps(PAN.build_preview(analysis))
    for sheet in workbook:
        assert not any("SENTINEL" in str(cell.value) for row in sheet for cell in row)
    response = client.post("/api/extract/excel", data={"source_vendor": "palo_alto", "excel_profile": profile,
        "file": (BytesIO(PAN_SNMP.encode()), "synthetic.xml")})
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        assert all(b"SENTINEL" not in archive.read(name) for name in archive.namelist() if name.endswith(".xml"))


HIERARCHY = '''interfaces {
 ge-0/0/0 { description "uplink #1 /* literal */";
  unit 100 { vlan-id 200; family inet { address 192.0.2.1/24; } }
 }
}
system { root-authentication { encrypted-password "JUNOS_SENTINEL"; } }
'''


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_junos_sanitizer_preserves_hierarchy_and_http_units(client, profile):
    safe = sanitize_source("juniper_srx", HIERARCHY)
    assert 'description "uplink #1 /* literal */";' in safe
    assert "JUNOS_SENTINEL" not in safe
    assert sanitize_source("juniper_srx", safe) == safe
    direct = JUNOS.analyze_source(HIERARCHY)
    sanitized = JUNOS.analyze_source(safe)
    assert direct.config.get_context().interfaces == sanitized.config.get_context().interfaces
    snapshot = make_snapshot(CollectedSource("juniper_srx", HIERARCHY, "synthetic.conf", "test", CollectionStatus.SUCCESS))
    restored = JUNOS.analyze_source(parse_snapshot(json.dumps(snapshot).encode()).source_text)
    assert restored.config.get_context().interfaces == sanitized.config.get_context().interfaces
    response = client.post("/api/extract/excel", data={"source_vendor": "juniper_srx", "excel_profile": profile,
        "file": (BytesIO(HIERARCHY.encode()), "synthetic.conf")})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    data = rows(workbook, "Interface Addresses")
    assert [(row["Unit"], row["Address"]) for row in data] == [("100", "192.0.2.1/24")]
    assert next(row for row in rows(workbook, "Interface Details") if row["Unit"] == "100")["VLAN ID"] == 200


@pytest.mark.parametrize("literal", ["literal /* marker", "literal */ marker", "}", "{", ";", "[", "]", 'escaped "quote" /* marker'])
def test_junos_quoted_syntax_does_not_remove_following_source(literal):
    quoted = json.dumps(literal)
    for source in (
        f'set interfaces ge-0/0/0 description {quoted}\nset security address-book global address KEEP 192.0.2.1/32',
        f'interfaces {{ ge-0/0/0 {{ description {quoted}; }} }}\nsecurity {{ address-book global {{ address KEEP 192.0.2.1/32; }} }}',
    ):
        analysis = JUNOS.analyze_source(sanitize_source("juniper_srx", source))
        context = analysis.config.get_context()
        assert context.interfaces["ge-0/0/0"].description == literal
        assert context.address_books["global"].addresses["KEEP"].prefix == "192.0.2.1/32"


def test_junos_multiline_hierarchy_strings_preserve_values_and_source_lines():
    source = 'interfaces {\n ge-0/0/0 { description "first\n/* literal # marker\nlast"; }\n}\nsystem { host-name edge; }'
    safe = sanitize_source("juniper_srx", source)
    assert safe == source
    analysis = JUNOS.analyze_source(safe)
    assert analysis.config.get_context().interfaces["ge-0/0/0"].description == "first\n/* literal # marker\nlast"
    assert analysis.config.hostname == "edge"
    assert analysis.inventory_items[-1].commands[0].line_number == 6


@pytest.mark.parametrize("source", [
    'interfaces { ge-0/0/0 { description "unterminated; } }',
    'interfaces { ge-0/0/0 { description uplink;',
    'interfaces { ge-0/0/0 { description uplink } }',
    'interfaces { ge-0/0/0 { description uplink; } } }',
    'applications { application-set X { application [ A B; } }',
    'interfaces { /* unterminated',
])
def test_junos_malformed_hierarchy_is_rejected_with_safe_line_diagnostic(client, source):
    with pytest.raises(ValueError, match=r"line \d+"):
        JUNOS.analyze_source(source)
    response = client.post("/api/extract/excel", data={"source_vendor": "juniper_srx",
        "file": (BytesIO(source.encode()), "malformed.conf")})
    assert response.status_code == 400


def test_junos_case_distinct_objects_groups_and_scopes_remain_distinct():
    source = '''set security address-book global address Foo 192.0.2.1/32
set security address-book global address foo 192.0.2.2/32
deactivate security address-book global address Foo
set groups G security address-book global address UPPER 192.0.2.3/32
set groups g security address-book global address lower 192.0.2.4/32
set apply-groups G
set logical-systems LS interfaces ge-0/0/0 description upper
set logical-systems ls interfaces ge-0/0/0 description lower
deactivate logical-systems LS interfaces ge-0/0/0'''
    analysis, workbook = export(JUNOS, source)
    states = {row["Name"]: row["Effective State"] for row in rows(workbook, "Addresses")}
    assert states == {"Foo": "INACTIVE", "foo": "EFFECTIVE"}
    paths = [row["target_path"] for row in analysis.derived.inheritance["effective_commands"]]
    assert any("UPPER" in path for path in paths)
    assert not any("lower" in path for path in paths)
    detail = {row["Context"]: row["Effective State"] for row in rows(workbook, "Interface Details")}
    assert detail["logical-system LS"] == "INACTIVE"
    assert detail["logical-system ls"] == "EFFECTIVE"


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_pan_feature_rows_retain_full_device_scope(profile):
    source = "<config><devices>" + "".join(
        f'<entry name="{device}" serial="{device}"><vsys><entry name="vsys1"><address>'
        '<entry name="X"><ip-netmask>192.0.2.1/32</ip-netmask></entry></address></entry></vsys></entry>'
        for device in ("A", "B")) + "</devices></config>"
    _, workbook = export(PAN, source, profile)
    data = rows(workbook, "Addresses", 3)
    assert {row["Device Serial"] for row in data} == {"A", "B"}
    assert len({row["Scope Identity"] for row in data}) == 2
    assert len({row["Source Record"] for row in data}) == 2


def test_pan_policy_reference_display_does_not_join_nat_rule_of_same_name():
    source = '''<config><shared><address><entry name="A"><ip-netmask>192.0.2.1</ip-netmask></entry>
<entry name="B"><ip-netmask>192.0.2.2</ip-netmask></entry></address><rulebase>
<security><rules><entry name="R"><source><member>A</member></source></entry></rules></security>
<nat><rules><entry name="R"><source><member>B</member></source><source-translation>
<static-ip><translated-address>192.0.2.3</translated-address></static-ip></source-translation></entry></rules></nat>
</rulebase></shared></config>'''
    _, workbook = export(PAN, source)
    references = rows(workbook, "Security Policies")[0]["Resolved Source References"]
    assert "A ->" in references
    assert "B ->" not in references


def test_pan_route_diagnostics_stay_with_nested_owner_in_preview_and_excel():
    source = '<config><devices><entry name="fw"><network><virtual-router>' + "".join(
        f'<entry name="{name}"><routing-table><ip><static-route><entry name="default"><destination>'
        f'{destination}</destination></entry></static-route></ip></routing-table></entry>'
        for name, destination in (("VR1", "0.0.0.0/0"), ("VR2", "BAD"))
    ) + '</virtual-router></network></entry></devices></config>'
    analysis, workbook = export(PAN, source)
    data = {row["Virtual Router"]: row for row in rows(workbook, "Virtual Router Routes")}
    assert data["VR1"]["Analysis Status"] == "NO_VALIDATION_ISSUES"
    assert data["VR2"]["Analysis Status"] == "REVIEW_REQUIRED"
    assert len([issue for issue in analysis.validation.issues if "BAD" in issue.message]) == 1
    preview = PAN.build_preview(analysis)
    route_rows = preview["sections"]["routes"]
    assert len([row for row in route_rows if row["review"]]) == 1


def test_pan_standalone_vsys_rules_have_explicit_source_only_order():
    source = '''<config><devices><entry name="fw"><vsys><entry name="vsys1"><rulebase><security><rules>
<entry name="first"/><entry name="second"/></rules></security></rulebase></entry></vsys></entry></devices></config>'''
    analysis, workbook = export(PAN, source)
    assert [row.rule_name for row in analysis.derived.policy_order] == ["first", "second"]
    assert all(not row.complete for row in analysis.derived.policy_order)
    assert all("source-only" in row["Effective Order"] for row in rows(workbook, "Security Policies"))


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_feature_settings_include_application_top_level_terms_and_explicit_vlan(profile):
    source = '''set applications application HTTPS protocol tcp
set applications application HTTPS destination-port 443
set applications application MULTI term DNS protocol udp
set applications application MULTI term DNS destination-port 53'''
    _, workbook = export(JUNOS, source, profile)
    settings = rows(workbook, "Application Settings")
    https = next(row for row in settings if row["Application"] == "HTTPS")
    assert (https["Location"], https["Term"], https["Protocol"], https["Destination Ports"]) == ("top-level", None, "tcp", "['443']")
    dns = next(row for row in settings if row["Application"] == "MULTI" and row["Term"] == "DNS")
    assert (dns["Protocol"], dns["Destination Ports"]) == ("udp", "['53']")
    source = '''<config><devices><entry name="fw"><network><interface><ethernet><entry name="ethernet1/1">
<layer3><units><entry name="ethernet1/1.100"><tag>200</tag><ip><entry name="192.0.2.1/24"/></ip></entry></units></layer3>
</entry></ethernet></interface></network></entry></devices></config>'''
    _, workbook = export(PAN, source, profile)
    unit = next(row for row in rows(workbook, "Interfaces") if row["Name"] == "ethernet1/1.100")
    assert (unit["Unit / Subinterface"], unit["VLAN Tag"]) == ("ethernet1/1.100", "200")


@pytest.mark.parametrize("distribution", ["round-robin", "source-ip-hash", "ip-modulo", "ip-hash", "least-sessions", "future-choice"])
def test_pan_scalar_nat_distribution_survives_with_accurate_coverage(distribution):
    source = f'''<config><shared><rulebase><nat><rules><entry name="N"><dynamic-destination-translation>
<translated-address><member>192.0.2.1</member></translated-address><distribution>{distribution}</distribution>
</dynamic-destination-translation></entry></rules></nat></rulebase></shared></config>'''
    analysis, workbook = export(PAN, source)
    assert analysis.config.nat_rules[0].dynamic_destination_translation.distribution == distribution
    assert rows(workbook, "NAT Rules")[0]["Dynamic Destination Distribution"] == distribution
    status = next(row["Extraction Status"] for row in rows(workbook, "PAN-OS Source Inventory") if row["Object"] == "N")
    assert status == ("PARTIAL" if distribution == "future-choice" else "EXTRACTED")


@pytest.mark.parametrize("reporter,header", [(PAN, 3), (JUNOS, 1)])
@pytest.mark.parametrize("profile", ["fast", "full"])
def test_upstream_overflow_fix_preserves_actual_vendor_workbook_values(reporter, header, profile):
    value = "x" * 32767 + "=1+1😀END"
    source = (f'<config><shared><address><entry name="X"><description>{value}</description></entry></address></shared></config>'
              if reporter is PAN else f'set security address-book global address X description "{value}"')
    _, workbook = export(reporter, source, profile)
    assert "See Text Overflow:" in rows(workbook, "Addresses", header)[0]["Description"]
    overflow = [row for row in list(workbook["Text Overflow"].values)[1:] if row[0] == "Addresses"]
    assert "".join(row[4] for row in overflow) == value
    assert all(cell.data_type != "f" for row in workbook["Text Overflow"] for cell in row)


@pytest.mark.parametrize("profile", ["fast", "full"])
def test_pan_nested_settings_ipv6_flags_and_service_group_fields_survive_download(client, profile):
    source = '''<config><shared>
      <service><entry name="svc"><future-root>ROOT_SERVICE</future-root><protocol>
        <tcp><port>443</port><future-control>TCP_EXTRA</future-control>
          <override><yes><timeout>60</timeout><future-override>OVERRIDE_EXTRA</future-override>
            <api-key>SECRET_SENTINEL</api-key></yes></override></tcp>
        <udp><port>53</port><future-control>UDP_EXTRA</future-control></udp>
      </protocol></entry></service>
      <service-group><entry name="group"><members><member>svc</member></members>
        <description>=GROUP_DESCRIPTION</description><disable-override>yes</disable-override>
      </entry></service-group>
      <rulebase><nat><rules><entry name="nat"><future-root>ROOT_NAT</future-root>
        <source-translation><dynamic-ip><translated-address><member>192.0.2.1</member></translated-address>
          <future-source>SOURCE_EXTRA</future-source><fallback><interface-address>
            <interface>ethernet1/2</interface><api-key>SECRET_SENTINEL</api-key>
          </interface-address></fallback></dynamic-ip></source-translation>
        <destination-translation><translated-address>192.0.2.2</translated-address>
          <future-destination>DESTINATION_EXTRA</future-destination>
          <dns-rewrite><direction>reverse</direction><future-dns>DNS_EXTRA</future-dns></dns-rewrite>
        </destination-translation>
      </entry></rules></nat></rulebase>
    </shared><devices><entry name="fw"><network><interface>
      <ethernet><entry name="ethernet1/1"><layer3><ipv6><address>
        <entry name="2001:db8::1/64"><enable>yes</enable></entry>
        <entry name="2001:db8::2/64"><enable>no</enable></entry>
        <entry name="2001:db8::3/64"/>
      </address></ipv6><units><entry name="ethernet1/1.10"><tag>10</tag><ipv6><address>
        <entry name="2001:db8:10::1/64"><enable>no</enable></entry>
      </address></ipv6></entry></units></layer3></entry></ethernet>
      <tunnel><units><entry name="tunnel.1"><ipv6><address>
        <entry name="2001:db8:20::1/64"><enable>yes</enable></entry>
      </address></ipv6></entry></units></tunnel>
    </interface></network></entry></devices></config>'''
    _, direct = export(PAN, source, profile)
    response = client.post("/api/extract/excel", data={"source_vendor": "palo_alto", "excel_profile": profile,
        "file": (BytesIO(source.encode()), "nested.xml")})
    assert response.status_code == 200
    downloaded = load_workbook(BytesIO(response.data))
    for workbook in (direct, downloaded):
        nat = rows(workbook, "NAT Rules")[0]["Additional Settings"]
        for value in ("ROOT_NAT", "SOURCE_EXTRA", "fallback", "ethernet1/2", "DESTINATION_EXTRA", "DNS_EXTRA"):
            assert value in nat
        services = {row["Protocol"]: row for row in rows(workbook, "Services", 3)}
        assert all("ROOT_SERVICE" in row["Additional Settings"] for row in services.values())
        assert "TCP_EXTRA" in services["tcp"]["Additional Settings"]
        assert "OVERRIDE_EXTRA" in services["tcp"]["Additional Settings"]
        assert "UDP_EXTRA" not in services["tcp"]["Additional Settings"]
        assert "UDP_EXTRA" in services["udp"]["Additional Settings"]
        group = rows(workbook, "Service Groups", 3)[0]
        assert group["Description"] == "'=GROUP_DESCRIPTION"
        assert group["Disable Override"] == "yes"
        interfaces = {row["Name"]: row for row in rows(workbook, "Interfaces")}
        physical = interfaces["ethernet1/1"]["Additional Settings"]
        assert literal_eval(physical.split("ipv6_addresses: ", 1)[1]) == (
            {"address": "2001:db8::1/64", "enable": "yes"},
            {"address": "2001:db8::2/64", "enable": "no"},
            {"address": "2001:db8::3/64"},
        )
        assert "'enable': 'no'" in interfaces["ethernet1/1.10"]["Additional Settings"]
        assert "'enable': 'yes'" in interfaces["tunnel.1"]["Additional Settings"]
        assert "SECRET_SENTINEL" not in str([value for sheet in workbook for value in sheet.values])
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


@pytest.mark.parametrize("profile", ["fast", "full"])
@pytest.mark.parametrize("section,settings,expected", [
    ("router static6", "set dst 2001:db8:100::/64\nset gateway 2001:db8::1", "Yes"),
    ("firewall addrgrp6", 'set member "IPv6-object"', "Yes"),
    ("firewall ippool6", "set startip 2001:db8::10\nset endip 2001:db8::20", "Yes"),
    ("firewall vipgrp6", 'set member "IPv6-vip"', "Yes"),
    ("system interface", "config ipv6\nset ip6-address 2001:db8::1/64\nend", "Yes"),
    ("router static", "set dst 192.0.2.0 255.255.255.0", "No"),
])
def test_fortigate_excel_ipv6_presence_includes_source_sections(profile, section, settings, expected):
    from fwmigrate.vendors.fortigate.source_report import FortiGateSourceReporter

    reporter = FortiGateSourceReporter()
    analysis = reporter.analyze_source(f"config {section}\nedit 1\n{settings}\nnext\nend")
    before = deepcopy(analysis.extracted.config.model_dump())
    output = BytesIO()
    reporter.export_excel(analysis, output, profile=profile)
    workbook = load_workbook(output)
    summary = {row[0]: row[1] for row in workbook["Summary"].values}
    assert summary["IPv6 Explicit Configuration Present"] == expected
    assert analysis.extracted.config.model_dump() == before


@pytest.mark.parametrize("members", ["<member>ethernet1/1</member>", ""])
def test_pan_non_entry_import_has_consistent_inventory_preview_and_coverage(members):
    source = f'''<config><devices><entry name="fw"><vsys><entry name="vsys1"><import><network>
<interface>{members}</interface></network></import></entry></vsys></entry></devices></config>'''
    analysis, workbook = export(PAN, source)
    coverage = next(row for row in rows(workbook, "Extraction Coverage") if row["Source Domain"] == "interface_import")
    assert (coverage["Found"], coverage["Source Records"], coverage["Typed Objects"], coverage["Status"]) == ("yes", 1, 1, "EXTRACTED")
    assert any(row["Kind"] == "interface" and row["Extraction Status"] == "EXTRACTED"
               for row in rows(workbook, "PAN-OS Source Inventory"))
    preview = next(row for row in PAN.build_preview(analysis)["extraction_coverage"] if row["domain"] == "interface_import")
    assert (preview["source_records"], preview["typed_objects"], preview["status"]) == (1, 1, "EXTRACTED")
    evidence = rows(workbook, "PAN-OS Source Appendix", 3)
    if members:
        assert sum(row["Evidence Value"] == "ethernet1/1" for row in evidence) == 1


def test_pan_failed_non_entry_extraction_retains_inventory_and_failure_status(monkeypatch):
    from fwmigrate.vendors.palo_alto import source_builder

    original = source_builder.extract_typed

    def extract(element, path, context, order):
        if path[-3:] == ("import", "network", "interface"):
            raise ValueError("synthetic typed extraction failure")
        return original(element, path, context, order)

    monkeypatch.setattr(source_builder, "extract_typed", extract)
    source = '<config><devices><entry name="fw"><vsys><entry name="vsys1"><import><network><interface><member>ethernet1/1</member></interface></network></import></entry></vsys></entry></devices></config>'
    analysis = PAN.analyze_source(source)
    coverage = next(row for row in analysis.accounting.coverage if row.domain == "interface_import")
    assert (coverage.source_records, coverage.typed_objects, coverage.status) == (1, 0, "UNSUPPORTED")
    record = next(row for row in analysis.config.source_inventory if row.source_path.endswith("import/network/interface"))
    assert "ethernet1/1" in record.raw_xml
