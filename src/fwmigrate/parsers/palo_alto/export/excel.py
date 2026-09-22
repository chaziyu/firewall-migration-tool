from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook

from ..source_report import PaloAltoSourceResult


def _safe(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return "\n".join(f"{key}: {item}" for key, item in value.items()) if isinstance(value, dict) else "\n".join(map(str, value))
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_panos_excel(analysis: PaloAltoSourceResult, output: BinaryIO | str | Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    config = analysis.config

    summary = workbook.create_sheet("Summary")
    summary.append(["PAN-OS Source Report"])
    summary.append(["Hostname", config.hostname])
    summary.append(["Source Version", config.source_version])
    summary.append(["Scopes", len(config.scopes)])
    summary.append(["Source Records", len(config.records)])

    scopes = workbook.create_sheet("Scopes")
    scopes.append(["Kind", "Name", "Device", "Serial", "Device Group", "Template Stack"])
    for scope in config.scopes:
        scopes.append([scope.kind, scope.name, scope.device_name, scope.device_serial, scope.device_group, scope.template_stack])

    records = workbook.create_sheet("Source Inventory")
    records.append(["Order", "Kind", "Source Path", "Name", "Scope", "Values"])
    for record in config.records:
        records.append([record.source_order, record.kind, record.source_path, record.name, record.scope.name if record.scope else None, _safe(record.values)])

    for sheet_name, needles in {
        "Interfaces": ("interface",),
        "Addresses": ("address",),
        "Policies": ("rule",),
        "NAT Rules": ("nat",),
        "Routes": ("route",),
    }.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Order", "Kind", "Source Path", "Name", "Scope", "Values"])
        for record in config.records:
            if any(needle in record.source_path.lower() for needle in needles):
                sheet.append([record.source_order, record.kind, record.source_path, record.name, record.scope.name if record.scope else None, _safe(record.values)])

    validation = workbook.create_sheet("Validation")
    validation.append(["Severity", "Domain", "Source Path", "Source Name", "Message"])
    for issue in analysis.validation.issues:
        validation.append([issue.severity, issue.domain, issue.source_path, issue.source_name, issue.message])

    if isinstance(output, (str, Path)):
        Path(output).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    else:
        workbook.save(output)
