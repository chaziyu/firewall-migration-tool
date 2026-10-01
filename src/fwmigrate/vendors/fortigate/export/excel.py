from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from dataclasses import asdict, is_dataclass
from datetime import datetime, timezone
from enum import Enum
from functools import lru_cache
from itertools import chain
from pathlib import Path
from time import perf_counter
from typing import Any, BinaryIO, Iterable, Iterator, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fwmigrate.source_reporting import ExcelExportProfile
from fwmigrate.source_reporting.metrics import ExcelExportMetrics

from ..config import ExtractionConfig
from ..derived import DerivedViews, build_derived_views
from ..extraction.result import ExtractionResult
from ..extraction.source_inventory import SourceObjectRecord
from ..extraction.coverage import (
    build_typed_source_identity_index,
    build_typed_source_inventory,
    extraction_status,
    find_typed_source_object,
    supports_path,
)
from ..security.extraction import sanitize_source_attributes, sanitize_source_value
from ..section_registry import registered_sections
from ..transform.policies import effective_policy_action
from ..validation.index import ValidationIssueIndex
from ..validation.models import ValidationIssue, ValidationResult
from .excel_schema import (
    HIDDEN_COLUMNS_BY_DEFAULT,
    SHEET_HEADERS,
    SHEET_ORDER,
)
from .excel_common import (
    _POLICY_ACTION_NOTE,
    _VISIBLE_MODEL_FIELDS_BY_SHEET,
    _additional_source_settings,
    _additional_settings,
    _build_additional_settings,
    _excel_safe,
    _interface_source_values,
    _lookup_source_header,
    _model_rows,
    _normalize_key,
    _normalized_source_values,
    _overlay_safe_raw,
)
from .excel_rows import (
    _iter_fortigate_source_inventory_rows,
    rows_for_sheet as _rows_for_sheet,
)


_TITLE_FILL = PatternFill("solid", fgColor="17324D")
_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_REVIEW_FILL = PatternFill("solid", fgColor="FEF3C7")
_ERROR_FILL = PatternFill("solid", fgColor="FEE2E2")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_MUTED_FONT = Font(color="667085", italic=True, size=9)
_LINK_FONT = Font(color="0563C1", underline="single")
_THIN = Side(style="thin", color="D9E2DF")
_BORDER = Border(bottom=_THIN)

_COLUMN_WIDTHS_BY_HEADER = {
    "Name": 28,
    "Description": 40,
    "Comments": 40,
    "VDOM": 18,
    "VDOMs": 28,
    "Interface": 24,
    "Interfaces": 32,
    "Members": 40,
    "Source Addresses": 36,
    "Destination Addresses": 36,
    "Services": 36,
    "Analysis Status": 20,
    "Review Reasons": 45,
    "Additional Settings": 50,
}
_COLUMN_WIDTH_OVERRIDES_BY_SHEET = {
    "Administrators": {"IPv4 Trusted Hosts": 36, "VDOMs": 28},
    "Policies": {"Source Interface": 32, "Destination Interface": 32},
}
_DEFAULT_COLUMN_WIDTH = 22
_SOURCE_INVENTORY_WIDTHS = {
    "Domain": 24,
    "VDOM": 14,
    "Scope Type": 14,
    "Source Path": 34,
    "Object": 32,
    "Parent / Subsection": 32,
    "Operation": 12,
    "Setting": 36,
    "Value": 60,
    "Primitive Registration": 20,
    "Extraction Status": 18,
    "Raw Extra Entries": 18,
}

_FAST_REQUIRED_SHEETS = frozenset({"Review Required"})
_FAST_EXCLUDED_SHEETS = frozenset({"FortiGate Source Inventory", "Extraction Coverage"})




def export_excel(
    *,
    extracted: ExtractionResult,
    validation: ValidationResult,
    output: BinaryIO | Path | str,
    config: ExtractionConfig | None = None,
    derived: DerivedViews | None = None,
    profile: ExcelExportProfile | str = ExcelExportProfile.FULL,
    source_name: str | None = None,
    metrics: ExcelExportMetrics | None = None,
) -> None:
    """Write the FortiGate configuration report."""

    del config

    profile = ExcelExportProfile(profile)
    derived = derived or build_derived_views(extracted.config)

    excel_started = perf_counter()
    context_started = perf_counter()
    context = _ExcelContext(
        extracted=extracted,
        derived=derived,
        validation=validation,
        profile=profile,
        source_name=source_name,
    )
    if metrics is not None and profile is ExcelExportProfile.FAST:
        metrics.add_stage("excel_context", (perf_counter() - context_started) * 1000)

    build_started = perf_counter()
    workbook = _build_workbook(
        context,
        metrics=metrics if profile is ExcelExportProfile.FAST else None,
    )
    if metrics is not None and profile is ExcelExportProfile.FAST:
        metrics.add_stage("workbook_build", (perf_counter() - build_started) * 1000)
    if isinstance(output, (str, Path)):
        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)
        save_target = path
    else:
        save_target = output

    save_started = perf_counter()
    workbook.save(save_target)
    if metrics is not None and profile is ExcelExportProfile.FAST:
        metrics.add_stage("workbook_save", (perf_counter() - save_started) * 1000)
        metrics.add_stage("excel_total", (perf_counter() - excel_started) * 1000)


class _ExcelContext:
    def __init__(
        self,
        *,
        extracted: ExtractionResult,
        derived: DerivedViews,
        validation: ValidationResult,
        profile: ExcelExportProfile,
        source_name: str | None,
    ) -> None:
        self.extracted = extracted
        self.config = extracted.config
        self.derived = derived
        self.validation = validation
        self.profile = profile
        self.source_name = source_name or ""
        self.source_by_path: dict[str, list[SourceObjectRecord]] = defaultdict(list)
        self.source_by_identity: dict[tuple[str, str, str | None], list[SourceObjectRecord]] = defaultdict(list)
        self.nested_source_by_parent: dict[tuple[str, str, str | None], list[SourceObjectRecord]] = defaultdict(list)
        vdoms: set[str] = set()
        ipv6_explicit = False

        for record in extracted.source_objects:
            self.source_by_path[record.source_path].append(record)
            vdoms.add(record.vdom)
            self.source_by_identity[(record.vdom, record.source_path, record.object_name)].append(record)
            if record.source_path.startswith("system interface "):
                parent = record.parent_objects[0] if record.parent_objects else None
                self.nested_source_by_parent[(record.vdom, "system interface", parent)].append(record)
            ipv6_explicit = ipv6_explicit or any(
                "ipv6" in _normalize_key(key) or "ip6" in _normalize_key(key)
                for key in record.values
            )

        self.source_paths = frozenset(self.source_by_path)
        self.vdoms = frozenset(vdoms)
        self.registered_sections = frozenset(registered_sections())
        self._typed_source_inventory = None
        self._typed_source_identity_index = None
        self.source_only_counts = {
            path: len(records)
            for path, records in self.source_by_path.items()
            if not supports_path(path)
        }
        self.ipv6_explicit = ipv6_explicit

        self.validation_index = ValidationIssueIndex(validation)
        self._issues_cache: dict[
            tuple[str, tuple[str, ...], tuple[str, ...]],
            tuple[ValidationIssue, ...],
        ] = {}
        self._interface_safe_source_cache: dict[tuple[str, str], dict[str, Any]] = {}
        self._interface_source_key_cache: dict[tuple[str, str], frozenset[str]] = {}

    @property
    def typed_source_inventory(self):
        if self._typed_source_inventory is None:
            self._typed_source_inventory = build_typed_source_inventory(self.config)
        return self._typed_source_inventory

    @property
    def typed_source_identity_index(self):
        if self._typed_source_identity_index is None:
            self._typed_source_identity_index = build_typed_source_identity_index(
                self.typed_source_inventory
            )
        return self._typed_source_identity_index

    def issues_for(
        self,
        *,
        vdom: str,
        names: Iterable[Any],
        domains: Iterable[str] | None = None,
    ) -> tuple[ValidationIssue, ...]:
        vdom = str(vdom or "")
        names = tuple(
            dict.fromkeys(
                str(name)
                for name in names
                if name not in (None, "")
            )
        )
        domains = tuple(sorted(set(domains or ())))
        cache_key = (vdom, names, domains)
        cached = self._issues_cache.get(cache_key)
        if cached is not None:
            return cached

        found: list[ValidationIssue] = []
        seen: set[tuple[str, str, str]] = set()

        for issue in self.validation_index.issues_for(
            vdom,
            names,
            domains,
        ):
            key = (issue.domain, issue.field or "", issue.message)
            if key in seen:
                continue
            seen.add(key)
            found.append(issue)

        result = tuple(found)
        self._issues_cache[cache_key] = result
        return result


def _build_workbook(
    context: _ExcelContext,
    *,
    metrics: ExcelExportMetrics | None = None,
) -> Workbook:
    if context.profile is ExcelExportProfile.FAST:
        return _build_fast_workbook(context, metrics=metrics)

    workbook = Workbook()
    workbook.remove(workbook.active)

    order = list(SHEET_ORDER)

    for sheet_name in order:
        if sheet_name == "Summary":
            _build_summary(workbook, context, order)
        elif sheet_name == "FortiGate Source Inventory":
            _write_source_inventory_sheet(workbook, context)
        else:
            headers = list(SHEET_HEADERS[sheet_name])
            _write_table_sheet(
                workbook,
                sheet_name,
                headers,
                _rows_for_sheet(sheet_name, context, headers),
                context=context,
            )

    return workbook


def _build_fast_workbook(
    context: _ExcelContext,
    *,
    metrics: ExcelExportMetrics | None = None,
) -> Workbook:
    workbook = Workbook(write_only=True)
    summary_started = perf_counter()
    _write_fast_summary(workbook, context)
    if metrics is not None:
        metrics.add_stage("fast_summary", (perf_counter() - summary_started) * 1000)
    no_rows = object()
    other_sheet_total = 0.0
    measured_sheets: set[str] = set()
    measured_names = {
        "Addresses", "Policies", "Interfaces", "NAT Rules", "Routes", "VPN Tunnels"
    }

    for sheet_name in SHEET_ORDER:
        if sheet_name == "Summary" or sheet_name in _FAST_EXCLUDED_SHEETS:
            continue

        sheet_started = perf_counter()
        headers = SHEET_HEADERS[sheet_name]
        rows = iter(_rows_for_sheet(sheet_name, context, headers))

        first = next(rows, no_rows)
        if first is no_rows and sheet_name not in _FAST_REQUIRED_SHEETS:
            sheet_elapsed = (perf_counter() - sheet_started) * 1000
            if metrics is not None and sheet_name in measured_names:
                metrics.add_sheet(sheet_name, 0, len(headers), sheet_elapsed)
                measured_sheets.add(sheet_name)
            if sheet_name not in measured_names:
                other_sheet_total += sheet_elapsed
            continue

        row_count = _write_table_sheet_fast(
            workbook,
            sheet_name,
            headers,
            () if first is no_rows else chain((first,), rows),
            context=context,
        )
        sheet_elapsed = (perf_counter() - sheet_started) * 1000
        if metrics is not None:
            metrics.add_sheet(sheet_name, row_count, len(headers), sheet_elapsed)
            measured_sheets.add(sheet_name)
        if sheet_name not in measured_names:
            other_sheet_total += sheet_elapsed

    if metrics is not None:
        for sheet_name in ("Addresses", "Policies", "Interfaces", "NAT Rules", "Routes", "VPN Tunnels"):
            if sheet_name not in measured_sheets:
                metrics.add_sheet(sheet_name, 0, len(SHEET_HEADERS[sheet_name]), 0.0)
        metrics.add_stage("other_sheet_total", other_sheet_total)

    return workbook


def _fast_cell(
    sheet,
    value: Any,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    return cell


def _write_fast_summary(workbook: Workbook, context: _ExcelContext) -> None:
    sheet = workbook.create_sheet("Summary")
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 90
    for column, width in (("A", 42), ("B", 60), ("C", 20), ("D", 4), ("E", 38), ("F", 18)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[10].height = 45

    width = 6
    band = lambda value, end: [
        _fast_cell(sheet, value, font=_TITLE_FONT if end == 6 else _WHITE_FONT, fill=_TITLE_FILL if end == 6 else _HEADER_FILL),
        *[_fast_cell(sheet, None, fill=_TITLE_FILL if end == 6 else _HEADER_FILL) for _ in range(end - 1)],
        *([None] * (width - end)),
    ]
    sheet.append([
        *band("FortiGate Configuration Report", 6),
    ])
    sheet.append([None] * width)
    metadata = [
        ("Source File", context.source_name),
        ("Hostname", _hostname(context)),
        ("FortiOS Version", context.extracted.source_metadata.fortios_version),
        ("VDOMs", "\n".join(_vdoms(context))),
        ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Validation Errors", len(context.validation.errors)),
        ("Validation Warnings", len(context.validation.warnings)),
        (
            "Export Profile",
            "FAST export omits command-level Source Inventory and Extraction Coverage. "
            "Use FULL export for complete source traceability.",
        ),
    ]
    counts = _inventory_counts(context)
    navigation = [
        target
        for _, count, target in counts
        if count and target not in _FAST_EXCLUDED_SHEETS
    ]
    nav_index = 0

    def append_row(values: list[Any], *, navigation_header: bool = False) -> None:
        nonlocal nav_index
        values.extend([None] * (width - len(values)))
        if navigation_header:
            values[4] = _fast_cell(sheet, "Workbook navigation", font=_WHITE_FONT, fill=_HEADER_FILL)
        elif nav_index < len(navigation):
            target = navigation[nav_index]
            nav_index += 1
            cell = _fast_cell(sheet, target, font=_LINK_FONT)
            cell.hyperlink = f"#'{target}'!A1"
            values[4] = cell
        sheet.append(values)

    for index, (label, value) in enumerate(metadata):
        row = [
            _fast_cell(sheet, label, font=Font(bold=True, color="41504C")),
            _fast_cell(sheet, _excel_safe(value), alignment=Alignment(wrap_text=True, vertical="top")),
            None,
            None,
        ]
        append_row(row, navigation_header=index == 0)

    sheet.append(band("Inventory", 3))
    for label, count, target in counts:
        row = [label, count]
        if count and target not in _FAST_EXCLUDED_SHEETS:
            cell = _fast_cell(sheet, "Open sheet", font=_LINK_FONT)
            cell.hyperlink = f"#'{target}'!A1"
            row.append(cell)
        else:
            row.append(None)
        append_row(row)

    sheet.append(band("Migration Indicators", 3))
    for label, value in _migration_indicators(context):
        append_row([label, _excel_safe(value)])

    while nav_index < len(navigation):
        append_row([])


def _build_summary(
    workbook: Workbook,
    context: _ExcelContext,
    order: Sequence[str],
) -> None:
    sheet = workbook.create_sheet("Summary")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "FortiGate Configuration Report"
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    metadata = [
        ("Source File", context.source_name),
        ("Hostname", _hostname(context)),
        ("FortiOS Version", context.extracted.source_metadata.fortios_version),
        ("VDOMs", "\n".join(_vdoms(context))),
        ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Validation Errors", len(context.validation.errors)),
        ("Validation Warnings", len(context.validation.warnings)),
    ]

    for row_number, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row_number, 1, label).font = Font(bold=True, color="41504C")
        sheet.cell(row_number, 2, _excel_safe(value))
        sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")

    section_row = 11
    sheet.cell(section_row, 1, "Inventory")
    sheet.cell(section_row, 1).fill = _HEADER_FILL
    sheet.cell(section_row, 1).font = _WHITE_FONT
    sheet.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=3)

    counts = _inventory_counts(context)

    row_number = section_row + 1
    for label, count, target in counts:
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, count)
        if target in order:
            cell = sheet.cell(row_number, 3, "Open sheet")
            cell.hyperlink = f"#'{target}'!A1"
            cell.font = _LINK_FONT
        row_number += 1

    indicator_row = row_number + 1
    sheet.cell(indicator_row, 1, "Migration Indicators")
    sheet.cell(indicator_row, 1).fill = _HEADER_FILL
    sheet.cell(indicator_row, 1).font = _WHITE_FONT
    sheet.merge_cells(
        start_row=indicator_row,
        start_column=1,
        end_row=indicator_row,
        end_column=3,
    )

    for row_number, (label, value) in enumerate(
        _migration_indicators(context),
        start=indicator_row + 1,
    ):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)

    nav_col = 5
    sheet.cell(3, nav_col, "Workbook navigation")
    sheet.cell(3, nav_col).fill = _HEADER_FILL
    sheet.cell(3, nav_col).font = _WHITE_FONT

    nav_row = 4
    for name in order:
        if name == "Summary":
            continue
        cell = sheet.cell(nav_row, nav_col, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = _LINK_FONT
        nav_row += 1

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["E"].width = 34


def _write_table_sheet(
    workbook: Workbook,
    sheet_name: str,
    headers: Sequence[str],
    rows: Iterable[Mapping[str, Any]],
    *,
    context: _ExcelContext,
) -> None:
    if context.profile is ExcelExportProfile.FAST:
        _write_table_sheet_fast(workbook, sheet_name, headers, rows, context=context)
        return

    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 90

    max_col = max(len(headers), 1)
    hidden_headers = set(HIDDEN_COLUMNS_BY_DEFAULT.get(sheet_name, ()))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    sheet.cell(1, 1, sheet_name)
    sheet.cell(1, 1).fill = _TITLE_FILL
    sheet.cell(1, 1).font = _TITLE_FONT
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24

    if max_col > 2:
        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=max_col - 1,
        )
    note = sheet.cell(2, 1)
    note.value = ""
    note.font = _MUTED_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = _BORDER

    _apply_widths(sheet, headers, sheet_name=sheet_name)
    for column, header in enumerate(headers, start=1):
        if header in hidden_headers:
            sheet.column_dimensions[get_column_letter(column)].hidden = True

    row_count = 0
    for row in rows:
        row_count += 1
        sheet.append([_excel_safe(row.get(header)) for header in headers])
        outline = row.get("__outline_level__")
        if isinstance(outline, int) and outline > 0:
            sheet.row_dimensions[row_count + 3].outlineLevel = min(outline, 7)

    note.value = _sheet_note(sheet_name, row_count)

    if row_count:
        sheet.auto_filter.ref = f"A3:{get_column_letter(max_col)}{row_count + 3}"

    sheet.freeze_panes = _freeze_pane(sheet_name, headers)
    _add_back_link(sheet)

    if sheet_name == "Review Required":
        _apply_review_colors(sheet, headers, row_count)


def _write_table_sheet_fast(
    workbook: Workbook,
    sheet_name: str,
    headers: Sequence[str],
    rows: Iterable[Mapping[str, Any]],
    *,
    context: _ExcelContext,
) -> int:
    del context
    sheet = workbook.create_sheet(sheet_name)
    max_col = max(len(headers), 1)
    hidden_headers = set(HIDDEN_COLUMNS_BY_DEFAULT.get(sheet_name, ()))
    sheet.freeze_panes = _freeze_pane(sheet_name, headers)
    _apply_widths(sheet, headers, wide=True)
    for column, header in enumerate(headers, start=1):
        if header in hidden_headers:
            sheet.column_dimensions[get_column_letter(column)].hidden = True
    sheet.append([
        _fast_cell(sheet, sheet_name, font=_TITLE_FONT, fill=_TITLE_FILL),
        *([None] * (max_col - 1)),
    ])
    sheet.append([
        _fast_cell(
            sheet,
            (
                "FAST lightweight export. " + _POLICY_ACTION_NOTE
                if sheet_name == "Policies"
                else "FAST lightweight export; values preserve explicit source data and analysis columns."
            ),
            font=_MUTED_FONT,
            alignment=Alignment(wrap_text=True, vertical="top"),
        ),
        *([None] * (max_col - 1)),
    ])
    sheet.append([_fast_cell(sheet, header, font=_HEADER_FONT, fill=_HEADER_FILL) for header in headers])

    row_count = 0
    for row in rows:
        row_count += 1
        sheet.append([_excel_safe(row.get(header)) for header in headers])

    if row_count:
        sheet.auto_filter.ref = f"A3:{get_column_letter(max_col)}{row_count + 3}"
    return row_count



def _write_source_inventory_sheet(
    workbook: Workbook,
    context: _ExcelContext,
    *,
    rows: Iterable[tuple[Any, ...]] | None = None,
) -> None:
    sheet_name = "FortiGate Source Inventory"
    headers = SHEET_HEADERS[sheet_name]
    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 90

    max_col = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    sheet.cell(1, 1, sheet_name)
    sheet.cell(1, 1).fill = _TITLE_FILL
    sheet.cell(1, 1).font = _TITLE_FONT
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col - 1)
    note = sheet.cell(2, 1)
    note.font = _MUTED_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = _BORDER

    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = _SOURCE_INVENTORY_WIDTHS[header]

    row_count = 0
    source_rows = rows if rows is not None else _iter_fortigate_source_inventory_rows(context)
    for values in source_rows:
        sheet.append(values)
        row_count += 1

    note.value = _sheet_note(sheet_name, row_count, context.profile)

    sheet.auto_filter.ref = f"A3:{get_column_letter(max_col)}{row_count + 3}"
    sheet.freeze_panes = "A4"
    _add_back_link(sheet)


def _add_back_link(sheet) -> None:
    if sheet.max_column <= 1:
        return

    cell = sheet.cell(2, sheet.max_column)
    cell.value = "Back to Summary"
    cell.hyperlink = "#'Summary'!A1"
    cell.font = _LINK_FONT
    cell.alignment = Alignment(horizontal="right")


def _apply_review_colors(
    sheet,
    headers: Sequence[str],
    row_count: int,
    *,
    fast: bool = False,
) -> None:
    if "Severity" not in headers:
        return

    severity_col = headers.index("Severity") + 1

    for row in range(4, row_count + 4):
        severity = str(sheet.cell(row, severity_col).value or "").lower()
        fill = _ERROR_FILL if severity == "error" else _REVIEW_FILL

        if fast:
            sheet.cell(row, severity_col).fill = fill
            continue

        for column in range(1, len(headers) + 1):
            sheet.cell(row, column).fill = fill


def _apply_widths(
    sheet,
    headers: Sequence[str],
    *,
    sheet_name: str | None = None,
    wide: bool = False,
) -> None:
    for index, header in enumerate(headers, start=1):
        if sheet_name is not None:
            width = (
                _COLUMN_WIDTH_OVERRIDES_BY_SHEET.get(sheet_name, {}).get(header)
                or _COLUMN_WIDTHS_BY_HEADER.get(header)
                or _DEFAULT_COLUMN_WIDTH
            )
        else:
            normalized = header.lower()
            if any(token in normalized for token in ("additional settings", "review reason", "raw capture", "notes")):
                width = 36
            elif any(token in normalized for token in ("description", "comment", "source path", "setting", "value")):
                width = 28
            elif any(token in normalized for token in ("members", "addresses", "interfaces", "services", "profiles", "references")):
                width = 26
            elif any(token in normalized for token in ("name", "object", "category", "status", "type")):
                width = 20
            elif "id" in normalized or "port" in normalized or "vlan" in normalized:
                width = 14
            else:
                width = 18
        if wide:
            width += 10

        sheet.column_dimensions[get_column_letter(index)].width = width


def _freeze_pane(sheet_name: str, headers: Sequence[str]) -> str:
    if sheet_name == "Interfaces":
        return "D4"
    if sheet_name == "Policies":
        return "D4"
    if len(headers) > 12:
        return "C4"
    return "A4"


def _sheet_note(
    sheet_name: str,
    row_count: int,
    profile: ExcelExportProfile = ExcelExportProfile.FULL,
) -> str:
    if sheet_name == "FortiGate Source Inventory":
        if profile is ExcelExportProfile.FAST:
            return (
                f"{row_count} selected source row(s). FAST retains source-only, model-gap, "
                "raw-extra, and unsupported evidence; fully typed command duplication is omitted."
            )
        return (
            "Explicit FortiGate source commands retained for traceability. "
            "Extraction Status identifies whether the source section has "
            "dedicated typed extraction support."
        )

    if sheet_name == "Policies" and row_count:
        return (
            f"{row_count} record(s). {_POLICY_ACTION_NOTE} "
            "Other values are explicit source data unless identified as derived or analysis output."
        )

    if row_count:
        return (
            f"{row_count} record(s). "
            "Values are explicit FortiGate source data unless the column is "
            "identified as derived or analysis output."
        )

    return "No matching explicit FortiGate source records were extracted."


def _inventory_counts(context: _ExcelContext) -> list[tuple[str, int, str]]:
    config = context.config
    source_only = context.source_only_counts
    counts = [
        ("Source Interfaces", len(config.interfaces), "Interfaces"),
        ("Zones", len(config.zones), "Zones"),
        ("Addresses", len(config.addresses), "Addresses"),
        ("Address Groups", len(config.address_groups), "Address Groups"),
        ("Services", len(context.derived.services.services), "Services"),
        ("Service Groups", len(context.derived.services.groups), "Service Groups"),
        ("Policies", len(config.policies), "Policies"),
        ("Security Policies", len(config.security_policies), "Security Policies"),
        ("NAT Rules", len(context.derived.nat), "NAT Rules"),
        ("IP Pools", len(config.ip_pools) + len(config.ip_pools6), "IP Pools"),
        ("Virtual IPs", len(config.vips) + len(config.vips6), "Virtual IPs"),
        ("Routes", len(config.static_routes), "Routes"),
        ("VPN Tunnels", len(config.ipsec_phase1), "VPN Tunnels"),
        ("Policy IPsec Phase 1", len(config.ipsec_policy_phase1), "Policy IPsec Phase 1"),
        ("Policy IPsec Phase 2", len(config.ipsec_policy_phase2), "Policy IPsec Phase 2"),
        ("VPN Phase 2", len(config.ipsec_phase2), "VPN Phase 2"),
        ("Local Users", len(config.local_users), "Local Users"),
        ("DHCP Servers", len(config.dhcp_servers), "DHCP Servers"),
        ("SD-WAN Members", sum(len(item.members) for item in config.sdwans), "SD-WAN Members"),
        ("SSL VPN Portals", len(config.ssl_vpn_portals), "SSL VPN Portals"),
        ("SSL VPN Realms", len(config.ssl_vpn_realms), "SSL VPN Realms"),
        ("SSL VPN Clients", len(config.ssl_vpn_clients), "SSL VPN Clients"),
        ("Administrators", len(config.administrators), "Administrators"),
        ("IPS Sensors", len(config.ips_sensors), "IPS Sensors"),
        ("External Resources", len(config.external_resources), "External Resources"),
        ("Review Required", len(context.validation.issues), "Review Required"),
        ("Source-only Sections", len(source_only), "Unsupported"),
    ]
    return counts


def _hostname(context: _ExcelContext) -> str | None:
    for record in context.source_by_path.get("system global", ()):
        for key, value in record.values.items():
            if _normalize_key(key) == "hostname":
                return value
    return None


def _migration_indicators(context: _ExcelContext) -> list[tuple[str, Any]]:
    config = context.config
    topology_issues = sum(
        len(item.issues)
        for item in (
            *context.derived.topology.interfaces,
            *context.derived.topology.vpns,
        )
    )
    nat_review_items = sum(
        issue.domain == "nat"
        for issue in context.validation.issues
    )
    dynamic_wan = any(
        (interface.mode or "").lower() in {"dhcp", "pppoe"}
        for interface in config.interfaces
    )
    return [
        ("IPv6 Explicit Configuration Present", "Yes" if context.ipv6_explicit else "No"),
        ("Dynamic WAN Addressing Present", "Yes" if dynamic_wan else "No"),
        ("SD-WAN Present", "Yes" if config.sdwans else "No"),
        ("NAT Review Items", nat_review_items),
        ("Broken References", len(context.derived.broken_references)),
        ("Topology Issues", topology_issues),
    ]


def _vdoms(context: _ExcelContext) -> list[str]:
    values = set(context.vdoms)
    values.update(
        getattr(item, "vdom", "root")
        for collection_name in (
            "interfaces",
            "zones",
            "addresses",
            "policies",
            "static_routes",
            "ipsec_phase1",
        )
        for item in getattr(context.config, collection_name, ())
    )
    return sorted(values)
