from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from fwmigrate.extraction.models import ExtractionStatus
from fwmigrate.extraction.sanitize import sanitize_raw_text, sanitize_source_attributes, sanitize_source_value
from ....source_reporting.options import ExcelExportProfile
from ....source_reporting.excel_style import append_report_row, finish_report_workbook, write_report_summary
from ....source_reporting.collection_excel import collection_summary, write_collection_evidence

from .excel_schema import SHEET_HEADERS, SHEET_ORDER
from ..transforms.effective_state import effective_path_state


def _value(value: Any) -> Any:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value) if isinstance(value, (list, tuple, dict)) else value


def _inheritance_rows(view: dict) -> tuple[dict, ...]:
    rows = [{"record_type": "statement", **item}
            for item in view.get("effective_statements", ())]
    for item in view.get("candidates", ()):
        provenance = item.get("provenance") or {}
        context = provenance.get("source_context") or {}
        rows.append({"record_type": "candidate", "context": str(context.get("name") or context.get("context_type") or "root"),
                     "origin": provenance.get("provenance_kind", "group"), "status": item.get("status"),
                     "target_path": item.get("target_path") or provenance.get("target_path"),
                     "source_path": provenance.get("source_path"),
                     "source_group": provenance.get("source_group_name"),
                     "group_chain": provenance.get("source_group_chain", ()),
                     "source_order": item.get("source_order"), "active": item.get("effective"),
                     "value": item.get("value")})
    rows.extend({"record_type": "issue", "context": item.get("context"), "origin": "group",
                 "status": item.get("status"), "target_path": item.get("target_path"),
                 "source_order": item.get("source_order"), "active": False}
                for item in view.get("issues", ()))
    return tuple(rows)


def _additional_settings(value: Any, context: str, path: tuple[str, ...] = ()) -> list[tuple[Any, ...]]:
    if hasattr(value, "model_dump"):
        value = value.model_dump(mode="json")
    rows = []
    if isinstance(value, dict):
        for key, child in value.items():
            if key in {"raw_extra", "settings", "source_attributes"} and isinstance(child, dict):
                for setting, setting_value in child.items():
                    safe = sanitize_source_value(str(setting), setting_value)
                    rows.append((context, ".".join(path[:-1]), path[-1] if path else "source",
                                 setting, json.dumps(safe, sort_keys=True, ensure_ascii=False, default=str)))
            elif key not in {"field_candidate_history", "non_effective_candidate_history", "field_provenance"}:
                rows.extend(_additional_settings(child, context, (*path, str(key))))
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            rows.extend(_additional_settings(child, context, (*path, str(index))))
    return rows


def _coverage_row(section: Any, inventory: dict[str, Any]) -> tuple[Any, ...]:
    commands = inventory.get(section.path).commands if section.path in inventory else ()
    counts = {status: sum(command.status == status for command in commands)
              for status in ExtractionStatus}
    return (section.source_context, section.path, len(commands), counts[ExtractionStatus.EXTRACTED],
            counts[ExtractionStatus.PARTIAL], counts[ExtractionStatus.SOURCE_ONLY],
            counts[ExtractionStatus.UNSUPPORTED], counts[ExtractionStatus.UNKNOWN], counts[ExtractionStatus.IGNORED],
            counts[ExtractionStatus.PARSE_ERROR],
            section.unresolved_dependencies, section.status.value, section.status != ExtractionStatus.EXTRACTED)


def _unresolved_reference_row(item: Any) -> dict[str, Any]:
    context_type, _, context_name = (item.source_context or "").partition(" ")
    return {"context_type": context_type or None, "context": context_name or item.source_context,
            "source_path": item.source_path, "source_object": item.source_object,
            "source_field": item.source_field, "reference": item.reference,
            "expected_type": item.expected_type, "resolution": item.result,
            "reason": item.notes or item.reason}


def _context_scope(context: Any) -> str:
    return "root" if context.context_type == "root" else f"{context.context_type} {context.name}"


def _native_source_rows(
    config: Any,
    inheritance_view: dict[str, Any],
    policy_relationships: Any = (),
) -> dict[str, tuple[dict[str, Any], ...]]:
    rows: dict[str, list[dict[str, Any]]] = {
        "Addresses": [], "Address Sets": [], "Application Sets": [], "Application Settings": [], "Schedulers": [],
        "Interface Details": [], "Interface Addresses": [],
        "Policy Details": [], "NAT Pools": [], "NAT Rules": [],
        "Administrators": [], "Login Classes": [], "IDP Policies": [], "IDP Rules": [],
        "UTM Policies": [], "UTM Profiles": [],
        "IKE Proposals": [], "IKE Policies": [], "IKE Gateways": [],
        "IPsec Proposals": [], "IPsec Policies": [], "IPsec VPNs": [],
        "Traffic Selectors": [], "Static Routes": [],
    }
    policy_orders = {}
    for view in policy_relationships:
        for group in view["zone_policy_sets"]:
            for item in group["policies"]:
                key = (item["context"], "zone", item["from_zone"], item["to_zone"], item["name"])
                policy_orders[key] = item["order"]
        for item in view["global_policies"]:
            policy_orders[(item["context"], "global", item["from_zone"], item["to_zone"], item["name"])] = item["order"]

    for item in config.admin_users.values():
        rows["Administrators"].append({"context": "root", "name": item.name,
                                       "login_class": item.login_class, "settings": dict(item.settings)})
    for item in config.login_classes.values():
        rows["Login Classes"].append({"context": "root", "name": item.name, "settings": dict(item.settings)})

    for context in config.iter_contexts():
        scope = _context_scope(context)
        state = lambda path: effective_path_state(inheritance_view, scope, path)

        for interface in context.interfaces.values():
            path = ("interfaces", interface.name)
            rows["Interface Details"].append({
                "context": scope, "record_type": "interface", "interface": interface.name,
                "description": interface.description, "disabled": interface.disabled, "mtu": interface.mtu,
                "speed": interface.speed, "link_mode": interface.link_mode, "encapsulation": interface.encapsulation,
                "effective_state": state(path),
            })
            for unit in interface.units.values():
                unit_path = (*path, "unit", unit.unit)
                rows["Interface Details"].append({
                    "context": scope, "record_type": "unit", "interface": interface.name, "unit": unit.unit,
                    "description": unit.description, "disabled": unit.disabled, "mtu": unit.mtu,
                    "encapsulation": unit.encapsulation, "vlan_id": unit.vlan_id,
                    "effective_state": state(unit_path),
                })
                for address in unit.addresses:
                    rows["Interface Addresses"].append({
                        "context": scope, "interface": interface.name, "unit": unit.unit,
                        "family": address.family, "address": address.address,
                        "primary": address.primary, "preferred": address.preferred,
                        "effective_state": state((*unit_path, "family", address.family, "address", address.address)),
                    })

        for application in context.applications.values():
            path = ("applications", "application", application.name)
            top = application.top_level
            has_top_level = application.description is not None or bool(top.source_attributes) or any(
                getattr(top, field) is not None for field in (
                    "protocol", "protocol_number", "icmp_type", "icmp_code", "application_protocol", "inactivity_timeout")
            ) or bool(top.source_ports or top.destination_ports)
            for term in (*((top,) if has_top_level else ()), *application.terms):
                term_name = getattr(term, "name", None)
                rows["Application Settings"].append({
                    "context": scope, "application": application.name,
                    "location": "term" if term_name is not None else "top-level", "term": term_name,
                    "protocol": term.protocol, "protocol_number": term.protocol_number,
                    "source_ports": list(term.source_ports) or None,
                    "destination_ports": list(term.destination_ports) or None,
                    "icmp_type": term.icmp_type, "icmp_code": term.icmp_code,
                    "application_protocol": term.application_protocol, "inactivity_timeout": term.inactivity_timeout,
                    "description": application.description if term_name is None else None,
                    "effective_state": state((*path, "term", term_name) if term_name is not None else path),
                    "additional_settings": sanitize_source_attributes(term.source_attributes) or None,
                })

        for book in context.address_books.values():
            for item in book.addresses.values():
                rows["Addresses"].append({
                    "context": scope, "address_book": book.name, "name": item.name, "type": item.type,
                    "prefix": item.prefix, "fqdn": item.fqdn, "range_start": item.range_start,
                    "range_end": item.range_end, "wildcard": item.wildcard, "zone": item.zone,
                    "description": item.description,
                    "effective_state": state(("security", "address-book", book.name, "address", item.name)),
                })
            for item in book.address_sets.values():
                rows["Address Sets"].append({
                    "context": scope, "address_book": book.name, "name": item.name,
                    "members": [member.name for member in item.members],
                    "member_types": [member.member_type for member in item.members],
                    "zone": item.zone, "description": item.description,
                    "effective_state": state(("security", "address-book", book.name, "address-set", item.name)),
                })

        for item in context.application_sets.values():
            rows["Application Sets"].append({
                "context": scope, "name": item.name, "applications": list(item.applications),
                "description": item.description,
                "effective_state": state(("applications", "application-set", item.name)),
            })

        for item in context.schedulers.values():
            rows["Schedulers"].append({
                "context": scope, "name": item.name, "start_date": item.start_date,
                "stop_date": item.stop_date, "daily": list(item.daily),
                "weekdays": dict(item.weekdays), "daily_windows": list(item.daily_windows),
                "weekday_windows": dict(item.weekday_windows), "exclusions": list(item.exclusions),
                "description": item.description,
                "effective_state": state(("schedulers", "scheduler", item.name)),
            })

        for item in (*context.policies, *context.global_policies):
            policy_scope = item.policy_scope
            from_zone = item.from_zone if policy_scope == "zone" else tuple(item.from_zones) or None
            to_zone = item.to_zone if policy_scope == "zone" else tuple(item.to_zones) or None
            order = policy_orders.get((scope, policy_scope, from_zone, to_zone, item.name))
            policy_path = (
                ("security", "policies", "global", "policy", item.name)
                if item.policy_scope == "global"
                else ("security", "policies", "from-zone", item.from_zone, "to-zone", item.to_zone, "policy", item.name)
            )
            rows["Policy Details"].append({
                "context": scope, "policy_scope": item.policy_scope, "name": item.name,
                "order": order, "from_zones": list(item.from_zones),
                "to_zones": list(item.to_zones), "source_addresses": list(item.source_addresses),
                "destination_addresses": list(item.destination_addresses),
                "applications": list(item.applications),
                "dynamic_applications": list(item.dynamic_applications),
                "source_identities": list(item.source_identities), "scheduler": item.scheduler_name,
                "action": item.action, "log_session_init": item.log_session_init,
                "log_session_close": item.log_session_close,
                "application_services": list(item.application_services),
                "security_profile_references": dict(item.security_profile_references),
                "description": item.description, "effective_state": state(policy_path),
            })

        for item in context.idp_policies.values():
            rows["IDP Policies"].append({"context": scope, "name": item.name,
                                         "rulebases": list(item.rulebase),
                                         "source_attributes": sanitize_source_attributes(item.source_attributes)})
            for rulebase, rules in item.rulebase.items():
                for order, rule in enumerate(rules):
                    rows["IDP Rules"].append({"context": scope, "policy": item.name, "rulebase": rulebase,
                                              "name": rule.name, "order": order, "match": dict(rule.match),
                                              "exceptions": list(rule.exceptions), "action": rule.action,
                                              "severity": list(rule.severity),
                                              "source_attributes": sanitize_source_attributes(rule.source_attributes)})
        for item in context.utm_policies.values():
            rows["UTM Policies"].append({"context": scope, "name": item.name,
                                         "settings": dict(item.settings)})
        for profile_type, profiles in (("anti-virus", context.antivirus_profiles),
                                       ("web-filtering", context.web_filtering_profiles),
                                       ("content-filtering", context.content_filtering_profiles),
                                       ("anti-spam", context.anti_spam_profiles)):
            for item in profiles.values():
                rows["UTM Profiles"].append({"context": scope, "profile_type": profile_type,
                                             "name": item.name,
                                             "settings": item.model_dump(mode="python", exclude={"name", "source_attributes"}),
                                             "source_attributes": sanitize_source_attributes(item.source_attributes)})

        for nat_type, pools in (("source", context.nat.source_pools), ("destination", context.nat.destination_pools)):
            for item in pools.values():
                rows["NAT Pools"].append({
                    "context": scope, "nat_type": nat_type, "name": item.name,
                    "routing_instance": item.routing_instance, "addresses": list(item.addresses),
                    "address_ranges": list(item.address_ranges), "ports": list(item.ports),
                    "options": dict(item.options),
                    "effective_state": state(("security", "nat", nat_type, "pool", item.name)),
                })

        for nat_type, rule_sets in (
            ("source", context.nat.source_rule_sets),
            ("destination", context.nat.destination_rule_sets),
            ("static", context.nat.static_rule_sets),
        ):
            for rule_set in rule_sets.values():
                to_context = rule_set.to_context
                for order, item in enumerate(rule_set.rules, 1):
                    rows["NAT Rules"].append({
                        "context": scope, "nat_type": nat_type, "rule_set": rule_set.name,
                        "rule": item.name, "order": order,
                        "from_zones": list(rule_set.from_context.zones),
                        "from_interfaces": list(rule_set.from_context.interfaces),
                        "from_routing_instances": list(rule_set.from_context.routing_instances),
                        "to_zones": list(to_context.zones) if to_context else [],
                        "to_interfaces": list(to_context.interfaces) if to_context else [],
                        "to_routing_instances": list(to_context.routing_instances) if to_context else [],
                        "source_addresses": list(item.match.source_addresses),
                        "destination_addresses": list(item.match.destination_addresses),
                        "source_address_names": list(item.match.source_address_names),
                        "destination_address_names": list(item.match.destination_address_names),
                        "source_ports": list(item.match.source_ports),
                        "destination_ports": list(item.match.destination_ports),
                        "protocols": list(item.match.protocols), "applications": list(item.match.applications),
                        "action": dict(item.action), "disabled": item.disabled, "description": item.description,
                        "effective_state": state(("security", "nat", nat_type, "rule-set", rule_set.name, "rule", item.name)),
                    })

        vpn = context.vpn
        rows["IKE Proposals"].extend({
            "context": scope, "name": item.name, "authentication_method": item.authentication_method,
            "dh_group": item.dh_group, "authentication_algorithm": item.authentication_algorithm,
            "encryption_algorithm": item.encryption_algorithm,
            "digital_signature_scheme": item.digital_signature_scheme, "prf_algorithm": item.prf_algorithm,
            "signature_hash_algorithm": item.signature_hash_algorithm,
            "lifetime_seconds": item.lifetime_seconds, "description": item.description,
            "effective_state": state(("security", "ike", "proposal", item.name)),
        } for item in vpn.ike_proposals.values())
        rows["IKE Policies"].extend({
            "context": scope, "name": item.name, "mode": item.mode, "proposal_set": item.proposal_set,
            "proposals": list(item.proposals), "pre_shared_key_configured": item.has_pre_shared_key,
            "certificate_reference": item.certificate_reference, "local_certificate": item.local_certificate,
            "effective_state": state(("security", "ike", "policy", item.name)),
        } for item in vpn.ike_policies.values())
        rows["IKE Gateways"].extend({
            "context": scope, "name": item.name, "ike_policy": item.ike_policy, "address": item.address,
            "external_interface": item.external_interface, "version": item.version,
            "local_address": item.local_address, "local_identity": item.local_identity,
            "remote_identity": item.remote_identity, "nat_traversal": item.nat_traversal,
            "dpd": dict(item.dpd), "certificate_reference": item.certificate_reference,
            "effective_state": state(("security", "ike", "gateway", item.name)),
        } for item in vpn.ike_gateways.values())
        rows["IPsec Proposals"].extend({
            "context": scope, "name": item.name, "protocol": item.protocol,
            "authentication_algorithm": item.authentication_algorithm,
            "encryption_algorithm": item.encryption_algorithm, "lifetime_seconds": item.lifetime_seconds,
            "lifetime_kilobytes": item.lifetime_kilobytes, "description": item.description,
            "effective_state": state(("security", "ipsec", "proposal", item.name)),
        } for item in vpn.ipsec_proposals.values())
        rows["IPsec Policies"].extend({
            "context": scope, "name": item.name, "proposal_set": item.proposal_set,
            "proposals": list(item.proposals), "pfs_group": item.pfs_group,
            "effective_state": state(("security", "ipsec", "policy", item.name)),
        } for item in vpn.ipsec_policies.values())
        for item in vpn.ipsec_vpns.values():
            rows["IPsec VPNs"].append({
                "context": scope, "name": item.name, "bind_interface": item.bind_interface,
                "ike_gateway": item.ike_gateway, "ipsec_policy": item.ipsec_policy,
                "establish_tunnels": item.establish_tunnels,
                "vpn_monitor_destination": item.vpn_monitor.destination_ip if item.vpn_monitor else None,
                "vpn_monitor_source_interface": item.vpn_monitor.source_interface if item.vpn_monitor else None,
                "effective_state": state(("security", "ipsec", "vpn", item.name)),
            })
            for selector in item.traffic_selectors.values():
                rows["Traffic Selectors"].append({
                    "context": scope, "vpn": item.name, "name": selector.name, "term": None,
                    "local_ip": list(selector.local_ip), "remote_ip": list(selector.remote_ip),
                    "protocol": selector.protocol, "local_port": list(selector.local_port),
                    "remote_port": list(selector.remote_port),
                    "effective_state": state(("security", "ipsec", "vpn", item.name, "traffic-selector", selector.name)),
                })
                for term in selector.terms.values():
                    rows["Traffic Selectors"].append({
                        "context": scope, "vpn": item.name, "name": selector.name, "term": term.name,
                        "local_ip": list(term.local_ip), "remote_ip": list(term.remote_ip),
                        "protocol": term.protocol, "local_port": list(term.local_port),
                        "remote_port": list(term.remote_port),
                        "effective_state": state(("security", "ipsec", "vpn", item.name, "traffic-selector", selector.name, "term", term.name)),
                    })

        for item in context.routes:
            route_path = (
                ("routing-instances", item.routing_instance, "routing-options", "static", "route", item.destination)
                if item.routing_instance else
                ("routing-options", "rib", item.rib, "static", "route", item.destination)
                if item.rib else
                ("routing-options", "static", "route", item.destination)
            )
            hops = item.next_hops or (None,)
            for hop in hops:
                rows["Static Routes"].append({
                    "context": scope, "routing_instance": item.routing_instance, "rib": item.rib,
                    "destination": item.destination, "next_hop": hop.value if hop else None,
                    "qualified": hop.qualified if hop else None, "next_table": item.next_table,
                    "route_preference": item.preference, "route_metric": item.metric, "route_tag": item.tag,
                    "next_hop_preference": hop.preference if hop else None,
                    "next_hop_metric": hop.metric if hop else None, "next_hop_tag": hop.tag if hop else None,
                    "action": item.action, "retain": item.retain, "installation": item.installation,
                    "disabled": item.disabled, "effective_state": state(route_path),
                })
    return {name: tuple(values) for name, values in rows.items()}


def export_juniper_excel(
    result: Any,
    output: Any,
    *,
    profile: ExcelExportProfile | str = ExcelExportProfile.FULL,
    source_name: str | None = None,
    collection_status: str | None = None,
    collection_parts: Any = None,
    collection_warnings: Any = None,
) -> Any:
    from openpyxl import Workbook

    profile = ExcelExportProfile(profile)
    workbook = Workbook(write_only=profile is ExcelExportProfile.FAST)
    if not workbook.write_only:
        workbook.remove(workbook.active)
    views = result.derived
    rows = {"Interfaces": views.interface_topology, "Zones": views.zone_memberships,
            "Routing Instances": views.routing_instances, "Address Books": views.address_books,
            "Applications": views.applications, "Policies": views.policies, "NAT": views.nat_rule_sets,
            "VPN": views.vpn_relationships,
            **_native_source_rows(result.config, views.inheritance_view, views.policy_relationships),
            "DHCP Local Servers": tuple(item for item in views.dhcp if item["kind"] == "local-server"),
            "DHCP Relay Groups": tuple(item for item in views.dhcp if item["kind"] == "relay-group"),
            "DHCP Pools": tuple(item for item in views.dhcp if item["kind"] == "address-assignment-pool"),
            "Access Profiles": views.access_profiles, "Firewall Users": views.firewall_users,
            "APBR": views.apbr, "Remote Access": views.remote_access,
            "Policy Relationships": tuple(row for item in views.policy_relationships
                                           for group in item["zone_policy_sets"] for row in group["policies"])
                                  + tuple(row for item in views.policy_relationships for row in item["global_policies"]),
            "Policy Reference Relationships": tuple(edge for item in views.policy_relationships
                                                      for edge in item["edges"]),
            "NAT Usage": views.nat_usage, "NAT Pool Usage": views.nat_pool_usage, "VPN Relationships": views.vpn_graph,
            "Secure Connect": views.secure_connect_graph, "APBR Relationships": views.apbr_graph,
            "Inheritance": _inheritance_rows(views.inheritance_view),
            "Effective Objects": views.effective_objects,
            "Review Required": result.review_required,
            "Unresolved References": tuple(_unresolved_reference_row(item) for item in views.dependencies
                                           if item.result == "UNRESOLVED")}
    inventory_by_path = {item.source_path: item for item in result.inventory_items}
    additional = []
    for context in result.config.iter_contexts():
        scope = "root" if context.context_type == "root" else f"{context.context_type} {context.name}"
        additional.extend(_additional_settings(context, scope, (context.name,)))
    additional.extend(_additional_settings(result.config.model_dump(
        mode="json", exclude={"contexts", "configuration_groups", "field_provenance",
                               "field_candidate_history", "non_effective_candidate_history",
                               "activation_directives", "unsupported_commands"}), "root"))
    rows["Additional Settings"] = tuple(additional)
    for name in SHEET_ORDER:
        if profile is ExcelExportProfile.FAST and name in {"Source Inventory", "Extraction Coverage"}:
            continue
        sheet = workbook.create_sheet(name)
        if name == "Summary":
            continue
        if name == "Source Inventory":
            append_report_row(sheet, SHEET_HEADERS[name])
            for item in result.inventory_items:
                if not item.commands:
                    append_report_row(sheet, (item.domain, item.source_path, item.source_context, None, None, None,
                                              None, item.status.value, None, item.requires_manual_review))
                    continue
                for command in item.commands:
                    append_report_row(sheet, (
                        item.domain, item.source_path, command.source_context or item.source_context,
                        command.line_number, command.operation, command.key, _value(command.values),
                        command.status.value if command.status else None, command.parser_handler,
                        command.requires_manual_review,
                    ))
            continue
        if name == "Review Required":
            append_report_row(sheet, SHEET_HEADERS[name])
            for item in result.review_required:
                append_report_row(sheet, (item.severity, item.code, item.category, item.message, item.context,
                                          item.source_path, item.object_name, item.field))
            continue
        if name == "Validation":
            append_report_row(sheet, SHEET_HEADERS[name])
            for issue in result.validation.issues:
                append_report_row(sheet, (issue.code, issue.severity, issue.category, issue.message, issue.context_type,
                              issue.context, issue.source_path, issue.object_type, issue.object_name,
                              issue.field, issue.reference, issue.expected_type))
            continue
        if name == "Extraction Coverage":
            append_report_row(sheet, SHEET_HEADERS[name])
            for section in result.source_sections:
                append_report_row(sheet, _coverage_row(section, inventory_by_path))
            continue
        if name == "Unsupported Source":
            append_report_row(sheet, SHEET_HEADERS[name])
            for item in result.unsupported_items:
                append_report_row(sheet, (item.source_context, item.source_path, item.source_name, item.reason, item.raw_capture))
            continue
        append_report_row(sheet, SHEET_HEADERS[name])
        for item in rows[name]:
            if isinstance(item, tuple):
                append_report_row(sheet, tuple(_value(value) for value in item))
                continue
            if hasattr(item, "model_dump"):
                item = item.model_dump(mode="python")
            elif hasattr(item, "__dict__"):
                item = item.__dict__
            append_report_row(sheet, tuple(_value(item.get(header.lower().replace(" ", "_"))) for header in SHEET_HEADERS[name]))
    coverage = []
    if collection_status is not None:
        parts = tuple(collection_parts or ())
        warnings = tuple(collection_warnings or ())
        write_collection_evidence(workbook, collection_status, parts, warnings)
        coverage.extend((*collection_summary(collection_status),
                         ("Incomplete Collection Parts", sum(
                             not (part.get("complete") if isinstance(part, dict) else getattr(part, "complete", False))
                             for part in parts))))
        if warnings:
            coverage.append(("Collection Warnings", " | ".join(sanitize_raw_text(str(item)) for item in warnings)))
    write_report_summary(
        workbook["Summary"], title="Juniper SRX Configuration Report", profile=profile.value.upper(),
        metadata=(("Hostname", result.config.hostname), ("Version", None),
                  ("Scopes", "\n".join(f"{context.context_type}: {context.name}"
                                       for context in result.config.iter_contexts()) or None),
                  ("Source File", source_name), ("Source Format", result.source_format)),
        validation=(("Errors", len(result.validation.errors)), ("Warnings", len(result.validation.warnings)),
                    ("Unresolved References", sum(item.result == "UNRESOLVED" for item in views.dependencies))),
        inventory=((name, workbook[name]._report_row_count - 1, name) for name in SHEET_ORDER
                   if name in workbook.sheetnames and name not in {
                       "Summary", "Review Required", "Validation", "Source Inventory",
                       "Extraction Coverage", "Unsupported Source", "Unresolved References", "Additional Settings",
                   }),
        coverage=(*coverage, ("Contexts", len(result.config.contexts)),
                  ("Configuration Groups", len(result.config.configuration_groups)),
                  ("Source Sections", len(result.source_sections)),
                  *((label, sum(item.status == status for item in result.source_sections)) for label, status in (
                      ("Partial Sections", ExtractionStatus.PARTIAL),
                      ("Source Only Sections", ExtractionStatus.SOURCE_ONLY),
                      ("Unsupported Sections", ExtractionStatus.UNSUPPORTED),
                      ("Parse Errors", ExtractionStatus.PARSE_ERROR))),
                  ("Export Limitations", "FAST omits Source Inventory and Extraction Coverage. Use FULL for traceability."
                   if profile is ExcelExportProfile.FAST else "FULL includes Source Inventory and Extraction Coverage.")),
    )
    finish_report_workbook(workbook)
    if hasattr(output, "write"):
        workbook.save(output)
        return output
    path = Path(output)
    workbook.save(path)
    return path


__all__ = ["export_juniper_excel"]
