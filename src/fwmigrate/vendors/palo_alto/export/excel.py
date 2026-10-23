from __future__ import annotations

from itertools import chain
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Mapping

from ....source_reporting.excel_lifecycle import report_workbook, save_report_workbook

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ....source_reporting.options import ExcelExportProfile
from ....source_reporting.collection_excel import collection_summary, write_collection_evidence
from ....source_reporting.excel_style import check_report_dimensions, REPORT_BODY_NOTE, excel_safe_value, report_cell_value, append_readable_report_row, report_column_width, style_report_header, write_report_summary

from .excel_rows import ROW_BUILDERS, _PANExcelContext
from .excel_schema import SHEET_ORDER, HIDDEN_COLUMNS_BY_DEFAULT, SHEET_HEADERS, OWNER_COLUMNS
from ..source_report import PaloAltoSourceResult

_TITLE_FILL = PatternFill("solid", fgColor="17324D")
_TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
_MUTED_FONT = Font(color="64748B", italic=True)
_LINK_FONT = Font(color="0563C1", underline="single")
_TITLE_ALIGNMENT = Alignment(vertical="center")
_LINK_ALIGNMENT = Alignment(horizontal="right")


def _excel_safe(value: Any) -> Any:
    return excel_safe_value(value)


def _write_summary(
    workbook: Workbook, context: _PANExcelContext, row_counts: Mapping[str, int], profile: ExcelExportProfile,
) -> None:
    sheet = workbook["Summary"] if "Summary" in workbook.sheetnames else workbook.create_sheet("Summary")
    coverage_sheets = {"Review Required", "Unresolved References", "Unsupported",
                       "PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"}
    write_report_summary(
        sheet, title="PAN-OS Configuration Report",
        metadata=(("Hostname", context.config.hostname), ("PAN-OS Version", context.config.source_version),
                  ("Scopes", "\n".join(context.derived.scope_identities) or None),
                  ("Source File", context.source_name)),
        profile=profile.value.upper(),
        validation=(("Errors", len(context.validation.errors)),
                    ("Warnings", len(context.validation.warnings)),
                    ("Unresolved References", len(context.derived.unresolved_references))),
        inventory=((name, count, name) for name, count in row_counts.items()
                   if name not in coverage_sheets),
        coverage=(*collection_summary(context.collection_status),
                  *((name, count, name) for name, count in row_counts.items() if name in coverage_sheets),
                  ("Source Inventory Records", len(context.config.source_inventory)),
                  ("Extraction Issues", len(context.config.extraction_issues)),
                  *((f"{status.replace('_', ' ').title()} Records", sum(
                      value == status for values in context.domain_statuses.values() for value in values))
                    for status in ("PARTIAL", "SOURCE_ONLY", "UNSUPPORTED")),
                  ("Export Limitations", (
                      "FAST omits PAN-OS Source Inventory, Extraction Coverage, and PAN-OS Source Appendix. "
                      "Use FULL for complete source traceability." if profile is ExcelExportProfile.FAST else
                      "FULL includes PAN-OS Source Inventory, Extraction Coverage, and PAN-OS Source Appendix."
                  ))),
    )


def _write_table(
    workbook: Workbook, name: str, rows: Iterable[Mapping[str, Any]], *, full: bool = False,
) -> int:
    headers = SHEET_HEADERS[name]
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_format.defaultRowHeight = 18
    sheet.row_dimensions[3].height = 30
    detail_columns = sum(header not in OWNER_COLUMNS for header in headers)
    sheet.freeze_panes = "D4" if name in {"Interfaces", "Security Policies", "NAT Rules"} else "C4" if detail_columns > 12 else "A4"
    for column, header in enumerate(headers, 1):
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.width = report_column_width(header)
        dimension.hidden = header in HIDDEN_COLUMNS_BY_DEFAULT.get(name, ())
    if full:
        sheet.merged_cells.add(f"A1:{get_column_letter(len(headers))}1")
        if len(headers) > 2:
            sheet.merged_cells.add(f"A2:{get_column_letter(len(headers) - 1)}2")
        sheet.row_dimensions[1].height = 24
    title = WriteOnlyCell(sheet, name)
    title.fill, title.font = _TITLE_FILL, _TITLE_FONT
    title.alignment = _TITLE_ALIGNMENT
    sheet.append([title])
    note = WriteOnlyCell(sheet, (
        "FULL report. Row counts are listed in Summary. " if full else "FAST lightweight export. "
    ) + "Analysis Status reports validation only; NO_VALIDATION_ISSUES does not imply complete extraction. "
        "Source Inventory reports extraction coverage." + (" Source Inventory is available in FULL." if not full else "") + " " + REPORT_BODY_NOTE)
    note.font = _MUTED_FONT
    values = [note]
    if full and len(headers) > 1:
        link = WriteOnlyCell(sheet, "Back to Summary")
        link.hyperlink, link.font, link.alignment = "#'Summary'!A1", _LINK_FONT, _LINK_ALIGNMENT
        values = [note, *([None] * (len(headers) - 2)), link]
    sheet.append(values)
    cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, header)
        style_report_header(cell)
        cells.append(cell)
    sheet.append(cells)
    row_count = 0
    for row_count, row in enumerate(rows, 1):
        check_report_dimensions(row_count + 3, len(headers))
        values = [report_cell_value(sheet, row_count + 3, column, row.get(header))
                  for column, header in enumerate(headers, 1)]
        if name not in {"PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"}:
            append_readable_report_row(sheet, row_count + 3, values)
        else:
            sheet.append(values)
    if row_count:
        sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{row_count + 3}"
    return row_count


def _build_workbook(context: _PANExcelContext, profile: ExcelExportProfile, workbook: Workbook | None = None) -> Workbook:
    workbook = workbook if workbook is not None else Workbook(write_only=True)
    workbook.create_sheet("Summary")
    counts = {}
    no_rows = object()
    for name in SHEET_ORDER:
        if name == "Summary":
            continue
        if profile is ExcelExportProfile.FULL:
            counts[name] = _write_table(workbook, name, ROW_BUILDERS[name](context), full=True)
            continue
        if name in {"PAN-OS Source Inventory", "Extraction Coverage", "PAN-OS Source Appendix"}:
            continue
        rows = iter(ROW_BUILDERS[name](context))
        first = next(rows, no_rows)
        if first is no_rows and name != "Review Required":
            continue
        counts[name] = _write_table(workbook, name, () if first is no_rows else chain((first,), rows))
    _write_summary(workbook, context, counts, profile)
    return workbook


def export_panos_excel(
    analysis: PaloAltoSourceResult, output: BinaryIO | str | Path, *,
    source_name: str | None = None, profile: ExcelExportProfile | str = ExcelExportProfile.FULL,
    collection_status: str | None = None, collection_parts: Iterable[Mapping[str, Any]] = (),
    collection_warnings: Iterable[str] = (),
) -> None:
    profile = ExcelExportProfile(profile)
    context = _PANExcelContext(analysis, source_name, collection_status=collection_status)
    with report_workbook(Workbook(write_only=True)) as workbook:
        _build_workbook(context, profile, workbook=workbook)
        write_collection_evidence(workbook, collection_status, collection_parts, collection_warnings)
        if isinstance(output, (str, Path)):
            Path(output).parent.mkdir(parents=True, exist_ok=True)
        save_report_workbook(workbook, output)
