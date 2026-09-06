"""Source-vendor visibility policy for Excel inventory workbooks.

This module intentionally operates only on the rendered workbook. It does not
parse vendor syntax, mutate canonical IR, or affect target generation.
"""

from __future__ import annotations

import io
from typing import Any

from fwmigrate.report.excel_exporter import IRExcelExporter as _BaseIRExcelExporter


_BASE_SHEET_ORDER = tuple(_BaseIRExcelExporter.SHEET_ORDER)


def _without_sheets(order: tuple[str, ...], excluded: frozenset[str]) -> tuple[str, ...]:
    return tuple(sheet_name for sheet_name in order if sheet_name not in excluded)


class VendorAwareIRExcelExporter(_BaseIRExcelExporter):
    """Hide source-vendor-inapplicable worksheets from the final workbook."""

    PALO_ALTO_ONLY_SHEETS = frozenset(
        {
            "Management Service Routes",
            "Security Profile Definitions",
            "Security Profile Rules",
            "Custom URL Categories",
            "GlobalProtect Portals",
            "GlobalProtect Gateways",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Root CAs",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
            "GlobalProtect Network Gateways",
            "PAN Log Servers",
            "PAN Log Forwarding",
            "PAN Log Forward Matches",
            "PAN DNS Proxies",
            "PAN DNS Proxy Domains",
            "PAN Monitor Profiles",
            "PAN QoS Profiles",
            "PAN QoS Classes",
            "PAN High Availability",
            "PAN HA Monitoring",
            "PAN Device Settings",
            "PAN VSYS Settings",
            "PAN Botnet Report",
            "PAN Custom Reports",
        }
    )

    # Keep this deliberately narrow. These sheets are explicitly FortiGate
    # source-detail views by name/implementation; shared canonical sheets stay
    # visible for other vendors.
    FORTIGATE_ONLY_SHEETS = frozenset(
        {
            "FortiGate Source Configuration",
            "Firewall Policy Source Settings",
            "Interface Nested Configuration",
            "FortiTokens",
        }
    )

    PALO_ALTO_SUMMARY_LABELS = frozenset(
        {
            "Management Service Routes",
            "Security Profile Definitions",
            "Security Profile Rules",
            "Custom URL Categories",
            "GlobalProtect Portals",
            "GlobalProtect Gateways",
            "GlobalProtect Network Gateways",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
        }
    )

    _VENDOR_ALIASES = {
        "fortigate": "fortigate",
        "fortinet": "fortigate",
        "fortios": "fortigate",
        "palo_alto": "palo_alto",
        "paloalto": "palo_alto",
        "panos": "palo_alto",
        "pan_os": "palo_alto",
        "cisco_asa": "cisco_asa",
        "cisco asa": "cisco_asa",
        "checkpoint": "checkpoint",
        "check_point": "checkpoint",
        "check point": "checkpoint",
        "juniper_srx": "juniper_srx",
        "juniper srx": "juniper_srx",
        "juniper": "juniper_srx",
    }
    _KNOWN_VENDORS = frozenset(_VENDOR_ALIASES.values())

    # Existing tests and callers historically inspect this class-level constant.
    # FortiGate is the legacy/default source vendor, so expose its active order
    # here while instance generation still resolves the actual source vendor.
    SHEET_ORDER = _without_sheets(_BASE_SHEET_ORDER, PALO_ALTO_ONLY_SHEETS)

    def _source_vendor(self) -> str:
        extraction_vendor = getattr(self.extraction, "source_vendor", None)
        metadata = getattr(self.ir, "metadata", None)
        ir_vendor = getattr(metadata, "source_vendor", None)
        raw_vendor = extraction_vendor or ir_vendor or ""
        normalized = str(raw_vendor).strip().lower().replace("-", "_")
        return self._VENDOR_ALIASES.get(normalized, normalized)

    def _active_sheet_order(self) -> tuple[str, ...]:
        vendor = self._source_vendor()

        # Unknown/legacy vendor identifiers keep the historical full workbook.
        # This avoids hiding source evidence when vendor identity is ambiguous.
        if vendor not in self._KNOWN_VENDORS:
            return _BASE_SHEET_ORDER

        excluded: set[str] = set()
        if vendor != "palo_alto":
            excluded.update(self.PALO_ALTO_ONLY_SHEETS)
        if vendor != "fortigate":
            excluded.update(self.FORTIGATE_ONLY_SHEETS)

        return tuple(
            sheet_name
            for sheet_name in _BASE_SHEET_ORDER
            if sheet_name not in excluded
        )

    def generate(self) -> bytes:
        """Generate the normal inventory, then apply source-vendor visibility."""
        # The base exporter currently builds every known worksheet before ordering.
        # Give it the historical complete order so its validation remains unchanged.
        self.SHEET_ORDER = _BASE_SHEET_ORDER
        workbook_bytes = super().generate()

        # Import only after base generation has confirmed the optional reports
        # dependency is available.
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(workbook_bytes))
        active_order = self._active_sheet_order()
        active_sheets = set(active_order)

        for worksheet in list(workbook.worksheets):
            if worksheet.title not in active_sheets:
                workbook.remove(worksheet)

        # Rebuild Summary after filtering so navigation and record counts only
        # reference worksheets that remain in this source-vendor workbook.
        if "Summary" in workbook.sheetnames:
            workbook.remove(workbook["Summary"])

        self.SHEET_ORDER = active_order
        self._build_summary(workbook)
        self._remove_inapplicable_summary_rows(workbook["Summary"])
        self._order_sheets(workbook)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _remove_inapplicable_summary_rows(self, summary: Any) -> None:
        vendor = self._source_vendor()
        if vendor not in self._KNOWN_VENDORS:
            return

        hidden_labels: set[str] = set()
        if vendor != "palo_alto":
            hidden_labels.update(self.PALO_ALTO_SUMMARY_LABELS)

        if not hidden_labels:
            return

        # Delete bottom-up so row indexes remain stable. Only Inventory Counts
        # labels match this set; navigation uses sheet names in column B.
        for row in range(summary.max_row, 1, -1):
            if summary.cell(row, 1).value in hidden_labels:
                summary.delete_rows(row, 1)
