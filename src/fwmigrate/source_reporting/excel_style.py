"""Lightweight presentation shared by vendor-native Excel reports."""

from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)


def style_fast_sheet(sheet, *, header_row: int = 1) -> None:
    """Style headers only; leave body cells on Excel's native grid."""
    sheet.sheet_view.showGridLines = True
    sheet.sheet_format.defaultColWidth = 28
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = 28
        cell = sheet.cell(header_row, column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"


def append_report_row(sheet, values) -> None:
    """Append vendor-owned values to either a normal or streaming worksheet."""
    if not sheet.parent.write_only:
        sheet.append(values)
        return
    count = getattr(sheet, "_report_row_count", 0)
    if count == 0:
        values = tuple(values)
        sheet._report_column_count = len(values)
        sheet.sheet_view.showGridLines = True
        sheet.sheet_format.defaultColWidth = 28
        sheet.freeze_panes = "A2"
        cells = []
        for column, value in enumerate(values, 1):
            sheet.column_dimensions[get_column_letter(column)].width = 28
            cell = WriteOnlyCell(sheet, value)
            cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
            cells.append(cell)
        sheet.append(cells)
    else:
        sheet.append(values)
    sheet._report_row_count = count + 1


def finish_report_workbook(workbook) -> None:
    """Finalize FAST filters or FULL headers and workbook navigation."""
    if workbook.write_only:
        for sheet in workbook.worksheets:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(sheet._report_column_count)}{sheet._report_row_count}"
            )
        return
    for sheet in workbook.worksheets:
        style_fast_sheet(sheet)
    summary = workbook["Summary"]
    summary.append(())
    summary.append(("Workbook navigation",))
    for sheet in workbook.worksheets:
        if sheet is summary:
            continue
        summary.append((sheet.title,))
        cell = summary.cell(summary.max_row, 1)
        cell.hyperlink = f"#'{sheet.title}'!A1"
        cell.font = Font(color="0563C1", underline="single")
        sheet.cell(1, 1).hyperlink = "#'Summary'!A1"
