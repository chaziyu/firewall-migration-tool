"""Presentation of acquisition evidence, separate from source inventory."""

from openpyxl.utils import get_column_letter

from ..extraction.sanitize import sanitize_raw_text
from .excel_style import append_report_row, style_fast_sheet


def collection_summary(status):
    if status is None:
        return ()
    rows = [("Collection Status", sanitize_raw_text(status))]
    if status != "SUCCESS":
        rows.append(("Collection Completeness", "Collection is incomplete. Missing configuration is unknown; see Collection Evidence."))
    return tuple(rows)


def write_collection_evidence(workbook, status, parts, warnings):
    if status is None:
        return
    sheet = workbook.create_sheet("Collection Evidence")
    headers = ("Part Name", "Status", "Complete", "Count", "Warnings")
    append_report_row(sheet, headers)
    for part in parts:
        append_report_row(sheet, (
            sanitize_raw_text(part.get("name")), sanitize_raw_text(part.get("status")),
            part.get("complete"), part.get("count"), None,
        ))
    for warning in warnings:
        append_report_row(sheet, (None, None, None, None, sanitize_raw_text(warning)))
    if workbook.write_only:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet._report_row_count}"
    else:
        style_fast_sheet(sheet)
